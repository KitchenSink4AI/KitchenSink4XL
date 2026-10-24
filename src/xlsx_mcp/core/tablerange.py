"""Converting a table to a plain range keeps the formulas that name it.

Excel's Convert to Range (ListObject.Unlist) rewrites every structured
reference to the table as a plain cell reference before the table goes, so
each formula keeps its value. Measured with a private hidden Excel on made-up
tables (Sales on sheet Data, header row 1, data rows 2 to 4, a totals row 5
when shown):

    Sales[Q2]                    Data!$C$2:$C$4    (any sheet, always qualified)
    Sales[]  Sales[#Data]        Data!$A$2:$D$4    (the data rows, every column)
    Sales[[Q1]:[Q2]]             Data!$B$2:$C$4
    Sales[#All]                  Data!$A$1:$D$5    (header, data and totals)
    Sales[#Headers]              Data!$A$1:$D$1
    Sales[[#Headers],[Q2]]       Data!$C$1         (one cell is written as one cell)
    Sales[[#Headers],[#Data],[Q1]]  Data!$B$1:$B$4
    Sales[[#Data],[#Totals],[Q1]]   Data!$B$2:$B$5
    Sales[#Totals]               Data!$A$5:$D$5
    Sales[[#This Row],[Q1]]      Data!$B2          (the formula's own row, relative)
    Sales[#This Row]             Data!$A2:$D2
    "Sales[Q1]" (text)           unchanged
    Sales (the name alone)       unchanged: Excel writes the whole table as
                                 Sales[] and reads a bare Sales in a file as
                                 an unknown name (#NAME?), before and after

Defined names are rewritten the same way. A reference this module cannot turn
into the same cells (a #This Row outside the table's data rows or not on its
sheet, a #This Row in a defined name or a rule, a column the table does not
have, a specifier Excel itself would refuse) refuses the conversion and names
where it stands, before anything changes. A pivot table whose source is the
table also refuses: its source would need rewriting too.
"""
from __future__ import annotations

from dataclasses import dataclass

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from . import structrefs as _sr
from .errors import XlMcpError


class Unconvertible(Exception):
    """A structured reference that cannot be written as the same cells."""


@dataclass
class TableSpan:
    name: str
    sheet: str
    min_row: int
    min_col: int
    max_row: int
    max_col: int
    headers: int
    totals: int
    columns: list[str]

    @property
    def data_top(self) -> int:
        return self.min_row + self.headers

    @property
    def data_bottom(self) -> int:
        return self.max_row - self.totals


def span_of(ws, table) -> TableSpan:
    c0, r0, c1, r1 = range_boundaries(table.ref)
    hrc = table.headerRowCount if table.headerRowCount is not None else 1
    trc = table.totalsRowCount if table.totalsRowCount is not None else 0
    cols = [str(c.name) for c in (table.tableColumns or ())]
    return TableSpan(table.displayName, ws.title, r0, c0, r1, c1, int(hrc), int(trc), cols)


def _items_with_seps(body: str) -> list[tuple[str, str]] | None:
    """[(separator before the item, raw item)] of a complex specifier; None
    when the body is not made only of bracketed items."""
    out, k, sep = [], 0, ""
    while k < len(body):
        c = body[k]
        if c in " ":
            k += 1
            continue
        if c in ",:":
            sep = c
            k += 1
            continue
        if c != "[":
            return None
        close = _sr._match(body, k)
        if close is None:
            return None
        out.append((sep, body[k + 1:close]))
        sep = ""
        k = close + 1
    return out


def _column_index(span: TableSpan, raw: str) -> int:
    name = _sr.unescape(raw)
    for i, c in enumerate(span.columns):
        if c.lower() == name.lower():
            return i
    raise Unconvertible(f"names a column {name!r} that table {span.name} does not have")


def _parse(body: str, span: TableSpan) -> tuple[frozenset, tuple[int, int] | None]:
    """(the special items, the columns as (first, last) index or None for
    all) of one structured reference's body."""
    b = body.strip()
    specials: set[str] = set()
    raw_cols: list[tuple[str, str]] = []
    if b.startswith("@"):
        specials.add("#This Row")
        rest = b[1:].strip()
        if rest.startswith("["):
            items = _items_with_seps(rest)
            if items is None:
                raise Unconvertible("cannot be read")
            raw_cols = items
        elif rest:
            raw_cols = [("", rest)]
    elif b.startswith("["):
        items = _items_with_seps(b)
        if items is None:
            raise Unconvertible("cannot be read")
        for sep, raw in items:
            item = raw.strip()
            if item.startswith("#"):
                sp = _sr.SPECIALS.get(item.lower())
                if sp is None:
                    raise Unconvertible(f"uses {item}, which Excel does not know")
                specials.add(sp)
            else:
                raw_cols.append((sep, raw))
    elif not b:
        pass
    elif b.startswith("#"):
        sp = _sr.SPECIALS.get(b.lower())
        if sp is None:
            raise Unconvertible(f"uses {b}, which Excel does not know")
        specials.add(sp)
    else:
        raw_cols = [("", body)]
    cols = None
    if raw_cols:
        if len(raw_cols) > 2 or (len(raw_cols) == 2 and raw_cols[1][0] != ":"):
            raise Unconvertible("names its columns in a form Excel does not accept")
        idx = [_column_index(span, raw) for _sep, raw in raw_cols]
        cols = (min(idx), max(idx))
    return frozenset(specials), cols


_ROWS = {
    frozenset(): "data", frozenset({"#Data"}): "data", frozenset({"#All"}): "all",
    frozenset({"#Headers"}): "headers", frozenset({"#Totals"}): "totals",
    frozenset({"#Headers", "#Data"}): "headers+data", frozenset({"#Data", "#Totals"}): "data+totals",
    frozenset({"#This Row"}): "this row",
}


def _a1(span: TableSpan, specials: frozenset, cols, row: int | None, sheet: str | None,
        qual: str) -> str:
    kind = _ROWS.get(specials)
    if kind is None:
        raise Unconvertible("combines " + ", ".join(sorted(specials)) + ", which Excel does not accept")
    c0, c1 = (span.min_col + cols[0], span.min_col + cols[1]) if cols else (span.min_col, span.max_col)
    if kind == "this row":
        if row is None:
            raise Unconvertible("uses #This Row where there is no row of its own")
        if sheet != span.sheet or not span.data_top <= row <= span.data_bottom:
            raise Unconvertible(f"uses #This Row outside the data rows of table {span.name}")
        a = f"${get_column_letter(c0)}{row}"
        b = f"${get_column_letter(c1)}{row}"
        return qual + (a if c0 == c1 else f"{a}:{b}")
    if kind in ("headers", "headers+data") and not span.headers:
        raise Unconvertible(f"names the header row of table {span.name}, which has none")
    if kind in ("totals", "data+totals") and not span.totals:
        raise Unconvertible(f"names the totals row of table {span.name}, which has none")
    r0, r1 = {
        "data": (span.data_top, span.data_bottom),
        "all": (span.min_row, span.max_row),
        "headers": (span.min_row, span.min_row),
        "totals": (span.max_row, span.max_row),
        "headers+data": (span.min_row, span.data_bottom),
        "data+totals": (span.data_top, span.max_row),
    }[kind]
    if r0 > r1:
        raise Unconvertible(f"names the data rows of table {span.name}, which has none")
    a = f"${get_column_letter(c0)}${r0}"
    b = f"${get_column_letter(c1)}${r1}"
    return qual + (a if (c0, r0) == (c1, r1) else f"{a}:{b}")


def in_formula(text: str, span: TableSpan, *, sheet: str | None = None,
               row: int | None = None, col: int | None = None) -> str:
    """text with every reference to the table written as the cells Excel's
    Convert to Range writes (module docstring). sheet/row/col: where the
    formula stands (a cell), for #This Row and for an unqualified reference
    inside the table. Raises Unconvertible."""
    from . import refs as _refs

    if span.name.lower() not in text.lower() and "[" not in text:
        return text
    qual = _refs.qualifier([span.sheet])
    inside = (sheet == span.sheet and row is not None and col is not None
              and span.min_row <= row <= span.max_row and span.min_col <= col <= span.max_col)
    out, pos = [], 0
    for ref in _sr.find_refs(text):
        if ref.table is not None:
            if ref.table.lower() != span.name.lower():
                continue
        elif not inside:
            continue
        specials, cols = _parse(ref.body, span)
        out.append(text[pos:ref.start])
        out.append(_a1(span, specials, cols, row, sheet, qual))
        pos = ref.end
    return "".join(out) + text[pos:]


def convert(wb, ws, table) -> dict:
    """Rewrite every reference to `table` across the workbook as Excel's
    Convert to Range does, or refuse (XlMcpError, nothing changed) naming
    each place that cannot be rewritten. Returns what changed."""
    from openpyxl.worksheet.formula import ArrayFormula

    from . import calc as _calc
    from . import refs as _refs

    span = span_of(ws, table)
    problems: list[str] = []
    cell_edits: list[tuple] = []
    for w in wb.worksheets:
        for cell in list(getattr(w, "_cells", {}).values()):
            val = cell.value
            text = None
            if isinstance(val, str) and _calc.is_formula_cell(cell):
                text = val
            elif isinstance(val, ArrayFormula) and val.text:
                text = val.text
            if text is None or span.name.lower() not in text.lower() and not (
                    "[" in text and w.title == span.sheet):
                continue
            where = f"{_refs.qualifier([w.title])}{cell.coordinate}"
            try:
                new = in_formula(text, span, sheet=w.title, row=cell.row, col=cell.column)
            except Unconvertible as exc:
                problems.append(f"{where} ({exc})")
                continue
            if new != text:
                cell_edits.append((cell, val, new))
    site_edits: list = []

    def other(text: str, home) -> str:
        if span.name.lower() not in text.lower():
            return text
        try:
            return in_formula(text, span)
        except Unconvertible as exc:
            problems.append(f"{site_desc[0]} ({exc})")
            return text

    site_desc = ["a defined name, rule or chart"]
    # defined names, conditional formats, data validation, charts and other
    # tables' own formulas, through the shared walk; cells were done above
    probe = _refs.RewriteReport()
    scopes = [(None, wb.defined_names)] + [
        (w.title, w.defined_names) for w in wb.worksheets if getattr(w, "defined_names", None) is not None]
    for home, dndict in scopes:
        for name, defn in list(dndict.items()):
            val = getattr(defn, "value", None) or getattr(defn, "attr_text", None)
            if not isinstance(val, str) or span.name.lower() not in val.lower():
                continue
            site_desc[0] = f"the defined name {name}"
            new = other(val, home)
            if new != val:
                site_edits.append((defn, new))
    for cache in _refs._pivot_caches(wb):
        src = _refs._worksheet_source(cache)
        nm = getattr(src, "name", None) if src is not None else None
        if isinstance(nm, str) and nm.lower() == span.name.lower():
            problems.append("a pivot table's source (it reads the table by name)")
    rule_hits: list[str] = []
    site = ["", ""]  # kind and cells (or chart reference) of the rule being read

    def names_table(text: str) -> bool:
        return span.name.lower() in text.lower() and any(
            r.table is not None and r.table.lower() == span.name.lower() for r in _sr.find_refs(text))

    def rules(text: str, home) -> str:
        if names_table(text):
            kind, where = site
            hit = (f"a chart that reads {where}" if kind == "chart"
                   else f"a {kind} rule on {where}")
            if hit not in rule_hits:
                rule_hits.append(hit)
        return text

    def at(kind: str, where: str) -> None:
        site[0], site[1] = kind, where

    _refs._rule_and_chart_sites(wb, rules, probe, on_site=at)
    for h in rule_hits:
        problems.append(f"{h} (Excel does not convert table references there)")
    for w in wb.worksheets:
        for t in list(getattr(w, "tables", {}).values()):
            if t is table:
                continue
            for tc in t.tableColumns or []:
                for attr in ("calculatedColumnFormula", "totalsRowFormula"):
                    tf = getattr(tc, attr, None)
                    text = getattr(tf, "attr_text", None)
                    if isinstance(text, str) and names_table(text):
                        problems.append(f"the column formula of table {t.displayName}")
    if problems:
        shown = "; ".join(problems[:12]) + (f"; and {len(problems) - 12} more" if len(problems) > 12 else "")
        exc = XlMcpError(
            f"to_range would leave formulas naming table {span.name} that no longer exists, and these "
            f"cannot be rewritten as plain cell references the way Excel's Convert to Range does: "
            f"{shown}. Nothing was changed. Rewrite them yourself, then convert the table.")
        exc.detail = {"table": span.name, "cannot_rewrite": problems}
        raise exc
    for cell, val, new in cell_edits:
        cell.value = ArrayFormula(val.ref, new) if isinstance(val, ArrayFormula) else new
    for defn, new in site_edits:
        try:
            defn.value = new
        except Exception:  # noqa: BLE001
            defn.attr_text = new
    return {"formulas": len(cell_edits), "names": len(site_edits),
            "cells": [f"{_refs.qualifier([c.parent.title])}{c.coordinate}" for c, _v, _n in cell_edits][:50]}


__all__ = ["TableSpan", "Unconvertible", "convert", "in_formula", "span_of"]
