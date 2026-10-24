"""core/surgical.py: the byte-preserving save for cell-level edits.

WHY THIS EXISTS. Every file-tier save used to go through openpyxl: load the
whole workbook into its object model, change the model, write a brand-new
package from the model. Anything the model does not hold was gone from that
new package: pictures, SharePoint and OneDrive customXml, x14 dropdowns,
slicers, printer settings, the cached result of every formula, the shared
string table, Excel's own row heights. The hazard scan refused the edit when
it could see the loss coming, so on four of seven ordinary real workbooks a
one-cell append either refused or offered to delete the pictures; on
the files it called clean it quietly stripped every cached formula result
and rewrote rows, strings and metadata.

THE ROUTE. A cell-level edit (values, formulas, number formats, cell
formatting, clearing, appending rows) changes a handful of <c> elements in
one or more worksheet parts, and occasionally adds an entry to the shared
string table, the style table, and the calculation flag. This module writes
exactly that and copies every other part of the package through unchanged:
the produced package holds the same parts, in the same order, and every part
the edit did not need is byte-for-byte the original. Inside the parts it does
touch, only the affected elements change; the rest of each part's bytes are
spliced through as they were (so x14 extension blocks, unknown attributes,
row heights, and every other cell's cached value survive).

HOW AN EDIT BECOMES SURGERY. The ops keep editing the openpyxl model exactly
as before; nothing in them changes. snapshot() records the model as it loaded
(every cell's value, type, style, hyperlink and comment, plus a fingerprint
of everything that is not a cell). At save, plan_edit() compares the model
against that snapshot: the cells that differ are the edit. If anything
outside the cells differs (a merge, a column width, a new sheet, a chart, a
defined name, a new named style ...) the edit is not cell-level and the save
takes the full openpyxl rewrite as before, where the hazard gate names any
loss and refuses without allow_loss. build() then reads the ORIGINAL package
and produces the new part bytes; NotSurgical from either stage means "use the
rewrite", never "lose something".

CACHED VALUES. Untouched formula cells keep their <v>. A formula the edit
wrote has no cached value (Excel computes it on open). The formulas whose
precedents the edit changed (directly or through other formulas) are reported
as possibly stale, and whenever anything may be stale the workbook's
calcPr gets fullCalcOnLoad="1" so Excel recalculates on open. A format-only
edit, or a value edit no formula refers to, leaves nothing stale and does not
set the flag. stale_formulas() is the shared dependency walk; it is
conservative (an INDIRECT/OFFSET formula, a name it cannot resolve, a 3D
reference, all count as depending on everything).
"""

from __future__ import annotations

import decimal
import html
import math
import posixpath
import re
import zipfile
from array import array
from dataclasses import dataclass, field
from typing import Any, Iterable

from . import partnames as _partnames

ROUTE_RAW = "raw_ooxml"
ROUTE_REWRITE = "openpyxl"


class NotSurgical(Exception):
    """The edit cannot be written as byte-preserving surgery. The message
    says why, in words a user can read (it is surfaced when the rewrite route
    then refuses)."""


# =================================================================== snapshot

_DEFAULT_STYLE = array("i", [0] * 9).tobytes()


def _value_key(v: Any) -> tuple:
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

    if isinstance(v, ArrayFormula):
        return ("ArrayFormula", str(v.ref), v.text)
    if isinstance(v, DataTableFormula):
        return ("DataTableFormula", tuple(sorted(dict(v).items())))
    t = type(v).__name__
    if t == "CellRichText":
        return (t, str(v), repr(list(v)))
    return (t, v)


def cell_key(cell) -> tuple:
    """Everything the save writes for one cell, as a comparable value."""
    st = getattr(cell, "_style", None)
    hl = getattr(cell, "_hyperlink", None)
    cm = getattr(cell, "_comment", None)
    hk = None if hl is None else (
        hl.target, hl.location, hl.tooltip, hl.display, getattr(hl, "ref", None))
    ck = None if cm is None else (
        cm.text, cm.author, getattr(cm, "width", None),
        getattr(cm, "height", None))
    return (_value_key(cell._value), cell.data_type,
            st.tobytes() if st is not None else _DEFAULT_STYLE, hk, ck)


EMPTY_KEY = (("NoneType", None), "n", _DEFAULT_STYLE, None, None)


def _ser(obj: Any) -> Any:
    """A comparable serialization of an openpyxl model object."""
    if obj is None or isinstance(obj, (str, int, float, bool)):
        return obj
    to_tree = getattr(obj, "to_tree", None)
    if callable(to_tree):
        try:
            el = to_tree()
        except Exception as exc:  # noqa: BLE001 - deterministic marker
            return ("unserializable", type(obj).__name__, type(exc).__name__)
        if el is None:
            return None
        from openpyxl.xml.functions import tostring
        try:
            return tostring(el)
        except Exception:  # noqa: BLE001
            return ("unserializable", type(obj).__name__)
    if isinstance(obj, (list, tuple)):
        return tuple(_ser(x) for x in obj)
    if isinstance(obj, dict):
        return tuple(sorted((str(k), _ser(v)) for k, v in obj.items()))
    return repr(obj)


def _dim_items(d) -> tuple:
    """A row or column dimension's attributes, read WITHOUT its style
    property: openpyxl's Dimension.style (and RowDimension.s) go through
    style_id, which ADDS the dimension's style array to the workbook's
    cell-style list as a side effect. Reading it here grew that list on
    LibreOffice files, so the style table no longer matched the model."""
    out = []
    for key in getattr(d, "__fields__", ()):
        if key in ("style", "s"):
            continue
        try:
            v = getattr(d, key, None)
        except Exception:  # noqa: BLE001
            v = None
        if v:
            out.append((key, str(v)))
    st = getattr(d, "_style", None)
    if st is not None and any(st):
        out.append(("_style", tuple(st)))
    return tuple(out)


def _dims(holder) -> tuple:
    out = []
    for k, d in list(holder.items()):
        try:
            items = _dim_items(d)
        except Exception:  # noqa: BLE001
            items = (repr(d),)
        if items:
            out.append((str(k), items))
    return tuple(sorted(out))


def _cf(ws) -> tuple:
    out = []
    for cf in ws.conditional_formatting:
        rules = []
        for r in cf.rules:
            rules.append((_ser(r), _ser(getattr(r, "dxf", None))))
        out.append((str(cf.sqref), tuple(rules)))
    return tuple(out)


def _tables(ws) -> tuple:
    """Each table with its column NAMES taken out: a header rename is the one
    table change the surgical route writes itself (table_renames)."""
    out = []
    for name in list(ws.tables):
        t = ws.tables[name]
        cols = tuple(_ser(_strip_name(tc)) for tc in (t.tableColumns or ()))
        out.append((name, t.ref, t.displayName, _ser(t.autoFilter),
                    _ser(t.tableStyleInfo), t.headerRowCount,
                    t.totalsRowCount, t.totalsRowShown, cols))
    return tuple(out)


def _strip_name(tc):
    from copy import copy
    c = copy(tc)
    c.name = "_"
    return c


def _anchors(objs) -> tuple:
    out = []
    for o in objs:
        anchor = getattr(o, "anchor", None)
        out.append((id(o), type(o).__name__,
                    anchor if isinstance(anchor, str) else _ser(anchor)))
    return tuple(out)


def ws_fingerprint(ws) -> tuple:
    """Everything about a worksheet that is not a cell. A difference here
    means the edit is not cell-level."""
    parts: list[Any] = [type(ws).__name__, ws.title,
                        getattr(ws, "sheet_state", None)]
    if type(ws).__name__ != "Worksheet":
        return tuple(parts)
    parts += [
        tuple(sorted(str(r) for r in ws.merged_cells.ranges)),
        _dims(ws.column_dimensions), _dims(ws.row_dimensions),
        _ser(ws.row_breaks), _ser(ws.col_breaks),
        _cf(ws), _ser(ws.data_validations), _tables(ws),
        _ser(ws.auto_filter), _ser(ws.views), _ser(ws.protection),
        _ser(ws.print_options), _ser(ws.page_margins), _ser(ws.page_setup),
        _ser(getattr(ws, "HeaderFooter", None)), _ser(ws.sheet_properties),
        _ser(ws.sheet_format), str(ws.print_title_rows),
        str(ws.print_title_cols), str(ws.print_area),
        _ser(dict(getattr(ws, "defined_names", {}) or {})),
        str(ws.legacy_drawing), _ser(ws.scenarios),
        _anchors(ws._images), _anchors(ws._charts), len(ws._pivots),
        _table_names(ws),
    ]
    return tuple(parts)


def _table_names(ws) -> tuple:
    return tuple((name, tuple(tc.name for tc in (ws.tables[name].tableColumns or ())))
                 for name in list(ws.tables))


def wb_fingerprint(wb) -> tuple:
    calc = wb.calculation
    # Every calculation property, the recalculate-on-open flag included: an
    # op that sets it (set_workbook_properties) is a workbook-level change
    # the cell splice would otherwise silently leave unwritten.
    calc_items = tuple(sorted(dict(calc).items())) if calc is not None else None
    return (
        tuple((type(s).__name__, s.title, getattr(s, "sheet_state", None))
              for s in wb._sheets),
        wb._active_sheet_index,
        _ser(dict(wb.defined_names)),
        calc_items,
        _ser(wb.properties), _ser(wb.custom_doc_props), _ser(wb.security),
        tuple(_ser(v) for v in wb.views),
        tuple(s.name for s in wb._named_styles),
        len(wb._differential_styles.styles),
        tuple(id(x) for x in wb._external_links),
        wb.code_name, wb.epoch, wb.vba_archive is not None, wb.is_template,
        len(wb._pivots),
    )


@dataclass
class SheetSnap:
    ws: Any
    title: str
    cells: dict
    fp: tuple


@dataclass
class WorkbookSnap:
    sheets: dict            # id(ws) -> SheetSnap
    order: tuple            # ids in load order
    wb_fp: tuple
    cell_styles: list       # at-load StyleArray tuples, index = original xf
    n_fonts: int
    n_fills: int
    n_borders: int
    n_numfmts: int
    named_styles: tuple
    cells_total: int = 0


def snapshot(wb) -> WorkbookSnap:
    """Record the model as it loaded, for plan_edit's comparison."""
    # The style list first: nothing below may touch it, but if a model
    # read ever did grow it, the snapshot must still be the file's own list.
    cell_styles = [tuple(s) for s in wb._cell_styles]
    sheets = {}
    total = 0
    for ws in wb._sheets:
        cells = {}
        if type(ws).__name__ == "Worksheet":
            for rc, cell in ws._cells.items():
                k = cell_key(cell)
                if k != EMPTY_KEY:
                    cells[rc] = k
            total += len(cells)
        sheets[id(ws)] = SheetSnap(ws, ws.title, cells, ws_fingerprint(ws))
    return WorkbookSnap(
        sheets=sheets, order=tuple(id(ws) for ws in wb._sheets),
        wb_fp=wb_fingerprint(wb),
        cell_styles=cell_styles,
        n_fonts=len(wb._fonts), n_fills=len(wb._fills),
        n_borders=len(wb._borders), n_numfmts=len(wb._number_formats),
        named_styles=tuple(s.name for s in wb._named_styles),
        cells_total=total)


# ======================================================================= plan

@dataclass
class TableRename:
    sheet: str
    table: str
    index: int
    old: str
    new: str


@dataclass
class SurgicalPlan:
    #: sheet title -> {(row, col): (new cell or None, old key, new key)}
    sheets: dict = field(default_factory=dict)
    table_renames: list = field(default_factory=list)

    @property
    def cell_count(self) -> int:
        return sum(len(v) for v in self.sheets.values())


def _describe_fp_change(old: tuple, new: tuple) -> str:
    labels = ["kind", "name", "visibility", "merged cells", "column widths",
              "row heights or row properties", "page breaks", "page breaks",
              "conditional formatting", "data validation", "tables",
              "autofilter", "sheet view", "sheet protection", "print options",
              "page margins", "page setup", "header and footer",
              "sheet properties", "sheet format", "print titles",
              "print titles", "print area", "sheet-scoped names",
              "comments drawing", "scenarios", "pictures", "charts",
              "pivot tables", "table column names"]
    for i, (a, b) in enumerate(zip(old, new)):
        if a != b:
            return labels[i] if i < len(labels) else "sheet structure"
    return "sheet structure"


def cell_changes(snap: WorkbookSnap, wb) -> dict:
    """{sheet title: {(row, col): (cell or None, old key, new key)}} for every
    cell whose value, type, style, hyperlink or comment differs from load.
    Sheets the snapshot does not know are skipped (plan_edit refuses them)."""
    out: dict = {}
    for ws in wb.worksheets:
        ss = snap.sheets.get(id(ws))
        if ss is None:
            continue
        changes: dict = {}
        old = ss.cells
        for rc, cell in ws._cells.items():
            k = cell_key(cell)
            ok = old.get(rc, EMPTY_KEY)
            if k != ok:
                changes[rc] = (cell, ok, k)
        for rc, ok in old.items():
            if rc not in ws._cells:
                changes[rc] = (None, ok, EMPTY_KEY)
        if changes:
            out[ws.title] = changes
    return out


def check_changed_formulas(changes: dict) -> None:
    """Excel's formula limits (core/limits.py) on every formula the edit
    wrote, whatever path wrote it."""
    from . import limits as _limits
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.formula import ArrayFormula

    for title, cells in changes.items():
        for (r, c), (cell, ok, nk) in cells.items():
            if cell is None or cell.data_type != "f" or ok[0] == nk[0]:
                continue
            v = cell._value
            text = v.text if isinstance(v, ArrayFormula) else v
            _limits.check_formula(
                text, what=f"the formula in {title}!{get_column_letter(c)}{r}")


def plan_edit(snap: WorkbookSnap, wb, changes: dict | None = None) -> SurgicalPlan:
    """The cell-level difference between the model now and as it loaded.
    Raises NotSurgical when the edit reaches beyond cells."""
    if wb_fingerprint(wb) != snap.wb_fp:
        old, new = snap.wb_fp, wb_fingerprint(wb)
        what = ["the sheet list", "the active sheet", "defined names",
                "calculation settings", "document properties",
                "custom properties", "workbook protection", "workbook views",
                "named styles", "conditional-format styles",
                "external links", "the code name", "the date system",
                "the macro project", "the template flag", "pivot tables"]
        for i, (a, b) in enumerate(zip(old, new)):
            if a != b:
                raise NotSurgical(f"this edit changes {what[i]}")
        raise NotSurgical("this edit changes the workbook structure")
    plan = SurgicalPlan()
    for ws in wb._sheets:
        ss = snap.sheets.get(id(ws))
        if ss is None:
            raise NotSurgical("this edit adds a sheet")
        fp = ws_fingerprint(ws)
        if fp != ss.fp:
            if type(ws).__name__ == "Worksheet" and _only_table_names(ss.fp, fp):
                plan.table_renames.extend(_table_renames(ss.ws, snap, ws))
                if not plan.table_renames:
                    raise NotSurgical(f"this edit changes a table on sheet {ws.title!r}")
            else:
                raise NotSurgical(
                    f"this edit changes the {_describe_fp_change(ss.fp, fp)} "
                    f"of sheet {ws.title!r}")
    plan.sheets = changes if changes is not None else cell_changes(snap, wb)
    return plan


_TABLE_NAMES_FP_INDEX = -1


def _only_table_names(old: tuple, new: tuple) -> bool:
    if len(old) != len(new):
        return False
    return old[:-1] == new[:-1] and old[-1] != new[-1]


def _table_renames(old_ws, snap, ws) -> list:
    old = dict(snap.sheets[id(ws)].fp[_TABLE_NAMES_FP_INDEX])
    out = []
    for name, cols in _table_names(ws):
        before = old.get(name)
        if before is None or len(before) != len(cols):
            raise NotSurgical(f"this edit changes table {name!r} on sheet {ws.title!r}")
        for i, (a, b) in enumerate(zip(before, cols)):
            if a != b:
                out.append(TableRename(ws.title, name, i, a, b))
    return out


# ================================================================ XML helpers

def xml_text(s: str) -> str:
    """Escape element text the way Excel writes it."""
    s = re.sub(r"_(x[0-9A-Fa-f]{4}_)", r"_x005F_\1", s)
    s = s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return s.replace("\r", "&#13;")


def xml_attr(s: str) -> str:
    return (s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;").replace("\r", "&#13;")
            .replace("\n", "&#10;").replace("\t", "&#9;"))


_ATTR_RE = re.compile(rb'([\w:.\-]+)\s*=\s*(?:"([^"]*)"|\'([^\']*)\')')


def _attrs(blob: bytes) -> dict:
    out = {}
    for m in _ATTR_RE.finditer(blob):
        v = m.group(2) if m.group(2) is not None else m.group(3)
        out[m.group(1).decode("ascii", "replace")] = v
    return out


def _set_attr(tag: bytes, name: bytes, value: bytes | None) -> bytes:
    """Set (or with None remove) one attribute inside a start tag, keeping
    every other byte of the tag."""
    rx = re.compile(rb'(\s)' + re.escape(name) + rb'\s*=\s*(?:"[^"]*"|\'[^\']*\')')
    m = rx.search(tag)
    if m:
        if value is None:
            return tag[:m.start()] + tag[m.end():]
        return tag[:m.start()] + m.group(1) + name + b'="' + value + b'"' + tag[m.end():]
    if value is None:
        return tag
    end = len(tag) - 1
    if tag.endswith(b"/>"):
        end = len(tag) - 2
    while end > 0 and tag[end - 1:end] in (b" ", b"\t", b"\r", b"\n"):
        end -= 1
    return tag[:end] + b" " + name + b'="' + value + b'"' + tag[end:]


_REL_NS_B = b"http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_WS_ROOT_RE = re.compile(
    rb"<((?:[A-Za-z_][\w.\-]*:)?)worksheet(?=[\s/>])((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>")
_XMLNS_ANY_RE = re.compile(rb"xmlns:([A-Za-z_][\w.\-]*)\s*=")


def worksheet_rel_prefix(data: bytes) -> tuple[bytes, bytes, int]:
    """The prefix a worksheet's root binds to the officeDocument
    relationships namespace, as (data, prefix, shift).

    Only the root's own declarations are in scope for a new top-level child
    (<legacyDrawing>, <hyperlinks>), so only they are read, in either quote
    style. With none, one is declared on the root start tag (shift is the
    bytes added, all of them before <sheetData>): under r, or when r is
    already declared anywhere in the sheet (a producer's own ignorable
    namespace, say) under the first of r0, r1, ... that nothing declares or
    lists as ignorable. Declaring a prefix the root already binds would make
    the sheet malformed XML."""
    root = _WS_ROOT_RE.search(data)
    if root is None:
        raise NotSurgical("a worksheet root is unreadable")
    root_attrs = _attrs(root.group(2))
    for k, v in root_attrs.items():
        if k.startswith("xmlns:") and v == _REL_NS_B:
            return data, k[6:].encode("ascii"), 0
    taken = {m.group(1) for m in _XMLNS_ANY_RE.finditer(data)}
    for k, v in root_attrs.items():
        if k.endswith(":Ignorable") or k == "Ignorable":
            taken.update(v.split())
    pfx, n = b"r", 0
    while pfx in taken:
        pfx = b"r%d" % n
        n += 1
    add = b" xmlns:" + pfx + b'="' + _REL_NS_B + b'"'
    at = root.start(2)
    return data[:at] + add + data[at:], pfx, len(add)


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def _col_letters(idx: int) -> str:
    s = ""
    while idx:
        idx, r = divmod(idx - 1, 26)
        s = chr(65 + r) + s
    return s


_COORD_RE = re.compile(rb"^\$?([A-Za-z]{1,3})\$?(\d+)$")


def _row_number(raw: bytes) -> int | None:
    """A worksheet row's r attribute as an int. openpyxl tolerates a decimal
    integer (<row r="1.0">, a real producer quirk) by reading it as that
    integer; this route used to run int() straight on the text and raise a
    bare ValueError that no caller caught, so a cell edit or diagnose on such
    a file failed with a Python error. Returns None for anything that is
    neither an integer nor an integer-valued decimal, so the caller can
    refuse it plainly."""
    text = raw.decode("ascii", "replace").strip()
    try:
        return int(text)
    except ValueError:
        pass
    try:
        f = float(text)
    except ValueError:
        return None
    return int(f) if f.is_integer() else None


def _split_coord(coord: bytes) -> tuple[int, int] | None:
    m = _COORD_RE.match(coord.strip())
    if not m:
        return None
    return int(m.group(2)), _col_index(m.group(1).decode("ascii"))


def _decode_part(data: bytes) -> None:
    """Refuse a part this module does not splice: anything but UTF-8."""
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        raise NotSurgical("a worksheet is stored as UTF-16")


# ============================================================ package reading

def rels_path(part: str) -> str:
    head, base = posixpath.split(part)
    return posixpath.join(head, "_rels", base + ".rels")


def resolve_target(source_part: str, target: str) -> str:
    """A relationship Target as a package part name."""
    t = _partnames.normalize(target.replace("\\", "/"))
    if t.startswith("/"):
        return posixpath.normpath(t.lstrip("/"))
    base = posixpath.dirname(source_part)
    return posixpath.normpath(posixpath.join(base, t))


_REL_RE = re.compile(rb"<(?:[\w.\-]+:)?Relationship\b((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>")


def parse_rels(data: bytes) -> list[dict]:
    out = []
    for m in _REL_RE.finditer(data):
        a = _attrs(m.group(1))
        out.append({k: html.unescape(v.decode("utf-8", "replace"))
                    for k, v in a.items()})
    return out


class Package:
    """Read access to the original package, by part name."""

    def __init__(self, zf: zipfile.ZipFile):
        self.zf = zf
        self.infos = zf.infolist()
        self.names = [i.filename for i in self.infos]
        self._lower = {n.lower(): n for n in self.names}
        for n in self.names:  # a part named with a plain-character percent code
            self._lower.setdefault(_partnames.normalize(n).lower(), n)
        self._cache: dict[str, bytes] = {}

    def has(self, name: str) -> bool:
        return self.real(name) is not None

    def real(self, name: str) -> str | None:
        return (self._lower.get(name.lower())
                or self._lower.get(_partnames.normalize(name).lower()))

    def read(self, name: str) -> bytes:
        real = self.real(name)
        if real is None:
            raise KeyError(name)
        if real not in self._cache:
            self._cache[real] = self.zf.read(real)
        return self._cache[real]

    def rels(self, part: str) -> list[dict]:
        rp = rels_path(part)
        if not self.has(rp):
            return []
        return parse_rels(self.read(rp))

    def workbook_part(self) -> str:
        for r in self.rels(""):
            if r.get("Type", "").endswith("/officeDocument"):
                return resolve_target("", r.get("Target", ""))
        return "xl/workbook.xml"

    def related(self, part: str, type_suffix: str) -> list[str]:
        return [resolve_target(part, r.get("Target", ""))
                for r in self.rels(part)
                if r.get("Type", "").endswith(type_suffix)
                and r.get("TargetMode") != "External"]

    def sheet_parts(self) -> dict[str, str]:
        wbp = self.workbook_part()
        data = self.read(wbp)
        ids = {r.get("Id"): resolve_target(wbp, r.get("Target", ""))
               for r in self.rels(wbp)}
        out = {}
        for m in re.finditer(rb"<(?:[\w.\-]+:)?sheet\b((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>",
                             data):
            a = _attrs(m.group(1))
            name = a.get("name")
            rid = next((v for k, v in a.items()
                        if k.endswith(":id") or k == "id"), None)
            if name is None or rid is None:
                continue
            part = ids.get(html.unescape(rid.decode("utf-8")))
            if part:
                out[html.unescape(name.decode("utf-8"))] = part
        return out


# ============================================================ sheet scanning

_TAG_RE = re.compile(
    rb"<(/?)((?:[A-Za-z_][\w.\-]*:)?)(row|c)(?=[\s/>])((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>")
_SHEETDATA_RE = re.compile(
    rb"<((?:[A-Za-z_][\w.\-]*:)?)sheetData\b((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>")


@dataclass
class CellLoc:
    col: int
    start: int
    tag_end: int
    end: int
    attrs: bytes
    selfclose: bool


@dataclass
class RowLoc:
    num: int
    start: int
    tag_end: int
    end: int
    attrs: bytes
    selfclose: bool
    cells: list = field(default_factory=list)


@dataclass
class SheetScan:
    prefix: bytes
    sd_open: tuple          # (start, end) of the <sheetData> start tag
    sd_selfclose: bool
    sd_close: int           # start of </sheetData> (or sd_open end)
    rows: dict              # num -> RowLoc for the target rows that exist
    insert_at: dict         # num -> offset for target rows that do not
    all_cells: list | None  # (start, end, col, row) of every cell (pass 2)


def scan_sheet(data: bytes, target_rows: Iterable[int], *,
               all_cells: bool = False) -> SheetScan:
    m = _SHEETDATA_RE.search(data)
    if m is None:
        raise NotSurgical("a worksheet has no cell table")
    prefix = m.group(1)
    selfclose = m.group(2).rstrip().endswith(b"/")
    targets = sorted(set(target_rows))
    scan = SheetScan(prefix, (m.start(), m.end()), selfclose, m.end(), {}, {},
                     [] if all_cells else None)
    if selfclose:
        for t in targets:
            scan.insert_at[t] = m.end()
        return scan
    close = data.find(b"</" + prefix + b"sheetData>", m.end())
    if close < 0:
        raise NotSurgical("a worksheet's cell table is not closed")
    scan.sd_close = close
    region_start = m.end()
    for marker in (b"<!--", b"<![CDATA[", b"<?"):
        if data.find(marker, region_start, close) >= 0:
            raise NotSurgical("a worksheet's cell table holds XML comments or CDATA")
    want = set(targets)
    ti = 0
    last_row = 0
    row: RowLoc | None = None
    row_num = 0
    keep = False
    last_col = 0
    cell: CellLoc | None = None
    for t in _TAG_RE.finditer(data, region_start, close):
        closing = t.group(1) == b"/"
        name = t.group(3)
        attrs = t.group(4)
        sc = attrs.rstrip().endswith(b"/")
        if name == b"row":
            if closing:
                if row is not None:
                    row.end = t.end()
                    if keep:
                        scan.rows[row.num] = row
                row = None
                continue
            a = _attrs(attrs)
            if "r" in a:
                num = _row_number(a["r"])
                if num is None:
                    raise NotSurgical("a worksheet has an unreadable row number")
            else:
                num = last_row + 1
            if num <= last_row:
                raise NotSurgical("a worksheet's rows are out of order")
            last_row = num
            while ti < len(targets) and targets[ti] < num:
                scan.insert_at[targets[ti]] = t.start()
                ti += 1
            keep = ti < len(targets) and targets[ti] == num
            if keep:
                ti += 1
            row = RowLoc(num, t.start(), t.end(), t.end(), attrs, sc)
            row_num = num
            last_col = 0
            if sc:
                row.end = t.end()
                if keep:
                    scan.rows[num] = row
                row = None
            continue
        # a <c> tag
        if row is None:
            raise NotSurgical("a worksheet has a cell outside any row")
        if closing:
            if cell is not None:
                cell.end = t.end()
                if keep:
                    row.cells.append(cell)
                if scan.all_cells is not None:
                    scan.all_cells.append((cell.start, cell.end, cell.col, row_num))
            cell = None
            continue
        a = _attrs(attrs)
        if "r" in a:
            rc = _split_coord(a["r"])
            if rc is None:
                raise NotSurgical("a worksheet has a cell with an unreadable address")
            if rc[0] != row_num:
                raise NotSurgical("a worksheet has a cell filed under the wrong row")
            col = rc[1]
        else:
            col = last_col + 1
        if col <= last_col:
            raise NotSurgical("a worksheet's cells are out of order")
        last_col = col
        c = CellLoc(col, t.start(), t.end(), t.end(), attrs, sc)
        if sc:
            if keep:
                row.cells.append(c)
            if scan.all_cells is not None:
                scan.all_cells.append((c.start, c.end, c.col, row_num))
        else:
            cell = c
    if row is not None:
        raise NotSurgical("a worksheet's last row is not closed")
    while ti < len(targets):
        scan.insert_at[targets[ti]] = close
        ti += 1
    return scan


_F_ELEM_RE = re.compile(
    rb"<((?:[A-Za-z_][\w.\-]*:)?)f\b((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>")


@dataclass
class OrigCell:
    s: int | None
    t: str | None
    has_formula: bool
    f_attrs: dict
    f_span: tuple | None     # (start, end) of the whole <f> element in data
    has_value: bool


def read_orig_cell(data: bytes, loc: CellLoc) -> OrigCell:
    a = _attrs(loc.attrs)
    s = int(a["s"]) if a.get("s", b"").strip().isdigit() else None
    t = a.get("t").decode("ascii", "replace") if "t" in a else None
    if loc.selfclose:
        return OrigCell(s, t, False, {}, None, False)
    body_start, body_end = loc.tag_end, loc.end
    fm = _F_ELEM_RE.search(data, body_start, body_end)
    f_attrs: dict = {}
    f_span = None
    if fm is not None:
        f_attrs = {k: v.decode("utf-8", "replace") for k, v in _attrs(fm.group(2)).items()}
        if fm.group(2).rstrip().endswith(b"/"):
            f_span = (fm.start(), fm.end())
        else:
            close = data.find(b"</" + fm.group(1) + b"f>", fm.end(), body_end)
            f_span = (fm.start(), close + len(fm.group(1)) + 4) if close >= 0 else None
            if f_span is None:
                raise NotSurgical("a formula element is not closed")
    has_v = re.search(rb"<(?:[A-Za-z_][\w.\-]*:)?v\b", data[body_start:body_end]) is not None
    return OrigCell(s, t, fm is not None, f_attrs, f_span, has_v)


# ============================================================ shared strings

_SI_RE = re.compile(rb"<((?:[A-Za-z_][\w.\-]*:)?)si\b(?:(?:[^>\"']|\"[^\"]*\"|'[^']*')*)>(.*?)</\1si>|<(?:[A-Za-z_][\w.\-]*:)?si\b(?:[^>]*)/>",
                    re.S)
_PLAIN_T_RE = re.compile(
    rb"^\s*<((?:[A-Za-z_][\w.\-]*:)?)t(?:\s+xml:space=\"preserve\")?\s*(?:/>|>(.*?)</\1t>)\s*$", re.S)


def _xml_unescape(raw: bytes) -> str:
    return html.unescape(raw.decode("utf-8"))


_XML_NUMREF_RE = re.compile(r"&#(x[0-9A-Fa-f]+|[0-9]+);")
_OOXML_ESCAPE_RE = re.compile(r"_x([0-9A-Fa-f]{4})_")


def _xml_entity_decode(text: str) -> str:
    """XML's own entity rules, not HTML's. XML defines the five predefined
    entities and numeric character references that name a code point
    directly; html.unescape adds the HTML named entities and, worse, applies
    HTML5's Windows-1252 remap of the C1 numeric references (&#x80; -> euro),
    which XML does not. Reading a shared string with html.unescape therefore
    read &#x80; as the euro sign and reused that entry when the euro sign was
    written, storing the wrong character."""
    def _num(m: re.Match) -> str:
        body = m.group(1)
        try:
            cp = int(body[1:], 16) if body[0] in "xX" else int(body)
            return chr(cp)
        except (ValueError, OverflowError):
            return m.group(0)
    text = _XML_NUMREF_RE.sub(_num, text)
    text = (text.replace("&lt;", "<").replace("&gt;", ">")
            .replace("&quot;", '"').replace("&apos;", "'"))
    return text.replace("&amp;", "&")


def _sst_text(raw: bytes) -> str:
    """The characters Excel reads from a shared string's <t> body: XML
    entities first (to recover the stored ST_Xstring text), then the OOXML
    _xHHHH_ escape. The escape names one character, so _x005F_ is a literal
    underscore and _x0041_ is 'A'; it is NOT the substring 'x005F_' deleted
    wherever it sits, which read the eight characters 'Ax005F_B' as 'AB' and
    reused that entry when 'AB' was written."""
    text = _xml_entity_decode(raw.decode("utf-8"))
    if "_x" in text:
        text = _OOXML_ESCAPE_RE.sub(lambda m: chr(int(m.group(1), 16)), text)
    return text


class SharedStrings:
    """The original shared string table, appended to only."""

    def __init__(self, part: str | None, data: bytes | None):
        self.part = part
        self.data = data
        self.prefix = b""
        self._index: dict[str, int] | None = None
        self.count = 0
        self.added: list[str] = []
        self.ref_delta = 0
        if data is not None:
            m = re.search(rb"<((?:[A-Za-z_][\w.\-]*:)?)sst\b", data)
            if m is None:
                raise NotSurgical("the shared string table is unreadable")
            self.prefix = m.group(1)

    @property
    def available(self) -> bool:
        return self.data is not None

    def _build(self) -> None:
        idx: dict[str, int] = {}
        n = 0
        for m in _SI_RE.finditer(self.data):
            body = m.group(2)
            if body is not None:
                pm = _PLAIN_T_RE.match(body)
                if pm is not None:
                    text = _sst_text(pm.group(2) or b"")
                    idx.setdefault(text, n)
            n += 1
        self._index = idx
        self.count = n

    def index_of(self, text: str) -> int:
        if self._index is None:
            self._build()
        hit = self._index.get(text)
        if hit is not None:
            return hit
        i = self.count + len(self.added)
        self.added.append(text)
        self._index[text] = i
        return i

    def serialize(self) -> bytes | None:
        if self.data is None or (not self.added and not self.ref_delta):
            return None
        data = self.data
        if self.added:
            p = self.prefix.decode("ascii")
            chunks = []
            for text in self.added:
                space = ' xml:space="preserve"' if _needs_preserve(text) else ""
                chunks.append(f"<{p}si><{p}t{space}>{xml_text(text)}</{p}t></{p}si>")
            close = data.rfind(b"</" + self.prefix + b"sst>")
            if close < 0:
                raise NotSurgical("the shared string table is not closed")
            data = data[:close] + "".join(chunks).encode("utf-8") + data[close:]
        m = re.search(rb"<" + re.escape(self.prefix) + rb"sst\b((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>", data)
        tag = m.group(0)
        a = _attrs(m.group(1))
        new_tag = tag
        total = self.count + len(self.added)
        if "uniqueCount" in a:
            new_tag = _set_attr(new_tag, b"uniqueCount", str(total).encode())
        if "count" in a:
            try:
                old = int(a["count"])
            except ValueError:
                old = total
            new_tag = _set_attr(new_tag, b"count",
                                str(max(total, old + self.ref_delta)).encode())
        return data[:m.start()] + new_tag + data[m.end():]


def _needs_preserve(text: str) -> bool:
    return bool(text) and (text[0].isspace() or text[-1].isspace()
                           or "\n" in text or "\t" in text)


# ==================================================================== styles

class Styles:
    """The original style table, appended to only. Maps a model StyleArray to
    an xf index of the ORIGINAL cellXfs."""

    def __init__(self, part: str | None, data: bytes | None, snap: WorkbookSnap, wb):
        self.part = part
        self.data = data
        self.snap = snap
        self.wb = wb
        self._by_tuple: dict[tuple, int] = {}
        for i, t in enumerate(snap.cell_styles):
            self._by_tuple.setdefault(t, i)
        self._new_xfs: list[str] = []
        self._new_fonts: list[str] = []
        self._new_fills: list[str] = []
        self._new_borders: list[str] = []
        self._new_numfmts: list[tuple[int, str]] = []
        self._font_map: dict[int, int] = {}
        self._fill_map: dict[int, int] = {}
        self._border_map: dict[int, int] = {}
        self._numfmt_map: dict[int, int] = {}
        self._orig_numfmts: dict[str, int] | None = None
        self._max_numfmt = 163
        self.prefix = ""
        self._counts: dict[str, int] = {}
        if data is not None:
            m = re.search(rb"<((?:[A-Za-z_][\w.\-]*:)?)styleSheet\b", data)
            if m is None:
                raise NotSurgical("the style table is unreadable")
            self.prefix = m.group(1).decode("ascii")

    def xf_index(self, sa, old_xf: int | None) -> int:
        t = tuple(sa)
        if old_xf is not None and old_xf < len(self.snap.cell_styles) and \
                self.snap.cell_styles[old_xf] == t:
            return old_xf
        hit = self._by_tuple.get(t)
        if hit is not None:
            return hit
        if self.data is None:
            raise NotSurgical("the workbook has no style table to extend")
        idx = len(self.snap.cell_styles) + len(self._new_xfs)
        self._new_xfs.append(self._xf_xml(sa, old_xf))
        self._by_tuple[t] = idx
        return idx

    # ---- components
    def _block(self, local: str) -> tuple[int, int, int] | None:
        """(start of open tag, end of open tag, start of close tag) of a
        top-level style collection, or None when absent."""
        p = re.escape(self.prefix.encode())
        m = re.search(rb"<" + p + local.encode() + rb"\b((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>", self.data)
        if m is None:
            return None
        if m.group(1).rstrip().endswith(b"/"):
            return (m.start(), m.end(), -1)
        close = self.data.find(b"</" + self.prefix.encode() + local.encode() + b">", m.end())
        if close < 0:
            raise NotSurgical(f"the style table's {local} is not closed")
        return (m.start(), m.end(), close)

    def _count_children(self, local: str, child: str) -> int:
        if local in self._counts:
            return self._counts[local]
        b = self._block(local)
        if b is None or b[2] < 0:
            n = 0
        else:
            n = len(re.findall(rb"<" + re.escape(self.prefix.encode()) + child.encode()
                               + rb"(?=[\s/>])", self.data[b[1]:b[2]]))
        self._counts[local] = n
        return n

    def _component(self, kind: str, idx: int, n_at_load: int, mapping: dict,
                   new_list: list, local: str, child: str) -> int:
        if idx < n_at_load:
            return idx
        if idx in mapping:
            return mapping[idx]
        obj = getattr(self.wb, kind)[idx]
        base = self._count_children(local, child)
        if base != n_at_load:
            raise NotSurgical(f"the style table's {local} does not match the model")
        new = base + len(new_list)
        new_list.append(_emit(obj.to_tree(), self.prefix))
        mapping[idx] = new
        return new

    def _numfmt(self, num_id: int) -> int:
        if num_id < 164:
            return num_id
        if num_id in self._numfmt_map:
            return self._numfmt_map[num_id]
        code = self.wb._number_formats[num_id - 164]
        if self._orig_numfmts is None:
            self._orig_numfmts = {}
            b = self._block("numFmts")
            if b is not None and b[2] >= 0:
                for m in re.finditer(
                        rb"<" + re.escape(self.prefix.encode())
                        + rb"numFmt\b((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>",
                        self.data[b[1]:b[2]]):
                    a = _attrs(m.group(1))
                    try:
                        nid = int(a.get("numFmtId", b"-1"))
                    except ValueError:
                        continue
                    fc = html.unescape(a.get("formatCode", b"").decode("utf-8"))
                    self._orig_numfmts.setdefault(fc, nid)
                    self._max_numfmt = max(self._max_numfmt, nid)
        hit = self._orig_numfmts.get(code)
        if hit is None:
            hit = self._max_numfmt + 1
            self._max_numfmt = hit
            self._new_numfmts.append((hit, code))
            self._orig_numfmts[code] = hit
        self._numfmt_map[num_id] = hit
        return hit

    def _xf_id(self, sa, old_xf: int | None) -> int:
        new = sa.xfId
        old = None
        if old_xf is not None and old_xf < len(self.snap.cell_styles):
            old = self.snap.cell_styles[old_xf][8]
        if new == (old if old is not None else 0):
            return new
        try:
            name = self.wb._named_styles[new].name
        except Exception:  # noqa: BLE001
            raise NotSurgical("a cell style could not be matched") from None
        orig = self._orig_named()
        if name not in orig:
            raise NotSurgical(f"the cell style {name!r} is new to this workbook")
        return orig[name]

    def _orig_named(self) -> dict:
        out: dict = {}
        b = self._block("cellStyles")
        if b is None or b[2] < 0:
            return out
        for m in re.finditer(rb"<" + re.escape(self.prefix.encode())
                             + rb"cellStyle\b((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>",
                             self.data[b[1]:b[2]]):
            a = _attrs(m.group(1))
            try:
                out.setdefault(html.unescape(a["name"].decode("utf-8")), int(a["xfId"]))
            except (KeyError, ValueError):
                continue
        return out

    def _xf_xml(self, sa, old_xf: int | None) -> str:
        wb = self.wb
        snap = self.snap
        font = self._component("_fonts", sa.fontId, snap.n_fonts, self._font_map,
                               self._new_fonts, "fonts", "font")
        fill = self._component("_fills", sa.fillId, snap.n_fills, self._fill_map,
                               self._new_fills, "fills", "fill")
        border = self._component("_borders", sa.borderId, snap.n_borders,
                                 self._border_map, self._new_borders, "borders",
                                 "border")
        numfmt = self._numfmt(sa.numFmtId)
        xfid = self._xf_id(sa, old_xf)
        p = self.prefix
        attrs = [f'numFmtId="{numfmt}"', f'fontId="{font}"', f'fillId="{fill}"',
                 f'borderId="{border}"', f'xfId="{xfid}"']
        if sa.quotePrefix:
            attrs.append('quotePrefix="1"')
        if sa.pivotButton:
            attrs.append('pivotButton="1"')
        attrs += ['applyNumberFormat="1"', 'applyFont="1"', 'applyFill="1"',
                  'applyBorder="1"']
        children = ""
        if sa.alignmentId:
            attrs.append('applyAlignment="1"')
            children += _emit(wb._alignments[sa.alignmentId].to_tree(), p)
        if sa.protectionId:
            attrs.append('applyProtection="1"')
            children += _emit(wb._protections[sa.protectionId].to_tree(), p)
        head = f"<{p}xf " + " ".join(attrs)
        return head + (f">{children}</{p}xf>" if children else "/>")

    @property
    def changed(self) -> bool:
        return bool(self._new_xfs)

    def serialize(self) -> bytes | None:
        if not self._new_xfs:
            return None
        data = self.data
        edits: list[tuple[int, int, bytes]] = []

        def append(local: str, child: str, items: list[str], base_n: int):
            if not items:
                return
            b = self._block(local)
            p = self.prefix
            new_n = base_n + len(items)
            payload = "".join(items).encode("utf-8")
            if b is None:
                raise NotSurgical(f"the style table has no {local} collection")
            start, open_end, close = b
            tag = data[start:open_end]
            if close < 0:
                tag2 = tag[:-2].rstrip() + b">"
                tag2 = _set_attr(tag2, b"count", str(new_n).encode())
                edits.append((start, open_end, tag2 + payload
                              + f"</{p}{local}>".encode()))
            else:
                if b"count=" in tag:
                    edits.append((start, open_end,
                                  _set_attr(tag, b"count", str(new_n).encode())))
                edits.append((close, close, payload))

        append("fonts", "font", self._new_fonts, self._count_children("fonts", "font"))
        append("fills", "fill", self._new_fills, self._count_children("fills", "fill"))
        append("borders", "border", self._new_borders,
               self._count_children("borders", "border"))
        n_xf = self._count_children("cellXfs", "xf")
        if n_xf != len(self.snap.cell_styles):
            raise NotSurgical("the style table's cellXfs does not match the model")
        append("cellXfs", "xf", self._new_xfs, n_xf)
        if self._new_numfmts:
            p = self.prefix
            items = [f'<{p}numFmt numFmtId="{i}" formatCode="{xml_attr(c)}"/>'
                     for i, c in self._new_numfmts]
            b = self._block("numFmts")
            if b is None:
                m = re.search(rb"<" + re.escape(p.encode())
                              + rb"styleSheet\b(?:[^>\"']|\"[^\"]*\"|'[^']*')*>", data)
                block = (f'<{p}numFmts count="{len(items)}">' + "".join(items)
                         + f"</{p}numFmts>").encode("utf-8")
                edits.append((m.end(), m.end(), block))
            else:
                base = self._count_children("numFmts", "numFmt")
                append("numFmts", "numFmt", items, base)
        return _apply_edits(data, edits)


def _emit(el, prefix: str) -> str:
    """Serialize an openpyxl-built (namespace-less) element with the target
    part's prefix for the main namespace."""
    tag = el.tag.rsplit("}", 1)[-1]
    out = [f"<{prefix}{tag}"]
    for k, v in el.attrib.items():
        out.append(f' {k}="{xml_attr(str(v))}"')
    kids = list(el)
    text = el.text
    if not kids and text is None:
        out.append("/>")
        return "".join(out)
    out.append(">")
    if text is not None:
        out.append(xml_text(text))
    for k in kids:
        out.append(_emit(k, prefix))
        if k.tail:
            out.append(xml_text(k.tail))
    out.append(f"</{prefix}{tag}>")
    return "".join(out)


def _apply_edits(data: bytes, edits: list[tuple[int, int, bytes]]) -> bytes:
    """Splice non-overlapping (start, end, replacement) edits into data."""
    edits = sorted(edits, key=lambda e: (e[0], e[1]))
    out = []
    pos = 0
    for start, end, rep in edits:
        if start < pos:
            raise NotSurgical("overlapping edits inside one part")
        out.append(data[pos:start])
        out.append(rep)
        pos = end
    out.append(data[pos:])
    return b"".join(out)


# ============================================================== cell writing

def number_text(v: Any) -> str:
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, decimal.Decimal):
        v = float(v)
    if isinstance(v, float):
        if not math.isfinite(v):
            raise NotSurgical("a number that is not finite")
        r = repr(v)
        if r.endswith(".0"):
            r = r[:-2]
        return r.upper() if "e" in r else r
    raise NotSurgical(f"a value of type {type(v).__name__}")


class CellWriter:
    def __init__(self, wb, sst: SharedStrings, styles: Styles, prefix: bytes):
        self.wb = wb
        self.sst = sst
        self.styles = styles
        self.p = prefix.decode("ascii")

    def element(self, coord: str, s: int | None, cell) -> str:
        """The <c> element for a cell, or '' for a cell that should not
        exist (empty and unstyled)."""
        from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

        p = self.p
        v = cell._value if cell is not None else None
        sattr = f' s="{s}"' if s else ""
        head = f'<{p}c r="{coord}"{sattr}'
        if v is None or (isinstance(v, str) and v == "" and cell.data_type != "f"):
            return f"{head}/>" if s else ""
        dt = cell.data_type
        if dt == "f":
            if isinstance(v, ArrayFormula):
                text = v.text or ""
                text = text[1:] if text.startswith("=") else text
                return (f'{head}><{p}f t="array" ref="{xml_attr(str(v.ref))}">'
                        f"{xml_text(text)}</{p}f></{p}c>")
            if isinstance(v, DataTableFormula):
                raise NotSurgical("a data table formula")
            text = str(v)
            text = text[1:] if text.startswith("=") else text
            return f"{head}><{p}f>{xml_text(text)}</{p}f></{p}c>"
        if dt == "s":
            if not isinstance(v, str):
                raise NotSurgical("rich text in a cell")
            if self.sst.available:
                i = self.sst.index_of(v)
                self.sst.ref_delta += 1
                return f'{head} t="s"><{p}v>{i}</{p}v></{p}c>'
            space = ' xml:space="preserve"' if _needs_preserve(v) else ""
            return (f'{head} t="inlineStr"><{p}is><{p}t{space}>{xml_text(v)}'
                    f"</{p}t></{p}is></{p}c>")
        if dt == "b":
            return f'{head} t="b"><{p}v>{"1" if v else "0"}</{p}v></{p}c>'
        if dt == "e":
            return f'{head} t="e"><{p}v>{xml_text(str(v))}</{p}v></{p}c>'
        if dt == "d":
            from openpyxl.utils.datetime import to_excel
            if getattr(v, "tzinfo", None) is not None:
                raise NotSurgical("a date with a time zone")
            serial = to_excel(v, self.wb.epoch)
            if serial is None:
                raise NotSurgical("a date Excel cannot store")
            return f"{head}><{p}v>{number_text(serial)}</{p}v></{p}c>"
        if dt == "n":
            return f"{head}><{p}v>{number_text(v)}</{p}v></{p}c>"
        raise NotSurgical(f"a cell of type {dt!r}")


# =========================================================== workbook.xml

_CALC_FOLLOWERS = ("oleSize", "customWorkbookViews", "pivotCaches",
                   "smartTagPr", "smartTagTypes", "webPublishing",
                   "fileRecoveryPr", "webPublishObjects", "extLst")


def set_full_calc(data: bytes) -> tuple[bytes, bool]:
    """Set fullCalcOnLoad="1" on calcPr, adding the element in its schema
    position when the workbook has none."""
    m = re.search(rb"<((?:[A-Za-z_][\w.\-]*:)?)calcPr\b((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>", data)
    if m is not None:
        tag = m.group(0)
        a = _attrs(m.group(2))
        if a.get("fullCalcOnLoad") in (b"1", b"true"):
            return data, False
        return data[:m.start()] + _set_attr(tag, b"fullCalcOnLoad", b"1") + data[m.end():], True
    root = re.search(rb"<((?:[A-Za-z_][\w.\-]*:)?)workbook\b", data)
    if root is None:
        raise NotSurgical("the workbook part is unreadable")
    p = root.group(1)
    el = b"<" + p + b'calcPr fullCalcOnLoad="1"/>'
    for follower in _CALC_FOLLOWERS:
        fm = re.search(rb"<" + re.escape(p) + follower.encode() + rb"[\s/>]", data)
        if fm is not None:
            return data[:fm.start()] + el + data[fm.start():], True
    close = data.rfind(b"</" + p + b"workbook>")
    if close < 0:
        raise NotSurgical("the workbook part is not closed")
    return data[:close] + el + data[close:], True


# ============================================================== staleness

#: Functions whose references cannot be read off the formula text.
_UNKNOWN_REF_FUNCS = re.compile(r"\b(?:INDIRECT|OFFSET)\s*\(", re.I)
#: The staleness walk reads formula text with the reference lexer
#: (core.refs: strings, quoted sheet names, structured-reference brackets and
#: other workbooks' prefixes kept apart), and reads names in any script, as
#: Excel does. Its old ASCII-only name pattern skipped Tasa_ñ, 売上Rate, a?b,
#: a\b and every sheet-qualified name (=Loc!LocalRate*5), so their dependents
#: were never reported and the report still said it was complete.


def _has_relative_ref(text: str) -> bool:
    """True when a defined name's formula holds a relative reference (a row
    or column without $). Such a name points somewhere different for every
    cell that uses it, so its precedents cannot be read off the text alone
    (=RowVal with RowVal = Inputs!$B1 read as B1 for every
    caller)."""
    from . import refs as _refs

    for m in _refs.find_references(text[1:] if text.startswith("=") else text):
        for g in ("c1", "c2"):
            tok = m.group(g)
            if tok and not re.fullmatch(r"\$[A-Za-z]+\$\d+", tok):
                return True
        for g in ("cs1", "cs2", "rs1", "rs2"):
            tok = m.group(g)
            if tok and not tok.startswith("$"):
                return True
    return False


EVERYTHING = ("*",)
#: A sheet qualifier just before a '!': quoted ('My Sheet'), or bare in any script.
_QUALIFIER_TAIL = None


def _qualifier_before(text: str, bang: int) -> str | None:
    """The sheet name written just before the '!' at text[bang] (quotes
    removed), or None when there is none."""
    global _QUALIFIER_TAIL
    from . import refs as _refs

    if _QUALIFIER_TAIL is None:
        _QUALIFIER_TAIL = re.compile(
            "(?:" + _refs._QUOTED_SHEET + "|" + _refs._NAME_CH + "+)$")
    m = _QUALIFIER_TAIL.search(text, 0, bang)
    if m is None:
        return None
    q = m.group(0)
    if q.startswith("'"):
        q = q[1:-1].replace("''", "'")
    return q


def _formula_regions(text: str, own_sheet: str, names: dict, tables: dict,
                     depth: int = 0) -> list:
    """The precedent rectangles of one formula: (sheet_lower, r0, c0, r1, c1),
    or EVERYTHING when they cannot be bounded."""
    from . import refs as _refs

    body = text[1:] if text.startswith("=") else text
    if _UNKNOWN_REF_FUNCS.search(body):
        return [EVERYTHING]
    read = _refs._Read(body)
    rest = list(read.view)
    regions: list = []
    for m in read.hits(_refs._REF_RE, "sheet"):
        for k in range(m.start(), m.end()):
            rest[k] = " "
        if m.before in _refs._NOT_OURS:
            continue  # another workbook, or the tail of an error: a local edit cannot change it
        sheet_tok = m.group("sheet")
        sheet = own_sheet
        if sheet_tok:
            name, is_3d = _refs._sheet_scope(sheet_tok)
            if is_3d:
                return [EVERYTHING]
            if "[" in sheet_tok:
                continue  # another workbook: local edits cannot change it
            sheet = name
        sl = (sheet or "").lower()
        if m.group("c1"):
            r0c0 = _split_coord(m.group("c1").encode())
            r1c1 = _split_coord(m.group("c2").encode()) if m.group("c2") else r0c0
            if r0c0 is None or r1c1 is None:
                return [EVERYTHING]
            regions.append((sl, min(r0c0[0], r1c1[0]), min(r0c0[1], r1c1[1]),
                            max(r0c0[0], r1c1[0]), max(r0c0[1], r1c1[1])))
        elif m.group("cs1"):
            c0 = _col_index(m.group("cs1").replace("$", ""))
            c1 = _col_index(m.group("cs2").replace("$", ""))
            regions.append((sl, 1, min(c0, c1), 1048576, max(c0, c1)))
        else:
            r0 = int(m.group("rs1").replace("$", ""))
            r1 = int(m.group("rs2").replace("$", ""))
            regions.append((sl, min(r0, r1), 1, max(r0, r1), 16384))
    # Structured references: Table[...] resolves to the whole table, in any
    # script (Ventas_ñ[Amt]); an unqualified one ([@Amt]) cannot be bounded.
    # Another workbook's prefix ([1]) holds no reference of this workbook.
    for start, end, kind in _refs._lex(body):
        if kind == _refs._EXT:
            for k in range(start, end):
                rest[k] = " "
            continue
        if kind != _refs._BRK:
            continue
        j = start
        while j > 0 and _refs._is_name_char(read.view[j - 1]):
            j -= 1
        tname = body[j:start]
        if j > 0 and body[j - 1] == "!":
            # a table of another workbook (Book.xlsx!T[v]): not ours
            for k in range(j, end):
                rest[k] = " "
            continue
        t = tables.get(tname.lower()) if tname else None
        if t is None:
            return [EVERYTHING]
        regions.append(t)
        for k in range(j, end):
            rest[k] = " "
    rest_s = "".join(rest)
    for m in re.finditer(_refs._NAME_CH + "+", rest_s):
        tok = m.group(0)
        if tok[0].isdigit() or tok[0] == ".":
            continue  # a number
        s_, e_ = m.span()
        after = rest_s[e_] if e_ < len(rest_s) else ""
        before = rest_s[s_ - 1] if s_ > 0 else ""
        if after == "!" or before == "#":
            continue  # a sheet qualifier, or the text of an error (#REF!, #N/A)
        up = tok.upper()
        if up in ("TRUE", "FALSE") or up.startswith("_XLFN") or up.startswith("_XLPM"):
            continue
        if before == "!":
            qual = _qualifier_before(body, s_ - 1)
            if qual is None or "[" in qual or "]" in qual:
                continue  # a name of another workbook
            if ":" in qual:
                return [EVERYTHING]
            target = names.get((qual.lower(), up)) or names.get((None, up))
        else:
            target = names.get((own_sheet.lower(), up)) or names.get((None, up))
        if target is None:
            continue  # a LAMBDA parameter, a function, or a literal word
        if depth > 8 or _has_relative_ref(target[1]):
            return [EVERYTHING]
        sub = _formula_regions(target[1], target[0] or own_sheet, names,
                               tables, depth + 1)
        if EVERYTHING in sub:
            return [EVERYTHING]
        regions.extend(sub)
    return regions


def _names_and_tables(wb) -> tuple[dict, dict]:
    names: dict = {}
    for name, dn in dict(wb.defined_names).items():
        names[(None, name.upper())] = (None, "=" + str(dn.attr_text or ""))
    tables: dict = {}
    for ws in wb.worksheets:
        for name, dn in dict(getattr(ws, "defined_names", {}) or {}).items():
            names[(ws.title.lower(), name.upper())] = (ws.title, "=" + str(dn.attr_text or ""))
        for tname in list(ws.tables):
            t = ws.tables[tname]
            rc = _ref_box(t.ref)
            if rc is not None:
                tables[tname.lower()] = (ws.title.lower(),) + rc
    return names, tables


def _ref_box(ref: str) -> tuple | None:
    parts = str(ref).replace("$", "").split(":")
    a = _split_coord(parts[0].encode())
    b = _split_coord(parts[-1].encode())
    if a is None or b is None:
        return None
    return (min(a[0], b[0]), min(a[1], b[1]), max(a[0], b[0]), max(a[1], b[1]))


@dataclass
class Staleness:
    cells: list            # [(sheet, "A1")] possibly stale, in workbook order
    complete: bool = True  # False when the walk hit its work ceiling


#: Work ceiling for the dependency walk (formula x dirty-point checks).
_STALE_BUDGET = 4_000_000


def stale_formulas(wb, dirty: dict, *, exclude: set | None = None) -> Staleness:
    """Formula cells whose result may have changed because the cells in
    `dirty` ({sheet title: {(row, col)}}) changed, directly or through other
    formulas. `exclude` holds (sheet, row, col) cells the edit itself wrote
    (a newly written formula is reported separately)."""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.formula import ArrayFormula

    if not any(dirty.values()):
        return Staleness([])
    names, tables = _names_and_tables(wb)
    formulas = []  # (sheet_title, row, col, regions)
    for ws in wb.worksheets:
        for (r, c), cell in ws._cells.items():
            v = cell._value
            if isinstance(v, ArrayFormula):
                text = v.text or ""
            elif type(v).__name__ == "DataTableFormula":
                # A what-if data table's results depend on its
                # model through Excel's substitution, which the text does not
                # show; any content edit may change them. Every cell of the
                # table's range, not just the anchor that holds the formula.
                box = _ref_box(getattr(v, "ref", "") or "") or (r, c, r, c)
                for rr in range(box[0], box[2] + 1):
                    for cc in range(box[1], box[3] + 1):
                        formulas.append((ws.title, rr, cc, [EVERYTHING]))
                continue
            elif cell.data_type == "f" and isinstance(v, str):
                text = v
            else:
                continue
            formulas.append((ws.title, r, c,
                             _formula_regions(text, ws.title, names, tables)))
    points: dict = {}
    for sheet, cells in dirty.items():
        points.setdefault(sheet.lower(), set()).update(cells)
    stale: dict = {}
    budget = _STALE_BUDGET
    queue = [(s, rc) for s, cs in points.items() for rc in cs]
    seen_points = set(queue)
    complete = True
    pending = list(range(len(formulas)))
    while queue:
        batch = queue
        queue = []
        by_sheet: dict = {}
        for s, rc in batch:
            by_sheet.setdefault(s, []).append(rc)
        boxes = {s: (min(r for r, _ in v), min(c for _, c in v),
                     max(r for r, _ in v), max(c for _, c in v))
                 for s, v in by_sheet.items()}
        still = []
        for fi in pending:
            sheet, r, c, regions = formulas[fi]
            hit = False
            for reg in regions:
                if reg is EVERYTHING or reg == EVERYTHING:
                    hit = True
                    break
                pts = by_sheet.get(reg[0])
                if not pts:
                    continue
                bx = boxes[reg[0]]
                if reg[3] < bx[0] or reg[1] > bx[2] or reg[4] < bx[1] or reg[2] > bx[3]:
                    continue
                budget -= len(pts)
                if budget < 0:
                    hit = True
                    complete = False
                    break
                for pr, pc in pts:
                    if reg[1] <= pr <= reg[3] and reg[2] <= pc <= reg[4]:
                        hit = True
                        break
                if hit:
                    break
            if hit:
                key = (sheet, r, c)
                if not exclude or key not in exclude:
                    stale[key] = True
                pt = (sheet.lower(), (r, c))
                if pt not in seen_points:
                    seen_points.add(pt)
                    queue.append(pt)
            else:
                still.append(fi)
        pending = still
    order = {ws.title: i for i, ws in enumerate(wb.worksheets)}
    cells = sorted(stale, key=lambda k: (order.get(k[0], 0), k[1], k[2]))
    return Staleness([(s, f"{get_column_letter(c)}{r}") for s, r, c in cells],
                     complete)


# ================================================================== building

@dataclass
class BuildResult:
    parts: dict                 # part name -> new bytes
    removed: set                # part names dropped (calcChain)
    sheets_changed: list
    cells_written: int
    formulas_written: list      # [(sheet, "A1")]
    stale: Staleness
    full_calc: bool
    strings_added: int
    styles_added: int


def _row_attr_tag(row: RowLoc, data: bytes) -> bytes:
    return data[row.start:row.tag_end]


def _widen_spans(tag: bytes, cols: list[int]) -> bytes:
    a = _attrs(tag[1:])
    sp = a.get("spans")
    if sp is None or not cols:
        return tag
    lo, hi = None, None
    for piece in sp.split():
        try:
            x, y = piece.split(b":")
            x, y = int(x), int(y)
        except ValueError:
            return _set_attr(tag, b"spans", None)
        lo = x if lo is None else min(lo, x)
        hi = y if hi is None else max(hi, y)
    if lo is None:
        return tag
    nlo, nhi = min([lo] + cols), max([hi] + cols)
    if (nlo, nhi) == (lo, hi):
        return tag
    return _set_attr(tag, b"spans", f"{nlo}:{nhi}".encode())


def _cols_styles(data: bytes, prefix: bytes) -> list[tuple[int, int, int]]:
    out = []
    m = re.search(rb"<" + re.escape(prefix) + rb"cols\b", data)
    if m is None:
        return out
    close = data.find(b"</" + prefix + b"cols>", m.end())
    if close < 0:
        return out
    for cm in re.finditer(rb"<" + re.escape(prefix) + rb"col\b((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>",
                          data[m.end():close]):
        a = _attrs(cm.group(1))
        try:
            st = int(a.get("style", b"0"))
            lo, hi = int(a["min"]), int(a["max"])
        except (KeyError, ValueError):
            continue
        if st:
            out.append((lo, hi, st))
    return out


def _dimension_edit(data: bytes, prefix: bytes, rows_cols: list[tuple[int, int]],
                    empty: bool = False):
    m = re.search(rb"<" + re.escape(prefix) + rb"dimension\b((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>", data)
    if m is None or not rows_cols:
        return None
    a = _attrs(m.group(1))
    ref = a.get("ref")
    if ref is None:
        return None
    box = _ref_box(ref.decode("ascii", "replace"))
    if box is None:
        return None
    r0, c0, r1, c1 = box
    rows = [r for r, _ in rows_cols]
    cols = [c for _, c in rows_cols]
    if empty:
        # An empty sheet's dimension is a placeholder ("A1"); the new cells
        # are the whole used range, as Excel writes it.
        nr0, nc0, nr1, nc1 = min(rows), min(cols), max(rows), max(cols)
    else:
        nr0, nc0 = min(rows + [r0]), min(cols + [c0])
        nr1, nc1 = max(rows + [r1]), max(cols + [c1])
    if (nr0, nc0, nr1, nc1) == (r0, c0, r1, c1):
        return None
    new = f"{_col_letters(nc0)}{nr0}"
    if (nr0, nc0) != (nr1, nc1):
        new += f":{_col_letters(nc1)}{nr1}"
    return (m.start(), m.end(), _set_attr(m.group(0), b"ref", new.encode()))


def _materialize_f(data: bytes, span: tuple, prefix: bytes, text: str, attrs: dict) -> bytes:
    keep = {k: v for k, v in attrs.items() if k not in ("t", "ref", "si")}
    extra = "".join(f' {k}="{xml_attr(v)}"' for k, v in keep.items())
    p = prefix.decode("ascii")
    body = text[1:] if text.startswith("=") else text
    return f"<{p}f{extra}>{xml_text(body)}</{p}f>".encode("utf-8")


def build(snap: WorkbookSnap, wb, plan: SurgicalPlan, src_path: str) -> BuildResult:
    """Produce the new bytes of every part the edit touches, reading the
    original package. Raises NotSurgical when the original's structure is
    one this module does not splice."""
    from openpyxl.utils import get_column_letter

    with zipfile.ZipFile(src_path) as zf:
        pkg = Package(zf)
        wbp = pkg.workbook_part()
        sheet_parts = pkg.sheet_parts()
        sst_parts = pkg.related(wbp, "/sharedStrings")
        sst_part = sst_parts[0] if sst_parts and pkg.has(sst_parts[0]) else None
        st_parts = pkg.related(wbp, "/styles")
        st_part = st_parts[0] if st_parts and pkg.has(st_parts[0]) else None
        sst = SharedStrings(sst_part, pkg.read(sst_part) if sst_part else None)
        styles = Styles(st_part, pkg.read(st_part) if st_part else None, snap, wb)
        out_parts: dict[str, bytes] = {}
        removed: set[str] = set()
        formulas_written: list = []
        ct_edits: list = []
        formula_removed = False
        dirty: dict = {}
        written = 0
        changed_sheets = []
        for title, changes in plan.sheets.items():
            part = sheet_parts.get(title)
            if part is None or not pkg.has(part):
                raise NotSurgical(f"sheet {title!r} has no worksheet part")
            part = pkg.real(part)
            data = pkg.read(part)
            _decode_part(data)
            rows = {r for (r, _c) in changes}
            scan = scan_sheet(data, rows)
            prefix = scan.prefix
            writer = CellWriter(wb, sst, styles, prefix)
            ws = wb[title]
            edits: list[tuple[int, int, bytes]] = []
            # original cells, by (row, col)
            orig_locs: dict = {}
            for rnum, row in scan.rows.items():
                for c in row.cells:
                    orig_locs[(rnum, c.col)] = c
            col_styles = None
            shared_groups: set[str] = set()
            new_cells_by_row: dict[int, list[tuple[int, bytes]]] = {}
            notes: dict = {}
            links: dict = {}
            replaced: dict = {}
            extent: list = []
            for (r, c), (cell, old_key, new_key) in sorted(changes.items()):
                coord = f"{get_column_letter(c)}{r}"
                loc = orig_locs.get((r, c))
                orig = read_orig_cell(data, loc) if loc is not None else None
                content_changed = (cell is None or old_key[0] != new_key[0]
                                   or old_key[1] != new_key[1])
                if old_key[3] != new_key[3]:
                    links[(r, c)] = (getattr(cell, "_hyperlink", None)
                                     if cell is not None else None)
                if old_key[4] != new_key[4]:
                    notes[(r, c)] = (getattr(cell, "_comment", None) if cell is not None
                                     else None, old_key[4])
                if cell is not None and old_key[:3] == new_key[:3]:
                    continue  # only a link or a note changed; the cell stays as it is
                if orig is not None and orig.has_formula and \
                        orig.f_attrs.get("t") == "shared":
                    si = orig.f_attrs.get("si")
                    if si is not None and content_changed:
                        shared_groups.add(si)
                if orig is not None and orig.has_formula and \
                        orig.f_attrs.get("t") == "dataTable":
                    raise NotSurgical(f"{title}!{coord} is part of a data table")
                # --- the style index
                if cell is None:
                    s_new = None
                else:
                    sa = cell._style
                    old_s = orig.s if orig is not None else None
                    if old_key[2] == new_key[2] and orig is not None:
                        s_new = orig.s
                    elif loc is None and new_key[2] == _DEFAULT_STYLE:
                        # A brand-new cell keeps the look Excel would give a
                        # value typed there: the row's format, else the
                        # column's.
                        s_new = None
                        row = scan.rows.get(r)
                        if row is not None:
                            ra = _attrs(row.attrs)
                            if ra.get("customFormat") in (b"1", b"true") and \
                                    ra.get("s", b"").isdigit():
                                s_new = int(ra["s"])
                        if s_new is None:
                            if col_styles is None:
                                col_styles = _cols_styles(data, prefix)
                            for lo, hi, st in col_styles:
                                if lo <= c <= hi:
                                    s_new = st
                                    break
                    else:
                        s_new = styles.xf_index(sa, old_s)
                # --- style-only change on an existing cell: patch s= only
                if loc is not None and not content_changed:
                    tag = data[loc.start:loc.tag_end]
                    new_tag = _set_attr(tag, b"s", str(s_new).encode() if s_new else None)
                    if new_tag != tag:
                        edits.append((loc.start, loc.tag_end, new_tag))
                        written += 1
                    continue
                # --- content change
                if orig is not None and orig.t == "s" and sst.available:
                    sst.ref_delta -= 1
                xml = writer.element(coord, s_new, cell).encode("utf-8")
                if cell is not None and cell.data_type == "f":
                    formulas_written.append((title, coord))
                if orig is not None and orig.has_formula and \
                        (cell is None or cell.data_type != "f"):
                    formula_removed = True
                dirty.setdefault(title, set()).add((r, c))
                written += 1
                if xml:
                    extent.append((r, c))
                if loc is not None:
                    edits.append((loc.start, loc.end, xml))
                    replaced[(r, c)] = True
                else:
                    if xml:
                        new_cells_by_row.setdefault(r, []).append((c, xml))
            # --- new cells into existing or new rows
            p = prefix.decode("ascii")
            for r, items in sorted(new_cells_by_row.items()):
                row = scan.rows.get(r)
                if row is None:
                    at = scan.insert_at.get(r)
                    if at is None:
                        raise NotSurgical("could not place a new row")
                    body = b"".join(x for _c, x in items)
                    row_xml = f'<{p}row r="{r}">'.encode() + body + f"</{p}row>".encode()
                    if scan.sd_selfclose:
                        continue  # handled below
                    edits.append((at, at, row_xml))
                    continue
                cols = [c for c, _x in items]
                tag = _row_attr_tag(row, data)
                if row.selfclose:
                    open_tag = _widen_spans(tag[:-2].rstrip() + b">", cols)
                    body = b"".join(x for _c, x in items)
                    edits.append((row.start, row.end,
                                  open_tag + body + f"</{p}row>".encode()))
                    continue
                new_tag = _widen_spans(tag, cols)
                if new_tag != tag:
                    edits.append((row.start, row.tag_end, new_tag))
                for c, x in items:
                    after = next((cl for cl in row.cells if cl.col > c), None)
                    if after is not None:
                        pos = after.start
                    else:
                        pos = data.rfind(b"</" + prefix + b"row>", row.tag_end, row.end)
                        if pos < 0:
                            raise NotSurgical("a row is not closed")
                    edits.append((pos, pos, x))
            if scan.sd_selfclose and new_cells_by_row:
                rows_xml = b""
                for r, items in sorted(new_cells_by_row.items()):
                    rows_xml += (f'<{p}row r="{r}">'.encode()
                                 + b"".join(x for _c, x in items) + f"</{p}row>".encode())
                s, e = scan.sd_open
                edits.append((s, e, f"<{p}sheetData>".encode() + rows_xml
                              + f"</{p}sheetData>".encode()))
            # --- shared formula groups the edit broke into
            if shared_groups:
                edits.extend(_materialize_groups(data, ws, scan, shared_groups,
                                                 replaced, prefix))
            was_empty = scan.sd_selfclose or re.search(
                rb"<" + re.escape(prefix) + rb"c[\s/>]",
                data[scan.sd_open[1]:scan.sd_close]) is None
            dim = _dimension_edit(data, prefix, extent, was_empty)
            if dim is not None:
                edits.append(dim)
            new_data = _apply_edits(data, edits) if edits else data
            if notes:
                from . import rawnotes as _rawnotes
                new_data = _rawnotes.apply(pkg, part, new_data, notes, out_parts,
                                           ct_edits, removed)
            if links:
                from . import rawlinks as _rawlinks
                new_data = _rawlinks.apply(pkg, part, new_data, links, out_parts)
            if new_data != data:
                out_parts[part] = new_data
            if edits or notes or links:
                changed_sheets.append(title)
        if plan.table_renames:
            _rename_table_columns(pkg, sheet_parts, plan.table_renames, out_parts)
        if ct_edits:
            from . import rawnotes as _rawnotes
            _rawnotes.apply_content_types(pkg, out_parts, ct_edits)
        # --- the calculation chain lists formula cells; one the edit turned
        # into a value would point Excel at a cell with no formula.
        if formula_removed:
            chain = pkg.related(wbp, "/calcChain")
            for cp in chain:
                if pkg.has(cp):
                    removed.add(pkg.real(cp))
        stale = stale_formulas(
            wb, dirty, exclude={(s, *_rc(c)) for s, c in formulas_written})
        full_calc = bool(formulas_written) or bool(stale.cells) or not stale.complete
        if full_calc:
            new_wb, changed = set_full_calc(pkg.read(wbp))
            if changed:
                out_parts[pkg.real(wbp)] = new_wb
        if removed:
            _drop_relationships(pkg, out_parts, wbp, removed)
        sst_bytes = sst.serialize()
        if sst_bytes is not None:
            out_parts[pkg.real(sst_part)] = sst_bytes
        st_bytes = styles.serialize()
        if st_bytes is not None:
            out_parts[pkg.real(st_part)] = st_bytes
        return BuildResult(out_parts, removed, changed_sheets, written,
                           formulas_written, stale, full_calc,
                           len(sst.added), len(styles._new_xfs))


_TABLE_ROOT_RE = re.compile(rb"<((?:[A-Za-z_][\w.\-]*:)?)table(?=[\s/>])((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>")
_TABLE_COL_RE = re.compile(rb"<((?:[A-Za-z_][\w.\-]*:)?)tableColumn(?=[\s/>])((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>")


def _rename_table_columns(pkg: Package, sheet_parts: dict, renames: list,
                          out_parts: dict) -> None:
    """A header cell typed over is a column rename (core/package.py
    _settle_tables follows Excel there); write the new column name into the
    table part, touching nothing else in it."""
    by_table: dict = {}
    for tr in renames:
        by_table.setdefault((tr.sheet, tr.table), []).append(tr)
    for (sheet, table), items in by_table.items():
        sp = sheet_parts.get(sheet)
        found = None
        for tp in pkg.related(sp, "/table") if sp else []:
            if not pkg.has(tp):
                continue
            data = pkg.read(tp)
            m = _TABLE_ROOT_RE.search(data)
            if m is None:
                continue
            a = _attrs(m.group(2))
            names = {html.unescape(a.get(k, b"").decode("utf-8")) for k in ("name", "displayName")}
            if table in names:
                found = (pkg.real(tp), data)
                break
        if found is None:
            raise NotSurgical(f"table {table!r} has no table part")
        tp, data = found
        cols = list(_TABLE_COL_RE.finditer(data))
        edits = []
        for tr in items:
            if tr.index >= len(cols):
                raise NotSurgical(f"table {table!r} does not match the model")
            m = cols[tr.index]
            cur = _attrs(m.group(2)).get("name")
            if cur is None or html.unescape(cur.decode("utf-8")) != tr.old:
                raise NotSurgical(f"table {table!r} does not match the model")
            edits.append((m.start(), m.end(),
                          _set_attr(m.group(0), b"name", xml_attr(tr.new).encode("utf-8"))))
        out_parts[tp] = _apply_edits(data, edits)


def _rc(coord: str) -> tuple[int, int]:
    rc = _split_coord(coord.encode())
    return rc if rc is not None else (0, 0)


def _materialize_groups(data: bytes, ws, scan: SheetScan, groups: set[str],
                        replaced: dict, prefix: bytes) -> list:
    """Rewrite every member of the shared-formula groups the edit touched as
    a plain formula (keeping its cached value), so no dependent is left
    pointing at a master that no longer exists."""
    full = scan_sheet(data, [], all_cells=True)
    edits = []
    for start, end, col, row in full.all_cells:
        if (row, col) in replaced:
            continue
        fm = _F_ELEM_RE.search(data, start, end)
        if fm is None:
            continue
        a = {k: v.decode("utf-8", "replace") for k, v in _attrs(fm.group(2)).items()}
        if a.get("t") != "shared" or a.get("si") not in groups:
            continue
        cell = ws._cells.get((row, col))
        text = cell._value if cell is not None else None
        if not isinstance(text, str) or not text.startswith("="):
            raise NotSurgical("a shared formula could not be expanded")
        loc = CellLoc(col, start, data.find(b">", start) + 1, end, b"", False)
        orig = read_orig_cell(data, loc)
        if orig.f_span is None:
            continue
        edits.append((orig.f_span[0], orig.f_span[1],
                      _materialize_f(data, orig.f_span, fm.group(1), text, a)))
    return edits


def _drop_relationships(pkg: Package, out_parts: dict, owner: str, removed: set) -> None:
    rp = pkg.real(rels_path(owner))
    if rp is not None:
        data = out_parts.get(rp, pkg.read(rp))
        edits = []
        for m in _REL_RE.finditer(data):
            a = _attrs(m.group(1))
            tgt = a.get("Target")
            if tgt is None:
                continue
            part = resolve_target(owner, html.unescape(tgt.decode("utf-8")))
            if pkg.real(part) in removed:
                end = m.end()
                edits.append((m.start(), end, b""))
        if edits:
            out_parts[rp] = _apply_edits(data, edits)
    ct = pkg.real("[Content_Types].xml")
    if ct is not None:
        data = out_parts.get(ct, pkg.read(ct))
        edits = []
        for m in re.finditer(rb"<(?:[\w.\-]+:)?Override\b((?:[^>\"']|\"[^\"]*\"|'[^']*')*)>", data):
            a = _attrs(m.group(1))
            pn = a.get("PartName")
            if pn is None:
                continue
            name = html.unescape(pn.decode("utf-8")).lstrip("/")
            if pkg.real(name) in removed:
                edits.append((m.start(), m.end(), b""))
        if edits:
            out_parts[ct] = _apply_edits(data, edits)


def splice_check(path: str) -> str | None:
    """None when every worksheet of the file can take a cell splice, else
    the first reason it cannot (for diagnose_workbook's route report)."""
    try:
        with zipfile.ZipFile(path) as zf:
            pkg = Package(zf)
            for title, part in pkg.sheet_parts().items():
                if not pkg.has(part):
                    return f"sheet {title!r} has no worksheet part"
                data = pkg.read(part)
                _decode_part(data)
                scan_sheet(data, [])
            wbp = pkg.workbook_part()
            for related in ("/sharedStrings", "/styles"):
                for p in pkg.related(wbp, related):
                    if pkg.has(p):
                        data = pkg.read(p)
                        if related == "/sharedStrings":
                            SharedStrings(p, data)
    except NotSurgical as exc:
        return str(exc)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        return f"the package could not be read ({type(exc).__name__})"
    return None


def write_package(src_path: str, dst_path: str, result: BuildResult) -> None:
    """Copy the original package to dst member by member, substituting the
    rebuilt parts and leaving out the removed ones. Member order, names,
    compression and every untouched part's bytes stay as they were."""
    with zipfile.ZipFile(src_path) as zin, \
            zipfile.ZipFile(dst_path, "w", allowZip64=True) as zout:
        zout.comment = zin.comment
        for info in zin.infolist():
            if info.filename in result.removed:
                continue
            data = result.parts.get(info.filename)
            if data is None:
                data = zin.read(info.filename)
            ni = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            ni.compress_type = info.compress_type
            ni.external_attr = info.external_attr
            ni.create_system = info.create_system
            ni.comment = info.comment
            ni.extra = b""
            zout.writestr(ni, data)
        have = {i.filename for i in zin.infolist()}
        for name, data in result.parts.items():
            if name not in have:
                ni = zipfile.ZipInfo(name)
                ni.compress_type = zipfile.ZIP_DEFLATED
                zout.writestr(ni, data)


__all__ = [
    "ROUTE_RAW", "ROUTE_REWRITE", "NotSurgical", "snapshot", "plan_edit",
    "cell_changes", "check_changed_formulas", "splice_check",
    "build", "write_package", "stale_formulas", "Staleness", "BuildResult",
    "SurgicalPlan", "WorkbookSnap", "cell_key", "set_full_calc", "Package",
    "parse_rels", "resolve_target", "rels_path", "scan_sheet",
]
