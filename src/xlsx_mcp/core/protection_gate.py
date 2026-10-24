"""Sheet and workbook protection is honoured by every write (the same
design as the Word editor).

Before, every write went straight through protection: set_cell changed a
locked cell of a protected sheet, manage_worksheet renamed a sheet of a
structure-protected workbook, with no check and no warning.

ONE check, at the one place every file-tier write passes, before anything
reaches the disk (`check`, called by WorkbookPackage.save): the workbook as
the call changed it is compared with the workbook as it loaded, and only
what Excel itself would refuse is refused, with the code PROTECTED.

  workbook structure locked   no sheet may be added, deleted, renamed,
                              moved, copied, hidden or unhidden
  a protected sheet           no LOCKED cell may change its value or
                              formula (unlocked cells stay editable, as in
                              Excel); and, unless the sheet's protection
                              allows it (set_protection options):
                                formatting of any cell, merges,
                                conditional formats      format_cells
                                row heights / hidden     format_rows
                                column widths / hidden   format_columns
                                insert / delete rows     insert_rows,
                                                         delete_rows (and no
                                                         locked value in a
                                                         deleted row)
                                same for columns         insert_columns,
                                                         delete_columns
                                hyperlinks               insert_hyperlinks
                                comments, images, charts objects
                                the autofilter           auto_filter
                                sort_range               sort
                              data validation and tables never change on a
                              protected sheet (Excel offers no option)

A formula whose text changed only because references were rewritten (rows
inserted on another sheet, a range moved) is not an edit of that cell; Excel
adjusts references on protected sheets too. Writes an op RECORDS (the
package's intended map and written regions) always count.

The check is armed only during a dispatched writing call (`dispatched_call`,
from the surface dispatcher), exactly like Word's backstop; in-process engine
use keeps its raw behaviour. set_protection (lifting or changing the
protection) and manage_backups (restoring a backup) stay callable. The xl_app
ops edit through Excel, which enforces protection itself.
"""

from __future__ import annotations

import contextlib
import contextvars
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator

from .errors import SurfaceRefusal

#: Ops always callable on a protected workbook (lift it; restore a backup).
ALWAYS_ALLOWED = frozenset({"set_protection", "manage_backups"})

#: set_protection's option names (True = allowed while protected) -> the
#: openpyxl SheetProtection attribute (True = blocked). The same table as
#: ops/protection.SHEET_OPTIONS, which this module cannot import (cycle).
OPTIONS = {
    "format_cells": "formatCells", "format_columns": "formatColumns",
    "format_rows": "formatRows", "insert_columns": "insertColumns",
    "insert_rows": "insertRows", "insert_hyperlinks": "insertHyperlinks",
    "delete_columns": "deleteColumns", "delete_rows": "deleteRows",
    "sort": "sort", "auto_filter": "autoFilter", "objects": "objects",
}

_DEFAULT_STYLE = (0,) * 9
_LIST_CELLS = 10


class _Gate:
    def __init__(self, op: str) -> None:
        self.op = op


_GATE: contextvars.ContextVar[_Gate | None] = contextvars.ContextVar(
    "ks4xl_protection_gate", default=None)


@contextlib.contextmanager
def dispatched_call(op: str) -> Iterator[_Gate]:
    """Arm the check for the duration of one dispatched writing call."""
    gate = _Gate(op)
    token = _GATE.set(gate)
    try:
        yield gate
    finally:
        _GATE.reset(token)


def armed() -> _Gate | None:
    gate = _GATE.get()
    if gate is None or gate.op in ALWAYS_ALLOWED:
        return None
    return gate


# ------------------------------------------------------------------ snapshot


def _vkey(v: Any) -> tuple:
    if v is None or isinstance(v, (str, int, float, bool)):
        return (type(v).__name__, v)
    return (type(v).__name__, str(getattr(v, "text", v)), str(getattr(v, "ref", "")))


def _cell_key(cell) -> tuple:
    style = cell._style
    hl = cell.hyperlink
    cm = cell.comment
    return (_vkey(cell.value),
            tuple(style) if style is not None else _DEFAULT_STYLE,
            (hl.target, hl.location, hl.tooltip) if hl is not None else None,
            (cm.text, cm.author) if cm is not None else None)


_MISSING = (_vkey(None), _DEFAULT_STYLE, None, None)


def _rows(ws) -> dict:
    out = {}
    for r, d in ws.row_dimensions.items():
        key = (d.ht, bool(d.hidden), d.outlineLevel or 0, bool(d.collapsed))
        if key != (None, False, 0, False):
            out[r] = key
    return out


def _cols(ws) -> dict:
    out = {}
    for k, d in ws.column_dimensions.items():
        key = (d.min, d.max, round(float(d.width or 0), 4), bool(d.hidden),
               d.outline_level or 0, bool(d.collapsed))
        if key[2:] != (13.0, False, 0, False):
            out[k] = key
    return out


def _xml(obj) -> str:
    from openpyxl.xml.functions import tostring

    try:
        return tostring(obj.to_tree()).decode("utf-8", "replace")
    except Exception:  # noqa: BLE001 (a model that will not serialize compares by repr)
        return repr(obj)


def _cf(ws) -> tuple:
    return tuple((str(cf.sqref), tuple(_xml(r) for r in rules))
                 for cf, rules in ws.conditional_formatting._cf_rules.items())


def _dv(ws) -> tuple:
    return tuple(_xml(dv) for dv in ws.data_validations.dataValidation)


def _autofilter(ws):
    af = ws.auto_filter
    return _xml(af) if af is not None and af.ref else None


@dataclass
class SheetSnap:
    blocked: dict[str, bool]
    cells: dict[tuple[int, int], tuple]
    locked: dict[tuple[int, int], bool]
    merges: frozenset
    rows: dict
    cols: dict
    cf: tuple
    dv: tuple
    tables: tuple
    objects: tuple
    autofilter: Any

    @classmethod
    def of(cls, ws) -> "SheetSnap":
        prot = ws.protection
        cells, locked = {}, {}
        for key, cell in list(ws._cells.items()):
            cells[key] = _cell_key(cell)
            locked[key] = bool(cell.protection.locked)
        return cls(
            blocked={opt: bool(getattr(prot, attr, True)) for opt, attr in OPTIONS.items()},
            cells=cells, locked=locked,
            merges=frozenset(str(r) for r in ws.merged_cells.ranges),
            rows=_rows(ws), cols=_cols(ws), cf=_cf(ws), dv=_dv(ws),
            tables=tuple(sorted(ws.tables.items())),  # TableList.items(): (name, ref)
            objects=(len(ws._images), len(ws._charts)),
            autofilter=_autofilter(ws))


@dataclass
class BookSnap:
    structure_locked: bool
    sheets: tuple
    protected: dict[int, tuple[Any, SheetSnap]] = field(default_factory=dict)


def _sheet_list(wb) -> tuple:
    return tuple((ws.title, ws.sheet_state) for ws in wb._sheets)


def snapshot(wb) -> BookSnap | None:
    """The protection-relevant state of a freshly loaded workbook, or None
    when nothing in it is protected (or no dispatched write is running)."""
    if armed() is None:
        return None
    sec = wb.security
    structure = bool(sec is not None and sec.lockStructure)
    protected = {id(ws): (ws, SheetSnap.of(ws)) for ws in wb.worksheets if ws.protection.sheet}
    if not structure and not protected:
        return None
    return BookSnap(structure, _sheet_list(wb), protected)


# ------------------------------------------------------------------- check


def _default_locked(ws, row: int, col: int) -> bool:
    """Whether a cell that does not exist yet is locked: its column's or
    row's style decides when one is set, else Excel's default (locked)."""
    for d in ws.column_dimensions.values():
        if d.min and d.max and d.min <= col <= d.max and d.has_style:
            return bool(d.protection.locked)
    rd = ws.row_dimensions.get(row) if hasattr(ws.row_dimensions, "get") else None
    if rd is not None and rd.has_style:
        return bool(rd.protection.locked)
    return True


def _a1(row: int, col: int) -> str:
    from openpyxl.utils import get_column_letter

    return f"{get_column_letter(col)}{row}"


def _intended(pkg, title: str, key: tuple[int, int]) -> bool:
    a1 = _a1(*key)
    if (title, a1) in pkg._intended:
        return True
    r, c = key
    return any(s == title and r0 <= r <= r1 and c0 <= c <= c1
               for s, r0, c0, r1, c1 in pkg._written_regions)


def _is_formula(vkey: tuple) -> bool:
    kind, value = vkey[0], vkey[1]
    return (kind == "str" and isinstance(value, str) and value.startswith("=")) or \
        kind in ("ArrayFormula", "DataTableFormula")


def refusal(pkg, sheet: str | None, what: str, op: str, *, cells: list[str] | None = None,
            allowed: list[str] | None = None) -> SurfaceRefusal:
    name = Path(pkg.path).name
    if sheet is None:
        message = (f"{name}'s workbook structure is protected, so {op} cannot add, delete, "
                   "rename, move, copy, hide or unhide sheets. Nothing was changed.")
        details = {"op": op, "protection": "structure", "file": str(pkg.path)}
    else:
        also = (f" This sheet's protection allows: {', '.join(allowed)}." if allowed else "")
        message = (f"sheet {sheet!r} of {name} is protected, so {op} cannot change {what}."
                   f"{also} Cells that are not locked can still be edited. Nothing was "
                   "changed.")
        details = {"op": op, "protection": "sheet", "sheet": sheet, "file": str(pkg.path)}
        if cells:
            details["cells"] = cells[:_LIST_CELLS]
            if len(cells) > _LIST_CELLS:
                details["cells_total"] = len(cells)
    return SurfaceRefusal("PROTECTED", message, details=details)


def _allowed(snap: SheetSnap) -> list[str]:
    return sorted(o for o, blocked in snap.blocked.items() if not blocked)


def check_sheet(pkg, ws, snap: SheetSnap, op: str) -> None:
    """Refuse PROTECTED when this call changed what the sheet's protection
    forbids (module docstring)."""
    title = ws.title
    allowed = _allowed(snap)

    def no(what, option=None, cells=None):
        if option is None or snap.blocked.get(option, True):
            raise refusal(pkg, title, what, op, cells=cells, allowed=allowed)

    now = {key: _cell_key(cell) for key, cell in list(ws._cells.items())}
    locked_edits, unlocked_style, links, comments, any_value = [], [], [], [], False
    for key in sorted(set(now) | set(snap.cells)):
        old = snap.cells.get(key, _MISSING)
        new = now.get(key, _MISSING)
        if old == new:
            continue
        value_changed = old[0] != new[0]
        if value_changed and _is_formula(old[0]) and _is_formula(new[0]) \
                and not _intended(pkg, title, key):
            value_changed = False  # a reference rewrite, which Excel makes too
        locked = snap.locked.get(key)
        if locked is None:
            locked = _default_locked(ws, *key)
        if value_changed:
            any_value = True
            if locked:
                locked_edits.append(_a1(*key))
        if old[1] != new[1]:
            unlocked_style.append(_a1(*key))
        if old[2] != new[2]:
            links.append(_a1(*key))
        if old[3] != new[3]:
            comments.append(_a1(*key))
    if locked_edits:
        no("locked cells (" + ", ".join(locked_edits[:_LIST_CELLS])
           + (", ..." if len(locked_edits) > _LIST_CELLS else "") + ")", cells=locked_edits)
    if unlocked_style:
        no("cell formatting", "format_cells", unlocked_style)
    if links:
        no("hyperlinks", "insert_hyperlinks", links)
    if comments:
        no("comments", "objects", comments)
    if frozenset(str(r) for r in ws.merged_cells.ranges) != snap.merges:
        no("merged cells", "format_cells")
    if _cf(ws) != snap.cf:
        no("conditional formats", "format_cells")
    if _dv(ws) != snap.dv:
        no("data validation")
    if tuple(sorted(ws.tables.items())) != snap.tables:
        no("tables")
    if (len(ws._images), len(ws._charts)) != snap.objects:
        no("images or charts", "objects")
    filter_changed = _autofilter(ws) != snap.autofilter
    if filter_changed:
        no("the autofilter", "auto_filter")
    rows = _rows(ws)
    if rows != snap.rows:
        def _no_hidden(d: dict) -> dict:
            out = {k: (v[0],) + v[2:] for k, v in d.items()}
            return {k: v for k, v in out.items() if v != (None, 0, False)}

        # Rows a filter hides or shows are the filter's change (auto_filter,
        # judged above), not a row format.
        by_filter = (filter_changed or op in ("set_filter", "clear_filter")) \
            and _no_hidden(rows) == _no_hidden(snap.rows)
        if by_filter:
            no("the rows the filter hides", "auto_filter")
        else:
            no("row heights or hidden rows", "format_rows")
    if _cols(ws) != snap.cols:
        no("column widths or hidden columns", "format_columns")
    if op == "sort_range" and any_value:
        no("the order of its rows", "sort")


def check(pkg) -> None:
    """Called by WorkbookPackage.save before any disk change."""
    gate = armed()
    snap = getattr(pkg, "_prot_snap", None)
    if gate is None or snap is None or pkg._workbook is None:
        return
    wb = pkg._workbook
    if snap.structure_locked and _sheet_list(wb) != snap.sheets:
        raise refusal(pkg, None, "", gate.op)
    live = {id(ws) for ws in wb.worksheets}
    for wid, (ws, ss) in snap.protected.items():
        if wid in live:
            check_sheet(pkg, ws, ss, gate.op)


def before_structure_edit(pkg, ws, kind: str, index: int, count: int) -> None:
    """Called by WorkbookPackage.modify_structure before rows or columns are
    inserted or deleted: judge what changed so far on this sheet, then the
    edit itself against the sheet's protection."""
    gate = armed()
    snap = getattr(pkg, "_prot_snap", None)
    if gate is None or snap is None:
        return
    entry = snap.protected.get(id(ws))
    if entry is None:
        return
    ss = entry[1]
    check_sheet(pkg, ws, ss, gate.op)
    option = {"insert_rows": "insert_rows", "delete_rows": "delete_rows",
              "insert_cols": "insert_columns", "delete_cols": "delete_columns"}.get(kind)
    words = {"insert_rows": "rows (insert)", "delete_rows": "rows (delete)",
             "insert_cols": "columns (insert)", "delete_cols": "columns (delete)"}.get(kind, kind)
    if option is None or ss.blocked.get(option, True):
        raise refusal(pkg, ws.title, words, gate.op, allowed=_allowed(ss))
    if kind in ("delete_rows", "delete_cols"):
        hit = []
        for (r, c), cell in list(ws._cells.items()):
            inside = (index <= r < index + count) if kind == "delete_rows" else \
                (index <= c < index + count)
            if inside and cell.value is not None and ss.locked.get((r, c), True):
                hit.append(_a1(r, c))
        if hit:
            raise refusal(pkg, ws.title, "locked cells in the "
                          + ("rows" if kind == "delete_rows" else "columns")
                          + " it would delete (" + ", ".join(hit[:_LIST_CELLS]) + ")",
                          gate.op, cells=hit, allowed=_allowed(ss))


def after_structure_edit(pkg, ws) -> None:
    """The edit moved this sheet's cells: judge later changes against the
    sheet as it now stands (its locks travelled with the cells)."""
    snap = getattr(pkg, "_prot_snap", None)
    if snap is None or id(ws) not in snap.protected:
        return
    snap.protected[id(ws)] = (ws, SheetSnap.of(ws))


__all__ = ["dispatched_call", "armed", "snapshot", "check", "before_structure_edit",
           "after_structure_edit", "ALWAYS_ALLOWED", "OPTIONS"]
