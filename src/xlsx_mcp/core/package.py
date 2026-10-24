"""core/package.py: WorkbookPackage, the central open/mutate/save object.

Every file-based tool opens a workbook through this object, mutates it, and
saves through it, so the safety core runs end to end on every write without
per-tool wiring. It integrates:

  - the HAZARD SCAN (core.hazard) on open, and the routing decision on save:
    a mutation of a workbook with drop-risk parts REFUSES with HAZARD_REFUSED,
    naming the at-risk parts and the remedy (the COM route when it exists, or
    an explicit allow_loss override), unless the mutation goes through a loss-
    safe path (a clean workbook, or a .xlsm whose only hazard is VBA, preserved
    via keep_vba). Drop-risk includes content with no part of its own: a
    worksheet extLst carrying x14 conditional formatting, sparklines or a
    slicer list refuses on the same terms, since openpyxl writes the worksheet
    back without it.
  - BACKUP-BEFORE-MUTATION (core.safesave two-slot backups) captured against
    the pre-mutation state immediately before promotion, plus the per-file
    write lock held across the whole read-modify-verify-save cycle.
  - ATOMIC VALIDATED SAVE with VERIFY-AFTER-WRITE (core.verify): the mutation
    is written to a temp path and verified (structural re-open, DEFAULT-FAIL
    part-inventory diff vs the pre-write scan so ANY unexplained part loss
    refuses, the fragile-part checks, content read-back against intent)
    BEFORE the original is touched. Ops that deliberately remove parts at the
    model level (sheet delete, table to_range, comment delete) declare it via
    expect_removal() so the inventory diff does not mistake the removal for
    silent loss. A failed verify leaves the original unmodified and refuses
    with VALIDATION_FAILED. After promotion a second verify runs; if it
    fails, the pre-mutation backup is restored.
  - FORMULA-WRITE SAFETY: any formula written through the package is
    normalized (core.calc._xlfn/_xlpm shim) and the workbook is marked
    fullCalcOnLoad so the next Excel/LibreOffice open recalculates.
  - STRUCTURAL-EDIT INTEGRITY: insert/delete rows/columns rewrite every
    reference (core.refs) so the workbook stays coherent, then pass verify.
  - DIRTY/LOCK HANDLING: a workbook held open in Excel surfaces as
    WORKBOOK_LOCKED (the PermissionError path) rather than a
    raw traceback.

The object addresses cells through core.locate (the grid location resolver).
"""

from __future__ import annotations

import math
import os
import re
import uuid
from dataclasses import replace as _dc_replace
from pathlib import Path
from typing import Any, Callable

from . import arrays as _arrays
from . import calc as _calc
from . import carry as _carry
from . import hazard as _hazard
from . import limits as _limits
from . import linkprops as _linkprops
from . import locate as _locate
from . import protection_gate as _protection_gate
from . import refs as _refs
from . import safesave as _safesave
from . import surgical as _surgical
from . import verify as _verify
from . import xmlsafe as _xmlsafe
from .errors import (
    HazardRefused,
    UnsupportedStructure,
    ValidationFailed,
    WorkbookCorrupt,
    WorkbookLocked,
    XlMcpError,
)
from .sandbox import check_path


# --------------------------------------------------- unwritable-target copy
# Two different situations produced one message for the life of 1.0; the
# diagnosis below tells them apart.


def _refuse_table(message: str, detail: dict | None = None) -> None:
    """The table refusal: the save would leave a table Excel
    will not open without a repair."""
    from .errors import UnsupportedStructure

    exc = UnsupportedStructure(message)
    if detail:
        exc.detail = detail
    raise exc


def _names_column(wb, column: str, table: str | None = None) -> list[str]:
    """Where a formula or defined name names a table column ([col], [@col],
    T[[#This Row],[col]], and the escaped forms Excel writes for names with
    [ ] # or ', like T[x'#]), for the table rename check. References
    are read through core.structrefs, so an escaped reference is
    seen; with table given, only references to that table count (qualified,
    or unqualified in a cell inside it)."""
    from . import structrefs as _sr

    tables = _sr.tables_of(wb)
    names = [table] if table else [t.name for t in tables.values()]
    hits = []
    for ws in wb.worksheets:
        for (r, c), cell in list(getattr(ws, "_cells", {}).items()):
            v = cell.value
            text = v if isinstance(v, str) else getattr(v, "text", None)
            if isinstance(text, str) and text.startswith("=") and any(
                    _sr.columns_named(text, t, column, sheet=ws.title, row=r, col=c,
                                      tables=tables) for t in names):
                hits.append(f"{ws.title}!{cell.coordinate}")
    for _scope, dname, defn in _locate._all_defined_names(wb):
        text = str(getattr(defn, "attr_text", "") or "")
        if any(_sr.columns_named(text, t, column, tables=tables) for t in names):
            hits.append(dname)
    return hits


#: A save's answer lists each written cell only up to
#: this many; past it the answer carries a summary (count, kinds, the
#: rectangle per sheet, the first few cells). A 200 x 7 write used to echo
#: 1,400 entries, about 17k tokens. echo_cells(True) asks for the full list.
ECHO_LIMIT = 20

import contextlib as _contextlib  # noqa: E402
import contextvars as _contextvars  # noqa: E402

#: Set by the dispatcher for a call that takes an
#: automatic snapshot ({"op": name, "label": None}); the first in-place save
#: of the call fills label at its point of no return.
_AUTO_SNAPSHOT: _contextvars.ContextVar[dict | None] = _contextvars.ContextVar(
    "ks4xl_auto_snapshot", default=None)


@_contextlib.contextmanager
def auto_snapshot(op: str | None):
    """Arm one call's automatic snapshot; yields the state the save fills
    (label stays None when no save happened)."""
    state = {"op": op, "label": None} if op else None
    token = _AUTO_SNAPSHOT.set(state)
    try:
        yield state
    finally:
        _AUTO_SNAPSHOT.reset(token)


_ECHO_CELLS: _contextvars.ContextVar[bool] = _contextvars.ContextVar(
    "ks4xl_echo_cells", default=False)


@_contextlib.contextmanager
def echo_cells(on: bool):
    """Within this block, a save lists every written cell."""
    token = _ECHO_CELLS.set(bool(on))
    try:
        yield
    finally:
        _ECHO_CELLS.reset(token)


def _cells_summary(entries: list[dict]) -> dict:
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.cell import coordinate_to_tuple

    kinds: dict[str, int] = {}
    boxes: dict[str, list[int]] = {}
    for e in entries:
        kinds[e["kind"]] = kinds.get(e["kind"], 0) + 1
        r, c = coordinate_to_tuple(e["cell"])
        b = boxes.setdefault(e["sheet"], [r, c, r, c])
        b[0], b[1], b[2], b[3] = min(b[0], r), min(b[1], c), max(b[2], r), max(b[3], c)
    ranges = [f"{sheet}!{get_column_letter(b[1])}{b[0]}:{get_column_letter(b[3])}{b[2]}"
              for sheet, b in boxes.items()]
    return {"count": len(entries), "kinds": kinds, "ranges": ranges,
            "first": entries[:5],
            "note": "the full list is not repeated; pass echo_cells=true to get it"}


def parse_iso_date(raw):
    """An ISO date ("2026-03-24") or date-time ("2026-03-24T09:30[:15]", a
    space works for the T) as a date or datetime, or None."""
    import datetime as _dt
    import re as _re

    if not isinstance(raw, str):
        return None
    text = raw.strip()
    if _re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        try:
            return _dt.date.fromisoformat(text)
        except ValueError:
            return None
    if _re.fullmatch(r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(:\d{2})?", text):
        try:
            return _dt.datetime.fromisoformat(text.replace(" ", "T"))
        except ValueError:
            return None
    return None


#: A plain numeral as typed: optional sign, digits with or
#: without thousands commas, optional decimals or exponent. No currency
#: sign, percent sign or unit: those are ambiguous, and Excel itself would
#: not take "12 kg" for a number.
_NUMERAL = re.compile(
    r"[+-]?(?:(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?|\.\d+)(?:[eE][+-]?\d+)?")


#: A whole number from here up cannot be kept digit for digit as
#: an Excel number (a double has 53 bits of integer).
EXACT_NUMBER_LIMIT = 2 ** 53


def _numbers_beside(ws, coord: str, reach: int = 3) -> bool:
    """Does the column hold a number within `reach` rows of this cell? Reads
    the cells already in the model (ws._cells) so that looking never creates
    one."""
    from openpyxl.utils.cell import coordinate_to_tuple

    row, col = coordinate_to_tuple(coord)
    cells = getattr(ws, "_cells", {})
    for r in [row + d for d in range(-reach, reach + 1) if d != 0]:
        cell = cells.get((r, col)) if r >= 1 else None
        v = cell.value if cell is not None else None
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return True
    return False


def _cell_is_new(ws, coord: str) -> bool:
    """The cell is not in the model yet. Reads ws._cells, so looking never
    creates one."""
    from openpyxl.utils.cell import coordinate_to_tuple

    return coordinate_to_tuple(coord.replace("$", "")) not in getattr(ws, "_cells", {})


def _formatted_as_text(ws, coord: str, new: bool) -> bool:
    """Is the cell formatted as Text, by its own number format or, for a
    cell this write creates, by the format the save gives it: its row's
    format when the row has one, otherwise its column's, as Excel does for
    a value typed there (core.surgical). Formatting a whole column as Text
    is how a sheet marks an ID or code column, so a numeral written there
    is text on purpose ("250" written into such a column was warned about
    as a number stored as text)."""
    from openpyxl.utils.cell import coordinate_to_tuple

    cell = ws[coord]
    if cell.number_format == "@":
        return True
    if not new or cell.has_style:
        return False
    row, col = coordinate_to_tuple(coord.replace("$", ""))
    row_dim = ws.row_dimensions.get(row)
    if row_dim is not None and row_dim.has_style:
        return row_dim.number_format == "@"
    for col_dim in ws.column_dimensions.values():
        lo = col_dim.min or 0
        hi = col_dim.max or lo
        if lo <= col <= hi and col_dim.has_style:
            return col_dim.number_format == "@"
    return False


def looks_like_number(text: str) -> bool:
    """Is this string a number somebody typed as text?

    Codes and IDs are not: a numeral with a leading zero ("00123", "007")
    keeps its zeros only as text, and one of sixteen or more significant
    digits cannot be a faithful number (the advice is to send it as a
    string), so neither earns the "pass a number instead" warning."""
    s = text.strip()
    if not _NUMERAL.fullmatch(s):
        return False
    if re.match(r"[+-]?0\d", s):
        return False
    mantissa = re.split(r"[eE]", s, maxsplit=1)[0]
    return len(re.sub(r"\D", "", mantissa).lstrip("0")) <= 15


def held_refusal(path: str) -> str:
    """Another process has the file open and will not share the write. It
    says "open in Excel" only when Excel's owner file (~$name) is beside the
    workbook; otherwise it names another program."""
    from . import access as _access
    return _access.in_use_text(path)


def read_only_refusal(path: str) -> str:
    """The file itself is marked read-only. Nothing has it open."""
    return (f"{Path(path).name}: this file's read-only attribute is set, so "
            "saving to it would fail. No program is holding it; the attribute "
            "is a property of the file itself. Clear it (right-click, "
            "Properties, untick Read-only) and try again, or save a copy "
            "elsewhere with copy_workbook")


def _fill_calculated_columns(ws, edit, tables_before: dict) -> list[str]:
    """Rows inserted inside a table take its calculated columns' formula, as
    Excel's own Insert does (measured: a sheet row inserted into a table's
    data rows, or at its totals row, gets the column formula in every
    calculated column and nothing in the others; a row inserted at the
    header or below the last row is not the table's and gets nothing).
    tables_before maps each table on the edited sheet to its range before
    the edit. Returns the cells written."""
    from openpyxl.utils.cell import range_boundaries
    from openpyxl.utils import get_column_letter

    filled: list[str] = []
    for name, old_ref in tables_before.items():
        if name not in ws.tables:
            continue
        try:
            _c1, top, _c2, bottom = range_boundaries(old_ref)
        except (TypeError, ValueError):
            continue
        if not top < edit.index <= bottom:
            continue  # at the header or past the last row: not the table's
        table = ws.tables[name]
        try:
            min_col, new_top, _mc, _mr = range_boundaries(table.ref)
        except (TypeError, ValueError):
            continue
        header = table.headerRowCount if table.headerRowCount is not None \
            else 1
        data_top = new_top + header
        for i, tc in enumerate(table.tableColumns or ()):
            calc = getattr(getattr(tc, "calculatedColumnFormula", None),
                           "attr_text", None)
            if not isinstance(calc, str) or not calc.strip():
                continue
            for r in range(edit.index, edit.index + edit.count):
                # the column formula is stored for the first data row; a
                # relative reference in it follows the row, as Excel fills
                text = _refs.offset_formula("=" + calc, r - data_top, 0,
                                            ws.title)
                coord = f"{get_column_letter(min_col + i)}{r}"
                ws[coord] = text
                filled.append(coord)
    return filled


def _inherited(dim):
    """What an inserted row or column takes from its neighbour. Measured with
    Excel's own Insert (CopyOrigin left or above, its default): the height or
    width, the style and the outline level carry over, the hidden state and a
    group's collapsed mark do not."""
    from copy import copy

    new = copy(dim)
    new._style = copy(dim._style)
    new.hidden = False
    new.collapsed = False
    return new


def _shift_row_dimensions(ws, edit) -> None:
    """openpyxl moves cells on insert_rows/delete_rows but not
    the row settings (height, hidden, outline level, style), so a tall row's
    height stayed at its old number and landed on whatever moved there.
    Excel moves them with their rows; so does this. An inserted row takes the
    settings of the row above it (_inherited), as Excel's own Insert does;
    one inserted at row 1 takes none."""
    dims = ws.row_dimensions
    old = list(dims.items())
    if not old:
        return
    lo, n = edit.index, edit.count
    moved = {}
    for r, rd in old:
        if r < lo:
            moved[r] = rd
        elif edit.kind == _refs.INSERT_ROWS:
            if r + n <= _refs.MAX_ROW:
                moved[r + n] = rd
        elif r >= lo + n:
            moved[r - n] = rd
        # rows inside a deleted span go with it
    if edit.kind == _refs.INSERT_ROWS and lo > 1 and (lo - 1) in dims:
        above = dims[lo - 1]
        for r in range(lo, min(lo + n, _refs.MAX_ROW + 1)):
            moved[r] = _inherited(above)
    dims.clear()
    for r, rd in sorted(moved.items()):
        rd.index = r
        dims[r] = rd


def _shift_column_dimensions(ws, edit) -> None:
    """openpyxl moves cells on insert_cols/delete_cols but not the column
    settings (width, hidden, outline level, style). A wide column's width, a
    hidden column and a group's outline stayed at their old letters and
    landed on whatever moved there. Excel moves them with their columns, and
    so does this, span by span (one <col> can cover many columns):
      - an insert inside a span splits it around the new columns;
      - a delete through a span shortens it;
      - a span that reaches the last column (16384) still reaches it after a
        delete, and is cut at the edge by an insert;
      - an inserted column takes the settings of the column to its left
        (_inherited), as Excel's own Insert does; one inserted at column A
        takes none."""
    from openpyxl.utils import column_index_from_string, get_column_letter

    dims = ws.column_dimensions
    spans = []
    for cd in list(dims.values()):
        lo_, hi_ = cd.min, cd.max
        if not lo_ or not hi_:
            lo_ = hi_ = column_index_from_string(cd.index)
        spans.append((lo_, hi_, cd))
    if not spans:
        return
    last = _refs.MAX_COL
    at, n = edit.index, edit.count
    out = []

    def piece(cd, a, b, inherit=False):
        b = min(b, last)
        if a > b:
            return
        if inherit:
            new = _inherited(cd)
        else:
            from copy import copy
            new = copy(cd)
            new._style = copy(cd._style)
        new.min, new.max = a, b
        new.index = get_column_letter(a)
        out.append(new)

    if edit.kind == _refs.INSERT_COLS:
        for a, b, cd in spans:
            if b < at:
                piece(cd, a, b)
            elif a >= at:
                piece(cd, a + n, b + n)
            else:
                piece(cd, a, at - 1)
                piece(cd, at + n, b + n)
        if at > 1:
            left = next((cd for a, b, cd in spans if a <= at - 1 <= b), None)
            if left is not None:
                piece(left, at, at + n - 1, inherit=True)
    else:
        end = at + n - 1
        for a, b, cd in spans:
            nb = last if b == last else b - n
            if b < at:
                piece(cd, a, b)
            elif a > end:
                piece(cd, a - n, nb)
            else:
                if a < at:
                    piece(cd, a, at - 1)
                if b > end:
                    piece(cd, at, nb)
    dims.clear()
    for cd in sorted(out, key=lambda d: d.min):
        dims[cd.index] = cd


def _table_growth(ws, index: int) -> tuple[tuple[int, int, int, int, int, int], ...]:
    """(top_row, first_row, last_row, bottom_row, min_col, max_col) of each
    table on the sheet that rows inserted at `index` grow: at its first data
    row, or at its totals row (the rows go above it)."""
    from openpyxl.utils.cell import range_boundaries
    out = []
    for name in list(ws.tables):
        table = ws.tables[name]
        try:
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        except (TypeError, ValueError):
            continue
        hrc = table.headerRowCount if table.headerRowCount is not None else 1
        trc = table.totalsRowCount if table.totalsRowCount is not None else 0
        first, last = min_row + hrc, max_row - trc
        if last < first:
            continue
        # at the first data row, or at the totals row (the new row goes above
        # it): either way the inserted rows become rows of the table's data
        if (hrc >= 1 and index == first) or (trc >= 1 and index == last + 1):
            out.append((min_row, first, last, max_row, min_col, max_col))
    return tuple(out)


class WorkbookPackage:
    """An open workbook, its hazard report, its pending edits, and the save
    pipeline. Construct with ``WorkbookPackage.open(path)``; mutate via the
    edit methods; persist with ``save()``. Hold one instance per mutation and
    let ``save()`` (which takes the write lock) serialize concurrent writers.
    """

    def __init__(self, path: str, *, com_manager: Any = None):
        self.path = os.fspath(path)
        self.com_manager = com_manager
        self._workbook = None            # lazy openpyxl load
        self._keep_vba = _xmlsafe.keeps_vba(self.path)
        self._hazard: _hazard.HazardReport | None = None
        self._pre_parts: list[str] = []
        self._pre_sizes: dict[str, int] = {}
        self._intended: dict[tuple[str, str], tuple[str, Any]] = {}
        #: (sheet, coord) -> text for each plain string a caller
        #: wrote that looks like a number; save() names them in a warning. A
        #: later write to the same cell takes its entry away.
        self._text_numbers: dict[tuple[str, str], str] = {}
        self._formula_written = False
        #: (sheet, coord) of formulas written this save that return arrays;
        #: followed through structural edits made after them.
        self._dynamic_written: set[tuple[str, str]] = set()
        self._structural: list[_refs.RefEdit] = []
        #: Sentences the next save adds to its warnings (set by the edits).
        self._notes: list[str] = []
        self._changed: dict[str, Any] = {}
        self._expected_removals: set[str] = set()
        self._expected_preserved: set[str] = set()
        #: (mtime_ns, size) of the file the model and the hazard scan were
        #: read from. save() re-stats under the write lock and refuses if
        #: another writer changed the file in between (see _check_unchanged).
        self._stamp: tuple[int, int] | None = None
        #: extensions openpyxl announced it was discarding at load time
        #: (x14 conditional formatting, sparkline and slicer-list extLst
        #: blocks); they live INSIDE surviving parts, so no part-level verify
        #: can see them. hazard.scan_path detects them up front by reading the
        #: worksheet extLst; this list is the independent second detector, and
        #: the gate refuses on either unless allow_loss covers it.
        self._dropped_extensions: list[str] = []
        #: hazard keys the caller was explicitly warned would be dropped when
        #: they passed allow_loss; the verify gate excuses these and nothing
        #: else (allow_loss is not a blanket amnesty for fragile parts).
        self._warned_loss_keys: set[str] = set()
        #: Worksheet parts whose ca="1" always-calculate flags and cached
        #: values the save put back (core.calc.restore_always_calc_cache).
        self._restored_calc_cache = 0
        #: Rectangles (sheet, min_row, min_col, max_row, max_col) whose cell
        #: VALUES an op rewrites wholesale without a per-cell intent record
        #: (copy/move destinations, sorted regions). The untouched-number
        #: preservation pass (core.calc.preserve_untouched_number_text) must
        #: never "restore" an original 17-digit double onto a cell such an op
        #: legitimately re-populated, so these regions are excluded from the
        #: repair alongside the intended map. See note_region_write().
        self._written_regions: list[tuple[str, int, int, int, int]] = []
        #: Cells whose original full-precision stored text the save restored
        #: (core.calc.preserve_untouched_number_text).
        self._preserved_number_texts = 0
        #: Every multi-cell array formula of
        #: every sheet, with the value and hyperlink of each of its cells, as
        #: the workbook was when its model loaded (id(ws) -> (ws, [(array,
        #: {coord: key})])), and the cells a writer recorded through set_cell
        #: or set_formula. save() compares every array cell against that
        #: snapshot, so a writer that assigns cell.value directly (a
        #: hyperlink's display text, import's neutralized text, any future
        #: one) cannot change part of an array either.
        self._array_snap: dict[int, tuple[Any, list]] = {}
        #: Each table's (ref, header cells, column names)
        #: as the workbook was when its model loaded, by (id(ws), name).
        self._table_snap: dict[tuple[int, str], tuple] = {}
        self._array_writes: dict[tuple[int, Any], set[str]] = {}
        #: The protection-relevant state at load, during a
        #: dispatched write (core.protection_gate); None when nothing in the
        #: workbook is protected.
        self._prot_snap = None
        #: The model as it loaded, for the byte-preserving save
        #: (core/surgical.py). None until the model loads.
        self._surgical_snap: _surgical.WorkbookSnap | None = None
        #: This save's route ("raw_ooxml" or "openpyxl"), the surgical build
        #: when the route is raw, and why the raw route was not taken.
        self._route: str | None = None
        self._route_build: _surgical.BuildResult | None = None
        self._route_note: str | None = None
        #: Broken internal links the ORIGINAL already had (verify fails only
        #: on new ones). Taken before the first verify of a save, from the
        #: file as it was; reset whenever a save lands.
        self._integrity_baseline: set[str] | None = None
        #: The cell-level difference this save writes (core/surgical.py
        #: cell_changes), what the rewrite route put back from the original
        #: (core/carry.py), and which formulas may now show stale results.
        self._route_changes: dict | None = None
        self._carry: _carry.CarryReport | None = None
        self._rewrite_stale: _surgical.Staleness | None = None
        self._route_structural = False
        #: Rewrites of formula TEXT that change no result (a sheet or table
        #: rename), in the order they ran. The save applies them to a
        #: formula as it loaded: when that gives the formula now stored, the
        #: formula is the same one and keeps its stored result (core.carry,
        #: core.calc). Reset whenever a save lands.
        self._text_rewrites: list[Callable[[str], str]] = []

    def note_text_rewrite(self, fn: Callable[[str], str]) -> None:
        """Register a rewrite of formula text that changes no result (a
        rename). Formulas it rewrote keep their stored results through the
        save; Excel's own rename keeps every value, and an iterative model
        seeds from them (without them it opened showing #VALUE!)."""
        self._text_rewrites.append(fn)

    def _same_formula(self, loaded: str, now: str) -> bool:
        """Whether formula text `now` is the formula `loaded` after this
        save's registered result-preserving rewrites (both with or without
        the leading '=')."""
        if loaded == now:
            return True
        if not self._text_rewrites:
            return False
        text = loaded
        for fn in self._text_rewrites:
            try:
                text = fn(text)
            except Exception:  # noqa: BLE001 - unreadable: not the same
                return False
        return text == now

    def _title_map(self) -> dict[str, str]:
        """Each sheet's title as it loaded -> its title now (a rename),
        for the sheets still in the workbook."""
        snap = self._surgical_snap
        wb = self._workbook
        if snap is None or wb is None:
            return {}
        live = {id(ws) for ws in wb._sheets}
        return {ss.title: ss.ws.title for key, ss in snap.sheets.items()
                if key in live and ss.title != ss.ws.title}

    # ------------------------------------------------------------- open

    @classmethod
    def open(cls, path: str, *, com_manager: Any = None) -> "WorkbookPackage":
        """Open a workbook for reading/mutation. Runs the hazard scan and
        captures the pre-write part list. Does not load the openpyxl model
        until a mutation or an explicit ``workbook`` access needs it."""
        p = check_path(path, "open workbook")
        # A workbook that a save which did not finish left aside
        # is named in the answer (reads never move it back).
        from . import swapwait as _swapwait
        _swapwait.require_workbook(p, f"no such workbook: {p}")
        if os.path.isdir(p):
            raise XlMcpError(
                f"{p} is a directory, not a workbook file; pass the path of "
                "an .xlsx/.xlsm file")
        # A legacy .xls / encrypted OLE container must refuse HERE, honestly
        # and in-envelope, before openpyxl's extension check leaks a raw
        # InvalidFileException or the hazard scan half-parses the embedded
        # theme fragment modern .xls files carry.
        _hazard.refuse_ole_container(p)
        pkg = cls(p, com_manager=com_manager)
        rep = _hazard.scan_path(p)
        if rep.error_kind == "missing":
            # Another save swapped it in at this very instant.
            _swapwait.require_workbook(p, f"no such workbook: {p}")
            rep = _hazard.scan_path(p)
            if rep.error_kind == "missing":
                raise _swapwait.workbook_not_found(p, f"no such workbook: {p}")
        if rep.error is not None:
            # A lock is not corruption: the wrong word sends a panicked
            # user toward a destructive "repair".
            if rep.error_kind in ("locked", "permissions"):
                # Neither is corruption. "permissions" is also not a lock:
                # the file is intact and this process simply may not read
                # it, which is a different sentence and a different fix.
                raise WorkbookLocked(f"{p}: {rep.error}")
            raise WorkbookCorrupt(f"{p}: {rep.error}")
        pkg._hazard = rep
        pkg._pre_parts = list(rep.parts)
        pkg._pre_sizes = dict(rep.sizes)
        pkg._stamp = _file_stamp(p)
        return pkg

    def _check_unchanged(self) -> None:
        """Refuse if the file changed since this package read it.

        The write lock serializes THIS server's writers, but Excel, another
        tool, or a second process can land a write between open() (which took
        the hazard scan and the part inventory) and save() (which writes the
        model loaded from the older bytes). Without this check that write is
        silently clobbered, and the pre-write inventory the verify gate
        compares against describes a file that no longer exists. Reproduced
        end to end: the concurrent cell simply vanished."""
        if self._stamp is None:
            return
        now = _file_stamp(self.path)
        if now is None or now == self._stamp:
            return
        raise _safesave.changed_on_disk_error(self.path)

    @property
    def hazard(self) -> _hazard.HazardReport:
        if self._hazard is None:
            self._hazard = _hazard.scan_path(self.path)
        return self._hazard

    @property
    def workbook(self):
        """The openpyxl workbook, loaded on first access. keep_vba is set for
        .xlsm and .xltm so the VBA project survives the round-trip."""
        if self._workbook is None:
            import zipfile as _zipfile
            import warnings as _warnings
            try:
                with _warnings.catch_warnings(record=True) as caught:
                    _warnings.simplefilter("always")
                    self._workbook = _xmlsafe.load_workbook(
                        self.path, keep_vba=self._keep_vba, data_only=False,
                        rich_text=False)
            except _zipfile.BadZipFile:
                # The promised conversion: a mutation-path load must refuse
                # as WorkbookCorrupt, never a raw BadZipFile, unless a
                # byte-range lock is the real cause.
                if _hazard.file_read_blocked(self.path):
                    raise WorkbookLocked(
                        f"{self.path}: {_hazard._LOCKED_ERROR}") from None
                raise WorkbookCorrupt(
                    f"{self.path} is not a valid .xlsx/.xlsm package (the "
                    "zip structure is damaged); the file was NOT modified"
                ) from None
            except (TypeError, ValueError, IndexError, SyntaxError) as exc:
                # openpyxl's reader raises a bare TypeError/ValueError on some
                # malformed-but-openable markup (a <row r="1.0"> that also
                # carries a height makes bind_row_dimensions do int("1.0")),
                # an IndexError for a cell pointing past the shared strings
                # and the XML parser's SyntaxError (ParseError) for a part
                # cut off mid-way. A mutation must refuse plainly, never let
                # a Python error escape as if the caller's arguments were
                # wrong (or as the server's own fault).
                raise WorkbookCorrupt(
                    f"{self.path} has markup this reader cannot load "
                    f"({type(exc).__name__}: {exc}); the file was NOT modified"
                ) from exc
            self._dropped_extensions = _extension_drops(caught)
            for ws in self._workbook.worksheets:
                self._snapshot_arrays(ws)
                for t in ws.tables.values():
                    self._table_snap[(id(ws), t.name)] = self._table_state(ws, t)
            # The byte-preserving save's snapshot first: it must record the
            # style table exactly as the file has it, before any other reader
            # of the model can touch it.
            self._surgical_snap = _surgical.snapshot(self._workbook)
            self._prot_snap = _protection_gate.snapshot(self._workbook)
        return self._workbook

    # ------------------------------------------------------------- address

    def resolve(self, location: Any, *,
                default_sheet: str | None = None) -> _locate.ResolvedGrid:
        """Resolve a location object to a normalized rectangle (core.locate)."""
        return _locate.resolve_location(
            self.workbook, location, default_sheet=default_sheet)

    def _ws(self, sheet: str | None):
        wb = self.workbook
        if sheet is None:
            return wb.active
        return wb[sheet]

    # ------------------------------------------------------------- mutate

    def set_cell(self, sheet: str | None, coord: str, value: Any, *,
                 from_caller: bool = False) -> None:
        """Write one cell. A string starting with '=' is treated as a formula
        (normalized through the _xlfn shim); anything else is a literal.

        A TYPED value object writes what a string cannot say:
        {"date": "2026-03-24"} (or "2026-03-24T09:30") is a
        real Excel date with a date number format; {"text": "=A1"} is text,
        whatever it starts with. Excel's own convention works too: a string
        that starts with an apostrophe and then =, +, - or @ is that text
        without the apostrophe, kept as text (quote prefix), as typing it in
        Excel does.

        from_caller marks a value the caller typed for this cell
        (set_cell, write_range, set_cells, apply_edits). A plain string among
        those that looks like a number is stored as the text it is, and the
        save says so; a value an op derives (import, replace, table
        headers) is not the caller's to be told about."""
        self._text_numbers.pop((self._ws(sheet).title, coord.upper()), None)
        if isinstance(value, dict):
            self._set_typed(sheet, coord, value)
            return
        if isinstance(value, str) and len(value) > 1 and value[0] == "'" \
                and value[1] in "=+-@":
            self._set_typed(sheet, coord, {"text": value[1:]})
            return
        if isinstance(value, str) and value.startswith("=") and value != "=":
            self.set_formula(sheet, coord, value)
            return
        if value == "=":
            # A bare "=" is not a formula; the file stores it as an inline
            # string either way, but the response used to REPORT it as
            # kind:"formula". Store and report it
            # honestly as text.
            ws = self._ws(sheet)
            self._note_array_write(ws, coord)
            cell = ws[coord]
            cell.value = value
            cell.data_type = "s"
            self._intended[(ws.title, coord.upper())] = ("value", value)
            return
        # Control characters, lone surrogates, and text past Excel's cell
        # ceiling: refuse in-envelope here rather than let openpyxl's
        # IllegalCharacterError escape the refusal shape (and rather than
        # accept a value the reply cannot encode, or one that only the
        # post-write verify would notice).
        _limits.check_cell_text(value, what=f"the value for {coord}")
        ws = self._ws(sheet)
        self._note_array_write(ws, coord)
        new_cell = _cell_is_new(ws, coord)
        ws[coord] = value
        self._intended[(ws.title, coord.upper())] = ("value", value)
        if from_caller and isinstance(value, str) and looks_like_number(value) \
                and not _formatted_as_text(ws, coord, new_cell):
            # A cell formatted as Text, by itself or by its row or column,
            # is the deliberate way to keep a numeral as text; only an
            # unformatted one is warned about.
            self._text_numbers[(ws.title, coord.upper())] = value

    def _set_typed(self, sheet: str | None, coord: str, value: dict) -> None:
        """A typed value object (set_cell's docstring)."""
        import datetime as _dt

        if len(value) != 1 or next(iter(value)) not in ("date", "text"):
            raise XlMcpError(
                f"a typed value for {coord} is one of {{\"date\": \"2026-03-24\"}} or "
                f"{{\"text\": \"...\"}}; got keys {sorted(value)}")
        kind, raw = next(iter(value.items()))
        ws = self._ws(sheet)
        self._note_array_write(ws, coord)
        cell = ws[coord]
        if kind == "text":
            if not isinstance(raw, str):
                raise XlMcpError(f"the text for {coord} must be a string")
            _limits.check_cell_text(raw, what=f"the value for {coord}")
            cell.value = raw
            cell.data_type = "s"
            if raw[:1] in ("=", "+", "-", "@"):
                # What Excel sets when such text is typed with an apostrophe:
                # editing the cell in Excel keeps it text.
                cell.quotePrefix = True
            self._intended[(ws.title, coord.upper())] = ("value", raw)
            return
        parsed = parse_iso_date(raw)
        if parsed is None:
            raise XlMcpError(
                f"the date for {coord} must be ISO text, \"2026-03-24\" or "
                f"\"2026-03-24T09:30\" (seconds optional); got {raw!r}")
        has_time = isinstance(parsed, _dt.datetime)
        if not has_time:
            # Stored as the serial number of midnight, as Excel stores a date;
            # openpyxl reads every date cell back as a datetime.
            parsed = _dt.datetime.combine(parsed, _dt.time())
        from openpyxl.styles.numbers import is_date_format
        prior_format = cell.number_format or "General"
        cell.value = parsed  # openpyxl sets its own datetime format here
        if is_date_format(prior_format):
            cell.number_format = prior_format  # the cell's own date format stays
        else:
            cell.number_format = ("yyyy-mm-dd" if not has_time else
                                  "yyyy-mm-dd hh:mm:ss" if parsed.second else
                                  "yyyy-mm-dd hh:mm")
        self._intended[(ws.title, coord.upper())] = ("value", parsed)

    @staticmethod
    def _array_cell_key(ws, row: int, col: int):
        """What a cell of an array formula holds, for the array comparison:
        its value and its hyperlink. Reads ws._cells so that looking never
        creates a cell; a cell that does not exist reads as an empty one
        (a reader that merely touches a cell changes nothing). The anchor's
        array formula counts by its ref only: rewriting the formula text of
        the WHOLE array (replace_cells keeps the array, core.arrays.retext)
        is not a write into part of it."""
        cell = ws._cells.get((row, col))
        v = cell.value if cell is not None else None
        hl = cell.hyperlink if cell is not None else None
        af = _arrays.array_formula_of_value(v)
        vk = ("array", str(af.ref)) if af is not None else (type(v).__name__, v)
        return (vk, (hl.target, hl.location) if hl is not None else None)

    def _snapshot_arrays(self, ws) -> None:
        from openpyxl.utils import get_column_letter

        arrays = []
        for a in _arrays.array_ranges(ws):
            if not a.multi_cell:
                continue
            keys = {f"{get_column_letter(c)}{r}": self._array_cell_key(ws, r, c)
                    for r in range(a.min_row, a.max_row + 1)
                    for c in range(a.min_col, a.max_col + 1)}
            arrays.append((a, keys))
        self._array_snap[id(ws)] = (ws, arrays)

    def _note_array_write(self, ws, coord: str) -> None:
        """Record a cell write that lands inside a multi-cell array formula.
        A write of the value a cell already held still counts as
        covering it, which the snapshot alone could not see."""
        entry = self._array_snap.get(id(ws))
        if not entry or not entry[1]:
            return
        key = coord.upper().replace("$", "")
        for a, keys in entry[1]:
            if key in keys:
                self._array_writes.setdefault((id(ws), a), set()).add(key)

    def _refuse_partial_array_writes(self, only_ws=None) -> None:
        """Before anything reaches the disk, refuse when the
        cells changed (against the snapshot) or written (recorded) inside a
        multi-cell array formula cover part of it and not all of it."""
        from openpyxl.utils.cell import coordinate_to_tuple

        if self._workbook is None:
            return
        live = {id(ws) for ws in self._workbook.worksheets}
        by_sheet: dict[str, list] = {}
        for wid, (ws, arrays) in self._array_snap.items():
            if wid not in live or (only_ws is not None and ws is not only_ws):
                continue
            for a, keys in arrays:
                touched = set(self._array_writes.get((wid, a), ()))
                for coord, key in keys.items():
                    if self._array_cell_key(ws, *coordinate_to_tuple(coord)) != key:
                        touched.add(coord)
                if touched and len(touched) < len(keys):
                    by_sheet.setdefault(ws.title, []).append((a, sorted(touched)))
        for sheet, partial in by_sheet.items():
            _arrays.refuse_partial_writes(sheet, partial)

    @staticmethod
    def _table_state(ws, t) -> tuple:
        from openpyxl.utils.cell import range_boundaries

        c0, r0, c1, _r1 = range_boundaries(t.ref)
        hrc = t.headerRowCount if t.headerRowCount is not None else 1
        headers = tuple(ws._cells[(r0, c)].value if (r0, c) in ws._cells else None
                        for c in range(c0, c1 + 1)) if hrc else ()
        names = tuple(tc.name for tc in (t.tableColumns or ()))
        return (t.ref, headers, names)

    def _settle_tables(self) -> None:
        """Before anything reaches the disk, refuse a save
        that would leave a table Excel will not open without a repair, and
        keep a renamed header and its column name together.

        Only what this edit changed is judged: a table whose header cells
        and column names are what they were at load passes (a workbook that
        arrived broken is not this edit's doing; a row insert that only
        moved the table changes nothing here), and within a table only the
        header cells that changed are checked, against the whole header
        row for duplicates. A changed header must be non-empty text, not a
        formula, and not repeat another header. New text in a header is a
        column rename, as when a header is typed over in Excel: the table's
        column name follows it, unless a formula or defined name refers to
        the old name in brackets, which would then point at nothing. No
        multi-cell array formula that was not there at load may sit inside
        a table (Excel does not allow one there)."""
        from openpyxl.utils import get_column_letter
        from openpyxl.utils.cell import range_boundaries

        if self._workbook is None:
            return
        for ws in self._workbook.worksheets:
            tables = list(ws.tables.values())
            if not tables:
                continue
            old_arrays = {a.ref for a, _k in self._array_snap.get(id(ws), (None, []))[1]}
            new_arrays = [a for a in _arrays.array_ranges(ws)
                          if a.multi_cell and a.ref not in old_arrays]
            for t in tables:
                c0, r0, c1, r1 = range_boundaries(t.ref)
                for a in new_arrays:
                    if not (a.max_row < r0 or a.min_row > r1 or a.max_col < c0
                            or a.min_col > c1):
                        _refuse_table(
                            f"a multi-cell array formula over {a.ref} would sit inside table "
                            f"{t.displayName} ({t.ref}) on {ws.title!r}. Excel does not allow an "
                            "array formula that spans cells in a table and will not open the "
                            "workbook without a repair. Write values into the table instead "
                            "(copy_range with what='values'), or put the array outside the "
                            "table.")
                _ref, headers, names = self._table_state(ws, t)
                old = self._table_snap.get((id(ws), t.name))
                if not headers or (old is not None and old[1:] == (headers, names)):
                    continue
                old_headers = old[1] if old is not None else ()
                aligned = len(old_headers) == len(headers)
                low = [h.strip().lower() if isinstance(h, str) else None for h in headers]
                for i, v in enumerate(headers):
                    if aligned and old_headers[i] == v and (
                            not names or i >= len(names) or names[i] == v):
                        continue
                    cell = f"{get_column_letter(c0 + i)}{r0}"
                    column = names[i] if i < len(names) else None
                    problem = None
                    if not isinstance(v, str):
                        problem = "it would hold a number or nothing"
                    elif not v.strip():
                        problem = "it would be empty"
                    elif v.startswith("="):
                        problem = "it would hold a formula"
                    elif low.count(low[i]) > 1:
                        j = next(k for k, x in enumerate(low) if x == low[i] and k != i)
                        problem = (f"it would repeat the header in "
                                   f"{get_column_letter(c0 + j)}{r0}")
                    elif column is not None and v != column:
                        users = _names_column(self._workbook, column, t.displayName)
                        if users:
                            problem = (f"renaming column {column!r} to {v!r} would leave "
                                       f"{', '.join(users[:5])} pointing at a column that no "
                                       "longer exists")
                        else:
                            t.tableColumns[i].name = v  # the rename Excel itself makes
                            continue
                    if problem:
                        _refuse_table(
                            f"cannot change {cell} on {ws.title!r}: it is a header cell of table "
                            f"{t.displayName}, and {problem}. A table header must stay unique, "
                            "non-empty text, or Excel will not open the workbook without a "
                            "repair. To rename the column, write its new text into the header "
                            "cell (once no formula names the old one); data goes in the rows "
                            "below the header.",
                            {"table": t.displayName, "cell": cell, "column": column})

    def set_formula(self, sheet: str | None, coord: str, formula: str) -> None:
        """Write a formula, normalized so modern functions do not land as
        #NAME? and the workbook recalculates on open."""
        ws = self._ws(sheet)
        if not formula.strip().lstrip("=").strip():
            raise XlMcpError(
                f"the formula for {coord} is empty (only an equals sign), so "
                "nothing was written; pass the formula itself")
        self._text_numbers.pop((ws.title, coord.upper()), None)
        self._note_array_write(ws, coord)
        _limits.check_text_storable(formula, what=f"the formula for {coord}")
        if "[" in formula:
            # Excel's typing shorthand for a structured
            # reference (T[@v], [v] inside a table) is stored in the file
            # form Excel itself writes; what Excel refuses is refused here.
            from openpyxl.utils.cell import coordinate_to_tuple

            from . import structrefs as _sr

            r, c = coordinate_to_tuple(coord.replace("$", ""))
            stored = _sr.to_file_form(formula, _sr.tables_of(self.workbook), ws.title, r, c)
            if stored != formula:
                seen = self._changed.setdefault("structured_references_stored_as", [])
                if len(seen) < 20:  # the first few say what happened; a fill repeats it
                    seen.append({"cell": f"{ws.title}!{coord.upper()}", "given": formula,
                                 "stored": stored})
                formula = stored
        _limits.check_formula(formula, what=f"the formula for {coord}")
        normalized, _prefixed = _calc.normalize_formula(formula)
        ws[coord] = normalized
        from . import dynarray as _da
        if _da.is_spilling(normalized):
            self._dynamic_written.add((ws.title, coord.upper().replace("$", "")))
        self._intended[(ws.title, coord.upper())] = ("formula", normalized)
        self._formula_written = True

    def modify_structure(self, edit: _refs.RefEdit) -> _refs.RewriteReport:
        """Insert or delete rows/columns and rewrite every reference so the
        workbook stays coherent (core.refs). openpyxl shifts the cell grid;
        core.refs repairs formulas, names, conditional formats, data
        validations, table refs, and merged ranges. The change is verified on
        save like any other mutation."""
        wb = self.workbook
        ws = wb[edit.sheet]
        # The snapshot is by address, and the edit moves addresses, so
        # the sheet is judged now and snapshotted again after the edit.
        self._refuse_partial_array_writes(only_ws=ws)
        # The same for a protected sheet, and the edit itself is judged.
        _protection_gate.before_structure_edit(self, ws, edit.kind, edit.index, edit.count)
        tables_before = {name: ws.tables[name].ref for name in list(ws.tables)} \
            if edit.kind == _refs.INSERT_ROWS else {}
        if edit.kind == _refs.INSERT_ROWS and not edit.grow:
            # Rows inserted at a table's first data row grow the table; what
            # covers exactly its data column is extended, as Excel does.
            grow = _table_growth(ws, edit.index)
            if grow:
                edit = _dc_replace(edit, grow=grow)
        if edit.kind == _refs.INSERT_ROWS:
            ws.insert_rows(edit.index, edit.count)
        elif edit.kind == _refs.DELETE_ROWS:
            ws.delete_rows(edit.index, edit.count)
        elif edit.kind == _refs.INSERT_COLS:
            ws.insert_cols(edit.index, edit.count)
        elif edit.kind == _refs.DELETE_COLS:
            ws.delete_cols(edit.index, edit.count)
        elif edit.kind == _refs.MOVE:
            raise ValueError(
                "modify_structure handles insert/delete; MOVE is a copy_range "
                "concern that rewrites refs directly via core.refs")
        if edit.kind in (_refs.INSERT_ROWS, _refs.DELETE_ROWS):
            _shift_row_dimensions(ws, edit)
        else:
            _shift_column_dimensions(ws, edit)
        report = _refs.rewrite_workbook(wb, edit)
        self._notes.extend(n for n in report.not_extended if n not in self._notes)
        if tables_before:
            filled = _fill_calculated_columns(ws, edit, tables_before)
            if filled:
                self._formula_written = True
                self._changed.setdefault("calculated_column_cells", []).extend(
                    filled)
        self._structural.append(edit)
        # The arrays moved with the grid (the edit itself is guarded by
        # refuse_if_split): snapshot them again at their new addresses.
        self._array_writes = {k: v for k, v in self._array_writes.items()
                              if k[0] != id(ws)}
        self._snapshot_arrays(ws)
        _protection_gate.after_structure_edit(self, ws)
        if report.formulas:
            self._formula_written = True
        # Structural edits move cells, so the tracked intents no longer point
        # at the right coordinates; content read-back is by report, not cells.
        self._intended.clear()
        moved = set()
        for sh, coord in self._dynamic_written:
            if sh != edit.sheet:
                moved.add((sh, coord))
                continue
            new = _refs.transpose_formula("=" + coord, edit, sh)[1:]
            if _refs.REF_ERROR not in new:
                moved.add((sh, new.replace("$", "")))
        self._dynamic_written = moved
        self._changed.setdefault("structural", []).append(
            {"kind": edit.kind, "sheet": edit.sheet, "index": edit.index,
             "count": edit.count, "rewrites": report.as_dict()})
        return report

    def note_region_write(self, sheet: str, min_row: int, min_col: int,
                          max_row: int, max_col: int) -> None:
        """Register a rectangle whose cell values this op rewrites WITHOUT a
        per-cell entry in the intended map (a copy/move destination, a sorted
        region). Untouched-number preservation skips every cell inside it, so
        a full-precision original can never be restored over a value the op
        deliberately placed there. Scoped to the next successful save."""
        self._written_regions.append(
            (sheet, min_row, min_col, max_row, max_col))

    def _number_preservation_maps(self):
        """(remap, skip) for core.calc.preserve_untouched_number_text.

        remap follows an original address through this save's recorded
        structural edits (insert/delete rows/cols shift, a deleted cell or a
        MOVE-covered cell returns None: untrackable is untouchable). skip is
        True for every cell the op wrote: the intended map plus every
        registered written region."""
        from openpyxl.utils import get_column_letter
        from openpyxl.utils.cell import coordinate_from_string
        from openpyxl.utils.cell import column_index_from_string

        edits = list(self._structural)
        intended = {(s, c.upper()) for (s, c) in self._intended}
        regions = list(self._written_regions)

        def remap(sheet: str, addr: str) -> str | None:
            try:
                col_s, row = coordinate_from_string(addr)
                col = column_index_from_string(col_s)
            except Exception:  # noqa: BLE001
                return None
            for e in edits:
                if e.sheet != sheet:
                    continue
                if e.kind == _refs.INSERT_ROWS:
                    if row >= e.index:
                        row += e.count
                elif e.kind == _refs.DELETE_ROWS:
                    if e.index <= row < e.index + e.count:
                        return None
                    if row >= e.index + e.count:
                        row -= e.count
                elif e.kind == _refs.INSERT_COLS:
                    if col >= e.index:
                        col += e.count
                elif e.kind == _refs.DELETE_COLS:
                    if e.index <= col < e.index + e.count:
                        return None
                    if col >= e.index + e.count:
                        col -= e.count
                else:
                    # A MOVE (or any future kind) relocates content this map
                    # does not model; conservative means hands off the sheet.
                    return None
            return f"{get_column_letter(col)}{row}"

        def skip(sheet: str, addr: str) -> bool:
            if (sheet, addr.upper()) in intended:
                return True
            try:
                col_s, row = coordinate_from_string(addr)
                col = column_index_from_string(col_s)
            except Exception:  # noqa: BLE001
                return True
            for (s, r0, c0, r1, c1) in regions:
                if s == sheet and r0 <= row <= r1 and c0 <= col <= c1:
                    return True
            return False

        return remap, skip

    def expect_removal(self, *prefixes: str) -> None:
        """Register part-name prefixes an op DELIBERATELY removes at the model
        level (sheet delete, table to_range, comment delete), so the default-
        fail inventory check in verify-after-write does not flag the removal
        as silent loss. Scoped to the next successful save only; cleared with
        the other pending-edit state. This is the narrow, per-operation escape
        hatch; nothing else excuses a non-fragile part from the inventory
        diff."""
        self._expected_removals.update(p.lower() for p in prefixes)

    def expect_preserved(self, *hazard_keys: str) -> None:
        """Register hazard KEYS (core.hazard spec keys, e.g. 'media') whose
        SEV_DROPS gate refusal this op downgrades to a warning, because the
        op's own save path has been OBSERVED to preserve those parts and the
        default-fail part checks in verify-after-write still refuse the save
        if they are in fact lost. The sanctioned user: ops/objects.py
        manage_image, whose Pillow-backed load round-trips picture drawings
        and xl/media/ intact (observed on this machine for Excel-authored
        and openpyxl-authored images alike; the op verifies the picture-only
        precondition before registering). This NEVER weakens verification:
        part_loss_check runs with allow_loss=False for these keys, so an
        actual loss still fails the write and the original stays untouched.
        Scoped to the next successful save; cleared with pending state."""
        self._expected_preserved.update(k.lower() for k in hazard_keys)

    # ------------------------------------------------------------- save

    def _lock_state(self) -> tuple[bool, bool]:
        """(lockfile_present, writable). The ~$ owner lockfile is a SIGNAL,
        not authority: a stale one survives an Excel crash (verified against
        a real Excel), so the write-probe
        decides. A real Excel hold denies the r+b open; a stale lockfile
        leaves the file writable."""
        present, writable, _why = self._lock_diagnosis()
        return present, writable

    def _lock_diagnosis(self) -> tuple[bool, bool, str | None]:
        """(lockfile_present, writable, why) where why names the CAUSE of an
        unwritable target: 'read_only_attribute' or 'held'.

        The probe alone cannot tell them apart, and the difference is the
        whole message. A file carrying the read-only attribute (a copy off
        old media, a OneDrive restore, anything that came off a CD or a
        locked share) used to refuse with "is open in Excel", which sends
        the user to close a program they never opened. The attribute is
        checked first because it is a property of the file rather than a
        race: st_mode's owner-write bit is exactly the read-only attribute
        on Windows and the ordinary permission bit elsewhere."""
        from . import access as _access
        p = Path(self.path)
        present = _access.owner_file(p).exists()
        cause = _access.write_block(p)
        if cause is None:
            return present, True, None
        if cause == _access.IN_USE:
            return present, False, "held"
        return present, False, cause

    def _extension_warnings(self) -> list[str]:
        """The announced-loss line for in-part extension blocks openpyxl
        discarded at load. Reported only on the allow_loss path, because
        without allow_loss the gate now REFUSES instead (see _hazard_gate).

        These blocks live INSIDE xl/worksheets/sheetN.xml, which survives the
        save at full size, so no part-level check can catch them: x14
        conditional formatting (every modern data bar and icon set),
        sparkline groups, slicer lists. The hazard scan reads the worksheet
        extLst and flags them up front; openpyxl's own load-time
        announcement, captured here, is the independent second detector and
        the one that names the exact extensions it threw away."""
        if not self._dropped_extensions:
            return []
        return ["openpyxl discarded in-part extension blocks on load, so "
                "they are NOT in the saved file: "
                + "; ".join(self._dropped_extensions)
                + ". These live inside the worksheet part (x14 conditional "
                  "formatting, sparklines, slicer lists), so no part-level "
                  "check can catch them. You asked for this with allow_loss, "
                  "and the pre-edit backup is the only copy that still has "
                  "them."]

    def _refuse_extension_loss(self) -> None:
        """Backstop refusal for extension blocks openpyxl announced dropping
        that the hazard scan did not flag.

        The scan is the primary detector and normally fires first, so this
        path is reached only when the two disagree: a worksheet extLst the
        scan could not read, or an extension openpyxl drops from somewhere
        the scan does not look. Drop-class content must not go through on a
        detector disagreement, so the announcement alone refuses."""
        exc = HazardRefused(
            "openpyxl discarded in-part extension blocks when it loaded this "
            "workbook, so a file-based save would write it back without them: "
            + "; ".join(self._dropped_extensions)
            + ". These live inside the worksheet part (x14 conditional "
              "formatting, sparklines, slicer lists), so no part-level check "
              "can catch the loss after the fact. Refusing the mutation "
              "rather than destroy them. Two routes, and they are the only "
              "two: pass allow_loss:true to accept the loss (the workbook is "
              "backed up first, and the loss is permanent in the saved "
              "file), or leave this workbook alone at the file tier. xl_app "
              "is NOT a route for this: none of its ops writes a "
              "cell, a format, or a row. Reads never touch the workbook, and "
              "xl_workbook copy_workbook branches it byte-for-byte.")
        exc.detail = {"dropped_extensions": list(self._dropped_extensions),
                      "routes": _hazard.refusal_routes()}
        raise exc

    def _hazard_gate(self, allow_loss: bool) -> list[str]:
        """Apply the routing decision. Returns advisory warnings; raises
        HazardRefused when a drop-risk mutation has no loss-safe path.

        In-part extension blocks are drop-class under the package
        policy (drop-risk refuses without allow_loss, degrade-risk warns and
        proceeds), so they raise here like any other SEV_DROPS hazard. They
        used to warn and proceed, which announced the loss but still let a
        data bar disappear from a workbook the user never agreed to sacrifice.
        With allow_loss the announcement survives as a warning."""
        rep = self.hazard
        warnings: list[str] = []
        # Families the rewrite puts back from the original (core/carry.py)
        # are not at risk on this route; verify-after-write still refuses
        # if one of them does go missing.
        carried = self._carried_keys()
        hazards = [h for h in rep.hazards if h.key not in carried]
        if rep.error is None and not hazards:
            if self._dropped_extensions:
                # The scan saw nothing but openpyxl still threw extensions
                # away; the backstop governs.
                if not allow_loss:
                    self._refuse_extension_loss()
                warnings.extend(self._extension_warnings())
            return warnings
        drop_keys = [h.key for h in hazards
                     if h.severity == _hazard.SEV_DROPS]
        # An Excel-made chart carries style and colour parts the rewrite does
        # not write back; verify-after-write then refused with a bare
        # VALIDATION_FAILED and no route. Name the loss here instead, where
        # the refusal can say what goes and how to accept it.
        if any(h.key == "charts" and any(
                p.rsplit("/", 1)[-1].lower().startswith(("style", "colors"))
                for p in h.parts) for h in hazards):
            drop_keys.append("charts")
        # The generic degrade sentence says a family is "re-serialized
        # through openpyxl's model", which is the truth about charts and
        # pivots and a lie about document properties: those are not
        # re-serialized at all, they are rewritten from an empty template
        # and the authored fields go. A hazard whose note says something
        # different gets its own line rather than being folded into a
        # sentence that misdescribes it.
        self_describing = {_hazard.DOC_PROPERTIES_KEY,
                           _hazard.BROKEN_SHARED_KEY}
        degradable = [h.label for h in hazards
                      if h.severity == _hazard.SEV_DEGRADES
                      and h.key not in self_describing]
        degrade_notes = [f"{h.label}: {h.note}" for h in hazards
                         if h.severity == _hazard.SEV_DEGRADES
                         and h.key in self_describing]
        degrade_keys = [h.key for h in hazards
                        if h.severity == _hazard.SEV_DEGRADES]
        conditional = [h for h in hazards
                       if h.severity == _hazard.SEV_CONDITIONAL]

        # A conditional hazard (VBA) survives ONLY because the openpyxl load
        # passes keep_vba, and keep_vba is set from the .xlsm / .xltm extension. A
        # vbaProject.bin inside a plain .xlsx (renamed or mis-authored file)
        # gets NO keep_vba, so the save drops it: treat it as a hard drop,
        # never claim it is preserved.
        if conditional and not self._keep_vba:
            drop_keys.extend(h.key for h in conditional)
            conditional = []

        # Keys the op registered via expect_preserved: the gate downgrades
        # them to a warning because the op's save path preserves them and
        # verify-after-write still hard-fails if they are actually lost.
        preserved = [k for k in drop_keys if k in self._expected_preserved]
        if preserved:
            drop_keys = [k for k in drop_keys if k not in preserved]
            labels = [next(h.label for h in rep.hazards if h.key == k)
                      for k in preserved]
            warnings.append(
                ", ".join(labels) + " are expected to survive this "
                "operation's save path (observed round-trip); "
                "verify-after-write refuses the save if they are lost")

        if drop_keys and not allow_loss:
            labels = [next(h.label for h in rep.hazards if h.key == k)
                      for k in drop_keys]
            parts = [p for h in rep.hazards if h.key in drop_keys
                     for p in h.parts]
            costs = _hazard.loss_costs(drop_keys)
            exc = HazardRefused(
                "This edit would drop content the file writer cannot "
                "preserve yet: " + "; ".join(costs)
                + ". Two honest routes: pass allow_loss=true to proceed "
                "while accepting exactly that loss, or make this edit in "
                "Excel itself, which preserves everything. A future release "
                "adds preservation for this content.")
            exc.detail = {"parts": parts, "labels": labels,
                          "losses": costs,
                          "routes": _hazard.refusal_routes()}
            raise exc
        if allow_loss and (drop_keys or degrade_keys):
            # Record exactly which hazard families the caller was told they
            # were sacrificing; verify-after-write excuses those and only
            # those (see verify.part_loss_check loss_keys). DEGRADE-warned
            # families are included: "may degrade" concretely means chart
            # styling sub-parts (colors1.xml, style1.xml) and anchors can
            # vanish, and scoping the amnesty to DROP families alone made
            # the refusal's advertised remedy a dead end on any file that
            # held both classes, while a file with ONLY degrade families
            # fell through to the blanket-None amnesty: same loss, opposite
            # outcomes depending on unrelated content. One consistent rule
            # now: allow_loss excuses exactly the families the gate warned
            # about, of either class.
            self._warned_loss_keys = set(drop_keys) | set(degrade_keys)
        if drop_keys and allow_loss:
            labels = [next(h.label for h in rep.hazards if h.key == k)
                      for k in drop_keys]
            warnings.append(
                "allow_loss: proceeding despite " + ", ".join(labels)
                + " which the file-based save drops; the backup is the safety "
                "net")
            if "query_tables" in drop_keys:
                warnings.append(
                    "legacy query tables: the table keeps its cells, name, "
                    "range and style and becomes an ordinary table; its "
                    "query definition and data connection are gone, so it "
                    "no longer refreshes from its source")
        # The in-part extension backstop, same rule as the clean path. Without
        # allow_loss, reaching this line means drop_keys was empty (a
        # populated one already raised above), so an announcement the scan
        # missed is the only thing standing between the caller and a lost data
        # bar: refuse on it. With allow_loss it becomes the itemized warning.
        if self._dropped_extensions:
            if not allow_loss:
                self._refuse_extension_loss()
            warnings.extend(self._extension_warnings())
        warnings.extend(degrade_notes)
        if degradable:
            warnings.append(
                ", ".join(degradable) + " are re-serialized through openpyxl's "
                "model; sub-features the model does not know may degrade"
                + (" (with allow_loss, sub-parts these families lose on the "
                   "round-trip, such as Excel-authored chart styling parts, "
                   "count as accepted losses)" if allow_loss else ""))
        if conditional:
            warnings.append(
                "VBA project preserved via keep_vba; note that form and "
                "ActiveX controls still break on a file-based edit")
        return warnings

    def _carried_keys(self) -> set[str]:
        """Hazard families the rewrite route puts back from the original
        (core/carry.py): customXml always; document properties when the
        sheet list is unchanged (app.xml lists the sheets)."""
        keys = set(_carry.CARRIED_KEYS)
        wb = self._workbook
        snap = self._surgical_snap
        if wb is not None and snap is not None and \
                _surgical.wb_fingerprint(wb)[0] == snap.wb_fp[0]:
            keys.add(_hazard.DOC_PROPERTIES_KEY)
        return keys

    def _route_gate(self, allow_loss: bool) -> list[str]:
        """Choose this save's route, then apply the gate that route needs.

        A cell-level edit (core/surgical.py) is written by splicing the
        changed cells into the original package; every other part is copied
        through as it was, so nothing can be lost and the hazard gate has
        nothing to judge. Anything else takes the full rewrite, and the
        hazard gate decides it exactly as before (with the reason the
        byte-preserving route was not available added to any refusal)."""
        self._route, self._route_build, self._route_note = None, None, None
        changes = None
        if self._workbook is not None and self._surgical_snap is not None:
            # Excel's formula limits on every formula this edit wrote,
            # whichever op and whichever route; nothing has touched the disk.
            changes = _surgical.cell_changes(self._surgical_snap, self._workbook)
            _surgical.check_changed_formulas(changes)
        self._route_changes = changes
        self._route_structural = bool(self._structural)
        if changes is not None and self._structural:
            # Rows or columns moved: every address below or right of the
            # edit changed, and so did row heights, extension-block ranges
            # and everything else keyed by position. A cell diff by address
            # cannot express that; only the rewrite (with core/refs.py) can.
            self._route_note = "this edit inserts or deletes rows or columns"
        elif changes is not None and surgical_enabled():
            try:
                plan = _surgical.plan_edit(self._surgical_snap, self._workbook,
                                           changes)
                build = _surgical.build(self._surgical_snap, self._workbook,
                                        plan, self.path)
            except _surgical.NotSurgical as exc:
                self._route_note = str(exc)
            else:
                self._route = _surgical.ROUTE_RAW
                self._route_build = build
                return []
        self._route = _surgical.ROUTE_REWRITE
        try:
            return self._hazard_gate(allow_loss)
        except HazardRefused as exc:
            if self._route_note:
                exc.args = (_rewrite_needed(self._route_note) + " "
                            + str(exc.args[0] if exc.args else exc),)
                detail = dict(getattr(exc, "detail", None) or {})
                detail["needs_full_rewrite"] = self._route_note
                exc.detail = detail
            raise

    def _precision_warnings(self) -> list[str]:
        """Numbers the file stores exactly but Excel reads as a
        different double (a whole number past 2^53, a subnormal) are written
        as asked and named in a warning, with the way to keep every digit.

        From 2^53 itself, not above it. A client rounds a
        16-digit number to a double before the call is made, so 9007199254740993
        arrives as exactly 9007199254740992, and a warning that starts above
        that value never fires for the very case it exists for. A float of
        that size is named too (an integer-valued 1e16 or 2^53 as a float is
        the same case), so the rule is the one stated: |x| >= 2^53."""
        big: list[str] = []
        tiny: list[str] = []
        for (sheet, coord), (kind, v) in self._intended.items():
            if kind != "value" or isinstance(v, bool):
                continue
            if isinstance(v, int) and abs(v) >= EXACT_NUMBER_LIMIT:
                big.append(f"{sheet}!{coord}")
            elif isinstance(v, float):
                if math.isfinite(v) and abs(v) >= EXACT_NUMBER_LIMIT:
                    big.append(f"{sheet}!{coord}")
                elif v != 0.0 and abs(v) < 2.2250738585072014e-308:
                    tiny.append(f"{sheet}!{coord}")
        out = []
        if big:
            out.append(
                f"{', '.join(big[:5])}{' and more' if len(big) > 5 else ''}: the "
                "exact digits can't be stored as a number. A number of "
                "9,007,199,254,740,992 (2^53) or more is rounded to the nearest "
                "value Excel can hold (9007199254740993 becomes "
                "9007199254740992), and a client often rounds it before this "
                "server sees it. To keep every digit, pass the value as a "
                "string and set the cell's number format to Text (xl_format "
                "format_cells with number_format \"@\").")
        if tiny:
            out.append(
                f"{', '.join(tiny[:5])}{' and more' if len(tiny) > 5 else ''}: "
                "a number this close to zero (below about 2.2E-308) is stored "
                "as written, but Excel reads it as 0.")
        return out

    def _text_number_warnings(self) -> list[str]:
        """Plain strings a caller wrote that look like numbers
        are stored as text, as asked. SUM and AVERAGE skip text, and a model
        that sent "250" meaning 250 would otherwise report a total that went
        down without knowing why. One warning for the
        whole call; the cells whose last write is still the string."""
        if not self._text_numbers:
            return []
        hits = [(s, c, t) for (s, c), t in self._text_numbers.items()
                if self._intended.get((s, c)) == ("value", t)]
        if not hits:
            return []
        wb = self._workbook
        near = wb is not None and any(
            _numbers_beside(wb[s], c) for s, c, _t in hits if s in wb.sheetnames)
        refs = [f"{s}!{c}" for s, c, _t in hits]
        text = hits[0][2]
        # The number as it is passed, without thousands separators
        # ("Pass 1250 as a number", not "Pass 1,250": a model copies it).
        number = text.strip().replace(",", "")
        if len(hits) == 1:
            lead = f'{refs[0]} now holds the TEXT "{text}", not the number {number}'
            if near:
                lead += ", and the cells next to it in the column hold numbers"
            skipped = "it"
            fix = f"Pass {number} as a number, without quotes, to store a number."
        else:
            shown = (", ".join(refs[:-1]) + " and " + refs[-1] if len(refs) <= 5
                     else ", ".join(refs[:5]) + f" and {len(refs) - 5} more")
            lead = (f'{shown} now hold TEXT that looks like numbers '
                    f'({refs[0]} is "{text}")')
            if near:
                lead += ", and cells next to them in the column hold numbers"
            skipped = "them"
            fix = (f'Pass numbers without quotes ({number}, not "{text}") to '
                   "store numbers; keep text for codes and IDs.")
        tail = (f"SUM and AVERAGE skip text, so a total over the column leaves "
                f"{skipped} out." if near else "SUM and AVERAGE skip text.")
        return [f"{lead}. {tail} {fix}"]

    def _reservation_warnings(self, before: dict | None) -> list[str]:
        """This file route edits a workbook that has a password
        to modify without asking for it, and keeps the password. That is how
        the file tools have always treated it (only the routes that open
        Excel refuse); what was missing is saying so in the answer, so a
        model that reports "done" has the fact in front of it. `before` is
        the file's reservation read before the write; the file is read again
        so that "the password is kept" is said only when it is."""
        if before is None or before.get("password_to_modify") is not True:
            return []
        after = _hazard.write_reservation(self.path)
        kept = after is not None and after.get("password_to_modify") is True
        return [_hazard.write_reserved_notice(self.path, kept=kept)]

    def _dynamic_cells(self) -> dict[str, set[str]]:
        """The formula cells this save writes that call an array-returning
        function, by sheet: they are stored as dynamic arrays."""
        from . import dynarray as _da
        out: dict[str, set[str]] = {}
        wb = self._workbook
        for sheet, coord in self._dynamic_written:
            # The cell must still hold the spilling formula at save time
            # (a later write in the same call may have replaced it).
            if wb is None or sheet not in wb.sheetnames:
                continue
            text = _calc.formula_text_of(wb[sheet][coord])
            if text is not None and _da.is_spilling(text):
                out.setdefault(sheet, set()).add(coord)
        return out

    def _mark_dynamic(self, tmp: str, cells: dict[str, set[str]]) -> None:
        """Without this, a formula such as =UNIQUE(A1:A5) was
        stored as a legacy formula, which Excel opens as =@UNIQUE(A1:A5)
        showing one value and never spilling."""
        if not cells:
            return
        from . import dynarray as _da
        try:
            n = _da.mark_dynamic(tmp, cells)
        except _da.NotDynamic as exc:
            raise UnsupportedStructure(
                f"the formula returns an array, which Excel stores as a "
                f"dynamic array, and this workbook could not be marked for "
                f"it ({exc}); nothing was changed. Write it in Excel, or "
                f"wrap it in a function that returns one value "
                f"(for example INDEX(...,1)).") from exc
        if n:
            self._changed["dynamic_array_formulas"] = n

    def _default_saver(self, tmp: str) -> None:
        dynamic = self._dynamic_cells()
        if self._route_build is not None:
            _surgical.write_package(self.path, tmp, self._route_build)
            self._mark_dynamic(tmp, dynamic)
            return
        wb = self.workbook
        self._rewrite_stale = self._rewrite_staleness(wb)
        stale = self._rewrite_stale
        if self._formula_written or self._structural or (
                stale is not None and (stale.cells or not stale.complete)):
            try:
                wb.calculation.fullCalcOnLoad = True
            except Exception:
                pass
        self._detach_query_tables(wb)
        self._sync_hyperlink_refs(wb)
        wb.save(tmp)
        # openpyxl leaves an EMPTY <v></v> on every formula cell it writes.
        # Excel reads it as an error value, which a normal recalculation
        # overwrites but an ITERATIVE one seeds from, leaving a circular
        # formula stuck at #VALUE! (numbers-safety gate). Strip it before the
        # verify gate sees the package.
        try:
            _calc.strip_empty_cached_values(tmp)
        except Exception:  # noqa: BLE001
            pass  # cosmetic-plus: never fail a save over this
        # The ADJACENT case: openpyxl also drops the always-calculate flag
        # and any REAL cached value from every formula cell. Harmless under
        # normal calculation (fullCalcOnLoad fills it back in), fatal under
        # ITERATIVE calculation, where the iteration seeds from the current
        # value and a seedless circular formula sticks on #VALUE! forever.
        # Untouched formulas get their ca="1" and their <v> back; an edited
        # one never does.
        titles = self._title_map()
        try:
            self._restored_calc_cache = _calc.restore_always_calc_cache(
                self.path, tmp, titles=titles, same=self._same_formula)
        except Exception:  # noqa: BLE001
            self._restored_calc_cache = 0
        # THE 1-ULP REPAIR: openpyxl's %.16g float
        # writer silently rewrites every 17-significant-digit double in the
        # file -- including cells this edit never touched -- 1 ulp off, which
        # can flip equality-gated logic in an Excel-authored workbook. The
        # original numeric TEXT of every untouched full-precision cell is
        # restored by raw surgery; only cells the operation actually wrote
        # (intended map + registered written regions, addresses followed
        # through the structural edits) keep openpyxl's re-serialization.
        try:
            remap, skip = self._number_preservation_maps()
            self._preserved_number_texts = \
                _calc.preserve_untouched_number_text(
                    self.path, tmp, remap=remap, skip=skip, titles=titles)
        except Exception:  # noqa: BLE001
            self._preserved_number_texts = 0
        # What the model cannot hold, put back from the original: custom
        # properties, customXml, app.xml, external links, printer settings,
        # every untouched formula's cached result, automatic row heights
        # (core/carry.py).
        try:
            remap, skip = self._number_preservation_maps()
            self._carry = _carry.carry_forward(
                self.path, tmp, remap=remap,
                keep=self._cache_keeper(skip), titles=titles,
                **self._carry_flags(wb))
        except Exception:  # noqa: BLE001
            self._carry = None
        self._mark_dynamic(tmp, dynamic)

    @staticmethod
    def _sync_hyperlink_refs(wb) -> int:
        """A hyperlink is written under the address it was given when it was
        attached, and a row or column insert or delete moves the cell, not
        that address: the link stayed behind on a cell that no longer held
        its text, and the cell that moved lost it. Each link goes back to
        the address of the cell that carries it. Returns how many moved."""
        moved = 0
        for ws in wb.worksheets:
            for cell in ws._cells.values():
                link = getattr(cell, "_hyperlink", None)
                if link is not None and link.ref != cell.coordinate:
                    link.ref = cell.coordinate
                    moved += 1
        return moved

    @staticmethod
    def _detach_query_tables(wb) -> list[str]:
        """The rewrite does not write a legacy query table's own part (or its
        data connection) back, and a table that still says it is a query
        table, with its query column ids, points at a part that is gone:
        Excel then refuses the whole file. Such a table becomes an ordinary
        table here (its cells, name, range and style stay); the gate has
        already named the loss. Returns the table names it changed."""
        changed: list[str] = []
        for ws in wb.worksheets:
            for table in ws.tables.values():
                if getattr(table, "tableType", None) != "queryTable":
                    continue
                table.tableType = None
                for col in table.tableColumns:
                    col.queryTableFieldId = None
                    col.uniqueName = None
                changed.append(table.displayName or table.name)
        return changed

    def _carry_flags(self, wb) -> dict:
        snap = self._surgical_snap
        if snap is None:
            return {"sheets_unchanged": False, "custom_props_unchanged": False,
                    "external_links_unchanged": False}
        now = _surgical.wb_fingerprint(wb)
        return {"sheets_unchanged": now[0] == snap.wb_fp[0],
                "custom_props_unchanged": now[5] == snap.wb_fp[5],
                "external_links_unchanged": now[10] == snap.wb_fp[10]}

    def _cache_keeper(self, skip):
        """keep(sheet, old_addr, new_addr) for carry_forward: is the formula
        at new_addr the one that was at old_addr? After a row or column
        insert or delete the references were rewritten, so any formula the
        op did not write counts (its value is then flagged by the
        recalculate-on-open flag); otherwise the formula must be exactly the
        one that loaded."""
        from openpyxl.utils.cell import coordinate_to_tuple

        snap = self._surgical_snap
        wb = self._workbook
        structural = bool(self._structural)
        titles = self._title_map()

        def keep(sheet: str, old: str, new: str) -> bool:
            # sheet is the title the original package has; a renamed sheet
            # is found under its title now
            sheet = titles.get(sheet, sheet)
            if skip(sheet, new):
                return False
            if structural:
                return True
            if snap is None or wb is None or sheet not in wb.sheetnames:
                return False
            ws = wb[sheet]
            ss = snap.sheets.get(id(ws))
            if ss is None:
                return False
            try:
                ok = ss.cells.get(coordinate_to_tuple(old))
                cell = ws._cells.get(coordinate_to_tuple(new))
            except Exception:  # noqa: BLE001
                return False
            if ok is None or cell is None:
                return False
            now = _surgical.cell_key(cell)[0]
            if now == ok[0]:
                return True
            # a rename rewrote the formula's text, not its result
            return (now[0] == ok[0][0] == "str"
                    and isinstance(now[1], str) and isinstance(ok[0][1], str)
                    and ok[0][1].startswith("=")
                    and self._same_formula(ok[0][1], now[1]))

        return keep

    def _rewrite_staleness(self, wb):
        """Formulas whose results this rewrite may have made stale. After a
        row/column insert or delete every formula is treated as possibly
        stale (references moved); otherwise the dependency walk over the
        cells whose content changed."""
        changes = self._route_changes
        if changes is None:
            return None
        if self._structural:
            return None
        dirty = {}
        for title, cells in changes.items():
            pts = {rc for rc, (cell, ok, nk) in cells.items()
                   if cell is None or ok[0] != nk[0] or ok[1] != nk[1]}
            if pts:
                dirty[title] = pts
        try:
            return _surgical.stale_formulas(wb, dirty)
        except Exception:  # noqa: BLE001
            return _surgical.Staleness([], complete=False)

    def _run_verify(self, target: str, allow_loss: bool) -> _verify.VerifyResult:
        """Run the verify gate and NEVER let it raise.

        Every check inside verify reads a package the server just produced,
        and a package that is broken in the right way makes those readers
        throw rather than return a verdict (a dangling
        chart relationship made content_readback's openpyxl load raise
        KeyError straight out of save). A raising verify is worse than a
        failing one: in the POST-promote position the exception skips the
        restore-from-backup branch entirely and leaves the broken file in
        place. An exception is therefore a verification FAILURE, reported as
        one."""
        # Content-verified comment-only vmlDrawing*.vml anchors: openpyxl
        # regenerates them under its own commentsDrawing
        # naming, so exactly those original names legitimately vanish. The
        # comment PARTS themselves stay guarded by the renumbered legacy-
        # comment family, so a note that actually got lost still fails.
        expected = set(self._expected_removals)
        if self._hazard is not None:
            expected.update(
                p.lower() for p in self._hazard.comment_anchor_vml)
        try:
            if self._integrity_baseline is None:
                self._integrity_baseline = _verify.package_problems(self.path)
            return _verify.verify_after_write(
                target, pre_parts=self._pre_parts, intended=self._intended,
                allow_loss=allow_loss, pre_sizes=self._pre_sizes,
                expected_removals=frozenset(expected),
                loss_keys=self._warned_loss_keys or None,
                integrity_baseline=self._integrity_baseline)
        except Exception as exc:  # noqa: BLE001
            return _verify.VerifyResult(
                ok=False,
                reasons=[f"the verification pass could not read the produced "
                         f"package ({type(exc).__name__}: {exc}); treating "
                         "that as a failed verify"])

    def _restore_from_backup(self) -> bool:
        import shutil
        slot = _safesave.slot_dir(self.path) / _safesave.PREV_SLOT
        if slot.exists():
            shutil.copy2(str(slot), self.path)
            return True
        return False

    def save(self, *, allow_loss: bool = False, backup: bool = True,
             saver: Callable[[str], None] | None = None,
             verify_com: bool | None = None) -> dict:
        """Persist pending mutations through the full safety pipeline and
        return the mutation-success envelope. Raises HazardRefused,
        ValidationFailed, or WorkbookLocked (all mapped to closed codes by the
        envelope). On any refusal the original file is left as it was (a failed
        pre-promote verify never touches it; a failed post-promote verify
        restores it from the backup).

        verify_com is the DEEP, authoritative check: after promotion the file
        is opened in a private hidden Excel worker and a repair-free open is
        required. It is the third layer of the "Excel will refuse this" gate
        (core.limits is the first, honest reporting the second), and it stays
        OPT-IN because it costs a COM round trip and needs Excel installed.
        None means "take the default", which is False unless KS4XL_VERIFY_COM
        is set."""
        if verify_com is None:
            verify_com = com_verify_default()
        path = self.path
        # A backup slot is never an edit target (before the lock, which
        # would itself create a slot folder inside the store).
        _safesave.refuse_backup_slot_target(path, "save edits to")
        _protection_gate.check(self)  # Protection: before any disk change
        self._refuse_partial_array_writes()  # Arrays: before any disk change
        self._settle_tables()  # Tables: before any disk change
        with _safesave.write_lock(path):
            self._check_unchanged()
            lock_present, writable, why = self._lock_diagnosis()
            if not writable:
                from . import access as _access
                raise WorkbookLocked(read_only_refusal(path)
                                     if why == "read_only_attribute"
                                     else _access.access_denied_text(path)
                                     if why == _access.ACCESS_DENIED
                                     else held_refusal(path))
            warnings = self._route_gate(allow_loss)
            warnings.extend(self._precision_warnings())
            warnings.extend(self._text_number_warnings())
            warnings.extend(self._notes)
            reserved_before = _hazard.write_reservation(path)
            if lock_present:
                warnings.append(
                    f"a stale owner lockfile (~${Path(path).name}) is "
                    "present but the file is writable; a prior Excel "
                    "session likely crashed. Proceeding; the ~$ file can "
                    "be deleted safely")

            # The sheet filter over a table that earlier
            # versions wrote makes a workbook Excel will not open. What the
            # workbook already carries is repaired in the payload below;
            # anything an edit would add is not, and the verify gate fails
            # on whatever is left (core.tablefilter, core.verify).
            from . import tablefilter as _tablefilter
            old_damage = _tablefilter.scan(path).keys()

            _sweep_stale_write_tmps(os.path.dirname(os.path.abspath(path)))
            # The backup folder's place taken by a file refuses
            # before anything is written, in plain words.
            if backup or _AUTO_SNAPSHOT.get() is not None:
                _safesave.check_backup_folder(path)
            tmp = os.path.join(
                os.path.dirname(os.path.abspath(path)) or ".",
                f"{_safesave.WRITE_TMP_PREFIX}{uuid.uuid4().hex}{Path(path).suffix}")
            try:
                (saver or self._default_saver)(tmp)
            except PermissionError:
                _silent_remove(tmp)
                # The cause, never "may be open in Excel"
                # with the raw system text.
                from . import access as _access
                raise WorkbookLocked(_access.write_refusal_text(path)) from None
            except OSError as exc:
                # Disk-full and friends: the refusal is environmental, not
                # the caller's parameters, and the half-written temp must
                # not linger.
                _silent_remove(tmp)
                raise _os_write_error(Path(path).name, exc)
            except BaseException:
                _silent_remove(tmp)
                raise

            # A custom document property linked to a cell is refreshed into
            # that cell when Excel opens the file: an edited cell must carry
            # its new value into the property (core/linkprops.py), on either
            # save route, before the verify step reads the package.
            try:
                linked = _linkprops.sync(path, tmp, allow_loss=allow_loss)
            except BaseException:
                _silent_remove(tmp)
                raise
            warnings.extend(linked.warnings)
            if linked.changed:
                self._changed["linked_properties"] = linked.changed

            if old_damage:
                try:
                    warnings.extend(_tablefilter.repair_package(
                        tmp, allowed=old_damage))
                except Exception:  # noqa: BLE001 - the gate fails on what is left
                    pass

            pre = self._run_verify(tmp, allow_loss)
            if not pre.ok:
                _silent_remove(tmp)
                raise ValidationFailed(
                    "verify-after-write failed; the original file was NOT "
                    "modified. " + "; ".join(pre.reasons))

            # Two-phase backup rotation: STAGE the pre-mutation content now,
            # COMMIT it onto the slots only after a successful promote. A
            # promote that fails must leave the slots untouched; rotating
            # first burned the prev undo of the previous successful mutation
            # even though nothing was written.
            # The point of no return. The active call's deadline commits
            # here or the call refuses TIMEOUT with nothing on disk changed
            # (the validated temp is removed, the workbook and its slots are
            # untouched).
            from . import deadline as _deadline

            try:
                _deadline.commit_or_raise("the save")
            except BaseException:
                _silent_remove(tmp)
                raise
            backup_slot = None
            ticket = None
            # A labelled copy of the workbook as it was, before a
            # destructive call replaces it (once per call, first save).
            snap = _AUTO_SNAPSHOT.get()
            snap_file = None
            try:
                if snap is not None and snap["label"] is None and os.path.isfile(path):
                    from ..ops import backups as _backups

                    snap["label"], snap_file = _backups.write_auto_snapshot(
                        path, snap["op"])
                if backup:
                    ticket = _safesave.prepare_rotation(path)
            except BaseException as exc:
                # Nothing changed, so nothing of this call stays
                # behind: not the payload, not the snapshot.
                _silent_remove(tmp)
                if snap_file is not None:
                    _silent_remove(str(snap_file))
                    snap["label"] = None
                if isinstance(exc, OSError):
                    raise _os_write_error(Path(path).name, exc) from exc
                raise
            try:
                # ReplaceFileW, so the workbook keeps its
                # permissions, Mark of the Web, attributes and creation time.
                # The staged rotation lands inside the promote,
                # after the swap, and is dropped there when the swap fails.
                warnings.extend(_safesave.promote(tmp, path, ticket=ticket))
            except BaseException as exc:
                if ticket is not None:
                    ticket.abort()
                _silent_remove(tmp)
                if snap_file is not None:
                    # Nothing changed, so the copy is not the state before
                    # an edit; it goes, and the label with it.
                    _silent_remove(str(snap_file))
                    snap["label"] = None
                if isinstance(exc, PermissionError):
                    from . import access as _access
                    raise WorkbookLocked(_access.replace_refusal_text(
                        path, getattr(exc, "winerror", None))) from None
                if isinstance(exc, OSError):
                    raise _os_write_error(Path(path).name, exc)
                raise
            if ticket is not None and ticket.rotated["prev"]:
                backup_slot = "prev"

            post = self._run_verify(path, allow_loss)
            if not post.ok:
                restored = False
                if backup:
                    restored = self._restore_from_backup()
                raise ValidationFailed(
                    "post-promote verify failed"
                    + (" and the file was restored from the backup"
                       if restored else "")
                    + ". " + "; ".join(post.reasons))

            warnings.extend(self._reservation_warnings(reserved_before))

            verified_com = None
            if verify_com:
                verified_com = self._verify_com_after_promote(
                    path, backup=backup, warnings=warnings)

            # The on-disk package changed; the cached model and scan are stale.
            self._release_model()
            self._hazard = _hazard.scan_path(path)
            self._pre_parts = list(self._hazard.parts)
            self._pre_sizes = dict(self._hazard.sizes)
            changed = dict(self._changed)
            if self._intended:
                entries = [{"sheet": s, "cell": c, "kind": k}
                           for (s, c), (k, _v) in self._intended.items()]
                if len(entries) <= ECHO_LIMIT or _ECHO_CELLS.get():
                    changed["cells"] = entries
                else:
                    changed["cells_summary"] = _cells_summary(entries)
            if self._preserved_number_texts:
                changed["full_precision_cells_preserved"] = \
                    self._preserved_number_texts
            self._intended = {}
            self._text_numbers = {}
            self._formula_written = False
            self._dynamic_written = set()
            self._structural = []
            self._notes = []
            self._changed = {}
            self._expected_removals = set()
            self._expected_preserved = set()
            self._warned_loss_keys = set()
            self._dropped_extensions = []
            self._written_regions = []
            self._text_rewrites = []
            # The file we just wrote is now the baseline a second save on this
            # package must compare against.
            self._stamp = _file_stamp(path)
            result = {
                "ok": True,
                "file": path,
                "changed": changed,
                "saved": True,
                "backup": backup_slot,
                "verified": True,
                "warnings": warnings,
            }
            result.update(self._route_report())
            self._route, self._route_build, self._route_note = None, None, None
            self._integrity_baseline = None
            if verified_com is not None:
                result["verified_com"] = verified_com
            return result

    def _route_report(self) -> dict:
        """What the save did beyond the cells: the route, and the truth
        about formula results."""
        build = self._route_build
        if build is None:
            out: dict = {"route": self._route or _surgical.ROUTE_REWRITE}
            if self._route_note:
                out["route_note"] = _rewrite_needed(self._route_note)
            carry = self._carry
            if carry is not None and carry.as_dict():
                out.update({k: v for k, v in carry.as_dict().items()
                            if k != "cached_values_restored"})
            cached: dict = {}
            if carry is not None:
                cached["kept"] = carry.cached_values
                if carry.cached_dropped:
                    cached["dropped"] = carry.cached_dropped
                    cached["dropped_note"] = (
                        "these formulas have no cached result in the saved "
                        "file until Excel recalculates it")
            stale = self._rewrite_stale
            if self._route_structural:
                cached["possibly_stale"] = "all"
                cached["possibly_stale_note"] = (
                    "rows or columns moved, so any formula result may be out "
                    "of date until Excel recalculates")
            elif stale is not None and (stale.cells or not stale.complete):
                cached["possibly_stale"] = len(stale.cells)
                cached["possibly_stale_cells"] = [
                    f"{s}!{c}" for s, c in stale.cells[:10]]
            if cached:
                out["cached_values"] = cached
            return out
        out = {"route": _surgical.ROUTE_RAW}
        stale = build.stale
        cached: dict = {"kept": "every formula result this edit did not "
                                "affect keeps its cached value"}
        if build.formulas_written:
            cached["formulas_written"] = len(build.formulas_written)
        if stale.cells or not stale.complete:
            cached["possibly_stale"] = len(stale.cells)
            cached["possibly_stale_cells"] = [
                f"{s}!{c}" for s, c in stale.cells[:10]]
            if not stale.complete:
                cached["possibly_stale_note"] = (
                    "the dependency check stopped at its work limit, so "
                    "more formulas than listed may depend on this edit")
        if build.full_calc:
            cached["recalculates_on_open"] = True
        out["cached_values"] = cached
        return out

    def _verify_com_after_promote(self, path: str, *, backup: bool,
                                  warnings: list[str]) -> bool | None:
        """The optional DEEP verification (verify_com:true): after promotion,
        open the produced file in a private hidden Excel worker and require a
        clean, repair-free open (com.session.opens_clean). On failure the
        backup is restored and the save refuses. When COM is unavailable the
        save stands on the file-tier verify and says so in warnings (None)."""
        from ..com import session as _com_session
        ok, why = _com_session.com_available()
        if not ok:
            warnings.append(
                f"verify_com was requested but Excel/COM is unavailable "
                f"({why}); the save stands on the file-tier verification "
                "only")
            return None
        try:
            res = _com_session.opens_clean(path)
        except Exception as exc:  # noqa: BLE001
            warnings.append(
                "verify_com could not run (" + f"{type(exc).__name__}: {exc}"
                + "); the save stands on the file-tier verification only")
            return None
        if not res.get("opens_clean"):
            restored = self._restore_from_backup() if backup else False
            raise ValidationFailed(
                "deep verification failed: Excel refused the produced file "
                "or demanded a repair (" + str(res.get("excel_says", ""))
                + ")" + (" and the file was restored from the backup"
                         if restored else ""))
        return True

    def _release_model(self) -> None:
        """Drop the loaded model and its snapshot. openpyxl keeps a .xlsm's
        VBA parts in an in-memory zip it never closes; closing it here keeps
        its finalizer from running against a buffer that is already gone."""
        wb = self._workbook
        self._workbook = None
        self._surgical_snap = None
        arch = getattr(wb, "vba_archive", None) if wb is not None else None
        if arch is not None:
            try:
                arch.close()
            except Exception:  # noqa: BLE001
                pass

    def close(self) -> None:
        if self._workbook is not None:
            try:
                self._workbook.close()
            except Exception:
                pass
            self._release_model()

    def __enter__(self) -> "WorkbookPackage":
        return self

    def __exit__(self, *exc) -> None:
        self.close()


#: Environment switch making the deep Excel verification the DEFAULT for
#: every save on this server, rather than a per-call opt-in.
#:
#: Verify-after-write has no notion of "Excel will refuse this": it proves
#: the produced package is valid OOXML
#: that reads back as intended, and a file Excel throws out can be exactly
#: that. core.limits closes the known routes cheaply, but only Excel can
#: answer for the unknown ones, and com_validate_opens_clean answers
#: correctly (it got all ten of the unopenable test files right). This is
#: the switch that puts it on the write path. It stays OFF by default because
#: it costs a COM round trip per save and needs Excel installed; a headless
#: CI run and a Linux install must keep working. Per-call verify_com:true
#: always wins over the default, and it is reachable from the surface: every
#: mutating tool takes verify_com, so a single risky save can ask for Excel's
#: verdict without turning it on server-wide.
VERIFY_COM_ENV = "KS4XL_VERIFY_COM"


#: Environment switch for the byte-preserving save. On unless set to off:
#: the escape hatch if a workbook ever needs the old full-rewrite behavior.
SURGICAL_ENV = "KS4XL_SURGICAL_SAVE"


def surgical_enabled() -> bool:
    return os.environ.get(SURGICAL_ENV, "").strip().lower() not in (
        "0", "off", "false", "no")


def _rewrite_needed(reason: str) -> str:
    """Why an edit could not keep the rest of the file byte-for-byte."""
    return (f"Cell edits leave everything else in the file untouched, but "
            f"{reason}, which needs the full rewrite.")


def com_verify_default() -> bool:
    """Whether saves deep-verify through Excel unless told otherwise. Never
    when the launch settings turn Excel off (KS4XL_FORBID=xl_app): that
    setting wins over KS4XL_VERIFY_COM."""
    from .. import runtime as _runtime

    if _runtime.excel_forbidden():
        return False
    return os.environ.get(VERIFY_COM_ENV, "").strip().lower() in (
        "1", "true", "yes", "on")


def _file_stamp(path: str) -> tuple[int, int] | None:
    """(mtime_ns, size) identity of the file on disk, or None if it is gone.
    The same identity the COM tier's saves compare. Kept here, not
    imported from the save module, because every read reaches it."""
    try:
        st = os.stat(path)
    except OSError:
        return None
    return (st.st_mtime_ns, st.st_size)


#: openpyxl's own wording when it discards an extLst block it cannot model
#: ("Conditional Formatting extension is not supported and will be removed",
#: "Slicer List extension ...", "Unknown extension ...").
_EXT_DROP_MARKER = "is not supported and will be removed"


def _extension_drops(caught) -> list[str]:
    """De-duplicated extension-drop notices from a captured warning list."""
    out: list[str] = []
    for w in caught or ():
        text = str(getattr(w, "message", w))
        if _EXT_DROP_MARKER in text and text not in out:
            out.append(text)
    return out


def _silent_remove(path: str) -> None:
    try:
        os.remove(path)
    except OSError:
        pass


def _os_write_error(name: str, exc: OSError) -> XlMcpError:
    """An honest refusal for OS-level write failures. Disk-full used to map
    to BAD_PARAMS ("the caller's mistake") with the raw errno text;
    it is the environment's state, so it carries
    CONFLICT and says so in plain words. The original file is intact in
    every case: the failure happened on the temp or at the promote, never
    mid-way through the workbook's own bytes."""
    if isinstance(exc, _safesave.WorkbookLeftAside):
        # Error 1177 and the move back failed; the workbook is safe,
        # aside, and comes back at the next change or server start.
        err = XlMcpError(
            f"{name} could not be saved, and the workbook could not be put "
            "back in its place: it is kept, as it was before this save, at "
            f"{exc.aside}; nothing in it was lost. The next change to it "
            "through this server, or the next start of the server, puts it "
            f"back; you can also rename that file to {name}.")
        err.code = "CONFLICT"
        return err
    err = XlMcpError(
        (f"{name}: the disk is full (no space left to write the new "
         "version). The original file is intact and was NOT modified; free "
         "some disk space and retry.")
        if getattr(exc, "errno", None) == 28 else
        (f"{name}: the operating system refused the write "
         f"({type(exc).__name__}: {exc}). The original file is intact and "
         "was NOT modified."))
    err.code = "CONFLICT"
    return err


#: A .ks4xl-write-* temp older than this is an orphan from a killed process
#: (no mutation legitimately runs this long; the write lock's stale window
#: is 10 minutes) and is swept on the next save in the same folder. Kill
#: storms left these forever; the write.lock self-heals, the temps did not.
_STALE_TMP_SECONDS = 60 * 60


def _sweep_stale_write_tmps(directory: str) -> None:
    """Best-effort janitor for aged .ks4xl-write-* temp files. Runs under
    the write lock on the save path; a FRESH temp (a concurrent writer's
    live work) is never touched.

    It never removes a workbook that a save set aside. The current aside
    name (.ks4xl-replaced-...) does not start with the scratch prefix;
    the name earlier versions gave it,
    .ks4xl-write-<id>.replaced, does, and is skipped by name: it can be the
    only copy of a workbook whose save was stopped mid-swap, and a rename
    keeps the workbook's own old modification time, so age says nothing.
    The new version such a save was moving in is skipped too while the save
    journal still records it: put_back_if_aside decides about it once the
    workbook is back."""
    try:
        now = __import__("time").time()
        protected = None
        with os.scandir(directory or ".") as it:
            for entry in it:
                if not entry.name.startswith(_safesave.WRITE_TMP_PREFIX):
                    continue
                if entry.name.endswith(".replaced"):
                    continue  # an earlier version's aside: never by age
                try:
                    if not entry.is_file():
                        continue
                    if now - entry.stat().st_mtime > _STALE_TMP_SECONDS:
                        if protected is None:
                            protected = _safesave.protected_new_versions()
                        if os.path.normcase(os.path.abspath(entry.path)) in protected:
                            continue
                        os.remove(entry.path)
                except OSError:
                    continue
    except OSError:
        pass


__all__ = ["WorkbookPackage", "com_verify_default", "VERIFY_COM_ENV"]
