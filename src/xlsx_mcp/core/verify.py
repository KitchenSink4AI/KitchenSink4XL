"""core/verify.py: verify-after-write.

The brand-defining guarantee, engineered as a real check rather than asserted.
After any file-based mutation, before the produced package is promoted onto the
user's path, the server re-opens it and confirms three things:

  1. STRUCTURAL: the produced zip re-opens as a valid OOXML package, its
     ``[Content_Types].xml`` and workbook rels resolve, at least one sheet
     stays visible, every worksheet part and workbook.xml parse as
     well-formed XML (a streaming parse, so a truncated or garbage sheet
     cannot pass just because openpyxl's lazy read-only load never touched
     it), and openpyxl can parse the package.
  2. NO UNEXPECTED PART LOSS OR REPLACEMENT, DEFAULT-FAIL: the produced
     package's part inventory (names and uncompressed sizes) is diffed
     against the pre-write scan, and the default is REFUSAL. ANY part that
     was present before and is missing or empty afterward fails the write
     unless its loss is explained by one of four documented excuses:
       (a) the legitimate-rewrite pass-list (_PASS_EXACT / _PASS_PREFIX),
           each entry justified by an observed, harmless round-trip
           behavior recorded next to it;
       (b) a model-regenerated family (_RENUMBERED_FAMILIES) whose part
           NAMES openpyxl reassigns on every save, judged by family COUNT
           instead of name identity (a renumber is not a loss; a shrinking
           count still fails unless explained);
       (c) an explicit allow_loss override, which excuses ONLY parts that
           match a named hazard family the user was warned about, never an
           unrelated surprise loss;
       (d) a deliberate model-level removal the mutating op registered via
           WorkbookPackage.expect_removal (sheet delete, table to_range,
           comment delete).
     Fragile parts (hazard table) additionally keep their stricter checks:
     a fragile part still PRESENT but replaced with empty content, or
     whose XML no longer parses, also fails (loss by replacement, not just
     by omission).
  3. CONTENT READ-BACK: the specific cells the tool claimed to write are
     re-read and compared against intent (literal values or formula strings).
     A mismatch fails the write.

On failure the caller does NOT promote the temp file (the original is never
touched), and raises ValidationFailed (VALIDATION_FAILED) stating what failed.
Where COM is available the caller additionally runs com_validate_opens_clean
as the authoritative corruption smoke test; this module stays file-only so it
runs in headless CI. Verify-after-write is a bounded, honest check on THIS
server's own writes; it cannot certify a file the server did not produce.
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from typing import Iterable

from . import calc as _calc
from . import limits as _limits
from . import hazard as _hazard
from . import xmlsafe as _xmlsafe

#: Parts openpyxl legitimately drops or regenerates on a normal save; their
#: absence after a write is expected, never a loss. (Parts under
#: xl/printerSettings/ are also expected drops, but they match no hazard spec,
#: so the loss check never sees them; the hazard module documents that policy.)
EXPECTED_DROPPABLE = frozenset({
    "xl/calcchain.xml",
    "docprops/thumbnail.jpeg",
})

# --------------------------------------------------------------------------
# THE LEGITIMATE-REWRITE PASS-LIST for the default-fail inventory check.
#
# Every entry below is justified by an OBSERVED openpyxl round-trip behavior,
# measured on this machine against the fixture corpus (and targeted synthetic
# injections) with the package's own load flags (keep_vba per extension,
# rich_text=False, data_only=False). An entry with no observation does not
# belong here; an uncataloged part type that vanishes must FAIL, that is the
# whole point of the default-fail posture.

#: Exact part names (compared lowercase) whose absence after a save is a
#: harmless, observed rewrite behavior.
_PASS_EXACT: dict[str, str] = {
    "xl/calcchain.xml": (
        "OBSERVED: injected xl/calcChain.xml into clean.xlsx, round-trip "
        "dropped it. It is a rebuildable calculation-order cache; Excel "
        "regenerates it on the next open, no user content lives in it."),
    "xl/sharedstrings.xml": (
        "OBSERVED: pivot.xlsx / pivot_slicer.xlsx / shape.xlsx (Excel-"
        "authored) round-trips drop xl/sharedStrings.xml; the strings are "
        "re-emitted INLINE in the sheet XML (t=\"inlineStr\", read back and "
        "confirmed intact). A representation change, not a loss."),
}

#: Part-name prefixes (compared lowercase) whose absence after a save is a
#: harmless, observed rewrite behavior.
_PASS_PREFIX: dict[str, str] = {
    "xl/printersettings/": (
        "OBSERVED: injected xl/printerSettings/printerSettings1.bin, round-"
        "trip dropped it. Device page-setup blobs; the pageSetup element in "
        "the sheet XML survives. Policy documented in core/hazard.py."),
    "docprops/thumbnail.": (
        "OBSERVED: injected docProps/thumbnail.jpeg, round-trip dropped it. "
        "A preview image Excel regenerates; openpyxl never writes one."),
}

#: Model-regenerated part FAMILIES whose names openpyxl reassigns from live
#: model ids on every save, so name identity is meaningless across a save and
#: the check compares family COUNTS instead. A member 'lost' while the family
#: count holds is a renumber, not a loss; a count that SHRANK fails unless an
#: expected_removals registration explains it.
_RENUMBERED_FAMILIES: tuple[tuple[str, re.Pattern, str], ...] = (
    ("worksheet", re.compile(r"^xl/worksheets/sheet\d+\.xml$"),
     "OBSERVED: a workbook carrying xl/worksheets/sheet5.xml round-trips to "
     "xl/worksheets/sheet1.xml, content intact (openpyxl assigns sequential "
     "ids at load, not original part names)."),
    ("chartsheet", re.compile(r"^xl/chartsheets/sheet\d+\.xml$"),
     "OBSERVED: xl/chartsheets/sheet9.xml round-trips to "
     "xl/chartsheets/sheet1.xml, same sequential-id mechanism."),
    ("legacy comment", re.compile(
        r"^(xl/comments\d+\.xml|xl/comments/comment\d+\.xml)$"),
     "OBSERVED: Excel-style xl/comments1.xml round-trips to openpyxl-style "
     "xl/comments/comment1.xml with the comment intact (a rename AND a "
     "relocation, so both spellings are one family)."),
    ("comment VML anchor",
     re.compile(r"^xl/drawings/commentsdrawing\d+\.vml$"),
     "OBSERVED: comment anchors are written with model-assigned sequential "
     "ids like the comment parts they accompany (removing a comment via the "
     "model dropped xl/drawings/commentsDrawing1.vml alongside its comment "
     "part). Excel-authored vmlDrawing*.vml stays under the drawings hazard "
     "spec and is not excused here."),
)


def _passlisted(low: str) -> bool:
    return (low in _PASS_EXACT
            or any(low.startswith(p) for p in _PASS_PREFIX))


def _is_rels_part(low: str) -> bool:
    """Relationship parts are derived metadata openpyxl regenerates from the
    model on every save; a rels part legitimately vanishes when its last
    relationship does (OBSERVED: shape.xlsx round-trip drops
    xl/worksheets/_rels/sheet1.xml.rels once its only target, the dropped
    drawing, is gone). Every rels TARGET is itself a part in this same
    inventory, so excusing the rels part loses no coverage, and a rels loss
    that breaks package resolution still fails the structural check."""
    return low.endswith(".rels") and (low == "_rels/.rels"
                                      or "/_rels/" in low)


def _family_label(low: str) -> str | None:
    for label, rex, _obs in _RENUMBERED_FAMILIES:
        if rex.match(low):
            return label
    return None

#: Ceiling for the per-part XML well-formedness re-parse in the replacement
#: check. Fragile parts are typically small; anything bigger is skipped there
#: (the worksheet streaming parse in structural_check has no such cap).
_XML_RECHECK_MAX_BYTES = 8 * 1024 * 1024


@dataclass
class VerifyResult:
    ok: bool
    reasons: list[str] = field(default_factory=list)
    lost_parts: list[str] = field(default_factory=list)
    mismatches: list[dict] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {
            "ok": self.ok,
            "reasons": self.reasons,
            "lost_parts": self.lost_parts,
            "mismatches": self.mismatches,
        }


#: The error values Excel opens as a plain stored value under t="e"
#: (measured with Excel's own open of each one, in any letter case).
_ERROR_LITERALS = frozenset({
    "#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A",
    "#GETTING_DATA", "#BUSY!"})
#: Errors Excel stores only as rich-value errors (#VALUE! with a value
#: metadata index): stored plainly, as a constant or as a formula's cached
#: result, each makes Excel refuse the whole file (measured).
_RICH_VALUE_ERRORS = frozenset({
    "#SPILL!", "#CALC!", "#FIELD!", "#BLOCKED!", "#CONNECT!", "#UNKNOWN!",
    "#PYTHON!", "#EXTERNAL!"})
#: A stored number as Excel reads one (measured with Excel's own open of
#: each form: " 5", "5 ", ".5", "5.", "1e5", "1E+05", "-.5", "007", "-0",
#: "\t5" and "-.5e1" open clean, and "-" and "." open as 0; "+5", "- 5",
#: "INF", "-INF", "NaN", "1,5", "0x10", "1_000", "1e", "1e+" and "1.5.2" do
#: not; an empty value opens as an empty cell).
_NUMBER_TEXT = re.compile(r"-?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?")
#: Cell types whose stored value is free text (or not in <v> at all).
_FREE_TYPES = frozenset({"str", "inlineStr"})
#: How many contradicting cells one refusal names.
_TYPE_SHOWN = 8


def _value_contradiction(ctype: str | None, value: str,
                         shared_count: int | None) -> str | None:
    """Why a stored <v> contradicts its cell type, or None when it fits.

    Excel reads <v> by the cell's t= attribute. A value that is not of that
    type makes the whole workbook unopenable (measured: a number cell
    holding text fails "Open method of Workbooks class" in normal, repair
    and extract-data modes), so this is a property of the produced file,
    not a cosmetic one."""
    if ctype in _FREE_TYPES:
        return None
    shown = value if len(value) <= 40 else value[:37] + "..."
    v = value.strip()
    if ctype == "s":
        # an empty index is not a string: Excel refuses it like any other
        if not v.isascii() or not v.isdigit():
            return f"holds {shown!r} where a shared-string number belongs"
        if shared_count is not None and int(v) >= shared_count:
            return (f"points at shared string {v}, and the workbook has "
                    f"{shared_count}")
        return None
    if not v:
        return None  # Excel opens an empty value as an empty cell
    if ctype == "b":
        if v in ("0", "1", "true", "false"):
            return None
        return f"holds {shown!r} in a TRUE/FALSE cell"
    if ctype == "e":
        if v.upper() in _ERROR_LITERALS:
            return None
        if v.upper() in _RICH_VALUE_ERRORS:
            return (f"holds {shown!r} as a plain error value, which Excel "
                    "stores only as a rich value")
        return f"holds {shown!r} in an error cell"
    if ctype == "d":
        import datetime as _dt
        try:
            _dt.datetime.fromisoformat(v.rstrip("Z"))
        except ValueError:
            return f"holds {shown!r} in a date cell"
        return None
    # a number cell; Excel reads an unknown t= the same way (measured)
    if (_NUMBER_TEXT.fullmatch(v) and v.isascii()) or v in ("-", "."):
        return None  # Excel reads a lone "-" or "." as 0 (measured)
    return f"holds the text {shown!r} in a number cell"


def _shared_string_count(zf: zipfile.ZipFile, names: set[str]) -> int:
    """The number of shared strings; 0 when the package has no shared
    string part (a t="s" cell then points at nothing)."""
    part = None
    try:
        rels = zf.read("xl/_rels/workbook.xml.rels")
        m = re.search(rb'<(?:[\w.\-]+:)?Relationship\b[^>]*?Type="[^"]*/'
                      rb'sharedStrings"[^>]*>', rels)
        if m:
            t = re.search(rb'Target="([^"]+)"', m.group(0))
            if t:
                target = t.group(1).decode("utf-8", "replace")
                part = _calc._resolve_part(target)
    except KeyError:
        pass
    if part is None:
        part = next((n for n in names
                     if n.lower() == "xl/sharedstrings.xml"), None)
    if part is None or part not in names:
        return 0
    data = zf.read(part)
    return len(re.findall(rb"<(?:[\w.\-]+:)?si[\s>/]", data))


def _scan_cell_values(fh, shared_count: int | None, found: list, *,
                      doc: str = "the workbook", part: str = "a worksheet"
                      ) -> None:
    """Stream one worksheet part, appending (cell, why) for every stored value
    that contradicts its cell's type, and every formula past Excel's length
    limit. The parse is xmlsafe's: a part that declares a DTD is refused, and
    bad XML raises ParseError."""
    for _event, elem in _xmlsafe.iterparse(fh, doc=doc, part=part):
        tag = elem.tag
        if not (tag.endswith("}c") or tag == "c"):
            if tag.endswith("}row") or tag == "row":
                elem.clear()  # its cells are read: keep the parse flat
            continue
        v = None
        for child in elem:
            ct = child.tag
            if ct.endswith("}v") or ct == "v":
                v = child.text
            elif (ct.endswith("}f") or ct == "f") and child.text and \
                    len(child.text) > _limits.MAX_FORMULA_CHARS:
                # counted as stored, as Excel counts it (core.limits)
                found.append((elem.get("r") or "?",
                              f"holds a formula of {len(child.text):,} "
                              "characters, past Excel's limit of "
                              f"{_limits.MAX_FORMULA_CHARS:,}"))
        if v is None:
            continue
        why = _value_contradiction(elem.get("t"), v, shared_count)
        if why is not None:
            found.append((elem.get("r") or "?", why))


def structural_check(path: str, *, cell_values: list | None = None
                     ) -> tuple[bool, list[str]]:
    """Re-open the produced package and confirm it is structurally valid.

    With cell_values (a list), the same streaming parse of every worksheet
    also reads each stored value against its cell's type, and appends one
    sentence per worksheet that holds a contradiction (see
    cell_value_check)."""
    reasons: list[str] = []
    try:
        with zipfile.ZipFile(path) as zf:
            bad = zf.testzip()
            if bad is not None:
                return False, [f"corrupt zip member: {bad}"]
            names = set(zf.namelist())
            # Every worksheet part and workbook.xml must parse as well-formed
            # XML. This is a streaming parse (no tree kept), and it exists
            # because the openpyxl read_only load below is LAZY: it never
            # touches sheet XML, so a truncated or garbage sheet part would
            # otherwise pass whenever no cell read-back runs (exactly the
            # structural-edit and raw-surgery saves).
            xml_parts = sorted(
                n for n in names
                if n.lower().startswith("xl/worksheets/")
                and n.lower().endswith(".xml")
                and "/_rels/" not in n.lower())
            if "xl/workbook.xml" in names:
                xml_parts.append("xl/workbook.xml")
            shared = titles = None
            if cell_values is not None:
                shared = _shared_string_count(zf, names)
                titles = {part: title for title, part
                          in _calc._sheet_parts(zf).items()}
            for member in xml_parts:
                try:
                    with zf.open(member) as fh:
                        if cell_values is not None and                                 member != "xl/workbook.xml":
                            found: list = []
                            _scan_cell_values(fh, shared, found, doc=path,
                                              part=member)
                            if found:
                                cell_values.append(_cell_value_sentence(
                                    titles.get(member, member), found))
                        else:
                            _xmlsafe.walk(fh, doc=path, part=member)
                except Exception as exc:  # noqa: BLE001
                    reasons.append(
                        f"{member} is not well-formed XML "
                        f"({type(exc).__name__})")
    except zipfile.BadZipFile:
        return False, ["produced file is not a valid zip / OOXML package"]
    except FileNotFoundError:
        return False, ["produced file is missing"]
    if reasons:
        return False, reasons
    if "[Content_Types].xml" not in names:
        reasons.append("missing [Content_Types].xml")
    if "xl/workbook.xml" not in names:
        reasons.append("missing xl/workbook.xml")
    if "xl/_rels/workbook.xml.rels" not in names:
        reasons.append("missing workbook relationships")
    # openpyxl must be able to parse it, and at least one sheet must stay
    # visible.
    try:
        keep_vba = _xmlsafe.keeps_vba(path)
        wb = _xmlsafe.load_workbook(path, read_only=True, keep_vba=keep_vba,
                                    data_only=False)
        try:
            visible = [ws for ws in wb.worksheets
                       if getattr(ws, "sheet_state", "visible") == "visible"]
            if not visible:
                reasons.append("no visible worksheet remains")
        finally:
            wb.close()
    except Exception as exc:  # noqa: BLE001
        reasons.append(f"openpyxl cannot parse the produced file: "
                       f"{type(exc).__name__}: {exc}")
    return (not reasons), reasons


def _cell_value_sentence(title: str, found: list) -> str:
    sheet = title if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", title) \
        else "'" + title.replace("'", "''") + "'"
    shown = "; ".join(f"{sheet}!{ref} {why}" for ref, why
                      in found[:_TYPE_SHOWN])
    more = f" (+{len(found) - _TYPE_SHOWN} more)" \
        if len(found) > _TYPE_SHOWN else ""
    return shown + more


def _workbook_part(zf: zipfile.ZipFile, names: set[str], rel_type: str,
                   default: str) -> str | None:
    """The part the workbook's relationship of type .../rel_type points at,
    else the default name when the package has it."""
    part = None
    try:
        rels = zf.read("xl/_rels/workbook.xml.rels")
        m = re.search(rb'<(?:[\w.\-]+:)?Relationship\b[^>]*?Type="[^"]*/'
                      + re.escape(rel_type.encode()) + rb'"[^>]*>', rels)
        if m:
            t = re.search(rb'Target="([^"]+)"', m.group(0))
            if t:
                part = _calc._resolve_part(t.group(1).decode("utf-8",
                                                             "replace"))
    except KeyError:
        pass
    if part is None or part not in names:
        part = next((n for n in names if n.lower() == default.lower()), None)
    return part


#: The first eight bytes of an OLE compound file, which a macro project is.
_CFB_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


_QT_REL_TYPES = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/queryTable",
    "http://purl.oclc.org/ooxml/officeDocument/relationships/queryTable",
)


def _query_table_orphans(zf: zipfile.ZipFile, names: set[str]) -> list[str]:
    """A table declared as a query table (tableType="queryTable") must have
    its queryTable part: the table's own relationships name it, with an
    internal relationship of the queryTable type that resolves to a part
    whose root element is a queryTable. A table that keeps the declaration
    and the column ids while the part is gone makes Excel refuse the whole
    file ("Open method of Workbooks class failed").

    The table part and its relationships are read as XML, by namespace and
    after XML decoding (single or double quotes, a prefixed namespace, UTF-16
    and character references read the same); a relationship that is external,
    of another type, or aimed at some other part (a worksheet, say) does not
    count."""
    from . import opcxml as _opc
    by_lower = {n.lower(): n for n in names}
    out: list[str] = []
    for name in sorted(names):
        low = name.lower()
        if (not (low.startswith("xl/tables/") and low.endswith(".xml"))
                or "/_rels/" in low):
            continue
        try:
            root = _xmlsafe.fromstring(zf.read(name), doc="the workbook", part=name)
        except Exception:  # noqa: BLE001
            continue
        if root.tag not in tuple("{%s}table" % ns for ns in _opc.MAIN_NAMESPACES):
            continue
        if (root.get("tableType") or "").strip() != "queryTable":
            continue
        has_part = False
        try:
            rels_name = _opc.rels_part_for(name)
            real_rels = by_lower.get(rels_name.lower())
            rels = _opc.relationships(zf.read(real_rels), part=rels_name) \
                if real_rels else []
        except Exception:  # noqa: BLE001
            rels = []
        for rel in rels:
            if rel.external or rel.type not in _QT_REL_TYPES:
                continue
            target = by_lower.get(_opc.resolve_part(name, rel.target).lower())
            if target is None:
                continue
            try:
                head = _opc.root_tag(zf.read(target), part=target)
            except Exception:  # noqa: BLE001
                continue
            if head is not None and head[0] in _opc.MAIN_NAMESPACES \
                    and head[1] == "queryTable":
                has_part = True
        if not has_part:
            label = root.get("displayName") or name
            out.append(
                f"the table {label!r} is declared as a query table but its "
                "query table part is missing")
    return out


_CT_ROOT = "application/vnd.openxmlformats-officedocument.spreadsheetml."
_CT_MS = "application/vnd.ms-excel."

#: file extension -> (workbook part content type, what Excel calls that kind).
#: Excel opens a file only when the two agree.
_WORKBOOK_KINDS = {
    ".xlsx": (_CT_ROOT + "sheet.main+xml", "a workbook"),
    ".xlsm": (_CT_MS + "sheet.macroEnabled.main+xml", "a macro-enabled workbook"),
    ".xltx": (_CT_ROOT + "template.main+xml", "a template"),
    ".xltm": (_CT_MS + "template.macroEnabled.main+xml", "a macro-enabled template"),
    ".xlam": (_CT_MS + "addin.macroEnabled.main+xml", "a macro-enabled add-in"),
}


def _workbook_type_mismatch(zf: zipfile.ZipFile, names: set[str],
                            path: str) -> list[str]:
    """The file name's extension and the workbook part's content type must
    agree: Excel refuses a .xltm that holds a plain template, a .xlsx that
    holds a macro-enabled workbook, and so on ("the file format or file
    extension is not valid"). The content type is read from
    [Content_Types].xml by namespace: the workbook part's Override, else the
    Default for its extension. [] when the extension is not one of the five
    Open XML workbook kinds, or the part's type cannot be determined (the
    structural check owns a package without one)."""
    import os as _os
    from . import opcxml as _opc
    ext = _os.path.splitext(path)[1].lower()
    want = _WORKBOOK_KINDS.get(ext)
    if want is None:
        return []
    by_lower = {n.lower(): n for n in names}
    wbpart = "xl/workbook.xml"
    try:
        for rel in _opc.relationships(zf.read(by_lower["_rels/.rels"]),
                                      part="_rels/.rels"):
            if rel.type.endswith("/officeDocument") and not rel.external:
                wbpart = _opc.resolve_part("", rel.target)
                break
    except (KeyError, *_xmlsafe.XML_ERRORS):
        pass
    try:
        root = _xmlsafe.fromstring(zf.read(by_lower["[content_types].xml"]),
                                   part="[Content_Types].xml")
    except (KeyError, *_xmlsafe.XML_ERRORS):
        return []
    have = None
    default = None
    ext_of_part = wbpart.rsplit(".", 1)[-1].lower() if "." in wbpart else ""
    for el in root:
        if not isinstance(el.tag, str):
            continue
        local = el.tag.rpartition("}")[2]
        if local == "Override" and (el.get("PartName") or "").lstrip("/").lower() \
                == wbpart.lower():
            have = (el.get("ContentType") or "").strip()
        elif local == "Default" and (el.get("Extension") or "").lower() == ext_of_part:
            default = (el.get("ContentType") or "").strip()
    have = have or default
    if not have or have.lower() == want[0].lower():
        return []
    kind = next((k for e, (ct, k) in _WORKBOOK_KINDS.items()
                 if ct.lower() == have.lower()), None)
    stored = kind or f"the type {have}"
    return [f"the file name ends in {ext}, which Excel opens only as {want[1]}, "
            f"but its workbook part is stored as {stored}"]


def package_refusals(path: str, *, macro_part: bool = False) -> list[str]:
    """What makes Excel refuse a package whatever its cells hold: a font
    name past Excel's 31 characters in the styles (measured: the whole file
    refuses) and, with macro_part, a macro part that is not a macro project
    (a compound file). One sentence each; [] when there is none, or when the
    package cannot be read at all (structural_check owns that verdict).

    The gate after every save asks for the macro-part rule too (Excel
    measured on a damaged file: the file and every edit of it fail to open),
    so an edit of a workbook that validate calls unopenable is refused with
    validate's own reason instead of being answered as verified."""
    out: list[str] = []
    try:
        with zipfile.ZipFile(path) as zf:
            names = set(zf.namelist())
            styles = _workbook_part(zf, names, "styles", "xl/styles.xml")
            if styles is not None:
                long_names: list[str] = []
                with zf.open(styles) as fh:
                    for _event, el in _xmlsafe.iterparse(
                            fh, doc=path, part=styles):
                        if el.tag.endswith("}name") or el.tag == "name":
                            val = el.get("val") or ""
                            if len(val) > _limits.MAX_FONT_NAME_CHARS and \
                                    val not in long_names:
                                long_names.append(val)
                for val in long_names[:_TYPE_SHOWN]:
                    out.append(
                        f"its styles name the font {val!r}, {len(val)} "
                        "characters long, and Excel allows "
                        f"{_limits.MAX_FONT_NAME_CHARS}")
            out.extend(_query_table_orphans(zf, names))
            out.extend(_workbook_type_mismatch(zf, names, path))
            for name in sorted(names) if macro_part else ():
                if name.lower().endswith("vbaproject.bin"):
                    with zf.open(name) as fh:
                        head = fh.read(len(_CFB_MAGIC))
                    if head != _CFB_MAGIC:
                        out.append(f"its macro part {name} is not a macro "
                                   "project (not a compound file)")
    except Exception:  # noqa: BLE001 - structural_check reports bad packages
        return out
    return out


def cell_value_check(path: str) -> tuple[bool, list[str]]:
    """Every stored cell value read against its cell's type (t=). A number
    cell holding text, an error cell holding something that is not an
    error, a TRUE/FALSE cell holding anything but 0/1/true/false, a
    shared-string cell pointing past the shared strings, or a formula past
    8,192 stored characters: Excel refuses to open any of them, so a
    produced file holding one must never be answered verified. This failed
    whatever its age: no save can promise a file Excel will open while the
    file keeps such a cell."""
    found: list[str] = []
    ok, reasons = structural_check(path, cell_values=found)
    if not ok:
        return False, reasons
    return (not found), found


def part_loss_check(pre_parts, path: str, *, allow_loss: bool,
                    expected_removals: Iterable[str] = (),
                    loss_keys: Iterable[str] | None = None,
                    ) -> tuple[bool, list[str]]:
    """Diff the produced part list against the pre-write scan. A fragile part
    that vanished unexpectedly (and was not allow_loss-covered) fails.

    expected_removals extends excuse (d) of the module contract to FRAGILE
    parts: an op that deliberately removes an object at the model level
    (manage_chart delete drops its xl/charts/ part and anchor drawing;
    manage_image delete drops xl/media/ and its drawing) registers the
    removal via WorkbookPackage.expect_removal, and the registered prefixes
    excuse exactly those losses here. Without this, a deliberate chart or
    image delete could never pass verify (OBSERVED: removing a model chart
    drops xl/charts/chart1.xml plus xl/drawings/drawing1.xml on save).
    Anything not covered by a registered prefix still fails by default.

    loss_keys SCOPES the allow_loss excuse to the hazard families the caller
    was actually warned about (the SEV_DROPS keys the gate named when it let
    the mutation through). Without it, allow_loss was a blanket amnesty for
    every fragile family: a caller who accepted the loss of slicers also,
    silently, accepted a chart part vanishing whole, even though the gate had
    only told them charts might DEGRADE. None keeps the old behavior for
    direct callers that cannot know the warned set."""
    try:
        with zipfile.ZipFile(path) as zf:
            post = set(zf.namelist())
    except Exception:  # noqa: BLE001
        return False, ["could not list produced package parts"]
    pre = set(pre_parts)
    lost = pre - post
    exp = tuple(e.lower() for e in expected_removals)
    keys = None if loss_keys is None else {k.lower() for k in loss_keys}
    fragile_lost: list[str] = []
    unexcused: list[str] = []
    for name in sorted(lost):
        low = name.lower()
        if low in EXPECTED_DROPPABLE:
            continue
        if any(low.startswith(e) for e in exp):
            continue  # a registered deliberate removal (excuse d)
        for spec in _hazard.HAZARD_SPECS:
            if spec.matcher(name):
                fragile_lost.append(name)
                if keys is not None and spec.key not in keys:
                    unexcused.append(f"{name} ({spec.label})")
                break
    if fragile_lost and not allow_loss:
        return False, fragile_lost
    if allow_loss and unexcused:
        return False, unexcused
    return True, fragile_lost if allow_loss else []


def _is_fragile(name: str) -> bool:
    return any(spec.matcher(name) for spec in _hazard.HAZARD_SPECS)


def part_inventory_check(pre_sizes, path: str, *,
                         expected_removals: Iterable[str] = (),
                         ) -> tuple[bool, list[str]]:
    """The DEFAULT-FAIL part-inventory diff: any part present before the save
    and missing or empty afterward fails unless its loss is explained.

    Jurisdiction: NON-fragile parts only. Fragile parts (hazard table) are
    part_loss_check's and part_content_check's job, where allow_loss governs;
    allow_loss deliberately does NOT reach this check, because the user who
    accepted the loss of a NAMED fragile family did not accept unrelated
    surprise losses.

    Excuses, in order: the legitimate-rewrite pass-list (_PASS_EXACT /
    _PASS_PREFIX), regenerated relationship parts (_is_rels_part), an
    expected_removals prefix registered by the mutating op for a deliberate
    model-level delete, and renumbered families (judged by count, below).

    pre_sizes maps part name to uncompressed size at open time; sizes of -1
    mean 'name known, size unknown' (the emptied check skips those). An empty
    mapping checks nothing and passes."""
    if not pre_sizes:
        return True, []
    try:
        with zipfile.ZipFile(path) as zf:
            post = {i.filename: i.file_size for i in zf.infolist()}
    except Exception:  # noqa: BLE001
        return False, ["could not re-read produced package parts"]

    exp = tuple(e.lower() for e in expected_removals)

    def _expected(low: str) -> bool:
        return any(low.startswith(e) for e in exp)

    problems: list[str] = []
    missing = [n for n in pre_sizes if n not in post]

    # --- missing parts (name-based, families deferred to the count stage) ---
    for name in sorted(missing):
        low = name.lower()
        if _is_fragile(name):
            continue  # part_loss_check's jurisdiction (allow_loss governs)
        if _passlisted(low) or _is_rels_part(low) or _expected(low):
            continue
        if _family_label(low) is not None:
            continue  # judged by family count below
        problems.append(
            f"{name} was present before the save and is missing after it, "
            "and its loss is not on the legitimate-rewrite list")

    # --- renumbered families: count comparison ------------------------------
    for label, rex, _obs in _RENUMBERED_FAMILIES:
        pre_n = sum(1 for n in pre_sizes if rex.match(n.lower()))
        if not pre_n:
            continue
        post_n = sum(1 for n in post if rex.match(n.lower()))
        if post_n >= pre_n:
            continue  # renumber or growth, never a loss
        lost = [n for n in missing if rex.match(n.lower())]
        if lost and all(_expected(n.lower()) for n in lost):
            continue  # a registered deliberate removal explains the deficit
        problems.append(
            f"{pre_n - post_n} {label} part(s) missing after the save "
            f"({pre_n} before, {post_n} after; lost: {', '.join(sorted(lost))})")

    # --- parts shrunk to empty ---------------------------------------------
    for name, pre_size in sorted(pre_sizes.items()):
        if name not in post or not isinstance(pre_size, int) or pre_size <= 0:
            continue
        if post[name] != 0:
            continue
        low = name.lower()
        if _is_fragile(name):
            continue  # part_content_check reports fragile replacement
        if _passlisted(low) or _is_rels_part(low) or _expected(low):
            continue
        problems.append(
            f"{name} was replaced with empty content "
            f"(was {pre_size} bytes)")

    return (not problems), problems


def part_content_check(pre_sizes, path: str) -> tuple[bool, list[str]]:
    """Catch loss by REPLACEMENT, which a presence diff cannot see: a fragile
    part that still exists but was written back empty when it had content
    before, or whose XML no longer parses. pre_sizes maps part name to its
    uncompressed size at open time (hazard scan sizes); an empty mapping
    (no pre-scan sizes available) checks nothing and passes."""
    if not pre_sizes:
        return True, []
    problems: list[str] = []
    try:
        with zipfile.ZipFile(path) as zf:
            post = {i.filename: i.file_size for i in zf.infolist()}
            for name, pre_size in sorted(pre_sizes.items()):
                if name not in post or not _is_fragile(name):
                    continue  # absence is part_loss_check's job
                post_size = post[name]
                if pre_size > 0 and post_size == 0:
                    problems.append(
                        f"{name} was replaced with empty content "
                        f"(was {pre_size} bytes)")
                    continue
                if (name.lower().endswith(".xml")
                        and 0 < post_size <= _XML_RECHECK_MAX_BYTES):
                    try:
                        _xmlsafe.fromstring(zf.read(name), doc=path, part=name)
                    except Exception:  # noqa: BLE001
                        problems.append(
                            f"{name} is no longer well-formed XML")
    except Exception:  # noqa: BLE001
        return False, ["could not re-read produced package parts"]
    return (not problems), problems


# ------------------------------------------------- relationship integrity

#: The relationships namespace in its transitional and strict spellings. An
#: attribute in this namespace (r:id, r:embed, r:link, r:pict ...) names a
#: relationship of the part it sits in.
_REL_NAMESPACES = (
    b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    b"http://purl.oclc.org/ooxml/officeDocument/relationships",
)
_XMLNS_RE = re.compile(rb'xmlns:([A-Za-z_][\w.\-]*)\s*=\s*"([^"]*)"')


def _rels_source(rels_name: str) -> str | None:
    """The part a .rels part describes ('' for the package itself)."""
    if rels_name == "_rels/.rels":
        return ""
    head, _sep, base = rels_name.rpartition("/_rels/")
    if not _sep or not base.endswith(".rels"):
        return None
    return (head + "/" if head else "") + base[:-5]


def package_problems(path: str, *, filter_out: set | None = None) -> set[str]:
    """Broken links inside a package, each as one sentence: a relationship
    whose target part is missing, an r:id (or r:embed, r:link ...) that names
    no relationship of its part, and a part with no content type. Excel
    repairs or refuses a file with any of these, and none of them shows up
    in a part inventory or a cell read-back (an external link's
    relationship vanished, the save said verified, Excel refused the file).

    The check reads every XML part's relationship-namespace attributes, so
    it is linear in package size. filter_out, when given, also receives the
    table-filter problems among them (integrity_check fails on those
    whatever their age)."""
    import posixpath
    from . import surgical as _surgical

    problems: set[str] = set()
    try:
        zf = zipfile.ZipFile(path)
    except Exception:  # noqa: BLE001
        return {"the package could not be opened"}
    with zf:
        names = [i.filename for i in zf.infolist()]
        lower = {n.lower() for n in names}
        rel_ids: dict[str, set[str]] = {}
        for name in names:
            if not name.lower().endswith(".rels"):
                continue
            source = _rels_source(name)
            if source is None:
                continue
            try:
                rels = _surgical.parse_rels(zf.read(name))
            except Exception:  # noqa: BLE001
                problems.add(f"{name} cannot be read")
                continue
            ids = rel_ids.setdefault(source.lower(), set())
            for r in rels:
                rid = r.get("Id")
                if rid:
                    ids.add(rid)
                if r.get("TargetMode") == "External":
                    continue
                target = r.get("Target", "")
                if not target or target.startswith("#"):
                    continue
                part = _surgical.resolve_target(source, target)
                if part.lower() not in lower and \
                        posixpath.normpath(part).lower() not in lower:
                    problems.add(
                        f"{name} points relationship {rid} at {part}, which "
                        "is not in the package")
        for name in names:
            low = name.lower()
            if low.endswith(".rels") or not (low.endswith(".xml")
                                             or low.endswith(".vml")):
                continue
            try:
                data = zf.read(name)
            except Exception:  # noqa: BLE001
                continue
            if not any(ns in data for ns in _REL_NAMESPACES):
                continue
            prefixes = {m.group(1) for m in _XMLNS_RE.finditer(data)
                        if m.group(2) in _REL_NAMESPACES}
            if not prefixes:
                continue
            have = rel_ids.get(low, set())
            for pfx in prefixes:
                rx = re.compile(rb"[\s<]" + re.escape(pfx)
                                + rb':([A-Za-z]+)\s*=\s*"([^"]*)"')
                for m in rx.finditer(data):
                    rid = m.group(2).decode("utf-8", "replace")
                    if rid and rid not in have:
                        problems.add(
                            f"{name} refers to relationship {rid}, which "
                            "its relationship part does not define")
        # content types
        try:
            ct = zf.read("[Content_Types].xml")
        except KeyError:
            problems.add("the package has no [Content_Types].xml")
            return problems
        defaults = {m.group(1).lower() for m in re.finditer(
            rb'<(?:[\w.\-]+:)?Default\b[^>]*?Extension="([^"]*)"', ct)}
        overrides = {m.group(1).decode("utf-8", "replace").lstrip("/").lower()
                     for m in re.finditer(
                         rb'<(?:[\w.\-]+:)?Override\b[^>]*?PartName="([^"]*)"', ct)}
        for name in names:
            if name == "[Content_Types].xml" or name.endswith("/"):
                continue
            low = name.lower()
            ext = low.rsplit(".", 1)[-1].encode() if "." in low.rsplit("/", 1)[-1] else b""
            if low not in overrides and ext not in defaults:
                problems.add(f"{name} has no content type")
        from . import tablefilter as _tablefilter
        damage = _tablefilter.scan_zip(zf, names)
        if filter_out is not None:
            filter_out |= damage.fatal_sentences()
        problems |= damage.sentences()
    return problems


def filter_table_problems(zf: zipfile.ZipFile, names: list[str]) -> set[str]:
    """A sheet filter over a table's range, or an
    _xlnm._FilterDatabase name pointing into a table, makes a workbook
    Excel will not open at all ("Open method of Workbooks class failed").
    A table is filtered by its own autofilter, in the table part. set_filter
    used to write the broken shape for a table and the verify gate passed
    it; it is now a problem every save is checked for, and one the gate
    fails on whatever its age (integrity_check). The scan
    itself lives in core.tablefilter, which also repairs it."""
    from . import tablefilter as _tablefilter
    return _tablefilter.scan_zip(zf, names).sentences()


def integrity_check(path: str, baseline: set[str] | None) -> tuple[bool, list[str]]:
    """Fail on any broken link the produced package has that the original
    did not (a file that arrived broken is not this save's doing).

    One exception: a sheet filter over a table fails the
    check even when the original already carried it. Excel will not open
    such a file, every file-tier save repairs the old damage before this
    check runs (core.tablefilter), and answering "verified" for a workbook
    Excel refuses is the silent failure this gate exists to stop. An
    _FilterDatabase name inside a table on its own still fails only when
    new: Excel opens it, and Excel itself keeps one when a filtered range
    becomes a table."""
    always: set[str] = set()
    found = package_problems(path, filter_out=always)
    new = sorted((found - (baseline or set())) | always)
    return (not new), new


def content_readback(path: str, intended: dict) -> tuple[bool, list[dict]]:
    """Re-read the cells the tool claimed to write and compare against intent.

    intended maps (sheet, "A1") -> ("value", literal) or ("formula", string).
    Formula strings are compared after normalizing a leading '='.
    """
    if not intended:
        return True, []
    mismatches: list[dict] = []
    keep_vba = _xmlsafe.keeps_vba(path)
    try:
        wb = _xmlsafe.load_workbook(path, keep_vba=keep_vba, data_only=False)
    except Exception as exc:  # noqa: BLE001
        # A produced package can be broken in ways that make the reader throw
        # rather than answer (a dangling chart relationship raises KeyError
        # deep inside openpyxl's reader). That is a failed read-back, not an
        # excuse to escape the gate.
        return False, [{"problem": "the produced file could not be re-opened "
                                   "for read-back",
                        "error": f"{type(exc).__name__}: {exc}"}]
    try:
        for (sheet, coord), (kind, expected) in intended.items():
            if sheet not in wb.sheetnames:
                mismatches.append({"sheet": sheet, "cell": coord,
                                   "problem": "sheet missing after write"})
                continue
            got = wb[sheet][coord].value
            if kind == "formula":
                exp = expected if str(expected).startswith("=") else "=" + str(expected)
                # An ARRAY formula's value is an openpyxl ArrayFormula OBJECT
                # carrying the text on .text; str() of it is the repr, which
                # made every array write fail read-back as a "mismatch".
                text = getattr(got, "text", None)
                if isinstance(text, str):
                    got = text
                gots = got if isinstance(got, str) else str(got)
                if not gots.startswith("="):
                    gots = "=" + gots
                if gots != exp:
                    mismatches.append({"sheet": sheet, "cell": coord,
                                       "expected": exp, "got": got})
            else:
                # An empty string and an empty cell are the same "blank" in
                # Excel: openpyxl serializes value "" as an absent cell, which
                # reads back as None. Treating that as a mismatch made it
                # IMPOSSIBLE to write a blank string through any write tool --
                # the safety net rejected a write that had in fact done exactly
                # the right thing.
                if (expected == "" and got is None) or \
                        (expected is None and got == ""):
                    continue
                if _serial_matches(expected, got, wb):
                    # A number written into a date-formatted
                    # cell reads back as a date; it is the same serial.
                    continue
                if got != expected:
                    entry = {"sheet": sheet, "cell": coord,
                             "expected": expected, "got": got}
                    # The 1-ulp precision class: a
                    # double that needs 17 significant digits does not
                    # survive openpyxl's %.16g float writer, so the readback
                    # is exactly the 16-digit snap of the intent. Name the
                    # cause so the refusal reads as a stated limit of the
                    # file tier, not a server bug.
                    snapped = _calc.full_precision_loss(expected)
                    if snapped is not None and got == snapped:
                        entry["cause"] = "precision"
                    mismatches.append(entry)
    finally:
        wb.close()
    return (not mismatches), mismatches


def _serial_matches(expected, got, wb) -> bool:
    """True when a NUMBER was written and the cell, formatted as a date or
    time, reads back as the date or time with that serial (such
    a write was refused as "did not read back as intended", though Excel
    stores and shows exactly what was asked)."""
    import datetime as _dt
    if isinstance(expected, bool) or not isinstance(expected, (int, float)):
        return False
    if not isinstance(got, (_dt.datetime, _dt.date, _dt.time, _dt.timedelta)):
        return False
    try:
        from openpyxl.utils.datetime import to_excel
        serial = to_excel(got, getattr(wb, "epoch", None) or _dt.datetime(1899, 12, 30))
    except Exception:  # noqa: BLE001
        return False
    if serial is None:
        return False
    want = float(expected)
    if isinstance(got, _dt.time):
        want = want % 1  # a time-only format reads back the fraction
    return abs(float(serial) - want) < 1e-8


def verify_after_write(path: str, *, pre_parts, intended: dict | None = None,
                       allow_loss: bool = False,
                       pre_sizes: dict | None = None,
                       expected_removals: Iterable[str] = (),
                       loss_keys: Iterable[str] | None = None,
                       integrity_baseline: set[str] | None = None,
                       ) -> VerifyResult:
    """Run the full verify gate on a produced (temp) package. Returns a
    VerifyResult; the caller raises ValidationFailed and refuses to promote on
    ``ok is False``."""
    result = VerifyResult(ok=True)

    contradictions: list[str] = []
    ok, reasons = structural_check(path, cell_values=contradictions)
    if not ok:
        result.ok = False
        result.reasons.extend(reasons)
        return result  # a structurally broken file is fatal; stop here
    if contradictions:
        # A stored value that contradicts its cell's type, or a formula past
        # Excel's length limit: Excel cannot open the file, so it is never
        # answered verified (the always-calculate cache once went back
        # without its t="str" and the save said verified on a file Excel
        # refused).
        result.ok = False
        result.reasons.append(
            "Excel would refuse to open the produced file, because of what "
            "these cells store: " + "; ".join(contradictions)
            + ". Write those cells again (or clear them) to repair them")
    for why in package_refusals(path, macro_part=True):
        # whatever its age: the file does not open in Excel
        result.ok = False
        result.reasons.append(
            f"Excel would refuse to open the produced file: {why}")

    ok, lost = part_loss_check(pre_parts, path, allow_loss=allow_loss,
                               expected_removals=expected_removals,
                               loss_keys=loss_keys)
    if not ok:
        result.ok = False
        result.lost_parts = lost
        result.reasons.append(
            "the write would drop fragile part(s) that were present before"
            + (" and were NOT among the families allow_loss accepted"
               if allow_loss else "")
            + ": " + ", ".join(lost))

    # default-fail inventory diff over everything the fragile checks do not
    # own; falls back to a name-only inventory when no pre-scan sizes exist.
    inventory = pre_sizes or {n: -1 for n in pre_parts}
    ok, unexplained = part_inventory_check(
        inventory, path, expected_removals=expected_removals)
    if not ok:
        result.ok = False
        result.reasons.append(
            "unexplained part loss (refusing by default): "
            + "; ".join(unexplained))

    if integrity_baseline is not None:
        ok, broken = integrity_check(path, integrity_baseline)
        if not ok:
            result.ok = False
            shown = "; ".join(broken[:8])
            more = f" (+{len(broken) - 8} more)" if len(broken) > 8 else ""
            result.reasons.append(
                "the produced file has broken internal links Excel would "
                f"repair or refuse: {shown}{more}")

    ok, replaced = part_content_check(pre_sizes or {}, path)
    if not ok:
        result.ok = False
        result.reasons.append(
            "fragile part(s) were damaged in place: " + "; ".join(replaced))

    if intended:
        ok, mism = content_readback(path, intended)
        if not ok:
            result.ok = False
            result.mismatches = mism
            result.reasons.append(_readback_reason(mism))

    return result


def _readback_reason(mismatches: list[dict]) -> str:
    """The read-back failure line, naming the cells and -- for the 1-ulp
    class -- the cause. The old text ('N written cell(s) did not read back as
    intended') named neither, so a full-precision refusal read as a server
    bug rather than a stated storage limit."""
    def _addr(m: dict) -> str:
        s, c = m.get("sheet"), m.get("cell")
        return f"{s}!{c}" if s and c else str(c or m.get("problem", "?"))

    precision = [m for m in mismatches if m.get("cause") == "precision"]
    other = [m for m in mismatches if m.get("cause") != "precision"]
    parts: list[str] = []
    if other:
        shown = ", ".join(_addr(m) for m in other[:10])
        more = f" (+{len(other) - 10} more)" if len(other) > 10 else ""
        parts.append(
            f"{len(other)} written cell(s) did not read back as intended: "
            f"{shown}{more}")
    if precision:
        shown = ", ".join(
            f"{_addr(m)} ({m.get('expected')!r})" for m in precision[:10])
        more = f" (+{len(precision) - 10} more)" \
            if len(precision) > 10 else ""
        parts.append(
            f"{len(precision)} written cell(s) hold double(s) that need 17 "
            "significant digits, which the full rewrite this edit needed "
            f"stores at 16 (a 1-ulp change): {shown}{more}. This is a "
            "storage limit, not a data error: round the value (e.g. to 15 "
            "significant digits) if the last digit is noise, or write it in "
            "a call of its own, without a structural change in the same "
            "call; a cell edit on its own stores all 17 digits")
    return "; ".join(parts)


__all__ = [
    "VerifyResult", "verify_after_write", "structural_check",
    "part_loss_check", "part_inventory_check", "part_content_check",
    "content_readback", "EXPECTED_DROPPABLE", "package_problems",
    "integrity_check", "cell_value_check", "package_refusals",
]
