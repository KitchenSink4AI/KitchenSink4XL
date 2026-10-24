"""core/tablefilter.py: the table-filter damage earlier versions wrote, found
and repaired.

Earlier versions wrote a SHEET filter over the table's range, with an
_xlnm._FilterDatabase name beside it, when set_filter was given a table (its
name, its range or a cell in it). Excel will not open such a workbook ("Open
method of Workbooks class failed"; measured: the sheet filter alone does it,
the name alone does not). 1.2.6 and before write the same shape. set_filter
no longer writes it and clear_filter repairs it, but every other edit of a
damaged workbook used to save it as it was and answer ok and verified.

Now every file-tier save repairs, in its payload and before the verify gate,
the damage the workbook ALREADY carried (core.package), and the gate fails
closed on any of it that is left, whatever its age (core.verify). Damage a
save would ADD is never repaired here, so a route that wrote it again still
fails closed. The repair:

- a sheet filter over a table goes. When it covers exactly the table's own
  filter range (header and data rows) and the table's filter holds no
  criteria of its own, its criteria move into the table's filter: the filter
  the user asked for, now stored where Excel reads it, over the same hidden
  rows. Otherwise they are dropped and the rows stay as they were;
- the _xlnm._FilterDatabase name that the same versions wrote beside it, the
  one pointing into the same table, goes with it. A name like that on its
  own is left alone: Excel opens it, and Excel itself keeps one when a
  filtered range becomes a table.

Raw-zip surgery on the payload, no model load, the same way core.calc
rewrites a payload's parts.
"""

from __future__ import annotations

import os
import re
import shutil
import zipfile
from dataclasses import dataclass, field

_SHEET_ENTRY_RE = re.compile(rb'<(?:[\w.\-]+:)?sheet\b[^>]*>')
_RID_RE = re.compile(rb'\b(?:[\w.\-]+:)?id="([^"]+)"')
_NAME_ATTR_RE = re.compile(rb'\bname="([^"]*)"')
_TABLE_REF_RE = re.compile(rb'<(?:[\w.\-]+:)?table\b[^>]*?\bref="([^"]+)"')
_TABLE_NAME_RE = re.compile(rb'\bdisplayName="([^"]+)"')
_TABLE_TOTALS_RE = re.compile(rb'<(?:[\w.\-]+:)?table\b[^>]*?\btotalsRowCount="(\d+)"')
#: The whole sheet-level autoFilter element (it never nests another).
_AUTOFILTER_EL_RE = re.compile(
    rb'<((?:[\w.\-]+:)?)autoFilter\b([^>]*?)(?:/>|>(.*?)</\1autoFilter\s*>)', re.S)
_REF_ATTR_RE = re.compile(rb'\bref="([^"]+)"')
_FILTER_DB_EL_RE = re.compile(
    rb'<((?:[\w.\-]+:)?)definedName\b([^>]*\bname="_xlnm\._FilterDatabase"[^>]*)>'
    rb'([^<]*)</\1definedName\s*>')
_LOCAL_SHEET_RE = re.compile(rb'\blocalSheetId="(\d+)"')
_EMPTY_NAMES_RE = re.compile(rb'<((?:[\w.\-]+:)?)definedNames\b[^>]*>\s*</\1definedNames\s*>')


def _ref_rect(text: str):
    """(min_col, min_row, max_col, max_row) of an A1 range, or None."""
    from openpyxl.utils import range_boundaries
    try:
        ref = text.rsplit("!", 1)[-1].replace("$", "").strip()
        return range_boundaries(ref) if ref else None
    except Exception:  # noqa: BLE001
        return None


def _meet(a, b) -> bool:
    return not (a[2] < b[0] or b[2] < a[0] or a[3] < b[1] or b[3] < a[1])


@dataclass
class Table:
    name: str
    part: str
    ref: str
    rect: tuple
    totals: int


@dataclass
class SheetFilter:
    """A sheet filter over one or more tables of a sheet."""
    title: str
    part: str
    ref: str
    span: tuple          # (start, end) of the element in the sheet part
    prefix: bytes
    inner: bytes         # the element's children (filterColumn, sortState)
    tables: list = field(default_factory=list)


@dataclass
class FilterName:
    """An _xlnm._FilterDatabase name pointing into a table."""
    title: str
    ref: str
    span: tuple          # (start, end) of the element in workbook.xml
    tables: list = field(default_factory=list)


@dataclass
class Damage:
    filters: list = field(default_factory=list)
    names: list = field(default_factory=list)

    def keys(self) -> set:
        """Stable keys of each piece: (kind, table name). Table names are
        unique in a workbook and survive a row insert or a sheet rename."""
        out = set()
        for f in self.filters:
            out |= {("sheet_filter", t.name) for t in f.tables}
        for n in self.names:
            out |= {("filter_name", t.name) for t in n.tables}
        return out

    def sentences(self) -> set:
        return self.fatal_sentences() | self.name_sentences()

    def fatal_sentences(self) -> set:
        """The sheet filters over tables: what makes Excel refuse the file."""
        out = set()
        for f in self.filters:
            for t in f.tables:
                out.add(
                    f"sheet {f.title!r} has a sheet filter ({f.ref}) over table "
                    f"{t.name} ({t.ref}); Excel will not open a workbook with a "
                    "sheet filter over a table (a table is filtered by its own "
                    f"filter); clear_filter on table {t.name} removes it")
        return out

    def name_sentences(self) -> set:
        """_FilterDatabase names inside tables. Excel opens a workbook with
        one (it keeps one itself when a filtered range becomes a table);
        it is what a sheet filter over a table writes beside it."""
        out = set()
        for n in self.names:
            for t in n.tables:
                out.add(
                    f"the _xlnm._FilterDatabase name of sheet {n.title!r} points into "
                    f"table {t.name} ({t.ref}), where only the table's own filter "
                    "belongs")
        return out


def _resolve(base: str, target: str) -> str:
    from . import surgical as _surgical
    return _surgical.resolve_target(base, target)


def _parse_rels(data: bytes):
    from . import surgical as _surgical
    return _surgical.parse_rels(data)


def scan_zip(zf: zipfile.ZipFile, names: list[str] | None = None) -> Damage:
    """Every sheet filter over a table and every _FilterDatabase name inside
    one, in an open package."""
    names = names if names is not None else zf.namelist()
    damage = Damage()
    lower = {n.lower(): n for n in names}
    try:
        wb_xml = zf.read(lower.get("xl/workbook.xml", "xl/workbook.xml"))
        wb_rels = _parse_rels(zf.read(
            lower.get("xl/_rels/workbook.xml.rels", "xl/_rels/workbook.xml.rels")))
    except Exception:  # noqa: BLE001
        return damage
    targets = {r.get("Id"): _resolve("xl/workbook.xml", r.get("Target", ""))
               for r in wb_rels if r.get("Id")}
    sheets = []  # (sheet part, display name) in workbook order
    for m in _SHEET_ENTRY_RE.finditer(wb_xml):
        tag = m.group(0)
        rid = _RID_RE.findall(tag)
        nm = _NAME_ATTR_RE.search(tag)
        part = targets.get(rid[-1].decode("utf-8", "replace")) if rid else None
        sheets.append((part, nm.group(1).decode("utf-8", "replace") if nm else "?"))

    def tables_of(part):
        out = []
        rels_name = part.rsplit("/", 1)[0] + "/_rels/" + part.rsplit("/", 1)[-1] + ".rels"
        real = lower.get(rels_name.lower())
        if real is None:
            return out
        try:
            rels = _parse_rels(zf.read(real))
        except Exception:  # noqa: BLE001
            return out
        for r in rels:
            if not str(r.get("Type", "")).endswith("/table") or r.get("TargetMode") == "External":
                continue
            tpart = lower.get(_resolve(part, r.get("Target", "")).lower())
            if tpart is None:
                continue
            try:
                data = zf.read(tpart)
            except Exception:  # noqa: BLE001
                continue
            ref = _TABLE_REF_RE.search(data)
            name = _TABLE_NAME_RE.search(data)
            totals = _TABLE_TOTALS_RE.search(data)
            rect = _ref_rect(ref.group(1).decode("utf-8", "replace")) if ref else None
            if rect is not None:
                out.append(Table(
                    name=name.group(1).decode("utf-8", "replace") if name else tpart,
                    part=tpart, ref=ref.group(1).decode("utf-8", "replace"), rect=rect,
                    totals=int(totals.group(1)) if totals else 0))
        return out

    tables_by_index = {}
    for i, (part, title) in enumerate(sheets):
        if not part or part.lower() not in lower:
            continue
        real = lower[part.lower()]
        tables = tables_of(real)
        tables_by_index[i] = (title, tables)
        if not tables:
            continue
        try:
            data = zf.read(real)
        except Exception:  # noqa: BLE001
            continue
        # The sheet's own filter sits after its cell data; a custom view's
        # filter (inside customSheetViews) is not the sheet filter.
        end = data.find(b"</sheetData>")
        start = end if end >= 0 else 0
        stop = data.find(b"customSheetView", start)
        region = data[start:stop] if stop >= 0 else data[start:]
        m = _AUTOFILTER_EL_RE.search(region)
        if m is None:
            continue
        ref_m = _REF_ATTR_RE.search(m.group(2))
        if ref_m is None:
            continue
        ref = ref_m.group(1).decode("utf-8", "replace")
        rect = _ref_rect(ref)
        hit = [t for t in tables if rect is not None and _meet(rect, t.rect)]
        if hit:
            damage.filters.append(SheetFilter(
                title=title, part=real, ref=ref,
                span=(start + m.start(), start + m.end()), prefix=m.group(1),
                inner=m.group(3) or b"", tables=hit))
    for m in _FILTER_DB_EL_RE.finditer(wb_xml):
        local = _LOCAL_SHEET_RE.search(m.group(2))
        if local is None:
            continue
        title, tables = tables_by_index.get(int(local.group(1)), ("?", []))
        ref = m.group(3).decode("utf-8", "replace")
        rect = _ref_rect(ref)
        hit = [t for t in tables if rect is not None and _meet(rect, t.rect)]
        if hit:
            damage.names.append(FilterName(title=title, ref=ref,
                                           span=(m.start(), m.end()), tables=hit))
    return damage


def scan(path: str) -> Damage:
    try:
        with zipfile.ZipFile(path) as zf:
            return scan_zip(zf)
    except Exception:  # noqa: BLE001
        return Damage()


def _table_filter_rect(t: Table) -> tuple:
    c0, r0, c1, r1 = t.rect
    return c0, r0, c1, r1 - t.totals


def _move_into_table(table_xml: bytes, sheet_prefix: bytes, inner: bytes):
    """The table part with `inner` (filterColumn / sortState children) as the
    content of its own autoFilter, or None when that cannot be done safely:
    no autoFilter of its own, criteria of its own already, or a different
    namespace prefix (the moved elements would change meaning)."""
    m = _AUTOFILTER_EL_RE.search(table_xml)
    if m is None or m.group(1) != sheet_prefix:
        return None
    if (m.group(3) or b"").strip():
        return None
    p = m.group(1)
    new = b"<" + p + b"autoFilter" + m.group(2).rstrip(b"/").rstrip() + b">" + \
        inner + b"</" + p + b"autoFilter>"
    return table_xml[:m.start()] + new + table_xml[m.end():]


def repair_package(path: str, allowed: set | None = None) -> list[str]:
    """Repair the damage in the package at `path` in place and return one
    sentence per repair. Only damage whose key is in `allowed` is repaired
    (None: all of it); the save passes the keys of the damage the workbook
    carried before the edit, so damage an edit would add is never hidden."""
    try:
        with zipfile.ZipFile(path) as zf:
            damage = scan_zip(zf)
            if not damage.filters and not damage.names:
                return []
            listed = zf.namelist()
            lower = {n.lower(): n for n in listed}
            wanted = {f.part for f in damage.filters}
            wanted |= {t.part for f in damage.filters for t in f.tables}
            if "xl/workbook.xml" in lower:
                wanted.add(lower["xl/workbook.xml"])
            parts = {n: zf.read(n) for n in wanted}
    except Exception:  # noqa: BLE001
        return []
    said: list[str] = []
    edits: dict[str, list] = {}  # part -> [(start, end, replacement)]
    moved_tables: dict[str, bytes] = {}

    def ok(kind, tables):
        return allowed is None or all((kind, t.name) in allowed for t in tables)

    for f in damage.filters:
        if not ok("sheet_filter", f.tables):
            continue
        moved = None
        if len(f.tables) == 1 and b"filterColumn" in f.inner:
            t = f.tables[0]
            if _ref_rect(f.ref) == _table_filter_rect(t):
                cur = moved_tables.get(t.part, parts[t.part])
                moved = _move_into_table(cur, f.prefix, f.inner)
                if moved is not None:
                    moved_tables[t.part] = moved
        edits.setdefault(f.part, []).append((f.span[0], f.span[1], b""))
        names = ", ".join(t.name for t in f.tables)
        if moved is not None:
            said.append(
                f"{f.title}!{f.ref}: an earlier version of this server stored the "
                f"filter of table {names} as a sheet filter, which Excel will not "
                "open; it is now the table's own filter, with the same criteria "
                "and the same hidden rows")
        else:
            said.append(
                f"{f.title}!{f.ref}: removed a sheet filter over table {names} "
                "that an earlier version of this server wrote and Excel will not "
                "open; the rows it hid stay hidden (clear_filter on the table "
                "shows them)")
    wb_part = lower.get("xl/workbook.xml")
    filtered = {t.name for f in damage.filters if ok("sheet_filter", f.tables)
                for t in f.tables}
    for n in damage.names:
        if wb_part is None or not ok("filter_name", n.tables):
            continue
        if not {t.name for t in n.tables} & filtered:
            continue  # a name on its own: Excel opens it, and keeps such names itself
        edits.setdefault(wb_part, []).append((n.span[0], n.span[1], b""))
        said.append(
            f"removed the _xlnm._FilterDatabase name of sheet {n.title!r}, which "
            f"pointed into table {', '.join(t.name for t in n.tables)}")
    if not edits and not moved_tables:
        return []
    out = dict(moved_tables)
    for part, cuts in edits.items():
        data = parts[part]
        for start, end, rep in sorted(cuts, reverse=True):
            data = data[:start] + rep + data[end:]
        if part == wb_part:
            data = _EMPTY_NAMES_RE.sub(b"", data)
        out[part] = data
    from .safesave import scratch_file
    folder = os.path.dirname(os.path.abspath(path)) or "."
    fd, tmp = scratch_file(folder, os.path.splitext(path)[1])
    os.close(fd)
    try:
        with zipfile.ZipFile(path) as zin, \
                zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = out.get(item.filename)
                zout.writestr(item, data if data is not None else zin.read(item.filename))
        shutil.move(tmp, path)
    except BaseException:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise
    return said


def problems(path: str) -> set[str]:
    """The damage sentences of the package at `path` (empty when none)."""
    return scan(path).sentences()


__all__ = ["Damage", "scan", "scan_zip", "repair_package", "problems"]
