"""ops/view.py: the grid view + batch layer.

get_grid_view is the get_document_view analog for grids: a compact, token-
efficient projection of a sheet or range (dimensions, the TRUE used range, a
markdown table with A1 addressing, formula and merge markers, hazard summary),
so an agent can see the grid without a per-cell JSON dump.

apply_edits is the batched writer: many addressed edits validated as ONE batch
(every location resolved before anything is applied), then a single backup, a
single save, and a single verify-after-write through WorkbookPackage. A failure
anywhere in validation or verify leaves the file byte-for-byte unchanged, so a
batch is all-or-nothing.
"""

from __future__ import annotations

from typing import Any

from ..core import calc as _calc
from ..core import hazard as _hazard
from ..core import locate as _locate
from ..core.errors import XlMcpError
from ..core.package import WorkbookPackage
from ..core.sandbox import check_path
from . import gridio

DEFAULT_VIEW_ROWS = 50
DEFAULT_VIEW_COLS = 30
_MAX_VIEW_ROWS = 200
_MAX_VIEW_COLS = 100
_CELL_CLIP = 40

EDIT_OPS = ("set_value", "set_formula", "clear", "write_range")


# --------------------------------------------------------------- grid view


def _render_cell(v: Any) -> str:
    if v is None:
        return ""
    s = str(gridio.compact_value(v))
    if len(s) > _CELL_CLIP:
        s = s[: _CELL_CLIP - 1] + "\u2026"
    return s.replace("|", "\\|").replace("\n", " ")


def _not_shown(path: str, fws, r0: int, c0: int, r1: int, c1: int) -> dict:
    """What the table leaves out, so a reader knows it is
    there: the notes and threaded comments on the shown cells, and the
    charts and images on the sheet, each with the op that reads it."""
    out: dict[str, dict] = {}
    cells = [gridio.a1(r, c) for (r, c), cell in (getattr(fws, "_cells", None) or {}).items()
             if r0 <= r <= r1 and c0 <= c <= c1 and getattr(cell, "comment", None) is not None]
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

    from .annotations import threaded_comments
    for (title, ref), _th in threaded_comments(path).items():
        if title != fws.title or ref in cells:
            continue
        col, row = coordinate_from_string(ref)
        if r0 <= row <= r1 and c0 <= column_index_from_string(col) <= c1:
            cells.append(ref)
    if cells:
        out["comments"] = {"count": len(cells), "cells": sorted(cells)[:20],
                           "read_with": "xl_format manage_comment action='list'"}
    from .objects import _anchor_cell, chart_kind
    charts = [{"chart_type": chart_kind(ch), "at_cell": _anchor_cell(ch)}
              for ch in getattr(fws, "_charts", [])]
    if charts:
        out["charts"] = {"count": len(charts), "items": charts[:20],
                         "read_with": "xl_format manage_chart action='list'"}
    # Pictures from the drawing parts themselves: openpyxl only loads
    # a picture when Pillow is installed, so the model can be blind to all of
    # them. The same read counts what no op lists cell by cell.
    try:
        from ..core import pkginfo as _pkginfo
        objs = _pkginfo.drawing_objects(check_path(path, "open workbook"),
                                        fws.title).get(fws.title, [])
        extra = _pkginfo.hidden_objects(check_path(path, "open workbook")).get(fws.title, {})
    except Exception:  # noqa: BLE001
        objs, extra = None, {}
    if objs is None:
        images = [_anchor_cell(im) for im in getattr(fws, "_images", [])]
    else:
        images = [o["anchor"]["from"] for o in objs if o["kind"] == "picture"]
    if images:
        out["images"] = {"count": len(images), "at_cells": images[:20],
                         "read_with": "xl_format manage_image action='list'"}
    shapes = sum(extra.get(k, 0) for k in ("shapes", "shape_groups", "connectors"))
    if shapes:
        out["shapes"] = {"count": shapes,
                         "read_with": "not listed by any op; Excel shows them"}
    if extra.get("x14_validations"):
        out["validations"] = {"count": extra["x14_validations"],
                              "read_with": "xl_structure manage_data_validation action='list'"}
    if extra.get("sparkline_groups"):
        out["sparklines"] = {"count": extra["sparkline_groups"],
                             "read_with": "not listed by any op; Excel shows them"}
    if extra.get("slicer_lists"):
        out["slicers"] = {"count": extra["slicer_lists"],
                          "read_with": "not listed by any op; Excel shows them"}
    return out


def get_grid_view(path: str, location: Any = None, sheet: str | None = None,
                  max_rows: int = DEFAULT_VIEW_ROWS,
                  max_cols: int = DEFAULT_VIEW_COLS,
                  values: str = "cached") -> dict:
    """Compact, token-efficient projection of a sheet or range: the used range,
    a markdown table with A1 addressing, formula and merge markers, and the
    hazard summary. Reads only; paginates with max_rows / max_cols.

    location defaults to the sheet's TRUE used range. values='cached' shows last
    calculated values (formula cells marked) and adds `shows` for the shown
    cells whose text in Excel differs from their value; 'formula' shows formula
    strings. The returned dimensions and truncated flags say when the view was
    clipped so the caller can page; not_shown names the comments, charts and
    images the table leaves out, with the op that reads each. Nothing is
    written.
    """
    if values not in ("cached", "formula"):
        raise XlMcpError(
            f"get_grid_view takes values 'cached' or 'formula', not {values!r}; "
            "its cached view adds shows, the text Excel shows where it differs "
            "from the value. read_range also takes 'both', which labels each "
            "cell cached, absent or formula, and values='displayed', the text "
            "of every cell")
    max_rows = max(1, min(int(max_rows), _MAX_VIEW_ROWS))
    max_cols = max(1, min(int(max_cols), _MAX_VIEW_COLS))
    disp_wb = gridio.open_wb(path, data_only=(values == "cached"))
    formula_wb = gridio.open_wb(path, data_only=False)
    try:
        loc = location if location is not None else {
            "used_range": sheet if sheet is not None else True}
        grid = gridio.resolve(formula_wb, loc, default_sheet=sheet)
        rep = _hazard.scan_path(check_path(path, "scan workbook"))
        if grid.empty:
            return {"sheet": grid.sheet, "used_range": None, "empty": True,
                    "table": "", "dimensions": {"rows": 0, "cols": 0},
                    "hazards": rep.labels()}
        full_rows = grid.max_row - grid.min_row + 1
        full_cols = grid.max_col - grid.min_col + 1
        end_row = min(grid.max_row, grid.min_row + max_rows - 1)
        end_col = min(grid.max_col, grid.min_col + max_cols - 1)
        truncated_rows = end_row < grid.max_row
        truncated_cols = end_col < grid.max_col

        from openpyxl.utils import get_column_letter
        dws = disp_wb[grid.sheet]
        fws = formula_wb[grid.sheet]
        cols = list(range(grid.min_col, end_col + 1))
        header = "| |" + "|".join(get_column_letter(c) for c in cols) + "|"
        sep = "|---|" + "|".join("---" for _ in cols) + "|"
        lines = [header, sep]
        formula_cells: dict[str, str] = {}
        differs = gridio.ShownText(path, gridio.is_1904(dws)) \
            if values == "cached" else None
        for r in range(grid.min_row, end_row + 1):
            cells = []
            for c in cols:
                fv = fws.cell(r, c).value
                is_formula = _calc.is_formula_cell(fws.cell(r, c))
                if values == "formula" and is_formula:
                    # The form read_range returns.
                    text = _render_cell(_calc.denormalize_formula(fv)
                                        if isinstance(fv, str) else fv)
                else:
                    dv = dws.cell(r, c).value
                    text = _render_cell(dv)
                    if differs is not None and dv is not None:
                        differs.add(grid.sheet, r, c, dv, dws.cell(r, c).number_format)
                    if is_formula:
                        formula_cells[gridio.a1(r, c)] = (
                            _calc.denormalize_formula(fv) if isinstance(fv, str) else fv)
                        if not text:
                            text = "\u0192"  # florin: formula, no cached value
                cells.append(text)
            lines.append(f"| {r} |" + "|".join(cells) + "|")

        merges = []
        mc = getattr(dws, "merged_cells", None)
        if mc is not None:
            for mr in mc.ranges:
                if mr.min_row <= end_row and mr.min_col <= end_col:
                    merges.append(str(mr))
        # View-stable addressing: one self-contained token for the SHOWN
        # rectangle, fingerprinted over the raw (formula-load) values so a
        # later {"anchor": token} location refuses STALE_ANCHOR if the
        # region changed after this view. Cells inside stay plain A1.
        anchor = _locate.make_anchor(
            fws, grid.min_row, grid.min_col, end_row, end_col)
        hidden = _not_shown(path, fws, grid.min_row, grid.min_col, end_row, end_col)
        out = {
            "sheet": grid.sheet,
            "used_range": grid.a1,
            "anchor": anchor,
            "dimensions": {"rows": full_rows, "cols": full_cols},
            "shown": {"rows": end_row - grid.min_row + 1,
                      "cols": end_col - grid.min_col + 1},
            "truncated": {"rows": truncated_rows, "cols": truncated_cols},
            "value_mode": values,
            "table": "\n".join(lines),
            **(differs.fields() if differs is not None else {}),
            "formula_cells": formula_cells,
            "merged_cells": merges,
            "hazards": rep.labels(),
        }
        if hidden:
            out["not_shown"] = hidden
        out.update(gridio.broken_shared_fields(
            path, grid.sheet, rect=(grid.min_row, grid.min_col, end_row, end_col)))
        if values == "cached" and formula_cells:
            # The shown results may predate a recalculation
            # the workbook is still owed.
            pending = gridio.pending_recalc_note(path)
            if pending:
                out["cached_note"] = pending
        return out
    finally:
        disp_wb.close()
        formula_wb.close()


# ---------------------------------------------------------------- apply_edits


def apply_edits(path: str, edits: list, allow_loss: bool = False,
                backup: bool = True,
                verify_com: bool | None = None) -> dict:
    """Apply many addressed edits as ONE atomic batch (validate every edit
    first, then one backup + one save + one verify).

    edits is a list of {op, location, ...}: set_value {value}, set_formula
    {formula}, clear {what: contents|formats|all}, write_range {data: 2D}. Every
    location is resolved and every op validated BEFORE anything is written, so a
    bad edit anywhere refuses the whole batch and the file is left unchanged;
    verify-after-write then guards the single save. Formula edits are normalized
    and flag recalculation.
    """
    if not isinstance(edits, list) or not edits:
        raise XlMcpError("edits must be a non-empty list of edit objects")
    pkg = WorkbookPackage.open(path)

    # ---- validation pass: resolve every location, check every op. Nothing is
    #      mutated here, so any refusal leaves the workbook untouched.
    plan: list[tuple[dict, Any]] = []
    for i, edit in enumerate(edits):
        if not isinstance(edit, dict):
            raise XlMcpError(f"edit #{i} is not an object")
        op = edit.get("op")
        if op not in EDIT_OPS:
            raise XlMcpError(
                f"edit #{i}: op must be one of {EDIT_OPS}, got {op!r}")
        loc = edit.get("location", edit.get("at"))
        if loc is None:
            raise XlMcpError(f"edit #{i}: missing 'location'")
        grid = pkg.resolve(loc, default_sheet=edit.get("sheet"))
        if op == "set_value":
            if "value" not in edit:
                raise XlMcpError(f"edit #{i}: set_value needs 'value'")
            if not grid.is_single:
                raise XlMcpError(
                    f"edit #{i}: set_value needs a single cell ({grid.a1} is a "
                    "range); use write_range")
        elif op == "set_formula":
            f = edit.get("formula")
            if not isinstance(f, str) or not f.strip():
                raise XlMcpError(f"edit #{i}: set_formula needs 'formula'")
            if not grid.is_single:
                raise XlMcpError(
                    f"edit #{i}: set_formula needs a single cell")
        elif op == "clear":
            what = edit.get("what", "contents")
            if what not in ("contents", "formats", "all"):
                raise XlMcpError(
                    f"edit #{i}: what must be contents|formats|all")
        else:  # write_range
            data = edit.get("data")
            if not isinstance(data, list) or not data or not all(
                    isinstance(r, list) for r in data):
                raise XlMcpError(f"edit #{i}: write_range needs 2D 'data'")
            widths = {len(r) for r in data}
            if len(widths) > 1:
                # A ragged block was written row by row, so short rows left
                # whatever sat under them in place (write_range has the same
                # guard).
                raise XlMcpError(
                    f"edit #{i}: write_range rows have different lengths "
                    f"{sorted(widths)}; the block must be rectangular (pad "
                    "short rows with null to clear those cells)")
        plan.append((edit, grid))

    # ---- apply pass: every op is now known-valid; mutate the in-memory model.
    from openpyxl.styles import Alignment, Border, Font, PatternFill
    touched = 0
    for edit, grid in plan:
        op = edit["op"]
        ws = pkg.workbook[grid.sheet]
        if op == "set_value":
            pkg.set_cell(grid.sheet, gridio.a1(grid.min_row, grid.min_col),
                         edit["value"], from_caller=True)
            touched += 1
        elif op == "set_formula":
            pkg.set_formula(grid.sheet,
                            gridio.a1(grid.min_row, grid.min_col),
                            edit["formula"])
            touched += 1
        elif op == "clear":
            what = edit.get("what", "contents")
            for r in range(grid.min_row, grid.max_row + 1):
                for c in range(grid.min_col, grid.max_col + 1):
                    cell = ws.cell(r, c)
                    if what in ("contents", "all") and cell.value is not None:
                        cell.value = None
                        pkg._intended[(grid.sheet, gridio.a1(r, c))] = (
                            "value", None)
                        touched += 1
                    if what in ("formats", "all"):
                        cell.font = Font(); cell.fill = PatternFill()
                        cell.border = Border(); cell.alignment = Alignment()
                        cell.number_format = "General"
        else:  # write_range
            data = edit["data"]
            top, left = grid.min_row, grid.min_col
            for i2, row in enumerate(data):
                for j2, val in enumerate(row):
                    pkg.set_cell(grid.sheet, gridio.a1(top + i2, left + j2),
                                 val, from_caller=True)
                    touched += 1

    result = pkg.save(allow_loss=allow_loss, backup=backup,
                      verify_com=verify_com)
    result["changed"]["edits_applied"] = len(plan)
    result["changed"]["cells_touched"] = touched
    return result


__all__ = ["get_grid_view", "apply_edits", "EDIT_OPS"]
