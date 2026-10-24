"""ops/dataio.py: import and export.

import_data writes CSV / TSV / JSON into a sheet at an anchor. export_range
reads a range, table, or sheet out to CSV / TSV / JSON. With query_range these
are the interop layer.

import_data runs the formula-injection lint the security posture requires: a
cell whose text begins with =, +, -, or @ is data that a spreadsheet would
execute as a formula, so by default such cells are written as TEXT (never as a
live formula) and the count is reported; pass formulas=true to write them as
formulas deliberately. Truncation is loud, never silent: an import past the
write ceiling refuses rather than dropping rows.

export_range states which value mode produced the output (cached, formula, or
both), so a cell holding a formula with no cached value is never passed off as
blank. import_data routes its write through WorkbookPackage for the backup and
verify; export_range is read-only.
"""

from __future__ import annotations

import codecs as _codecs
import csv as _csv
import io as _io
import json as _json
import os
import re
from typing import Any

from openpyxl.utils import get_column_letter

from ..core import limits as _limits
from ..core import safesave as _safesave
from ..core.errors import CeilingExceeded, RangeOutOfBounds, XlMcpError
from ..core.outguard import guard_out_dir, guard_out_file
from ..core.package import WorkbookPackage
from ..core.sandbox import check_path
from . import gridio

_FORMATS = ("auto", "csv", "tsv", "json")
_EXPORT_FORMATS = ("csv", "tsv", "json")
_INJECTION = ("=", "+", "-", "@")


#: Byte-order marks, longest first (a UTF-32 mark starts like a UTF-16 one).
_BOMS = ((b"\xff\xfe\x00\x00", "utf-32"), (b"\x00\x00\xfe\xff", "utf-32"),
         (b"\xef\xbb\xbf", "utf-8-sig"), (b"\xff\xfe", "utf-16"),
         (b"\xfe\xff", "utf-16"))


def _decode_text_file(raw: bytes, encoding: str) -> tuple[str, str]:
    """(the file's text, the codec used). A byte-order mark decides: Excel's
    "Unicode Text" is UTF-16 with one, and "CSV UTF-8" carries the UTF-8
    mark. Without one the caller's encoding is used (utf-8 as utf-8-sig), and
    bytes that are not in it refuse in plain words, naming what to pass."""
    for mark, codec in _BOMS:
        if raw.startswith(mark):
            return raw.decode(codec), codec  # the codec drops the mark
    enc = "utf-8-sig" if _codecs.lookup(encoding).name == "utf-8" else encoding
    try:
        return raw.decode(enc), _codecs.lookup(encoding).name
    except UnicodeDecodeError as exc:
        raise XlMcpError(
            f"the file is not {encoding} text: byte 0x{raw[exc.start]:02x} at "
            f"position {exc.start:,} is not valid {encoding}. Pass the "
            "file's encoding, for example encoding='cp1252' (Windows "
            "Western), 'latin-1' or 'utf-16'") from None


def _detect_format(fmt: str, source_file: str | None) -> str:
    if fmt != "auto":
        return fmt
    if source_file:
        low = source_file.lower()
        if low.endswith(".json"):
            return "json"
        if low.endswith(".tsv"):
            return "tsv"
    return "csv"


_SNIFF_DELIMITERS = (",", ";", "\t", "|")


def _sniff_delimiter(text: str) -> str:
    """The delimiter a CSV's own lines use. A semicolon
    file (the usual European export, with decimal commas) was split on its
    decimal commas and returned ok. The delimiter that appears the same
    number of times (at least once) on the first lines wins; a comma wins a
    tie, as before."""
    lines = [ln for ln in text.splitlines()[:20] if ln.strip()][:10]
    if not lines:
        return ","
    best, best_n = ",", 0
    for d in _SNIFF_DELIMITERS:
        counts = []
        for ln in lines:
            try:
                counts.append(len(next(_csv.reader([ln], delimiter=d))) - 1)
            except (_csv.Error, StopIteration):
                counts.append(0)
        if counts and min(counts) > 0 and len(set(counts)) == 1 and counts[0] > best_n:
            best, best_n = d, counts[0]
    return best


def _parse(text: str, fmt: str, delimiter: str | None,
           header: bool, coerce: bool = True) -> list[list[Any]]:
    if fmt == "json":
        obj = _json.loads(text)
        if isinstance(obj, dict):
            obj = [obj]
        if not isinstance(obj, list):
            raise XlMcpError("JSON import expects a list of objects or rows")
        if obj and isinstance(obj[0], dict):
            cols: list[str] = []
            for rec in obj:
                for k in rec:
                    if k not in cols:
                        cols.append(k)
            rows = [[rec.get(c) for c in cols] for rec in obj]
            return [cols] + rows if header else rows
        return [list(r) if isinstance(r, list) else [r] for r in obj]
    delim = delimiter or ("\t" if fmt == "tsv" else ",")
    reader = _csv.reader(_io.StringIO(text), delimiter=delim)
    try:
        if not coerce:
            return [list(row) for row in reader]
        return [[_coerce(v) for v in row] for row in reader]
    except _csv.Error as exc:
        # The text is the caller's argument, so a CSV the
        # reader cannot take is a refusal of that argument, not a server
        # fault (it used to reach INTERNAL_ERROR, "the arguments are not the
        # cause"). The usual case is one field over the reader's 131,072-
        # character limit; an Excel cell holds at most 32,767.
        raise XlMcpError(
            f"the {fmt.upper()} data could not be read at line {reader.line_num}: {exc}. "
            "A single field may not exceed 131,072 characters here, and an Excel "
            "cell holds at most 32,767.") from None


def _coerce(v: str) -> Any:
    if v == "":
        return None
    try:
        if v.strip().lstrip("-").isdigit():
            return int(v)
        f = float(v)
    except ValueError:
        return v
    # float() accepts "nan"/"inf"/"infinity"; Excel has no such numbers and
    # a NaN cell value serializes as garbage, so those stay text.
    if f != f or f in (float("inf"), float("-inf")):
        return v
    return f


# ------------------------------------------------- import typing

#: A value that LOOKS like a number but is a code Excel's own text import
#: would keep whole: a leading zero (000123, ZIP 02134), E-notation text
#: (1E5, a part number), or more digits than a double holds (an 18-digit
#: ID). One such value makes its whole column text, so a column never
#: mixes 02134 (text) with 10001 (number).
_LEADING_ZERO = re.compile(r"[+-]?0\d+(\.\d+)?")
_E_NOTATION = re.compile(r"[+-]?\d+(\.\d+)?[eE][+-]?\d+")
_THOUSANDS = re.compile(r"[+-]?\d{1,3}(,\d{3})+(\.\d+)?")
_THOUSANDS_DC = re.compile(r"[+-]?\d{1,3}(\.\d{3})+(,\d+)?")
_DECIMAL_COMMA = re.compile(r"[+-]?\d+,\d+")
COLUMN_TYPES = ("auto", "text", "number", "date")


def _code_reason(v: str) -> str | None:
    t = v.strip()
    if _LEADING_ZERO.fullmatch(t):
        return "a leading zero"
    if _E_NOTATION.fullmatch(t):
        return "E-notation text"
    digits = sum(ch.isdigit() for ch in t)
    if digits > 15 and re.fullmatch(r"[+-]?[\d.,]+", t):
        return "more than 15 digits"
    return None


def _number(v: str, decimal_comma: bool):
    """A CSV field as a number, Excel's way, or the text unchanged."""
    t = v.strip()
    if decimal_comma:
        if _THOUSANDS_DC.fullmatch(t) or _DECIMAL_COMMA.fullmatch(t):
            t = t.replace(".", "").replace(",", ".")
    elif _THOUSANDS.fullmatch(t):
        t = t.replace(",", "")
    out = _coerce(t)
    return out if not isinstance(out, str) else v


def _column_key(key, names: list[str]) -> int:
    """A column_types key (header name, letter, or 1-based number) as an
    index into the imported columns."""
    if isinstance(key, int) or (isinstance(key, str) and key.isdigit()):
        i = int(key) - 1
        if 0 <= i < len(names):
            return i
    elif isinstance(key, str):
        for i, n in enumerate(names):
            if n == key:
                return i
        from openpyxl.utils import column_index_from_string
        try:
            i = column_index_from_string(key.upper()) - 1
        except ValueError:
            i = -1
        if 0 <= i < len(names):
            return i
    raise XlMcpError(
        f"column_types names {key!r}, which is not one of the imported columns "
        f"({', '.join(names)}); use a header, a letter or a 1-based number")


def _type_columns(matrix: list[list[Any]], header: bool, column_types: dict | None,
                  parse_dates: bool, decimal_comma: bool) -> tuple[list[list[Any]], dict]:
    """Decide each column's type and convert the raw CSV text."""
    from openpyxl.utils import get_column_letter
    from ..core.package import parse_iso_date

    ncols = max((len(r) for r in matrix), default=0)
    first = 1 if header and matrix else 0
    names = [(str(matrix[0][i]) if header and matrix and i < len(matrix[0])
              and matrix[0][i] not in (None, "") else get_column_letter(i + 1))
             for i in range(ncols)]
    wanted: dict[int, str] = {}
    for key, kind in (column_types or {}).items():
        k = _vocab_fold(kind)
        if k not in COLUMN_TYPES:
            raise XlMcpError(f"column_types values are {', '.join(COLUMN_TYPES)}; got {kind!r}")
        wanted[_column_key(key, names)] = k
    report: dict = {}
    as_text, dates, not_numbers, not_dates = [], [], [], []
    for c in range(ncols):
        cells = [(r, row[c]) for r, row in enumerate(matrix[first:], first)
                 if c < len(row) and isinstance(row[c], str) and row[c] != ""]
        kind = wanted.get(c, "auto")
        if kind == "auto":
            reasons = [(v, _code_reason(v)) for _r, v in cells]
            hit = next(((v, why) for v, why in reasons if why), None)
            if hit:
                kind = "text"
                as_text.append({"column": names[c], "reason": hit[1], "example": hit[0]})
            elif parse_dates and cells and all(parse_iso_date(v) is not None for _r, v in cells):
                kind = "date"
        for r, v in cells:
            if kind == "text":
                matrix[r][c] = v
            elif kind == "date":
                if parse_iso_date(v) is not None:
                    matrix[r][c] = {"date": v}
                else:
                    not_dates.append(f"{get_column_letter(c + 1)}{r + 1}")
            else:
                out = _number(v, decimal_comma)
                if kind == "number" and isinstance(out, str):
                    not_numbers.append(f"{get_column_letter(c + 1)}{r + 1}")
                matrix[r][c] = out
        if kind == "date" and cells:
            dates.append(names[c])
    if as_text:
        report["columns_kept_as_text"] = as_text
    if dates:
        report["date_columns"] = dates
    if not_numbers:
        report["not_numbers"] = not_numbers[:10]
    if not_dates:
        report["not_dates"] = not_dates[:10]
    return matrix, report


def _vocab_fold(v) -> str:
    return re.sub(r"[\s_\-]+", "", str(v)).lower()


def import_data(path: str, source: str | None = None,
                source_file: str | None = None, fmt: str = "auto",
                location: Any = None, sheet: str | None = None,
                header: bool = True, delimiter: str | None = None,
                encoding: str = "utf-8", formulas: bool = False,
                column_types: dict | None = None, parse_dates: bool = False,
                allow_loss: bool = False, backup: bool = True,
                verify_com: bool | None = None) -> dict:
    """Import CSV / TSV / JSON into a sheet at an anchor. Backup + verify."""
    if fmt not in _FORMATS:
        raise XlMcpError(f"fmt must be one of {_FORMATS}")
    if not source and not source_file:
        raise XlMcpError("pass source (inline text) or source_file (a path)")
    # Check the codec name before either branch. The file branch used to
    # discover a bad name inside open(), which surfaced as a generic lookup
    # failure, and the inline branch never looked at all, so "utf-99" was
    # accepted in silence and the caller learned nothing about the encoding
    # it thought it had asked for.
    try:
        _codecs.lookup(encoding)
    except LookupError:
        raise XlMcpError(
            f"unknown encoding {encoding!r}; encoding must name a codec "
            "Python knows, for example utf-8, utf-8-sig, cp1252, or "
            "latin-1") from None
    info: dict = {}
    if source_file:
        sp = check_path(source_file, "read import file")
        with open(sp, "rb") as fh:
            raw = fh.read()
        text, used = _decode_text_file(raw, encoding)
        asked = _codecs.lookup(encoding).name
        if used != asked and not (used == "utf-8-sig" and asked == "utf-8"):
            info["encoding"] = (f"{used} (the file starts with its "
                                "byte-order mark)")
    else:
        text = source
    if text.startswith("\ufeff"):
        text = text[1:]  # A byte-order mark is not data
    fmt = _detect_format(fmt, source_file)
    if fmt == "json":
        matrix = _parse(text, fmt, delimiter, header)
    else:
        if delimiter is None and fmt == "csv":
            delimiter = _sniff_delimiter(text)
            if delimiter != ",":
                info["delimiter"] = delimiter
        decimal_comma = (delimiter == ";")
        if decimal_comma:
            info["decimal_comma"] = True
        matrix = _parse(text, fmt, delimiter, header, coerce=False)
        matrix, typing = _type_columns(matrix, header, column_types, parse_dates,
                                       decimal_comma)
        info.update(typing)
    rows = len(matrix)
    ncols = max((len(r) for r in matrix), default=0)
    if rows == 0 or ncols == 0:
        raise XlMcpError("the import produced no cells")
    if rows * ncols > gridio.MAX_WRITE_CELLS:
        raise CeilingExceeded(
            f"the import is {rows * ncols:,} cells, over the "
            f"{gridio.MAX_WRITE_CELLS:,}-cell write ceiling; split the source")

    pkg = WorkbookPackage.open(path)
    grid = pkg.resolve(location if location is not None else {"cell": "A1"},
                       default_sheet=sheet)
    top, left = grid.min_row, grid.min_col
    if top + rows - 1 > gridio._locate.MAX_ROW or \
            left + ncols - 1 > gridio._locate.MAX_COL:
        raise RangeOutOfBounds("the import would extend past the grid limits")
    ws = pkg.workbook[grid.sheet]
    neutralized = 0
    for i, row in enumerate(matrix):
        for j, val in enumerate(row):
            r, c = top + i, left + j
            if isinstance(val, str) and val[:1] in _INJECTION and not formulas:
                _limits.check_cell_text(
                    val, what=f"the imported value for {gridio.a1(r, c)}")
                cell = ws.cell(r, c)
                cell.value = val
                cell.data_type = "s"
                cell.quotePrefix = True  # stays text if edited in Excel
                neutralized += 1
                pkg._intended[(grid.sheet, gridio.a1(r, c))] = ("value", val)
            else:
                pkg.set_cell(grid.sheet, gridio.a1(r, c), val)
    result = pkg.save(allow_loss=allow_loss, backup=backup,
                      verify_com=verify_com)
    result["changed"]["imported"] = {
        "sheet": grid.sheet, "top_left": gridio.a1(top, left),
        "rows": rows, "cols": ncols, "format": fmt, **info}
    warnings = list(result.get("warnings", []))
    if neutralized:
        warnings.append(
            f"{neutralized} imported cell(s) began with a formula character "
            "(=, +, -, @) and were written as TEXT to block formula injection; "
            "pass formulas=true to import them as live formulas")
    result["warnings"] = warnings
    return result


def _csv_cell(v):
    """A CSV/TSV field the way Excel writes one: TRUE and
    FALSE, numbers in General form (0.00001, not 1e-05), empty for nothing."""
    from ..core import numfmt as _numfmt
    if v is None:
        return ""
    if isinstance(v, (bool, int, float)):
        return _numfmt.general(v)
    return v


def _text_encoding(fmt: str, bom: bool) -> str:
    """A CSV or TSV file starts with a UTF-8 byte-order mark by
    default, as Excel's own "CSV UTF-8" does, so Excel opens accented text
    as written instead of as mojibake; bom=false leaves it out."""
    return "utf-8-sig" if bom and fmt in ("csv", "tsv") else "utf-8"


def export_range(path: str, location: Any = None, sheet: str | None = None,
                 fmt: str = "csv", header: bool = True, values: str = "cached",
                 records: bool = False, out_file: str | None = None,
                 overwrite: bool = False, bom: bool = True) -> dict:
    """Export a range / table / sheet to CSV / TSV / JSON. Read-only of the
    WORKBOOK; out_file is guarded by core.outguard (never the source
    workbook, never a workbook extension, never inside the backup store,
    and an existing file refuses unless overwrite=true, which first keeps a
    timestamped .bak of what it replaces).

    location defaults to the sheet's true used range. values is cached |
    formula | both. out_file, when given, writes the serialized text to a
    file; otherwise the text is returned inline."""
    if fmt not in _EXPORT_FORMATS:
        raise XlMcpError(f"fmt must be one of {_EXPORT_FORMATS}")
    if values not in gridio.VALUE_MODES:
        raise XlMcpError(f"values must be one of {gridio.VALUE_MODES}")
    formula_wb = gridio.open_wb(path, data_only=False) \
        if values in ("formula", "both") else None
    cached_wb = gridio.open_wb(path, data_only=True) \
        if values in gridio.CACHED_MODES else None
    base = formula_wb if formula_wb is not None else cached_wb
    try:
        loc = location if location is not None else {
            "used_range": sheet if sheet is not None else True}
        grid = gridio.resolve(base, loc, default_sheet=sheet, path=path)
        if grid.empty:
            content = "[]" if fmt == "json" else ""
            return {"sheet": grid.sheet, "range": grid.a1, "format": fmt,
                    "rows": 0, "content": content, "value_mode": values}
        gridio.guard_cell_count(grid)
        shown = gridio.DisplayNotes(gridio.is_1904(base[grid.sheet]), path) \
            if values == "displayed" else None
        vals, labels, _hf = gridio.read_matrix(
            grid, mode=values, formula_wb=formula_wb, cached_wb=cached_wb,
            path=path, display_notes=shown)
        vals = [[gridio.compact_value(v) for v in row] for row in vals]
        if fmt == "json":
            if header and vals:
                cols = [str(h) if h is not None else get_column_letter(
                    grid.min_col + i) for i, h in enumerate(vals[0])]
                body = vals[1:]
                payload = ([dict(zip(cols, r)) for r in body] if records
                           else {"columns": cols, "rows": body})
            else:
                payload = vals
            content = _json.dumps(payload, ensure_ascii=False, default=str,
                                  indent=2)
        else:
            delim = "\t" if fmt == "tsv" else ","
            buf = _io.StringIO()
            writer = _csv.writer(buf, delimiter=delim, lineterminator="\n")
            for row in vals:
                writer.writerow([_csv_cell(v) for v in row])
            content = buf.getvalue()
        out: dict[str, Any] = {
            "sheet": grid.sheet, "range": grid.a1, "format": fmt,
            "rows": len(vals), "cols": len(vals[0]) if vals else 0,
            "value_mode": values}
        # An export is the one read whose output leaves the server entirely,
        # so an uncalculated formula cell silently exported as an empty field
        # is a wrong number in someone else's spreadsheet.
        note = gridio.absent_note(labels)
        if note:
            out["warning"] = note
            out["absent_cells"] = [
                gridio.a1(grid.min_row + i, grid.min_col + j)
                for i, row in enumerate(labels)
                for j, lab in enumerate(row) if lab == "absent"][:100]
        # Cached results from a workbook still owed a
        # recalculation are exported with a note saying so.
        if any(lab == "cached" for row in labels for lab in row):
            pending = gridio.pending_recalc_note(path)
            if pending:
                out["cached_note"] = pending
        if shown is not None and shown.note():
            out["display_note"] = shown.note()
        if out_file:
            from ..core import deadline as _deadline
            _deadline.commit_or_raise("the export")  # point of no return
            op, ginfo = guard_out_file(
                out_file, source=path, overwrite=overwrite,
                what="export file")
            with open(op, "w", encoding=_text_encoding(fmt, bom),
                      newline="") as fh:
                fh.write(content)
            out["written_to"] = op
            out.update(ginfo)
        else:
            out["content"] = content
        return out
    finally:
        for wb in (formula_wb, cached_wb):
            if wb is not None:
                wb.close()


_FILE_FORMATS = ("csv", "tsv", "json")
_UNSAFE_NAME = re.compile(r'[\\/:*?"<>|]+')


def _serialize_matrix(vals, fmt: str, header: bool, records: bool,
                      min_col: int) -> Any:
    """Shared range-to-text/payload serialization (export_range's logic,
    factored so export_file emits identical output per sheet). For json the
    PAYLOAD is returned (the caller decides bundling); csv/tsv return
    text."""
    if fmt == "json":
        if header and vals:
            cols = [str(h) if h is not None else get_column_letter(
                min_col + i) for i, h in enumerate(vals[0])]
            body = vals[1:]
            return ([dict(zip(cols, r)) for r in body] if records
                    else {"columns": cols, "rows": body})
        return vals
    delim = "\t" if fmt == "tsv" else ","
    buf = _io.StringIO()
    writer = _csv.writer(buf, delimiter=delim, lineterminator="\n")
    for row in vals:
        writer.writerow([_csv_cell(v) for v in row])
    return buf.getvalue()


def export_file(path: str, fmt: str = "csv", sheets: list | None = None,
                out_dir: str | None = None, out_file: str | None = None,
                header: bool = True, values: str = "cached",
                records: bool = False, overwrite: bool = False,
                bom: bool = True) -> dict:
    """Export whole sheets (or the whole workbook) to a CSV/TSV set or one
    JSON bundle. Read-only of the WORKBOOK; every output target is guarded
    by core.outguard (never the source workbook, never a workbook
    extension, never inside the backup store, and an existing file refuses
    unless overwrite=true, which first keeps a timestamped .bak of what it
    replaces).

    csv/tsv: one file per sheet into out_dir (named <stem>_<sheet>.csv), or
    inline per-sheet text when out_dir is omitted. json: a single bundle
    keyed by sheet name, written to out_file or returned inline."""
    if fmt not in _FILE_FORMATS:
        raise XlMcpError(f"fmt must be one of {_FILE_FORMATS}")
    if values not in gridio.VALUE_MODES:
        raise XlMcpError(f"values must be one of {gridio.VALUE_MODES}")
    if out_dir and fmt == "json":
        raise XlMcpError("json produces ONE bundle; pass out_file, not "
                         "out_dir")
    if out_file and fmt != "json":
        raise XlMcpError(f"{fmt} produces one file per sheet; pass "
                         "out_dir, not out_file")

    formula_wb = gridio.open_wb(path, data_only=False) \
        if values in ("formula", "both") else None
    cached_wb = gridio.open_wb(path, data_only=True) \
        if values in gridio.CACHED_MODES else None
    base = formula_wb if formula_wb is not None else cached_wb
    try:
        wanted = sheets if sheets else base.sheetnames
        missing = [s for s in wanted if s not in base.sheetnames]
        if missing:
            raise XlMcpError(
                f"no sheet(s) named {missing}; sheets: {base.sheetnames}")
        per_sheet: dict[str, Any] = {}
        exported: list[dict] = []
        absent_total = 0
        any_cached = False
        shown = gridio.DisplayNotes(gridio.is_1904(base.worksheets[0]), path) \
            if values == "displayed" else None
        for name in wanted:
            grid = gridio.resolve(base, {"used_range": name})
            if grid.empty:
                per_sheet[name] = [] if fmt == "json" else ""
                exported.append({"sheet": name, "rows": 0, "empty": True})
                continue
            gridio.guard_cell_count(grid)
            vals, labels, _hf = gridio.read_matrix(
                grid, mode=values, formula_wb=formula_wb,
                cached_wb=cached_wb, path=path, display_notes=shown)
            vals = [[gridio.compact_value(v) for v in row] for row in vals]
            per_sheet[name] = _serialize_matrix(
                vals, fmt, header, records, grid.min_col)
            any_cached = any_cached or any("cached" in row for row in labels)
            n_absent = sum(row.count("absent") for row in labels)
            absent_total += n_absent
            entry = {"sheet": name, "range": grid.a1, "rows": len(vals)}
            if n_absent:
                entry["absent_cells"] = n_absent
            exported.append(entry)

        out: dict[str, Any] = {"format": fmt, "value_mode": values,
                               "sheets": exported,
                               "sheet_count": len(exported)}
        if absent_total:
            out["warning"] = (f"{absent_total} exported cell(s): "
                              + gridio.ABSENT_WARNING)
        if any_cached:  # a note for cached results owed a recalculation
            pending = gridio.pending_recalc_note(path)
            if pending:
                out["cached_note"] = pending
        if shown is not None and shown.note():
            out["display_note"] = shown.note()
        if fmt == "json":
            bundle = {"workbook": os.path.basename(path),
                      "value_mode": values, "sheets": per_sheet}
            content = _json.dumps(bundle, ensure_ascii=False, default=str,
                                  indent=2)
            if out_file:
                from ..core import deadline as _deadline
                _deadline.commit_or_raise("the export")  # point of no return
                op, ginfo = guard_out_file(
                    out_file, source=path, overwrite=overwrite,
                    what="export file")
                with open(op, "w", encoding="utf-8", newline="") as fh:
                    fh.write(content)
                out["written_to"] = op
                out.update(ginfo)
            else:
                out["content"] = content
        else:
            if out_dir:
                od = guard_out_dir(out_dir, what="export files")
                if not os.path.isdir(od):
                    raise XlMcpError(
                        f"out_dir is not an existing directory: {od}")
                stem = os.path.splitext(os.path.basename(path))[0]
                dests = []
                for name, text in per_sheet.items():
                    safe = _UNSAFE_NAME.sub("_", name)
                    # workbook stem + sheet name + extension is built, not
                    # given, so it can exceed a filesystem's per-component
                    # limit (255 BYTES on ext4, 255 characters on Windows)
                    # even though both halves were legal on their own.
                    dests.append((
                        os.path.join(
                            od, _safesave.bound_name(f"{stem}_{safe}.{fmt}")),
                        text))
                lows = [os.path.normcase(d) for d, _t in dests]
                if len(set(lows)) < len(lows):
                    raise XlMcpError(
                        "two sheet names sanitize to the same output file "
                        "name; export them separately (export_range with "
                        "distinct out_file paths)")
                # Refuse-before-write: no partial file set on a collision.
                clashes = [d for d, _t in dests if os.path.exists(d)]
                if clashes and not overwrite:
                    raise FileExistsError(
                        "refusing to overwrite existing file(s): "
                        + ", ".join(clashes)
                        + ". Pass overwrite:true to replace them (each "
                        "replaced file is first copied to a timestamped "
                        ".bak beside it).")
                written = []
                replaced = []
                from ..core import deadline as _deadline
                _deadline.commit_or_raise("the export")  # point of no return
                for dest, text in dests:
                    op, ginfo = guard_out_file(
                        dest, source=path, overwrite=overwrite,
                        what="export file")
                    with open(op, "w", encoding=_text_encoding(fmt, bom),
                              newline="") as fh:
                        fh.write(text)
                    written.append(op)
                    if ginfo.get("backup_of_replaced"):
                        replaced.append(ginfo["backup_of_replaced"])
                out["written_to"] = written
                if replaced:
                    out["overwrote"] = True
                    out["backups_of_replaced"] = replaced
            else:
                out["content"] = per_sheet
        return out
    finally:
        for wb in (formula_wb, cached_wb):
            if wb is not None:
                wb.close()


__all__ = ["import_data", "export_range", "export_file"]
