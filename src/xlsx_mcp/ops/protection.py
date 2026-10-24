"""ops/protection.py: sheet and workbook protection.

set_protection is the file-tier protection tool: sheet protection with the
per-option allow flags, the workbook structure lock, and range-level unlock
exceptions, plus removal and a read-only status report.

HONESTY, stated everywhere it matters: xlsx protection is ADVISORY tamper
discouragement, not security. The password is stored as a hash Excel checks
before allowing UI edits; the file itself stays a readable zip and any tool
(including this one) can strip the protection. Real encryption of the
package is a different mechanism entirely and lives in the COM tier
(com_save_with_password), because no pure-Python path writes the CDF/agile
encryption wrapper safely.

Password hashing uses the legacy 16-bit hash from ECMA-376 Part 1 (the
password attribute), via openpyxl's implementation, which every Excel
version verifies. The stronger SHA-512 + salt + spin-count agile hash is
deliberately NOT hand-rolled here: writing it without an Excel round-trip
check risks a file Excel rejects, and the advisory nature of the feature
means the legacy hash gives the same practical protection.

OPTION SEMANTICS: in the file format each protection attribute TRUE means
the action is BLOCKED while the sheet is protected. This tool speaks the
Excel UI's language instead ("allow all users of this worksheet to..."):
the options dict maps option names to True = allowed, and the tool inverts
them for storage. Unlisted options keep Excel's defaults (everything
blocked except selecting cells).
"""

from __future__ import annotations


from ..core.errors import TargetNotFound, XlMcpError
from ..core.package import WorkbookPackage
from . import gridio

PROTECTION_ACTIONS = ("sheet", "workbook", "unlock", "remove", "status")

#: user-facing option name (True = allowed while protected) -> the openpyxl
#: SheetProtection attribute it inverts into.
SHEET_OPTIONS: dict[str, str] = {
    "select_locked_cells": "selectLockedCells",
    "select_unlocked_cells": "selectUnlockedCells",
    "format_cells": "formatCells",
    "format_columns": "formatColumns",
    "format_rows": "formatRows",
    "insert_columns": "insertColumns",
    "insert_rows": "insertRows",
    "insert_hyperlinks": "insertHyperlinks",
    "delete_columns": "deleteColumns",
    "delete_rows": "deleteRows",
    "sort": "sort",
    "auto_filter": "autoFilter",
    "pivot_tables": "pivotTables",
    "objects": "objects",
    "scenarios": "scenarios",
}


#: Which optional arguments each action reads (an argument
#: the chosen action would ignore is refused rather than dropped silently).
_ACTION_ARGS = {
    "sheet": {"sheet", "password", "options", "unlock_ranges"},
    "workbook": {"password"},
    "unlock": {"sheet", "unlock_ranges"},
    "remove": {"sheet"},
    "status": set(),
}

#: The most unlocked ranges status lists per sheet.
_MAX_UNLOCKED_LISTED = 50


def _unlocked_ranges(ws) -> list[str]:
    """The sheet's unlocked cells as A1 ranges: runs along each row, then
    runs with the same columns on consecutive rows joined into rectangles."""
    from openpyxl.utils import get_column_letter

    cells = sorted((r, c) for (r, c), cell in ws._cells.items()
                   if cell.has_style and not cell.protection.locked)
    runs: dict[int, list[tuple[int, int]]] = {}
    for r, c in cells:
        row = runs.setdefault(r, [])
        if row and row[-1][1] == c - 1:
            row[-1] = (row[-1][0], c)
        else:
            row.append((c, c))
    rects: list[list[int]] = []  # [r0, r1, c0, c1]
    open_: dict[tuple[int, int], list[int]] = {}
    for r in sorted(runs):
        nxt: dict[tuple[int, int], list[int]] = {}
        for span in runs[r]:
            rect = open_.get(span)
            if rect is not None and rect[1] == r - 1:
                rect[1] = r
            else:
                rect = [r, r, span[0], span[1]]
                rects.append(rect)
            nxt[span] = rect
        open_ = nxt
    out = []
    for r0, r1, c0, c1 in rects:
        a, b = f"{get_column_letter(c0)}{r0}", f"{get_column_letter(c1)}{r1}"
        out.append(a if a == b else f"{a}:{b}")
    return out


def _mark(pkg: WorkbookPackage, ranges, sheet, locked: bool) -> list[dict]:
    from openpyxl.styles import Protection as _CellProt

    touched = []
    for item in ranges:
        loc = {"range": item} if isinstance(item, str) else item
        grid = pkg.resolve(loc, default_sheet=sheet)
        ws = pkg.workbook[grid.sheet]
        for row in ws.iter_rows(min_row=grid.min_row, max_row=grid.max_row,
                                min_col=grid.min_col, max_col=grid.max_col):
            for cell in row:
                cell.protection = _CellProt(locked=bool(locked))
        touched.append({"sheet": grid.sheet, "range": grid.a1})
    return touched


def _require_sheet(pkg: WorkbookPackage, sheet: str | None):
    wb = pkg.workbook
    if sheet is None:
        return wb.active
    if sheet not in wb.sheetnames:
        raise TargetNotFound(
            f"no sheet named {sheet!r}; sheets: {wb.sheetnames}")
    return wb[sheet]


def set_protection(path: str, action: str, sheet: str | None = None,
                   password: str | None = None, options: dict | None = None,
                   unlock_ranges: list | None = None, locked: bool = False,
                   structure: bool = True, windows: bool = False,
                   scope: str = "all", allow_loss: bool = False,
                   backup: bool = True,
                   verify_com: bool | None = None) -> dict:
    """Sheet/workbook protection, unlock exceptions, removal, status.
    Backup + verify on write; status is read-only."""
    if action not in PROTECTION_ACTIONS:
        raise XlMcpError(
            f"action must be one of {PROTECTION_ACTIONS}, got {action!r}")
    given = {name for name, value in (("sheet", sheet), ("password", password),
                                      ("options", options),
                                      ("unlock_ranges", unlock_ranges)) if value}
    ignored = sorted(given - _ACTION_ARGS[action])
    if ignored:
        raise XlMcpError(
            f"action '{action}' does not use {', '.join(ignored)}; it takes "
            + (", ".join(sorted(_ACTION_ARGS[action])) or "no optional arguments")
            + ". To protect a sheet with some cells left editable, pass unlock_ranges "
            "with action='sheet'.")

    if action == "status":
        wb = gridio.open_wb(path)
        try:
            sec = wb.security
            # A password is stored either as the legacy 16-bit hash or, as
            # every current Excel writes it, a salted hash (hashValue).
            book = {
                "structure_locked": bool(sec and sec.lockStructure),
                "windows_locked": bool(sec and sec.lockWindows),
                "has_password": bool(sec and (sec.workbookPassword
                                              or sec.workbookHashValue)),
            }
            sheets = []
            for ws in wb.worksheets:
                prot = ws.protection
                blocked = {
                    opt: bool(getattr(prot, attr, False))
                    for opt, attr in SHEET_OPTIONS.items()}
                unlocked = _unlocked_ranges(ws)
                entry = {
                    "sheet": ws.title,
                    "protected": bool(prot.sheet),
                    "has_password": bool(prot.password
                                         or getattr(prot, "hashValue", None)),
                    "allowed_while_protected": sorted(
                        o for o, b in blocked.items() if not b),
                    # The cells that stay editable under protection.
                    "unlocked_ranges": unlocked[:_MAX_UNLOCKED_LISTED],
                }
                if len(unlocked) > _MAX_UNLOCKED_LISTED:
                    entry["unlocked_ranges_total"] = len(unlocked)
                sheets.append(entry)
            return {"workbook": book, "sheets": sheets}
        finally:
            wb.close()

    pkg = WorkbookPackage.open(path)

    if action == "sheet":
        from openpyxl.worksheet.protection import SheetProtection
        ws = _require_sheet(pkg, sheet)
        opts = options or {}
        unknown = sorted(set(opts) - set(SHEET_OPTIONS))
        if unknown:
            raise XlMcpError(
                f"unknown protection option(s) {unknown}; valid: "
                f"{sorted(SHEET_OPTIONS)}")
        # DECLARATIVE re-protect: start from a FRESH SheetProtection so options
        # from a prior protect() do not survive. The reused object let a first
        # call's sort/auto_filter leak past a second, narrower call -- the
        # response said allowed:[format_cells] while status still allowed sort
        # and auto_filter, contradicting this module's own "unlisted options
        # keep Excel's defaults" contract.
        # Excel's own defaults: Protect Sheet blocks objects and
        # scenarios unless allowed. openpyxl's defaults left both open, so
        # comments, images and charts stayed editable on a "protected" sheet,
        # against this module's own contract.
        prot = SheetProtection(sheet=True, objects=True, scenarios=True)
        ws.protection = prot
        for opt, allowed in opts.items():
            # file semantics: TRUE blocks; user semantics: TRUE allows
            setattr(prot, SHEET_OPTIONS[opt], not bool(allowed))
        if password:
            prot.password = password  # legacy ECMA-376 hash via openpyxl
        # unlock_ranges with action='sheet' were dropped
        # silently; they are the cells the protection leaves editable.
        unlocked = _mark(pkg, unlock_ranges, ws.title, False) if unlock_ranges else []
        pkg._changed["protection"] = {
            "sheet": ws.title, "protected": True,
            "allowed": sorted(o for o, a in opts.items() if a),
            "password": bool(password)}
        if unlocked:
            pkg._changed["protection"]["cells_unlocked"] = unlocked
    elif action == "workbook":
        from openpyxl.workbook.protection import WorkbookProtection
        sec = pkg.workbook.security or WorkbookProtection()
        sec.lockStructure = bool(structure)
        sec.lockWindows = bool(windows)
        if password:
            sec.workbookPassword = password
        pkg.workbook.security = sec
        pkg._changed["protection"] = {
            "workbook": True, "structure": bool(structure),
            "windows": bool(windows), "password": bool(password)}
    elif action == "unlock":
        if not unlock_ranges:
            raise XlMcpError(
                "unlock needs unlock_ranges (a list of location objects or "
                "A1 range strings to mark unlocked)")
        touched = _mark(pkg, unlock_ranges, sheet, locked)
        pkg._changed["protection"] = {
            "cells_locked" if locked else "cells_unlocked": touched}
    else:  # remove
        if scope not in ("sheet", "workbook", "all"):
            raise XlMcpError(
                f"scope must be sheet, workbook, or all, got {scope!r}")
        removed = []
        if scope in ("sheet", "all"):
            from openpyxl.worksheet.protection import SheetProtection
            targets = ([_require_sheet(pkg, sheet)] if sheet is not None
                       else pkg.workbook.worksheets)
            for ws in targets:
                if ws.protection.sheet:
                    ws.protection = SheetProtection(sheet=False)
                    removed.append(f"sheet:{ws.title}")
        if scope in ("workbook", "all"):
            sec = pkg.workbook.security
            # openpyxl auto-creates an empty WorkbookProtection on new
            # files; only an object that actually locks something counts
            if sec is not None and (sec.lockStructure or sec.lockWindows
                                    or sec.workbookPassword):
                pkg.workbook.security = None
                removed.append("workbook")
        if not removed:
            raise TargetNotFound(
                "nothing to remove: no matching protection is set")
        pkg._changed["protection"] = {"removed": removed}

    return pkg.save(allow_loss=allow_loss, backup=backup,
                    verify_com=verify_com)


__all__ = ["set_protection", "PROTECTION_ACTIONS", "SHEET_OPTIONS"]
