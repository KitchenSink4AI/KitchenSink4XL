"""ops/sortfilter.py: sort and filter.

sort_range reorders whole rows of a range (or table body) by one or more keys,
writing the reordered values and styles back. A moved formula has its relative
references shifted by its row displacement (core.refs.offset_formula), Excel's
own sort semantics, so a formula that referred to its own row keeps referring to
its own row after the move; absolute ($) anchors stay put.

set_filter is the honest answer to a file-format landmine: an .xlsx autofilter
stores the filter CRITERIA, not which rows are hidden. A criteria-only write
would leave every row visible until Excel re-applies the filter. So set_filter
EVALUATES the criteria here and hides the non-matching data rows (row-level
hidden flag), and also records the autofilter range so Excel shows the filter
UI. The hidden state is a snapshot over the current cached / literal values;
Excel re-evaluates on its own when the user touches the filter. Formula cells
with no cached value cannot be evaluated and are reported rather than guessed.
clear_filter removes the autofilter and unhides the rows.

All three route mutations through WorkbookPackage for the backup and verify.
"""

from __future__ import annotations

from copy import copy as _copy
from typing import Any

from openpyxl.utils import get_column_letter

from ..core import arrays as _arrays
from ..core import calc as _calc
from ..core import refs as _refs
from ..core.errors import TargetNotFound, UnsupportedStructure, XlMcpError
from ..core.package import WorkbookPackage
from . import cells as _cells
from . import gridio


def _column_index(col: Any, header_names: list[str], min_col: int) -> int:
    """Resolve a key/criteria column to a 0-based offset within the block."""
    if isinstance(col, int):
        return col - 1
    s = str(col)
    if s in header_names:
        return header_names.index(s)
    from openpyxl.utils import column_index_from_string
    try:
        return column_index_from_string(s.upper()) - min_col
    except Exception:
        raise XlMcpError(f"no column {col!r} in the range; headers: "
                         f"{header_names}")


def _header_names(ws, row: int, min_col: int, max_col: int) -> list[str]:
    return [str(ws.cell(row, c).value) if ws.cell(row, c).value is not None
            else get_column_letter(c)
            for c in range(min_col, max_col + 1)]


def _refuse_merges(ws, grid, what: str) -> None:
    """Refuse when the target rectangle overlaps a merged range.

    Rewriting rows under a merge is not a thing openpyxl can do: only the
    top-left of a merge holds a value and the rest are read-only MergedCell
    objects, so the write half of a sort died mid-flight with
    "'MergedCell' object attribute 'value' is read-only"
    after the reads had already run. Refuse up front, name the merges, and
    say the remedy."""
    overlapping = []
    for mr in list(getattr(ws, "merged_cells", None).ranges
                   if getattr(ws, "merged_cells", None) else []):
        if (mr.min_row <= grid.max_row and mr.max_row >= grid.min_row
                and mr.min_col <= grid.max_col and mr.max_col >= grid.min_col):
            overlapping.append(str(mr))
    if overlapping:
        raise UnsupportedStructure(
            f"cannot {what} {grid.a1} on {grid.sheet!r}: it overlaps merged "
            f"range(s) {', '.join(sorted(overlapping))}, and merged cells "
            f"cannot be rewritten row by row. Unmerge them first "
            f"(set_merge action='unmerge'), then {what}.")


def sort_range(path: str, location: Any, keys: list, has_header: bool = True,
               sheet: str | None = None, allow_loss: bool = False,
               backup: bool = True,
               verify_com: bool | None = None) -> dict:
    """Sort a range or table body by one or more keys. Backup + verify."""
    if not keys or not isinstance(keys, list):
        raise XlMcpError("keys must be a non-empty list of {column, order}")
    pkg = WorkbookPackage.open(path)
    grid = pkg.resolve(location, default_sheet=sheet)
    if grid.empty:
        raise XlMcpError("cannot sort an empty range")
    gridio.guard_cell_count(grid)
    ws = pkg.workbook[grid.sheet]
    _refuse_merges(ws, grid, "sort")
    # An array formula cannot be reordered row by row: its ref anchor stays
    # behind and the workbook stops opening. Excel refuses the identical Sort
    # with "You can't change part of an array" (core.arrays).
    _arrays.refuse_if_touched(ws, grid, "sort")
    cached = gridio.open_wb(path, data_only=True)
    cws = cached[grid.sheet]

    min_row, min_col, max_row, max_col = (
        grid.min_row, grid.min_col, grid.max_row, grid.max_col)
    data_top = min_row + 1 if has_header else min_row
    header_names = _header_names(ws, min_row, min_col, max_col) if has_header \
        else [get_column_letter(c) for c in range(min_col, max_col + 1)]

    key_specs = []
    for k in keys:
        # The shared direction reader (query_range uses it
        # too): {"descending": true} and order "down" now sort descending, and
        # an unknown direction or key refuses instead of sorting ascending.
        col, desc = _cells.sort_spec(k, "keys")
        if col is None:
            raise XlMcpError("each key needs a column")
        key_specs.append((_column_index(col, header_names, min_col), desc))

    # Hidden rows, Excel's way (COM ground truth): sorting a
    # range with an ACTIVE AUTOFILTER sorts only the visible rows and
    # leaves filter-hidden rows pinned in place with their data; a sort
    # over MANUALLY hidden rows (no filter) moves every row while the
    # hidden flags stay with the row POSITIONS. The old code always did
    # the latter, so a sort after set_filter left the filter hiding the
    # WRONG rows with zero warnings.
    hidden_rows = [r for r in range(data_top, max_row + 1)
                   if ws.row_dimensions[r].hidden]
    # A table's own filter (set_filter writes it to the table part
    # now) is an active autofilter too, so its hidden rows stay pinned.
    target = (min_col, min_row, max_col, max_row)
    table_filtered = any(
        _intersects(_rect(t.ref), target) and _table_filtered(ws, t)
        for t in ws.tables.values())
    pin_hidden = bool(hidden_rows) and (bool(ws.auto_filter.ref)
                                        or table_filtered)
    if pin_hidden:
        hidden_set = set(hidden_rows)
        sortable_rows = [r for r in range(data_top, max_row + 1)
                         if r not in hidden_set]
    else:
        sortable_rows = list(range(data_top, max_row + 1))

    key_offsets = {o for o, _rev in key_specs}
    uncached_keys = 0
    #: Formula cells in the sorted block that DID have a cached value. The
    #: openpyxl round-trip does not carry a cache through, so every one of
    #: them reads back 'absent' after this save. uncached_keys counts the
    #: opposite population (cells with no cache at read time), which is why a
    #: fully-computed workbook produced zero warnings while losing all of its
    #: numbers.
    dropped_cache = 0
    rows = []
    for r in sortable_rows:
        vals, styles, sortvals, texts = [], [], [], []
        for c in range(min_col, max_col + 1):
            cell = ws.cell(r, c)
            vals.append(cell.value)
            styles.append(_copy(cell._style))
            # A TEXT cell that merely starts with '=' is data, not a formula
            # (import_data's neutralized injection text); it must keep its
            # string type through the rewrite instead of going live.
            texts.append(_calc.looks_like_formula_text(cell))
            cv = cws.cell(r, c).value
            if _calc.is_formula_cell(cell):
                if cv is None:
                    if (c - min_col) in key_offsets:
                        uncached_keys += 1  # sorts by formula TEXT, not value
                else:
                    dropped_cache += 1
            sortvals.append(cv if cv is not None else cell.value)
        rows.append({"src": r, "vals": vals, "styles": styles,
                     "sort": sortvals, "text": texts})
    cached.close()

    for offset, reverse in reversed(key_specs):
        rows.sort(key=lambda row, o=offset, rv=reverse: _cells._sort_key(
            row["sort"][o], reverse=rv), reverse=reverse)

    wrote_formula = False
    for i, row in enumerate(rows):
        dest_r = sortable_rows[i]
        dr = dest_r - row["src"]
        for j, (val, style, is_text) in enumerate(
                zip(row["vals"], row["styles"], row["text"])):
            cell = ws.cell(dest_r, min_col + j)
            out = val
            if is_text:
                cell.value = out
                cell.data_type = "s"   # stays TEXT, never re-armed
                cell._style = _copy(style)
                continue
            if isinstance(val, str) and val.startswith("="):
                if dr != 0:
                    out = _refs.offset_formula(val, dr, 0, grid.sheet)
                wrote_formula = True
            cell.value = out
            cell._style = _copy(style)
    if wrote_formula:
        pkg._formula_written = True
    # The sort re-populates the whole data block without per-cell intents;
    # the untouched-number preservation pass must keep its hands off it
    # (a moved 17-digit double is a WRITTEN cell here, not an untouched one).
    pkg.note_region_write(grid.sheet, data_top, min_col, max_row, max_col)
    pkg._changed["sorted"] = {
        "sheet": grid.sheet, "range": grid.a1,
        "rows_sorted": len(rows),
        **({"rows_pinned_hidden": len(hidden_rows)} if pin_hidden else {}),
        "keys": [
            {"column": header_names[o] if o < len(header_names) else o,
             "order": "desc" if rev else "asc"} for o, rev in key_specs]}
    result = pkg.save(allow_loss=allow_loss, backup=backup,
                      verify_com=verify_com)
    notes: list[str] = []
    if pin_hidden:
        notes.append(
            f"{len(hidden_rows)} filter-hidden row(s) were left in place "
            "and excluded from the sort, matching Excel's semantics for "
            "sorting a filtered range; clear_filter first to sort every "
            "row.")
    elif hidden_rows:
        notes.append(
            f"{len(hidden_rows)} manually hidden row(s) were sorted along "
            "with the visible rows (Excel's own behavior without a "
            "filter); hidden flags stay with the row POSITIONS, so "
            "different rows may be hidden now.")
    if uncached_keys:
        notes.append(
            f"{uncached_keys} sort-key cell(s) hold a formula with no cached "
            "value and were ordered by their formula TEXT, not their result. "
            "Run recalculate or open in Excel first for a value-accurate "
            "sort.")
    if dropped_cache:
        notes.append(
            f"{dropped_cache} formula cell(s) in the sorted range had a "
            "cached value before this sort and have none after it (label "
            "'absent'): a file-based save cannot carry Excel's cache through "
            "a formula rewrite. " + gridio.ABSENT_WARNING)
    if notes:
        result["warnings"] = list(result.get("warnings", [])) + notes
    return result


def _rect(ref: str) -> tuple[int, int, int, int]:
    """(min_col, min_row, max_col, max_row) of an A1 range."""
    from openpyxl.utils import range_boundaries
    return range_boundaries(str(ref).replace("$", ""))


def _intersects(a: tuple, b: tuple) -> bool:
    return not (a[2] < b[0] or b[2] < a[0] or a[3] < b[1] or b[3] < a[1])


def _contains(outer: tuple, inner: tuple) -> bool:
    return (outer[0] <= inner[0] and outer[1] <= inner[1]
            and inner[2] <= outer[2] and inner[3] <= outer[3])


def _table_for_filter(ws, grid, what: str):
    """The table a filter target lies in, or None when it touches none.

    Excel filters a table with the table's OWN autofilter, stored in the
    table part. A sheet filter over a table's range (and the
    _xlnm._FilterDatabase name that goes with it) makes a workbook Excel
    will not open at all, which is what set_filter wrote for a table, its
    range or a cell in it (silent corruption, live since
    1.2.6). A target inside a table means that table, as in Excel; a range
    that reaches outside the table it overlaps refuses with the remedy."""
    target = (grid.min_col, grid.min_row, grid.max_col, grid.max_row)
    hits = [t for t in ws.tables.values() if _intersects(_rect(t.ref), target)]
    if not hits:
        return None
    names = [t.displayName or t.name for t in hits]
    if len(hits) > 1:
        raise XlMcpError(
            f"{grid.a1} on {ws.title!r} overlaps the tables {names}; Excel "
            f"filters one table at a time. Nothing was changed. {what.capitalize()} "
            f"one table by name (location {{\"table\": \"{names[0]}\"}}).")
    table = hits[0]
    name = names[0]
    if not _contains(_rect(table.ref), target):
        raise XlMcpError(
            f"{grid.a1} on {ws.title!r} overlaps table {name} ({table.ref}) "
            "but reaches outside it, and Excel will not open a workbook with "
            "a sheet filter over a table. Nothing was changed. Use location "
            f"{{\"table\": \"{name}\"}} (or a cell inside it) to {what} the "
            "table, or a range outside the table.")
    return table


def _table_filter_rect(ws, table) -> tuple[int, int, int, int]:
    """The rows a table's own filter covers: the header row and the data
    rows, never the totals row (Excel's table autofilter stops above it)."""
    hrc = table.headerRowCount if table.headerRowCount is not None else 1
    if not hrc:
        raise XlMcpError(
            f"table {table.displayName or table.name} has no header row, so "
            "it has no filter; turn the header row on in Excel (Table Design, "
            "Header Row) first. Nothing was changed.")
    c0, r0, c1, r1 = _rect(table.ref)
    return c0, r0, c1, r1 - int(table.totalsRowCount or 0)


def _a1(rect: tuple) -> str:
    return (f"{get_column_letter(rect[0])}{rect[1]}:"
            f"{get_column_letter(rect[2])}{rect[3]}")


#: How each criterion is stored in the autofilter, the
#: way Excel stores it, so Excel's own Reapply shows the same rows. Only eq
#: and in were written before; a second criterion (Amount gt 1000) hid rows
#: here but was missing from the file.
_CUSTOM_OPS = {"ne": "notEqual", "gt": "greaterThan", "ge": "greaterThanOrEqual",
               "lt": "lessThan", "le": "lessThanOrEqual"}


def _escape_wildcards(text: str) -> str:
    """Excel's custom filters read * and ? as wildcards and ~ as their
    escape (contains "?" was stored as *?*, and Excel's
    Reapply then matched every row; measured: *~?* matches only the "?")."""
    return text.replace("~", "~~").replace("*", "~*").replace("?", "~?")


def _custom(op: str, value) -> tuple[str, str] | None:
    """(operator, val) of one Excel customFilter, or None."""
    text = "" if value is None else str(value)
    if op in _CUSTOM_OPS:
        if op == "ne" and _numeric_criterion(value) is None:
            # Excel's notEqual reads * and ? as wildcards too (measured).
            return _CUSTOM_OPS[op], _escape_wildcards(text)
        return _CUSTOM_OPS[op], text
    if op == "contains":
        return "equal", f"*{_escape_wildcards(text)}*"
    if op == "startswith":
        return "equal", f"{_escape_wildcards(text)}*"
    if op == "endswith":
        return "equal", f"*{_escape_wildcards(text)}"
    if op == "not_blank":
        return "notEqual", " "
    return None


def _filter_column(offset: int, crits: list[tuple[str, Any]]):
    """The FilterColumn Excel writes for one column's criteria, or (None,
    why) when Excel's autofilter cannot hold them."""
    from openpyxl.worksheet.filters import (
        CustomFilter, CustomFilters, FilterColumn, Filters)

    if len(crits) == 1:
        op, value = crits[0]
        if op in ("eq", "in"):
            vals = value if isinstance(value, list) else [value]
            return FilterColumn(colId=offset, filters=Filters(
                filter=[str(v) for v in vals if v is not None],
                blank=any(v is None or v == "" for v in vals) or None)), None
        if op == "is_blank":
            return FilterColumn(colId=offset, filters=Filters(blank=True)), None
        if op == "not_in":
            vals = value if isinstance(value, list) else [value]
            if len(vals) <= 2:
                return FilterColumn(colId=offset, customFilters=CustomFilters(
                    _and=True, customFilter=[CustomFilter(
                        operator="notEqual",
                        val=str(v) if _numeric_criterion(v) is not None
                        else _escape_wildcards("" if v is None else str(v)))
                        for v in vals])), None
            return None, "not_in with more than two values"
    customs = [_custom(op, value) for op, value in crits]
    if all(c is not None for c in customs) and len(customs) <= 2:
        return FilterColumn(colId=offset, customFilters=CustomFilters(
            _and=len(customs) == 2 or None,
            customFilter=[CustomFilter(operator=o, val=v) for o, v in customs])), None
    return None, ("a regex criterion" if any(op == "regex" for op, _ in crits)
                  else "more than two conditions, or conditions Excel cannot combine, on one "
                  "column")


def _is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _value_list_columns(preds) -> set[int]:
    """Columns whose only criteria are eq / in: stored as
    Excel stores a ticked value list, and matched the way Excel matches it."""
    ops: dict[int, set] = {}
    for offset, op, _value in preds:
        ops.setdefault(offset, set()).add(op)
    return {o for o, s in ops.items() if s <= {"eq", "in"}}


def _shown(value, cell) -> str:
    """The text a cell shows (core.numfmt), which is what an Excel value
    list matches."""
    from ..core import numfmt as _numfmt
    from . import gridio
    return _numfmt.display(value, cell.number_format,
                           date1904=gridio.is_1904(cell.parent))


def _target_hits(value, shown: str, target) -> bool:
    """A value-list criterion matches a cell by value when both are numbers
    (1000 matches a cell showing 1,000), else by the text the cell shows,
    ignoring case: "007" never matches "7", as in Excel."""
    if _is_number(value) and _is_number(target):
        return float(value) == float(target)
    return shown.casefold() == ("" if target is None else str(target)).casefold()


def _numeric_criterion(target) -> float | None:
    """The number a criterion stands for in Excel's notEqual: a number, or
    text that reads as one ("7", "007"); None for other text."""
    if _is_number(target):
        return float(target)
    if isinstance(target, str):
        try:
            return float(target)
        except ValueError:
            return None
    return None


def _ne_keeps(value, cell, target) -> bool:
    """ne / not_in the way Excel's notEqual re-applies (measured): a
    criterion that reads as a number removes only cells holding that number
    (ne "7" and ne "007" both remove the number 7 and keep the texts "7" and
    "007"; ne 1000 removes a cell showing 1,000); other text removes the
    cells showing it, ignoring case."""
    for t in (target if isinstance(target, list) else [target]):
        n = _numeric_criterion(t)
        if n is not None:
            if _is_number(value) and float(value) == n:
                return False
        elif _shown(value, cell).casefold() == ("" if t is None else str(t)).casefold():
            return False
    return True


def set_filter(path: str, location: Any, criteria: list | None = None,
               sheet: str | None = None, allow_loss: bool = False,
               backup: bool = True,
               verify_com: bool | None = None) -> dict:
    """Apply an autofilter and hide the non-matching rows. Backup + verify."""
    pkg = WorkbookPackage.open(path)
    grid = pkg.resolve(location, default_sheet=sheet)
    if grid.empty:
        raise XlMcpError("cannot filter an empty range")
    ws = pkg.workbook[grid.sheet]
    min_row, min_col, max_row, max_col = (
        grid.min_row, grid.min_col, grid.max_row, grid.max_col)
    # A target inside a table filters that table, with its own
    # autofilter; never a sheet filter over it (Excel will not open that).
    table = _table_for_filter(ws, grid, "filter")
    filter_ref = grid.a1
    if table is not None:
        rect = _table_filter_rect(ws, table)
        min_col, min_row, max_col, max_row = rect
        filter_ref = _a1(rect)
    header_names = _header_names(ws, min_row, min_col, max_col)

    preds = []
    for cr in (criteria or []):
        if not isinstance(cr, dict) or "column" not in cr:
            raise XlMcpError(
                "each criterion is {column, op, value}; op defaults to eq")
        # Unknown keys used to be IGNORED: the natural LLM shape
        # {"column": "Category", "equals": "Freight"} parsed as op=eq
        # value=None and hid every data row with ok:true. A stray key or a
        # value-less binary op refuses like a wrong op does.
        unknown = set(cr) - {"column", "op", "value"}
        if unknown:
            raise XlMcpError(
                f"unknown criterion key(s) {sorted(unknown)}; each "
                "criterion is {column, op, value} (op defaults to eq; "
                "for equality put the wanted value under 'value')")
        op = cr.get("op", "eq")
        if op not in (_cells._OPS_BINARY | _cells._OPS_SET | _cells._OPS_UNARY):
            raise XlMcpError(f"unknown filter op {op!r}")
        if op not in _cells._OPS_UNARY and "value" not in cr:
            raise XlMcpError(
                f"criterion for column {cr['column']!r} has op {op!r} but "
                "no value; only is_blank/not_blank need none")
        offset = _column_index(cr["column"], header_names, min_col)
        preds.append((offset, op, cr.get("value")))

    cached = gridio.open_wb(path, data_only=True)
    cws = cached[grid.sheet]
    hidden = shown = uneval = 0
    warnings: list[str] = []

    def cell_value(r, c):
        """(value, evaluable): the cached result, or the literal; a formula
        with no cached result cannot be evaluated."""
        cv = cws.cell(r, c).value
        if cv is None and _calc.is_formula_cell(ws.cell(r, c)):
            return None, False
        return (cv if cv is not None else ws.cell(r, c).value), True

    # Excel matches a value list (eq, in) against the text a
    # cell SHOWS: <filter val="1000"/> hid every row of a column formatted
    # #,##0 (it shows 1,000), and "007" never matches "7" (this server
    # compared them as numbers). The rows a value-list criterion picks are
    # found here by value (numbers) or shown text, the list stored is the
    # shown text of those rows, and a row is kept when its shown text is on
    # the list: exactly what Excel's Reapply then shows (measured). A column
    # holding dates keeps the old comparison (Excel stores date groups).
    import datetime as _dt
    value_lists: dict[int, set[str]] = {}
    for offset in _value_list_columns(preds):
        c = min_col + offset
        seen = [cell_value(r, c) for r in range(min_row + 1, max_row + 1)]
        if any(isinstance(v, (_dt.datetime, _dt.date, _dt.time)) for v, _ok in seen):
            continue
        targets = []
        for o, op, value in preds:
            if o == offset:
                targets.extend(value if isinstance(value, list) else [value])
        picked: set[str] = set()
        for (v, ok), r in zip(seen, range(min_row + 1, max_row + 1)):
            if not ok:
                continue
            text = _shown(v, ws.cell(r, c))
            if any(_target_hits(v, text, t) for t in targets):
                picked.add(text)
        value_lists[offset] = picked
    lists_folded = {o: {t.casefold() for t in s} for o, s in value_lists.items()}

    for r in range(min_row + 1, max_row + 1):
        keep = True
        for offset, op, value in preds:
            c = min_col + offset
            v, ok = cell_value(r, c)
            if not ok:
                uneval += 1
                continue  # cannot evaluate a formula with no cached value
            if offset in value_lists:
                if _shown(v, ws.cell(r, c)).casefold() not in lists_folded[offset]:
                    keep = False
                    break
                continue
            if op in ("ne", "not_in") and not _ne_keeps(v, ws.cell(r, c), value):
                keep = False
                break
            if op not in ("ne", "not_in") and not _cells._cmp(v, op, value):
                keep = False
                break
        ws.row_dimensions[r].hidden = not keep
        if keep:
            shown += 1
        else:
            hidden += 1
    cached.close()
    if uneval:
        warnings.append(
            f"{uneval} filtered cell(s) hold a formula with no cached value; "
            "those rows were kept visible. Run recalculate or open in Excel "
            "first for an exact filter.")

    stray = None
    if table is not None:
        from openpyxl.worksheet.filters import AutoFilter
        if table.autoFilter is None:
            table.autoFilter = AutoFilter()
        af = table.autoFilter
        # A sheet filter over this table, as earlier versions
        # wrote it, is the shape Excel will not open; filtering the table
        # removes it, as clear_filter does.
        stray = ws.auto_filter.ref
        if stray and _intersects(_rect(stray), _rect(table.ref)):
            ws.auto_filter.ref = None
            ws.auto_filter.filterColumn = []
        else:
            stray = None
    else:
        af = ws.auto_filter
    af.ref = filter_ref
    af.filterColumn = []
    by_column: dict[int, list] = {}
    for offset, op, value in preds:
        by_column.setdefault(offset, []).append((op, value))
    for offset, crits in sorted(by_column.items()):
        if offset in value_lists:
            from openpyxl.worksheet.filters import FilterColumn, Filters
            texts = value_lists[offset]
            af.filterColumn.append(FilterColumn(colId=offset, filters=Filters(
                filter=sorted(t for t in texts if t != ""),
                blank=True if "" in texts else None)))
            continue
        fc, why = _filter_column(offset, crits)
        if fc is not None:
            af.filterColumn.append(fc)
        else:
            warnings.append(
                f"the rows were filtered here, but {why} on column "
                f"{header_names[offset]!r} cannot be stored in an Excel autofilter; "
                "Excel's Reapply will not repeat that part of the filter")

    pkg._changed["filter"] = {
        "sheet": grid.sheet, "range": filter_ref,
        **({"table": table.displayName or table.name} if table is not None else {}),
        # Named for what they count. "rows_shown": 4 was read as
        # "the header and 3 data rows" when it was 4 data rows.
        "data_rows_visible": shown, "data_rows_hidden": hidden,
        "counted": "data rows below the header row; the header is not counted",
        "evaluated": "cached-and-literal values, non-matching rows hidden",
        **({"removed_sheet_filter_over_table": str(stray)} if stray else {})}
    result = pkg.save(allow_loss=allow_loss, backup=backup,
                      verify_com=verify_com)
    if warnings:
        result["warnings"] = list(result.get("warnings", [])) + warnings
    return result


def _table_filtered(ws, table) -> bool:
    """A table filter is on when it holds criteria or hides a data row."""
    af = table.autoFilter
    if af is not None and af.filterColumn:
        return True
    try:
        c0, r0, c1, r1 = _table_filter_rect(ws, table)
    except XlMcpError:
        return False
    return any(ws.row_dimensions[r].hidden for r in range(r0 + 1, r1 + 1))


def clear_filter(path: str, location: Any = None, sheet: str | None = None,
                 allow_loss: bool = False, backup: bool = True,
                 verify_com: bool | None = None) -> dict:
    """Remove the autofilter and unhide the rows it hid. Backup + verify.

    A location inside a table (or no location, on a sheet whose only
    filter is one table's) clears that table's own filter: its criteria go,
    its data rows are shown, its filter buttons stay. A sheet filter that
    overlaps the table (what set_filter used to write, which Excel will not
    open) is removed with it, which repairs such a file."""
    pkg = WorkbookPackage.open(path)
    wb = pkg.workbook
    ws = wb[sheet] if sheet is not None else wb.active
    grid = None
    if location is not None:
        grid = pkg.resolve(location, default_sheet=sheet)
        ws = wb[grid.sheet]
    table = None
    if grid is not None:
        table = _table_for_filter(ws, grid, "clear the filter of")
    elif not ws.auto_filter.ref:
        filtered = [t for t in ws.tables.values() if _table_filtered(ws, t)]
        if len(filtered) > 1:
            names = [t.displayName or t.name for t in filtered]
            raise XlMcpError(
                f"{ws.title!r} has filters on the tables {names}; name the one "
                f"to clear (location {{\"table\": \"{names[0]}\"}}). Nothing "
                "was changed.")
        table = filtered[0] if filtered else None
    if table is not None:
        name = table.displayName or table.name
        c0, r0, c1, r1 = _table_filter_rect(ws, table)
        unhidden = 0
        for r in range(r0 + 1, r1 + 1):
            if ws.row_dimensions[r].hidden:
                ws.row_dimensions[r].hidden = False
                unhidden += 1
        if table.autoFilter is not None:
            table.autoFilter.filterColumn = []
            table.autoFilter.sortState = None
        changed = {"sheet": ws.title, "table": name, "range": _a1((c0, r0, c1, r1)),
                   "rows_unhidden": unhidden}
        stray = ws.auto_filter.ref
        if stray and _intersects(_rect(stray), _rect(table.ref)):
            ws.auto_filter.ref = None
            ws.auto_filter.filterColumn = []
            changed["removed_sheet_filter_over_table"] = str(stray)
        pkg._changed["filter_cleared"] = changed
        return pkg.save(allow_loss=allow_loss, backup=backup,
                        verify_com=verify_com)
    ref = ws.auto_filter.ref
    if grid is not None:
        ref = ws.auto_filter.ref or grid.a1
    if not ref:
        # NOT_FOUND, matching the documented contract (the tool description
        # says "refuses (NOT_FOUND)"); a bare XlMcpError mapped to BAD_PARAMS
        # and contradicted the docs.
        raise TargetNotFound(
            f"no autofilter on sheet {ws.title!r} to clear")
    from openpyxl.utils import range_boundaries
    _c0, r0, _c1, r1 = range_boundaries(str(ref))
    unhidden = 0
    for r in range(r0, r1 + 1):
        if ws.row_dimensions[r].hidden:
            ws.row_dimensions[r].hidden = False
            unhidden += 1
    ws.auto_filter.ref = None
    ws.auto_filter.filterColumn = []
    pkg._changed["filter_cleared"] = {
        "sheet": ws.title, "range": str(ref), "rows_unhidden": unhidden}
    return pkg.save(allow_loss=allow_loss, backup=backup,
                    verify_com=verify_com)


__all__ = ["sort_range", "set_filter", "clear_filter"]
