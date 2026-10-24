"""ops/lifecycle.py: workbook and worksheet lifecycle + discovery.

The orient-before-editing family: create and copy
workbooks, read workbook metadata (sheets, the TRUE used range per sheet, named
ranges, defined tables, and the round-trip hazard summary), manage the sheet
lifecycle (add / delete / rename / copy / reorder / hide), and surface the
hazard scan as a user-facing health readout.

Every mutation routes through WorkbookPackage, so the hazard gate, the
backup-before-mutation, and verify-after-write all run without per-op wiring.
create_workbook writes a brand-new file (no existing content to endanger, so no
scan) but still honors the sandbox and never clobbers an existing path unless
told to. Reads load the workbook directly (no mutation, no backup).
"""

from __future__ import annotations

import os
from typing import Any

from ..core import hazard as _hazard
from ..core import locate as _locate
from ..core import xmlsafe as _xmlsafe
from ..core import calc as _calc
from ..core.errors import TargetNotFound, XlMcpError
from ..core.package import WorkbookPackage
from ..core.sandbox import check_path
from . import format as _format
from . import gridio

WS_ACTIONS = ("add", "delete", "rename", "copy", "reorder", "hide", "unhide")
_HIDE_STATES = {"hidden": "hidden", "very_hidden": "veryHidden",
                "veryhidden": "veryHidden"}


def _check_sheet_name(name: str) -> None:
    """Excel's sheet-name grammar, applied BEFORE openpyxl sees the title.

    This is the single entry point for every sheet title the file tier
    writes (create_workbook, and manage_worksheet add / rename / copy), and
    it delegates to core.limits.check_sheet_title, which owns the measured
    grammar: 1-31 characters, none of : \\ / ? * [ ], no apostrophe at
    either end, and no character Excel cannot store.

    The delegation matters: check_sheet_title was once implemented, exported
    and docstringed but wired to NOTHING, so sheet naming rode on whatever
    openpyxl happened to enforce. openpyxl validates only length and the
    banned charset, and it raises a bare ValueError doing it; the apostrophe
    rule it has no notion of at all, which is how a name starting with one
    used to produce a workbook Excel refuses to OPEN while verify-after-write
    still passed. Refusing here means one message, one exception type, and
    a refusal before any part of the package is touched."""
    from ..core.limits import check_sheet_title
    check_sheet_title(name)


def _refuse_taken(wb, name: str, but: str | None = None) -> None:
    """Excel compares sheet names without regard to case
    ("summary" is "Summary"), and renames a clash on open; refuse it here.
    but is the sheet being renamed, which may change only its own case."""
    for title in wb.sheetnames:
        if title.casefold() == name.casefold() and title != but:
            raise XlMcpError(
                f"a sheet named {title!r} already exists, and Excel treats sheet names "
                f"without regard to case, so {name!r} would clash with it")


# ----------------------------------------------------------------- create/copy


def _refuse_backup_store_target(p: str, what: str) -> None:
    """A workbook written INTO a .ks4xl-backups store overwrites somebody's
    undo; no creation tool may target it."""
    from ..core.safesave import BACKUP_DIR_NAME, in_backup_store
    # Lexical AND junction-resolved.
    if in_backup_store(p):
        raise XlMcpError(
            f"refusing to {what} inside a {BACKUP_DIR_NAME} backup store: "
            "overwriting a backup slot destroys the undo it holds. Choose "
            "a path outside the backup store.")


def _refuse_while_left_aside(p: str, what: str) -> None:
    """A save that did not finish left the workbook at this path aside, and
    it could not be put back (a program has that file open). A new file must
    not take the name meanwhile: the user's workbook, and anything they save
    into that file, would be left stranded beside it."""
    if os.path.exists(p):
        return
    from ..core import swapwait as _swapwait
    found = _swapwait.asides(p)
    if not found:
        return
    name = os.path.basename(p)
    err = XlMcpError(
        f"{name} was not {what}: a save that did not finish left the workbook "
        f"as it was at {found[0]}, and it could not be put back, so a program "
        f"may have that file open. Close it there and retry, or rename that "
        f"file to {name}. Nothing was changed.")
    err.code = "CONFLICT"
    raise err


def _write_atomic(p: str, writer, ticket=None) -> list[str]:
    """Write via a sibling temp file + os.replace. This matters beyond crash
    atomicity: the backup slots are HARDLINKS of the pre-overwrite file, so
    an in-place wb.save/copy2 onto the same inode would rewrite the backup
    it just took. os.replace points the path at a NEW inode and the slot
    keeps the old bytes."""
    import uuid
    tmp = os.path.join(
        os.path.dirname(os.path.abspath(p)) or ".",
        f".ks4xl-write-{uuid.uuid4().hex}{os.path.splitext(p)[1]}")
    try:
        writer(tmp)
        from ..core import deadline as _deadline
        from ..core.safesave import promote
        _deadline.commit_or_raise("the save")  # point of no return
        # A replaced file keeps its identity. The answer
        # carries what the promote reports.
        return promote(tmp, p, ticket=ticket)
    except BaseException as exc:
        if ticket is not None:
            ticket.abort()
        try:
            os.remove(tmp)
        except OSError:
            pass
        if isinstance(exc, PermissionError) and os.path.exists(p):
            # A holder at the swap, or a folder that refuses it: answered in
            # words, as an edit and a restore are, never the raw OS text.
            from ..core import access as _access
            from ..core.errors import WorkbookLocked
            raise WorkbookLocked(_access.replace_refusal_text(
                p, getattr(exc, "winerror", None))) from None
        raise


def _rotate_before_clobber(p: str):
    """Backup-before-mutation for the overwrite=true paths: the workbook
    about to be replaced is staged for its prev/anchor slots, so one
    habitual overwrite:true is no longer unrecoverable ruin. It is
    staged, and landed by the promote only once the replacement went
    through, so a refused overwrite leaves the slots as they were. Returns
    (extra result fields, the ticket)."""
    from ..core import deadline as _deadline
    from ..core import safesave as _safesave
    # The rotation replaces the existing prev slot, so it
    # is the call's point of no return, not the later atomic write (whose
    # own commit is then already made; commits are idempotent).
    _deadline.commit_or_raise("the replacement")
    ticket = _safesave.prepare_rotation(p)
    return {
        "backup": "prev",
        "replaced_existing": True,
        "note": ("the replaced workbook was backed up first: "
                 "manage_backups(action='restore', path=..., source='prev') "
                 "brings it back"),
    }, ticket


def create_workbook(path: str, sheets: list[str] | None = None,
                    overwrite: bool = False) -> dict:
    """Create a new .xlsx at `path` with the named sheets (default one sheet
    'Sheet1'). Refuses an existing path unless overwrite=True; overwrite
    rotates the existing workbook into its backup slots before replacing
    it (backup-before-mutation covers this last bypass too)."""
    p = check_path(path, "create workbook")
    _refuse_backup_store_target(p, "create a workbook")
    # A workbook a stopped save left aside comes back first, so
    # a new file never takes its name.
    from ..core.safesave import put_back_before_write
    put_back_before_write(p)
    _refuse_while_left_aside(p, "created")
    if os.path.exists(p) and not overwrite:
        raise FileExistsError(
            f"{p} already exists; pass overwrite=true to replace it (the "
            "existing file is backed up to its prev slot first)")
    names = [str(s) for s in (sheets or ["Sheet1"]) if str(s).strip()]
    if not names:
        names = ["Sheet1"]
    seen: set[str] = set()
    for n in names:
        low = n.casefold()
        if low in seen:
            raise XlMcpError(f"duplicate sheet name {n!r}")
        seen.add(low)
        _check_sheet_name(n)
    extra: dict = {}
    ticket = None
    if os.path.exists(p) and overwrite:
        extra, ticket = _rotate_before_clobber(p)
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = names[0]
    for n in names[1:]:
        wb.create_sheet(title=n)
    # A new workbook carries neutral metadata, as the Word editor's does.
    # openpyxl stamps every new file's author as "openpyxl"; the
    # author and last-modified-by are left empty, the dates are now.
    import datetime as _dt
    now = _dt.datetime.now(_dt.timezone.utc).replace(tzinfo=None, microsecond=0)
    props = wb.properties
    props.creator = None
    props.lastModifiedBy = None
    props.created = now
    props.modified = now
    promoted = _write_atomic(p, wb.save, ticket)
    result = {"file": p, "created": True, "sheets": names, "saved": True,
              **extra}
    if promoted:
        result["warnings"] = promoted
    return result


def copy_workbook(src: str, dst: str, overwrite: bool = False) -> dict:
    """Copy a workbook file byte-for-byte to a new path (no round-trip through
    openpyxl, so nothing is degraded). Refuses an existing dst unless
    overwrite=True; overwrite rotates the existing dst into its backup
    slots before replacing it."""
    import shutil
    s = check_path(src, "read source workbook")
    d = check_path(dst, "write copy")
    from ..core import swapwait as _swapwait
    if not _swapwait.exists_after_swap(s):
        raise _swapwait.workbook_not_found(s, f"no such workbook: {s}")
    _refuse_backup_store_target(d, "copy a workbook")
    from ..core.safesave import put_back_before_write
    put_back_before_write(d)  # as in create_workbook
    _refuse_while_left_aside(d, "written")
    from ..core.safesave import canonical_key as _ck
    if os.path.exists(d) and _ck(s) == _ck(d):
        raise XlMcpError(
            "src and dst are the same file; nothing to copy")
    if os.path.exists(d) and not overwrite:
        raise FileExistsError(
            f"{d} already exists; pass overwrite=true to replace it (the "
            "existing file is backed up to its prev slot first)")
    extra: dict = {}
    ticket = None
    if os.path.exists(d) and overwrite:
        extra, ticket = _rotate_before_clobber(d)
    promoted = _write_atomic(
        d, lambda tmp: _swapwait.through_swap(s, lambda: shutil.copy2(s, tmp)), ticket)
    result = {"file": d, "copied_from": s, "saved": True, **extra}
    if promoted:
        result["warnings"] = promoted
    return result


# -------------------------------------------------------------- read metadata


def get_workbook_metadata(path: str) -> dict:
    """Read structural metadata without mutating: every sheet with its state and
    TRUE used range, defined names, tables, and the hazard summary."""
    wb = gridio.open_wb(path, data_only=False)
    try:
        sheets = []
        for ws in wb.worksheets:
            bounds = _locate.true_used_range(ws)
            if bounds is None:
                used = None
                dims = {"rows": 0, "cols": 0}
            else:
                r0, c0, r1, c1 = bounds
                used = gridio.a1(r0, c0) + ":" + gridio.a1(r1, c1)
                dims = {"rows": r1 - r0 + 1, "cols": c1 - c0 + 1}
            entry = {
                "name": ws.title,
                "state": ws.sheet_state,
                "used_range": used,
                "dimensions": dims,
                "merged_cells": len(list(getattr(
                    ws, "merged_cells", []).ranges)) if getattr(
                    ws, "merged_cells", None) else 0,
            }
            # Row and column grouping, the read side of the outline
            # surface. Reported only when the sheet actually has groups, so
            # the ordinary workbook's metadata does not grow a key that is
            # always empty; a budget outline is invisible to a caller that
            # only reads values, and a collapsed group hides rows for a
            # reason that has nothing to do with a filter.
            outline = _format.sheet_outline(ws)
            if outline:
                entry["outline"] = outline
            sheets.append(entry)
        # A chart sheet holds one chart and no cells; Excel counts it among
        # the sheets, in its tab position.
        for cs in getattr(wb, "chartsheets", []):
            sheets.append({"name": cs.title, "state": cs.sheet_state,
                           "kind": "chart sheet (one chart, no cells)",
                           "used_range": None,
                           "dimensions": {"rows": 0, "cols": 0},
                           "merged_cells": 0})
        order = {t: i for i, t in enumerate(wb.sheetnames)}
        sheets.sort(key=lambda e: order.get(e["name"], len(order)))
        names = []
        for scope, name, defn in _locate._all_defined_names(wb):
            value = getattr(defn, "value", None)
            # As written, without Excel's storage prefixes (_xlfn., _xlpm.).
            names.append({"name": name, "scope": scope,
                          "value": _calc.denormalize_formula(value)
                          if isinstance(value, str) else value})
        notes = _repeated_sheet_names(names)
        tables = []
        for ws in wb.worksheets:
            tmap = getattr(ws, "tables", {})
            for tname in list(tmap):
                # index by name: TableList.items() yields ref STRINGS, not
                # Table objects (a latent AttributeError: this path crashed
                # on any workbook containing a table)
                t = tmap[tname]
                tables.append({"name": tname, "sheet": ws.title, "ref": t.ref})
        rep = _hazard.scan_path(check_path(path, "scan workbook"))
        # What the cell grid does not show: sheet kinds the
        # model flattens to worksheets, and the drawn objects, notes and
        # extension-block rules on each sheet.
        from ..core import pkginfo as _pkginfo
        try:
            kinds = _pkginfo.sheet_kinds(check_path(path, "open workbook"))
            hidden = _pkginfo.hidden_objects(check_path(path, "open workbook"))
        except Exception:  # noqa: BLE001
            kinds, hidden = {}, {}
        for entry in sheets:
            kind = kinds.get(entry["name"])
            if kind == "macrosheet":
                entry["kind"] = "Excel 4.0 macro sheet"
            elif kind == "dialogsheet":
                entry["kind"] = "Excel 5 dialog sheet"
            if entry["name"] in hidden:
                entry["objects"] = hidden[entry["name"]]
        out = {
            "file": os.path.abspath(path),
            "sheet_count": len(sheets),
            "sheets": sheets,
            "defined_names": names,
            "tables": tables,
            "active_sheet": wb.active.title if wb.active is not None else None,
            "hazards": {
                "clean": rep.clean,
                "would_lose": rep.would_lose,
                "labels": rep.labels(),
            },
        }
        if notes:
            out["warnings"] = notes
        return out
    finally:
        wb.close()


def _repeated_sheet_names(names: list[dict]) -> list[str]:
    """A name scoped to one sheet that exists on several sheets
    points at a different cell on each (Rate on US is US!$B$1, on EU is
    EU!$B$1). Nothing refuses an edit that addresses a cell by sheet, so a
    model that finds one Rate edits it and one that finds both edits both, each
    without asking. The flag says it before the edit."""
    by_name: dict[str, list[dict]] = {}
    for entry in names:
        if entry["scope"] != "workbook":
            by_name.setdefault(str(entry["name"]).casefold(), []).append(entry)
    out = []
    for entries in by_name.values():
        if len(entries) < 2:
            continue
        where = ", ".join(str(e["value"]) for e in entries[:6])
        out.append(
            f"The name {entries[0]['name']} is defined separately on "
            f"{len(entries)} sheets ({where}), so it means a different cell on "
            "each sheet. If the user means one of them, ask which before "
            "changing it.")
    return out


def diagnose_workbook(path: str) -> dict:
    """The hazard scan surfaced as a health readout: which fragile parts the
    file holds, whether an openpyxl edit would drop any of them, the routing
    recommendation, and a light integrity summary."""
    scan_path = check_path(path, "scan workbook")
    from ..core import swapwait as _swapwait
    if not _swapwait.exists_after_swap(scan_path):
        # Another save's swap is waited out, as every read does; a workbook
        # a save that did not finish left aside is named, never moved.
        missing = _swapwait.workbook_not_found(scan_path, "file not found")
        return {"file": os.path.abspath(path), "readable": False,
                "error": str(missing)}
    rep = _hazard.scan_path(scan_path)
    if getattr(rep, "error_kind", None) == "missing" and \
            _swapwait.exists_after_swap(scan_path):
        rep = _hazard.scan_path(scan_path)  # a swap began during the scan
    if rep.error is not None:
        return {"file": os.path.abspath(path), "readable": False,
                "error": rep.error}
    from ..core import routes as _routes
    from ..core import surgical as _surgical

    # The structural answer is the save gate's own verdict on this file
    # (run, not re-derived), and the cell answer is whether the splice can
    # read its worksheets.
    routing = _routes.predict(
        rep, verdict=_routes.gate_verdict(scan_path),
        splice_problem=_surgical.splice_check(scan_path))
    wb = gridio.open_wb(path, data_only=False)
    try:
        visible = [ws for ws in wb.worksheets
                   if ws.sheet_state == "visible"]
        formula_cells = 0
        for ws in wb.worksheets:
            cells = getattr(ws, "_cells", None)
            if cells is None:
                continue
            for cell in cells.values():
                if _calc.is_formula_cell(cell):
                    formula_cells += 1
        health = {
            "sheet_count": len(wb.worksheets),
            "visible_sheets": len(visible),
            "formula_cells": formula_cells,
            "keep_vba": _xmlsafe.keeps_vba(scan_path),
        }
    finally:
        wb.close()
    return {
        "file": os.path.abspath(path),
        "readable": True,
        "hazards": rep.as_dict(),
        "routing": {
            # The two original keys, now answered by the same decision the
            # save path makes (core/routes.py): what a cell edit and what a
            # structural edit will actually do to THIS file.
            "surgical_edit": {"route": routing["cell_edits"]["route"],
                              "reason": routing["cell_edits"]["reason"]},
            "structural_edit": {
                "route": routing["structural_edits"]["route"],
                "reason": routing["structural_edits"]["reason"]},
            **routing,
        },
        "health": health,
    }


# --------------------------------------------------------------- sheet manage


def manage_worksheet(path: str, action: str, sheet: str | None = None,
                     new_name: str | None = None, index: int | None = None,
                     state: str | None = None, allow_loss: bool = False,
                     backup: bool = True,
                     verify_com: bool | None = None) -> dict:
    """Add / delete / rename / copy / reorder / hide a worksheet. One backup +
    one verified save per call through WorkbookPackage."""
    if action not in WS_ACTIONS:
        raise XlMcpError(
            f"action must be one of {WS_ACTIONS}, got {action!r}")
    pkg = WorkbookPackage.open(path)
    wb = pkg.workbook

    def _need(name: str | None, what: str) -> str:
        if not name:
            raise XlMcpError(f"{action} needs {what}")
        if name not in wb.sheetnames:
            raise TargetNotFound(
                f"no sheet named {name!r} (sheets: "
                f"{', '.join(wb.sheetnames)})")
        return name

    detail: dict[str, Any] = {"action": action}
    if action == "add":
        name = new_name or sheet
        if not name:
            # It used to default to "Sheet", so a call that forgot new_name
            # silently created a differently named sheet than the caller
            # meant, and only the SECOND such call ever complained.
            raise XlMcpError(
                "add needs new_name (the title for the new sheet)")
        _refuse_taken(wb, name)
        _check_sheet_name(name)
        pos = index if index is not None else len(wb.sheetnames)
        ws = wb.create_sheet(title=name, index=pos)
        detail.update(sheet=name, index=wb.sheetnames.index(ws.title))
    elif action == "delete":
        name = _need(sheet, "the sheet to delete")
        if len(wb.sheetnames) == 1:
            raise XlMcpError(
                "cannot delete the only sheet; a workbook needs one visible "
                "sheet")
        visible = [w for w in wb.worksheets if w.sheet_state == "visible"]
        if len(visible) == 1 and visible[0].title == name:
            raise XlMcpError(
                f"cannot delete {name!r}: it is the only VISIBLE sheet (the "
                "others are hidden) and a workbook needs one visible sheet; "
                "unhide another sheet first")
        order = list(wb.sheetnames)
        del wb[name]
        # References to the deleted sheet take Excel's own
        # result: =North!F5 becomes =#REF!F5, and a 3D span whose end it was
        # shrinks to the next sheet inside the span.
        from ..core import refs as _refs
        rewrites = _refs.delete_sheet_refs(wb, name, order)
        detail["references_rewritten"] = rewrites.as_dict()
        # A deleted sheet takes its own part plus any per-sheet companion
        # parts (comments, comment VML anchors, tables) with it; declare the
        # removal so the default-fail inventory verify does not read it as
        # silent loss. Fragile per-sheet parts (charts, pivots) stay under
        # the hazard gate and still need allow_loss or COM.
        pkg.expect_removal("xl/worksheets/", "xl/chartsheets/",
                           "xl/comments", "xl/drawings/commentsdrawing",
                           "xl/tables/")
        detail.update(sheet=name)
    elif action == "rename":
        name = _need(sheet, "the sheet to rename")
        if not new_name:
            raise XlMcpError("rename needs new_name")
        _check_sheet_name(new_name)
        _refuse_taken(wb, new_name, but=name)
        if new_name != name and new_name.casefold() == name.casefold():
            # openpyxl's title setter compares without case against every
            # sheet, this one included, and would append a 1; a case-only
            # rename goes through a name no sheet has.
            wb[name].title = f"ks4xl-rename-{id(wb)}"[:31]
            wb[f"ks4xl-rename-{id(wb)}"[:31]].title = new_name
        else:
            wb[name].title = new_name
        # Every reference to the sheet follows the new
        # name, as Excel's rename does: formulas (3D spans included), defined
        # names, conditional-format and validation formulas, chart references
        # and pivot sources; quoted where Excel quotes ('West Coast'!F2).
        # Hyperlink locations are left alone, as Excel leaves them.
        from ..core import refs as _refs
        rewrites = _refs.rename_sheet_refs(wb, name, new_name)
        # The rename changes formula text, never a result: every formula it
        # rewrote keeps its stored result, as in Excel's own rename (an
        # iterative model seeds from them).
        pkg.note_text_rewrite(
            lambda t, o=name, n=new_name: _refs.rename_sheet_in_formula(t, o, n))
        detail.update(sheet=name, new_name=new_name,
                      references_rewritten=rewrites.as_dict())
    elif action == "copy":
        name = _need(sheet, "the sheet to copy")
        src = wb[name]
        if new_name:
            _check_sheet_name(new_name)
            _refuse_taken(wb, new_name)
        dup = wb.copy_worksheet(src)
        if new_name:
            dup.title = new_name
        detail.update(sheet=name, new_name=dup.title)
    elif action == "reorder":
        name = _need(sheet, "the sheet to move")
        if index is None:
            raise XlMcpError("reorder needs index (0-based target position)")
        target = max(0, min(int(index), len(wb.sheetnames) - 1))
        cur = wb.sheetnames.index(name)
        wb.move_sheet(name, offset=target - cur)
        detail.update(sheet=name, index=target)
    elif action == "hide":
        name = _need(sheet, "the sheet to hide")
        st = _HIDE_STATES.get((state or "hidden").lower())
        if st is None:
            raise XlMcpError("state must be 'hidden' or 'very_hidden'")
        visible = [w for w in wb.worksheets if w.sheet_state == "visible"]
        if len(visible) == 1 and visible[0].title == name:
            raise XlMcpError(
                "cannot hide the only visible sheet; a workbook needs one")
        wb[name].sheet_state = st
        detail.update(sheet=name, state=st)
    else:  # unhide
        name = _need(sheet, "the sheet to unhide")
        wb[name].sheet_state = "visible"
        detail.update(sheet=name, state="visible")

    pkg._changed["worksheet"] = detail
    result = pkg.save(allow_loss=allow_loss, backup=backup,
                      verify_com=verify_com)
    return result


__all__ = [
    "create_workbook", "copy_workbook", "get_workbook_metadata",
    "diagnose_workbook", "manage_worksheet", "WS_ACTIONS",
]
