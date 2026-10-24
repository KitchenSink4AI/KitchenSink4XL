"""A picture's hyperlink survives a full rewrite.

openpyxl keeps a picture's link id in the drawing XML (a:hlinkClick r:id)
but writes only the image relationships, so after a full rewrite the id
named nothing and the save was refused ("drawing refers to relationship
rId3, which its relationship part does not define"), or the id named the
image itself. The rewrite now puts each link back under a fresh id, with the
target the original gave it. All content is made up.
"""

from __future__ import annotations

import io
import re
import zipfile

import openpyxl
import pytest
from openpyxl.drawing.image import Image

from xlsx_mcp.ops import structure

LINK = "https://www.example.com/catalogue/item-7"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
HL_TYPE = R_NS + "/hyperlink"


def _png() -> bytes:
    from PIL import Image as PILImage
    buf = io.BytesIO()
    PILImage.new("RGB", (8, 8), (200, 40, 40)).save(buf, "PNG")
    return buf.getvalue()


def _book(tmp_path, link_id: str):
    p = tmp_path / "pictures.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pictures"
    ws["A1"] = "Item"
    ws["A2"] = "Widget"
    ws.add_image(Image(io.BytesIO(_png())), "C2")
    wb.save(p)
    with zipfile.ZipFile(p) as zin:
        items = [(i, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            if info.filename == "xl/drawings/drawing1.xml":
                x = data.decode("utf-8")
                x = re.sub(
                    r"(<(\w+:)?cNvPr\b[^>]*?)/>",
                    lambda m: (f'{m.group(1)}><a:hlinkClick xmlns:a="{A_NS}" '
                               f'xmlns:r="{R_NS}" r:id="{link_id}"/>'
                               f"</{m.group(2) or ''}cNvPr>"), x, count=1)
                data = x.encode("utf-8")
            elif info.filename == "xl/drawings/_rels/drawing1.xml.rels":
                x = data.decode("utf-8")
                # the image keeps the other id, the link takes link_id
                other = "rId2" if link_id == "rId1" else "rId1"
                x = re.sub(r'Id="rId1"', f'Id="{other}"', x)
                d = data.decode("utf-8")
                x = x.replace(
                    "</Relationships>",
                    f'<Relationship Id="{link_id}" Type="{HL_TYPE}" '
                    f'Target="{LINK}" TargetMode="External"/></Relationships>')
                data = x.encode("utf-8")
            zout.writestr(info, data)
    # the drawing's image reference follows its relationship id
    with zipfile.ZipFile(p) as zin:
        items = [(i, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            if info.filename == "xl/drawings/drawing1.xml":
                other = "rId2" if link_id == "rId1" else "rId1"
                data = re.sub(rb'r:embed="rId1"', f'r:embed="{other}"'.encode(),
                              data)
            zout.writestr(info, data)
    return str(p)


def _link_after(p: str) -> list[tuple[str, str]]:
    """(relationship type, target) of every picture link in the drawing."""
    with zipfile.ZipFile(p) as z:
        dr = next(n for n in z.namelist()
                  if re.fullmatch(r"xl/drawings/drawing\d+\.xml", n))
        x = z.read(dr).decode("utf-8")
        rels = z.read(dr.replace("drawings/", "drawings/_rels/") + ".rels"
                      ).decode("utf-8")
    out = []
    for rid in re.findall(r'hlinkClick\b[^>]*?r:id="([^"]+)"', x):
        m = re.search(r'<Relationship\b[^>]*Id="%s"[^>]*>' % re.escape(rid),
                      rels)
        assert m, f"{rid} is not defined in the drawing's relationships"
        t = re.search(r'Type="([^"]+)"', m.group(0)).group(1)
        tg = re.search(r'Target="([^"]+)"', m.group(0)).group(1)
        out.append((t, tg))
    return out


@pytest.mark.parametrize("link_id", ["rId5", "rId1"])
def test_the_link_is_kept_through_a_structural_rewrite(tmp_path, link_id):
    p = _book(tmp_path, link_id)
    assert _link_after(p) == [(HL_TYPE, LINK)]   # the made-up book is sound
    structure.modify_grid_structure(p, "insert_rows", 1, 1, sheet="Pictures",
                                    allow_loss=True)
    assert _link_after(p) == [(HL_TYPE, LINK)]
    wb = openpyxl.load_workbook(p)
    assert wb["Pictures"]["A3"].value == "Widget"
