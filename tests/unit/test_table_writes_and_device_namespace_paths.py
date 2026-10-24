"""Pre-existing defect classes found while attacking the array-formula fix
with Excel as the oracle: a table that must stay openable after an accepted
write, and device-namespace paths that refuse instead of hanging. Red before
the fix, and identically red on released 1.2.5.
"""

from __future__ import annotations

import json
import sys

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp import server as srv
from xlsx_mcp.com import session as com_session

_COM = com_session.com_available()[0]


def _table_book(p) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B1"], ws["B2"] = 1, 2
    ws["D1"], ws["E1"], ws["D2"], ws["E2"], ws["D3"], ws["E3"] = "k", "v", "a", 1, "b", 2
    wb.save(p)
    srv.create_table(str(p), "D1:E3", "T")
    return str(p)


# F-T (MAJOR, pre-existing on 1.2.5; P1) a write that changes a table's header
# cell (any value, even other text) or puts a multi-cell array formula inside
# a table saves "verified" with verify_com off (the default), and Excel then
# refuses to open the workbook without a repair (com_validate_opens_clean:
# opens_clean false; with verify_com:true the same save is refused and
# restored). Judged by the product's own opens-clean check in the private pool.

_ROUTES = {
    "header text": ("xl_edit", "set_cell", {"location": "E1", "value": "w"}),
    "header number": ("xl_edit", "set_cell", {"location": "E1", "value": 5}),
    "header formula": ("xl_edit", "set_formula", {"location": "E1", "formula": "=1+1"}),
    "plain moved onto header": ("xl_edit", "move_range", {"source": "B1", "dest": "E1"}),
}


@pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool) as the judge")
@pytest.mark.timeout(300)
@pytest.mark.parametrize("route", sorted(_ROUTES))
def test_ft_a_write_the_product_accepts_leaves_the_table_openable(route, tmp_path):
    p = _table_book(tmp_path / "t.xlsx")
    tool, op, args = _ROUTES[route]
    payload, err, _ = fh.start().call(tool, {"op": op, "args": {"path": p, **args}})
    if err:
        return  # a refusal is an acceptable answer
    verdict = srv.com_validate_opens_clean(p)
    assert verdict["opens_clean"], f"{route}: answered {payload.get('verified')=} but Excel says {verdict}"


# A device-namespace path that is not a reserved name but names a
# console device (\.\GLOBALROOT\Device\ConDrv\Console) still blocks on a
# read until the call deadline. Reserved names in every spelling now refuse.

@pytest.mark.skipif(sys.platform != "win32", reason="Windows device namespace")
@pytest.mark.timeout(120)
def test_fd3_a_device_namespace_input_refuses_instead_of_hanging():
    b = chr(92)
    path = b * 2 + "." + b + "GLOBALROOT" + b + "Device" + b + "ConDrv" + b + "Console"
    res = fh.stdio_session([("xl_read", {"op": "read_range", "args": {"path": path, "location": "A1"}})],
                           env={"KS4XL_CALL_TIMEOUT_S": "3"}, timeout=90)
    assert json.loads(res["results"][0]["text"])["error"]["code"] != "TIMEOUT"
