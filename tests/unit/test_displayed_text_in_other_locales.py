"""Displayed dates in other locales: Excel's own names, or the answer says so.

A date format with a locale code ([$-412] Korean, [$-411] Japanese era,
[$-407] German, [$-40C] French) was shown with English day and month names,
and the Japanese era code came back as the letters "ggge". The names are
now Excel's own, measured with Excel's Range.Text for fifteen locales; for
a locale or calendar not measured, a read gives the cell's value and a
display_note naming the cell, never a guessed text. Dates are made up.
"""

from __future__ import annotations

import datetime as dt

import openpyxl
import pytest

from xlsx_mcp.core import numfmt
from xlsx_mcp.ops import cells

FRI = 46073  # Friday, 20 February 2026

#: (value, format, Excel's text) measured in Excel
EXCEL = [
    (FRI, '[$-412]yyyy"년" m"월" d"일" dddd',
     "2026년 2월 20일 금요일"),
    (FRI, '[$-411]ggge"年"m"月"d"日"',
     "令和8年2月20日"),
    (FRI, "[$-407]dddd, d. mmmm yyyy", "Freitag, 20. Februar 2026"),
    (FRI, "[$-40C]d mmmm yyyy", "20 février 2026"),
    (FRI, '[$-804]yyyy"年"m"月"d"日" dddd',
     "2026年2月20日 星期五"),
    (FRI, "[$-409]mmmm d, yyyy", "February 20, 2026"),
    (FRI, "[$-809]ddd d mmm", "Fri 20 Feb"),
    (FRI, "[$-40C]ddd mmm", "ven. févr."),
    (FRI, "[$-407]ddd mmm mmmmm", "Fr Feb F"),
    (FRI, "[$-419]dddd mmmm", "пятница "
          "Февраль"),
    (FRI + 0.75, "[$-412]AM/PM h:mm", "오후 6:00"),
    (FRI + 0.75, "[$-407]h:mm AM/PM", "6:00 PM"),
    (FRI, "[$-411]ge", "R8"), (FRI, "[$-411]gggee", "令和08"),
    (32515, "[$-411]ggge", "昭和64"),   # 7 January 1989
    (32516, "[$-411]ggge", "平成1"),    # 8 January 1989
    (4595, "[$-411]gge", "大1"),            # 30 July 1912
    (FRI, "[$-ja-JP]yyyy/m/d", "2026/2/20"),
    (FRI, "[$-F800]dddd, mmmm dd, yyyy", "Friday, February 20, 2026"),
]


@pytest.mark.parametrize("value,fmt,excel", EXCEL,
                         ids=[f"{f}" for _v, f, _x in EXCEL])
def test_names_are_excels_own(value, fmt, excel):
    assert numfmt.display_checked(value, fmt) == (excel, None)


@pytest.mark.parametrize("fmt", [
    "[$-42A]dddd d mmmm yyyy",       # a locale whose names were not measured
    "[$-40C]h:mm AM/PM",             # AM/PM in a locale not measured for it
    "[$-1010409]dddd",               # another calendar
    "[$-409]ggge",                   # an era outside Japanese
    "[$-xx-YY]mmmm",                 # an unknown locale name
])
def test_a_locale_that_was_not_measured_is_named_not_guessed(fmt):
    text, why = numfmt.display_checked(FRI, fmt)
    assert text is None
    assert why


def test_numbers_and_digits_need_no_locale():
    assert numfmt.display_checked(1234.5, "[$-407]#,##0.00") == \
        ("1,234.50", None)
    assert numfmt.display_checked(42055, "[$-42A]yyyy-mm-dd") == \
        ("2015-02-20", None)  # 42055: Friday, 20 February 2015


def test_a_read_says_so_and_gives_the_value(tmp_path):
    p = tmp_path / "loc.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = dt.datetime(2015, 2, 20)
    ws["A1"].number_format = "[$-42A]dddd d mmmm yyyy"
    ws["A2"] = dt.datetime(2015, 2, 20)
    ws["A2"].number_format = "[$-407]dddd"
    wb.save(p)
    out = cells.read_range(str(p), "A1:A2", sheet="Data", values="displayed")
    assert out["values"] == [["2015-02-20"], ["Freitag"]]
    assert "A1" in out["display_note"] and "A2" not in out["display_note"]
    got = cells.get_cells(str(p), ["A1"], sheet="Data", values="displayed")
    assert "display_note" in got
