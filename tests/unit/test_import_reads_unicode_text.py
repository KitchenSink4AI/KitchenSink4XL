"""import_data reads Excel's own "Unicode Text" files.

Excel saves "Unicode Text" as tab-separated UTF-16 with a byte-order mark.
import_data opened every file as UTF-8 and failed with the raw codec text
("'utf-8' codec can't decode byte 0xff in position 0"). The byte-order mark
now decides the encoding, and a file that is not in the encoding asked for
is refused in plain words that say what to pass. All content is made up.
"""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core.errors import XlMcpError
from xlsx_mcp.ops import dataio

ROWS = "Region\tUnits\r\nNorth\t12\r\nSüd\t7\r\n"


def _book(tmp_path):
    p = tmp_path / "target.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Data"
    wb.save(p)
    return str(p)


def _cells(p):
    ws = openpyxl.load_workbook(p)["Data"]
    return [[c.value for c in row] for row in ws.iter_rows(min_row=1,
                                                         max_row=3)]


@pytest.mark.parametrize("bom,codec", [
    (b"\xff\xfe", "utf-16-le"),     # what Excel writes on Windows
    (b"\xfe\xff", "utf-16-be"),
])
def test_a_utf16_text_file_imports_by_its_byte_order_mark(tmp_path, bom,
                                                          codec):
    src = tmp_path / "unicode.txt"
    src.write_bytes(bom + ROWS.encode(codec))
    p = _book(tmp_path)
    dataio.import_data(p, source_file=str(src))
    assert _cells(p) == [["Region", "Units"], ["North", 12], ["Süd", 7]]


def test_a_utf8_bom_file_still_imports(tmp_path):
    src = tmp_path / "utf8.txt"
    src.write_bytes(b"\xef\xbb\xbf" + ROWS.encode("utf-8"))
    p = _book(tmp_path)
    dataio.import_data(p, source_file=str(src))
    assert _cells(p)[2] == ["Süd", 7]


def test_a_file_not_in_the_encoding_asked_for_is_refused_in_plain_words(
        tmp_path):
    src = tmp_path / "legacy.txt"
    src.write_bytes(ROWS.encode("cp1252"))
    p = _book(tmp_path)
    with pytest.raises(XlMcpError) as ei:
        dataio.import_data(p, source_file=str(src))
    msg = str(ei.value)
    assert "encoding" in msg and "cp1252" in msg
    assert "codec can't decode" not in msg
