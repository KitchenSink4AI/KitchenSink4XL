"""Every COM route that writes the workbook keeps what Windows keeps on the
file itself: its Mark of the Web, its permissions and its creation time.

Each route has Excel save a copy, checks the copy, and puts it in place with
safesave.promote, which uses ReplaceFileW: the replaced file's permissions,
creation time and named streams (the Mark of the Web is the Zone.Identifier
stream) carry over to the new version. A workbook downloaded from the
internet (ZoneId=3) must stay marked, so Excel keeps opening it in
Protected View and Office's block on macros from the internet still
applies, and a workbook only its owner may open must stay that way.

Checked here for each kind of route, with an owner-only permission list
(inheritance removed) and ZoneId=3: an edit (recalculate), an encrypted
save and its removal, the library recalculation (calc.recalc_via_com), a
copy that fails its check (nothing is replaced), and the fallback used
where ReplaceFileW is not supported (an os.replace, which gives the path a
new file: the identity is carried over by hand).
"""
from __future__ import annotations

import os
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl
import pytest

from xlsx_mcp.com import session as com_session
from xlsx_mcp.core import safesave
from xlsx_mcp.core.errors import ValidationFailed
from xlsx_mcp.ops import comtier

pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="Windows file identity")

NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)
SYS32 = os.path.join(os.environ.get("SystemRoot", r"C:\Windows"), "System32")
CFB = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _run(*a):
    return subprocess.run(list(a), capture_output=True, text=True, creationflags=NO_WINDOW,
                          stdin=subprocess.DEVNULL, timeout=60).stdout


def _user_sid() -> str:
    out = _run(os.path.join(SYS32, "whoami.exe"), "/user", "/fo", "csv", "/nh")
    return out.strip().split(",")[-1].strip().strip('"')


def _aces(p) -> list:
    """Every entry of the file's permission list, inherited ones marked (I)."""
    lines = [ln.strip() for ln in _run(os.path.join(SYS32, "icacls.exe"), p).splitlines()
             if ln.strip() and "Successfully" not in ln]
    lines[0] = lines[0].replace(p, "").strip()
    return sorted(lines)


def _identity(p) -> dict:
    try:
        with open(p + ":Zone.Identifier", encoding="utf-8") as f:
            motw = f.read()
    except OSError:
        motw = None
    return {"zone": motw, "aces": _aces(p), "born": os.stat(p).st_birthtime}


def _mark(p):
    """Downloaded from the internet, and only its owner may open it."""
    with open(p + ":Zone.Identifier", "w", encoding="utf-8") as f:
        f.write("[ZoneTransfer]\nZoneId=3\n")
    _run(os.path.join(SYS32, "icacls.exe"), p, "/inheritance:r", "/grant:r", f"*{_user_sid()}:(F)")
    ident = _identity(p)
    assert "ZoneId=3" in ident["zone"]
    # Owner only: nothing inherited, and the user's own entry is there. An
    # elevated account's new files also carry explicit entries for
    # Administrators, SYSTEM and OWNER RIGHTS; they stay in the comparison
    # (an edit must keep them too) and are not removed here.
    assert not any("(I)" in ace for ace in ident["aces"]), ident["aces"]
    assert any(ace.endswith(":(F)") and not ace.startswith(("BUILTIN", "NT AUTHORITY", "OWNER RIGHTS"))
               for ace in ident["aces"]), ident["aces"]
    return ident


@pytest.fixture
def book(tmp_path, monkeypatch):
    folder = tmp_path / "work"
    folder.mkdir()
    p = str(folder / "book.xlsx")
    wb = openpyxl.Workbook()
    wb.active["A1"] = 2
    wb.active["A2"] = "=A1*21"
    wb.save(p)
    temp = tmp_path / "temp"
    temp.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(temp))
    yield p
    _run(os.path.join(SYS32, "icacls.exe"), p, "/reset")


def _beside(book) -> list:
    return sorted(n for n in os.listdir(os.path.dirname(book))
                  if n not in (os.path.basename(book), safesave.BACKUP_DIR_NAME)
                  and not n.startswith("~$"))


def _com_ok() -> bool:
    try:
        return bool(com_session.com_available()[0])
    except Exception:  # noqa: BLE001
        return False


def _shut_down_excel():
    try:
        com_session.get_executor().shutdown()
    except Exception:  # noqa: BLE001
        pass
    com_session._EXECUTOR = None


# ------------------------------------------------------------- real Excel


@pytest.mark.com
@pytest.mark.skipif(not _com_ok(), reason="needs Excel (the product's private hidden pool)")
@pytest.mark.timeout(300)
def test_live_an_edit_keeps_the_identity(book):
    """recalculate(engine='com'), the _run_mutation route: Excel's copy is
    checked, then put in place."""
    before = _mark(book)
    try:
        res = comtier.recalculate(book, engine="com", timeout_seconds=90)
        assert res["ok"] and res["verified"]
        assert openpyxl.load_workbook(book, data_only=True).active["A2"].value == 42
        assert _identity(book) == before
        assert _beside(book) == []
    finally:
        _shut_down_excel()


@pytest.mark.com
@pytest.mark.skipif(not _com_ok(), reason="needs Excel (the product's private hidden pool)")
@pytest.mark.timeout(300)
def test_live_an_encrypted_save_and_its_removal_keep_the_identity(book):
    before = _mark(book)
    try:
        res = comtier.com_save_with_password(book, "pw-rc7", timeout_seconds=90)
        assert res["ok"] and Path(book).read_bytes()[:8] == CFB
        assert _identity(book) == before
        res = comtier.com_save_with_password(book, "", current_password="pw-rc7",
                                             timeout_seconds=90)
        assert res["ok"] and res["verified"] and zipfile.is_zipfile(book)
        assert _identity(book) == before
        assert _beside(book) == []
    finally:
        _shut_down_excel()


@pytest.mark.com
@pytest.mark.skipif(not _com_ok(), reason="needs Excel (the product's private hidden pool)")
@pytest.mark.timeout(300)
def test_live_the_library_recalculation_keeps_the_identity(book, tmp_path):
    from xlsx_mcp.com.instances import ExcelInstanceManager
    from xlsx_mcp.core import calc

    before = _mark(book)
    mgr = ExcelInstanceManager(str(tmp_path / "journal.json"))
    try:
        rr = calc.recalc_via_com(book, mgr)
        assert rr.ok, rr.note
        assert openpyxl.load_workbook(book, data_only=True).active["A2"].value == 42
        assert _identity(book) == before
    finally:
        for w in list(mgr._workers):
            try:
                mgr.quit(w)
            except Exception:  # noqa: BLE001
                pass
        mgr.force_reclaim()


# ------------------------------------------------------------- fake Excel


class _Worker:
    pid = 0
    app = object()


class _Manager:
    def acquire(self):
        return _Worker()

    def release_to_pool(self, w):
        pass


class _FakeWB:
    def __init__(self, path, garbage=False):
        self.path = path
        self.FullName = os.path.abspath(path)
        self.ReadOnly = False
        self.Password = ""
        self.garbage = garbage

    def SaveCopyAs(self, copy):
        if self.garbage:
            Path(copy).write_bytes(b"PK\x03\x04 not a workbook")
            return
        wb = openpyxl.load_workbook(self.path)
        wb.active["B1"] = "edited by Excel"
        wb.save(copy)

    def Close(self, save):
        pass


def _fake_excel(monkeypatch, garbage=False):
    monkeypatch.setattr(com_session, "run_com",
                        lambda label, job, timeout=None: job(_Manager()))
    monkeypatch.setattr(com_session, "_assert_hygiene", lambda app: ())
    monkeypatch.setattr(com_session, "_restore_hygiene", lambda app, *a: None)
    monkeypatch.setattr(com_session, "com_retry", lambda fn, label=None, **k: fn())
    monkeypatch.setattr(com_session, "open_workbook",
                        lambda app, p, password=None, **k: _FakeWB(p, garbage))


def test_a_copy_that_fails_its_check_leaves_the_workbook_and_its_identity(book, monkeypatch):
    before = _mark(book)
    data = Path(book).read_bytes()
    _fake_excel(monkeypatch, garbage=True)
    with pytest.raises(ValidationFailed):
        comtier._run_mutation("autofit(book.xlsx)", book, lambda app, wb: {"a": 1},
                              backup=False)
    assert Path(book).read_bytes() == data
    assert _identity(book) == before
    assert _beside(book) == []


def test_where_replacefile_is_unsupported_the_identity_is_carried_over(book, monkeypatch):
    """Some network and WebDAV drives answer ReplaceFileW with 'not
    supported'; the promote then renames the new version over the
    workbook, which gives the path a new file. The Mark of the Web, the
    permissions and the creation time are carried over to it first."""
    before = _mark(book)
    fn, move, last_error, text = safesave._replace_file()
    calls = []

    def unsupported(ldst, lsrc, lbak, flags, a, b):
        calls.append(ldst)
        return False

    monkeypatch.setattr(safesave, "_REPLACE_FILE", (unsupported, move, lambda: 50, text))
    _fake_excel(monkeypatch)
    res = comtier._run_mutation("autofit(book.xlsx)", book, lambda app, wb: {"a": 1})
    assert res["ok"] and res["verified"] and calls, "ReplaceFileW was not tried"
    assert openpyxl.load_workbook(book).active["B1"].value == "edited by Excel"
    assert _identity(book) == before
    assert _beside(book) == []


def test_the_fallback_keeps_explicit_entries_on_an_inheriting_list(book, monkeypatch):
    """A permission list that still takes the folder's entries, plus an
    explicit grant and an explicit deny: both explicit entries are carried
    over and the folder's entries apply as before."""
    _run(os.path.join(SYS32, "icacls.exe"), book, "/grant", "*S-1-5-32-545:(R)")
    _run(os.path.join(SYS32, "icacls.exe"), book, "/deny", "*S-1-5-7:(R)")
    with open(book + ":Zone.Identifier", "w", encoding="utf-8") as f:
        f.write("[ZoneTransfer]\nZoneId=3\n")
    before = _identity(book)
    assert any("(I)" in a for a in before["aces"]) and len(before["aces"]) >= 4, before
    fn, move, last_error, text = safesave._replace_file()
    monkeypatch.setattr(safesave, "_REPLACE_FILE",
                        (lambda *a: False, move, lambda: 50, text))
    _fake_excel(monkeypatch)
    assert comtier._run_mutation("autofit(book.xlsx)", book, lambda app, wb: {"a": 1})["ok"]
    assert _identity(book) == before
