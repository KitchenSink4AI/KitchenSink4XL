"""Unit tests for the COM tier's headless-safe surface.

NO test here spawns Excel (the suite must stay green in headless CI and must
never touch the author's machine's Excel outside the com_gates). Covered:
argument validation and refusal paths that fire BEFORE any COM call, the
executor's serialization/timeout/re-arm mechanics driven with plain Python
functions, the stale-lockfile guard, com_status shape, the verify_com wiring
in WorkbookPackage.save (with com availability monkeypatched), and the
iterative-calc file-tier remainder of the calc-settings surface.
"""

from __future__ import annotations

import threading
import time
from pathlib import Path

import openpyxl
import pytest

from xlsx_mcp.com import session as com_session
from xlsx_mcp.core.errors import (
    CalcUnavailable,
    ExcelBlocked,
    ValidationFailed,
    WorkbookLocked,
    XlMcpError,
)
from xlsx_mcp.core.package import WorkbookPackage
from xlsx_mcp.ops import comtier, properties


def _book(path: Path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 1
    ws["A2"] = 2
    ws["A3"] = "=A1+A2"
    wb.save(path)
    wb.close()
    return str(path)


# ------------------------------------------------------ argument refusals


def test_sparkline_source_qualified_to_group_sheet():
    # Field finding (MAJOR): an unqualified source bound to the ACTIVE sheet,
    # not the group's sheet, so a sparkline on "Sparklines" charted (and listed)
    # the wrong sheet's data. An unqualified ref now gets the group's own sheet.
    assert comtier._qualify_sparkline_source("B2:F10", "Sparklines") \
        == "Sparklines!B2:F10"
    # A sheet name needing quoting is quoted.
    assert comtier._qualify_sparkline_source("B2:F10", "My Data") \
        == "'My Data'!B2:F10"
    # A deliberate cross-sheet source already carries its sheet: untouched.
    assert comtier._qualify_sparkline_source("Data!B2:F10", "Sparklines") \
        == "Data!B2:F10"


def test_pivot_bad_action(tmp_path):
    p = _book(tmp_path / "b.xlsx")
    with pytest.raises(XlMcpError, match="action must be"):
        comtier.com_manage_pivot(p, "explode")


def test_pivot_create_needs_source_and_values(tmp_path):
    p = _book(tmp_path / "b.xlsx")
    with pytest.raises(XlMcpError, match="source_range"):
        comtier.com_manage_pivot(p, "create")
    with pytest.raises(XlMcpError, match="values"):
        comtier.com_manage_pivot(p, "create", source_range="A1:B4")
    with pytest.raises(XlMcpError, match="aggregation"):
        comtier.com_manage_pivot(
            p, "create", source_range="A1:B4",
            values=[{"field": "x", "func": "median"}])


def test_pivot_delete_needs_name(tmp_path):
    p = _book(tmp_path / "b.xlsx")
    with pytest.raises(XlMcpError, match="name"):
        comtier.com_manage_pivot(p, "delete")


def test_recalculate_bad_engine(tmp_path):
    p = _book(tmp_path / "b.xlsx")
    with pytest.raises(XlMcpError, match="engine"):
        comtier.recalculate(p, engine="prayer")


def test_recalculate_com_refuses_when_unavailable(tmp_path, monkeypatch):
    p = _book(tmp_path / "b.xlsx")
    monkeypatch.setattr(com_session, "com_available",
                        lambda refresh=False: (False, "test says no"))
    with pytest.raises(CalcUnavailable, match="test says no"):
        comtier.recalculate(p, engine="com")


def test_export_pdf_bad_scope_and_existing_output(tmp_path):
    p = _book(tmp_path / "b.xlsx")
    with pytest.raises(XlMcpError, match="scope"):
        comtier.com_export_pdf(p, str(tmp_path / "o.pdf"), scope="galaxy")
    with pytest.raises(XlMcpError, match="range_a1"):
        comtier.com_export_pdf(p, str(tmp_path / "o.pdf"), scope="range")
    out = tmp_path / "exists.pdf"
    out.write_bytes(b"x")
    with pytest.raises(FileExistsError):
        comtier.com_export_pdf(p, str(out))


def test_render_requires_png(tmp_path):
    p = _book(tmp_path / "b.xlsx")
    with pytest.raises(XlMcpError, match=".png"):
        comtier.com_render_sheet(p, str(tmp_path / "o.bmp"))


def test_convert_bad_format_and_same_path(tmp_path):
    p = _book(tmp_path / "b.xlsx")
    with pytest.raises(XlMcpError, match="format"):
        comtier.com_convert_format(p, str(tmp_path / "o.numbers"))
    with pytest.raises(XlMcpError, match="differ"):
        comtier.com_convert_format(p, p, format="xlsx", overwrite=True)


def test_password_removal_needs_current(tmp_path):
    p = _book(tmp_path / "b.xlsx")
    with pytest.raises(XlMcpError, match="current_password"):
        comtier.com_save_with_password(p, "")


def test_sparkline_bad_args(tmp_path):
    p = _book(tmp_path / "b.xlsx")
    with pytest.raises(XlMcpError, match="action"):
        comtier.com_set_sparkline(p, action="paint")
    with pytest.raises(XlMcpError, match="location"):
        comtier.com_set_sparkline(p, action="create")
    with pytest.raises(XlMcpError, match="type"):
        comtier.com_set_sparkline(p, action="create", location="G1",
                                  source="A1:F1", type="pie")
    with pytest.raises(XlMcpError, match="source"):
        comtier.com_set_sparkline(p, action="create", location="G1")


def test_goal_seek_bad_value(tmp_path):
    p = _book(tmp_path / "b.xlsx")
    with pytest.raises(XlMcpError, match="number"):
        comtier.com_goal_seek(p, "A3", "not-a-number", "A1")


def test_missing_workbook_refuses(tmp_path):
    from xlsx_mcp.core.errors import WorkbookNotFound
    with pytest.raises(WorkbookNotFound):
        comtier.recalculate(str(tmp_path / "ghost.xlsx"))


# --------------------------------------------------------- lockfile guard


def test_guard_stale_lockfile_degrades_to_warning(tmp_path):
    p = Path(_book(tmp_path / "b.xlsx"))
    (p.parent / ("~$" + p.name)).write_bytes(b"stale owner record")
    warnings = com_session.guard_target_closed(str(p))
    assert warnings and "stale" in warnings[0]


def test_guard_truly_locked_refuses(tmp_path, monkeypatch):
    p = Path(_book(tmp_path / "b.xlsx"))

    real_open = open

    def deny(*args, **kwargs):
        if args and str(args[0]) == str(p) and "r+" in str(args[1:2]):
            raise PermissionError("locked by Excel")
        return real_open(*args, **kwargs)

    monkeypatch.setattr("builtins.open", deny)
    with pytest.raises(WorkbookLocked, match="open in Excel"):
        com_session.guard_target_closed(str(p))


# ------------------------------------------------------------- executor


@pytest.mark.skipif(not com_session.com_available()[0],
                    reason="COM plumbing absent (executor needs pywin32)")
class TestExecutor:
    def test_serialization_and_result(self, tmp_path):
        ex = com_session.ComExecutor(journal_path=tmp_path / "test-excel-pids.json")
        order: list[int] = []

        def op(i):
            def fn(manager):
                order.append(i)
                time.sleep(0.05)
                return i
            return fn

        threads = []
        results: dict[int, int] = {}
        for i in range(4):
            def call(i=i):
                results[i] = ex.submit(f"op{i}", op(i), timeout=10)
            t = threading.Thread(target=call)
            threads.append(t)
            t.start()
        for t in threads:
            t.join()
        assert sorted(results.values()) == [0, 1, 2, 3]
        # serialized: every op ran on ONE worker thread, one at a time.
        assert len(order) == 4
        ex.shutdown()

    def test_timeout_poisons_and_rearms(self, tmp_path):
        ex = com_session.ComExecutor(journal_path=tmp_path / "test-excel-pids.json")
        with pytest.raises(ExcelBlocked, match="timeout"):
            ex.submit("slowpoke", lambda m: time.sleep(5), timeout=0.2)
        st = ex.status()
        assert st["timeouts"] == 1
        # re-armed: the next submit runs on a fresh generation.
        assert ex.submit("quick", lambda m: 42, timeout=10) == 42
        ex.shutdown()

    def test_error_propagates(self, tmp_path):
        ex = com_session.ComExecutor(journal_path=tmp_path / "test-excel-pids.json")

        def boom(manager):
            raise XlMcpError("inner failure")

        with pytest.raises(XlMcpError, match="inner failure"):
            ex.submit("boom", boom, timeout=10)
        ex.shutdown()


def test_com_status_shape_never_spawns_excel():
    from xlsx_mcp.com.instances import list_excel_pids

    before = list_excel_pids()
    st = comtier.com_status()
    after = list_excel_pids()
    assert after == before, "com_status must never spawn Excel"
    for key in ("ok", "com_available", "busy", "current_op", "last_op",
                "ops_completed", "timeouts", "contention_waits",
                "journaled_pids", "policy"):
        assert key in st
    assert st["ok"] is True


def test_retry_backoff_on_busy_hresult():
    calls = {"n": 0}

    class FakeComError(Exception):
        hresult = -2147418111  # RPC_E_CALL_REJECTED

    def flaky():
        calls["n"] += 1
        if calls["n"] < 3:
            raise FakeComError()
        return "done"

    assert com_session.com_retry(flaky, base_delay=0.01) == "done"
    assert calls["n"] == 3


# -------------------------------------------------------- verify_com wiring


def test_verify_com_unavailable_degrades_to_warning(tmp_path, monkeypatch):
    p = _book(tmp_path / "b.xlsx")
    monkeypatch.setattr(com_session, "com_available",
                        lambda refresh=False: (False, "no excel here"))
    pkg = WorkbookPackage.open(p)
    pkg.set_cell("Data", "B1", 7)
    result = pkg.save(verify_com=True)
    assert result["ok"] is True
    assert "verified_com" not in result
    assert any("verify_com" in w for w in result["warnings"])


def test_verify_com_failure_restores_backup(tmp_path, monkeypatch):
    p = _book(tmp_path / "b.xlsx")
    monkeypatch.setattr(com_session, "com_available",
                        lambda refresh=False: (True, "test"))
    monkeypatch.setattr(
        com_session, "opens_clean",
        lambda path, timeout=None: {"opens_clean": False,
                                    "excel_says": "repair needed"})
    pkg = WorkbookPackage.open(p)
    pkg.set_cell("Data", "B1", 7)
    with pytest.raises(ValidationFailed, match="repair needed"):
        pkg.save(verify_com=True)
    # restored: the promoted-then-refused write was rolled back.
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["B1"].value is None
    wb.close()


def test_verify_com_pass_reports_true(tmp_path, monkeypatch):
    p = _book(tmp_path / "b.xlsx")
    monkeypatch.setattr(com_session, "com_available",
                        lambda refresh=False: (True, "test"))
    monkeypatch.setattr(
        com_session, "opens_clean",
        lambda path, timeout=None: {"opens_clean": True, "worksheets": 1})
    pkg = WorkbookPackage.open(p)
    pkg.set_cell("Data", "B1", 7)
    result = pkg.save(verify_com=True)
    assert result["verified_com"] is True


# ------------------------------------------- iterative calc (file tier)


def test_iterative_calc_settings_roundtrip(tmp_path):
    p = _book(tmp_path / "b.xlsx")
    result = properties.set_workbook_properties(
        p, iterative_calc=True, max_iterations=50, max_change=0.01)
    assert result["ok"] is True
    assert result["changed"]["properties"]["iterative_calc"] is True
    assert any("iterative" in w for w in result["warnings"])
    report = properties.set_workbook_properties(p)
    assert report["calc"]["iterative_calc"] is True
    assert report["calc"]["max_iterations"] == 50
    assert report["calc"]["max_change"] == 0.01


def test_iterative_calc_bounds(tmp_path):
    p = _book(tmp_path / "b.xlsx")
    with pytest.raises(XlMcpError, match="max_iterations"):
        properties.set_workbook_properties(p, max_iterations=0)
    with pytest.raises(XlMcpError, match="max_change"):
        properties.set_workbook_properties(p, max_change=-1)


# --------------------------------------------------- formulas fallback


def test_recalculate_formulas_engine_reads_without_writing(tmp_path):
    pytest.importorskip("formulas")
    p = _book(tmp_path / "b.xlsx")
    before = Path(p).read_bytes()
    result = comtier.recalculate(p, engine="formulas")
    assert result["ok"] is True
    assert result["saved"] is False
    assert result["engine"] == "formulas"
    assert Path(p).read_bytes() == before, \
        "the formulas engine must never modify the file"


# ------------------------------------------------ blank render detection


def _png(width: int, height: int, palette_entries: int | None,
         *, colour_type: int = 3) -> bytes:
    """A minimal, valid PNG with a palette of a chosen size.

    Built from struct and zlib rather than Pillow because the detector runs
    in the shipped server, where Pillow is not a dependency, and a test that
    needs one would be testing a different environment.
    """
    import struct
    import zlib

    def chunk(kind: bytes, data: bytes) -> bytes:
        return (struct.pack(">I", len(data)) + kind + data
                + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF))

    ihdr = struct.pack(">IIBBBBB", width, height, 8, colour_type, 0, 0, 0)
    out = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
    if palette_entries is not None:
        out += chunk(b"PLTE", bytes(palette_entries * 3))
    per_row = width * (3 if colour_type == 2 else 1)
    raw = b"".join(b"\x00" + bytes(per_row) for _ in range(height))
    return out + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def _write_png(path: Path, *args, **kwargs) -> str:
    path.write_bytes(_png(*args, **kwargs))
    return str(path)


def test_png_palette_size_reads_the_declared_entry_count(tmp_path):
    """The PLTE chunk length is an exact upper bound on how many colours the
    picture can hold, which is the whole cheap check."""
    assert comtier._png_palette_size(
        _write_png(tmp_path / "six.png", 274, 116, 6)) == 6
    assert comtier._png_palette_size(
        _write_png(tmp_path / "eighty.png", 138, 40, 80)) == 80


def test_png_palette_size_has_no_bound_for_truecolour_or_rubbish(tmp_path):
    """None means "no cheap bound", never "blank": a truecolour export is
    rich by construction, and an unreadable file is not evidence of
    anything."""
    assert comtier._png_palette_size(
        _write_png(tmp_path / "rgb.png", 40, 40, None, colour_type=2)) is None
    junk = tmp_path / "junk.png"
    junk.write_bytes(b"not a png at all")
    assert comtier._png_palette_size(str(junk)) is None
    assert comtier._png_palette_size(str(tmp_path / "absent.png")) is None


def test_render_is_blank_splits_the_artefact_from_a_thin_real_picture(
        tmp_path):
    """The measured numbers: Excel's empty chart canvas comes back with six
    palette entries at every canvas size, a range of EMPTY cells (gridlines
    only) with thirty, and a single cell holding a number with eighty."""
    assert comtier._render_is_blank(
        _write_png(tmp_path / "artefact.png", 2177, 348, 6)) is True
    assert comtier._render_is_blank(
        _write_png(tmp_path / "edge.png", 100, 100, 8)) is True
    assert comtier._render_is_blank(
        _write_png(tmp_path / "just_over.png", 100, 100, 9)) is False
    assert comtier._render_is_blank(
        _write_png(tmp_path / "gridlines.png", 546, 306, 30)) is False
    assert comtier._render_is_blank(
        _write_png(tmp_path / "one_cell.png", 138, 40, 80)) is False
    assert comtier._render_is_blank(
        _write_png(tmp_path / "rgb.png", 40, 40, None,
                   colour_type=2)) is False


def test_blank_render_refusal_names_the_route_that_works():
    msg = comtier._BLANK_RENDER_REFUSAL
    assert "blank" in msg
    assert "com_export_pdf" in msg, (
        "a refusal that names no working route sends the caller back to "
        "the same broken one")


def test_render_refuses_and_removes_the_file_when_the_png_is_blank(
        tmp_path, monkeypatch):
    """The honesty gate. com_render_sheet used to hand back ok:true with a
    populated rendered path over a white rectangle; it now refuses, and it
    does not leave the artefact on disk for a caller to find later."""
    p = _book(tmp_path / "b.xlsx")
    out = tmp_path / "render.png"

    def fake_run(label, path, body, timeout=None):
        out.write_bytes(_png(2177, 348, 6))
        return {"rendered": str(out), "range": "$A$1:$A$3", "sheet": "Data"}

    monkeypatch.setattr(comtier, "_clipboard_available", lambda: True)
    monkeypatch.setattr(comtier, "_run_readonly", fake_run)
    with pytest.raises(ValidationFailed) as exc:
        comtier.com_render_sheet(p, str(out))
    assert "com_export_pdf" in str(exc.value)
    assert not out.exists(), "the blank artefact was left on disk"


def test_render_accepts_a_picture_with_real_content(tmp_path, monkeypatch):
    """The gate is a gate, not a wall: a real export still comes back
    ok."""
    p = _book(tmp_path / "b.xlsx")
    out = tmp_path / "render.png"

    def fake_run(label, path, body, timeout=None):
        out.write_bytes(_png(546, 154, 80))
        return {"rendered": str(out), "range": "$A$1:$A$3", "sheet": "Data"}

    monkeypatch.setattr(comtier, "_clipboard_available", lambda: True)
    monkeypatch.setattr(comtier, "_run_readonly", fake_run)
    result = comtier.com_render_sheet(p, str(out))
    assert result["ok"] is True
    assert out.exists()
