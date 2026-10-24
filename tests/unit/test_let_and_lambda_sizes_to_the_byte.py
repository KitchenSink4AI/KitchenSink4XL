"""LET and LAMBDA formulas priced to the byte, as Excel stores them.

A formula whose compiled size passes Excel's limit makes Excel refuse to open
the whole workbook, and Excel's repair then drops the formula. The size count
had LET and LAMBDA below Excel: LET(f,LAMBDA(a,a+1),f(A1)), LAMBDA(a,a+1)(A1)
and LAMBDA(a,b,a+b)(A1,A1) at k = 25 copies were accepted 41, 50 and 6
references past what Excel opens, and two LET shapes 1 past; set_formula
answered ok and verified, and Excel then refused the workbook.

MEASURED in a hidden Excel 16 on a private desktop: k copies of a shape plus a
filler of n references to A1, the last j of them written A1% (a percent sign
is one byte), so the filler costs 8n + j bytes and Excel's capacity is
bisected to ONE byte. The plain filler opens at 16,383 bytes and not 16,384.
X below is the smallest number of bytes the k copies may take for the file to
open (its formula comes back intact): the k copies plus the filler at
16,383 - (k - 1) - X bytes open, one byte more does not.

Every shape is refused one byte past Excel's limit (never written when Excel
would refuse the workbook). Every shape priced exactly is also accepted at
the limit. Three shapes use a captured name more than once; Excel took a
byte less for each further use than for a plain use, and the count takes a
plain use, so those are refused a few bytes early, by design. No Excel here
except the tests marked LIVE, which start the product's hidden Excel (run
them on a private desktop).
"""
from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.com import session as com_session
from xlsx_mcp.core import calc, limits
from xlsx_mcp.core.errors import XlMcpError

LIVE = pytest.mark.skipif(not com_session.com_available()[0], reason="needs Excel (COM)")

#: (shape as typed, k, X): Excel opens the k copies plus the filler when the
#: copies take X bytes and the filler 16,383 - (k - 1) - X; one byte more,
#: and it refuses the workbook.
EXACT = [
    ("ABS(A1)", 25, 250),
    ("LET(x,A1,x)", 25, 454),
    ("LET(x,A1,x)", 12, 220),
    ("LAMBDA(a,1)(1)", 25, 800),
    ("LAMBDA(a,1)(1)", 12, 384),
    ("LAMBDA(a,b,1)(1,1)", 25, 875),
    ("LAMBDA(a,b,1)(1,1)", 12, 420),
    ("LAMBDA(a,a)(1)", 25, 875),
    ("LAMBDA(a,a+a)(1)", 25, 1050),
    ("LAMBDA(a,b,c,1)(1,1,1)", 25, 950),
    ("LET(f,LAMBDA(a,1),1)", 25, 829),
    ("LET(f,LAMBDA(a,1),1)", 12, 400),
    ("LET(f,LAMBDA(a,b,1),1)", 25, 829),
    ("LET(f,LAMBDA(a,a),f(1))", 25, 1154),
    ("LET(f,LAMBDA(a,a),f(1)+f(1))", 25, 1504),
    ("LET(f,LAMBDA(a,a),MAP(A1:B2,f))", 25, 1404),
    ("LAMBDA(a,LAMBDA(b,a+b)(1))(1)", 25, 1900),
    ("LAMBDA(a,ISOMITTED(a))(1)", 25, 950),
    ("MAP(A1:B2,LAMBDA(a,a))", 25, 1125),
    ("REDUCE(0,A1:B2,LAMBDA(s,v,s+v))", 25, 1350),
    ("SCAN(0,A1:B2,LAMBDA(s,v,s+v))", 25, 1350),
    ("BYROW(A1:B2,LAMBDA(r,SUM(r)))", 25, 1225),
    ("BYCOL(A1:B2,LAMBDA(r,SUM(r)))", 25, 1225),
    ("MAKEARRAY(2,2,LAMBDA(r,c,r+c))", 25, 1100),
    ("LET(x,A1:B2,SUM(x))", 25, 704),
    ("LET(x,A1,y,x+1,y)", 25, 704),
    ("LET(f,LAMBDA(a,a+1),f(A1))", 25, 1354),
    ("LET(f,LAMBDA(a,a+1),f(A1))", 12, 652),
    ("LAMBDA(a,a+1)(A1)", 25, 1075),
    ("LAMBDA(a,b,a+b)(A1,A1)", 25, 1325),
    ("LAMBDA(a,a)(1)+NOW()", 25, 979),
    ("LAMBDA(a,a)(1)+NOW()", 12, 472),
    ("LAMBDA(a,LAMBDA(b,b)(1))(1)", 25, 1600),
    ("LAMBDA(a,LAMBDA(b,a)(1))(1)", 25, 1725),
    ("LET(x,1,f,LAMBDA(a,a+x),f(1))", 25, 1529),
    ("LET(x,1,f,LAMBDA(a,a+x),f(1))", 12, 736),
    ("LET(x,1,MAP(A1:B2,LAMBDA(a,a+x)))", 25, 1629),
    ("LAMBDA(a,LET(x,a,x))(1)", 25, 1250),
    ("LAMBDA(a,LET(x,a,x))(1)", 12, 600),
    ("LAMBDA(a,LET(x,1,x+a))(1)", 25, 1350),
    ("LET(x,A1,x)+NOW()", 25, 558),
    ("LET(x,A1,x)+NOW()", 12, 272),
    ("LAMBDA(g,g(1))(1)", 25, 1050),
    ("BYROW(A1:B2,SUM)", 25, 550),
    ("LAMBDA(a,b,a)(1,)", 25, 900),
    ("LAMBDA( a , a )( 1 )", 25, 1275),
    ("LET(x,1,x)+LET(y,2,y)", 25, 729),
    ("LAMBDA(a,a+NOW())(1)", 25, 979),
    ("LAMBDA(a,a+NOW())(1)", 12, 472),
    ("LAMBDA(a,a+NOW())(1)+NOW()", 25, 1079),
    ("LAMBDA(a,a+NOW())(1)+NOW()", 12, 520),
    ("LAMBDA(a,LAMBDA(b,a+b))(1)(2)", 25, 1900),
    ("LAMBDA(a,b,LAMBDA(c,a+b)(1))(1,1)", 25, 2100),
    ("LAMBDA(a,LAMBDA(b,LAMBDA(c,a)(1))(1))(1)", 25, 2575),
    ("LAMBDA(a,LAMBDA(b,a+b+b)(1))(1)", 25, 2075),
    ("LET(y,1,LAMBDA(a,LET(x,a,x))(y))", 25, 1529),
    ("LET(y,1,LAMBDA(a,LET(x,a,x))(y))", 12, 736),
    ("LAMBDA(a,IF(a,1,2))(1)", 25, 1425),
    ("IF(A1,1,2)", 25, 725),
]

#: A captured name used two or three times: counted high, so refused a few
#: bytes before Excel's limit (never after it).
COUNTED_HIGH = [
    ("LAMBDA(a,LAMBDA(b,a+a)(1))(1)", 25, 1875),
    ("LAMBDA(a,LAMBDA(b,a+a+a)(1))(1)", 25, 2025),
    ("LET(x,1,f,LAMBDA(a,x+x),f(1))", 25, 1504),
]

LIMIT = limits.MAX_FORMULA_COMPILED_BYTES


def _filler(total):
    n, j = divmod(total, 8)
    return ["A1%"] * j + ["A1"] * (n - j)


def _formula(shape, k, filler_bytes):
    return "=" + "+".join([shape] * k + _filler(filler_bytes))


def _stored(typed):
    s = calc.normalize_formula(typed)[0]
    return s if s.startswith("=") else "=" + s


def _accepted(typed):
    try:
        limits.check_formula(typed)
        limits.check_formula(_stored(typed))
        return True
    except limits.ExcelWouldRefuse:
        return False


def _edge(k, x):
    return LIMIT - (k - 1) - x


def test_the_filler_itself_is_priced_to_the_byte():
    assert limits.compiled_size_floor(_formula("A1", 0, 16384)[1:].lstrip("+")) == LIMIT


@pytest.mark.parametrize("shape,k,x", EXACT, ids=[f"{s} k{k}" for s, k, _ in EXACT])
def test_excels_limit_is_kept_to_the_byte(shape, k, x):
    t = _edge(k, x)
    assert _accepted(_formula(shape, k, t)), "refused a formula Excel opens"
    assert not _accepted(_formula(shape, k, t + 1)), "accepted a formula Excel will not open"


@pytest.mark.parametrize("shape,k,x", COUNTED_HIGH, ids=[s for s, _, _ in COUNTED_HIGH])
def test_a_repeated_captured_name_is_counted_high_never_low(shape, k, x):
    t = _edge(k, x)
    assert not _accepted(_formula(shape, k, t + 1)), "accepted a formula Excel will not open"
    over = limits.compiled_size_floor(_stored(_formula(shape, k, t))) - LIMIT
    assert 0 < over <= 2 * k, over


@pytest.mark.parametrize("shape,k,x", EXACT + COUNTED_HIGH, ids=[f"{s} k{k}" for s, k, _ in EXACT + COUNTED_HIGH])
def test_typed_and_stored_text_are_priced_alike(shape, k, x):
    typed = _formula(shape, k, _edge(k, x))
    assert limits.compiled_size_floor(typed) == limits.compiled_size_floor(_stored(typed))


# ------------------------------------------------------------ the figures


def _size(f):
    return limits.compiled_size_floor(f)


def test_a_lambda_costs_twenty_two_whatever_it_declares():
    one = _size("=LET(f,LAMBDA(a,1),1)")
    assert one == 5 + 0 + 22 + 0 + 3 + 3 + 4
    assert _size("=LET(f,LAMBDA(a,b,c,1),1)") == one


def test_a_lambda_called_where_it_is_written_adds_four_and_its_arguments():
    assert _size("=LAMBDA(a,1)(1)") == 22 + 3 + 4 + 3
    assert _size("=LAMBDA(a,b,1)(1,2)") == 22 + 3 + 4 + 3 + 3
    assert _size("=LAMBDA(a,b,a+b)(A1,A1)") == 22 + 13 + 4 + 7 + 7
    # a call on what a call returns is a call too
    assert _size("=LAMBDA(a,LAMBDA(b,b))(1)(2)") == 22 + (22 + 6) + 4 + 3 + 4 + 3


def test_a_missing_argument_costs_one():
    assert _size("=LAMBDA(a,b,a)(1,)") == _size("=LAMBDA(a,b,a)(1)") + 1
    assert _size("=IF(A1,,1)") == _size("=IF(A1,0,1)") - 3 + 1
    assert _size("=VLOOKUP(A1,B1:C9,2,)") == _size("=VLOOKUP(A1,B1:C9,2)") + 1
    assert _size("=NOW()") == 3 + limits.ONCE_PER_FORMULA_BYTES  # no argument is not a missing one


def test_lets_flag_is_paid_once_per_body():
    assert _size("=LET(x,1,x)+LET(y,2,y)") == 2 * (5 + 3 + 6) + 1 + 4
    # a LET only inside a LAMBDA: that LAMBDA's flag, none for the formula
    assert _size("=LAMBDA(a,LET(x,a,x))(1)") == 22 + 5 + 6 + 6 + 4 + 4 + 3
    # and both when the formula has one outside too
    assert _size("=LET(y,1,LAMBDA(a,LET(x,a,x))(y))") == 5 + 3 + 22 + 5 + 6 + 6 + 4 + 4 + 6 + 4


def test_lets_flag_and_the_volatile_flag_are_two_flags():
    assert _size("=LET(x,A1,x)+NOW()") == (5 + 7 + 6) + 1 + 3 + 4 + 4
    # the volatile flag is paid once per formula, inside a LAMBDA too
    assert _size("=LAMBDA(a,a+NOW())(1)+NOW()") == 22 + 6 + 1 + 3 + 4 + 3 + 1 + 3 + 4


def test_a_name_a_lambda_captures_costs_five_more_per_lambda():
    plain = _size("=LAMBDA(a,LAMBDA(b,b)(1))(1)")
    assert _size("=LAMBDA(a,LAMBDA(b,a)(1))(1)") == plain + 5
    assert _size("=LAMBDA(a,b,LAMBDA(c,a+b)(1))(1,1)") == 22 + 22 + 6 + 1 + 6 + 4 + 3 + 4 + 3 + 3 + 2 * 5
    # passed on through a LAMBDA that does not use it: paid by both
    assert _size("=LAMBDA(a,LAMBDA(b,LAMBDA(c,a)(1))(1))(1)") == 3 * 22 + 6 + 3 * (4 + 3) + 2 * 5
    # a LET name used inside a LAMBDA is captured too; a parameter of the
    # LAMBDA's own is not, also inside a LET in its body
    assert _size("=LET(x,1,MAP(A1:B2,LAMBDA(a,a+x)))") == 5 + 3 + 4 + 13 + 22 + 6 + 1 + 6 + 5 + 4
    assert _size("=LAMBDA(a,LET(x,1,x+a))(1)") == 22 + 5 + 3 + 6 + 1 + 6 + 4 + 4 + 3


def test_spaces_before_a_declared_name_cost_nothing():
    assert _size("=LAMBDA( a , a )( 1 )") == _size("=LAMBDA(a,a)(1)") + 4 * 4


def test_lambda_helpers_are_priced_as_measured():
    assert _size("=MAP(A1:B2,LAMBDA(a,a))") == 4 + 13 + 22 + 6
    assert _size("=BYROW(A1:B2,LAMBDA(r,SUM(r)))") == 4 + 13 + 22 + 4 + 6
    assert _size("=REDUCE(0,A1:B2,LAMBDA(s,v,s+v))") == 3 + 3 + 13 + 22 + 13
    assert _size("=MAKEARRAY(2,2,LAMBDA(r,c,r+c))") == 3 + 3 + 3 + 22 + 13


# ---------------------------------------------------------- through the product

REPORTED_SHAPES = [
    ("LET(x,A1:B2,SUM(x))", 704),
    ("LET(x,A1,y,x+1,y)", 704),
    ("LET(f,LAMBDA(a,a+1),f(A1))", 1354),
    ("LAMBDA(a,a+1)(A1)", 1075),
    ("LAMBDA(a,b,a+b)(A1,A1)", 1325),
]


def _realistic(tmp_path, name="Book.xlsx"):
    p = tmp_path / name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, 201):
        ws.cell(r, 1, r)
        ws.cell(r, 2, f"row {r}")
    ws["E1"] = "=SUM(A1:A200)"
    wb.create_sheet("Notes")["A1"] = "notes"
    wb.save(p)
    return p


@pytest.mark.parametrize("shape,x", REPORTED_SHAPES, ids=[s for s, _ in REPORTED_SHAPES])
def test_set_formula_refuses_one_byte_past_excel_and_leaves_the_file(tmp_path, shape, x):
    from xlsx_mcp.ops import formulas
    p = _realistic(tmp_path)
    before = p.read_bytes()
    with pytest.raises(XlMcpError, match="too large for Excel to store"):
        formulas.set_formula(str(p), "D1", _formula(shape, 25, _edge(25, x) + 1), sheet="Data")
    assert p.read_bytes() == before


@LIVE
@pytest.mark.parametrize("shape,x", REPORTED_SHAPES, ids=[s for s, _ in REPORTED_SHAPES])
def test_the_largest_formula_written_opens_clean_in_excel(tmp_path, shape, x):
    """The product writes the largest such formula it accepts, and real Excel
    opens the workbook with the formula intact."""
    from xlsx_mcp.ops import formulas
    p = _realistic(tmp_path)
    typed = _formula(shape, 25, _edge(25, x))
    out = formulas.set_formula(str(p), "D1", typed, sheet="Data")
    assert out.get("ok", True)
    try:
        verdict = com_session.opens_clean(str(p), timeout=120)
        assert verdict["opens_clean"] is True, verdict

        def read(manager):
            w = manager.acquire()
            wb = com_session.open_workbook(w.app, str(p), read_only=True)
            try:
                return str(wb.Worksheets("Data").Range("D1").Formula), \
                    float(wb.Worksheets("Data").Range("E1").Value)
            finally:
                wb.Close(False)
                manager.release_to_pool(w)
        formula, total = com_session.run_com("read-back", read, timeout=120)
        assert formula.startswith("=") and formula.count("A1") >= 1000, formula[:80]
        assert total == sum(range(1, 201))
    finally:
        com_session.get_executor().shutdown()
