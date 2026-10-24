"""A formula's cached result keeps the cell type it was stored with.

An always-calculate formula (ca="1": anything built on RAND, TODAY, NOW,
INDIRECT, OFFSET or CELL) keeps its cached result through a save that
rewrites the whole workbook, so an iterative model can seed from it. The
result went back WITHOUT its cell type: a text result came back as
<v>label</v> in a cell with no t="str", which is a number cell holding text.
Excel cannot open such a file at all ("Open method of Workbooks class
failed" in normal, repair and extract-data modes, measured), and the save
still answered verified. An error result lost t="e" the same way, and a
TRUE result lost t="b" and came back as the number 1.

Two fixes, both pinned here. The cached result goes back with its type. And
the verify gate now reads every stored value against its cell's type, so a
produced file in which they contradict each other is refused instead of
being answered verified.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import openpyxl
import pytest

from xlsx_mcp.core import verify
from xlsx_mcp.core.errors import ValidationFailed
from xlsx_mcp.core.package import WorkbookPackage
from xlsx_mcp.ops import cells, condformat, names, structure

#: Made-up always-calculate formulas whose cached results are text, an
#: error, a boolean and a number, as Excel stores them.
_VOLATILE = (
    '<c r="B1" t="str"><f ca="1">IF(RAND()&gt;2,"x","label")</f>'
    '<v>label</v></c>'
    '<c r="C1" t="e"><f ca="1">IF(RAND()&gt;2,1,NA())</f><v>#N/A</v></c>'
    '<c r="D1" t="b"><f ca="1">RAND()&lt;2</f><v>1</v></c>'
    '<c r="E1"><f ca="1">RAND()*0+3</f><v>3</v></c>')

_SHEET = "xl/worksheets/sheet1.xml"


def _patch_sheet(path: Path, fn) -> None:
    with zipfile.ZipFile(path) as zin:
        parts = {i.filename: zin.read(i.filename) for i in zin.infolist()}
    parts[_SHEET] = fn(parts[_SHEET].decode("utf-8")).encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)


def _book(tmp_path: Path, extra_cells: str = _VOLATILE) -> str:
    p = tmp_path / "volatile.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 5
    ws["A2"] = 7
    wb.save(p)

    def add(x: str) -> str:
        m = re.search(r'<c r="A1"[^>]*>.*?</c>', x, re.S)
        assert m, x
        return x[:m.end()] + extra_cells + x[m.end():]

    _patch_sheet(p, add)
    return str(p)


def _cell(path: str, ref: str) -> tuple[dict, str | None]:
    """(attributes, stored <v> text) of one cell of the first sheet."""
    with zipfile.ZipFile(path) as zf:
        x = zf.read(_SHEET).decode("utf-8")
    m = re.search(r'<c r="%s"((?:[^>/"]|"[^"]*")*)(?:/>|>(.*?)</c>)' % ref,
                  x, re.S)
    assert m, f"{ref} is not stored"
    attrs = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
    v = re.search(r"<v>(.*?)</v>", m.group(2) or "", re.S)
    return attrs, (v.group(1) if v else None)


_EDITS = {
    "name": lambda p: names.manage_name(
        p, "add", name="Two", refers_to="=Data!$A$1:$A$2"),
    "cf": lambda p: condformat.manage_conditional_format(
        p, "add", location="A1:A2", sheet="Data", cf_type="cell_is",
        params={"operator": "greater_than", "formula": ["1"],
                "fill": "FFC7CE"}),
    "insert_cols": lambda p: structure.modify_grid_structure(
        p, "insert_cols", "G", sheet="Data"),
}


@pytest.mark.parametrize("edit", sorted(_EDITS))
def test_a_full_rewrite_puts_each_cached_result_back_with_its_type(
        tmp_path, edit):
    p = _book(tmp_path)
    out = _EDITS[edit](p)
    assert out["verified"] is True
    assert out["route"] == "openpyxl"  # the whole workbook was rewritten
    expect = {"B1": ("str", "label"), "C1": ("e", "#N/A"),
              "D1": ("b", "1"), "E1": (None, "3")}
    for ref, (t, v) in expect.items():
        attrs, got = _cell(p, ref)
        assert attrs.get("t") == t, (ref, attrs)
        assert got == v, (ref, got)


def test_the_cached_results_read_back_as_text_error_boolean_and_number(
        tmp_path):
    p = _book(tmp_path)
    names.manage_name(p, "add", name="Two", refers_to="=Data!$A$1:$A$2")
    got = cells.get_cells(p, ["B1", "C1", "D1", "E1"], sheet="Data")
    values = [c[1] for c in got["cells"]]
    assert values == ["label", "#N/A", True, 3]


def test_a_cell_the_edit_wrote_never_takes_the_old_cache_or_type(tmp_path):
    """The restore is keyed on the formula being unchanged: a rewritten
    always-calculate formula gets neither its old result nor its old type."""
    p = _book(tmp_path)
    from xlsx_mcp.ops import formulas
    formulas.set_formula(p, {"cell": "B1"}, '=IF(RAND()>2,"x",7)',
                         sheet="Data")
    names.manage_name(p, "add", name="Two", refers_to="=Data!$A$1:$A$2")
    attrs, v = _cell(p, "B1")
    assert v != "label"
    assert attrs.get("t") != "str" or v is None


# ------------------------------------------------- the verify gate itself

def _with_cell(tmp_path: Path, cell: str, shared: list[str] | None = None
               ) -> str:
    """A made-up workbook whose sheet holds `cell` at B1 (raw XML)."""
    p = tmp_path / "typed.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "one shared string"  # openpyxl writes it inline
    wb.save(p)
    _patch_sheet(p, lambda x: re.sub(
        r"(<c r=\"A1\".*?</c>)", lambda m: m.group(1) + cell, x, count=1,
        flags=re.S))
    with zipfile.ZipFile(p) as zin:
        parts = {i.filename: zin.read(i.filename) for i in zin.infolist()}
    # one shared-string part with exactly `shared` (or none at all)
    sst_name = next((n for n in parts if n.lower() == "xl/sharedstrings.xml"),
                    None)
    if sst_name is not None:
        del parts[sst_name]
    ct = re.sub(rb'<Override PartName="/xl/sharedStrings.xml"[^>]*/>', b"",
                parts["[Content_Types].xml"])
    rels = re.sub(rb'<Relationship [^>]*sharedStrings[^>]*/>', b"",
                  parts["xl/_rels/workbook.xml.rels"])
    if shared is not None:
        sst = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/'
               '2006/main" count="%d" uniqueCount="%d">' % (len(shared),
                                                            len(shared))
               + "".join("<si><t>%s</t></si>" % s for s in shared) + "</sst>")
        parts["xl/sharedStrings.xml"] = sst.encode("utf-8")
        ct = ct.replace(
            b"</Types>",
            b'<Override PartName="/xl/sharedStrings.xml" ContentType="'
            b'application/vnd.openxmlformats-officedocument.spreadsheetml.'
            b'sharedStrings+xml"/></Types>')
        rels = rels.replace(
            b"</Relationships>",
            b'<Relationship Id="rIdSst" Type="http://schemas.openxmlformats.'
            b'org/officeDocument/2006/relationships/sharedStrings" '
            b'Target="sharedStrings.xml"/></Relationships>')
    parts["[Content_Types].xml"] = ct
    parts["xl/_rels/workbook.xml.rels"] = rels
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)
    return str(p)


#: Each form below was opened by Excel itself (the oracle for this gate):
#: every "refused" form fails "Open method of Workbooks class" in normal,
#: repair and extract-data modes; every "passes" form opens clean.
@pytest.mark.parametrize("cell", [
    # a formula's text result in a number cell
    '<c r="B1"><f ca="1">IF(RAND()&gt;2,"x","label")</f><v>label</v></c>',
    '<c r="B1"><f>IF(A2&gt;2,"x","label")</f><v>label</v></c>',
    '<c r="B1" t="n"><v>label</v></c>',
    # an error result that lost t="e"
    '<c r="B1"><f ca="1">NA()</f><v>#N/A</v></c>',
    # number spellings Excel does not read
    '<c r="B1"><v>+5</v></c>',
    '<c r="B1"><v>INF</v></c>',
    '<c r="B1"><v>1,5</v></c>',
    '<c r="B1"><v>- 5</v></c>',
    '<c r="B1"><v>1e</v></c>',
    '<c r="B1"><v>1e+</v></c>',
    '<c r="B1"><v>1.5.2</v></c>',
    # other types holding what they cannot hold
    '<c r="B1" t="e"><v>#BOGUS!</v></c>',
    '<c r="B1" t="b"><v>yes</v></c>',
    '<c r="B1" t="d"><v>label</v></c>',
    '<c r="B1" t="x"><v>label</v></c>',
])
def test_the_verify_gate_refuses_a_value_its_cell_type_cannot_hold(
        tmp_path, cell):
    p = _with_cell(tmp_path, cell)
    res = verify.verify_after_write(p, pre_parts=[])
    assert not res.ok
    text = " ".join(res.reasons)
    assert "Data!B1" in text
    assert "Excel" in text


@pytest.mark.parametrize("cell", [
    '<c r="B1" t="str"><f>IF(A2&gt;2,"x","label")</f><v>label</v></c>',
    '<c r="B1" t="e"><f>NA()</f><v>#N/A</v></c>',
    '<c r="B1" t="b"><f>A2&gt;1</f><v>1</v></c>',
    '<c r="B1"><f>A2*2</f><v>1.5E-3</v></c>',
    '<c r="B1"><v>-0.25</v></c>',
    '<c r="B1" t="inlineStr"><is><t>label</t></is></c>',
    '<c r="B1"><v> 5</v></c>',
    '<c r="B1"><v>.5</v></c>',
    '<c r="B1"><v>5.</v></c>',
    '<c r="B1"><v>1E+05</v></c>',
    '<c r="B1"><v></v></c>',
    '<c r="B1" t="x"><v>5</v></c>',
    '<c r="B1" t="e"><v>#n/a</v></c>',
    '<c r="B1" t="b"><v>false</v></c>',
    '<c r="B1" t="d"><v>2026-02-03T00:00:00Z</v></c>',
    # more number spellings Excel opens (Batch A4): "-" and "." read as 0
    '<c r="B1"><v>-.5</v></c>',
    '<c r="B1"><v>-5.</v></c>',
    '<c r="B1"><v>007</v></c>',
    '<c r="B1"><v>-</v></c>',
    '<c r="B1"><v>.</v></c>',
    '<c r="B1"><v>-.5e1</v></c>',
])
def test_the_verify_gate_passes_values_that_match_their_type(tmp_path, cell):
    p = _with_cell(tmp_path, cell)
    res = verify.verify_after_write(p, pre_parts=[])
    assert res.ok, res.reasons


#: Error values as a plain stored value (t="e"), opened by Excel itself: the
#: classic seven, #GETTING_DATA and #BUSY! open clean. Excel stores the errors
#: of dynamic arrays, data types and Python in Excel only as rich-value errors
#: (#VALUE! with a value-metadata index); a plain stored #SPILL!, #CALC!,
#: #FIELD!, #BLOCKED!, #CONNECT!, #UNKNOWN!, #PYTHON! or #EXTERNAL! makes it
#: refuse the whole file, as a constant or as a formula's cached result.
_ERRORS_EXCEL_OPENS = ["#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?",
                       "#NUM!", "#N/A", "#GETTING_DATA", "#BUSY!"]
_ERRORS_EXCEL_REFUSES = ["#SPILL!", "#CALC!", "#FIELD!", "#BLOCKED!",
                         "#CONNECT!", "#UNKNOWN!", "#PYTHON!", "#EXTERNAL!"]


@pytest.mark.parametrize("err", _ERRORS_EXCEL_OPENS)
def test_the_verify_gate_passes_every_error_value_excel_stores(tmp_path, err):
    p = _with_cell(tmp_path, f'<c r="B1" t="e"><v>{err}</v></c>')
    res = verify.verify_after_write(p, pre_parts=[])
    assert res.ok, res.reasons


@pytest.mark.parametrize("shape", ["constant", "cached"])
@pytest.mark.parametrize("err", _ERRORS_EXCEL_REFUSES)
def test_the_verify_gate_refuses_an_error_excel_stores_only_as_rich_value(
        tmp_path, err, shape):
    cell = (f'<c r="B1" t="e"><v>{err}</v></c>' if shape == "constant" else
            f'<c r="B1" t="e" cm="1"><f t="array" ref="B1">_xlfn.SEQUENCE(3)'
            f'</f><v>{err}</v></c>')
    p = _with_cell(tmp_path, cell)
    res = verify.verify_after_write(p, pre_parts=[])
    assert not res.ok
    text = " ".join(res.reasons)
    assert "Data!B1" in text and err in text
    assert "rich value" in text


@pytest.mark.parametrize("cell,shared", [
    ('<c r="B1" t="s"><v>1</v></c>', ["one"]),      # past the last string
    ('<c r="B1" t="s"><v>1.0</v></c>', ["one", "two"]),
    ('<c r="B1" t="s"><v>0</v></c>', None),          # no shared strings
])
def test_the_verify_gate_refuses_a_shared_string_that_is_not_there(
        tmp_path, cell, shared):
    p = _with_cell(tmp_path, cell, shared=shared)
    res = verify.verify_after_write(p, pre_parts=[])
    assert not res.ok
    assert "Data!B1" in " ".join(res.reasons)


def test_a_save_that_would_store_text_in_a_number_cell_is_refused(tmp_path):
    """End to end: whatever route produced it, a package holding such a cell
    never promotes, and the workbook on disk is left as it was."""
    p = _book(tmp_path, extra_cells="")
    before = Path(p).read_bytes()
    pkg = WorkbookPackage.open(p)
    pkg.set_cell("Data", "A2", 8)

    def saver(tmp):
        pkg._default_saver(tmp)
        _patch_sheet(Path(tmp), lambda x: re.sub(
            r'(<c r="A1"[^>]*>.*?</c>)',
            lambda m: m.group(1) + '<c r="B1"><f>A1&amp;"x"</f>'
                                   '<v>5x</v></c>', x, count=1, flags=re.S))

    with pytest.raises(ValidationFailed) as ei:
        pkg.save(saver=saver)
    assert "Data!B1" in str(ei.value)
    assert Path(p).read_bytes() == before
