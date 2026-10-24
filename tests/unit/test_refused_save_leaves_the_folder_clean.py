"""A refused save leaves the user's folder as it
was, and every scratch file a save makes beside the workbook is one the
stale-temp sweeper can clean.

-   When .ks4xl-backups beside the workbook is a FILE, every edit refused
    with the raw "[WinError 183] ..." text and left a hidden
    .ks4xl-write-<hex> payload behind per attempt. It now refuses before
    writing anything, with a plain message naming the folder.
-   The dynamic-array marker (and four calc helpers and the carry step) made
    their scratch copies with tempfile.mkstemp in the workbook's folder,
    named tmpXXXX, which the sweeper (.ks4xl-write-* only) never cleans if
    the process dies mid-write. They now use the swept prefix.
"""
from __future__ import annotations

import os
import re
from pathlib import Path

import openpyxl
import pytest

from xlsx_mcp.core.errors import XlMcpError

SRC = Path(__file__).resolve().parents[2] / "src" / "xlsx_mcp"


def _book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A2"] = 2
    ws["A3"] = "=A1+A2"
    wb.save(path)
    return str(path)


def _blocked(tmp_path):
    p = _book(tmp_path / "f.xlsx")
    (tmp_path / ".ks4xl-backups").write_text("x")
    return p, sorted(os.listdir(tmp_path)), Path(p).read_bytes()


def test_a_blocked_backup_folder_refuses_plainly_and_leaves_nothing(tmp_path):
    from xlsx_mcp import envelope
    from xlsx_mcp.ops import cells
    p, before, data = _blocked(tmp_path)
    for _ in range(2):
        with pytest.raises(XlMcpError) as ei:
            cells.set_cell(p, "D1", "x")
        msg = str(ei.value)
        assert ".ks4xl-backups" in msg and "WinError" not in msg and "Errno" not in msg
        assert "was not changed" in msg
        assert envelope.refusal(ei.value)["error"]["code"] == "CONFLICT"
    assert sorted(os.listdir(tmp_path)) == before
    assert Path(p).read_bytes() == data


def test_a_destructive_edit_with_a_snapshot_refuses_the_same_way(tmp_path):
    from xlsx_mcp.ops import structure
    p, before, data = _blocked(tmp_path)
    with pytest.raises(XlMcpError) as ei:
        structure.modify_grid_structure(p, action="delete_rows", at=1, count=1)
    assert ".ks4xl-backups" in str(ei.value) and "WinError" not in str(ei.value)
    assert sorted(os.listdir(tmp_path)) == before
    assert Path(p).read_bytes() == data


def test_without_a_backup_the_blocked_folder_does_not_matter(tmp_path):
    from xlsx_mcp.ops import cells
    p, before, _data = _blocked(tmp_path)
    out = cells.set_cell(p, "D1", "x", backup=False)
    assert out["ok"] and out.get("backup") is None
    assert sorted(os.listdir(tmp_path)) == before


def test_the_com_route_refuses_before_excel_with_the_same_words(tmp_path, monkeypatch):
    from xlsx_mcp.com import session
    from xlsx_mcp.ops import comtier
    monkeypatch.setattr(session, "run_com", lambda *a, **k: (_ for _ in ()).throw(
        AssertionError("Excel was started")))
    p, before, data = _blocked(tmp_path)
    with pytest.raises(XlMcpError) as ei:
        comtier.com_autofit(p)
    assert ".ks4xl-backups" in str(ei.value) and "WinError" not in str(ei.value)
    assert sorted(os.listdir(tmp_path)) == before



def test_every_scratch_file_beside_the_workbook_uses_the_swept_prefix():
    """No tempfile call in the save path may create an unswept name."""
    def call_text(text, start):
        depth, i = 0, text.index("(", start)
        for j in range(i, len(text)):
            depth += {"(": 1, ")": -1}.get(text[j], 0)
            if depth == 0:
                return text[start:j + 1]
        return text[start:]

    offenders = []
    for f in SRC.rglob("*.py"):
        text = f.read_text(encoding="utf-8")
        for m in re.finditer(r"tempfile\.(mkstemp|NamedTemporaryFile)\(", text):
            call = call_text(text, m.start())
            if "dir=" in call and "WRITE_TMP_PREFIX" not in call:
                line = text.count("\n", 0, m.start()) + 1
                offenders.append(f"{f.relative_to(SRC).as_posix()}:{line}")
    assert offenders == [], offenders


def test_the_sweeper_takes_an_aged_scratch_file(tmp_path):
    import time
    from xlsx_mcp.core import package, safesave
    fd, scratch = safesave.scratch_file(str(tmp_path), ".xlsx")
    os.close(fd)
    old = time.time() - 2 * 60 * 60
    os.utime(scratch, (old, old))
    assert os.path.basename(scratch).startswith(".ks4xl-write-")
    package._sweep_stale_write_tmps(str(tmp_path))
    assert not os.path.exists(scratch)
