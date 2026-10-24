"""Defects on the family surface, one test per defect.

Two kinds of test live here:
- RED before the fix: the product was wrong; the fix makes it green.
- GAP guards (said so in the docstring): green before the fix, but red under
  the mutant or the weakened product that the existing suite lets through.

COM tests use only the product's own private hidden Excel pool
(DispatchEx, Visible=False); they skip where Excel is not available.
"""

from __future__ import annotations

import asyncio
import hashlib
import json
import os
import platform
import re
import subprocess
import sys
import threading
import time
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp import envelope, registry
from xlsx_mcp import server as srv
from xlsx_mcp.com import session as com_session
from xlsx_mcp.core import safesave
from xlsx_mcp.settings import Settings
from xlsx_mcp.surface import dispatch, groups

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture
def xlsx(tmp_path):
    p = tmp_path / "v.xlsx"
    srv.create_workbook(str(p))
    srv.write_range(str(p), "A1", [["Name", "Qty"], ["Alpha", 1], ["Beta", 2]])
    return str(p)


def _sha(p) -> str | None:
    p = Path(p)
    return hashlib.sha256(p.read_bytes()).hexdigest() if p.exists() else None


def _wait_for_worker_threads(limit: float = 60.0) -> None:
    """Poll until no asyncio worker thread is alive (no timing margin, T6)."""
    end = time.monotonic() + limit
    while time.monotonic() < end:
        if not any(t.name.startswith("asyncio") and t.is_alive() for t in threading.enumerate()):
            return
        time.sleep(0.2)


def _slow(monkeypatch, target, attr, seconds=3.0):
    real = getattr(target, attr)

    def slow(*a, **k):
        time.sleep(seconds)
        return real(*a, **k)
    monkeypatch.setattr(target, attr, slow)


# --------------------------------------------------------------------------
# "a TIMEOUT answer means nothing was saved": manage_backups
# snapshot copies the workbook, and purge deletes backups, with no commit
# point, so both change the disk AFTER the TIMEOUT answer that says
# "nothing was saved" (details.saved false).

@pytest.mark.timeout(120)
@pytest.mark.parametrize("case", ["snapshot", "purge"])
def test_f1_timeout_answer_means_the_disk_was_not_changed(case, xlsx, tmp_path, monkeypatch):
    monkeypatch.setattr(dispatch, "BACKSTOP_GRACE_S", 0.2)
    srv.set_cell(xlsx, "C1", "x")  # a prev slot exists
    slot_dir = safesave.slot_dir(xlsx)
    snaps = tmp_path / "snaps"
    snaps.mkdir()
    from xlsx_mcp.ops import backups
    if case == "snapshot":
        _slow(monkeypatch, backups.safesave, "bound_name")
        args = {"action": "snapshot", "path": xlsx, "dest_dir": str(snaps)}
    else:
        _slow(monkeypatch, backups, "_collect_purge_targets")
        args = {"action": "purge", "path": xlsx, "scope": "slots", "dry_run": False}
    before = sorted(p.name for p in snaps.iterdir()), sorted(p.name for p in slot_dir.iterdir())
    h = fh.start(Settings(call_timeout_s=1.0, max_call_timeout_s=1.0))
    payload, err, _ = h.call("xl_workbook", {"op": "manage_backups", "args": args})
    assert err and payload["error"]["code"] == "TIMEOUT", payload
    assert payload["error"]["details"]["saved"] is False
    _wait_for_worker_threads()
    after = (sorted(p.name for p in snaps.iterdir()),
             sorted(p.name for p in slot_dir.iterdir()) if slot_dir.exists() else [])
    assert after == before, f"{case}: TIMEOUT said nothing was saved, but the disk changed: " \
                            f"{before} -> {after}"


# --------------------------------------------------------------------------
# The same invariant on the COM tier: every xl_app op that saves
# (the workbook itself through _run_mutation's wb.Save, or a new output file
# through SaveAs / ExportAsFixedFormat) has no commit point, so Excel writes
# after the dispatcher answered TIMEOUT "nothing was saved". Reachable with
# shipped settings whenever the call bound is shorter than the COM timeout
# (KS4XL_CALL_TIMEOUT_S / KS4XL_MAX_CALL_TIMEOUT_S below KS4XL_COM_TIMEOUT, or
# a caller timeout_seconds above the 900 s ceiling). Private hidden pool only.

_COM = com_session.com_available()[0]


@pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool)")
@pytest.mark.timeout(300)
@pytest.mark.parametrize("case", ["com_autofit", "com_export_pdf"])
def test_fw3_timeout_on_an_excel_save_means_nothing_was_written(case, xlsx, tmp_path,
                                                                 monkeypatch):
    monkeypatch.setattr(dispatch, "BACKSTOP_GRACE_S", 0.2)
    from xlsx_mcp.ops import comtier
    _slow(monkeypatch, comtier._session, "open_workbook", seconds=6.0)
    out = tmp_path / "v.pdf"
    before = _sha(xlsx)
    if case == "com_autofit":
        args = {"path": xlsx}
    else:
        args = {"path": xlsx, "output": str(out)}
    h = fh.start(Settings(call_timeout_s=2.0, max_call_timeout_s=2.0))
    payload, err, _ = h.call("xl_app", {"op": case, "args": args})
    assert err and payload["error"]["code"] == "TIMEOUT", payload
    assert payload["error"]["details"]["saved"] is False
    _wait_for_worker_threads(120)
    end = time.monotonic() + 120  # the COM worker finishes on its own thread
    while time.monotonic() < end and com_session.get_executor().status().get("busy"):
        time.sleep(0.5)
    time.sleep(2.0)
    assert _sha(xlsx) == before, f"{case}: Excel saved the workbook after TIMEOUT"
    assert not out.exists(), f"{case}: Excel wrote {out.name} after TIMEOUT"


# --------------------------------------------------------------------------
# An unknown op refuses NOT_FOUND with the WORKBOOK hint ("re-run
# get_workbook_metadata or get_grid_view ... sheets ... anchors"), which sends
# the caller to reread the file instead of to xl_help.

def test_f5_unknown_op_hint_points_at_help_not_the_workbook():
    h = fh.start()
    payload, err, _ = h.call("xl_read", {"op": "read_rnage", "args": {}})
    assert err and payload["error"]["code"] == "NOT_FOUND"
    hint = payload["error"]["hint"]
    assert "xl_help" in hint and "get_grid_view" not in hint, hint


# --------------------------------------------------------------------------
# The hints a caller receives name ops without the tool
# that holds them, and HINTS = {**NEW_HINTS, **_SHIPPED_HINTS} means the copy
# pass can never replace a shipped hint.

_IDENT = re.compile(r"\b[a-z][a-z0-9]*(?:_[a-z0-9]+)+\b")


def test_fw2_every_op_a_hint_names_comes_with_its_tool():
    ops = set(registry.public_ops())
    bad = []
    for code, text in envelope.HINTS.items():
        for word in set(_IDENT.findall(text)):
            if word in ops and groups.OP_GROUP[word] not in text:
                bad.append(f"{code}: {word} (tool {groups.OP_GROUP[word]} not named)")
    assert bad == [], bad


def test_fw2_a_copy_hint_for_a_shipped_code_wins(monkeypatch):
    import importlib

    from xlsx_mcp.surface import copy
    monkeypatch.setitem(copy.NEW_HINTS, "NOT_FOUND", "COPY-WINS")
    env = importlib.reload(envelope)
    try:
        assert env.HINTS["NOT_FOUND"] == "COPY-WINS"
    finally:
        monkeypatch.undo()
        importlib.reload(envelope)


# --------------------------------------------------------------------------
# An array-formula refusal tells the caller to "edit
# the workbook through xl_app, where Excel applies its own array semantics";
# xl_app has no op that writes a cell (the neighbouring refusals and the
# HAZARD_REFUSED hint say so). A refusal must not name a route that does not
# exist.

def _array_book(p: Path) -> str:
    from openpyxl.worksheet.formula import ArrayFormula
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B1"], ws["B2"] = 1, 2
    ws["A1"] = ArrayFormula("A1:A2", "=B1:B2*2")
    wb.save(p)
    return str(p)


def test_fx1_array_refusal_names_no_route_that_does_not_exist(tmp_path):
    p = _array_book(tmp_path / "arr.xlsx")
    h = fh.start()
    payload, err, _ = h.call("xl_structure", {"op": "sort_range",
                                              "args": {"path": p, "location": "A1:B2",
                                                       "keys": [{"column": "B"}]}})
    assert err, payload
    text = payload["error"]["message"] + " " + payload["error"].get("hint", "")
    assert "through xl_app, where Excel applies" not in text, text


# --------------------------------------------------------------------------
# F-A (MAJOR, numbers safety, pre-existing on 1.2.5) set_cell (and the other
# single-cell writers) never consult the array guard: writing into ONE cell
# of a legacy array formula A1:A2 {=B1:B2*2} succeeds and "verifies". At A1
# the whole array is silently deleted (A2's result vanishes); at A2 a static 5
# lands inside the array's ref. Excel's own answer to either edit is the
# refusal "You can't change part of an array" (core/arrays.py docstring). The
# movers, clear and sort already refuse partial array edits; the writers do not.

@pytest.mark.parametrize("cell", ["A1", "A2"])
def test_fa_writing_one_cell_of_an_array_refuses(cell, tmp_path):
    p = _array_book(tmp_path / "arr.xlsx")
    before = _sha(p)
    h = fh.start()
    payload, err, _ = h.call("xl_edit", {"op": "set_cell",
                                         "args": {"path": p, "location": cell, "value": 5}})
    assert err, f"set_cell {cell} inside array A1:A2 succeeded: {payload}"
    assert _sha(p) == before


# --------------------------------------------------------------------------
# --------------------------------------------------------------------------
# The public-copy release gate runs only with
# KS4XL_RELEASE_GATE=1 "which the release ship gate sets"; nothing sets it
# (publish.yml runs no tests, tests.yml never sets it, the edition scripts
# never run pytest), so a release can ship with the gate never run.

def test_f6_the_release_path_runs_the_public_copy_gate():
    hits = []
    for p in list((ROOT / ".github" / "workflows").glob("*.yml")) + \
            list((ROOT / "scripts").glob("*.py")):
        if re.search(r"KS4XL_RELEASE_GATE\s*[:=]\s*['\"]?1", p.read_text("utf-8", "replace")):
            hits.append(p.name)
    assert hits, "no release-path file sets KS4XL_RELEASE_GATE=1"


# --------------------------------------------------------------------------
# get_server_info calls platform.platform(), which on Windows
# Python 3.12+ queries WMI (platform._wmi_query) with no bound; a hung WMI
# turns the orient call into TIMEOUT and parks a thread every retry (uname is
# cached only on success). Simulated with a WMI stand-in that never answers.

@pytest.mark.skipif(sys.platform != "win32" or not hasattr(platform, "_wmi_query"),
                    reason="Windows WMI path")
@pytest.mark.timeout(120)
def test_f8_get_server_info_does_not_wait_on_wmi(monkeypatch):
    monkeypatch.setattr(dispatch, "BACKSTOP_GRACE_S", 0.2)
    release = threading.Event()

    def hung(*a, **k):
        release.wait(60)
        raise OSError("stand-in")
    monkeypatch.setattr(platform, "_wmi_query", hung)
    monkeypatch.setattr(platform, "_uname_cache", None)
    from fastmcp import Client
    mcp = srv.create_server(Settings(call_timeout_s=2.0, max_call_timeout_s=2.0))

    async def go():
        async with Client(mcp) as c:
            res = await c.call_tool_mcp("xl_read", {"op": "get_server_info"})
            release.set()
            return res
    res = asyncio.run(go())
    payload = json.loads(res.content[0].text)
    assert not res.isError, f"get_server_info waited on WMI: {payload}"


# --------------------------------------------------------------------------
# The dispatcher runs every op with asyncio.to_thread on the
# default executor (min(32, cpu + 4) threads). A timed-out thread cannot be
# killed, so after that many stuck calls every later call, even a trivial
# one, answers TIMEOUT for the life of the session.

@pytest.mark.timeout(300)
def test_f9_stuck_calls_do_not_starve_every_later_call(monkeypatch):
    monkeypatch.setattr(dispatch, "BACKSTOP_GRACE_S", 0.2)
    release = threading.Event()
    spec = registry.REGISTRY["inspect_vba"]

    def stuck(**_k):
        release.wait(240)
        return {}
    monkeypatch.setitem(registry.REGISTRY, "inspect_vba", replace(spec, handler=stuck))
    n = min(32, (os.cpu_count() or 1) + 4) + 2
    from fastmcp import Client
    mcp = srv.create_server(Settings(call_timeout_s=1.0, max_call_timeout_s=1.0))

    async def go():
        async with Client(mcp) as c:
            for _ in range(n):
                await c.call_tool_mcp("xl_read", {"op": "inspect_vba",
                                                  "args": {"path": "x.xlsm"}})
            res = await c.call_tool_mcp("xl_read", {"op": "get_workflows"})
            release.set()
            return res
    res = asyncio.run(go())
    payload = json.loads(res.content[0].text)
    assert not res.isError, f"a trivial call after {n} stuck calls: {payload}"


# --------------------------------------------------------------------------
# One malformed request before initialize makes
# the server write several KB (mcp's pydantic validation dump) to stderr. A
# client that pipes stderr without draining it blocks the server on that
# write: the server "goes silent".

@pytest.mark.timeout(120)
def test_f10_one_bad_pre_initialize_request_writes_little_to_stderr():
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    proc = subprocess.Popen(fh.launch_cmd(), stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                            stderr=subprocess.PIPE, env=fh._env(None), creationflags=flags)
    err_chunks: list[bytes] = []
    t = threading.Thread(target=lambda: err_chunks.append(proc.stderr.read()), daemon=True)
    t.start()
    try:
        proc.stdin.write(b'{"jsonrpc":"2.0","id":1,"method":"nope/unknown","params":{}}\n')
        proc.stdin.flush()
        answer = proc.stdout.readline()
        proc.stdin.write(json.dumps({"jsonrpc": "2.0", "id": 2, "method": "initialize",
                                     "params": {"protocolVersion": "2025-06-18",
                                                "capabilities": {},
                                                "clientInfo": {"name": "v", "version": "0"}}}
                                    ).encode() + b"\n")
        proc.stdin.flush()
        init = proc.stdout.readline()
    finally:
        proc.stdin.close()
        try:
            proc.wait(timeout=30)
        except subprocess.TimeoutExpired:
            proc.kill()
            proc.wait(timeout=30)
        t.join(30)
    assert answer and init, "the server did not answer"
    size = len(b"".join(err_chunks))
    assert size <= 1024, f"one bad request wrote {size} bytes to stderr"


# --------------------------------------------------------------------------
# test_op_sweep's output-path test calls every output site
# with the source or the prev slot as target but WITHOUT overwrite, so an
# existing target refuses CONFLICT whether or not the outguard exists; the
# test cannot fail if the guard is removed. Here each site is called with
# overwrite:true (where the op takes it), so only the outguard stands between
# the call and the slot. Green on 40d1b78; red with the guard removed.

_OVERWRITE_SITES = {
    "copy_workbook": ("xl_workbook", "dst", {}),  # slot only: dst=a normal workbook is legal
    "export_range": ("xl_workbook", "out_file", {"location": "A1:B3", "fmt": "csv"}),
    "export_file": ("xl_workbook", "out_file", {"fmt": "json"}),
    "com_export_pdf": ("xl_app", "output", {}),
    "com_convert_format": ("xl_app", "output", {}),
}


@pytest.mark.parametrize("op", sorted(_OVERWRITE_SITES))
@pytest.mark.parametrize("target", ["source", "slot"])
def test_f3_the_outguard_itself_refuses_with_overwrite(op, target, xlsx, monkeypatch):
    if op == "copy_workbook" and target == "source":
        pytest.skip("copying another workbook over a normal workbook with overwrite is legal")
    # No Excel is started: the guard fires before any COM call; if it did not,
    # the stand-in below refuses instead of reaching Excel.
    monkeypatch.setattr(com_session, "run_com",
                        lambda *a, **k: (_ for _ in ()).throw(AssertionError("reached COM")))
    srv.set_cell(xlsx, "C1", "x")
    slot = safesave.slot_dir(xlsx) / safesave.PREV_SLOT
    tool, key, extra = _OVERWRITE_SITES[op]
    dest = xlsx if target == "source" else str(slot)
    args = {"path": xlsx, **extra, key: dest, "overwrite": True}
    if op == "copy_workbook":
        other = Path(xlsx).with_name("other.xlsx")
        srv.create_workbook(str(other))
        args = {"src": str(other), "dst": dest, "overwrite": True}
    before = _sha(xlsx), _sha(slot)
    h = fh.start()
    payload, err, _ = h.call(tool, {"op": op, "args": args})
    assert err, payload
    msg = payload["error"]["message"].lower()
    assert "backup" in msg or "source" in msg or "same file" in msg, payload
    assert (_sha(xlsx), _sha(slot)) == before


# --------------------------------------------------------------------------
# The backup-store guards (outguard._in_backup_store,
# lifecycle._refuse_backup_store_target) are LEXICAL: they look for a path
# part named .ks4xl-backups in the unresolved path, so a junction (or an 8.3
# short name) into the store passes, and copy_workbook overwrite:true replaces
# the prev slot, destroying the undo it held.

@pytest.mark.skipif(sys.platform != "win32", reason="junctions are Windows")
def test_fj_a_junction_into_the_backup_store_is_still_the_backup_store(xlsx, tmp_path):
    srv.set_cell(xlsx, "C1", "x")
    slot_dir = safesave.slot_dir(xlsx)
    slot = slot_dir / safesave.PREV_SLOT
    link = tmp_path / "innocent"
    subprocess.run(["cmd", "/c", "mklink", "/J", str(link), str(slot_dir)], check=True,
                   capture_output=True, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    other = tmp_path / "other.xlsx"
    srv.create_workbook(str(other))
    before = _sha(slot)
    h = fh.start()
    payload, err, _ = h.call("xl_workbook", {"op": "copy_workbook", "args": {
        "src": str(other), "dst": str(link / safesave.PREV_SLOT), "overwrite": True}})
    try:
        assert err, f"copy_workbook overwrote the prev slot through a junction: {payload}"
        assert _sha(slot) == before
    finally:
        os.rmdir(link)


# --------------------------------------------------------------------------
# A refused --owner-test build leaves every built
# artifact in --output, where the next step (or a person) can pick up a
# wheel the gate refused.

@pytest.mark.timeout(900)
def test_f11_a_refused_owner_test_build_leaves_nothing_shippable(tmp_path):
    # This test was first written against a tree whose copy was unfilled.
    # Now that the real copy is filled, the build no longer refuses and the
    # original premise is gone. It builds a copy of the tracked tree with one
    # copy slot unfilled (the same thing the old tree was), so it always runs;
    # both assertions are unchanged.
    from test_deadline_sites_and_served_hints import build_env, tracked_tree_with_an_unfilled_slot
    tree = tracked_tree_with_an_unfilled_slot(tmp_path / "tree")
    out = tmp_path / "out"
    res = subprocess.run([sys.executable, "-X", "utf8", str(tree / "scripts" / "build_editions.py"),
                          "--edition", "community", "--owner-test", "--release-date", "2026-10-01",
                          "--output", str(out)],
                         capture_output=True, text=True, cwd=tree, timeout=880, env=build_env(),
                         stdin=subprocess.DEVNULL,
                         creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    assert res.returncode != 0, "the owner-test gate did not refuse the unfilled copy"
    left = [p.name for p in out.rglob("*") if p.suffix in (".whl", ".mcpb")]
    assert left == [], f"the refused build left {left}"


# --------------------------------------------------------------------------
# GAP guards for survivors of a mutation pass (green on the shipped product,
# each red under its mutant).
# (a) a deadline that has EXPIRED (not only one the backstop cancelled)
#     refuses the commit; mutant: try_commit checks only _cancelled.
# (b) an answer the output cap cut always carries its page block; mutant:
#     enforce drops `page` (test_output_cap_cuts_with_a_page_block accepts
#     page missing).
# (c) xl_help answers obey the output cap too; mutant: run_help skips enforce.

def test_fm_a_an_expired_deadline_refuses_the_commit():
    from xlsx_mcp.core import deadline as dl
    d = dl.Deadline(0.05)
    time.sleep(0.2)
    assert not d.cancelled
    assert d.try_commit() is False


def test_fm_b_a_cut_answer_always_says_what_it_left_out(tmp_path):
    p = tmp_path / "big.xlsx"
    srv.create_workbook(str(p))
    srv.write_range(str(p), "A1", [["x" * 50] * 10 for _ in range(300)])
    h = fh.start(Settings(max_output_chars=4000))
    out, err, res = h.call("xl_read", {"op": "read_range",
                                       "args": {"path": str(p), "location": "A1:J300"}})
    assert not err and len(res.content[0].text) <= 4000
    assert out["page"]["truncated"] is True and out["page"]["omitted"] > 0


def test_fm_c_help_answers_obey_the_output_cap():
    h = fh.start(Settings(max_output_chars=2000))
    out, err, res = h.call("xl_help", {"op": "apply_edits"})
    assert len(res.content[0].text) <= 2000


# --------------------------------------------------------------------------
# F-D (MINOR, pre-existing hang, now bounded) an INPUT path naming a Windows
# console device ("CON", "CONIN$") blocks the op on a console read until the
# call's deadline: every path-taking op answers TIMEOUT after the full bound
# (300 s by default) and leaves a thread stuck for the life of the server,
# so a few retries pile up. The outguard already refuses device names as
# OUTPUT targets; inputs are not checked. Over stdio so the stuck thread dies
# with the server process.

@pytest.mark.skipif(sys.platform != "win32", reason="Windows device names")
@pytest.mark.timeout(120)
@pytest.mark.parametrize("device", ["CON", "CONIN$"])
def test_fd_a_console_device_input_path_refuses_instead_of_hanging(device):
    res = fh.stdio_session([("xl_read", {"op": "read_range",
                                         "args": {"path": device, "location": "A1"}})],
                           env={"KS4XL_CALL_TIMEOUT_S": "3"}, timeout=90)
    body = json.loads(res["results"][0]["text"])
    assert body["error"]["code"] != "TIMEOUT", body


# --------------------------------------------------------------------------
# manage_image action='extract'
# writes image<N>.<ext> into out_dir and silently replaces an existing file of
# that name: no overwrite flag, no refusal, no .bak. Every other output site
# (core/outguard.guard_out_file) refuses an existing target without
# overwrite:true and keeps a timestamped .bak when it replaces one.

def test_fo_extract_never_silently_replaces_an_existing_file(tmp_path):
    import base64
    png = tmp_path / "p.png"
    png.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="))
    book = tmp_path / "i.xlsx"
    srv.create_workbook(str(book))
    srv.manage_image(str(book), "insert", image_file=str(png), location="B2")
    out = tmp_path / "out"
    out.mkdir()
    first = srv.manage_image(str(book), "extract", out_dir=str(out))
    victim = Path(first["extracted"][0])
    victim.write_text("a user's own file", encoding="utf-8")
    h = fh.start()
    payload, err, _ = h.call("xl_format", {"op": "manage_image", "args": {
        "path": str(book), "action": "extract", "out_dir": str(out)}})
    kept = victim.read_text("utf-8", errors="replace") == "a user's own file"
    backed_up = any(p.name.endswith(".bak") for p in out.iterdir())
    assert err or kept or backed_up, "extract replaced an existing file with no flag and no .bak"


# --------------------------------------------------------------------------
# import_data with an inline CSV whose one field exceeds the
# csv module's 131,072-character field limit answers INTERNAL_ERROR "... The
# arguments are not the cause", which is false: the argument is exactly the
# cause (csv.Error falls through CODE_MAP). Fuzzer find, 2 of 2.

def test_f4_an_oversized_csv_field_is_not_an_internal_error(xlsx):
    h = fh.start()
    payload, err, _ = h.call("xl_edit", {"op": "import_data", "args": {
        "path": xlsx, "source": "A" * 300000, "fmt": "csv", "location": "J1"}})
    assert err, payload
    assert payload["error"]["code"] != "INTERNAL_ERROR", payload


# F-X1 (second instance) a workbook held open refuses WORKBOOK_LOCKED with
# "close it, or use a live/COM route, then retry": Excel has no live route,
# and every xl_app op refuses the same open workbook (verified with a private
# hidden Excel holding it, 2 of 2). Held here with an exclusive file lock.

@pytest.mark.skipif(sys.platform != "win32", reason="Windows file locks")
def test_fx1_the_locked_refusal_names_no_route_that_does_not_exist(xlsx):
    import win32file
    Path(xlsx).with_name("~$" + Path(xlsx).name).write_bytes(b"owner")  # Excel's owner lockfile
    handle = win32file.CreateFile(xlsx, win32file.GENERIC_READ, win32file.FILE_SHARE_READ, None,
                                  win32file.OPEN_EXISTING, 0, None)
    try:
        h = fh.start()
        payload, err, _ = h.call("xl_edit", {"op": "set_cell",
                                             "args": {"path": xlsx, "location": "E1", "value": 1}})
    finally:
        handle.Close()
    assert err and payload["error"]["code"] == "WORKBOOK_LOCKED", payload
    assert "live/COM route" not in payload["error"]["message"], payload["error"]["message"]


# --------------------------------------------------------------------------
# com_render_sheet in a session without a Windows clipboard
# refuses BAD_PARAMS, whose hint says "the arguments did not match this op;
# call xl_help": the environment is the cause, not the arguments (the message
# itself is accurate). Seen for real in this verification's session (stdio,
# 2 of 2); simulated here.

def test_fp8_a_missing_clipboard_does_not_blame_the_arguments(xlsx, tmp_path, monkeypatch):
    from xlsx_mcp.ops import comtier
    monkeypatch.setattr(comtier, "_clipboard_available", lambda: False)
    h = fh.start()
    payload, err, _ = h.call("xl_app", {"op": "com_render_sheet", "args": {
        "path": xlsx, "output": str(tmp_path / "r.png")}})
    assert err, payload
    assert payload["error"]["code"] != "BAD_PARAMS", payload["error"]


# F-J (second route, pre-existing, same class on Word) any mutating op given
# a backup slot itself as `path` edits the slot in place, destroying the undo.

def test_fj_a_mutating_op_never_edits_a_backup_slot_in_place(xlsx):
    srv.set_cell(xlsx, "C1", "x")
    slot = safesave.slot_dir(xlsx) / safesave.PREV_SLOT
    before = _sha(slot)
    h = fh.start()
    payload, err, _ = h.call("xl_edit", {"op": "set_cell",
                                         "args": {"path": str(slot), "location": "A1", "value": 999}})
    assert err, f"set_cell edited the prev slot in place: {payload}"
    assert _sha(slot) == before


# --------------------------------------------------------------------------
# com_export_pdf with an output path
# that lacks ".pdf": the outguard checks the path as given, Excel's
# ExportAsFixedFormat then writes "<output>.pdf", and an existing file of that
# name is silently replaced (no overwrite flag, no .bak); the tool answers
# VALIDATION_FAILED "the PDF is missing or empty" and leaves the new file.
# Naming the same file directly refuses CONFLICT. stdio 2 of 2, private pool.

@pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool)")
@pytest.mark.timeout(300)
def test_fe_an_extensionless_pdf_output_never_replaces_a_users_file(xlsx, tmp_path):
    victim = tmp_path / "report.pdf"
    victim.write_bytes(b"a user's own pdf")
    h = fh.start()
    payload, err, _ = h.call("xl_app", {"op": "com_export_pdf", "args": {
        "path": xlsx, "output": str(tmp_path / "report")}})
    assert victim.read_bytes() == b"a user's own pdf", \
        f"Excel replaced report.pdf through output='report': {payload}"
