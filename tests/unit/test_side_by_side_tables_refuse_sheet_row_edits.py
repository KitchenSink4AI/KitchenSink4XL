"""A sheet-row insert or delete through tables side by side refuses, as in Excel.

Excel refuses to insert or delete a whole sheet row inside the rows of two
tables that sit side by side ("Insert method of Range class failed"),
measured with Excel's own tables, with and without filter buttons, with
values only and with unequal heights. It takes the same edit above or below
both tables, through one table, through stacked tables and through one
table beside plain data. The product did the refused edit and saved a book
Excel would never have produced; it now refuses it with the way round.
All content is made up.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.table import Table

from xlsx_mcp.core.errors import UnsupportedStructure
from xlsx_mcp.ops import structure


def _book(tmp_path: Path, layout: str) -> str:
    p = tmp_path / f"{layout}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    left = [("Item", "Qty"), ("a", 1), ("b", 2), ("c", 3)]
    for r, row in enumerate(left, start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    ws.add_table(Table(displayName="Left", ref="A1:B4"))
    second = [("Code", "Units"), ("k", 4), ("l", 5), ("m", 6)]
    if layout == "stacked":
        for r, row in enumerate(second, start=7):
            for c, v in enumerate(row, start=1):
                ws.cell(r, c, v)
        ws.add_table(Table(displayName="Right", ref="A7:B10"))
    else:
        for r, row in enumerate(second, start=1):
            for c, v in enumerate(row, start=7):
                ws.cell(r, c, v)
        if layout == "side":
            ws.add_table(Table(displayName="Right", ref="G1:H4"))
    wb.save(p)
    return str(p)


@pytest.mark.parametrize("action,at", [("insert_rows", 2), ("insert_rows", 3),
                                       ("insert_rows", 4), ("delete_rows", 3)])
def test_a_row_edit_through_two_tables_side_by_side_refuses(tmp_path, action,
                                                            at):
    p = _book(tmp_path, "side")
    before = Path(p).read_bytes()
    with pytest.raises(UnsupportedStructure) as ei:
        structure.modify_grid_structure(p, action, at, 1, sheet="Data")
    msg = str(ei.value)
    assert "Left" in msg and "Right" in msg and "Excel" in msg
    assert Path(p).read_bytes() == before


@pytest.mark.parametrize("layout,action,at", [
    ("side", "insert_rows", 1),        # above both headers: both move down
    ("side", "insert_rows", 5),        # below both
    ("side", "insert_cols", "F"),      # between them
    ("stacked", "insert_rows", 3),     # through one table only
    ("stacked", "insert_rows", 8),
    ("onedata", "insert_rows", 3),     # one table beside plain data
])
def test_the_edits_excel_takes_still_work(tmp_path, layout, action, at):
    p = _book(tmp_path, layout)
    structure.modify_grid_structure(p, action, at, 1, sheet="Data")
