"""A COM recalculation never reaches the network unless the caller asks.

Excel's automation security stops macros, not worksheet functions, and a few
functions fetch from the network as Excel calculates them (WEBSERVICE, often
inside FILTERXML, IMAGE, STOCKHISTORY, RTD, the CUBE functions, Python in
Excel); an external data connection can refresh from its server. recalculate
and com_goal_seek refuse such a workbook before Excel starts, naming the
cells and the functions, unless allow_network=true.

NO test here starts Excel: Excel availability is stubbed, and the COM run
itself is replaced by a stand-in that records whether it was reached. Every
workbook is made up and built by code.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl
import pytest

from xlsx_mcp.com import session as com_session
from xlsx_mcp.core.errors import FormulaRejected
from xlsx_mcp.ops import comtier

CONNECTIONS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
    '<connections xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    '<connection id="1" name="Rates feed" type="4" refreshedVersion="8" '
    'refreshOnLoad="1"><webPr url="https://example.com/rates.html"/></connection>'
    "</connections>")


def _book(tmp_path: Path, name: str, *, network: bool = True,
          connections: bool = False) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws["A1"] = "https://example.com/rate?cur=KRW"
    ws["B1"] = 1250
    # a function name inside text is not a call
    ws["C1"] = '="WEBSERVICE(" & A1 & ")"'
    ws["D1"] = "=B1*2"
    if network:
        ws["B2"] = '=FILTERXML(WEBSERVICE(A1),"//rate")'
        other = wb.create_sheet("Board Pics")
        other["A1"] = '=_xlfn.IMAGE("https://example.com/logo.png")'
        wb.defined_names["Feed"] = openpyxl.workbook.defined_name.DefinedName(
            "Feed", attr_text='_xlfn.WEBSERVICE("https://example.com/feed")')
    p = tmp_path / name
    wb.save(p)
    wb.close()
    if connections:
        with zipfile.ZipFile(p) as z:
            parts = {i.filename: z.read(i.filename) for i in z.infolist()}
        parts["xl/connections.xml"] = CONNECTIONS.encode("utf-8")
        with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
            for n, d in parts.items():
                z.writestr(n, d)
    return str(p)


@pytest.fixture
def no_excel(monkeypatch):
    """Excel 'available', and the COM run replaced by a recorder."""
    calls: list[str] = []
    monkeypatch.setattr(com_session, "com_available",
                        lambda **_: (True, "stubbed for the test"))

    def fake_run(label, path, body, **_kw):
        calls.append(label)
        return {"ok": True, "file": path, "saved": True, "verified": True}

    monkeypatch.setattr(comtier, "_run_mutation", fake_run)
    monkeypatch.setattr(comtier, "_run_readonly", fake_run)
    monkeypatch.setattr(comtier, "_count_absent_cached", lambda _p: 0)
    return calls


def test_recalc_refuses_network_formulas_before_excel_starts(tmp_path, no_excel):
    p = _book(tmp_path, "rates.xlsx")
    with pytest.raises(FormulaRejected) as exc:
        comtier.recalculate(p, engine="com")
    msg = str(exc.value)
    assert no_excel == [], "Excel must not be started"
    assert "Rates!B2 (WEBSERVICE)" in msg
    assert "'Board Pics'!A1 (IMAGE)" in msg
    assert "the defined name Feed (WEBSERVICE)" in msg
    assert "allow_network=true" in msg
    assert "C1" not in msg, "a function name inside text is not a call"
    assert exc.value.detail["override"] == "allow_network=true"


def test_auto_engine_refuses_too(tmp_path, no_excel):
    p = _book(tmp_path, "rates.xlsx")
    with pytest.raises(FormulaRejected):
        comtier.recalculate(p)
    assert no_excel == []


def test_external_data_connection_refuses(tmp_path, no_excel):
    p = _book(tmp_path, "conn.xlsx", network=False, connections=True)
    with pytest.raises(FormulaRejected) as exc:
        comtier.recalculate(p, engine="com")
    assert "Rates feed" in str(exc.value) and no_excel == []


def test_allow_network_runs_and_says_what_it_allowed(tmp_path, no_excel):
    p = _book(tmp_path, "rates.xlsx", connections=True)
    r = comtier.recalculate(p, engine="com", allow_network=True)
    assert no_excel and r["recalculated"] is True
    wheres = [f["where"] for f in r["network_reach"]["formulas"]]
    assert "Rates!B2" in wheres and r["network_reach"]["connections"] == ["Rates feed"]
    assert r["network_allowed"] is True and "allow_network was set" in r["network_note"]
    assert any("allow_network=true" in w for w in r["warnings"])


def test_a_workbook_that_cannot_reach_the_network_recalculates(tmp_path, no_excel):
    p = _book(tmp_path, "plain.xlsx", network=False)
    r = comtier.recalculate(p, engine="com")
    assert no_excel and "network_allowed" not in r


def test_goal_seek_refuses_the_same_way(tmp_path, no_excel):
    p = _book(tmp_path, "rates.xlsx")
    with pytest.raises(FormulaRejected) as exc:
        comtier.com_goal_seek(p, "D1", 3000, "B1", sheet="Rates")
    assert "Rates!B2 (WEBSERVICE)" in str(exc.value) and no_excel == []
    r = comtier.com_goal_seek(p, "D1", 3000, "B1", sheet="Rates", save=False,
                              allow_network=True)
    assert no_excel and r["network_allowed"]


@pytest.mark.parametrize("formula, calls", [
    ('FILTERXML(WEBSERVICE("https://example.com/x"),"//a")', ["WEBSERVICE"]),
    ('_xlfn.WEBSERVICE(A1)', ["WEBSERVICE"]),
    ('_xlfn.STOCKHISTORY("MADEUP",DATE(1999,1,1))', ["STOCKHISTORY"]),
    ('RTD("madeup.rtdserver",,"topic")', ["RTD"]),
    ('_xlfn._xlws.PY(0,1)', ["PY"]),
    ('PY(1)', []),  # a function of one's own named PY
    ('"IMAGE(" & A1', []),
    ("'IMAGE(1)'!A1+1", []),  # a sheet named like a call
    ("MYIMAGE(1)+IMAGEX(2)", []),
    ('CUBEVALUE("Sales","[Measures].[Total]")', ["CUBEVALUE"]),
    ("SUM(A1:A3)", []),
])
def test_calls_in_reads_calls_not_text(formula, calls):
    from xlsx_mcp.core import netreach
    assert netreach.calls_in(formula) == calls


def test_shared_formula_is_reported_by_its_range(tmp_path):
    p = _book(tmp_path, "shared.xlsx", network=False)
    with zipfile.ZipFile(p) as z:
        parts = {i.filename: z.read(i.filename) for i in z.infolist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode("utf-8")
    shared = ('<row r="5"><c r="E5"><f t="shared" ref="E5:E9" si="0">'
              '_xlfn.IMAGE(A1)</f></c></row><row r="6"><c r="E6">'
              '<f t="shared" si="0"/></c></row>')
    sheet = sheet.replace("</sheetData>", shared + "</sheetData>")
    parts["xl/worksheets/sheet1.xml"] = sheet.encode("utf-8")
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in parts.items():
            z.writestr(n, d)
    from xlsx_mcp.core import netreach
    reach = netreach.network_reach(p)
    assert reach["formulas"] == [{"where": "Rates!E5:E9", "functions": ["IMAGE"]}]


def test_each_place_is_reported_once_and_rules_are_found(tmp_path):
    # an empty cell written as <c/> ahead of the formula used to end the
    # sheetData strip early, so the cell came back a second time as a "rule"
    p = _book(tmp_path, "rules.xlsx", network=False)
    with zipfile.ZipFile(p) as z:
        parts = {i.filename: z.read(i.filename) for i in z.infolist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode("utf-8")
    cells = ('<row r="7"><c r="A7"/><c r="B7"><f>WEBSERVICE(A1)</f></c></row>')
    rule = ('<conditionalFormatting sqref="D1:D9"><cfRule type="expression" '
            'priority="1"><formula>LEN(_xlfn.WEBSERVICE($A$1))&gt;0</formula>'
            "</cfRule></conditionalFormatting>")
    sheet = sheet.replace("</sheetData>", cells + "</sheetData>" + rule)
    parts["xl/worksheets/sheet1.xml"] = sheet.encode("utf-8")
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in parts.items():
            z.writestr(n, d)
    from xlsx_mcp.core import netreach
    reach = netreach.network_reach(p)
    assert reach["formulas"] == [
        {"where": "Rates!B7", "functions": ["WEBSERVICE"]},
        {"where": "a conditional format or data validation rule on Rates",
         "functions": ["WEBSERVICE"]}]


def test_the_served_hint_is_true_for_the_network_refusal(tmp_path, no_excel):
    """The refusal reaches the caller with FORMULA_REJECTED's one recovery
    hint, so that hint, and the code's line in docs/llms.txt, must also cover
    a recalculation the network gate stopped, not only a blocked write."""
    from xlsx_mcp import envelope
    p = _book(tmp_path, "rates.xlsx")
    with pytest.raises(FormulaRejected) as exc:
        comtier.recalculate(p, engine="com")
    err = envelope.refusal(exc.value)["error"]
    assert err["code"] == "FORMULA_REJECTED" and no_excel == []
    assert "network" in err["hint"], err["hint"]
    llms = (Path(__file__).resolve().parents[2] / "docs" / "llms.txt").read_text(
        encoding="utf-8")
    line = next(ln for ln in llms.splitlines() if ln.startswith("- `FORMULA_REJECTED`"))
    assert "network" in line, line


def test_help_documents_allow_network():
    from xlsx_mcp import registry
    from xlsx_mcp import server  # noqa: F401  (the import registers ops)
    import inspect
    for op in ("recalculate", "com_goal_seek"):
        spec = registry.public_ops()[op]
        assert "allow_network" in spec.args.model_fields
        assert spec.args.model_fields["allow_network"].default is False
        assert ("Set it only after the user has agreed, for this workbook."
                in " ".join(inspect.getdoc(spec.handler).split()))
