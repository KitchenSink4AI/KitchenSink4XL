"""Renaming a sheet or a table keeps every formula's stored result.

A rename rewrites formula text (=Data!A1/2+10 becomes =Model!A1/2+10) but
changes no result, and Excel's own rename keeps every value. The save keyed
the stored results on the formula text being unchanged, and looked the
renamed sheet up by its old name, so every formula on the renamed sheet,
and every formula elsewhere naming it, lost its stored result. Under
iterative calculation that is not cosmetic: Excel seeds a circular model
from those results, and an iterative model whose sheet was renamed opened
showing #VALUE! where it showed 20 (measured in Excel).

The results now follow a rename, and so do the stored digits of
17-significant-digit numbers on the renamed sheet. All data is made up.
"""

from __future__ import annotations

import re
import zipfile

import openpyxl
import pytest

from xlsx_mcp.ops import lifecycle, tables


def _patch(path, part, fn):
    with zipfile.ZipFile(path) as zin:
        parts = {i.filename: zin.read(i.filename) for i in zin.infolist()}
    parts[part] = fn(parts[part].decode("utf-8")).encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)


def _iterative_model(tmp_path):
    """A circular model as Excel stores it: iterate="1", every formula with
    its converged result (the seed Excel iterates from)."""
    p = tmp_path / "loop.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "=Data!A1/2+10"
    ws["B1"] = "=Data!A1*2"
    ws["C1"] = "=A1+1"
    ws["A2"] = 5
    ws["A3"] = 0.30000000000000004  # needs 17 significant digits
    out = wb.create_sheet("Out")
    out["A1"] = "=Data!B1+1"
    wb.calculation.iterate = True
    wb.calculation.iterateCount = 100
    wb.calculation.iterateDelta = 0.001
    wb.save(p)
    cached = {"A1": "20", "B1": "40", "C1": "21"}

    def seed(x):
        for ref, v in cached.items():
            x = re.sub(r'(<c r="%s"[^>]*>)(<f>[^<]*</f>)(<v>\s*</v>|<v/>)?'
                       % ref, r"\1\2<v>%s</v>" % v, x)
        return x

    _patch(p, "xl/worksheets/sheet1.xml", seed)
    _patch(p, "xl/worksheets/sheet2.xml", lambda x: re.sub(
        r'(<c r="A1"[^>]*>)(<f>[^<]*</f>)(<v>\s*</v>|<v/>)?',
        r"\1\2<v>41</v>", x))
    _patch(p, "xl/worksheets/sheet1.xml", lambda x: x.replace(
        "<v>0.3</v>", "<v>0.30000000000000004</v>"))
    return str(p)


def _stored(path, sheet_title):
    """{address: (formula, stored value)} of one sheet, found by title."""
    wb = openpyxl.load_workbook(path)
    try:
        idx = wb.sheetnames.index(sheet_title) + 1
    finally:
        wb.close()
    with zipfile.ZipFile(path) as zf:
        x = zf.read(f"xl/worksheets/sheet{idx}.xml").decode("utf-8")
    out = {}
    for m in re.finditer(r'<c r="([A-Z]+\d+)"[^>]*>(.*?)</c>', x, re.S):
        f = re.search(r"<f>(.*?)</f>", m.group(2))
        v = re.search(r"<v>(.*?)</v>", m.group(2))
        out[m.group(1)] = (f.group(1) if f else None, v.group(1) if v else None)
    return out


def test_renaming_the_sheet_of_an_iterative_model_keeps_its_seeds(tmp_path):
    p = _iterative_model(tmp_path)
    lifecycle.manage_worksheet(p, "rename", sheet="Data", new_name="Model")
    model = _stored(p, "Model")
    assert model["A1"] == ("Model!A1/2+10", "20")
    assert model["B1"] == ("Model!A1*2", "40")
    assert model["C1"] == ("A1+1", "21")
    assert _stored(p, "Out")["A1"] == ("Model!B1+1", "41")


def test_the_renamed_sheet_keeps_its_full_precision_numbers(tmp_path):
    p = _iterative_model(tmp_path)
    lifecycle.manage_worksheet(p, "rename", sheet="Data", new_name="Model")
    assert _stored(p, "Model")["A3"] == (None, "0.30000000000000004")


def test_renaming_a_table_keeps_the_results_of_formulas_naming_it(tmp_path):
    p = tmp_path / "tbl.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Item", "Qty"])
    ws.append(["Widget", 4])
    ws.append(["Gadget", 3])
    from openpyxl.worksheet.table import Table
    ws.add_table(Table(displayName="Stock", ref="A1:B3"))
    ws["D1"] = "=SUM(Stock[Qty])"
    wb.calculation.iterate = True
    wb.save(p)
    _patch(p, "xl/worksheets/sheet1.xml", lambda x: re.sub(
        r'(<c r="D1"[^>]*>)(<f>[^<]*</f>)(<v>\s*</v>|<v/>)?', r"\1\2<v>7</v>", x))
    tables.manage_table(str(p), "Stock", "rename", new_name="Inventory")
    assert _stored(str(p), "Data")["D1"] == ("SUM(Inventory[Qty])", "7")


def test_a_rename_never_keeps_the_result_of_a_formula_it_did_not_rename(
        tmp_path):
    """Only the rename's own rewrite is excused: a formula whose text
    changed in any other way still loses its stored result."""
    p = _iterative_model(tmp_path)
    from xlsx_mcp.core.package import WorkbookPackage
    pkg = WorkbookPackage.open(p)
    wb = pkg.workbook
    wb["Data"].title = "Model"
    from xlsx_mcp.core import refs
    refs.rename_sheet_refs(wb, "Data", "Model")
    pkg.note_text_rewrite(
        lambda t: refs.rename_sheet_in_formula(t, "Data", "Model"))
    pkg.set_formula("Model", "B1", "=Model!A1*3")
    pkg.save()
    model = _stored(p, "Model")
    assert model["B1"][0] == "Model!A1*3"
    assert model["B1"][1] != "40"
    assert model["A1"] == ("Model!A1/2+10", "20")
