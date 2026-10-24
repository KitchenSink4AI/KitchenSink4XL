"""The name selector on the name shapes real workbooks hold, on a made-up
workbook: a whole column (Data!$A:$A, common for open-ended lists), whole
rows, an intersection (two ranges joined by a space), a union, a 3D name,
a name looked up in another letter case, and a sheet's own name beside a
workbook name that differs only in case.

What Excel answers for each (Names(n).RefersToRange through a private hidden
Excel): $A:$A, $1:$3, $B$2 for the intersection, two areas for the union, no
range for the 3D name, and the same name whatever the case. A whole column
or row is refused here with the bounded range that reads the same data, as
the same range given in A1 is; before, it raised TypeError.
"""

from __future__ import annotations

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName

from xlsx_mcp import server as srv
from xlsx_mcp.core.errors import XlMcpError


def _name(wb, name, text):
    wb.defined_names[name] = DefinedName(name, attr_text=text)


@pytest.fixture
def book(tmp_path):
    wb = openpyxl.Workbook()
    d = wb.active
    d.title = "Data"
    for r in range(1, 11):
        for c in range(1, 4):
            d.cell(r, c, r * 10 + c)
    for nm in ("Jan", "Feb", "Mar"):
        wb.create_sheet(nm)["B2"] = len(nm)
    wb.create_sheet("Summ")["A1"] = 3
    _name(wb, "ColA", "Data!$A:$A")
    _name(wb, "Rows13", "Data!$1:$3")
    _name(wb, "Isect", "Data!$A$1:$B$2 Data!$B$2:$C$3")
    _name(wb, "Apart", "Data!$A$1:$A$2 Data!$C$5:$C$6")
    _name(wb, "Uni", "Data!$A$1:$A$3,Data!$C$1:$C$3")
    _name(wb, "ThreeD", "Jan:Mar!$B$2")
    _name(wb, "RATE", "Data!$C$5")
    _name(wb, "Plain", "Data!$B$2:$C$4")
    d.defined_names["Rate"] = DefinedName("Rate", attr_text="Data!$C$6")
    path = tmp_path / "names.xlsx"
    wb.save(path)
    return str(path)


def _read(path, name, **extra):
    return srv.read_range(path, {"name": name, **extra})


def test_a_whole_column_name_is_refused_with_the_bounded_form(book):
    with pytest.raises(XlMcpError) as e:
        _read(book, "ColA")
    msg = str(e.value)
    assert "whole-column" in msg and "A1:A10" in msg


def test_a_whole_row_name_is_refused_with_the_bounded_form(book):
    with pytest.raises(XlMcpError) as e:
        _read(book, "Rows13")
    msg = str(e.value)
    assert "whole-row" in msg and "A1:C3" in msg


def test_an_intersection_name_reads_its_overlap(book):
    r = _read(book, "Isect")
    assert (r["sheet"], r["range"], r["values"]) == ("Data", "B2", [[22]])


def test_an_intersection_that_does_not_overlap_is_refused(book):
    with pytest.raises(XlMcpError) as e:
        _read(book, "Apart")
    assert "#NULL!" in str(e.value)


def test_a_union_name_is_still_two_ranges(book):
    with pytest.raises(XlMcpError) as e:
        _read(book, "Uni")
    assert "2 disjoint ranges" in str(e.value)


def test_a_3d_name_is_refused_by_what_it_is(book):
    with pytest.raises(XlMcpError) as e:
        _read(book, "ThreeD")
    assert "3D reference" in str(e.value)


def test_names_ignore_letter_case_as_in_excel(book):
    assert _read(book, "plain")["range"] == "B2:C4"
    assert _read(book, "isect")["range"] == "B2"


def test_the_exact_case_wins_and_a_sheet_reads_its_own_name(book):
    assert _read(book, "RATE")["range"] == "C5"          # the workbook's RATE
    assert _read(book, "Rate")["range"] == "C6"          # the sheet's Rate
    assert _read(book, "rate", sheet="Data")["range"] == "C6"   # on Data, its own name
    assert _read(book, "rate", scope="workbook")["range"] == "C5"
