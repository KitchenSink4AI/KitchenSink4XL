"""COM saves keep the workbook's relative hyperlinks and links.

Earlier versions had Excel write the edited workbook with SaveCopyAs into a private
folder in the system's temp area, and the copy was then put in place of the
workbook. Excel writes every relative hyperlink relative to the folder it
saves INTO, so "Source.xlsx" came back as "../../../../Documents/.../
Source.xlsx", a path that does not exist, on every COM route, while each
call answered ok and verified. The copy is now written beside the workbook,
where Excel writes every link exactly as its own Save in place does
(MEASURED: hyperlinks and external links byte-identical). A folder whose
path is too long for that copy is refused before anything changes.

The first tests start the product's private hidden Excel (run them on a
private desktop); the rest need no Excel.
"""
from __future__ import annotations

import os
import re
import time
import zipfile

import openpyxl
import pytest

from xlsx_mcp.com import session as com_session
from xlsx_mcp.core import safesave, swapwait
from xlsx_mcp.core.errors import XlMcpError
from xlsx_mcp.ops import cells, comtier

LIVE = pytest.mark.skipif(not com_session.com_available()[0], reason="needs Excel (COM)")
LINKS = {"B1": "Source.xlsx", "B2": "Sub/Source2.xlsx", "B3": "../Up.xlsx"}


def _book(tmp_path):
    work = tmp_path / "proj" / "work"
    (work / "Sub").mkdir(parents=True)
    for p in (work / "Source.xlsx", work / "Sub" / "Source2.xlsx", tmp_path / "proj" / "Up.xlsx"):
        wb = openpyxl.Workbook()
        wb.active["A1"] = 1
        wb.save(p)
    main = work / "Main.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "index"
    ws["A2"] = 5
    ws["A3"] = "=A2*2"
    for cell, target in LINKS.items():
        ws[cell] = target
        ws[cell].hyperlink = target
    wb.save(main)
    return str(main)


def _targets(p):
    with zipfile.ZipFile(p) as z:
        rels = z.read("xl/worksheets/_rels/sheet1.xml.rels").decode("utf8")
    return sorted(re.findall(r'Target="([^"]*)"[^>]*TargetMode="External"', rels)
                  + re.findall(r'TargetMode="External"[^>]*Target="([^"]*)"', rels))


def _beside(p):
    folder = os.path.dirname(p)
    return sorted(n for n in os.listdir(folder)
                  if n.startswith((".ks4xl-write-", "~$", ".ks4xl-replaced-")))


def _save_folders_of_this_process():
    import tempfile
    prefix = f"{com_session.SAVE_FOLDER_PREFIX}{os.getpid()}-"
    return [n for n in os.listdir(tempfile.gettempdir()) if n.startswith(prefix)]


ROUTES = {
    "com_autofit": lambda p: comtier.com_autofit(p, timeout_seconds=90),
    "recalculate_com": lambda p: comtier.recalculate(p, engine="com", timeout_seconds=90),
    "goal_seek_save": lambda p: comtier.com_goal_seek(
        p, target_cell="A3", target_value=20, changing_cell="A2", save=True,
        timeout_seconds=90),
}


@LIVE
@pytest.mark.parametrize("route", sorted(ROUTES))
def test_relative_hyperlinks_survive_a_com_save(tmp_path, route):
    p = _book(tmp_path)
    before = _targets(p)
    assert before == sorted(LINKS.values())
    try:
        r = ROUTES[route](p)
    finally:
        com_session.get_executor().shutdown()
    assert r["ok"] is True and r["verified"] is True
    assert _targets(p) == before, f"{route} moved the workbook's relative hyperlinks"
    base = os.path.dirname(p)
    assert all(os.path.exists(os.path.normpath(os.path.join(base, t))) for t in before)
    assert _beside(p) == [], "the save left something beside the workbook"
    assert _save_folders_of_this_process() == []


@LIVE
def test_a_password_set_and_removed_keeps_the_links(tmp_path):
    p = _book(tmp_path)
    before = _targets(p)
    try:
        comtier.com_save_with_password(p, "secret", timeout_seconds=90)
        comtier.com_save_with_password(p, "", current_password="secret",
                                       timeout_seconds=90)
    finally:
        com_session.get_executor().shutdown()
    assert _targets(p) == before
    assert _beside(p) == []


def _deep_workbook(tmp_path, folder_length):
    """A workbook in a folder whose full path is `folder_length` characters."""
    folder = str(tmp_path)
    while len(folder) < folder_length:
        room = folder_length - len(folder) - 1  # after the separator
        step = min(40, room)
        if room - step == 1:  # one character left cannot take a separator and a name
            step -= 1
        folder = os.path.join(folder, "d" * max(step, 1))
    os.makedirs(folder, exist_ok=True)
    assert len(folder) == folder_length, len(folder)
    p = os.path.join(folder, "Q3.xlsx")
    wb = openpyxl.Workbook()
    wb.active["A1"] = "kept"
    wb.save(p)
    return p


def test_a_folder_too_deep_for_excels_copy_is_refused_before_anything_changes(tmp_path, monkeypatch):
    p = _deep_workbook(tmp_path, 200)  # Q3.xlsx fits in 218; its copy would not
    assert len(p) <= com_session.EXCEL_PATH_LIMIT
    before = open(p, "rb").read()
    monkeypatch.setattr(com_session, "run_com",
                        lambda *a, **k: pytest.fail("Excel was started"))
    with pytest.raises(XlMcpError) as info:
        comtier.com_autofit(p)
    assert getattr(info.value, "code", None) == "LIMIT_EXCEEDED"
    msg = str(info.value)
    assert "was not changed" in msg and "218" in msg
    assert open(p, "rb").read() == before
    assert sorted(os.listdir(os.path.dirname(p))) == ["Q3.xlsx"], "something was made beside it"
    with pytest.raises(XlMcpError):
        comtier.com_save_with_password(p, "secret")
    assert sorted(os.listdir(os.path.dirname(p))) == ["Q3.xlsx"]


def test_the_copy_goes_beside_the_workbook_under_a_short_swept_name(tmp_path):
    p = _deep_workbook(tmp_path, 150)
    copy = com_session.copy_path_for(p)
    assert os.path.dirname(copy) == os.path.dirname(p)
    name = os.path.basename(copy)
    assert re.fullmatch(r"\.ks4xl-write-[0-9a-f]{6}\.xlsx", name)
    assert name.startswith(safesave.WRITE_TMP_PREFIX)


def test_discard_removes_only_a_copy_by_its_own_name(tmp_path):
    keep_dir = tmp_path / "folder"
    keep_dir.mkdir()
    (keep_dir / "inside.txt").write_text("x")
    other = tmp_path / "Budget.xlsx"
    other.write_bytes(b"user file")
    com_session.discard_saved(str(keep_dir / "inside.txt"))
    com_session.discard_saved(str(other))
    com_session.discard_saved(str(keep_dir))
    assert (keep_dir / "inside.txt").exists() and other.exists()
    copy = tmp_path / ".ks4xl-write-abc123.xlsx"
    copy.write_bytes(b"")
    com_session.discard_saved(str(copy))
    assert not copy.exists()


def test_a_leftover_copy_beside_the_workbook_is_swept_safely(tmp_path):
    book = tmp_path / "Budget.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "v1"
    wb.save(book)
    old_copy = tmp_path / ".ks4xl-write-0a1b2c.xlsx"   # a killed COM save, two hours ago
    young_copy = tmp_path / ".ks4xl-write-3d4e5f.xlsx"  # a save running now
    aside = tmp_path / swapwait.aside_name("Plan.xlsx", "0123456789ab")  # a stopped swap
    for f in (old_copy, young_copy, aside):
        f.write_bytes(b"")
    old = time.time() - 2 * 3600
    for f in (old_copy, aside):
        os.utime(f, (old, old))
    assert cells.set_cell(str(book), "A1", "v2")["ok"]
    assert not old_copy.exists()
    assert young_copy.exists() and aside.exists()
