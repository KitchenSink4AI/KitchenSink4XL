"""Minor behaviours: a range reaching the last row survives an insert, a
number written into a date-formatted cell, numbers Excel reads as a
different double are named, and the COM save path runs the integrity
check."""
from __future__ import annotations

import datetime
import os
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

from xlsx_mcp.core.errors import ValidationFailed
from xlsx_mcp.ops import cells, structure



def test_insert_rows_keeps_a_whole_column_conditional_format(tmp_path):
    p = str(tmp_path / "cf.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 21):
        ws.cell(r, 1, r)
    ws.conditional_formatting.add("A1:A1048576", CellIsRule(
        operator="greaterThan", formula=["10"], fill=PatternFill("solid", start_color="FFFF00")))
    ws["C1"] = "=SUM(A2:A1048576)"
    wb.save(p)
    structure.modify_grid_structure(p, "insert_rows", 3, count=2)
    wb2 = openpyxl.load_workbook(p)
    assert [str(cf.sqref) for cf in wb2.active.conditional_formatting] == ["A1:A1048576"]
    assert wb2.active["C1"].value == "=SUM(A2:A1048576)"


def test_a_range_pushed_wholly_off_the_grid_is_ref():
    from xlsx_mcp.core import refs
    edit = refs.RefEdit(kind=refs.INSERT_ROWS, sheet="S", index=3, count=5) \
        if hasattr(refs, "INSERT_ROWS") else None
    if edit is None:
        pytest.skip("RefEdit constants not exposed")
    assert refs.transpose_formula("=SUM(A1048575:A1048576)", edit, "S") == "=SUM(#REF!)"
    assert refs.transpose_formula("=SUM(A2:A1048576)", edit, "S") == "=SUM(A2:A1048576)"



@pytest.mark.parametrize("route", ["on", "off"])
def test_a_number_written_into_a_date_cell_is_accepted(tmp_path, monkeypatch, route):
    monkeypatch.setenv("KS4XL_SURGICAL_SAVE", route)
    p = str(tmp_path / "d.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = datetime.datetime(2026, 1, 1)
    ws["A1"].number_format = "yyyy-mm-dd"
    ws["A2"] = datetime.time(6, 0)
    ws["A2"].number_format = "hh:mm"
    wb.save(p)
    assert cells.set_cell(p, "A1", 46100)["verified"]
    assert cells.set_cell(p, "A2", 0.75)["verified"]
    got = openpyxl.load_workbook(p).active
    assert got["A1"].value == datetime.datetime(2026, 3, 19)
    assert got["A2"].value == datetime.time(18, 0)


def test_a_different_number_still_fails_read_back():
    from xlsx_mcp.core.verify import _serial_matches
    wb = openpyxl.Workbook()
    assert _serial_matches(46100, datetime.datetime(2026, 3, 19), wb)
    assert not _serial_matches(46101, datetime.datetime(2026, 3, 19), wb)
    assert not _serial_matches(True, datetime.datetime(1899, 12, 31), wb)
    assert not _serial_matches("46100", datetime.datetime(2026, 3, 19), wb)



def test_numbers_excel_reads_differently_are_named(tmp_path):
    p = str(tmp_path / "n.xlsx")
    openpyxl.Workbook().save(p)
    out = cells.set_cell(p, "A1", 9007199254740993)
    assert any("9,007,199,254,740,992" in w and "Sheet!A1" in w for w in out["warnings"])
    out = cells.set_cell(p, "A2", 5e-324)
    assert any("reads it as 0" in w for w in out["warnings"])
    # 2^53 itself warns, since a client rounds a 16-digit number
    # to exactly that value before the call is made; one below it does not.
    out = cells.set_cell(p, "A3", 9007199254740992)
    assert any("9,007,199,254,740,992" in w and "Sheet!A3" in w for w in out["warnings"])
    out = cells.set_cell(p, "A4", 9007199254740991)
    assert not any("9,007,199" in w for w in out["warnings"])



def _inject(path):
    tmp = path + ".inj"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            b = zin.read(i.filename)
            if i.filename == "xl/_rels/workbook.xml.rels":
                b = b.replace(b"</Relationships>",
                              b'<Relationship Id="rIdInj9" Type="http://schemas.openxmlformats.org/'
                              b'officeDocument/2006/relationships/externalLink" '
                              b'Target="externalLinks/externalLink9.xml"/></Relationships>')
            zout.writestr(i, b)
    os.replace(tmp, path)


def test_com_save_path_runs_the_integrity_check(tmp_path, monkeypatch):
    """Excel's save is simulated (no Excel starts): the copy Excel saves
    gains a dangling relationship. The op must refuse, and the
    copy is checked before it replaces the workbook, so the workbook is
    never replaced at all."""
    import shutil

    from xlsx_mcp.com import session as _session
    from xlsx_mcp.ops import comtier
    p = str(tmp_path / "c.xlsx")
    wb = openpyxl.Workbook()
    wb.active["A1"] = 2
    wb.save(p)
    before = Path(p).read_bytes()

    def save_workbook(wb, *, label, path):
        folder = tmp_path / "excel-copy"
        folder.mkdir()
        copy = str(folder / "book.xlsx")
        shutil.copyfile(p, copy)
        _inject(copy)
        return copy

    def fake_run_com(label, job, timeout=None):
        return {"changed": {"autofit": True}, "instance_pid": 0,
                "staged": save_workbook(None, label=label, path=p)}
    monkeypatch.setattr(_session, "run_com", fake_run_com)
    monkeypatch.setattr(_session, "guard_target_closed", lambda path: [])
    with pytest.raises(ValidationFailed, match="broken internal links"):
        comtier.com_autofit(p)
    assert Path(p).read_bytes() == before



def _heights(tmp_path):
    p = str(tmp_path / "h.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 11):
        ws.cell(r, 1, f"row {r}")
    ws.row_dimensions[2].height = 40
    ws.row_dimensions[5].height = 60
    ws.row_dimensions[8].hidden = True
    wb.save(p)
    return p


def test_row_heights_move_with_their_rows_on_insert(tmp_path):
    p = _heights(tmp_path)
    structure.modify_grid_structure(p, "insert_rows", 3, count=2)
    ws = openpyxl.load_workbook(p).active
    assert ws["A7"].value == "row 5" and ws.row_dimensions[7].height == 60
    assert ws.row_dimensions[2].height == 40
    assert ws.row_dimensions[5].height in (None, 15) and not ws.row_dimensions[5].customHeight
    # The new rows take the height of the row above, as Excel's own
    # Rows.Insert does (measured: a row inserted under a 30-point row is 30).
    assert ws.row_dimensions[3].height == 40 and ws.row_dimensions[4].height == 40
    assert ws["A10"].value == "row 8" and ws.row_dimensions[10].hidden
    assert not ws.row_dimensions[8].hidden


def test_row_heights_move_with_their_rows_on_delete(tmp_path):
    p = _heights(tmp_path)
    structure.modify_grid_structure(p, "delete_rows", 3, count=2)
    ws = openpyxl.load_workbook(p).active
    assert ws["A3"].value == "row 5" and ws.row_dimensions[3].height == 60
    assert ws["A6"].value == "row 8" and ws.row_dimensions[6].hidden
    assert ws.row_dimensions[2].height == 40



_FOLDER = "D:/Reports/2026/".replace("/", chr(92)).encode()
_ALT = (b'<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        b'<mc:Choice Requires="x15"><x15ac:absPath url="' + _FOLDER + b'" '
        b'xmlns:x15ac="http://schemas.microsoft.com/office/spreadsheetml/2010/11/ac"/></mc:Choice>'
        b'</mc:AlternateContent>')


def test_rewrite_keeps_the_saved_folder_for_relative_links(tmp_path):
    base = str(tmp_path / "b.xlsx")
    wb = openpyxl.Workbook()
    for r in range(1, 6):
        wb.active.cell(r, 1, r)
    wb.save(base)
    p = str(tmp_path / "a.xlsx")
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            b = zin.read(i.filename)
            if i.filename == "xl/workbook.xml":
                b = b.replace(b'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"',
                              b'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                              b'xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main"', 1)
                k = b.index(b"<workbookPr")
                end = b.index(b">", k) + 1
                b = b[:end] + _ALT + b[end:]
            zout.writestr(i, b)
    out = structure.modify_grid_structure(p, "insert_rows", 2, count=1)
    assert out["route"] == "openpyxl"
    wbx = zipfile.ZipFile(p).read("xl/workbook.xml")
    assert b'absPath url="' + _FOLDER + b'"' in wbx
    assert wbx.index(b"<workbookPr") < wbx.index(b"AlternateContent") < wbx.index(b"<sheets")
    assert b'xmlns:x15="' in wbx
    openpyxl.load_workbook(p)
