"""A save stopped between ReplaceFileW's two moves.

ReplaceFileW moves the workbook aside and then moves the new version in. A
server killed between the two moves (MEASURED: 2 of 81
aimed kills with a backup, 10 of 31 without) left the workbook's path empty;
a read answered "no such workbook" without saying where it went, and the
next save of ANY workbook in that folder deleted the original (the stale-temp
sweep removed every .ks4xl-write-* older than an hour, and a rename keeps the
workbook's own old modification time). With backup=false both versions were
then gone.

Now the aside name says what it holds and no sweep ever matches it; each
swap is written down in a journal first; the next write to the workbook and
the next start of a server put it back; reads name it and never move it.

The kill here is real: a saving process is paused after ReplaceFileW's first
move (its ReplaceFileW stand-in moves the workbook aside the way ReplaceFileW
does, then waits) and killed with TerminateProcess. No Excel.
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import time
from pathlib import Path

import openpyxl
import pytest

from xlsx_mcp.core import safesave, swapwait
from xlsx_mcp.core.errors import WorkbookNotFound
from xlsx_mcp.ops import backups, cells, lifecycle

pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="ReplaceFileW")

TWO_DAYS = 2 * 86400

_VICTIM = r'''
import os, sys, time
from xlsx_mcp.core import safesave
from xlsx_mcp.ops import cells
book, mark, backup = sys.argv[1], sys.argv[2], sys.argv[3] == "1"
real = safesave._replace_file()

def first_move_only(dst, src, bak, flags, a, b):
    os.rename(dst, bak)          # ReplaceFileW's first move
    open(mark, "w").close()      # tell the test we are between the two moves
    time.sleep(120)              # killed here
    return 0

safesave._REPLACE_FILE = (first_move_only, real[1], real[2], real[3])
cells.set_cell(book, "A1", "the edit being saved", backup=backup)
'''


def _book(p, v):
    wb = openpyxl.Workbook()
    wb.active["A1"] = v
    for r in range(2, 40):
        wb.active.cell(r, 1, f"row {r}")
    wb.save(p)


def _a1(p):
    import io
    wb = openpyxl.load_workbook(io.BytesIO(Path(p).read_bytes()), read_only=True)
    try:
        return wb.active["A1"].value
    finally:
        wb.close()


@pytest.fixture
def journal(tmp_path, monkeypatch):
    folder = tmp_path / "journal"
    monkeypatch.setenv("KS4XL_SAVE_JOURNAL_DIR", str(folder))
    return folder


def _kill_mid_swap(tmp_path, book: Path, *, backup: bool) -> dict:
    """Run a real save of `book` in another process and kill it between the
    two moves. Returns what the folder holds afterwards."""
    mark = tmp_path / "between-the-moves"
    victim = tmp_path / "victim.py"
    victim.write_text(_VICTIM, encoding="utf-8")
    env = dict(os.environ, PYTHONPATH=os.pathsep.join(sys.path))
    proc = subprocess.Popen([sys.executable, "-X", "utf8", str(victim), str(book), str(mark),
                             "1" if backup else "0"], env=env, stdin=subprocess.DEVNULL,
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    try:
        deadline = time.monotonic() + 60
        while not mark.exists():
            assert proc.poll() is None, "the saving process ended before its swap"
            assert time.monotonic() < deadline, "the saving process never reached its swap"
            time.sleep(0.05)
    finally:
        proc.kill()  # TerminateProcess
        proc.wait(30)
    folder = book.parent
    aside = swapwait.asides(book)
    new = [p for p in folder.iterdir() if p.name.startswith(".ks4xl-write-")
           and p.suffix == ".xlsx"]
    return {"aside": aside, "new": new}


def _age(path, seconds):
    old = time.time() - seconds
    os.utime(path, (old, old))


def test_a_killed_save_leaves_the_workbook_aside_under_a_name_that_says_so(tmp_path, journal):
    book = tmp_path / "Budget.xlsx"
    _book(book, "the user's workbook")
    state = _kill_mid_swap(tmp_path, book, backup=True)
    assert not book.exists()
    assert len(state["aside"]) == 1 and len(state["new"]) == 1
    name = os.path.basename(state["aside"][0])
    assert name.startswith(swapwait.aside_prefix("Budget.xlsx"))
    assert name.endswith(" Budget.xlsx (as it was before a save).tmp")
    assert _a1(state["aside"][0]) == "the user's workbook"
    entries = list(journal.glob("*.json"))
    assert len(entries) == 1
    record = json.loads(entries[0].read_text("utf-8"))
    assert os.path.normcase(record["workbook"]) == os.path.normcase(str(book))


def test_a_read_names_where_the_workbook_is_and_moves_nothing(tmp_path, journal):
    book = tmp_path / "Budget.xlsx"
    _book(book, "the user's workbook")
    state = _kill_mid_swap(tmp_path, book, backup=True)
    with pytest.raises(WorkbookNotFound) as info:
        cells.read_range(str(book), "A1:A2")
    assert state["aside"][0] in str(info.value)
    assert "puts it back" in str(info.value)
    assert not book.exists() and os.path.exists(state["aside"][0])


@pytest.mark.parametrize("backup", [True, False], ids=["backup", "backup_false"])
def test_saves_of_other_workbooks_never_delete_either_version(tmp_path, journal, backup):
    book = tmp_path / "Budget.xlsx"
    _book(book, "the user's workbook")
    state = _kill_mid_swap(tmp_path, book, backup=backup)
    aside, new = state["aside"][0], state["new"][0]
    _age(aside, TWO_DAYS)      # a rename keeps the workbook's own old time
    _age(new, 2 * 3600)        # and the new version is two hours old
    other = tmp_path / "Other.xlsx"
    _book(other, "another workbook")
    for i in range(3):
        assert cells.set_cell(str(other), "B1", f"edit {i}", backup=backup)["ok"]
    assert os.path.exists(aside), "a save of another workbook deleted the set-aside original"
    assert os.path.exists(new), "a save of another workbook deleted the interrupted new version"


@pytest.mark.parametrize("backup", [True, False], ids=["backup", "backup_false"])
def test_the_next_write_puts_the_workbook_back_first(tmp_path, journal, backup):
    book = tmp_path / "Budget.xlsx"
    _book(book, "the user's workbook")
    state = _kill_mid_swap(tmp_path, book, backup=backup)
    assert safesave.put_back_before_write(str(book)) is True
    assert _a1(book) == "the user's workbook"
    assert not swapwait.asides(book)
    assert not os.path.exists(state["new"][0]), "the unanswered edit stays behind"
    assert not list(journal.glob("*.json"))
    r = cells.set_cell(str(book), "B1", "next edit", backup=backup)
    assert r["ok"] is True and _a1(book) == "the user's workbook"


def test_a_server_start_puts_back_every_interrupted_workbook(tmp_path, journal):
    book = tmp_path / "Budget.xlsx"
    _book(book, "the user's workbook")
    _kill_mid_swap(tmp_path, book, backup=False)
    from xlsx_mcp import server
    from xlsx_mcp.settings import Settings
    assert server._recover_at_start(Settings(read_only=True)) == []
    assert not book.exists(), "a read-only server moved a file"
    assert safesave.recover_interrupted_saves() == [str(book)]
    assert _a1(book) == "the user's workbook"


def test_restore_and_the_write_lock_put_it_back_too(tmp_path, journal):
    book = tmp_path / "Budget.xlsx"
    _book(book, "v1")
    cells.set_cell(str(book), "A1", "v2")
    _kill_mid_swap(tmp_path, book, backup=True)
    r = backups.manage_backups("restore", path=str(book), source="prev")
    assert r["restored"] and _a1(book) == "v1"


def test_create_and_copy_never_take_the_name_of_a_workbook_that_is_aside(tmp_path, journal):
    book = tmp_path / "Budget.xlsx"
    _book(book, "the user's workbook")
    _kill_mid_swap(tmp_path, book, backup=True)
    with pytest.raises(FileExistsError):
        lifecycle.create_workbook(str(book))
    assert _a1(book) == "the user's workbook"
    book2 = tmp_path / "Plan.xlsx"
    _book(book2, "plan")
    _kill_mid_swap(tmp_path, book2, backup=True)
    src = tmp_path / "src.xlsx"
    _book(src, "copied")
    with pytest.raises(FileExistsError):
        lifecycle.copy_workbook(str(src), str(book2))
    assert _a1(book2) == "plan"


def test_purging_orphans_keeps_the_backups_of_a_workbook_that_is_aside(tmp_path, journal):
    book = tmp_path / "Budget.xlsx"
    _book(book, "v1")
    cells.set_cell(str(book), "A1", "v2")
    _kill_mid_swap(tmp_path, book, backup=True)
    report = backups.purge_backups("orphans", directory=str(tmp_path), dry_run=False)
    assert report["count"] == 0, report
    assert (safesave.slot_dir(book) / "prev.xlsx").exists()


def test_a_live_servers_swap_is_never_touched(tmp_path, journal):
    book = tmp_path / "Budget.xlsx"
    _book(book, "the user's workbook")
    aside = tmp_path / swapwait.aside_name("Budget.xlsx", "abc")
    os.rename(book, aside)
    other_server = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(60)"],
                                    creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    try:
        journal.mkdir()
        (journal / "live.json").write_text(json.dumps({
            "workbook": str(book), "aside": str(aside), "new": str(tmp_path / "x.xlsx"),
            "staged": [], "pid": other_server.pid, "token": "someone-else",
            "pid_created": safesave._process_create_time(other_server.pid),
            "host": safesave._local_host(), "time": time.time()}), "utf-8")
        assert safesave.recover_interrupted_saves() == []
        assert (journal / "live.json").exists()
    finally:
        other_server.kill()
        other_server.wait()


def test_an_earlier_versions_aside_is_named_never_swept_and_put_back_when_proven(tmp_path, journal):
    folder = tmp_path / "work"
    folder.mkdir()
    book = folder / "Budget.xlsx"
    _book(book, "the user's workbook")
    _age(book, TWO_DAYS)
    # the state an earlier version's save left when killed between the two moves
    staged = safesave.slot_dir(book, create=True) / ".slot-0123456789abcdef.tmp"
    os.link(book, staged)
    legacy = folder / ".ks4xl-write-0123456789abcdef0123456789abcdef.replaced"
    os.rename(book, legacy)
    other = folder / "Other.xlsx"
    _book(other, "another workbook")
    cells.set_cell(str(other), "B1", "unrelated")
    assert legacy.exists(), "the sweep deleted an earlier version's aside"
    with pytest.raises(WorkbookNotFound) as info:
        cells.read_range(str(book), "A1")
    assert legacy.name in str(info.value)
    assert safesave.put_back_before_write(str(book)) is True
    assert _a1(book) == "the user's workbook" and not legacy.exists()
    assert not staged.exists()


def test_an_unproven_earlier_aside_is_only_named(tmp_path, journal):
    folder = tmp_path / "work"
    folder.mkdir()
    book = folder / "Budget.xlsx"
    _book(book, "the user's workbook")
    legacy = folder / ".ks4xl-write-0123456789abcdef0123456789abcdef.replaced"
    os.rename(book, legacy)
    assert safesave.put_back_before_write(str(book)) is False
    assert legacy.exists() and not book.exists()


def test_the_aside_name_fits_for_a_long_korean_name():
    name = "가" * 200 + ".xlsx"
    aside = swapwait.aside_name(name, "0123456789ab")
    assert len(aside.encode("utf-16-le")) // 2 <= 255
    assert aside.startswith(swapwait.aside_prefix(name))
