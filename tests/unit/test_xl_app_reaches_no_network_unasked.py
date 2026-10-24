"""LIVE: no xl_app operation lets the hidden Excel reach the network unasked.

A made-up workbook's WEBSERVICE formula points at a listener on this machine
(127.0.0.1, never a real network address) and puts a cell's text into the
address. Excel recalculates that workbook as it opens (openpyxl's files ask
for it), so any route that opened it would send the request. Every operation
that opens a workbook in Excel is called without allow_network: each must
refuse and the listener must hear nothing. The same workbook as an .xlsb
(which this server cannot read, so Excel checks it with calculation off) and
as a CSV are refused too. One call with allow_network=true is the control:
the listener hears it.

The tests marked LIVE start the product's hidden Excel: run them on a
private desktop.
"""
from __future__ import annotations

import http.server
import re
import threading
import time

import openpyxl
import pytest

from xlsx_mcp.com import session as com_session
from xlsx_mcp.ops import comtier

LIVE = pytest.mark.skipif(not com_session.com_available()[0], reason="needs Excel (COM)")


@pytest.fixture(scope="module")
def listener():
    hits: list[str] = []

    class H(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            hits.append(self.path)
            self.send_response(200)
            self.send_header("Content-Length", "2")
            self.end_headers()
            self.wfile.write(b"42")

        def log_message(self, *a):
            pass

    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield srv.server_address[1], hits
    srv.shutdown()


def _book(p, port, tag):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws["A1"] = "made-up-" + tag
    ws["A2"] = 3
    ws["B1"] = f'=_xlfn.WEBSERVICE("http://127.0.0.1:{port}/{tag}?v="&A1)'
    ws["B2"] = "=A2*2"
    wb.save(p)
    return str(p)


def _calls(tmp_path):
    out = tmp_path / "out"
    out.mkdir(exist_ok=True)
    return {
        "com_validate_opens_clean": lambda p: comtier.com_validate_opens_clean(p),
        "com_export_pdf": lambda p: comtier.com_export_pdf(p, str(out / "a.pdf")),
        "com_render_sheet": lambda p: comtier.com_render_sheet(p, str(out / "a.png")),
        "com_convert_format_csv": lambda p: comtier.com_convert_format(p, str(out / "a.csv")),
        "com_convert_format_xlsb": lambda p: comtier.com_convert_format(p, str(out / "b.xlsb")),
        "com_autofit": lambda p: comtier.com_autofit(p),
        "com_save_with_password": lambda p: comtier.com_save_with_password(p, "made-up-pw"),
        "com_manage_pivot_list": lambda p: comtier.com_manage_pivot(p, "list"),
        "com_manage_pivot_create": lambda p: comtier.com_manage_pivot(
            p, "create", source_range="A1:B2", source_sheet="Rates", rows=["made-up-x"],
            values=[{"field": "made-up-y", "func": "sum"}]),
        "com_set_sparkline_list": lambda p: comtier.com_set_sparkline(p, "list"),
        "com_set_sparkline_create": lambda p: comtier.com_set_sparkline(
            p, "create", location="D2", source="A2:B2", sheet="Rates"),
        "recalculate": lambda p: comtier.recalculate(p, engine="com"),
        "com_goal_seek": lambda p: comtier.com_goal_seek(p, "B2", 10, "A2", sheet="Rates"),
        "verify_com_opens_clean": lambda p: com_session.opens_clean(p),
    }


OPS = ["com_validate_opens_clean", "com_export_pdf", "com_render_sheet", "com_convert_format_csv",
       "com_convert_format_xlsb", "com_autofit", "com_save_with_password", "com_manage_pivot_list",
       "com_manage_pivot_create", "com_set_sparkline_list", "com_set_sparkline_create", "recalculate",
       "com_goal_seek", "verify_com_opens_clean"]


def _quiet(hits, before, wait=1.5):
    time.sleep(wait)
    return hits[before:]


@LIVE
@pytest.mark.parametrize("op", OPS)
def test_no_xl_app_op_reaches_the_network_without_allow_network(tmp_path, listener, op):
    port, hits = listener
    p = _book(tmp_path / "rates.xlsx", port, op)
    before = len(hits)
    with pytest.raises(Exception) as ei:
        _calls(tmp_path)[op](p)
    assert "allow_network=true" in str(ei.value), str(ei.value)
    assert _quiet(hits, before) == [], f"{op} let Excel fetch without allow_network"


@LIVE
def test_a_binary_copy_is_checked_by_excel_and_refused_without_a_request(tmp_path, listener):
    port, hits = listener
    p = _book(tmp_path / "rates.xlsx", port, "binary")
    xlsb = tmp_path / "rates.xlsb"
    comtier.com_convert_format(p, str(xlsb), allow_network=True)  # making the copy may fetch: allowed
    _quiet(hits, 0)  # let any request of the copy's own making land first
    before = len(hits)
    for call in (lambda: comtier.com_export_pdf(str(xlsb), str(tmp_path / "b.pdf")),
                 lambda: comtier.com_convert_format(str(xlsb), str(tmp_path / "b.xlsx"))):
        with pytest.raises(Exception) as ei:
            call()
        assert "calculation turned off" in str(ei.value), str(ei.value)
    assert _quiet(hits, before) == []


@LIVE
def test_a_csv_with_a_fetching_formula_is_refused_without_a_request(tmp_path, listener):
    port, hits = listener
    p = tmp_path / "rates.csv"
    p.write_text(f'name,value\nmade-up,"=WEBSERVICE(""http://127.0.0.1:{port}/csv"")"\n', encoding="utf-8")
    before = len(hits)
    with pytest.raises(Exception) as ei:
        comtier.com_convert_format(str(p), str(tmp_path / "rates.xlsx"))
    assert "allow_network=true" in str(ei.value)
    assert _quiet(hits, before) == []


@LIVE
def test_allow_network_is_the_one_way_through(tmp_path, listener):
    port, hits = listener
    p = _book(tmp_path / "rates.xlsx", port, "control")
    before = len(hits)
    comtier.com_validate_opens_clean(p, allow_network=True)
    assert any("/control" in h for h in _quiet(hits, before))


@LIVE
def test_a_spreadsheet_excel_saved_as_opendocument_is_refused_without_a_request(tmp_path, listener):
    # Excel writes the function into an .ods as COM.MICROSOFT.WEBSERVICE( and fetches from it when it opens the file
    port, hits = listener
    p = _book(tmp_path / "rates.xlsx", port, "ods")
    ods = tmp_path / "rates.ods"
    comtier.com_convert_format(p, str(ods), allow_network=True)  # making the copy may fetch: allowed
    _quiet(hits, 0)  # let any request of the copy's own making land first
    import zipfile
    with zipfile.ZipFile(ods) as z:
        assert "WEBSERVICE" in z.read("content.xml").decode("utf-8")
    before = len(hits)
    for call in (lambda: comtier.com_validate_opens_clean(str(ods)),
                 lambda: comtier.com_export_pdf(str(ods), str(tmp_path / "o.pdf")),
                 lambda: comtier.com_convert_format(str(ods), str(tmp_path / "o.xlsx"))):
        with pytest.raises(Exception) as ei:
            call()
        assert "allow_network=true" in str(ei.value), str(ei.value)
    assert _quiet(hits, before) == []
    comtier.com_validate_opens_clean(str(ods), allow_network=True)  # the control: asked, so it fetches
    assert any("/ods" in h for h in _quiet(hits, before, 4))


@LIVE
@pytest.mark.parametrize("how", ["decimal reference", "comment inside the name", "cdata section", "utf-16"])
def test_a_function_name_written_another_way_in_the_markup_is_refused_without_a_request(tmp_path, listener, how):
    # Excel reads these as WEBSERVICE and fetches (measured); the check must read them the same way
    import zipfile
    port, hits = listener
    tag = "markup-" + how.split()[0]  # a name of its own: Excel keeps what a web address answered
    src = _book(tmp_path / "rates.xlsx", port, tag)
    spell = {
        "decimal reference": lambda t: t.replace("_xlfn.WEBSERVICE(", "_xlfn.WEB&#83;ERVICE("),
        "comment inside the name": lambda t: t.replace("_xlfn.WEBSERVICE(", "_xlfn.WEB<!-- x -->SERVICE("),
        "cdata section": lambda t: re.sub(
            r"<f>(_xlfn\.WEBSERVICE\(.*?)</f>",
            lambda m: "<f><![CDATA[" + m.group(1).replace("&amp;", "&") + "]]></f>", t),
        "utf-16": lambda t: t.replace("UTF-8", "UTF-16").encode("utf-16-le"),
    }[how]
    dst = str(tmp_path / "other.xlsx")
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                out = spell(data.decode("utf-8"))
                data = out if isinstance(out, bytes) else out.encode("utf-8")
            zo.writestr(item, data)
    before = len(hits)
    with pytest.raises(Exception) as ei:
        comtier.com_validate_opens_clean(dst)
    assert "allow_network=true" in str(ei.value), str(ei.value)
    assert _quiet(hits, before) == []
    comtier.com_validate_opens_clean(dst, allow_network=True)  # the control: Excel does fetch from this spelling
    assert any("/" + tag in h for h in _quiet(hits, before, 4))


@LIVE
@pytest.mark.parametrize("how", ["cdata cells", "processing instructions", "csv cells", "opendocument instructions"])
def test_markers_that_look_like_a_comment_hide_no_call_from_the_check(tmp_path, listener, how):
    # text that only looks like the two ends of a comment, on either side of the call: Excel still fetches (measured)
    import zipfile
    port, hits = listener
    tag = "marks-" + how.split()[0]  # a name of its own: Excel keeps what a web address answered
    url = f"http://127.0.0.1:{port}/{tag}"
    dst = str(tmp_path / ("other.csv" if how == "csv cells" else "other.ods" if how.startswith("opendoc")
                          else "other.xlsx"))
    if how == "csv cells":
        (tmp_path / "other.csv").write_text(f'<!--,"=WEBSERVICE(""{url}"")",-->\n', encoding="utf-8")
    elif how.startswith("opendoc"):
        src = _book(tmp_path / "rates.xlsx", port, tag)
        comtier.com_convert_format(src, str(tmp_path / "base.ods"), allow_network=True)  # may fetch: allowed
        _quiet(hits, 0)
        with zipfile.ZipFile(tmp_path / "base.ods") as zi, zipfile.ZipFile(dst, "w") as zo:
            for item in zi.infolist():
                data = zi.read(item.filename)
                if item.filename == "content.xml":
                    t = data.decode("utf-8")
                    t = t.replace("<office:spreadsheet>", "<office:spreadsheet><?x <!--?>", 1)
                    t = t.replace("</office:spreadsheet>", "<?x -->?></office:spreadsheet>", 1)
                    data = t.encode("utf-8")
                zo.writestr(item, data)
    else:
        src = _book(tmp_path / "rates.xlsx", port, tag)
        with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
            for item in zi.infolist():
                data = zi.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    t = data.decode("utf-8")
                    if how == "cdata cells":
                        # the markers in A1 and C1, around the call in B1: cells in column order, or Excel
                        # refuses the file and the control below would never fetch
                        left = '<c r="A1" t="inlineStr"><is><t><![CDATA[<!--]]></t></is></c>'
                        right = '<c r="C1" t="inlineStr"><is><t><![CDATA[-->]]></t></is></c>'
                        t = re.sub(r"<c r=\"A1\".*?</c>", lambda m: left, t, count=1, flags=re.S)
                        t = re.sub(r"(<c r=\"B1\".*?</c>)", lambda m: m.group(1) + right, t, count=1, flags=re.S)
                    else:
                        t = t.replace("<row ", "<?x <!--?><row ", 1).replace("</sheetData>", "<?x -->?></sheetData>", 1)
                    data = t.encode("utf-8")
                zo.writestr(item, data)
    before = len(hits)
    with pytest.raises(Exception) as ei:
        comtier.com_validate_opens_clean(dst)
    assert "allow_network=true" in str(ei.value), str(ei.value)
    assert _quiet(hits, before) == []
    comtier.com_validate_opens_clean(dst, allow_network=True)  # the control: asked, so Excel fetches
    assert any("/" + tag in h for h in _quiet(hits, before, 4))


def _rewrite_package(src, dst, fn):
    """Copy the package at src to dst, each member through fn(name, data) -> [(name, data), ...]."""
    import zipfile
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            for name, data in fn(item.filename, zi.read(item.filename)):
                zo.writestr(name, data)
    return dst


def _wrong_crc(path, member):
    import struct
    import zipfile
    data = bytearray(open(path, "rb").read())
    with zipfile.ZipFile(path) as z:
        info = z.getinfo(member)
    bad = struct.pack("<I", info.CRC ^ 0xFFFFFFFF)
    data[info.header_offset + 14:info.header_offset + 18] = bad
    i = data.find(b"PK\x01\x02")
    while i >= 0:
        n = struct.unpack("<H", data[i + 28:i + 30])[0]
        if bytes(data[i + 46:i + 46 + n]) == member.encode():
            data[i + 16:i + 20] = bad
            break
        i = data.find(b"PK\x01\x02", i + 4)
    open(path, "wb").write(bytes(data))


_OFFICE_DOC = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"


def _no_declaration(t):
    return re.sub(r"^\s*<\?xml[^>]*\?>\s*", "", t)


def _doubtful(src, dst, how):
    """The made-up workbook at src, rewritten to hide its call from a reading of the text alone (measured: Excel 16
    opens each one and fetches)."""
    sheet = "xl/worksheets/sheet1.xml"
    if how == "instruction named xml-x":
        return _rewrite_package(src, dst, lambda n, d: [(n, _no_declaration(d.decode()).replace(
            "_xlfn.WEBSERVICE(", "_xlfn.WEB<?xml-x?>SERVICE(").encode()) if n == sheet else (n, d)])
    if how == "declared utf-7":
        return _rewrite_package(src, dst, lambda n, d: [(n, ('<?xml version="1.0" encoding="UTF-7"?>' + _no_declaration(
            d.decode()).replace("_xlfn.WEBSERVICE(", "_xlfn.+AFc-EBSERVICE(")).encode("ascii")) if n == sheet
            else (n, d)])
    if how == "commented-out workbook relationship":
        decoy = f'<!--<Relationship Id="rIdD" Type="{_OFFICE_DOC}" Target="xl/decoy.xml"/>-->'

        def fn(n, d):
            if n == "_rels/.rels":
                return [(n, re.sub(r"(<Relationships\b[^>]*>)", lambda m: m.group(1) + decoy, d.decode(),
                                   count=1).encode())]
            if n == "xl/workbook.xml":
                return [(n, d), ("xl/decoy.xml", d)]
            return [(n, d)]
        return _rewrite_package(src, dst, fn)
    if how == "workbook part in another case":
        def fn(n, d):
            if n == "_rels/.rels":
                return [(n, d.replace(b"xl/workbook.xml", b"xl/book.xml"))]
            if n == "[Content_Types].xml":
                return [(n, d.replace(b"/xl/workbook.xml", b"/xl/book.xml"))]
            if n == "xl/workbook.xml":
                return [("xl/Book.xml", d), (n, d)]  # the real part, and a decoy with no relationships
            if n == "xl/_rels/workbook.xml.rels":
                return [("xl/_rels/Book.xml.rels", d)]
            return [(n, d)]
        return _rewrite_package(src, dst, fn)
    if how == "wrong checksum":
        _rewrite_package(src, dst, lambda n, d: [(n, d)])
        _wrong_crc(dst, sheet)
        return dst
    if how in ("sheet part not named xml, led by whitespace", "sheet part given a picture type"):
        new = "xl/worksheets/sheet1.dat"

        def fn(n, d):
            if n == sheet:
                return [(new, (" " * 5000 + _no_declaration(d.decode()).replace(
                    "_xlfn.WEBSERVICE(", "_xlfn.WEB&#83;ERVICE(")).encode())]
            if n == "xl/_rels/workbook.xml.rels":
                return [(n, d.replace(b"worksheets/sheet1.xml", b"worksheets/sheet1.dat"))]
            if n == "[Content_Types].xml":
                d = d.replace(b"/" + sheet.encode(), b"/" + new.encode())
                if how == "sheet part given a picture type":
                    d = re.sub(rb'(<Override PartName="/xl/worksheets/sheet1.dat" ContentType=")[^"]*"',
                               lambda m: m.group(1) + b'image/png"', d)
                return [(n, d)]
            return [(n, d)]
        return _rewrite_package(src, dst, fn)
    raise ValueError(how)


@LIVE
@pytest.mark.parametrize("how", ["instruction named xml-x", "declared utf-7", "commented-out workbook relationship",
                                 "workbook part in another case", "wrong checksum",
                                 "sheet part not named xml, led by whitespace", "sheet part given a picture type"])
def test_a_workbook_built_to_hide_a_call_is_refused_without_a_request(tmp_path, listener, how):
    # Excel 16 opens each of these and fetches (measured), so each must be refused unless asked
    port, hits = listener
    tag = "doubt-" + re.sub(r"\W+", "-", how)  # a name of its own: Excel keeps what a web address answered
    src = _book(tmp_path / "rates.xlsx", port, tag)
    dst = _doubtful(src, str(tmp_path / "other.xlsx"), how)
    before = len(hits)
    for call in (lambda: comtier.com_validate_opens_clean(dst),
                 lambda: comtier.com_export_pdf(dst, str(tmp_path / "d.pdf"))):
        with pytest.raises(Exception) as ei:
            call()
        assert "allow_network=true" in str(ei.value), str(ei.value)
    assert _quiet(hits, before) == []
    comtier.com_validate_opens_clean(dst, allow_network=True)  # the control: asked, so Excel fetches
    assert any("/" + tag in h for h in _quiet(hits, before, 4))
