"""What sits at a position on a sheet moves with its cells when rows or
columns are inserted or deleted, on a made-up model: a what-if data table
(its range and its input cell), a chart anchored over cells, manual page
breaks, a table whose calculated column and custom totals formula hold A1
references, a what-if scenario's input cell, and a pivot table.

openpyxl moves the cells only, so before this each of those kept its old
address: the data table pointed at the wrong input, the chart sat over other
cells, the page broke at the wrong row, the next row Excel added to the table
took the old formula, and the pivot's definition no longer matched its
cells. Each expected position is what Excel itself stored after the same
edit, made by a private hidden Excel on the same model (Rows.Insert,
Columns.Insert, Rows.Delete). Excel refuses an insert that cuts a pivot
table or a data table in two; so do the tools.
"""

from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.worksheet.formula import DataTableFormula
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.scenario import InputCells, Scenario, ScenarioList
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

import family_harness as fh

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "preserve"


def _call(tool, op, **args):
    return fh.start().call(tool, {"op": op, "args": args})


def _ok(tool, op, **args):
    payload, err, _ = _call(tool, op, **args)
    assert not err, (op, payload)
    return payload


def _model(path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["B3"], ws["C3"] = "Rate", 0.05
    ws["B4"], ws["C4"] = "Base", 1000
    ws["B5"], ws["C5"] = "Result", "=C4*(1+C3)"
    ws["F6"] = "=C5"
    for i, v in enumerate([0.01, 0.02, 0.03, 0.04]):
        ws.cell(7 + i, 5, v)
    ws["F7"] = DataTableFormula(ref="F7:F10", r1="C3")
    for r in range(8, 11):
        ws.cell(r, 6, 1000 + r)
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=5, min_row=7, max_row=10))
    chart.anchor = TwoCellAnchor(_from=AnchorMarker(col=7, row=11),
                                 to=AnchorMarker(col=13, row=22))
    ws.add_chart(chart)
    ws.row_breaks.append(Break(id=33))
    ws.col_breaks.append(Break(id=14))
    for j, h in enumerate(["Q1", "Q2", "Calc"]):
        ws.cell(40, 2 + j, h)
    for r in range(41, 45):
        ws.cell(r, 2, r - 40)
        ws.cell(r, 3, (r - 40) * 2)
        ws.cell(r, 4, "=Lines[[#This Row],[Q1]]*$C$4")
    ws["B45"] = "Total"
    ws["C45"] = "=SUM(Lines[Q2])+$C$3"
    ws["D45"] = "=SUBTOTAL(109,Lines[Calc])"
    t = Table(displayName="Lines", ref="B40:D45", totalsRowCount=1)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    t.tableColumns = [
        TableColumn(id=1, name="Q1", totalsRowLabel="Total"),
        TableColumn(id=2, name="Q2", totalsRowFunction="custom",
                    totalsRowFormula=TableFormula(attr_text="SUM(Lines[Q2])+$C$3")),
        TableColumn(id=3, name="Calc", totalsRowFunction="sum",
                    calculatedColumnFormula=TableFormula(attr_text="Lines[[#This Row],[Q1]]*$C$4")),
    ]
    ws.add_table(t)
    ws.scenarios = ScenarioList(scenario=[Scenario(
        name="High", user="Analyst", inputCells=[InputCells(r="C3", val="0.1")])])
    wb.save(path)
    return str(path)


def _positions(path) -> dict:
    wb = openpyxl.load_workbook(path)
    ws = wb["Model"]
    dts = {c.coordinate: (c.value.ref, c.value.r1) for row in ws.iter_rows() for c in row
           if isinstance(c.value, DataTableFormula)}
    a = ws._charts[0].anchor
    t = ws.tables["Lines"]
    return {
        "data_table": dts,
        "chart": ((a._from.col, a._from.row), (a.to.col, a.to.row)),
        "row_breaks": [b.id for b in ws.row_breaks.brk],
        "col_breaks": [b.id for b in ws.col_breaks.brk],
        "calc": t.tableColumns[2].calculatedColumnFormula.attr_text,
        "totals": t.tableColumns[1].totalsRowFormula.attr_text,
        "scenario": [ic.r for ic in ws.scenarios.scenario[0].inputCells],
        "table": t.ref,
    }


BASE = {
    "data_table": {"F7": ("F7:F10", "C3")},
    "chart": ((7, 11), (13, 22)),
    "row_breaks": [33], "col_breaks": [14],
    "calc": "Lines[[#This Row],[Q1]]*$C$4", "totals": "SUM(Lines[Q2])+$C$3",
    "scenario": ["C3"], "table": "B40:D45",
}


def test_the_model_reads_back_as_built(tmp_path):
    assert _positions(_model(tmp_path / "m.xlsx")) == BASE


def test_two_rows_inserted_above_everything_move_everything_down(tmp_path):
    p = _model(tmp_path / "m.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_rows", at=1, count=2,
        sheet="Model", allow_loss=True)
    assert _positions(p) == {
        "data_table": {"F9": ("F9:F12", "C5")},
        "chart": ((7, 13), (13, 24)),
        "row_breaks": [35], "col_breaks": [14],
        "calc": "Lines[[#This Row],[Q1]]*$C$6", "totals": "SUM(Lines[Q2])+$C$5",
        "scenario": ["C5"], "table": "B42:D47",
    }


def test_a_column_inserted_before_everything_moves_everything_right(tmp_path):
    p = _model(tmp_path / "m.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_cols", at="A",
        sheet="Model", allow_loss=True)
    assert _positions(p) == {
        "data_table": {"G7": ("G7:G10", "D3")},
        "chart": ((8, 11), (14, 22)),
        "row_breaks": [33], "col_breaks": [15],
        "calc": "Lines[[#This Row],[Q1]]*$D$4", "totals": "SUM(Lines[Q2])+$D$3",
        "scenario": ["D3"], "table": "C40:E45",
    }


def test_rows_inserted_through_a_chart_stretch_it(tmp_path):
    p = _model(tmp_path / "m.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_rows", at=15, count=2,
        sheet="Model", allow_loss=True)
    pos = _positions(p)
    assert pos["chart"] == ((7, 11), (13, 24))      # move and size: the bottom moves
    assert pos["data_table"] == {"F7": ("F7:F10", "C3")}
    assert pos["row_breaks"] == [35]


def test_deleting_the_row_that_starts_a_page_deletes_its_break(tmp_path):
    p = _model(tmp_path / "m.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="delete_rows", at=34,
        sheet="Model", allow_loss=True)
    assert _positions(p)["row_breaks"] == []


@pytest.mark.parametrize("at", [8, 10])
def test_an_insert_that_cuts_a_data_table_is_refused(tmp_path, at):
    p = _model(tmp_path / "m.xlsx")
    before = Path(p).read_bytes()
    payload, err, _ = _call("xl_structure", "modify_grid_structure", path=p, action="insert_rows",
                            at=at, sheet="Model", allow_loss=True)
    assert err and "data table" in str(payload), payload
    assert Path(p).read_bytes() == before


def test_a_pivot_table_moves_with_an_insert_above_it_and_refuses_a_cut(tmp_path):
    p = str(tmp_path / "pivot.xlsx")
    shutil.copy2(FIXTURES / "pivot.xlsx", p)
    ws = openpyxl.load_workbook(p)["Pivot"]
    assert ws._pivots[0].location.ref == "A3:B8"
    payload, err, _ = _call("xl_structure", "modify_grid_structure", path=p, action="insert_rows",
                            at=5, sheet="Pivot", allow_loss=True)
    assert err and "pivot table" in str(payload), payload
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_rows", at=1, count=2,
        sheet="Pivot", allow_loss=True)
    assert openpyxl.load_workbook(p)["Pivot"]._pivots[0].location.ref == "A5:B10"


def test_renames_reach_the_formulas_a_table_keeps_for_its_next_row(tmp_path):
    """A sheet rename and a table rename rewrite the table's own copy of its
    calculated-column formula too, as Excel does, so the next row Excel adds
    gets a formula that still works."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    rates = wb.create_sheet("Rates")
    rates["A1"] = 2
    ws.append(["v", "w"])
    for r in (2, 3):
        ws.cell(r, 1, r)
        ws.cell(r, 2, "=T[[#This Row],[v]]*Rates!$A$1")
    t = Table(displayName="T", ref="A1:B3")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    t.tableColumns = [TableColumn(id=1, name="v"), TableColumn(
        id=2, name="w", calculatedColumnFormula=TableFormula(attr_text="T[[#This Row],[v]]*Rates!$A$1"))]
    ws.add_table(t)
    p = str(tmp_path / "t.xlsx")
    wb.save(p)
    _ok("xl_structure", "manage_worksheet", path=p, action="rename", sheet="Rates", new_name="Fx 2")
    _ok("xl_structure", "manage_table", path=p, name="T", action="rename", new_name="Prices")
    wb = openpyxl.load_workbook(p)
    tc = wb["Data"].tables["Prices"].tableColumns[1]
    assert tc.calculatedColumnFormula.attr_text == "Prices[[#This Row],[v]]*'Fx 2'!$A$1"
    assert wb["Data"]["B2"].value == "=Prices[[#This Row],[v]]*'Fx 2'!$A$1"
