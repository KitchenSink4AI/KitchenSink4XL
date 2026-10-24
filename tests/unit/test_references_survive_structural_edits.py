"""Structural edits on made-up workbooks whose references a plain scan for A1
text misreads: a table whose columns are named like cells (Q1, FY24,
FY2024, A1), a link to another workbook that has a sheet named like a local
one, sheets named in other scripts, a sheet name holding double quotes, and
defined names whose insides look like cells.

Each expected formula is the text Excel itself stores after the same edit:
the edit was made once through the tools and once by a private hidden Excel
(Rows.Insert, Columns.Insert, ListRows.Add, ListColumns.Delete,
Worksheet.Name, Worksheet.Delete, Range.Copy) on the same workbook, and the
two files held the same formulas, names and rule formulas. The COM test at
the end makes some of the same edits in Excel again, where Excel exists.
"""

from __future__ import annotations

import shutil

import openpyxl
import pytest
from openpyxl.formatting.rule import FormulaRule
from openpyxl.packaging.relationship import Relationship
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.external_link.external import (
    ExternalBook, ExternalCell, ExternalLink, ExternalRow, ExternalSheetData,
    ExternalSheetDataSet, ExternalSheetNames)
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

import family_harness as fh
from xlsx_mcp import server as srv
from xlsx_mcp.com import session as com_session

_COM = com_session.com_available()[0]
FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


def _ok(tool, op, **args):
    payload, err, _ = fh.start().call(tool, {"op": op, "args": args})
    assert not err, (op, payload)
    return payload


def _name(wb, name, text):
    wb.defined_names[name] = DefinedName(name, attr_text=text)


def _formulas(path) -> dict:
    wb = openpyxl.load_workbook(path)
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    out[f"{ws.title}!{c.coordinate}"] = c.value
    return out


def _names(path) -> dict:
    return {k: v.attr_text for k, v in openpyxl.load_workbook(path).defined_names.items()}


# ------------------------------------------------------------------ workbooks

def _sales(path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Region", "Q1", "Q2", "FY24", "FY2024", "A1", "x#A1", "Total"])
    for r in (["North", 10, 20, 30, 40, 1, 2], ["South", 11, 21, 31, 41, 3, 4],
              ["East", 12, 22, 32, 42, 5, 6]):
        ws.append(r + ["=Sales[[#This Row],[Q1]]+Sales[[#This Row],[Q2]]"])
    table = Table(displayName="Sales", ref="A1:H4")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws["J2"] = 5
    ws["J3"] = "=SUM(Sales[Q1])+J2"
    s = wb.create_sheet("Summ")
    for i, f in enumerate(["=SUM(Sales[Q1])", "=SUM(Sales[FY24])", "=SUM(Sales[FY2024])",
                           "=SUM(Sales[[Q1]:[Q2]])", "=Sales[[#Headers],[Q1]]",
                           "=SUM(Sales[x'#A1])", "=SUM(Sales[A1])", "=ROWS(Sales[#All])",
                           "=SUM(Sales[Q1])+Data!B2", "=Q1Sum+FYName",
                           "=SUM(Sales[[#Data],[FY24]])", "=SUM(Sales[Q2])*A20",
                           "=Data!J2+SUM(Sales[Q1])"], start=1):
        s.cell(i, 1, f)
    s["A20"] = 2
    _name(wb, "Q1Sum", "SUM(Sales[Q1])")
    _name(wb, "FYName", "SUM(Sales[FY2024])")
    s.conditional_formatting.add("B1:B3", FormulaRule(formula=["Q1Sum>0"], fill=FILL))
    wb.save(path)
    return str(path)


SALES_UNCHANGED = {
    "Summ!A1": "=SUM(Sales[Q1])", "Summ!A2": "=SUM(Sales[FY24])",
    "Summ!A3": "=SUM(Sales[FY2024])", "Summ!A4": "=SUM(Sales[[Q1]:[Q2]])",
    "Summ!A5": "=Sales[[#Headers],[Q1]]", "Summ!A6": "=SUM(Sales[x'#A1])",
    "Summ!A7": "=SUM(Sales[A1])", "Summ!A8": "=ROWS(Sales[#All])",
    "Summ!A10": "=Q1Sum+FYName", "Summ!A11": "=SUM(Sales[[#Data],[FY24]])",
}
SALES_NAMES = {"FYName": "SUM(Sales[FY2024])", "Q1Sum": "SUM(Sales[Q1])"}
THIS_ROW = "=Sales[[#This Row],[Q1]]+Sales[[#This Row],[Q2]]"


def _linked(path) -> str:
    """A local sheet Data, and a link to another workbook that also has a
    sheet Data (Excel stores its references as [1]Data!A1)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, 6):
        ws.cell(r, 1, r)
        ws.cell(r, 2, r * 2)
    s = wb.create_sheet("Summ")
    for i, f in enumerate(["=[1]Data!A1", "=SUM([1]Data!A1:B3)+Data!A1", "=[1]Data!$A$2",
                           "='[1]My Data'!A1", "=[1]!ExtTotal+Data!B2", "=[1]Data!A3*Data!A3",
                           "=Data!A4"], start=1):
        s.cell(i, 1, f)
    s["A10"] = "=ExtRate*2"
    _name(wb, "ExtRate", "[1]Data!$B$1")
    rows = [ExternalRow(r=r, cell=[ExternalCell(r=f"A{r}", v=str(r * 7)),
                                   ExternalCell(r=f"B{r}", v=str(r * 11))]) for r in range(1, 6)]
    book = ExternalBook(sheetNames=ExternalSheetNames(sheetName=["Data", "My Data"]),
                        sheetDataSet=ExternalSheetDataSet(sheetData=[
                            ExternalSheetData(sheetId=0, row=rows)]))
    book.id = "rId1"
    link = ExternalLink(externalBook=book)
    link.file_link = Relationship(type="externalLinkPath", Target="other.xlsx",
                                  TargetMode="External")
    wb._external_links.append(link)
    wb.save(path)
    return str(path)


def _scripts(path) -> str:
    """Sheets whose names Excel writes without quotes in other scripts, and a
    sheet named es (the tail of Données)."""
    wb = openpyxl.Workbook()
    summ = wb.active
    summ.title = "Summ"
    for nm in ("Données", "売上", "ยอดขาย", "Cost€", "es", "Q１"):
        ws = wb.create_sheet(nm)
        for r in range(1, 8):
            ws.cell(r, 1, r * 10 + len(nm))
            ws.cell(r, 2, r * 100 + len(nm))
    for i, f in enumerate(["=Données!A5*2", "=SUM(売上!B2:B4)", "=ยอดขาย!A2+ยอดขาย!B3",
                           "=Cost€!A1", "=es!A5", "=Données!A5+es!A5", "=SUM(Données!A:A)",
                           "=SUM(売上!2:3)", "=Q１!A2", "=DonnéesTot+1",
                           "=COUNT(Données!A1:B7)"], start=1):
        summ.cell(i, 1, f)
    _name(wb, "DonnéesTot", "Données!$A$5")
    _name(wb, "売上Range", "売上!$B$2:$B$4")
    summ["C1"] = 5
    summ.conditional_formatting.add("C1", FormulaRule(formula=["Données!$A$5>0"], fill=FILL))
    dv = DataValidation(type="list", formula1="売上!$B$2:$B$4", allow_blank=True)
    summ.add_data_validation(dv)
    dv.add("D1")
    wb.save(path)
    return str(path)


def _quoted(path) -> str:
    """A sheet name holding double quotes, strings beside it holding the same
    characters, and names whose insides look like cells."""
    wb = openpyxl.Workbook()
    summ = wb.active
    summ.title = "Summ"
    nq, other = wb.create_sheet('North "Q"'), wb.create_sheet("North Q")
    for r in range(1, 6):
        nq.cell(r, 1, r)
        other.cell(r, 1, r * 1000)
    for i, f in enumerate(["='North \"Q\"'!A1", "=SUM('North \"Q\"'!A1:A3)",
                           "=\"North \"\"Q\"\"!A1\"&'North \"Q\"'!A2",
                           "='North Q'!A2+'North \"Q\"'!A3",
                           "=LEN(\"'North \"\"Q\"\"'!A1\")+'North \"Q\"'!A4"], start=1):
        summ.cell(i, 1, f)
    for i, nm in enumerate(["αA1", "A1α", "売上A1", "ÀB12", "x.A1", "Q1β"], start=1):
        summ.cell(i, 3, 100 + i)
        _name(wb, nm, f"Summ!$C${i}")
    for i, f in enumerate(["=αA1+A1", "=A1α*2+B2", "=売上A1+A1", "=ÀB12+B12", "=x.A1+A1",
                           "=Q1β+A1"], start=1):
        summ.cell(i, 5, f)
    summ["A12"], summ["B2"], summ["B12"] = 7, 3, 9
    wb.save(path)
    return str(path)


# ------------------------------------------------------------------ tables

def test_a_row_inserted_above_a_table_keeps_its_column_names(tmp_path):
    p = _sales(tmp_path / "s.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_rows", at=1,
        sheet="Data")
    got = _formulas(p)
    for k, v in SALES_UNCHANGED.items():
        assert got[k] == v, k
    assert [got[f"Data!H{r}"] for r in (3, 4, 5)] == [THIS_ROW] * 3
    assert got["Data!J4"] == "=SUM(Sales[Q1])+J3"
    assert got["Summ!A9"] == "=SUM(Sales[Q1])+Data!B3"
    assert got["Summ!A13"] == "=Data!J3+SUM(Sales[Q1])"
    assert _names(p) == SALES_NAMES


def test_a_column_inserted_beside_the_formulas_keeps_the_table_references(tmp_path):
    p = _sales(tmp_path / "s.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_cols", at="A",
        sheet="Summ")
    got = _formulas(p)
    for k, v in SALES_UNCHANGED.items():
        assert got["Summ!B" + k.split("!A")[1]] == v, k
    assert got["Summ!B12"] == "=SUM(Sales[Q2])*B20"
    assert [got[f"Data!H{r}"] for r in (2, 3, 4)] == [THIS_ROW] * 3
    assert _names(p) == SALES_NAMES


def test_a_table_row_added_leaves_the_column_names_alone(tmp_path):
    # The new row goes in at row 5, which once moved FY2024 (a column name)
    # to FY2025 and broke the name built on it.
    p = _sales(tmp_path / "s.xlsx")
    _ok("xl_structure", "manage_table", path=p, name="Sales", action="add_row",
        values=[["West", 13, 23, 33, 43, 7, 8]])
    got = _formulas(p)
    for k, v in SALES_UNCHANGED.items():
        assert got[k] == v, k
    assert _names(p) == SALES_NAMES


def test_a_table_column_deleted_leaves_the_other_column_names_alone(tmp_path):
    p = _sales(tmp_path / "s.xlsx")
    _ok("xl_structure", "manage_table", path=p, name="Sales", action="delete_column",
        column="FY24")
    got = _formulas(p)
    assert [got[f"Data!G{r}"] for r in (2, 3, 4)] == [THIS_ROW] * 3
    assert got["Data!I3"] == "=SUM(Sales[Q1])+I2"
    assert got["Summ!A1"] == "=SUM(Sales[Q1])"
    assert got["Summ!A3"] == "=SUM(Sales[FY2024])"
    assert got["Summ!A13"] == "=Data!I2+SUM(Sales[Q1])"
    assert _names(p) == SALES_NAMES


# ------------------------------------------------------------------ links

def test_a_local_insert_never_moves_a_reference_to_another_workbook(tmp_path):
    p = _linked(tmp_path / "l.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_rows", at=1,
        sheet="Data")
    assert _formulas(p) == {
        "Summ!A1": "=[1]Data!A1", "Summ!A2": "=SUM([1]Data!A1:B3)+Data!A2",
        "Summ!A3": "=[1]Data!$A$2", "Summ!A4": "='[1]My Data'!A1",
        "Summ!A5": "=[1]!ExtTotal+Data!B3", "Summ!A6": "=[1]Data!A3*Data!A4",
        "Summ!A7": "=Data!A5", "Summ!A10": "=ExtRate*2"}
    assert _names(p) == {"ExtRate": "[1]Data!$B$1"}


def test_a_local_delete_never_deletes_a_reference_to_another_workbook(tmp_path):
    # Deleting row 1 here once turned the name's [1]Data!$B$1 into
    # [1]Data!$#REF!, which Excel will not open.
    p = _linked(tmp_path / "l.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="delete_rows", at=1,
        sheet="Data")
    got = _formulas(p)
    assert got["Summ!A1"] == "=[1]Data!A1"
    assert got["Summ!A2"] == "=SUM([1]Data!A1:B3)+Data!#REF!"
    assert got["Summ!A6"] == "=[1]Data!A3*Data!A2"
    assert _names(p) == {"ExtRate": "[1]Data!$B$1"}


# ------------------------------------------------------------------ scripts

def test_a_sheet_named_in_another_script_follows_a_row_insert(tmp_path):
    p = _scripts(tmp_path / "u.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_rows", at=1,
        sheet="Données")
    got = _formulas(p)
    assert got["Summ!A1"] == "=Données!A6*2"
    assert got["Summ!A6"] == "=Données!A6+es!A5"
    assert got["Summ!A11"] == "=COUNT(Données!A2:B8)"
    assert got["Summ!A7"] == "=SUM(Données!A:A)"
    assert _names(p)["DonnéesTot"] == "Données!$A$6"
    ws = openpyxl.load_workbook(p)["Summ"]
    assert [r.formula for rs in ws.conditional_formatting._cf_rules.values() for r in rs] \
        == [["Données!$A$6>0"]]


def test_each_script_follows_its_own_sheet(tmp_path):
    p = _scripts(tmp_path / "u.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_cols", at="A",
        sheet="ยอดขาย")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_cols", at="A",
        sheet="Cost€")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_rows", at=1,
        sheet="Q１")
    _ok("xl_structure", "modify_grid_structure", path=p, action="delete_rows", at=2,
        sheet="売上")
    got = _formulas(p)
    assert got["Summ!A3"] == "=ยอดขาย!B2+ยอดขาย!C3"
    assert got["Summ!A4"] == "=Cost€!B1"
    assert got["Summ!A9"] == "=Q１!A3"
    assert got["Summ!A2"] == "=SUM(売上!B2:B3)"
    assert got["Summ!A8"] == "=SUM(売上!2:2)"
    assert _names(p)["売上Range"] == "売上!$B$2:$B$3"
    dv = openpyxl.load_workbook(p)["Summ"].data_validations.dataValidation[0]
    assert dv.formula1 == "売上!$B$2:$B$3"


def test_an_insert_on_the_formulas_own_sheet_leaves_other_scripts_alone(tmp_path):
    # Q１ (a full-width digit) once read as cell Q1 of this sheet: Q2!A2.
    p = _scripts(tmp_path / "u.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_rows", at=1,
        sheet="Summ")
    got = _formulas(p)
    assert got["Summ!A10"] == "=Q１!A2"
    assert got["Summ!A3"] == "=SUM(売上!B2:B4)"
    assert _names(p)["売上Range"] == "売上!$B$2:$B$4"


def test_renaming_a_short_sheet_never_touches_a_longer_name_that_ends_with_it(tmp_path):
    p = _scripts(tmp_path / "u.xlsx")
    _ok("xl_structure", "manage_worksheet", path=p, action="rename", sheet="es",
        new_name="Zed")
    got = _formulas(p)
    assert got["Summ!A1"] == "=Données!A5*2"
    assert got["Summ!A6"] == "=Données!A5+Zed!A5"
    assert _names(p)["DonnéesTot"] == "Données!$A$5"


def test_renaming_and_deleting_sheets_named_in_other_scripts(tmp_path):
    p = _scripts(tmp_path / "u.xlsx")
    _ok("xl_structure", "manage_worksheet", path=p, action="rename", sheet="Données",
        new_name="Résumé 2")
    _ok("xl_structure", "manage_worksheet", path=p, action="delete", sheet="売上")
    got = _formulas(p)
    assert got["Summ!A1"] == "='Résumé 2'!A5*2"
    assert got["Summ!A11"] == "=COUNT('Résumé 2'!A1:B7)"
    assert got["Summ!A2"] == "=SUM(#REF!)"
    assert got["Summ!A8"] == "=SUM(#REF!)"
    names = _names(p)
    assert names["DonnéesTot"] == "'Résumé 2'!$A$5"
    assert names["売上Range"] == "#REF!"


# ------------------------------------------------------------------ quotes and names

def test_a_sheet_name_holding_double_quotes_follows_an_insert(tmp_path):
    p = _quoted(tmp_path / "q.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_rows", at=1,
        sheet='North "Q"')
    got = _formulas(p)
    assert [got[f"Summ!A{r}"] for r in range(1, 6)] == [
        "='North \"Q\"'!A2", "=SUM('North \"Q\"'!A2:A4)",
        "=\"North \"\"Q\"\"!A1\"&'North \"Q\"'!A3", "='North Q'!A2+'North \"Q\"'!A4",
        "=LEN(\"'North \"\"Q\"\"'!A1\")+'North \"Q\"'!A5"]


def test_a_sheet_name_holding_double_quotes_follows_a_rename_and_a_delete(tmp_path):
    p = _quoted(tmp_path / "q.xlsx")
    _ok("xl_structure", "manage_worksheet", path=p, action="rename", sheet='North "Q"',
        new_name='South "Q"')
    got = _formulas(p)
    assert got["Summ!A3"] == "=\"North \"\"Q\"\"!A1\"&'South \"Q\"'!A2"
    assert got["Summ!A4"] == "='North Q'!A2+'South \"Q\"'!A3"
    _ok("xl_structure", "manage_worksheet", path=p, action="delete", sheet='South "Q"')
    got = _formulas(p)
    assert [got[f"Summ!A{r}"] for r in range(1, 6)] == [
        "=#REF!", "=SUM(#REF!)", "=\"North \"\"Q\"\"!A1\"&#REF!", "='North Q'!A2+#REF!",
        "=LEN(\"'North \"\"Q\"\"'!A1\")+#REF!"]


def test_names_whose_insides_look_like_cells_survive_an_insert_and_a_copy(tmp_path):
    p = _quoted(tmp_path / "q.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_cols", at="A",
        sheet="Summ")
    got = _formulas(p)
    assert [got[f"Summ!F{r}"] for r in range(1, 7)] == [
        "=αA1+B1", "=A1α*2+C2", "=売上A1+B1", "=ÀB12+C12", "=x.A1+B1", "=Q1β+B1"]
    assert got["Summ!B2"] == "=SUM('North \"Q\"'!A1:A3)"
    assert _names(p)["Q1β"] == "Summ!$D$6"
    q = _quoted(tmp_path / "q2.xlsx")
    _ok("xl_edit", "copy_range", path=q, source={"range": "E1:E6", "sheet": "Summ"},
        dest={"cell": "F2", "sheet": "Summ"})
    got = _formulas(q)
    assert [got[f"Summ!F{r}"] for r in range(2, 8)] == [
        "=αA1+B2", "=A1α*2+C3", "=売上A1+B2", "=ÀB12+C13", "=x.A1+B2", "=Q1β+B2"]


# ------------------------------------------------------------------ Excel as the judge

def _excel_edit(path: str, action) -> None:
    """Make the same edit in the product's own private hidden Excel."""
    def job(manager):
        w = manager.acquire()
        prior = com_session._assert_hygiene(w.app)
        wb = None
        try:
            wb = com_session.open_workbook(w.app, path)
            action(wb)
            wb.Save()
        finally:
            if wb is not None:
                wb.Close(False)
            com_session._restore_hygiene(w.app, *prior)
            manager.release_to_pool(w)
    com_session.run_com("oracle", job, timeout=120)


EXCEL_EDITS = [
    (_sales, ("modify_grid_structure", {"action": "insert_rows", "at": 1, "sheet": "Data"}),
     lambda wb: wb.Worksheets("Data").Rows("1:1").Insert()),
    (_sales, ("modify_grid_structure", {"action": "insert_cols", "at": "A", "sheet": "Data"}),
     lambda wb: wb.Worksheets("Data").Columns("A:A").Insert()),
    (_scripts, ("modify_grid_structure", {"action": "insert_rows", "at": 1, "sheet": "Données"}),
     lambda wb: wb.Worksheets("Données").Rows("1:1").Insert()),
    (_scripts, ("manage_worksheet", {"action": "rename", "sheet": "es", "new_name": "Zed"}),
     lambda wb: setattr(wb.Worksheets("es"), "Name", "Zed")),
    (_quoted, ("modify_grid_structure", {"action": "insert_rows", "at": 1, "sheet": "Summ"}),
     lambda wb: wb.Worksheets("Summ").Rows("1:1").Insert()),
    (_quoted, ("manage_worksheet", {"action": "delete", "sheet": 'North "Q"'}),
     lambda wb: wb.Worksheets('North "Q"').Delete()),
]


@pytest.mark.com
@pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool)")
@pytest.mark.timeout(900)
def test_excel_stores_the_same_references_after_the_same_edit(tmp_path):
    for i, (build, (op, args), excel_action) in enumerate(EXCEL_EDITS):
        ours = build(tmp_path / f"ours{i}.xlsx")
        theirs = str(tmp_path / f"excel{i}.xlsx")
        shutil.copy2(ours, theirs)
        _ok("xl_structure", op, path=ours, **args)
        _excel_edit(theirs, excel_action)
        assert _formulas(ours) == _formulas(theirs), (op, args)
        assert _names(ours) == _names(theirs), (op, args)
        assert srv.com_validate_opens_clean(ours)["opens_clean"] is True, (op, args)
