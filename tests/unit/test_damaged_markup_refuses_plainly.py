"""A workbook with damaged markup refuses plainly, as damaged.

A cell pointing past the shared strings escaped as INTERNAL_ERROR
("list index out of range"), and a worksheet whose XML is cut off refused as
BAD_PARAMS with "the message says what to change", as if the caller's
arguments were wrong. Both now refuse as a damaged workbook
(UNSUPPORTED_CONTENT) on every read and write route. All content is made up.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl
import pytest

from xlsx_mcp import envelope
from xlsx_mcp.core.errors import WorkbookCorrupt
from xlsx_mcp.ops import cells, lifecycle


def _book(tmp_path: Path, fix) -> str:
    p = tmp_path / "damaged.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Region"
    ws["A2"] = "North"
    ws["B2"] = 12
    wb.save(p)
    with zipfile.ZipFile(p) as zin:
        items = [(i, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            if info.filename == "xl/worksheets/sheet1.xml":
                data = fix(data.decode("utf-8")).encode("utf-8")
            zout.writestr(info, data)
    return str(p)


def _sst_past_end(x: str) -> str:
    # A2 becomes a shared-string cell pointing far past the last string
    import re
    return re.sub(r'<c r="A2"[^>]*>.*?</c>', '<c r="A2" t="s"><v>99</v></c>',
                  x, count=1, flags=re.S)


def _cut_off(x: str) -> str:
    return x[: x.index("</sheetData>")]


DAMAGE = {"shared string past the end": _sst_past_end,
          "worksheet XML cut off": _cut_off}
ROUTES = {
    "read_range": lambda p: cells.read_range(p, "A1:B2"),
    "metadata": lambda p: lifecycle.get_workbook_metadata(p),
    "write_range": lambda p: cells.write_range(p, "C1", [[1]]),
}


def test_the_shared_string_damage_is_really_there(tmp_path):
    p = _book(tmp_path, _sst_past_end)
    with zipfile.ZipFile(p) as z:
        assert "<v>99</v>" in z.read("xl/worksheets/sheet1.xml").decode()


@pytest.mark.parametrize("route", sorted(ROUTES))
@pytest.mark.parametrize("damage", sorted(DAMAGE))
def test_damaged_markup_refuses_as_a_damaged_workbook(tmp_path, damage,
                                                      route):
    p = _book(tmp_path, DAMAGE[damage])
    before = Path(p).read_bytes()
    with pytest.raises(WorkbookCorrupt) as ei:
        ROUTES[route](p)
    assert envelope.classify(ei.value) == "UNSUPPORTED_CONTENT"
    assert Path(p).read_bytes() == before
