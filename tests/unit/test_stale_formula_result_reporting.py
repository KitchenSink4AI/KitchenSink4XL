"""Stale formula results are reported, and said so by
every reader.

First: the dependency walk missed formulas that reach an edited cell through a
LAMBDA name used as a function (=Twice(5)) or a relative name (RowVal =
Inputs!$B1); the edit reported nothing stale and left fullCalcOnLoad unset,
so Excel itself showed the old result. What-if data table results were never
walked either.

Second: after an edit that leaves a result stale, read_range and get_cells said so,
but query_range, export_range, export_file, get_grid_view and find_cells used
the stale cached value with no note.
"""
from __future__ import annotations

import re
import zipfile

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName

from xlsx_mcp.ops import cells, dataio, gridio, search, view


def _calcpr(p):
    m = re.search(rb"<calcPr[^>]*>", zipfile.ZipFile(p).read("xl/workbook.xml"))
    return m.group(0).decode() if m else ""


def _patch(src, out, edits):
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            b = zin.read(i.filename)
            for part, fn in edits:
                if i.filename == part:
                    b = fn(b)
            zout.writestr(i, b)


def _no_pending(b):
    return re.sub(rb"<calcPr[^>]*/>", b'<calcPr calcId="191029"/>', b)



def _names_book(tmp_path):
    base = str(tmp_path / "base.xlsx")
    wb = openpyxl.Workbook()
    inp = wb.active
    inp.title = "Inputs"
    for i in range(1, 11):
        inp.cell(i, 2, i * 10)
    inp["D1"] = 7
    calc = wb.create_sheet("Calc")
    wb.defined_names["Twice"] = DefinedName(
        "Twice", attr_text="_xlfn.LAMBDA(_xlpm.x, _xlpm.x*2*Inputs!$B$2)")
    wb.defined_names["RowVal"] = DefinedName("RowVal", attr_text="Inputs!$B1")
    wb.defined_names["Fixed"] = DefinedName("Fixed", attr_text="Inputs!$B$3")
    calc["A1"] = "=Twice(5)"
    calc["D5"] = "=RowVal"
    calc["E1"] = "=Fixed"
    wb.save(base)
    p = str(tmp_path / "names.xlsx")
    _patch(base, p, [("xl/workbook.xml", _no_pending)])
    return p


@pytest.mark.parametrize("cell,watch", [("B2", "Calc!A1"), ("B5", "Calc!D5")],
                         ids=["lambda_name", "relative_name"])
def test_dependents_through_names_are_reported_and_recalculated(tmp_path, cell, watch):
    p = _names_book(tmp_path)
    out = cells.set_cell(p, {"cell": cell, "sheet": "Inputs"}, 999)
    cv = out.get("cached_values") or {}
    assert watch in (cv.get("possibly_stale_cells") or []), cv
    assert "fullCalcOnLoad" in _calcpr(p)


def test_an_absolute_name_stays_precise(tmp_path):
    """Control: an edit nothing depends on reports nothing stale and leaves
    calcPr alone; an edit to the absolute name's cell reports its user."""
    p = _names_book(tmp_path)
    out = cells.set_cell(p, {"cell": "D1", "sheet": "Inputs"}, 8)
    cv = out.get("cached_values") or {}
    assert "Calc!E1" not in (cv.get("possibly_stale_cells") or [])
    assert "Calc!A1" not in (cv.get("possibly_stale_cells") or [])
    (tmp_path / "q").mkdir()
    q = _names_book(tmp_path / "q")  # fresh copy
    out = cells.set_cell(q, {"cell": "B3", "sheet": "Inputs"}, 8)
    assert "Calc!E1" in (out.get("cached_values") or {}).get("possibly_stale_cells", [])


def test_name_used_as_function_resolves_its_body():
    from xlsx_mcp.core import surgical
    names = {(None, "TWICE"): (None, "=_xlfn.LAMBDA(_xlpm.x, _xlpm.x*2*Inputs!$B$2)")}
    regs = surgical._formula_regions("=Twice(5)", "Calc", names, {})
    assert ("inputs", 2, 2, 2, 2) in regs


@pytest.mark.parametrize("text,rel", [("=Inputs!$B1", True), ("=Inputs!B$1", True),
                                      ("=Inputs!$B$1", False), ("=Inputs!$A:$A", False),
                                      ("=Inputs!A:A", True), ("=Inputs!$B$1:$C$4", False),
                                      ("=Inputs!$B$1:C4", True)])
def test_relative_reference_detection(text, rel):
    from xlsx_mcp.core import surgical
    assert surgical._has_relative_ref(text) is rel


def test_data_table_results_are_walked(tmp_path):
    from xlsx_mcp.core import surgical
    from openpyxl.worksheet.formula import DataTableFormula
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["B1"] = "=A1*2"
    ws["C22"] = DataTableFormula(ref="C22:C24", r1="A1", dt2D=False, dtr=False)
    st = surgical.stale_formulas(wb, {ws.title: {(1, 1)}})
    got = {a for _s, a in st.cells}
    assert {"B1", "C22", "C23", "C24"} <= got



def _cached_book(tmp_path, pending: bool):
    base = str(tmp_path / "cbase.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Item", "Qty", "Amount"])
    ws.append(["a", 2, "=B2*10"])
    ws.append(["b", 3, "=B3*10"])
    wb.save(base)

    def caches(b):
        b = b.replace(b"<f>B2*10</f><v></v>", b"<f>B2*10</f><v>20</v>")
        b = b.replace(b"<f>B3*10</f><v></v>", b"<f>B3*10</f><v>30</v>")
        b = re.sub(rb"<c r=\"C2\"([^>]*)><f>B2\*10</f>(?:<v\s*/>)?</c>",
                   rb'<c r="C2"\1><f>B2*10</f><v>20</v></c>', b)
        b = re.sub(rb"<c r=\"C3\"([^>]*)><f>B3\*10</f>(?:<v\s*/>)?</c>",
                   rb'<c r="C3"\1><f>B3*10</f><v>30</v></c>', b)
        return b
    p = str(tmp_path / ("pending.xlsx" if pending else "fresh.xlsx"))
    edits = [("xl/worksheets/sheet1.xml", caches)]
    if not pending:
        edits.append(("xl/workbook.xml", _no_pending))
    _patch(base, p, edits)
    assert cells.read_range(p, "C2:C3")["values"] == [[20], [30]]
    return p


def _reads(p, tmp_path):
    return {
        "query_range": cells.query_range(p, "A1:C3", aggregate=[{"column": "Amount", "func": "sum"}]),
        "query_rows": cells.query_range(p, "A1:C3"),
        "export_range": dataio.export_range(p, "A1:C3", fmt="json"),
        "export_file": dataio.export_file(p, fmt="json"),
        "get_grid_view": view.get_grid_view(p),
        "find_cells": search.find_cells(p, "20"),
    }


def test_every_reader_says_when_cached_results_are_owed_a_recalculation(tmp_path):
    p = _cached_book(tmp_path, pending=False)
    out = cells.set_cell(p, "B2", 99)
    assert "Sales!C2" in (out.get("cached_values") or {}).get("possibly_stale_cells", []) \
        or "C2" in str(out.get("cached_values"))
    for name, r in _reads(p, tmp_path).items():
        assert r.get("cached_note") == gridio.PENDING_NOTE, name


def test_no_note_on_a_workbook_that_is_not_owed_a_recalculation(tmp_path):
    p = _cached_book(tmp_path, pending=False)
    for name, r in _reads(p, tmp_path).items():
        assert "cached_note" not in r, name


def test_no_note_when_the_range_holds_no_formula(tmp_path):
    p = _cached_book(tmp_path, pending=True)
    assert "cached_note" not in cells.query_range(p, "A1:B3")
    assert "cached_note" not in dataio.export_range(p, "A1:B3")
