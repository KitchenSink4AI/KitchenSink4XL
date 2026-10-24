"""A file-tier save keeps the workbook's
identity.

A file-tier edit put the new version in place with os.replace, which gives the
path a NEW file: one ordinary set_cell dropped the workbook's explicit
permissions (a DENY entry and an explicit grant vanished, so a restricted
workbook fell back to whatever its folder hands down), its Mark of the Web
(no Protected View afterwards, and Office's block on macros from the internet
no longer applied), its attributes (OneDrive's pinned and offline states, not
content indexed) and its creation time. Excel's own save keeps all of them.
safesave.promote now uses ReplaceFileW, which keeps them (measured), for the
save, the backup restore and the overwrite paths.
"""
from __future__ import annotations

import ctypes
import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from xlsx_mcp.core import safesave
from xlsx_mcp.core.errors import WorkbookLocked

pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="Windows file identity")

GENERIC_READ = 0x80000000
SHARE_READ, SHARE_WRITE, SHARE_DELETE = 0x1, 0x2, 0x4


def _icacls(*a):
    return subprocess.run(["icacls", *a], capture_output=True, text=True,
                          creationflags=0x08000000).stdout


def _explicit_aces(p):
    lines = [ln.strip() for ln in _icacls(p).splitlines() if ln.strip() and "Successfully" not in ln]
    lines[0] = lines[0].replace(p, "").strip()
    return sorted(ln for ln in lines if "(I)" not in ln)


def _attrs(p):
    return ctypes.windll.kernel32.GetFileAttributesW(p)


@pytest.fixture
def book(tmp_path):
    p = str(tmp_path / "book.xlsx")
    wb = openpyxl.Workbook()
    wb.active["A1"] = "before"
    wb.save(p)
    yield p
    _icacls(p, "/reset")
    ctypes.windll.kernel32.SetFileAttributesW(p, 0x80)


def _leftovers(folder):
    return sorted(n for n in os.listdir(folder) if n.startswith(".ks4xl-write-"))


def _edit(p, value="edited"):
    from xlsx_mcp.ops import cells
    r = cells.set_cell(p, "A1", value)
    assert r["ok"] and r["verified"]
    assert openpyxl.load_workbook(p).active["A1"].value == value
    return r


def test_explicit_permissions_survive(book):
    _icacls(book, "/grant", "*S-1-5-32-545:(R)")
    _icacls(book, "/deny", "*S-1-5-7:(R)")
    before = _explicit_aces(book)
    # The two entries just added are there; an elevated account's new files
    # carry a few explicit entries of their own, which must survive too.
    assert {"BUILTIN\\Users:(R)", "NT AUTHORITY\\ANONYMOUS LOGON:(DENY)(R)"} <= set(before), before
    _edit(book)
    assert _explicit_aces(book) == before


def test_the_mark_of_the_web_survives(book):
    with open(book + ":Zone.Identifier", "w") as f:
        f.write("[ZoneTransfer]\r\nZoneId=3\r\nHostUrl=https://example.invalid/b.xlsx\r\n")
    _edit(book)
    with open(book + ":Zone.Identifier") as f:
        assert "ZoneId=3" in f.read()


@pytest.mark.parametrize("attr", [0x2000, 0x1000, 0x80000, 0x100000],
                         ids=["not_content_indexed", "offline", "pinned", "unpinned"])
def test_attributes_survive(book, attr):
    if not ctypes.windll.kernel32.SetFileAttributesW(book, 0x20 | attr):
        pytest.skip("this file system does not take that attribute")
    _edit(book)
    assert _attrs(book) & attr


def test_the_creation_time_survives(book):
    before = os.stat(book).st_birthtime
    _edit(book)
    assert os.stat(book).st_birthtime == before


def test_the_prev_slot_still_holds_the_old_content(book):
    _edit(book, "one")
    _edit(book, "two")
    slot = safesave.slot_dir(book) / safesave.PREV_SLOT
    assert openpyxl.load_workbook(slot).active["A1"].value == "one"
    assert _leftovers(os.path.dirname(book)) == []


def test_an_edit_beside_a_reader_that_shares_delete(book):
    """Excel saves beside such a reader; so does the file tier now (os.replace
    could not replace a file another program had open)."""
    k32 = ctypes.WinDLL("kernel32", use_last_error=True)
    k32.CreateFileW.restype = ctypes.c_void_p
    k32.CreateFileW.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p,
                                ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p]
    h = k32.CreateFileW(book, GENERIC_READ, SHARE_READ | SHARE_WRITE | SHARE_DELETE, None, 3, 0x80, None)
    assert h and h != ctypes.c_void_p(-1).value
    try:
        _edit(book)
    finally:
        k32.CloseHandle(ctypes.c_void_p(h))


def test_a_restore_keeps_the_identity_too(book):
    """Putting a backup back is a replace as well: the workbook keeps the
    permissions it has now."""
    from xlsx_mcp.ops import backups
    _edit(book)
    _icacls(book, "/grant", "*S-1-5-32-545:(R)")
    _icacls(book, "/deny", "*S-1-5-7:(R)")
    before = _explicit_aces(book)
    backups.restore_backup(book, "prev")
    assert openpyxl.load_workbook(book).active["A1"].value == "before"
    assert _explicit_aces(book) == before


def test_a_holder_at_the_moment_of_the_swap_changes_nothing(book, monkeypatch):
    """A program that takes the file between the probe and the swap: the
    workbook stays as it was, nothing is left beside it, and the refusal
    names the cause."""
    from xlsx_mcp.core.package import WorkbookPackage
    monkeypatch.setattr(WorkbookPackage, "_lock_diagnosis", lambda self: (False, True, None))
    before = Path(book).read_bytes()
    with open(book, "rb"):  # the C runtime's default: shares read and write, not delete
        with pytest.raises(WorkbookLocked) as ei:
            _edit(book)
    msg = str(ei.value)
    assert "in use by another program" in msg
    assert "WinError" not in msg and "may be open in Excel" not in msg
    assert Path(book).read_bytes() == before
    assert _leftovers(os.path.dirname(book)) == []


def test_a_failed_swap_that_moved_the_original_aside_puts_it_back(tmp_path, monkeypatch):
    """ReplaceFileW's ERROR_UNABLE_TO_MOVE_REPLACEMENT_2: the original was
    renamed to the backup name; promote moves it back before refusing."""
    dst = tmp_path / "w.xlsx"
    dst.write_bytes(b"original")
    src = tmp_path / ".ks4xl-write-new.xlsx"
    src.write_bytes(b"new")
    state = {}

    def fake_replace(ldst, lsrc, lbak, flags, a, b):
        os.replace(ldst, lbak)
        state["bak"] = lbak
        return False

    def fake_move(a, b, flags):
        os.replace(a, b)
        return True
    monkeypatch.setattr(safesave, "_REPLACE_FILE",
                        (fake_replace, fake_move, lambda: 1177, lambda e: "could not move"))
    with pytest.raises(OSError):
        safesave.promote(str(src), str(dst))
    assert dst.read_bytes() == b"original"
    assert not os.path.exists(state["bak"])
