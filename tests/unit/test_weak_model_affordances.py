"""Weak-model affordances. The rule: judge a product by what a weak model
TELLS THE USER, so the product's own words make the truth hard to miss. One
section per finding from running a weak model as the user, each red before
its fix:

  * an edit on a workbook with a password to modify says so
  * numbers of 2^53 or more say their digits cannot be kept as a number
  * a numeric-looking string says it is stored as TEXT
  * the server instructions tell a model to find_cells before it guesses
  * get_workbook_metadata flags a sheet-scoped name on several sheets
  * set_filter counts data rows, visible and hidden
  17   the BAD_PARAMS misses name the valid form

The tests pin facts (which form, which words a model must see), not whole
sentences.
"""
from __future__ import annotations

import json
import zipfile

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp import license_policy, server as srv
from xlsx_mcp.core import hazard
from xlsx_mcp.ops import cells, format as fmt, lifecycle, structure

LEGACY = b'<fileSharing userName="Owner" reservationPassword="CC3D"/>'
MODERN = (b'<fileSharing readOnlyRecommended="1" userName="Owner" '
          b'algorithmName="SHA-512" hashValue="aGFzaA==" saltValue="c2FsdA==" '
          b'spinCount="100000"/>')
RO_ONLY = b'<fileSharing readOnlyRecommended="1"/>'


def _book(tmp_path, sharing: bytes | None, name="reserved.xlsx"):
    base = str(tmp_path / ("base_" + name))
    wb = openpyxl.Workbook()
    wb.active.title = "Plan"
    for r in range(1, 6):
        wb.active.cell(r, 1, r * 10)
    wb.save(base)
    out = str(tmp_path / name)
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            b = zin.read(i.filename)
            if i.filename == "xl/workbook.xml" and sharing:
                k = b.index(b"<workbookPr")
                b = b[:k] + sharing + b[k:]
            zout.writestr(i, b)
    return out


def _plain(tmp_path, name="plain.xlsx", rows=5):
    p = str(tmp_path / name)
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet"
    for r in range(1, rows + 1):
        wb.active.cell(r, 1, r * 10)
    wb.save(p)
    return p


def _joined(out) -> str:
    return " | ".join(out.get("warnings") or [])




@pytest.mark.parametrize("surgical", ["on", "off"])
@pytest.mark.parametrize("sharing", [LEGACY, MODERN], ids=["legacy", "modern"])
def test_t10_a_file_edit_on_a_password_to_modify_workbook_says_so(
        tmp_path, monkeypatch, surgical, sharing):
    monkeypatch.setenv("KS4XL_SURGICAL_SAVE", surgical)
    p = _book(tmp_path, sharing)
    out = cells.set_cell(p, "A1", 500)
    text = _joined(out)
    # The three facts: it has a password to modify, the edit was made without
    # it, and the password is kept.
    assert "password to modify" in text, text
    assert "without" in text and "kept" in text, text
    assert out["ok"] is True and out["saved"] is True
    assert openpyxl.load_workbook(p).active["A1"].value == 500
    assert hazard.write_reservation(p)["password_to_modify"] is True


def test_t10_the_warning_reaches_every_file_write_not_one_op(tmp_path):
    p = _book(tmp_path, LEGACY)
    results = [
        cells.write_range(p, "B1", [[1, 2], [3, 4]]),
        fmt.format_cells(p, "A1:A2", number_format="0.00"),
        structure.modify_grid_structure(p, "insert_rows", 2),
    ]
    for out in results:
        assert "password to modify" in _joined(out), out


def test_t10_the_warning_names_the_file(tmp_path):
    p = _book(tmp_path, LEGACY, "plan_2026.xlsx")
    assert "plan_2026.xlsx" in _joined(cells.set_cell(p, "A1", 1))


def test_t10_it_does_not_refuse_it_still_edits_and_keeps_the_password(tmp_path):
    p = _book(tmp_path, MODERN)
    out = cells.set_cell(p, "A2", 7)
    assert out["ok"] is True and out["verified"] is True
    again = hazard.write_reservation(p)
    assert again["password_to_modify"] and again["reserved_by"] == "Owner"


@pytest.mark.parametrize("sharing", [None, RO_ONLY], ids=["plain", "read-only-recommended"])
def test_t10_a_workbook_without_a_password_to_modify_gets_no_such_warning(tmp_path, sharing):
    p = _book(tmp_path, sharing, "other.xlsx")
    assert "password to modify" not in _joined(cells.set_cell(p, "A1", 1))


def test_t10_kept_is_said_only_when_the_saved_file_still_has_the_password(tmp_path, monkeypatch):
    """The file tools are tested never to drop the password, but "the password
    is kept" is a claim about the saved file, so it is read back, not assumed."""
    p = _book(tmp_path, LEGACY)
    real = hazard.write_reservation
    calls = {"n": 0}

    def vanishing(path):
        calls["n"] += 1
        return real(path) if calls["n"] == 1 else None  # present before, gone after

    monkeypatch.setattr(hazard, "write_reservation", vanishing)
    text = _joined(cells.set_cell(p, "A1", 1))
    assert "no longer carries it" in text and "is kept" not in text, text
    assert "manage_backups" in text and "'prev'" in text, text


def test_t10_the_warning_survives_the_envelope(tmp_path):
    p = _book(tmp_path, LEGACY)
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_cell", "args": {
        "path": p, "location": "B2", "value": 500, "sheet": "Plan"}})
    assert not err, payload
    assert any("password to modify" in w for w in payload["warnings"])



TWO53 = 2 ** 53


@pytest.mark.parametrize("value", [TWO53, TWO53 + 1, -TWO53, -(TWO53 + 1),
                                   12345678901234567890, float(TWO53), -float(TWO53), 1e16, 1e20],
                         ids=["2^53", "2^53+1", "-2^53", "-(2^53+1)", "20 digits",
                              "float 2^53", "float -2^53", "float 1e16", "float 1e20"])
def test_t15_a_number_of_2_53_or_more_says_its_digits_cannot_be_kept(tmp_path, value):
    p = _plain(tmp_path)
    out = cells.set_cell(p, "A1", value)
    text = _joined(out)
    assert "Sheet!A1" in text, text
    assert "exact digits" in text, text
    # the remedy: a string, with the Text number format
    assert "string" in text and "Text" in text, text
    assert "format_cells" in text and '"@"' in text, text


@pytest.mark.parametrize("value", [TWO53 - 1, -(TWO53 - 1), 123456789012345, 1e15, 3.14, 0, True,
                                   float(TWO53 - 2)],
                         ids=["2^53-1", "-(2^53-1)", "15 digits", "1e15", "pi-ish", "zero",
                              "bool", "float below 2^53"])
def test_t15_a_smaller_number_says_nothing(tmp_path, value):
    p = _plain(tmp_path)
    assert "exact digits" not in _joined(cells.set_cell(p, "A1", value))


def test_t15_a_block_names_the_cells_it_hit(tmp_path):
    p = _plain(tmp_path)
    out = cells.write_range(p, "B1", [[TWO53, 5], [TWO53 + 1, 6], [7, 8]])
    text = _joined(out)
    assert "Sheet!B1" in text and "Sheet!B2" in text and "Sheet!B3" not in text, text


def test_t15_following_the_advice_is_quiet_and_keeps_every_digit(tmp_path):
    """The remedy works: the digits as a string, then the Text number format."""
    p = _plain(tmp_path)
    out = cells.set_cell(p, "A1", "9007199254740993")
    assert "exact digits" not in _joined(out)
    assert "TEXT" not in _joined(out)  # sixteen digits cannot be a faithful number
    fmt.format_cells(p, "A1", number_format="@")
    cell = openpyxl.load_workbook(p).active["A1"]
    assert cell.value == "9007199254740993" and cell.number_format == "@"




def _numbers_column(tmp_path, name="col.xlsx"):
    p = str(tmp_path / name)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    for r, v in enumerate([100, 200, 300, 400], start=1):
        ws.cell(r, 2, v)
    ws["B5"] = "=SUM(B1:B4)"
    wb.save(p)
    return p


def test_t04_a_numeric_string_is_named_as_text_with_what_it_costs(tmp_path):
    p = _numbers_column(tmp_path)
    out = cells.set_cell(p, "B3", "250")
    text = _joined(out)
    assert "Sheet!B3" in text and '"250"' in text, text
    assert "TEXT" in text, text
    assert "SUM" in text and "AVERAGE" in text, text
    assert "as a number" in text, text


def test_t04_neighbours_that_are_numbers_make_it_louder(tmp_path):
    p = _numbers_column(tmp_path)
    near = _joined(cells.set_cell(p, "B3", "250"))
    assert "hold numbers" in near, near
    lone = _joined(cells.set_cell(_plain(tmp_path, "lone.xlsx", rows=0), "D9", "250"))
    assert "TEXT" in lone and "hold numbers" not in lone, lone


@pytest.mark.parametrize("text", ["250", "-3.5", "1,250", "1e3", " 42 ", ".5"])
def test_t04_what_counts_as_numeric_looking(tmp_path, text):
    p = _numbers_column(tmp_path)
    assert "TEXT" in _joined(cells.set_cell(p, "B3", text))


def test_t04_the_advice_names_the_number_without_separators(tmp_path):
    """A cosmetic note: "1,250" is advised as 1250,
    the form a model can pass back as a number."""
    p = _numbers_column(tmp_path)
    text = _joined(cells.set_cell(p, "B3", "1,250"))
    assert "Pass 1250 as a number" in text and "not the number 1250" in text, text
    assert "Pass 1,250" not in text


@pytest.mark.parametrize("text", ["abc", "A1", "2026-03-24", "12%", "$250", "1.2.3", "",
                                  "007", "00123", "-0123",
                                  "1234567890123456", "12345678901234567890"])
def test_t04_codes_ids_and_ordinary_text_are_left_alone(tmp_path, text):
    p = _numbers_column(tmp_path)
    assert "TEXT" not in _joined(cells.set_cell(p, "B3", text))


def test_t04_a_deliberate_text_value_is_quiet(tmp_path):
    p = _numbers_column(tmp_path)
    out = cells.set_cell(p, "B3", {"text": "250"})
    assert "TEXT" not in _joined(out)
    assert openpyxl.load_workbook(p).active["B3"].value == "250"


def test_t04_a_cell_already_formatted_as_text_is_quiet(tmp_path):
    p = _numbers_column(tmp_path)
    fmt.format_cells(p, "B3", number_format="@")
    assert "TEXT" not in _joined(cells.set_cell(p, "B3", "250"))


def test_t04_a_formula_and_a_real_number_are_quiet(tmp_path):
    p = _numbers_column(tmp_path)
    assert "TEXT" not in _joined(cells.set_cell(p, "B3", "=1+249"))
    assert "TEXT" not in _joined(cells.set_cell(p, "B3", 250))


def test_t04_write_range_lists_the_cells_and_counts_the_rest(tmp_path):
    p = _numbers_column(tmp_path)
    rows = [[str(100 + i), f"x{i}"] for i in range(9)]
    out = cells.write_range(p, "D1", rows)
    text = _joined(out)
    assert "Sheet!D1" in text and "4 more" in text, text
    assert text.count("TEXT") == 1, "one warning for the block, not one per cell"


def test_t04_set_cells_and_apply_edits_say_it_too(tmp_path):
    p = _numbers_column(tmp_path)
    one = cells.set_cells(p, [{"cell": "B3", "value": "250"}, {"cell": "C1", "value": "x"}])
    assert "Sheet!B3" in _joined(one) and "Sheet!C1" not in _joined(one)
    two = srv.apply_edits(p, [{"op": "set_value", "location": "B2", "value": "75"}])
    assert "Sheet!B2" in _joined(two) and "TEXT" in _joined(two)
    three = srv.apply_edits(p, [{"op": "write_range", "location": "E1", "data": [["19"]]}])
    assert "Sheet!E1" in _joined(three)


def test_t04_the_last_write_to_a_cell_decides(tmp_path):
    p = _numbers_column(tmp_path)
    out = srv.apply_edits(p, [{"op": "set_value", "location": "B3", "value": "250"},
                              {"op": "set_value", "location": "B3", "value": 250}])
    assert "TEXT" not in _joined(out)


def test_t04_table_rows_and_columns_typed_by_the_caller_say_it_too(tmp_path):
    from openpyxl.worksheet.table import Table
    p = str(tmp_path / "tb.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws.append(["Item", "Qty"])
    ws.append(["a", 2])
    ws.append(["b", 3])
    ws.add_table(Table(displayName="Tb", ref="A1:B3"))
    wb.save(p)
    out = srv.manage_table(p, "Tb", "add_row", values=[["c", "5"]])
    assert "Sheet!B4" in _joined(out) and "TEXT" in _joined(out), out.get("warnings")
    out = srv.manage_table(p, "Tb", "add_column", column="Price", values=["10", "20", "30"])
    assert "TEXT" in _joined(out) and "Sheet!C2" in _joined(out), out.get("warnings")
    # the header cell and a plain label are not numbers somebody typed as text
    assert "Sheet!C1" not in _joined(out)


def test_t04_import_and_replace_write_text_on_purpose_and_stay_quiet(tmp_path):
    p = _numbers_column(tmp_path)
    out = srv.import_data(p, source="zip,city\n01234,Oslo\n98765,Rome\n", location="D1")
    assert "TEXT" not in _joined(out)
    q = _plain(tmp_path, "rep.xlsx")
    cells.set_cell(q, "B1", "abc")
    out = srv.replace_cells(q, "abc", "250", location="B1")
    assert "TEXT" not in _joined(out)


def test_t04_the_text_is_stored_as_asked_the_warning_does_not_coerce(tmp_path):
    p = _numbers_column(tmp_path)
    cells.set_cell(p, "B3", "250")
    assert openpyxl.load_workbook(p).active["B3"].value == "250"




def test_t08_the_instructions_tell_a_model_to_check_before_it_guesses():
    from xlsx_mcp.surface import copy
    text = copy.INSTRUCTIONS
    assert "find_cells" in text and "several sheets" in text
    assert "ask the user" in text


def test_t08_the_assembled_instructions_stay_inside_claude_codes_cap():
    """Claude Code cuts server instructions at 2,048 characters, and the
    cut falls on the licensing guidance that follows the product text."""
    text = srv.instructions_for(license_policy.active())
    assert len(text) <= 2048, len(text)
    assert "AGPL-3.0 §13" in text, "the licensing notice must not be cut"
    assert text.rstrip().endswith("for licensing."), text[-80:]


def _two_rates(tmp_path, extra=None):
    from openpyxl.workbook.defined_name import DefinedName
    p = str(tmp_path / "names.xlsx")
    wb = openpyxl.Workbook()
    us = wb.active
    us.title = "US"
    eu = wb.create_sheet("EU")
    us["B1"], eu["B1"] = 0.05, 0.19
    us.defined_names["Rate"] = DefinedName("Rate", attr_text="US!$B$1")
    eu.defined_names["Rate"] = DefinedName("Rate", attr_text="EU!$B$1")
    if extra:
        extra(wb)
    wb.save(p)
    return p


def test_t09_metadata_flags_a_sheet_scoped_name_on_several_sheets(tmp_path):
    out = lifecycle.get_workbook_metadata(_two_rates(tmp_path))
    text = _joined(out)
    assert "Rate" in text and "2 sheets" in text, text
    assert "US!$B$1" in text and "EU!$B$1" in text, text
    assert "ask" in text, text


def test_t09_the_flag_comes_through_the_tool(tmp_path):
    payload, err, _ = fh.start().call("xl_read", {"op": "get_workbook_metadata", "args": {
        "path": _two_rates(tmp_path)}})
    assert not err and any("Rate" in w for w in payload["warnings"]), payload


def test_t09_a_name_defined_once_or_at_workbook_level_is_not_flagged(tmp_path):
    from openpyxl.workbook.defined_name import DefinedName

    def distinct(wb):
        wb["EU"].defined_names.pop("Rate")
        wb["EU"].defined_names["Rate_eu"] = DefinedName("Rate_eu", attr_text="EU!$B$1")

    out = lifecycle.get_workbook_metadata(_two_rates(tmp_path, distinct))
    assert "warnings" not in out, out

    def shadow(wb):
        wb["EU"].defined_names.pop("Rate")
        wb.defined_names["Rate"] = DefinedName("Rate", attr_text="EU!$B$1")

    out = lifecycle.get_workbook_metadata(_two_rates(tmp_path, shadow))
    assert "warnings" not in out, out




def _filter_book(tmp_path):
    p = str(tmp_path / "f.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["Region", "Item", "Amount"])
    regions = ["West", "East", "North", "West", "South", "East", "West", "North",
               "West", "South", "East", "North"]
    for i, region in enumerate(regions, start=1):
        ws.append([region, f"Item{i:02d}", 100 + i * 10])
    wb.save(p)
    return p


def test_t07_set_filter_counts_data_rows_visible_and_hidden(tmp_path):
    p = _filter_book(tmp_path)
    out = srv.set_filter(p, "A1:C13", criteria=[{"column": "Region", "op": "eq",
                                                 "value": "West"}])
    f = out["changed"]["filter"]
    assert f["data_rows_visible"] == 4 and f["data_rows_hidden"] == 8, f
    assert "header" in json.dumps(f), "the counts say the header row is not in them"


def test_t07_the_counts_hold_for_a_filter_that_hides_nothing(tmp_path):
    p = _filter_book(tmp_path)
    out = srv.set_filter(p, "A1:C13", criteria=[{"column": "Amount", "op": "gt", "value": 0}])
    f = out["changed"]["filter"]
    assert f["data_rows_visible"] == 12 and f["data_rows_hidden"] == 0, f


# =========================================================== BAD_PARAMS

#: The 17 BAD_PARAMS misses a weak model made in a trial run, as the model
#: sent them. (id, tool, op, args builder, words the message must carry)
_SET_CELLS_FORM = ["set_cells", "list of", '"cell"', '"value"']
_SET_FILTER_FORM = ["criteria", "list of", '"column"', '"op"', '"value"']
_COLUMN_A = ["'A:A'", "whole-column", "A1:A13"]

MISSES = [
    ("T03_r2 set_cells, cells as JSON text of an object", "xl_edit", "set_cells",
     lambda p: {"path": p, "sheet": "Sales",
                "cells": '{"D1": "Total", "D2": "=B2*C2", "D3": "=B3*C3"}'},
     _SET_CELLS_FORM),
    ("T18_r2 set_cells, cells as an object", "xl_edit", "set_cells",
     lambda p: {"path": p, "cells": {"C1": "Bonus", "C2": "=B2*0.1", "C3": "=B3*0.1"}},
     _SET_CELLS_FORM),
    ("T05_r1 modify_grid_structure, action insert", "xl_structure", "modify_grid_structure",
     lambda p: {"path": p, "action": "insert", "at": "4", "count": "2"},
     ["'insert'", "rows or columns", "insert_rows", "insert_cols"]),
    ("T05_r2 modify_grid_structure, action insert", "xl_structure", "modify_grid_structure",
     lambda p: {"path": p, "action": "insert", "at": "4", "count": "2"},
     ["'insert'", "rows or columns", "insert_rows", "insert_cols"]),
    ("T20_r2 modify_grid_structure, action delete", "xl_structure", "modify_grid_structure",
     lambda p: {"path": p, "action": "delete", "at": "B:B", "sheet": "Sales"},
     ["'delete'", "rows or columns", "delete_rows", "delete_cols"]),
    ("T20_r2 modify_grid_structure, at B:B", "xl_structure", "modify_grid_structure",
     lambda p: {"path": p, "action": "delete_cols", "at": "B:B", "sheet": "Sales"},
     ["'B:B'", "column letter", "'B'", "count"]),
    ("T07_r1 set_filter, criteria an object of lists", "xl_structure", "set_filter",
     lambda p: {"path": p, "location": "A1:C13", "sheet": "Sales",
                "criteria": {"Region": ["West"]}},
     _SET_FILTER_FORM),
    ("T07_r2 set_filter, criteria an object", "xl_structure", "set_filter",
     lambda p: {"path": p, "location": "A1:C13", "sheet": "Sales",
                "criteria": {"Region": "West"}},
     _SET_FILTER_FORM),
    ("T12_r1 get_grid_view, A:A", "xl_read", "get_grid_view",
     lambda p: {"path": p, "sheet": "Sales", "location": "A:A"}, _COLUMN_A),
    ("T12_r2 read_range, A:A", "xl_read", "read_range",
     lambda p: {"path": p, "location": "A:A", "sheet": "Sales"}, _COLUMN_A),
    ("T22_r1 replace_cells, A:A", "xl_edit", "replace_cells",
     lambda p: {"path": p, "find": "Acme Corp", "replace": "Acme Inc", "location": "A:A",
                "match": "exact"}, _COLUMN_A),
    ("T22_r2 replace_cells, A:A with a sheet", "xl_edit", "replace_cells",
     lambda p: {"path": p, "find": "Acme Corp", "replace": "Acme Inc", "location": "A:A",
                "match": "exact", "sheet": "Sales"}, _COLUMN_A),
    ("T13_r1 manage_backups restore, scope anchor", "xl_workbook", "manage_backups",
     lambda p: {"path": p, "action": "restore", "scope": "anchor"},
     ["restore takes source", "scope is for purge", 'source="anchor"']),
    ("T13_r2 manage_backups restore, scope prev", "xl_workbook", "manage_backups",
     lambda p: {"path": p, "action": "restore", "scope": "prev"},
     ["restore takes source", "scope is for purge", 'source="prev"']),
    ("T14_r1 read_range without location", "xl_read", "read_range",
     lambda p: {"path": p}, ["read_range needs location", "A1:D5", "used_range"]),
    ("T14_r2 read_range without location", "xl_read", "read_range",
     lambda p: {"path": p}, ["read_range needs location", "A1:D5", "used_range"]),
    ("T19_r2 get_grid_view, values both", "xl_read", "get_grid_view",
     lambda p: {"path": p, "sheet": "Sales", "values": "both"},
     ["get_grid_view", "'cached' or 'formula'", "read_range", "both"]),
]


@pytest.fixture
def sales(tmp_path):
    p = str(tmp_path / "sales.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Customer", "Qty", "Price", "Total", "Note"])
    for i in range(1, 13):
        ws.append([f"Acme Corp {i}", i, 10 + i, None, "x"])
    wb.save(p)
    return p


def test_there_are_seventeen_recorded_misses():
    assert len(MISSES) == 17


@pytest.mark.parametrize("case", MISSES, ids=[m[0] for m in MISSES])
def test_a_bad_params_miss_names_the_valid_form(case, sales):
    _id, tool, op, build, words = case
    before = open(sales, "rb").read()
    payload, err, _ = fh.start().call(tool, {"op": op, "args": build(sales)})
    assert err and payload["error"]["code"] == "BAD_PARAMS", payload
    message = payload["error"]["message"]
    for word in words:
        assert word in message, (word, message)
    assert open(sales, "rb").read() == before, "a refused call changes nothing"


def test_the_whole_row_form_names_the_row_range(sales):
    payload, err, _ = fh.start().call("xl_read", {"op": "read_range", "args": {
        "path": sales, "location": "2:2", "sheet": "Sales"}})
    assert err and "A2:E2" in payload["error"]["message"], payload


def test_a_whole_column_on_an_empty_sheet_says_there_is_no_data(tmp_path):
    p = str(tmp_path / "empty.xlsx")
    openpyxl.Workbook().save(p)
    payload, err, _ = fh.start().call("xl_read", {"op": "read_range", "args": {
        "path": p, "location": "A:A"}})
    assert err and "no data" in payload["error"]["message"], payload


def test_insert_with_a_column_letter_points_at_the_column_form(sales):
    payload, err, _ = fh.start().call("xl_structure", {"op": "modify_grid_structure", "args": {
        "path": sales, "action": "insert", "at": "C"}})
    msg = payload["error"]["message"]
    assert err and 'action="insert_cols"' in msg and 'at="C"' in msg, msg


def test_insert_with_a_row_number_points_at_the_row_form(sales):
    payload, err, _ = fh.start().call("xl_structure", {"op": "modify_grid_structure", "args": {
        "path": sales, "action": "insert", "at": 4, "count": 2}})
    msg = payload["error"]["message"]
    assert err and 'action="insert_rows"' in msg and "at=4" in msg, msg


def test_the_right_forms_still_work(sales):
    """The new wording must not have made the valid forms stricter."""
    ok = fh.start().calls([
        ("xl_edit", {"op": "set_cells", "args": {"path": sales, "sheet": "Sales",
                                                 "cells": [{"cell": "D2", "value": 1}]}}),
        ("xl_structure", {"op": "set_filter", "args": {
            "path": sales, "location": "A1:E13", "sheet": "Sales",
            "criteria": [{"column": "Qty", "op": "gt", "value": 3}]}}),
        ("xl_read", {"op": "read_range", "args": {"path": sales, "sheet": "Sales",
                                                  "location": {"used_range": True}}}),
        ("xl_structure", {"op": "modify_grid_structure", "args": {
            "path": sales, "sheet": "Sales", "action": "delete_cols", "at": "B"}}),
        ("xl_workbook", {"op": "manage_backups", "args": {
            "path": sales, "action": "restore", "source": "prev"}}),
    ])
    for payload, err, _ in ok:
        assert not err and payload["ok"] is True, payload
