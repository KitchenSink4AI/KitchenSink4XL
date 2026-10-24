"""The network check reads each part of a workbook package the way Excel does: by its content type.

Excel takes the reader for a part from [Content_Types].xml (an Override for the part's name first, then the Default
for its extension), not from the part's name or its first bytes, and XML allows any amount of whitespace before the
root element. So the check reads every part of an XML type with the XML parser and the raw text, whatever the part
is called and however it starts. Only a part of a known binary type (a picture, a font, printer settings, a VBA
project, an embedded object) is not parsed, and its bytes are still searched. A part with no content type or with
one the check does not know, and a part that does not decode or parse, is a refusal with the reason named, never
"nothing found".

Every file is made up by code. No test here starts Excel.
"""
from __future__ import annotations

import re
import zipfile

import openpyxl
import pytest

from xlsx_mcp.core.netreach import READ_IN_EXCEL, file_reach

URL = "http://127.0.0.1:9/made-up"
SHEET = "xl/worksheets/sheet1.xml"


def _items(tmp_path):
    p = tmp_path / "src.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws["A1"] = "made-up"
    ws["B1"] = "=1"
    wb.save(p)
    with zipfile.ZipFile(p) as z:
        return [(i.filename, z.read(i.filename)) for i in z.infolist()]


def _write(tmp_path, items, name="book.xlsx"):
    out = tmp_path / name
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in items:
            z.writestr(n, d.encode("utf-8") if isinstance(d, str) else d)
    return str(out)


def _no_declaration(t):
    return re.sub(r"^\s*<\?xml[^>]*\?>\s*", "", t)


def _with_formula(t, formula_xml):
    t = re.sub(r'<c r="B1"[^>]*>.*?</c>', f'<c r="B1"><f>{formula_xml}</f></c>', t, count=1, flags=re.S)
    assert formula_xml in t
    return t


def _sheet_as(tmp_path, formula_xml, *, new_name, pad=5000, content_type=None, listed=True, encode=None):
    """The made-up workbook with its sheet part stored under new_name (its relationship and its Override follow
    it), the formula in B1 spelled formula_xml, pad spaces before the root element and no declaration. With
    content_type the Override names that type instead; with listed=False the Override is left out."""
    out = []
    for n, d in _items(tmp_path):
        if n == SHEET:
            t = " " * pad + _with_formula(_no_declaration(d.decode("utf-8")), formula_xml)
            out.append((new_name, encode(t) if encode else t.encode("utf-8")))
        elif n == "xl/_rels/workbook.xml.rels":
            out.append((n, d.replace(b"worksheets/sheet1.xml", new_name.split("xl/", 1)[1].encode())))
        elif n == "[Content_Types].xml":
            t = d.decode("utf-8").replace("/" + SHEET, "/" + new_name)
            if content_type is not None:
                t = re.sub(r'(<Override PartName="/' + re.escape(new_name) + r'" ContentType=")[^"]*"',
                           lambda m: m.group(1) + content_type + '"', t)
            if not listed:
                t = re.sub(r'<Override PartName="/' + re.escape(new_name) + r'"[^>]*/>', "", t)
            out.append((n, t))
        else:
            out.append((n, d))
    return _write(tmp_path, out)


def _with_member(tmp_path, name, data, default=None):
    """The made-up workbook plus one member; default=(extension, content type) adds a Default for it."""
    out = []
    for n, d in _items(tmp_path):
        if n == "[Content_Types].xml" and default:
            d = d.decode("utf-8").replace(
                "<Default ", f'<Default Extension="{default[0]}" ContentType="{default[1]}"/><Default ', 1)
        out.append((n, d))
    out.append((name, data))
    return _write(tmp_path, out)


def _found(reach, fn="WEBSERVICE"):
    return bool(reach) and any(fn in f["functions"] for f in reach["formulas"])


def _unread(reach, words):
    return bool(reach) and any(words in u for u in reach.get("unread", []))


HIDDEN = {
    "character reference": f'_xlfn.WEB&#83;ERVICE("{URL}")',
    "comment inside the formula": f'_xlfn.WEB<!-- -->SERVICE("{URL}")',
    "cdata split": f'<![CDATA[_xlfn.WEB]]><![CDATA[SERVICE("{URL}")]]>',
    "plain": f'_xlfn.WEBSERVICE("{URL}")',
}


@pytest.mark.parametrize("case", sorted(HIDDEN))
def test_a_sheet_part_not_named_xml_and_led_by_whitespace_is_read_as_xml(tmp_path, case):
    how, reach = file_reach(_sheet_as(tmp_path, HIDDEN[case], new_name="xl/worksheets/sheet1.dat"))
    assert _found(reach), (how, reach)
    assert any(f["where"] == "Rates!B1" for f in reach["formulas"]), reach
    assert not reach.get("unread"), reach


def test_the_same_sheet_part_named_xml_is_read_as_before(tmp_path):
    how, reach = file_reach(_sheet_as(tmp_path, HIDDEN["character reference"], new_name=SHEET))
    assert _found(reach) and not reach.get("unread"), reach


def test_a_part_of_an_xml_type_with_no_name_extension_is_read_as_xml(tmp_path):
    how, reach = file_reach(_sheet_as(tmp_path, HIDDEN["comment inside the formula"],
                                      new_name="xl/worksheets/sheet1"))
    assert _found(reach), reach


def test_a_sheet_part_given_a_picture_type_is_refused_and_its_call_still_named(tmp_path):
    p = _sheet_as(tmp_path, HIDDEN["plain"], new_name="xl/worksheets/sheet1.dat", content_type="image/png")
    how, reach = file_reach(p)
    assert _unread(reach, "is loaded as a sheet"), reach
    assert _found(reach), reach


def test_any_part_a_relationship_loads_as_xml_is_read_as_xml_whatever_its_type(tmp_path):
    # Excel loads a part through the relationship that reaches it, so a data connection stored as a "picture"
    # is still a data connection: listed, and the file refused as unread
    sml = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/connections"
    out = []
    for n, d in _items(tmp_path):
        if n == "xl/_rels/workbook.xml.rels":
            d = d.replace(b"</Relationships>",
                          f'<Relationship Id="rIdC9" Type="{rel}" Target="conn.dat"/></Relationships>'.encode())
        elif n == "[Content_Types].xml":
            d = d.replace(b"</Types>", b'<Override PartName="/xl/conn.dat" ContentType="image/png"/></Types>')
        out.append((n, d))
    out.append(("xl/conn.dat", " " * 5000 + f'<connections xmlns="{sml}"><connection id="1" name="made-up-feed"'
                f' type="4"/></connections>'))
    how, reach = file_reach(_write(tmp_path, out))
    assert _unread(reach, "the part xl/conn.dat is loaded as a part of the type connections"), reach
    assert "made-up-feed" in reach["connections"], reach


def test_a_sheet_part_of_an_unknown_type_is_refused(tmp_path):
    p = _sheet_as(tmp_path, HIDDEN["comment inside the formula"], new_name="xl/worksheets/sheet1.dat",
                  content_type="application/x-made-up")
    how, reach = file_reach(p)
    assert _unread(reach, "which this check does not know"), reach


def test_a_sheet_part_missing_from_the_content_types_is_refused(tmp_path):
    p = _sheet_as(tmp_path, HIDDEN["comment inside the formula"], new_name="xl/worksheets/sheet1.dat", listed=False)
    how, reach = file_reach(p)
    assert _unread(reach, "has no content type"), reach


def test_a_sheet_part_that_does_not_decode_is_refused(tmp_path):
    p = _sheet_as(tmp_path, HIDDEN["plain"], new_name="xl/worksheets/sheet1.dat",
                  encode=lambda t: t.encode("utf-8").replace(b"made-up", b"made\xff-up"))
    how, reach = file_reach(p)
    assert _unread(reach, "does not decode"), reach


def test_a_sheet_part_stored_in_a_legacy_encoding_is_refused_not_skipped(tmp_path):
    p = _sheet_as(tmp_path, HIDDEN["plain"] + "é", new_name="xl/worksheets/sheet1.dat",
                  encode=lambda t: t.encode("latin-1"))
    how, reach = file_reach(p)
    assert reach and reach.get("unread"), reach


@pytest.mark.parametrize("spelling", [f'_xlfn.WEBSERVICE("{URL}")', f'_xlfn.WEB&#83;ERVICE("{URL}")'])
def test_a_call_in_the_plain_bytes_of_a_binary_part_is_found(tmp_path, spelling):
    data = b"\x89PNG\r\n\x1a\n\x00\x00\xff\xfe" + spelling.encode("ascii") + b"\x00\xff\x80"
    how, reach = file_reach(_with_member(tmp_path, "xl/media/image1.png", data, ("png", "image/png")))
    assert _found(reach), reach


def test_a_binary_part_with_no_call_is_not_refused(tmp_path):
    data = b"\x89PNG\r\n\x1a\n\x00\x00\xff\xfe made-up image (1) \x00\xff\x80" * 200
    how, reach = file_reach(_with_member(tmp_path, "xl/media/image1.png", data, ("png", "image/png")))
    assert reach is None, reach


def test_printer_settings_and_a_vba_project_are_binary_parts_not_refusals(tmp_path):
    out = []
    for n, d in _items(tmp_path):
        if n == "[Content_Types].xml":
            d = d.decode("utf-8").replace("<Default ", (
                '<Default Extension="bin" ContentType="application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.printerSettings"/><Override PartName="/xl/vbaProject.bin" '
                'ContentType="application/vnd.ms-office.vbaProject"/><Default '), 1)
        out.append((n, d))
    out += [("xl/printerSettings/printerSettings1.bin", b"\x00\x01\xfe\xff" * 50),
            ("xl/vbaProject.bin", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00\xff" * 100)]
    how, reach = file_reach(_write(tmp_path, out))
    assert reach is None, reach


def test_a_part_of_an_unknown_content_type_is_refused(tmp_path):
    p = _with_member(tmp_path, "xl/extra.foo", b"made-up", ("foo", "application/x-made-up"))
    how, reach = file_reach(p)
    assert _unread(reach, "the part xl/extra.foo has the content type application/x-made-up"), reach


def test_a_part_with_no_content_type_is_refused(tmp_path):
    p = _with_member(tmp_path, "xl/extra.zzz", b"made-up")
    how, reach = file_reach(p)
    assert _unread(reach, "the part xl/extra.zzz has no content type"), reach


def test_a_part_with_no_content_type_still_has_its_call_named(tmp_path):
    p = _with_member(tmp_path, "xl/extra.zzz", f'_xlfn.WEBSERVICE("{URL}")'.encode())
    how, reach = file_reach(p)
    assert _unread(reach, "has no content type") and _found(reach), reach


def test_a_folder_entry_that_holds_data_is_refused(tmp_path):
    out = [(n, d) for n, d in _items(tmp_path)]
    p = tmp_path / "book.xlsx"
    with zipfile.ZipFile(p, "w") as z:
        for n, d in out:
            z.writestr(n, d)
        z.writestr(zipfile.ZipInfo("xl/hidden/"), b"made-up data")
    how, reach = file_reach(str(p))
    assert _unread(reach, "is named like a folder and holds data"), reach


def test_a_binary_workbook_is_known_by_its_content_type(tmp_path):
    out = []
    for n, d in _items(tmp_path):
        if n == "[Content_Types].xml":
            d = d.replace(b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                          b"application/vnd.ms-excel.sheet.binary.macroEnabled.main")
        out.append((n, d))
    assert file_reach(_write(tmp_path, out))[0] == READ_IN_EXCEL


def test_a_workbook_part_of_a_picture_type_is_refused(tmp_path):
    out = []
    for n, d in _items(tmp_path):
        if n == "[Content_Types].xml":
            d = d.replace(b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml", b"image/png")
        out.append((n, d))
    how, reach = file_reach(_write(tmp_path, out))
    assert _unread(reach, "is loaded as the workbook"), reach


def test_a_content_type_list_that_names_a_part_twice_is_refused(tmp_path):
    out = []
    for n, d in _items(tmp_path):
        if n == "[Content_Types].xml":
            d = d.decode("utf-8").replace("<Default ", '<Override PartName="/xl/worksheets/sheet1.xml" '
                                          'ContentType="image/png"/><Default ', 1)
        out.append((n, d))
    how, reach = file_reach(_write(tmp_path, out))
    assert _unread(reach, "two content types"), reach


def test_an_opendocument_member_that_is_not_text_is_still_searched(tmp_path):
    p = tmp_path / "book.ods"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        z.writestr("content.xml", '<?xml version="1.0"?><office:document-content xmlns:office="urn:oasis:names:'
                   'tc:opendocument:xmlns:office:1.0"/>')
        z.writestr("Pictures/a.bin", b"\x89PNG\x00\x80" + b'of:=COM.MICROSOFT.WEBSERVICE("' + URL.encode() + b'")')
    how, reach = file_reach(str(p))
    assert _found(reach), reach


@pytest.mark.parametrize("lead", [5000, 20000])
def test_a_text_file_led_by_whitespace_that_is_xml_goes_through_the_parser(tmp_path, lead):
    p = tmp_path / "book.xml"
    p.write_text(" " * lead + '<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"><Bad></Workbook>',
                 encoding="utf-8")
    how, reach = file_reach(str(p))
    assert _unread(reach, "not well-formed"), reach


class _Cell:
    def __init__(self, i, formula):
        self.Address, self.Formula2, self.Formula = f"$A${i}", formula, formula


class _Used:
    """Excel's Find over formulas, made up: every cell holds the name as text, none calls it."""

    def __init__(self, n):
        self.cells = [_Cell(i + 1, '="WEBSERVICE(" & "made-up"') for i in range(n)]

    def Find(self, What, **kw):  # noqa: N802, N803
        return self.cells[0] if What.startswith("WEBSERVICE") else None

    def FindNext(self, cell):  # noqa: N802
        i = self.cells.index(cell) + 1
        return self.cells[i % len(self.cells)]


class _Sheet:
    def __init__(self, n):
        self.Name, self.UsedRange = "Rates", _Used(n)

        class _Fcs:
            Count = 0
        self.Cells = type("C", (), {"FormatConditions": _Fcs()})()


class _Book:
    def __init__(self, n):
        self.Worksheets, self.Names = [_Sheet(n)], []
        self.Connections = type("K", (), {"Count": 0})()


@pytest.mark.parametrize("n, refused", [(3, False), (12, True)])
def test_excel_s_look_counts_a_name_it_cannot_confirm_past_its_limit(monkeypatch, n, refused):
    from xlsx_mcp.com import session
    from xlsx_mcp.core import netreach
    monkeypatch.setattr(session, "_FIND_LIMIT", 10)
    reach = session._reach_of_open_workbook(_Book(n), netreach)
    assert bool(reach) is refused, reach
    if refused:
        assert reach["formulas"][0]["functions"] == ["WEBSERVICE"], reach


def test_a_file_that_is_there_but_cannot_be_read_is_refused(tmp_path, monkeypatch):
    from xlsx_mcp.core import netreach
    p = _write(tmp_path, _items(tmp_path))

    def locked(*a, **kw):
        raise PermissionError(13, "made-up lock")
    monkeypatch.setattr(netreach, "open", locked, raising=False)
    how, reach = file_reach(p)
    assert reach and reach.get("unread"), reach


def test_a_missing_file_is_left_to_the_open_that_reports_it(tmp_path):
    assert file_reach(str(tmp_path / "made-up-missing.xlsx"))[1] is None
