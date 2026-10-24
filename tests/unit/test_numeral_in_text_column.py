"""A numeral written into a column or row formatted as Text is text on purpose.

Formatting a whole column as Text is how a sheet marks an ID or code column.
The save gives a new cell there the column's Text format (its row's format
first, when the row has one), as Excel does for a value typed there. The
"now holds the TEXT ... Pass 250 as a number" warning used to fire anyway,
which pointed a model the wrong way. A cell with its own format keeps it,
so a General cell inside a Text column is still warned about.
"""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.ops import cells


def _book(tmp_path):
    p = tmp_path / "ids.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Name", "Amount"
    for r in range(2, 5):
        ws.cell(r, 1, f"n{r}")
        ws.cell(r, 2, r * 100)
    ws.column_dimensions["F"].number_format = "@"   # one column as Text
    ws.column_dimensions.group("H", "J")             # three columns in one <col>
    ws.column_dimensions["H"].number_format = "@"
    ws.row_dimensions[20].number_format = "@"        # one row as Text
    ws.row_dimensions[30].number_format = "0.00"     # a row with a number format
    ws["F3"] = "x"                                   # a cell of its own in the Text column
    ws["F3"].number_format = "General"
    wb.save(p)
    return str(p)


def _text_warnings(res):
    return [w for w in (res.get("warnings") or []) if "TEXT" in str(w)]


def _stored(p, cell):
    return openpyxl.load_workbook(p).active[cell]


@pytest.mark.parametrize("cell", ["F7", "I9", "H20", "Z20"])
def test_a_numeral_in_a_text_column_or_row_is_quiet_and_stays_text(tmp_path, cell):
    p = _book(tmp_path)
    res = cells.set_cell(p, {"cell": cell}, "250", backup=False)
    assert res.get("ok")
    assert not _text_warnings(res), _text_warnings(res)
    stored = _stored(p, cell)
    assert stored.value == "250" and stored.number_format == "@"


def test_write_range_into_a_text_column_is_quiet(tmp_path):
    p = _book(tmp_path)
    res = cells.write_range(p, {"cell": "F10"}, [["101"], ["102"], ["103"]], backup=False)
    assert res.get("ok")
    assert not _text_warnings(res), _text_warnings(res)
    assert [_stored(p, f"F{r}").number_format for r in (10, 11, 12)] == ["@", "@", "@"]


def test_a_cell_with_its_own_format_in_a_text_column_still_warns(tmp_path):
    p = _book(tmp_path)
    res = cells.set_cell(p, {"cell": "F3"}, "250", backup=False)
    assert _text_warnings(res)


def test_the_rows_format_comes_before_the_columns(tmp_path):
    """Excel gives a new cell its row's format when the row has one, so a
    numeral in a number-formatted row of a Text column is not text on
    purpose."""
    p = _book(tmp_path)
    res = cells.set_cell(p, {"cell": "F30"}, "250", backup=False)
    assert _stored(p, "F30").number_format == "0.00"
    assert _text_warnings(res)


def test_a_general_cell_beside_numbers_still_warns(tmp_path):
    p = _book(tmp_path)
    res = cells.set_cell(p, {"cell": "B6"}, "250", backup=False)
    assert _text_warnings(res)
