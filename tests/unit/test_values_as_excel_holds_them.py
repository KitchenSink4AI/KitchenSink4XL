"""Values the way Excel holds them.

  * a formula whose cached result is the empty string is PRESENT, not
    absent: Excel stores <c t="str"><f>..</f><v/></c>, and every reader
    used to call it "no cached value" and send the caller to recalculate
    forever, even after Excel's own CalculateFull.
"""

from __future__ import annotations

import re
import shutil
import zipfile

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp import server as srv
from xlsx_mcp.com import session as com_session

_COM = com_session.com_available()[0]


def _excel_form_book(p) -> str:
    """A workbook holding what Excel writes after a calculation: A1 is
    =IF(B1="","",B1) with the empty-string result, A2 a real number result,
    A3 a formula with no cached value at all."""
    p = str(p)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = '=IF(B1="","",B1)'
    ws["A2"] = "=1+1"
    ws["A3"] = "=2+2"
    wb.save(p)
    tmp = p + ".tmp"
    with zipfile.ZipFile(p) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                x = data.decode()
                x = re.sub(r'<c r="A1"[^>]*>.*?</c>',
                           '<c r="A1" t="str"><f>IF(B1="","",B1)</f><v/></c>', x)
                x = re.sub(r'<c r="A2"[^>]*>.*?</c>', '<c r="A2"><f>1+1</f><v>2</v></c>', x)
                x = re.sub(r'<c r="A3"[^>]*>.*?</c>', '<c r="A3"><f>2+2</f></c>', x)
                data = x.encode()
            zout.writestr(item, data)
    shutil.move(tmp, p)
    return p


def _call(tool, op, **args):
    payload, err, _ = fh.start().call(tool, {"op": op, "args": args})
    assert not err, payload
    return payload


def _label(payload, cell):
    for label, cells in (payload.get("labels") or {}).items():
        if cell in cells:
            return label
    return payload.get("labels_default")


def test_x11_an_empty_text_result_reads_as_cached(tmp_path):
    p = _excel_form_book(tmp_path / "e.xlsx")
    out = _call("xl_read", "read_range", path=p, location="A1:A3", values="cached")
    assert out["values"][0][0] == ""
    assert _label(out, "A1") == "cached"
    assert _label(out, "A2") == "cached"
    assert _label(out, "A3") == "absent"  # the real absent case still says so


def test_x11_the_formula_audit_does_not_count_it_missing(tmp_path):
    p = _excel_form_book(tmp_path / "e.xlsx")
    out = _call("xl_read", "audit_formulas", path=p)
    missing = out["missing_cached_values"]
    cells = [m.get("cell") for m in missing.get("items", missing)] \
        if isinstance(missing, dict) else [m.get("cell") for m in missing]
    assert "A3" in str(cells) and "A1" not in str(cells), missing


@pytest.mark.com
@pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool)")
@pytest.mark.timeout(300)
def test_x11_after_excels_own_calculation_nothing_is_absent(tmp_path):
    p = str(tmp_path / "x.xlsx")
    srv.create_workbook(p)
    srv.write_range(p, "A1", [['=IF(B1="","",B1)'], ['=B1&""']])
    srv.recalculate(p, engine="com")
    out = _call("xl_read", "read_range", path=p, location="A1:A2", values="cached")
    assert out["values"] == [[""], [""]]
    assert "absent" not in (out.get("labels") or {})
    assert "warning" not in out
