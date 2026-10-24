"""Data in and out the way Excel's own import and export treat
it.

  * a UTF-8 byte-order mark is not data
  * codes stay codes: leading zeros, E-notation text and long digit
    strings import as text (the whole column, reported); column_types
    chooses per column
  * a semicolon CSV is detected, its decimal commas read as numbers
  * real dates: {"date": ...} and parse_dates
  * text that starts like a formula: {"text": ...} and Excel's apostrophe
  * a big write answers with a summary, not every cell
  * the write and scan ceilings no longer stop a 50,000-row job
  * values="displayed", Excel's TRUE/FALSE and dates, the BOM
"""

from __future__ import annotations

import csv
import io
import json

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp import server as srv
from xlsx_mcp.com import session as com_session
from xlsx_mcp.core import numfmt

_COM = com_session.com_available()[0]


@pytest.fixture
def book(tmp_path):
    p = str(tmp_path / "i.xlsx")
    srv.create_workbook(p)
    return p


def _call(tool, op, **args):
    payload, err, _ = fh.start().call(tool, {"op": op, "args": args})
    assert not err, payload
    return payload


CUSTOMERS = ("﻿CustomerID,Name,Zip,Code,Phone,Account,Amount\n"
             "000123,Zoë Müller,02134,1E5,555.0106,123456789012345678,\"1,234.56\"\n"
             "000124,Ann Lee,10001,2E7,555.0107,123456789012345679,50\n")


def test_x13_the_bom_is_not_part_of_the_first_header(book, tmp_path):
    src = tmp_path / "c.csv"
    src.write_bytes(CUSTOMERS.encode("utf-8"))
    _call("xl_edit", "import_data", path=book, source_file=str(src))
    assert openpyxl.load_workbook(book).active["A1"].value == "CustomerID"
    out = _call("xl_read", "query_range", path=book,
                where=[{"column": "CustomerID", "op": "eq", "value": "000123"}])
    assert out["matched"] == 1


def test_x14_codes_stay_codes_and_the_columns_are_reported(book):
    out = _call("xl_edit", "import_data", path=book, source=CUSTOMERS)
    ws = openpyxl.load_workbook(book).active
    assert [ws[c].value for c in ("A2", "C2", "D2", "F2")] == [
        "000123", "02134", "1E5", "123456789012345678"]
    assert ws["C3"].value == "10001"  # the whole column is text, not mixed
    assert ws["E2"].value == 555.0106 and ws["G2"].value == 1234.56
    kept = {k["column"]: k["reason"] for k in out["changed"]["imported"]["columns_kept_as_text"]}
    assert kept == {"CustomerID": "a leading zero", "Zip": "a leading zero",
                    "Code": "E-notation text", "Account": "more than 15 digits"}


def test_x14_column_types_choose(book):
    out = _call("xl_edit", "import_data", path=book, source=CUSTOMERS,
                column_types={"Phone": "text", "Zip": "number", "B": "text"})
    ws = openpyxl.load_workbook(book).active
    assert ws["E2"].value == "555.0106" and ws["C2"].value == 2134
    payload, err, _ = fh.start().call("xl_edit", {"op": "import_data", "args": {
        "path": book, "source": CUSTOMERS, "column_types": {"Nope": "text"}}})
    assert err and "not one of the imported columns" in payload["error"]["message"]


def test_x64_a_semicolon_file_with_decimal_commas(book):
    out = _call("xl_edit", "import_data", path=book,
                source="Item;Price;Qty\nA;1,5;3\nB;1.234,75;4\n")
    ws = openpyxl.load_workbook(book).active
    assert [ws[c].value for c in ("A2", "B2", "B3", "C3")] == ["A", 1.5, 1234.75, 4]
    info = out["changed"]["imported"]
    assert info["delimiter"] == ";" and info["decimal_comma"] is True


def test_x24_typed_dates_and_parse_dates(book):
    _call("xl_edit", "write_range", path=book, location="A1", data=[
        [{"date": "2026-03-24"}, {"date": "2026-03-24T09:30"}, "2026-03-25"]])
    _call("xl_edit", "write_range", path=book, location="A2", data=[["2026-03-26"]],
          parse_dates=True)
    ws = openpyxl.load_workbook(book).active
    assert ws["A1"].is_date and ws["A1"].number_format == "yyyy-mm-dd"
    assert ws["B1"].is_date and ws["B1"].number_format == "yyyy-mm-dd hh:mm"
    assert ws["C1"].value == "2026-03-25"  # without parse_dates, text stays text
    assert ws["A2"].is_date
    out = _call("xl_read", "read_range", path=book, location="A1:B1")
    assert out["values"] == [["2026-03-24", "2026-03-24T09:30:00"]]


def test_x24_import_reads_iso_date_columns_with_parse_dates(book):
    out = _call("xl_edit", "import_data", path=book,
                source="When,Qty\n2026-01-15,3\n2026-01-16,4\n", parse_dates=True)
    ws = openpyxl.load_workbook(book).active
    assert ws["A2"].is_date and out["changed"]["imported"]["date_columns"] == ["When"]


def test_x24_a_malformed_date_refuses(book):
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_cell", "args": {
        "path": book, "location": "A1", "value": {"date": "24/03/2026"}}})
    assert err and "ISO" in payload["error"]["message"]


@pytest.mark.parametrize("value", [{"text": "=HYPERLINK(\"x\")"}, "'=HYPERLINK(\"x\")"])
def test_x17_text_that_starts_like_a_formula(value, book):
    _call("xl_edit", "set_cell", path=book, location="A1", value=value)
    c = openpyxl.load_workbook(book).active["A1"]
    assert c.value == '=HYPERLINK("x")' and c.data_type == "s" and c.quotePrefix


def test_x23_a_big_write_answers_with_a_summary(book):
    data = [[r * 10 + c for c in range(7)] for r in range(200)]
    out = _call("xl_edit", "write_range", path=book, location="A1", data=data)
    assert "cells" not in out["changed"]
    summ = out["changed"]["cells_summary"]
    assert summ["count"] == 1400 and summ["ranges"] == ["Sheet1!A1:G200"]
    assert len(json.dumps(out)) < 2000
    full = _call("xl_edit", "write_range", path=book, location="A1", data=data,
                 echo_cells=True)
    assert len(full["changed"]["cells"]) == 1400


@pytest.mark.timeout(600)
def test_x54_a_fifty_thousand_row_import_and_its_summary(book):
    rows = ["Day,Region,Units"] + [f"{i},{'NS'[i % 2]},{i % 7}" for i in range(50_000)]
    out = _call("xl_edit", "import_data", path=book, source="\n".join(rows))
    assert out["changed"]["imported"]["rows"] == 50_001
    agg = _call("xl_read", "query_range", path=book, group_by="Region",
                aggregate=[{"column": "Units", "func": "sum"}])
    assert {g["group"]["Region"] for g in agg["groups"]} == {"N", "S"}


# -------------------------------------------------------------- export


def test_numfmt_matches_excels_text_for_common_formats():
    # Excel's own Range.Text for each pair (fix4-scratch/oracle_numfmt.py).
    cases = [(2134, "00000", "02134"), (1234.5, "#,##0.00", "1,234.50"),
             (-1234.5, "#,##0.00;(#,##0.00)", "(1,234.50)"), (0.125, "0.0%", "12.5%"),
             (-5, '"$"#,##0.00', "-$5.00"), (12, "0.#", "12."),
             (0.000123, "0.00E+00", "1.23E-04"), (46036, "d-mmm-yy", "14-Jan-26"),
             (46036.75, "h:mm AM/PM", "6:00 PM"), (46036.5208333, "hh:mm:ss", "12:30:00"),
             (1234.5, "[$€-407] #,##0.00", "€ 1,234.50"), (5, '0 "kg"', "5 kg")]
    for value, fmt, excel in cases:
        assert numfmt.display(value, fmt) == excel, (value, fmt)


def _formatted(p) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Zip", "Amount", "When", "Paid"])
    ws.append([2134, 1234.5, None, True])
    ws["A2"].number_format = "00000"
    ws["B2"].number_format = "#,##0.00"
    import datetime
    ws["C2"] = datetime.datetime(2026, 1, 15)
    ws["C2"].number_format = "yyyy-mm-dd"
    wb.save(p)
    return str(p)


def test_x16_raw_export_writes_true_false_and_bare_dates(tmp_path):
    p = _formatted(tmp_path / "f.xlsx")
    out = _call("xl_workbook", "export_range", path=p, fmt="csv")
    assert out["content"].splitlines()[1] == "2134,1234.5,2026-01-15,TRUE"


def test_x16_x19_displayed_is_what_the_cells_show(tmp_path):
    p = _formatted(tmp_path / "f.xlsx")
    out = _call("xl_workbook", "export_range", path=p, fmt="csv", values="displayed")
    assert list(csv.reader(io.StringIO(out["content"])))[1] == [
        "02134", "1,234.50", "2026-01-15", "TRUE"]
    read = _call("xl_read", "read_range", path=p, location="A2", values="displayed")
    assert read["values"] == [["02134"]]


def test_x16_a_csv_file_starts_with_excels_bom_unless_asked_not_to(tmp_path):
    p = _formatted(tmp_path / "f.xlsx")
    a, b = tmp_path / "a.csv", tmp_path / "b.csv"
    _call("xl_workbook", "export_range", path=p, fmt="csv", out_file=str(a))
    _call("xl_workbook", "export_range", path=p, fmt="csv", out_file=str(b), bom=False)
    assert a.read_bytes()[:3] == b"\xef\xbb\xbf" and b.read_bytes()[:3] != b"\xef\xbb\xbf"


@pytest.mark.com
@pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool)")
@pytest.mark.timeout(300)
def test_x56_displayed_export_matches_excels_own_csv(tmp_path):
    """Excel's own CSV (com_convert_format) and export_range values=
    'displayed' agree field for field."""
    p = _formatted(tmp_path / "f.xlsx")
    excel_csv = tmp_path / "excel.csv"
    srv.com_convert_format(p, str(excel_csv), format="csv")
    theirs = list(csv.reader(io.StringIO(excel_csv.read_text(encoding="utf-8-sig"))))
    ours = list(csv.reader(io.StringIO(_call(
        "xl_workbook", "export_range", path=p, fmt="csv", values="displayed")["content"])))
    assert ours == theirs


@pytest.mark.com
@pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool)")
@pytest.mark.timeout(300)
def test_x24_excel_reads_a_written_date_as_a_date(book):
    _call("xl_edit", "write_range", path=book, location="A1", data=[
        [{"date": "2026-03-24"}, "=A1+1"]])
    srv.recalculate(book, engine="com")
    out = _call("xl_read", "read_range", path=book, location="A1:B1", values="displayed")
    # Excel computed A1+1 from a real date (46106 = 2026-03-25); B1 keeps
    # General, as a formula written to a file does.
    assert out["values"] == [["2026-03-24", "46106"]]
