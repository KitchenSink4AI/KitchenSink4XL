"""The package checks decide from the XML, not from how the XML is spelled.

Three checks used to read package parts with regular expressions, so a valid
file that spelled the same XML another way slipped past them: single-quoted
attributes, a namespace prefix on the spreadsheet elements, UTF-16 with a byte
order mark, a character reference inside a value. Each check below is run on
the SAME made-up workbook respelled each of those ways; the answer must not
change. (Every workbook is built by the test; nothing is a real document.)

- Content hidden under a merged range (a named loss): an edit that rewrites the
  sheet must be refused without allow_loss, never answer ok and delete it.
- Strict Open XML: refused by name however its root relationship and workbook
  part are spelled, and wherever the root relationship points.
- A table declared as a query table: refused when its query table part is gone,
  and also when the relationship is external, of another type, or aimed at a
  part that is not a query table.
"""
from __future__ import annotations

import re
import zipfile

import openpyxl
import pytest
from lxml import etree

from xlsx_mcp.core import hazard, verify

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
STRICT_MAIN = "http://purl.oclc.org/ooxml/spreadsheetml/main"
STRICT_REL = "http://purl.oclc.org/ooxml/officeDocument/relationships"
TRANS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_QT = TRANS_REL + "/queryTable"
CT_QT = "application/vnd.openxmlformats-officedocument.spreadsheetml.queryTable+xml"

#: Valid ways to spell the same XML. quote: attribute quote character; prefix:
#: bind the namespace to a prefix instead of the default; enc: the encoding.
SPELLINGS = {
    "plain": {},
    "single-quotes": {"quote": "'"},
    "prefixed": {"prefix": "x"},
    "prefixed-single": {"prefix": "s", "quote": "'"},
    "utf16le": {"enc": "utf-16-le"},
    "utf16be": {"enc": "utf-16-be"},
    "utf16le-prefixed-single": {"enc": "utf-16-le", "prefix": "q", "quote": "'"},
    "charref": {"charref": True},
    "charref-single": {"charref": True, "quote": "'"},
}


def _reprefix(el, prefix, ns_map):
    ns, _sep, local = el.tag[1:].partition("}") if el.tag.startswith("{") else ("", "", el.tag)
    tag = "{%s}%s" % (ns, local) if ns in ns_map else el.tag
    new = etree.Element(tag, nsmap={prefix: ns} if ns in ns_map else None)
    for k, v in el.attrib.items():
        new.set(k, v)
    new.text, new.tail = el.text, el.tail
    for child in el:
        new.append(_reprefix(child, prefix, ns_map))
    return new


def respell(xml: bytes, *, quote='"', prefix=None, enc="utf-8", charref=None,
            ns=(NS_MAIN,), refs=()):
    """The same XML document, spelled differently. `refs` are attribute VALUES
    to rewrite with a character reference for their first character."""
    root = etree.fromstring(xml)
    if prefix:
        root = _reprefix(root, prefix, set(ns))
    body = etree.tostring(root, encoding="unicode")
    if quote == "'":
        body = re.sub(r'(\s[\w:.\-]+)="([^"]*)"', r"\1='\2'", body)
    if charref:
        for value in refs:
            body = re.sub(r"([\"'])%s\1" % re.escape(value),
                          lambda m, v=value: m.group(1) + "&#%d;%s" % (ord(v[0]), v[1:])
                          + m.group(1), body)
    declared = {"utf-8": "UTF-8", "utf-16-le": "UTF-16", "utf-16-be": "UTF-16"}[enc]
    text = '<?xml version="1.0" encoding="%s" standalone="yes"?>\n%s' % (declared, body)
    if enc == "utf-16-le":
        return b"\xff\xfe" + text.encode("utf-16-le")
    if enc == "utf-16-be":
        return b"\xfe\xff" + text.encode("utf-16-be")
    return text.encode("utf-8")


def _rewrite(src, dst, edits, drop=(), extra=None):
    """Copy a package, replacing members: edits maps a member name to a
    function (bytes -> bytes); extra adds members."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            if info.filename in drop:
                continue
            data = zin.read(info.filename)
            if info.filename in edits:
                data = edits[info.filename](data)
            zout.writestr(info, data)
        for name, data in (extra or {}).items():
            zout.writestr(name, data)
    return str(dst)


# ---------------------------------------------- content under a merged range

def _merge_base(tmp_path, *, kind="inline"):
    """Data!A1:B1 merged; B1 (not the top-left cell) holds `kind`:
    inline text, a formula, a shared-style number, or nothing."""
    base = tmp_path / "base.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Shown"
    ws["A2"] = 1
    ws["D1"] = 5
    wb.save(base)
    cell = {
        "inline": '<c r="B1" t="inlineStr"><is><t>Hidden under the merge</t></is></c>',
        "formula": '<c r="B1"><f>A2+1</f></c>',
        "number": '<c r="B1"><v>41</v></c>',
        "empty": '<c r="B1"/>',
        "blank-value": '<c r="B1"><v></v></c>',
    }[kind]

    def edit(data):
        text = data.decode("utf-8")
        text = text.replace('<c r="D1"', cell + '<c r="D1"', 1)
        text = text.replace("</sheetData>", "</sheetData><mergeCells count=\"1\">"
                            "<mergeCell ref=\"A1:B1\"/></mergeCells>", 1)
        return text.encode("utf-8")
    return _rewrite(base, tmp_path / "m.xlsx", {"xl/worksheets/sheet1.xml": edit})


def _respell_sheet(path, tmp_path, spelling, name="v.xlsx", refs=("B1", "A1:B1")):
    return _rewrite(path, tmp_path / name, {
        "xl/worksheets/sheet1.xml":
            lambda d: respell(d, refs=refs, **SPELLINGS[spelling])})


@pytest.mark.parametrize("spelling", SPELLINGS)
@pytest.mark.parametrize("kind", ["inline", "formula", "number"])
def test_content_under_a_merge_is_found_however_the_sheet_is_spelled(tmp_path, spelling, kind):
    p = _respell_sheet(_merge_base(tmp_path, kind=kind), tmp_path, spelling)
    assert hazard.MERGED_HIDDEN_KEY in hazard.scan_path(p).lossy_keys


@pytest.mark.parametrize("spelling", SPELLINGS)
@pytest.mark.parametrize("kind", ["empty", "blank-value"])
def test_an_empty_cell_under_a_merge_is_no_hazard_in_any_spelling(tmp_path, spelling, kind):
    p = _respell_sheet(_merge_base(tmp_path, kind=kind), tmp_path, spelling)
    assert hazard.MERGED_HIDDEN_KEY not in hazard.scan_path(p).lossy_keys


@pytest.mark.parametrize("spelling", SPELLINGS)
def test_a_rewrite_that_would_delete_hidden_text_is_refused_in_every_spelling(tmp_path, spelling):
    """The exact failure: insert_cols without allow_loss answered ok and the
    hidden text was gone. It must refuse, and leave the file as it was."""
    from xlsx_mcp.ops import structure
    p = _respell_sheet(_merge_base(tmp_path), tmp_path, spelling)
    before = open(p, "rb").read()
    with pytest.raises(Exception) as exc:
        structure.modify_grid_structure(p, "insert_cols", 3, count=1, sheet="Data",
                                        verify_com=False, backup=False)
    msg = str(exc.value).lower()
    assert "merged" in msg and "allow_loss" in msg
    assert open(p, "rb").read() == before


def test_a_cell_with_no_coordinate_sits_where_the_cell_order_puts_it(tmp_path):
    """<c> without r follows the previous cell in its row, as Excel reads it."""
    base = _merge_base(tmp_path, kind="empty")

    def edit(d):
        t = d.decode("utf-8")
        t = t.replace('<c r="B1"/>', '<c t="inlineStr"><is><t>Implicit B1</t></is></c>')
        t = t.replace('<c r="A1"', '<c r="A1"', 1)
        return t.encode("utf-8")
    p = _rewrite(base, tmp_path / "implicit.xlsx", {"xl/worksheets/sheet1.xml": edit})
    assert hazard.MERGED_HIDDEN_KEY in hazard.scan_path(p).lossy_keys


@pytest.mark.parametrize("spelling", SPELLINGS)
def test_a_hyperlink_and_a_note_under_a_merge_are_found_in_every_spelling(tmp_path, spelling):
    base = _merge_base(tmp_path, kind="empty")
    link = _rewrite(base, tmp_path / "link.xlsx", {
        "xl/worksheets/sheet1.xml": lambda d: d.replace(
            b"</mergeCells>", b'</mergeCells><hyperlinks><hyperlink ref="B1" '
            b'location="Data!A2" display="x"/></hyperlinks>')})
    p = _respell_sheet(link, tmp_path, spelling, name="link_v.xlsx", refs=("B1", "A1:B1"))
    assert hazard.MERGED_HIDDEN_KEY in hazard.scan_path(p).lossy_keys

    comments = respell(
        b'<comments xmlns="' + NS_MAIN.encode() + b'"><authors><author>A</author></authors>'
        b'<commentList><comment ref="B1" authorId="0"><text><r><t>made-up note</t></r></text>'
        b'</comment></commentList></comments>', refs=("B1",), **SPELLINGS[spelling])
    rels = respell(
        b'<Relationships xmlns="' + NS_REL.encode() + b'"><Relationship Id="rId2" Type="'
        + TRANS_REL.encode() + b'/comments" Target="../comments1.xml"/></Relationships>',
        ns=(NS_REL,), **SPELLINGS[spelling])
    note = _rewrite(base, tmp_path / "note_v.xlsx", {
        "[Content_Types].xml": lambda d: d.replace(
            b"</Types>", b'<Override PartName="/xl/comments1.xml" ContentType="application/vnd.'
            b'openxmlformats-officedocument.spreadsheetml.comments+xml"/></Types>')},
        extra={"xl/comments1.xml": comments, "xl/worksheets/_rels/sheet1.xml.rels": rels})
    assert hazard.MERGED_HIDDEN_KEY in hazard.scan_path(note).lossy_keys


def test_a_lookalike_element_in_another_namespace_is_not_a_merge(tmp_path):
    """A mergeCell / c that is not in the spreadsheet namespace means nothing."""
    base = _merge_base(tmp_path)

    def edit(d):
        return d.replace(b"<mergeCell ", b'<mergeCell xmlns="urn:made-up" ')
    p = _rewrite(base, tmp_path / "other.xlsx", {"xl/worksheets/sheet1.xml": edit})
    assert hazard.MERGED_HIDDEN_KEY not in hazard.scan_path(p).lossy_keys


# ------------------------------------------------------------ Strict Open XML

def _strict_base(tmp_path, *, strict=True):
    base = tmp_path / "sbase.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Strict"
    ws["A1"] = "Total"
    ws["B1"] = 41
    wb.save(base)
    if not strict:
        return str(base)

    def to_strict(data):
        return (data.replace(NS_MAIN.encode(), STRICT_MAIN.encode())
                .replace(TRANS_REL.encode(), STRICT_REL.encode()))
    edits = {}
    with zipfile.ZipFile(base) as z:
        for n in z.namelist():
            if n.endswith((".xml", ".rels")) and n != "[Content_Types].xml":
                edits[n] = to_strict
    return _rewrite(base, tmp_path / "strict.xlsx", edits)


def _respell_roots(path, tmp_path, spelling, *, name="sv.xlsx", workbook_ns=STRICT_MAIN):
    """Respell the root relationships and the workbook part."""
    kw = SPELLINGS[spelling]
    return _rewrite(path, tmp_path / name, {
        "_rels/.rels": lambda d: respell(d, ns=(NS_REL,), **kw),
        "xl/workbook.xml": lambda d: respell(d, ns=(workbook_ns,), **kw)})


@pytest.mark.parametrize("spelling", SPELLINGS)
def test_strict_is_recognised_in_every_spelling(tmp_path, spelling):
    p = _respell_roots(_strict_base(tmp_path), tmp_path, spelling)
    assert hazard.is_strict_package(p) is True


@pytest.mark.parametrize("spelling", SPELLINGS)
def test_a_transitional_workbook_is_not_strict_in_any_spelling(tmp_path, spelling):
    plain = _strict_base(tmp_path, strict=False)
    p = _respell_roots(plain, tmp_path, spelling, workbook_ns=NS_MAIN)
    assert hazard.is_strict_package(p) is False


@pytest.mark.parametrize("spelling", ["plain", "single-quotes", "utf16le", "utf16be",
                                      "utf16le-prefixed-single"])
def test_every_read_and_edit_refuses_a_respelled_strict_workbook_by_name(tmp_path, spelling):
    from xlsx_mcp.core.errors import UnsupportedStructure
    from xlsx_mcp.ops import cells, gridio, lifecycle, validation
    p = _respell_roots(_strict_base(tmp_path), tmp_path, spelling)
    before = open(p, "rb").read()
    for call in (
            lambda: gridio.open_wb(p),
            lambda: cells.read_range(p, "A1:B1", sheet="Strict"),
            lambda: cells.set_cell(p, "A1", "x", sheet="Strict"),
            lambda: lifecycle.manage_worksheet(p, "rename", sheet="Strict", new_name="O")):
        with pytest.raises(UnsupportedStructure) as exc:
            call()
        assert "Strict Open XML" in str(exc.value)
    assert open(p, "rb").read() == before
    assert "Strict Open XML" in str(validation._check_structure(p))


@pytest.mark.parametrize("spelling", ["plain", "single-quotes", "utf16le", "utf16be"])
@pytest.mark.parametrize("transitional_type", [False, True])
def test_strict_is_found_through_the_root_relationship_to_a_renamed_workbook(
        tmp_path, spelling, transitional_type):
    """The workbook part is wherever _rels/.rels says; xl/workbook.xml is only
    the default. With a transitional relationship Type only the workbook's own
    namespace says Strict."""
    base = _strict_base(tmp_path)
    kw = SPELLINGS[spelling]

    def rels(d):
        d = d.replace(b"xl/workbook.xml", b"xl/book.xml")
        if transitional_type:
            d = d.replace(STRICT_REL.encode(), TRANS_REL.encode())
        return respell(d, ns=(NS_REL,), **kw)

    def renamed_ct(d):
        return d.replace(b"/xl/workbook.xml", b"/xl/book.xml")
    with zipfile.ZipFile(base) as zin:
        book = zin.read("xl/workbook.xml")
        book_rels = zin.read("xl/_rels/workbook.xml.rels")
    p = _rewrite(base, tmp_path / "renamed.xlsx",
                 {"_rels/.rels": rels, "[Content_Types].xml": renamed_ct},
                 drop=("xl/workbook.xml", "xl/_rels/workbook.xml.rels"),
                 extra={"xl/book.xml": respell(book, ns=(STRICT_MAIN,), **kw),
                        "xl/_rels/book.xml.rels": book_rels})
    assert hazard.is_strict_package(p) is True


# ----------------------------------------------------------- query tables

def _query_base(tmp_path, *, keep_part=True, rel_target="../queryTables/queryTable1.xml",
                rel_type=REL_QT, external=False, name="qt.xlsx"):
    base = tmp_path / ("base_" + name)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feeds"
    ws["A1"] = "Orders feed"
    ws.append([])
    ws.append(["Site", "Pallets", "Shipped"])
    ws.append(["North", 12, 10])
    ws.append(["South", 7, 7])
    from openpyxl.worksheet.table import Table
    ws.add_table(Table(displayName="Orders_feed", ref="A3:C5"))
    wb.save(base)

    def table(d):
        t = d.decode("utf-8")
        t = re.sub(r"<table ", '<table tableType="queryTable" ', t, count=1)
        for i in (1, 2, 3):
            t = t.replace('<tableColumn id="%d" ' % i,
                          '<tableColumn id="%d" uniqueName="%d" queryTableFieldId="%d" ' % (i, i, i))
        return t.encode("utf-8")
    qt = (b'<queryTable xmlns="' + NS_MAIN.encode() + b'" name="ExternalData_1" connectionId="1">'
          b'<queryTableRefresh nextId="4"><queryTableFields count="3">'
          b'<queryTableField id="1" name="Site" tableColumnId="1"/></queryTableFields>'
          b'</queryTableRefresh></queryTable>')
    mode = ' TargetMode="External"' if external else ""
    rels = (b'<Relationships xmlns="' + NS_REL.encode() + b'"><Relationship Id="rId1" Type="'
            + rel_type.encode() + b'" Target="' + rel_target.encode() + b'"' + mode.encode()
            + b"/></Relationships>")
    extra = {"xl/tables/_rels/table1.xml.rels": rels}
    edits = {"xl/tables/table1.xml": table}
    if keep_part:
        extra["xl/queryTables/queryTable1.xml"] = qt
        edits["[Content_Types].xml"] = lambda d: d.replace(
            b"</Types>", b'<Override PartName="/xl/queryTables/queryTable1.xml" ContentType="'
            + CT_QT.encode() + b'"/></Types>')
    return _rewrite(base, tmp_path / name, edits, extra=extra)


def _respell_query(path, tmp_path, spelling, name="qv.xlsx"):
    kw = SPELLINGS[spelling]
    qt_ref = ("queryTable",)
    return _rewrite(path, tmp_path / name, {
        "xl/tables/table1.xml": lambda d: respell(d, refs=qt_ref, **kw),
        "xl/tables/_rels/table1.xml.rels": lambda d: respell(d, ns=(NS_REL,), **kw)})


@pytest.mark.parametrize("spelling", SPELLINGS)
def test_a_missing_query_table_part_is_refused_in_every_spelling(tmp_path, spelling):
    p = _respell_query(_query_base(tmp_path, keep_part=False), tmp_path, spelling)
    why = verify.package_refusals(p)
    assert why and "query table" in why[0].lower() and "Orders_feed" in why[0]


@pytest.mark.parametrize("spelling", SPELLINGS)
def test_a_query_table_with_its_part_is_clean_in_every_spelling(tmp_path, spelling):
    p = _respell_query(_query_base(tmp_path), tmp_path, spelling)
    assert verify.package_refusals(p) == []


@pytest.mark.parametrize("what, kw", [
    ("a relationship aimed at the worksheet", {"rel_target": "../worksheets/sheet1.xml"}),
    ("an external relationship", {"external": True}),
    ("an external relationship to a local-looking target", {"external": True,
                                                            "rel_target": "../queryTables/queryTable1.xml"}),
    ("a relationship of another type", {"rel_type": TRANS_REL + "/hyperlink"}),
    ("a relationship whose target is absent", {"rel_target": "../queryTables/queryTable9.xml"}),
    ("a relationship aimed at another table", {"rel_target": "../tables/table1.xml"}),
])
@pytest.mark.parametrize("spelling", ["plain", "single-quotes", "utf16le", "prefixed"])
def test_a_query_table_relationship_must_be_internal_and_reach_a_query_table(
        tmp_path, what, kw, spelling):
    p = _respell_query(_query_base(tmp_path, **kw), tmp_path, spelling)
    why = verify.package_refusals(p)
    assert why and "query table" in why[0].lower(), what


def test_an_ordinary_table_is_never_a_query_table_orphan(tmp_path):
    base = tmp_path / "ord.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["a", "b"])
    ws.append([1, 2])
    from openpyxl.worksheet.table import Table
    ws.add_table(Table(displayName="Plain_one", ref="A1:B2"))
    wb.save(base)
    for spelling in SPELLINGS:
        p = _rewrite(base, tmp_path / ("o_%s.xlsx" % spelling), {
            "xl/tables/table1.xml": lambda d, s=spelling: respell(d, **SPELLINGS[s])})
        assert verify.package_refusals(p) == []
