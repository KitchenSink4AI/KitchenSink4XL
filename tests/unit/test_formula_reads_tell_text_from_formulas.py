"""Formula reads tell a text cell that starts with "=" from a formula.

read_range and get_table in formula mode returned a text cell such as
=[1]Rates!$B$2 (text, not a formula) exactly like a formula, with no label,
so a caller that read a range and wrote it back with set_formula turned the
text into a live external-link formula. Both reads now list those cells and
say how to write them back. All content is made up.
"""

from __future__ import annotations

import openpyxl
from openpyxl.worksheet.table import Table

from xlsx_mcp.ops import cells, tables

TEXT = "=[1]Rates!$B$2"


def _book(tmp_path):
    p = tmp_path / "text_like_formulas.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Item"
    ws["B1"] = "Note"
    ws["A2"] = 4
    ws["B2"] = "=A2*2"            # a formula
    ws["A3"] = 6
    c = ws["B3"]                  # text that starts with "="
    c.value = TEXT
    c.data_type = "s"
    ws["A4"] = 8
    ws["B4"] = "plain words"
    ws.add_table(Table(displayName="Notes", ref="A1:B4"))
    wb.save(p)
    return str(p)


def test_the_text_cell_is_still_text_in_the_file(tmp_path):
    p = _book(tmp_path)
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["B3"].data_type == "s"
    assert wb["Data"]["B2"].data_type == "f"


def test_read_range_formula_mode_names_text_that_looks_like_a_formula(
        tmp_path):
    p = _book(tmp_path)
    out = cells.read_range(p, "A1:B4", sheet="Data", values="formula")
    assert out["values"][2][1] == TEXT
    assert out["text_not_formulas"] == ["B3"]
    note = out["text_note"]
    assert "set_formula" in note and "text" in note


def test_read_range_without_such_text_adds_nothing(tmp_path):
    p = _book(tmp_path)
    out = cells.read_range(p, "A1:B2", sheet="Data", values="formula")
    assert "text_not_formulas" not in out and "text_note" not in out


def test_get_table_formula_mode_names_text_that_looks_like_a_formula(
        tmp_path):
    p = _book(tmp_path)
    out = tables.get_table(p, "Notes", values="formula")
    assert out["rows"][1][1] == TEXT
    assert out["text_not_formulas"] == ["B3"]
    assert "set_formula" in out["text_note"]
