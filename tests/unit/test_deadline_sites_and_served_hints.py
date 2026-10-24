"""The Excel family surface: the permanent tests for the write-site census,
the deadline contract, the release gate and the served hints. Some mirror
tests the Word editor has; the rest are written for Excel's own behaviour.

Sections:
  * the static write-site census (writes AND deletions) and
    the deadline's contract, COM tier included
  * one daemon thread per call, and no hidden Excel left
    behind when the server exits after a timed-out COM call
  * no platform/WMI call on the orient path
  * the release gate is wired upstream of the PyPI publish
  * a refused owner-test build leaves nothing
  * the hint a caller receives over stdio is the table's
  * every engine string names only real ops, tools, routes
  * partial array writes refuse; whole-array writes work
    (Excel as the oracle)
  * an extensionless PDF output is settled before Excel
  * the backup store by resolved path, and the other guards below
"""

from __future__ import annotations

import ast
import asyncio
import hashlib
import json
import os
import queue
import re
import shutil
import subprocess
import sys
import threading
import time
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp import envelope, registry
from xlsx_mcp import server as srv
from xlsx_mcp.com import session as com_session
from xlsx_mcp.core import deadline, safesave
from xlsx_mcp.core.deadline import Deadline
from xlsx_mcp.settings import Settings
from xlsx_mcp.surface import copy as copy_mod
from xlsx_mcp.surface import dispatch
from xlsx_mcp.surface.groups import TOOL_NAMES

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "src" / "xlsx_mcp"
NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)
_COM = com_session.com_available()[0]


def _sha(p) -> str | None:
    p = Path(p)
    return hashlib.sha256(p.read_bytes()).hexdigest() if p.exists() else None


@pytest.fixture
def xlsx(tmp_path):
    p = tmp_path / "f.xlsx"
    srv.create_workbook(str(p))
    srv.write_range(str(p), "A1", [["Name", "Qty"], ["Alpha", 1], ["Beta", 2]])
    return str(p)


# ======================================================================
# Every final write site commits the deadline first
# ======================================================================

#: Attribute calls that change a file on disk: writes, Excel's own saves
#: and exports, replaces, AND deletions (a
#: purge deleted the backups after answering TIMEOUT "nothing was saved").
WRITE_ATTRS = {"write_text", "write_bytes", "copy2", "copyfile", "copytree",
               "Save", "SaveAs", "SaveCopyAs", "ExportAsFixedFormat", "Export",
               "replace_with_retry", "unlink", "rmtree", "rmdir", "remove",
               "rotate_slots", "_place_onto_slot",
               # Every wb.Save goes through session.save_workbook,
               # so each caller of it is a write site this census checks
               "save_workbook",
               # A user's workbook is put in place by promote
               "promote"}

#: (module path, function) -> why its disk changes are not a caller-visible
#: publish the deadline must gate. Every entry is a temp, lock, cache,
#: journal or backup-rotation file, a change that happens only AFTER the
#: call committed, or a clean-up of the call's own refused work.
SITE_EXEMPT: dict[tuple[str, str], str] = {
    ("com/instances.py", "_save"): "the owned-Excel PID journal in the temp folder, not an output",
    ("com/session.py", "save_workbook"): "the one Excel save, a copy into a private temp folder; "
        "every caller commits before calling it, and this census counts each call as a "
        "write site",
    ("com/session.py", "discard_saved"): "removes the private temp folder of Excel's copy, never a "
        "workbook or an output",
    ("com/session.py", "sweep_stale_save_folders"): "removes private temp save folders that "
        "killed or hour-old saves left, never a workbook or an output",
    ("com/session.py", "_stand_in_ok"): "saves and removes its own check workbook in a private "
        "temp folder, never a user's file or an output",
    ("core/calc.py", "strip_empty_cached_values"): "rewrites the save's own .ks4xl-write temp "
        "payload (package._default_saver) before the commit; the temp is removed on refusal",
    ("core/calc.py", "restore_always_calc_cache"): "rewrites the save's own temp payload before "
        "the commit, as above",
    ("core/calc.py", "preserve_untouched_number_text"): "rewrites the save's own temp payload "
        "before the commit, as above",
    ("core/package.py", "_default_saver"): "writes the save's temp payload; save() commits before "
        "the replace and removes the temp when the commit refuses",
    ("core/surgical.py", "write_package"): "writes the byte-preserving save's temp payload "
        "(package._default_saver); save() commits before the replace",
    ("core/carry.py", "write"): "rewrites the rewrite route's own .ks4xl-write temp payload "
        "before the commit, as strip_empty_cached_values does",
    ("core/dynarray.py", "mark_dynamic"): "marks dynamic-array formulas in the save's own "
        ".ks4xl-write temp payload (package._default_saver) before the commit",
    ("core/package.py", "_silent_remove"): "removes the call's own temp payload",
    ("core/package.py", "_sweep_stale_write_tmps"): "removes only .ks4xl-write-* temps a killed "
        "earlier save left behind, never a workbook or an output",
    ("core/package.py", "_restore_from_backup"): "reached only after the commit, when the "
        "post-promote verify fails and the backup is put back",
    ("core/outguard.py", "_timestamped_backup"): "a NEW .bak copy of a file an authorized "
        "overwrite will replace; it changes and removes nothing, and the output itself is "
        "written only after the commit",
    ("core/safesave.py", "slot_dir"): "the slot folder's breadcrumb, written during a rotation "
        "that follows the commit",
    ("core/safesave.py", "replace_with_retry"): "the replace helper itself; every caller commits "
        "first (this census checks each caller)",
    ("core/safesave.py", "promote"): "the workbook replace helper itself (ReplaceFileW); "
        "every caller commits first (this census checks each caller)",
    ("core/safesave.py", "_carry_identity"): "gives promote's new version the workbook's Mark of "
        "the Web where ReplaceFileW is not supported; reached only inside promote",
    ("core/safesave.py", "_swap"): "promote's ReplaceFileW and its fallback; reached "
        "only inside promote, whose callers commit first",
    ("core/safesave.py", "_journal_open"): "the save journal entry in the server's own folder, "
        "written before a swap; not an output",
    ("core/safesave.py", "_journal_close"): "removes the save journal entry of a finished swap; "
        "not an output",
    ("core/safesave.py", "_remove_quietly"): "put_back_if_aside under the workbook's write lock: "
        "removes what a save that did not finish left, never a call's output",
    ("core/safesave.py", "_land_staged"): "put_back_if_aside under the workbook's write lock: "
        "lands the backup a save that did finish had staged",
    ("core/safesave.py", "_stage"): "prepare_rotation removes its own half-made staged temp when "
        "linking or copying it failed; not an output",
    ("core/access.py", "_folder_takes_new_files"): "a probe file in the workbook's folder, "
        "created and removed at once on a refusal path, to name why the write failed",
    ("core/tablefilter.py", "repair_package"): "rewrites the save's own .ks4xl-write temp payload "
        "(package.save) before the commit; the temp is removed on refusal",
    ("core/safesave.py", "_link_or_copy"): "backup slot rotation, reached only after the commit",
    ("core/safesave.py", "_place_onto_slot"): "backup slot rotation, reached only after the commit",
    ("core/safesave.py", "rotate_slots"): "backup slot rotation; its callers commit first "
        "(lifecycle._rotate_before_clobber)",
    ("core/safesave.py", "commit"): "RotationTicket.commit promotes staged slots, called only "
        "after the call's save committed (package.save, comtier)",
    ("core/safesave.py", "abort"): "removes the call's own staged slot temps",
    ("core/safesave.py", "_publish_lockfile"): "the per-file write lock, not an output",
    ("core/lockrelease.py", "break_stale"): "removes a dead writer's lock file (one winner), "
        "not an output",
    ("core/lockrelease.py", "break_guard"): "removes the guard file it made beside a lockfile while "
        "a dead writer's lock file is broken, not an output",
    ("core/lockrelease.py", "unlink_once"): "removes the per-file write lock's own lock file when "
        "its holder releases it, not an output",
    ("core/update_check.py", "*"): "the update-check cache in the user's cache folder",
    ("core/update_flow.py", "*"): "the update notice's state file and lock, and the downloaded "
        "Desktop bundle, all in the server's own state folder; never a workbook, so no workbook "
        "call's deadline applies",
    ("ops/backups.py", "restore_backup"): "the restore's temp copy inside the slot folder; the "
        "commit precedes its placement onto the workbook",
    ("ops/lifecycle.py", "create_workbook"): "wb.save is handed to _write_atomic as the temp "
        "writer; _write_atomic commits before its replace",
    ("ops/lifecycle.py", "copy_workbook"): "shutil.copy2 is handed to _write_atomic as the temp "
        "writer; _write_atomic commits before its replace",
}


def _write_calls(func: ast.AST):
    """Every place in func that changes a file, called directly or handed to
    a helper (com_retry(wb.Save, ...), _with_retry(shutil.copy2, ...))."""
    for node in ast.walk(func):
        if isinstance(node, ast.Attribute):
            name = node.attr
            owner = ast.unparse(node.value)
            if name == "save":
                # openpyxl Workbook.save; WorkbookPackage.save commits itself.
                if owner in ("wb", "workbook", "book", "self.workbook", "_wb"):
                    yield node, "Workbook.save"
            elif name in ("copy", "move") and owner == "shutil":
                yield node, f"shutil.{name}"
            elif name == "dump" and owner == "json":
                yield node, "json.dump"
            elif name == "replace" and owner == "os":
                yield node, "os.replace"
            elif name in ("remove", "rmdir") and owner != "os":
                continue  # list.remove, set.remove and friends
            elif name in WRITE_ATTRS:
                yield node, name
        elif isinstance(node, ast.Call):
            f = node.func
            if isinstance(f, ast.Name) and f.id in WRITE_ATTRS:
                yield node, f.id  # a helper imported by name
            elif isinstance(f, ast.Name) and f.id == "open" and len(node.args) >= 2:
                mode = node.args[1]
                if isinstance(mode, ast.Constant) and isinstance(mode.value, str) \
                        and any(c in mode.value for c in "wax"):
                    yield node, f"open({mode.value!r})"
            elif isinstance(f, ast.Attribute) and f.attr == "ZipFile" and len(node.args) >= 2:
                mode = node.args[1]
                if isinstance(mode, ast.Constant) and mode.value in ("w", "a", "x"):
                    yield node, f"ZipFile({mode.value!r})"


def _commits_before(func: ast.AST, line: int) -> bool:
    for node in ast.walk(func):
        if isinstance(node, ast.Call) and node.lineno < line:
            if ast.unparse(node.func).endswith("commit_or_raise"):
                return True
    return False


def _functions(tree: ast.AST):
    """Each function with its OWN write calls (a nested def is its own
    function; a lambda belongs to the function it is written in)."""
    for node in ast.walk(tree):
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            nested: set[int] = set()
            for sub in ast.walk(node):
                if sub is not node and isinstance(sub, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    nested.update(id(n) for n in ast.walk(sub))
            yield node, [(c, k) for c, k in _write_calls(node) if id(c) not in nested]


def test_every_final_write_site_commits_the_deadline_first():
    ungated, stale = [], set(SITE_EXEMPT)
    for path in sorted(SRC.rglob("*.py")):
        rel = path.relative_to(SRC).as_posix()
        for func, calls in _functions(ast.parse(path.read_text("utf-8"))):
            for call, kind in calls:
                key = (rel, func.name)
                if key in SITE_EXEMPT or (rel, "*") in SITE_EXEMPT:
                    stale.discard(key)
                    stale.discard((rel, "*"))
                    continue
                if not _commits_before(func, call.lineno):
                    ungated.append(f"{rel}:{call.lineno} {func.name}: {kind}")
    assert ungated == [], (
        "these sites change a file with no commit_or_raise before them in the same function; "
        "commit first (core/deadline.py) or add a reasoned SITE_EXEMPT entry:\n"
        + "\n".join(ungated))
    assert stale == set(), f"SITE_EXEMPT names sites that no longer write: {sorted(stale)}"


def test_the_census_sees_every_deadline_site():
    """The scan is not vacuous: it sees every write site that must commit
    the deadline first, and each of them does."""
    found = {}
    for path in sorted(SRC.rglob("*.py")):
        rel = path.relative_to(SRC).as_posix()
        for func, calls in _functions(ast.parse(path.read_text("utf-8"))):
            for call, _kind in calls:
                found.setdefault((rel, func.name), []).append(_commits_before(func, call.lineno))
    for site in [("ops/backups.py", "create_snapshot"), ("ops/backups.py", "purge_backups"),
                 ("ops/comtier.py", "job"), ("ops/comtier.py", "body"),
                 ("ops/dataio.py", "export_range"), ("ops/objects.py", "manage_image"),
                 ("ops/lifecycle.py", "_write_atomic"), ("core/package.py", "save"),
                 ("core/calc.py", "recalc_via_com")]:
        assert site in found, site
        assert all(found[site]), site


def test_deletions_are_write_sites():
    """An unlink, rmtree, os.remove or rmdir counts."""
    tree = ast.parse("def f(p):\n    import os, shutil\n    p.unlink()\n    shutil.rmtree(p)\n"
                     "    os.remove(p)\n    os.rmdir(p)\n    [].remove(1)\n")
    kinds = [k for _f, calls in _functions(tree) for _c, k in calls]
    assert kinds.count("unlink") == 1 and kinds.count("rmtree") == 1
    assert kinds.count("remove") == 1 and kinds.count("rmdir") == 1


# ------------------------------------------------------------ the contract


def test_a_commit_is_idempotent_even_past_the_bound():
    d = Deadline(0.2)
    assert d.try_commit()
    time.sleep(0.3)
    assert d.try_commit(), "a committed call is past its point of no return"
    assert not d.cancel()


def test_a_cancelled_deadline_refuses_every_commit():
    d = Deadline(30)
    assert d.cancel()
    assert not d.try_commit()
    with deadline.active(d), pytest.raises(Exception) as info:
        deadline.commit_or_raise("the export")
    assert info.value.code == "TIMEOUT" and info.value.details["saved"] is False


def test_a_committed_call_that_never_finishes_is_not_told_nothing_was_saved(monkeypatch):
    monkeypatch.setattr(dispatch, "BACKSTOP_GRACE_S", 0.1)
    monkeypatch.setattr(deadline, "MIN_POST_COMMIT_WAIT_S", 0.5)
    release = threading.Event()
    d = Deadline(0.3)

    def work():
        with deadline.active(d):
            deadline.commit_or_raise("the save")
        release.wait(20)  # Excel hung inside its own save
        return {"saved": True}

    async def go():
        return await dispatch.run_bounded(work, d)
    try:
        with pytest.raises(Exception) as info:
            asyncio.run(go())
    finally:
        release.set()
    assert info.value.code == "TIMEOUT"
    assert info.value.details["saved"] is None
    assert "nothing was saved" not in info.value.message
    assert "check the file" in info.value.message


@pytest.mark.skipif(not _COM, reason="needs Excel (the COM executor refuses without it)")
def test_the_com_worker_thread_runs_under_the_callers_deadline():
    """The one long-lived COM worker thread runs each job under the
    submitting call's context, so Excel's saves commit against its
    deadline. The probe job never acquires an Excel instance."""
    d = Deadline(30)
    with deadline.active(d):
        seen = com_session.run_com("probe", lambda _manager: deadline.current(), timeout=30)
    assert seen is d
    seen_outside = com_session.run_com("probe", lambda _manager: deadline.current(), timeout=30)
    assert seen_outside is not d and not seen_outside.bounded


# ======================================================================
# One daemon thread per call, and no hidden Excel left behind
# ======================================================================


def test_calls_do_not_run_on_the_default_executor():
    """Each call gets its own daemon thread: a stuck call can never
    hold a pool slot another call waits for."""
    names = []

    async def go():
        d = Deadline(10)
        return await dispatch.run_bounded(lambda: names.append(
            (threading.current_thread().name, threading.current_thread().daemon)) or 1, d)
    assert asyncio.run(go()) == 1
    name, daemon = names[0]
    assert name.startswith(dispatch.WORKER_PREFIX) and daemon


def _read_shared(path: Path) -> bytes:
    """Read a file the way that cannot block its writer's replace of it: every
    share flag, delete included (Path.read_text passes none of them, and a
    reader like that makes the server's own os.replace fail for a moment)."""
    import win32file
    handle = win32file.CreateFile(
        str(path), win32file.GENERIC_READ,
        win32file.FILE_SHARE_READ | win32file.FILE_SHARE_WRITE | win32file.FILE_SHARE_DELETE,
        None, win32file.OPEN_EXISTING, 0, None)
    try:
        size = win32file.GetFileSize(handle)
        return bytes(win32file.ReadFile(handle, size)[1]) if size else b""
    finally:
        handle.Close()


def _journaled_excel(journal: Path) -> dict[int, float]:
    """pid -> creation time of every Excel the server that owns this journal
    has started, from that server's own record of them. The journal sits in a
    temp folder only that server uses, so it names no other process's Excel."""
    try:
        records = json.loads(_read_shared(journal).decode("utf-8") or "{}")
    except Exception:  # noqa: BLE001 (not there yet, or caught mid-replace)
        return {}
    found: dict[int, float] = {}
    for pid_text, rec in records.items():
        created = rec.get("create_time") if isinstance(rec, dict) else None
        if isinstance(created, (int, float)) and not isinstance(created, bool):
            found[int(pid_text)] = float(created)
    return found


def _is_that_excel(pid: int, created: float) -> bool:
    """True while this PID is still the EXCEL.EXE that was created at that
    instant: PID and creation time together, so a recycled PID is no match."""
    import psutil
    try:
        proc = psutil.Process(pid)
        return proc.name().upper() == "EXCEL.EXE" and abs(proc.create_time() - created) <= 0.05
    except psutil.Error:
        return False


def _end_that_excel(pid: int, created: float) -> None:
    """End one Excel by its exact PID, re-checking its creation time first."""
    import psutil
    try:
        if _is_that_excel(pid, created):
            psutil.Process(pid).kill()
    except psutil.Error:
        pass


@pytest.mark.com
@pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool)")
@pytest.mark.timeout(600)
def test_r1_server_exit_after_a_timed_out_com_call_leaves_no_excel(tmp_path):
    """A
    call that timed out while its COM job was still inside the private Excel
    keeps running on a daemon thread. When the client then goes away (stdin
    closes, the normal end of a session), the server must not exit leaving
    that hidden EXCEL.EXE holding the workbook. Excel's remedy is the COM
    executor's shutdown at exit: a bounded wait for the in-flight job, then
    every Excel this server's own pool started is ended by owned PID
    (journal, PID plus creation time; never by image name).

    The test names its Excel through the server's own journal, never by
    "whatever EXCEL.EXE appeared meanwhile" (that would also catch another
    session's Excel started in the same minutes, and end it). The server's
    temp folder is the test's, so its PID journal holds only the Excel
    processes this server started, with their creation times. The test ends
    one of them only if it is still there after the server exited, and only
    after checking again that the PID still names the process the journal
    recorded."""
    pytest.importorskip("psutil")
    book = tmp_path / "o.xlsx"
    srv.create_workbook(str(book))
    server_tmp = tmp_path / "server-temp"
    server_tmp.mkdir()
    journal = server_tmp / "ks4xl_pid_journal.json"
    prelude = tmp_path / "slow_open.py"
    prelude.write_text(
        "import time\nfrom xlsx_mcp.com import session as _s\n_o = _s.open_workbook\n"
        "def _slow(*a, **k):\n    wb = _o(*a, **k); time.sleep(20); return wb\n"
        "_s.open_workbook = _slow\n", encoding="utf-8")
    launch = fh.LAUNCH.replace("from ks4xl_community_launch import main",
                               f"exec(open({str(prelude)!r}, encoding='utf-8').read()); "
                               "from ks4xl_community_launch import main")
    cmd = fh.launch_cmd()
    cmd[-1] = launch.format(paths=fh.PKG_SRC)
    proc = subprocess.Popen(cmd, stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                            stderr=subprocess.DEVNULL, creationflags=NO_WINDOW,
                            env=fh._env({"KS4XL_CALL_TIMEOUT_S": "2",
                                         "KS4XL_MAX_CALL_TIMEOUT_S": "2",
                                         "TEMP": str(server_tmp), "TMP": str(server_tmp)}))
    lines: queue.Queue = queue.Queue()
    threading.Thread(target=lambda: [lines.put(ln) for ln in iter(proc.stdout.readline, b"")],
                     daemon=True).start()
    # Every Excel the server journals, noted the moment it appears (a record
    # is forgotten once its process is reaped, so it is read early and often).
    seen: dict[int, float] = {}
    stop = threading.Event()

    def watch_journal():
        while not stop.is_set():
            seen.update(_journaled_excel(journal))
            stop.wait(0.1)

    threading.Thread(target=watch_journal, daemon=True).start()

    def send(msg):
        proc.stdin.write((json.dumps(msg) + "\n").encode())
        proc.stdin.flush()

    try:
        send({"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {
            "protocolVersion": "2025-06-18", "capabilities": {},
            "clientInfo": {"name": "fix1", "version": "0"}}})
        lines.get(timeout=120)
        send({"jsonrpc": "2.0", "method": "notifications/initialized"})
        send({"jsonrpc": "2.0", "id": 2, "method": "tools/call", "params": {
            "name": "xl_app", "arguments": {"op": "com_autofit",
                                            "args": {"path": str(book)}}}})
        answer = json.loads(lines.get(timeout=180))
        assert "TIMEOUT" in json.dumps(answer), answer
        # The hidden Excel of this call is up and inside its (slowed) open.
        # Polled, not asserted at once (on a loaded machine the
        # TIMEOUT can arrive before Excel's process is listed).
        mine: dict[int, float] = {}
        end = time.monotonic() + 60
        while time.monotonic() < end:
            mine = {p: c for p, c in dict(seen).items() if _is_that_excel(p, c)}
            if mine:
                break
            time.sleep(0.5)
        assert mine, "the call never started its private Excel within 60 s (nothing to test)"
        proc.stdin.close()  # the client goes away
        proc.wait(120)
        orphans: dict[int, float] = {}
        end = time.monotonic() + 30  # poll, no fixed margin
        while time.monotonic() < end:
            orphans = {p: c for p, c in dict(seen).items() if _is_that_excel(p, c)}
            if not orphans:
                break
            time.sleep(1)
        assert not orphans, f"hidden Excel left running after the server exited: {sorted(orphans)}"
        # and the workbook is free again: a write lock can be taken
        with open(book, "r+b"):
            pass
    finally:
        stop.set()
        if proc.poll() is None:
            proc.kill()
        for pid, created in dict(seen).items():  # only what this server's journal named
            _end_that_excel(pid, created)


# ======================================================================
# Nothing on the orient path asks WMI
# ======================================================================


@pytest.mark.skipif(sys.platform != "win32", reason="the WMI path exists on Windows only")
def test_get_server_info_never_calls_platform_on_windows(monkeypatch):
    import platform

    def boom(*a, **k):
        raise AssertionError("platform reached: WMI may hang here")
    for name in ("platform", "uname", "win32_ver", "machine", "release", "version", "node",
                 "system"):
        monkeypatch.setattr(platform, name, boom)
    monkeypatch.setattr(com_session, "_AVAILABLE", None)
    line = srv._os_line()
    assert line.startswith("Windows-") and line.count("-") >= 2
    assert com_session.com_available(refresh=True)[0] in (True, False)
    out = srv.get_server_info()
    assert out["platform"].startswith("Windows-")


# ======================================================================
# The public_copy release gate is wired upstream of the PyPI publish
# ======================================================================


def _publish_jobs() -> dict:
    import yaml
    return yaml.safe_load((ROOT / ".github/workflows/publish.yml").read_text("utf-8"))["jobs"]


def _needs(job: dict) -> set[str]:
    needs = job.get("needs") or []
    return {needs} if isinstance(needs, str) else set(needs)


def _upstream(jobs: dict, name: str) -> set[str]:
    seen, todo = set(), [name]
    while todo:
        for dep in _needs(jobs[todo.pop()]):
            if dep not in seen:
                seen.add(dep)
                todo.append(dep)
    return seen


def test_no_release_reaches_pypi_without_the_public_copy_gate():
    jobs = _publish_jobs()
    gates = []
    for name, job in jobs.items():
        for step in job.get("steps", []):
            env = step.get("env") or {}
            run = step.get("run") or ""
            if str(env.get("KS4XL_RELEASE_GATE")) == "1" and "-m public_copy" in run:
                gates.append((name, run))
    assert gates, "no publish.yml job runs the public_copy tests with KS4XL_RELEASE_GATE=1"
    # One job per distribution: EVERY job that calls the PyPI publish action must
    # have the gate upstream of it.
    publishers = [n for n, j in jobs.items() if any(
        str(s.get("uses", "")).startswith("pypa/gh-action-pypi-publish")
        for s in j.get("steps", []))]
    assert publishers, "publish.yml has no PyPI publishing job"
    for name in publishers:
        assert {n for n, _ in gates} & _upstream(jobs, name), (
            f"the PyPI publish job {name} does not depend on the public_copy gate job")
    for _name, run in gates:
        assert "skipped" in run and "tests == 0" in run, (
            "the gate must fail closed: on a skip and on zero tests as well as on a failure")


@pytest.mark.timeout(900)
def test_the_gate_command_runs_the_public_copy_tests_instead_of_skipping_them(tmp_path):
    xml = tmp_path / "gate.xml"
    env = {**os.environ, "KS4XL_RELEASE_GATE": "1", "PYTHONDONTWRITEBYTECODE": "1"}
    subprocess.run([sys.executable, "-X", "utf8", "-m", "pytest", "tests/unit", "-q",
                    "-p", "no:cacheprovider", "-p", "no:randomly", "-m", "public_copy",
                    f"--junitxml={xml}"], cwd=ROOT, env=env, capture_output=True,
                   stdin=subprocess.DEVNULL, creationflags=NO_WINDOW, timeout=850)
    suite = ET.parse(xml).getroot()
    suite = suite if suite.tag == "testsuite" else suite.find("testsuite")
    assert int(suite.get("tests")) >= 7
    assert int(suite.get("skipped")) == 0


# ======================================================================
# A refused owner-test build leaves nothing
# ======================================================================


def tracked_tree_with_an_unfilled_slot(dest: Path) -> Path:
    """The worktree's tracked files under dest, with one copy slot unfilled
    in surface/copy.py (shared with the refused owner-test build test)."""
    names = subprocess.run(["git", "ls-files", "-z"], cwd=ROOT, capture_output=True,
                           check=True, creationflags=NO_WINDOW).stdout.decode().split("\0")
    for name in filter(None, names):
        src = ROOT / name
        if src.is_file():
            target = dest / name
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, target)
    cp = dest / "src" / "xlsx_mcp" / "surface" / "copy.py"
    text = cp.read_text("utf-8")
    anchor = "def slots_in(text: str) -> bool:"
    assert anchor in text
    cp.write_text(text.replace(anchor, 'UNFILLED_FOR_TEST = _slot("a fact")\n\n\n' + anchor),
                  "utf-8")
    return dest


def build_env() -> dict:
    """The adapters come from the pinned source, materialised outside any
    tree, so a copied tree (no git) can build."""
    import adapter_source
    return {**os.environ, "KS4_LICENSE_ADAPTERS": str(adapter_source.materialise())}


@pytest.mark.timeout(900)
def test_a_refused_owner_test_build_removes_what_it_wrote(tmp_path):
    tree = tracked_tree_with_an_unfilled_slot(tmp_path / "tree")
    out = tmp_path / "out"
    res = subprocess.run(
        [sys.executable, "-X", "utf8", str(tree / "scripts" / "build_editions.py"),
         "--edition", "both", "--output", str(out), "--release-date", "2026-10-01",
         "--owner-test"],
        capture_output=True, text=True, stdin=subprocess.DEVNULL, cwd=tree, env=build_env(),
        creationflags=NO_WINDOW, timeout=850)
    said = res.stderr + res.stdout
    assert res.returncode != 0 and "owner-test ship gate refused" in said, said[-2000:]
    assert "unfilled copy slot" in said, said[-2000:]
    left = sorted(p.name for p in out.rglob("*")) if out.exists() else []
    assert left == [], f"a refused build left files behind: {left}"


# ------------------------------------------------ settling at exit


def test_exit_waits_for_finishing_workers_and_stops_uncommitted_ones_at_their_commit():
    """At exit a call that never committed is cancelled (it stops at its
    commit point instead of saving), and every call still running is waited
    for until it finishes."""
    release = threading.Event()
    stuck, saving = Deadline(30), Deadline(30)
    outcome = {}

    def never_committed():
        release.wait(20)
        try:
            with deadline.active(stuck):
                deadline.commit_or_raise("the save")
            outcome["stuck"] = "published"
        except Exception as exc:  # noqa: BLE001
            outcome["stuck"] = getattr(exc, "code", repr(exc))

    def mid_save():
        with deadline.active(saving):
            deadline.commit_or_raise("the save")
        time.sleep(1.0)
        outcome["saving"] = "finished"

    async def go():
        futs = [dispatch._start_worker(never_committed, stuck),
                dispatch._start_worker(mid_save, saving)]
        await asyncio.sleep(0.3)
        threading.Timer(0.5, release.set).start()
        report = await asyncio.to_thread(dispatch.drain_inflight, 30)
        await asyncio.gather(*futs, return_exceptions=True)
        return report
    report = asyncio.run(go())
    assert report == {"cancelled": 1, "waited": 2, "unfinished": 0}, report
    assert outcome == {"saving": "finished", "stuck": "TIMEOUT"}, outcome


def test_the_exit_wait_is_bounded_ends_nothing_and_names_the_pids(monkeypatch, capsys):
    from xlsx_mcp.com import instances
    ended = []
    monkeypatch.setattr(instances, "taskkill", lambda pid: ended.append(pid) or True)
    monkeypatch.setattr(com_session, "owned_excel_pids", lambda: [4242])
    release = threading.Event()
    d = Deadline(30)

    async def go():
        fut = dispatch._start_worker(lambda: release.wait(30), d)
        await asyncio.sleep(0.2)
        t0 = time.monotonic()
        report = await asyncio.to_thread(dispatch.drain_inflight, 0.5)
        took = time.monotonic() - t0
        alive = [t for t in threading.enumerate()
                 if t.name.startswith(dispatch.WORKER_PREFIX) and t.is_alive()]
        release.set()
        await asyncio.gather(fut, return_exceptions=True)
        return report, took, alive
    report, took, alive = asyncio.run(go())
    assert report["unfinished"] == 1 and took < 5, (report, took)
    assert alive, "the drain must not end the worker"
    assert ended == [], "the exit wait ended a process"
    err = capsys.readouterr().err
    assert err.count("\n") == 1 and "4242" in err and "still running" in err


def test_a_normal_exit_with_nothing_in_flight_is_immediate():
    t0 = time.monotonic()
    assert dispatch.drain_inflight(60) == {"cancelled": 0, "waited": 0, "unfinished": 0}
    assert time.monotonic() - t0 < 0.5


def test_no_caller_can_reach_the_exit_wait():
    """Only server.main calls drain_inflight: no op, tool or help path."""
    callers = []
    for path in sorted(SRC.rglob("*.py")):
        tree = ast.parse(path.read_text("utf-8"))
        for fn in ast.walk(tree):
            if isinstance(fn, (ast.FunctionDef, ast.AsyncFunctionDef)):
                for n in ast.walk(fn):
                    if isinstance(n, ast.Call) and ast.unparse(n.func).endswith("drain_inflight"):
                        callers.append((path.relative_to(SRC).as_posix(), fn.name))
    assert callers == [("server.py", "main")], callers
    assert "drain_inflight" not in registry.public_ops()


def test_main_settles_inflight_calls_when_the_transport_closes():
    tree = ast.parse((SRC / "server.py").read_text("utf-8"))
    fn = next(n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef) and n.name == "main")
    calls = [ast.unparse(n.func) for n in ast.walk(fn) if isinstance(n, ast.Call)]
    assert "_dispatch.drain_inflight" in calls


# ======================================================================
# The hint a caller receives is the table's, copy.py first
# ======================================================================


def test_every_hint_a_caller_can_receive_is_the_tables_over_stdio(tmp_path):
    """One stdio session reaching each code the server can be driven to
    without Excel or a timing race; the hint on the wire is the table's,
    or the raise site's own for the unknown op."""
    book = tmp_path / "h.xlsx"
    srv.create_workbook(str(book))
    srv.write_range(str(book), "A1", [["Revenue"], ["Revenue"]])
    junk = tmp_path / "junk.xlsx"
    junk.write_bytes(b"this is not a zip")
    b = str(book)
    cases = [
        # A missing workbook gets
        # the missing-file hint, not "re-read the metadata".
        ("NOT_FOUND", copy_mod.MISSING_FILE_HINT, ("xl_read", {"op": "read_range", "args": {
            "path": str(tmp_path / "no.xlsx"), "location": "A1"}})),
        ("NOT_FOUND", copy_mod.UNKNOWN_OP_HINT, ("xl_read", {"op": "no_such_op", "args": {}})),
        ("NOT_FOUND", copy_mod.UNKNOWN_OP_HINT, ("xl_help", {"op": "no_such_op"})),
        ("AMBIGUOUS_LOCATION", None, ("xl_read", {"op": "read_range", "args": {
            "path": b, "location": {"search": {"text": "Revenue"}}}})),
        ("CONFLICT", None, ("xl_workbook", {"op": "create_workbook", "args": {"path": b}})),
        ("UNSUPPORTED_CONTENT", None, ("xl_read", {"op": "read_range", "args": {
            "path": str(junk), "location": "A1"}})),
        ("BAD_PARAMS", None, ("xl_read", {"op": "read_range", "args": {"path": b, "nope": 1}})),
    ]
    res = fh.stdio_session([c for _, _, c in cases])
    for (code, hint, _call), r in zip(cases, res["results"]):
        err = json.loads(r["text"]).get("error") or {}
        assert r["is_error"], (code, r["text"][:300])
        assert err.get("code") == code, (code, r["text"][:300])
        want = hint or envelope.HINTS[code]
        assert err.get("hint", "").startswith(want), (code, err.get("hint"))
    res = fh.stdio_session([("xl_read", {"op": "read_range", "args": {"path": b,
                                                                     "location": "A1"}})],
                           env={"KS4XL_FORBID": "read_range"})
    err = json.loads(res["results"][0]["text"])["error"]
    assert err["code"] == "NOT_AVAILABLE" and err["hint"] == envelope.HINTS["NOT_AVAILABLE"]


def test_the_shipped_hints_name_each_op_with_its_tool():
    assert "xl_read get_workbook_metadata" in envelope.HINTS["NOT_FOUND"]
    assert "xl_read get_grid_view" in envelope.HINTS["STALE_ANCHOR"]
    assert "xl_read query_range" in envelope.HINTS["LIMIT_EXCEEDED"]


# ======================================================================
# Engine strings name only real ops, tools and routes
# ======================================================================

#: com_* identifiers in engine strings that are not op names: modules,
#: functions, result keys and the pywin32 exception type.
_NOT_OPS = {"com_ground_truth", "com_error", "com_available", "com_availability",
            "com_retry", "com_status", "com_verify_default", "com_run_macro",
            "com_tier", "com_gates", "com_verify"}

#: Routes that do not exist on this server, however they are spelled.
_DEAD_ROUTES = ("live/COM", "live='auto'", "live=\"auto\"", "com pack", "enable_tools",
                "through xl_app, where")


def _served_strings():
    """(where, text) for every string constant in the package except the
    docstrings of modules, classes and private functions, which never reach
    a caller. Public function docstrings in server.py are op help."""
    for path in sorted(SRC.rglob("*.py")):
        rel = path.relative_to(SRC).as_posix()
        tree = ast.parse(path.read_text("utf-8"))
        docs = set()
        for n in ast.walk(tree):
            if isinstance(n, (ast.Module, ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)) \
                    and n.body and isinstance(n.body[0], ast.Expr) \
                    and isinstance(n.body[0].value, ast.Constant):
                served = rel == "server.py" and isinstance(n, ast.FunctionDef) \
                    and not n.name.startswith("_")
                if not served:
                    docs.add(id(n.body[0].value))
        for node in ast.walk(tree):
            if id(node) in docs:
                continue
            if isinstance(node, ast.Constant) and isinstance(node.value, str):
                yield f"{rel}:{node.lineno}", node.value


def test_engine_strings_name_only_real_ops_tools_and_routes():
    """The permanent sweep: every op name in an
    engine string is in the registry, every xl_ tool name is a real tool,
    and no string offers a route this server does not have."""
    ops = set(registry.public_ops())
    assert ops, "registry empty: the sweep would be vacuous"
    com = re.compile(r"\b(com_[a-z][a-z_]*[a-z])\b")
    tool = re.compile(r"\b(xl_[a-z][a-z_]*[a-z])\b")
    from xlsx_mcp.ops.workflows import REMOVED_IN_V2

    bad = []
    for where, text in _served_strings():
        bad += [f"{where}: op {n}" for n in com.findall(text) if n not in ops | _NOT_OPS]
        bad += [f"{where}: tool {n}" for n in tool.findall(text) if n not in TOOL_NAMES]
        # The migrate-from-v1 map names the
        # removed 1.x tools themselves, as rows with no op; that names a
        # route as gone, not as one to take.
        if text in REMOVED_IN_V2:
            continue
        bad += [f"{where}: dead route {r!r}" for r in _DEAD_ROUTES if r in text]
    assert bad == [], bad


def test_the_sweep_catches_the_two_dead_routes(monkeypatch):
    """Not vacuous: the two dead-route sentences, put back, are caught."""
    old = ["Convert the array formula(s) to values first, or edit the workbook through "
           "xl_app, where Excel applies its own array semantics.",
           "b.xlsx is open in Excel (or another process holds it); close it, or use a "
           "live/COM route, then retry"]
    hits = [r for text in old for r in _DEAD_ROUTES if r in text]
    assert len(hits) == 2, hits


# ======================================================================
# A write into part of an array formula refuses; the whole array works
# ======================================================================


def _array_book(p: Path) -> str:
    from openpyxl.worksheet.formula import ArrayFormula
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B1"], ws["B2"] = 1, 2
    ws["A1"] = ArrayFormula("A1:A2", "=B1:B2*2")
    wb.save(p)
    return str(p)


@pytest.mark.parametrize("tool,op,args", [
    ("xl_edit", "set_cell", {"location": "A2", "value": 5}),
    ("xl_edit", "set_cells", {"cells": [{"cell": "A1", "value": 5}]}),
    ("xl_edit", "write_range", {"location": "A2", "data": [[5], [6]]}),
    ("xl_edit", "set_formula", {"location": "A1", "formula": "=1+1"}),
    ("xl_edit", "import_data", {"source": "x\n5", "fmt": "csv", "location": "A2"}),
    ("xl_edit", "apply_edits", {"edits": [{"op": "set_value", "location": "A2", "value": 5}]}),
], ids=["set_cell", "set_cells", "write_range_overlapping", "set_formula", "import_data",
        "apply_edits"])
def test_every_writer_refuses_part_of_an_array(tool, op, args, tmp_path):
    p = _array_book(tmp_path / "arr.xlsx")
    before = _sha(p)
    payload, err, _ = fh.start().call(tool, {"op": op, "args": {"path": p, **args}})
    assert err, f"{op} wrote part of array A1:A2: {payload}"
    msg = payload["error"]["message"]
    assert "A1:A2" in msg and "part of an array" in msg, msg
    assert "xl_app is not a route" in msg
    assert _sha(p) == before


@pytest.mark.parametrize("op,args", [
    ("write_range", {"location": "A1", "data": [[5], [6]]}),
    ("set_cells", {"cells": [{"cell": "A1", "value": 5}, {"cell": "A2", "value": 6}]}),
])
def test_writing_the_whole_array_replaces_it_with_the_values(op, args, tmp_path):
    p = _array_book(tmp_path / "arr.xlsx")
    payload, err, _ = fh.start().call("xl_edit", {"op": op, "args": {"path": p, **args}})
    assert not err, payload
    wb = openpyxl.load_workbook(p)
    assert (wb.active["A1"].value, wb.active["A2"].value) == (5, 6)


@pytest.mark.com
@pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool)")
@pytest.mark.timeout(300)
def test_excel_agrees_with_the_whole_array_write_and_the_refusal(tmp_path):
    """Excel is the oracle: after a refused
    partial write Excel still shows the array and its own results; after a
    whole-array write Excel shows exactly the values written, with no array
    left behind."""
    from xlsx_mcp.ops import comtier

    def excel_sees(path):
        def body(_app, wb):
            ws = wb.Worksheets(1)
            return {"values": [ws.Range(a).Value for a in ("A1", "A2")],
                    "array": [bool(ws.Range(a).HasArray) for a in ("A1", "A2")]}
        return comtier._run_readonly("oracle", path, body, timeout=120)

    p = _array_book(tmp_path / "arr.xlsx")
    h = fh.start()
    _payload, err, _ = h.call("xl_edit", {"op": "set_cell",
                                          "args": {"path": p, "location": "A2", "value": 5}})
    assert err
    seen = excel_sees(p)
    assert seen["array"] == [True, True] and seen["values"] == [2, 4], seen
    q = _array_book(tmp_path / "arr2.xlsx")
    payload, err, _ = h.call("xl_edit", {"op": "write_range", "args": {
        "path": q, "location": "A1", "data": [[5], [6]]}})
    assert not err, payload
    seen = excel_sees(q)
    assert seen["array"] == [False, False] and seen["values"] == [5, 6], seen


# ======================================================================
# The PDF output path is settled before Excel sees it
# ======================================================================


@pytest.mark.com
@pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool)")
@pytest.mark.timeout(300)
def test_an_extensionless_pdf_output_is_written_exactly_where_the_answer_says(xlsx, tmp_path):
    payload, err, _ = fh.start().call("xl_app", {"op": "com_export_pdf", "args": {
        "path": xlsx, "output": str(tmp_path / "report")}})
    assert not err, payload
    assert payload["exported"].endswith("report.pdf")
    assert (tmp_path / "report.pdf").read_bytes()[:5] == b"%PDF-"
    assert not (tmp_path / "report").exists()


@pytest.mark.parametrize("given", ["{src}", "{slot}", "{dir}/named.xlsx"])
def test_an_extensionless_pdf_output_is_judged_on_the_name_given_too(given, xlsx, tmp_path,
                                                                    monkeypatch):
    """The name the caller gave is refused when it is the source, a backup
    slot, or a workbook name, even though Excel would have written a
    sibling .pdf (no Excel is started)."""
    monkeypatch.setattr(com_session, "run_com", lambda *a, **k: (_ for _ in ()).throw(
        AssertionError("reached COM")))
    srv.set_cell(xlsx, "C1", "x")
    slot = safesave.slot_dir(xlsx) / safesave.PREV_SLOT
    out = given.format(src=xlsx, slot=slot, dir=tmp_path)
    payload, err, _ = fh.start().call("xl_app", {"op": "com_export_pdf", "args": {
        "path": xlsx, "output": out}})
    assert err and payload["error"]["code"] == "BAD_PARAMS", payload
    assert not Path(out + ".pdf").exists()


def test_convert_adds_the_formats_extension_before_the_guard(xlsx, tmp_path, monkeypatch):
    """The same class on com_convert_format: SaveAs adds the extension to a
    bare name, so the guard judges the name with it (an existing file of
    that name refuses CONFLICT; no Excel is started)."""
    monkeypatch.setattr(com_session, "run_com", lambda *a, **k: (_ for _ in ()).throw(
        AssertionError("reached COM")))
    (tmp_path / "copy.csv").write_text("a user's own csv", encoding="utf-8")
    payload, err, _ = fh.start().call("xl_app", {"op": "com_convert_format", "args": {
        "path": xlsx, "output": str(tmp_path / "copy"), "format": "csv"}})
    assert err and payload["error"]["code"] == "CONFLICT", payload
    assert (tmp_path / "copy.csv").read_text("utf-8") == "a user's own csv"


# ======================================================================
# The backup store by resolved path; a slot is never an edit target
# ======================================================================


@pytest.mark.skipif(sys.platform != "win32", reason="junctions are Windows")
def test_an_export_through_a_junction_into_the_store_refuses(xlsx, tmp_path):
    srv.set_cell(xlsx, "C1", "x")
    slot_dir = safesave.slot_dir(xlsx)
    link = tmp_path / "innocent"
    subprocess.run(["cmd", "/c", "mklink", "/J", str(link), str(slot_dir)], check=True,
                   capture_output=True, stdin=subprocess.DEVNULL, creationflags=NO_WINDOW)
    try:
        payload, err, _ = fh.start().call("xl_workbook", {"op": "export_range", "args": {
            "path": xlsx, "location": "A1:B3", "out_file": str(link / "x.csv")}})
        assert err and "backup store" in payload["error"]["message"], payload
        assert not (slot_dir / "x.csv").exists()
    finally:
        os.rmdir(link)


@pytest.mark.parametrize("tool,op,args", [
    ("xl_workbook", "manage_backups", {"action": "restore", "source": "prev"}),
    ("xl_format", "format_cells", {"location": "A1", "font": {"bold": True}}),
    ("xl_app", "com_autofit", {}),
    ("xl_app", "com_save_with_password", {"password": "p"}),
])
def test_no_op_edits_a_backup_slot_in_place(tool, op, args, xlsx, monkeypatch):
    monkeypatch.setattr(com_session, "run_com", lambda *a, **k: (_ for _ in ()).throw(
        AssertionError("reached COM")))
    srv.set_cell(xlsx, "C1", "x")
    slot = safesave.slot_dir(xlsx) / safesave.PREV_SLOT
    before = _sha(slot)
    payload, err, _ = fh.start().call(tool, {"op": op, "args": {"path": str(slot), **args}})
    assert err and "backup" in payload["error"]["message"], payload
    assert _sha(slot) == before


def test_reading_a_backup_slot_still_works(xlsx):
    srv.set_cell(xlsx, "C1", "x")
    slot = safesave.slot_dir(xlsx) / safesave.PREV_SLOT
    payload, err, _ = fh.start().call("xl_read", {"op": "read_range", "args": {
        "path": str(slot), "location": "A1:B3"}})
    assert not err, payload


# ======================================================================
# more served-hint and deadline cases
# ======================================================================


@pytest.mark.skipif(sys.platform != "win32", reason="Windows device names")
@pytest.mark.parametrize("device", ["CON", "conin$", "CONOUT$", "aux.xlsx", "COM1", "nul"])
def test_a_device_path_refuses_at_once_for_every_path_kind(device, tmp_path):
    h = fh.start(Settings(call_timeout_s=20.0, max_call_timeout_s=20.0))
    t0 = time.monotonic()
    for tool, op, args in [
            ("xl_read", "read_range", {"path": device, "location": "A1"}),
            ("xl_edit", "import_data", {"path": str(tmp_path / "i.xlsx"),
                                        "source_file": device, "fmt": "csv"}),
            ("xl_workbook", "copy_workbook", {"src": device,
                                              "dst": str(tmp_path / "c.xlsx")})]:
        payload, err, _ = h.call(tool, {"op": op, "args": args})
        assert err and payload["error"]["code"] == "BAD_PARAMS", payload
        assert "device" in payload["error"]["message"]
    assert time.monotonic() - t0 < 15


def test_ordinary_names_that_contain_a_device_name_still_open(tmp_path):
    for name in ("console.xlsx", "con_log.xlsx", "auxiliary.xlsx", "com10.xlsx"):
        p = tmp_path / name
        srv.create_workbook(str(p))
        assert srv.read_range(str(p), "A1") is not None


def test_extract_refuses_before_writing_anything_when_one_name_is_taken(tmp_path):
    import base64
    png = tmp_path / "p.png"
    png.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="))
    book = tmp_path / "i.xlsx"
    srv.create_workbook(str(book))
    srv.manage_image(str(book), "insert", image_file=str(png), location="B2")
    srv.manage_image(str(book), "insert", image_file=str(png), location="D4")
    fresh = tmp_path / "fresh"
    fresh.mkdir()
    first = srv.manage_image(str(book), "extract", out_dir=str(fresh))
    assert first["count"] >= 1
    taken = tmp_path / "taken"
    taken.mkdir()
    victim = taken / Path(first["extracted"][-1]).name
    victim.write_text("a user's own file", encoding="utf-8")
    payload, err, _ = fh.start().call("xl_format", {"op": "manage_image", "args": {
        "path": str(book), "action": "extract", "out_dir": str(taken)}})
    assert err and payload["error"]["code"] == "CONFLICT", payload
    assert victim.read_text("utf-8") == "a user's own file"
    assert sorted(p.name for p in taken.iterdir()) == [victim.name], "a partial extract"


def test_an_oversized_csv_field_is_refused_as_the_argument_it_is(xlsx):
    payload, err, _ = fh.start().call("xl_edit", {"op": "import_data", "args": {
        "path": xlsx, "source": "A" * 300000, "fmt": "csv", "location": "J1"}})
    assert err and payload["error"]["code"] == "BAD_PARAMS", payload
    assert "131,072" in payload["error"]["message"]
    # An engine's own refusal carries the
    # short engine hint, not the generic "did not fit this op" line.
    assert payload["error"]["hint"] == envelope.ENGINE_REFUSAL_HINT


def test_a_missing_clipboard_names_the_session_not_the_arguments(xlsx, tmp_path, monkeypatch):
    from xlsx_mcp.ops import comtier
    monkeypatch.setattr(comtier, "_clipboard_available", lambda: False)
    payload, err, _ = fh.start().call("xl_app", {"op": "com_render_sheet", "args": {
        "path": xlsx, "output": str(tmp_path / "r.png")}})
    assert err and payload["error"]["code"] == "APP_NOT_RUNNING", payload
    assert payload["error"]["hint"] == comtier.CLIPBOARD_HINT
    assert "clipboard" in payload["error"]["message"]


# ======================================================================
# Direct writers and array cells
# ======================================================================


def test_fa2_a_future_direct_writer_cannot_change_part_of_an_array(tmp_path):
    """F-A2 at the root: save() compares every array cell with the snapshot
    taken when the model loaded, so a writer that assigns cell.value (or a
    hyperlink) directly, bypassing set_cell, is refused too."""
    from xlsx_mcp.core.errors import UnsupportedStructure
    from xlsx_mcp.core.package import WorkbookPackage
    from openpyxl.worksheet.hyperlink import Hyperlink
    for change in ("value", "hyperlink"):
        p = _array_book(tmp_path / f"d_{change}.xlsx")
        before = _sha(p)
        pkg = WorkbookPackage.open(p)
        cell = pkg.workbook.active["A2"]
        if change == "value":
            cell.value = "typed over"
        else:
            cell.hyperlink = Hyperlink(ref="A2", target="https://example.com")
        with pytest.raises(UnsupportedStructure) as info:
            pkg.save()
        assert "A1:A2" in str(info.value)
        assert _sha(p) == before


def test_fa2_a_direct_writer_covering_the_whole_array_is_allowed(tmp_path):
    from xlsx_mcp.core.package import WorkbookPackage
    p = _array_book(tmp_path / "whole.xlsx")
    pkg = WorkbookPackage.open(p)
    ws = pkg.workbook.active
    ws["A1"].value, ws["A2"].value = "x", "y"
    pkg.save()
    wb = openpyxl.load_workbook(p)
    assert (wb.active["A1"].value, wb.active["A2"].value) == ("x", "y")


def test_fa2_a_structural_edit_resnapshots_the_moved_array(tmp_path):
    """In one package: a row inserted above the array moves it to A2:A3;
    a direct write into part of it at its NEW address is still refused, and
    the cells it left are not mistaken for a change."""
    from xlsx_mcp.core import refs
    from xlsx_mcp.core.errors import UnsupportedStructure
    from xlsx_mcp.core.package import WorkbookPackage
    p = _array_book(tmp_path / "moved.xlsx")
    before = _sha(p)
    pkg = WorkbookPackage.open(p)
    ws = pkg.workbook.active
    pkg.modify_structure(refs.RefEdit(sheet=ws.title, kind="insert_rows", index=1, count=1))
    assert getattr(ws["A2"].value, "ref", None) == "A2:A3"
    ws["A3"].value = "typed over"
    with pytest.raises(UnsupportedStructure) as info:
        pkg.save()
    assert "A2:A3" in str(info.value)
    assert _sha(p) == before


@pytest.mark.parametrize("spelling", ["CON:", "conin$:", "\\\\.\\CON", "\\\\.\\CONIN$",
                                      "\\\\?\\CONIN$", "//./COM1", "C:\\x\\aux.xlsx", "nul:"])
def test_fd2_every_device_spelling_is_named(spelling):
    from xlsx_mcp.core.sandbox import device_name
    assert device_name(spelling) is not None, spelling


@pytest.mark.parametrize("ordinary", ["console.xlsx", "C:\\con\\a.xlsx", "\\\\?\\C:\\x.xlsx",
                                      "com10.xlsx", "C:", "D:\\data\\book.xlsx"])
def test_fd2_ordinary_paths_are_not_devices(ordinary):
    from xlsx_mcp.core.sandbox import device_name
    assert device_name(ordinary) is None, ordinary


def _tests_workflow() -> dict:
    import yaml
    return yaml.safe_load((ROOT / ".github/workflows/tests.yml").read_text("utf-8"))


def test_ci_throttle_interim_matrix_and_full_matrix_on_request_only():
    """The CI throttle: a plain push or PR runs one Windows and
    one Ubuntu leg on the newest Python; the full 3 x 2 matrix only on
    workflow_dispatch full=true or a release-candidate tag; no macOS; one
    run per ref with cancel-in-progress."""
    wf = _tests_workflow()
    on = wf.get("on", wf.get(True))
    assert on["workflow_dispatch"]["inputs"]["full"]["type"] == "boolean"
    assert on["workflow_dispatch"]["inputs"]["full"]["default"] is False
    assert on["push"]["branches"] == ["main"]
    assert all("rc" in t for t in on["push"].get("tags", []))
    assert wf["concurrency"] == {"group": "${{ github.workflow }}-${{ github.ref }}",
                                 "cancel-in-progress": True}
    for name, job in wf["jobs"].items():
        matrix = job.get("strategy", {}).get("matrix", {})
        oses = matrix.get("os", [])
        assert not any("macos" in str(o) for o in oses), name
        assert sorted(oses) == ["ubuntu-latest", "windows-latest"], (name, oses)
        expr = matrix["python"]
        m = re.fullmatch(r"\$\{\{ fromJSON\((.*) && '(\[.*?\])' \|\| '(\[.*?\])'\) \}\}", expr)
        assert m, expr
        cond, full, interim = m.group(1), json.loads(m.group(2)), json.loads(m.group(3))
        assert interim == ["3.14"]
        assert full == ["3.12", "3.13", "3.14"]
        assert "github.event_name == 'workflow_dispatch' && inputs.full" in cond
        assert "startsWith(github.ref, 'refs/tags/')" in cond
        assert "push" not in cond and "pull_request" not in cond


# ======================================================================
# Table writes
# ======================================================================


def _table_book(p: Path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B1"], ws["B2"] = 1, 2
    ws["D1"], ws["E1"], ws["D2"], ws["E2"], ws["D3"], ws["E3"] = "k", "v", "a", 1, "b", 2
    wb.save(p)
    srv.create_table(str(p), "D1:E3", "T")
    return str(p)


@pytest.mark.parametrize("tool,op,args,why", [
    ("xl_edit", "set_cell", {"location": "E1", "value": 5}, "number or nothing"),
    ("xl_edit", "set_cell", {"location": "E1", "value": None}, "number or nothing"),
    ("xl_edit", "set_cell", {"location": "E1", "value": "K"}, "repeat the header"),
    ("xl_edit", "set_formula", {"location": "E1", "formula": "=1+1"}, "formula"),
    ("xl_edit", "move_range", {"source": "B1", "dest": "E1"}, "number or nothing"),
    ("xl_edit", "write_range", {"location": "D1", "data": [["k", "k"]]}, "repeat the header"),
], ids=["number", "empty", "duplicate", "formula", "moved", "write_range_dup"])
def test_ft_a_header_change_refuses_before_any_write(tool, op, args, why, tmp_path):
    p = _table_book(tmp_path / "t.xlsx")
    before = _sha(p)
    payload, err, _ = fh.start().call(tool, {"op": op, "args": {"path": p, **args}})
    assert err, payload
    msg = payload["error"]["message"]
    assert "header cell of table T" in msg and why in msg, msg
    assert "write its new text into the header cell" in msg
    assert _sha(p) == before


@pytest.mark.parametrize("what", ["all", "formulas"])
def test_ft_a_multi_cell_array_copied_into_a_table_refuses(what, tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula
    p = tmp_path / "ta.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"] = 1, 2
    ws["B1"] = ArrayFormula("B1:B2", "=A1:A2*2")
    ws["D1"], ws["E1"], ws["D2"], ws["E2"], ws["D3"], ws["E3"] = "k", "v", "a", 1, "b", 2
    wb.save(p)
    srv.create_table(str(p), "D1:E3", "T")
    before = _sha(p)
    payload, err, _ = fh.start().call("xl_edit", {"op": "copy_range", "args": {
        "path": str(p), "source": "A1:B2", "dest": "D2", "what": what}})
    assert err, payload
    assert "inside table T" in payload["error"]["message"], payload
    assert _sha(p) == before


def test_ft_new_header_text_renames_the_column_with_it(tmp_path):
    """Typing new text over a header renames the column, as in Excel: the
    table's column name follows the header cell."""
    p = _table_book(tmp_path / "rc.xlsx")
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_cell", "args": {
        "path": p, "location": "E1", "value": "w"}})
    assert not err, payload
    wb = openpyxl.load_workbook(p)
    t = wb.active.tables["T"]
    assert wb.active["E1"].value == "w"
    assert [c.name for c in t.tableColumns] in ([], ["k", "w"])


def test_ft_a_header_rename_refuses_while_a_formula_names_the_column(tmp_path):
    p = _table_book(tmp_path / "rf.xlsx")
    srv.set_formula(p, "G1", "=SUM(T[v])")
    before = _sha(p)
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_cell", "args": {
        "path": p, "location": "E1", "value": "w"}})
    assert err and "G1" in payload["error"]["message"], payload
    assert _sha(p) == before


def test_ft_table_ops_that_change_headers_still_work(tmp_path):
    p = _table_book(tmp_path / "ops.xlsx")
    h = fh.start()
    for args in ({"action": "add_column", "column": "x", "values": [7, 8]},
                 {"action": "delete_column", "column": "v"},
                 {"action": "add_row", "values": ["c", 3]},
                 {"action": "toggle_totals", "on": True},
                 {"action": "rename", "new_name": "Sales"}):
        payload, err, _ = h.call("xl_structure", {"op": "manage_table", "args": {
            "path": p, "name": "T", **args}})
        assert not err, (args, payload)
    wb = openpyxl.load_workbook(p)
    assert "Sales" in wb.active.tables


def test_ft_a_workbook_that_arrived_with_a_bad_header_still_takes_unrelated_edits(tmp_path):
    """A table that is as it was at load is not judged."""
    p = _table_book(tmp_path / "pre.xlsx")
    wb = openpyxl.load_workbook(p)
    wb.active["E1"] = 7  # broken outside this server
    wb.save(p)
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_cell", "args": {
        "path": p, "location": "A10", "value": 1}})
    assert not err, payload


@pytest.mark.com
@pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool) as the judge")
@pytest.mark.timeout(600)
def test_ft_the_controls_and_the_rename_still_save_and_open_clean_in_excel(tmp_path):
    """No edit, a far-away cell, a value in the table body, values under
    the header row and a header renamed by new text: each saves, and Excel
    opens it clean."""
    h = fh.start()
    cases = [
        None,
        ("xl_edit", "set_cell", {"location": "A20", "value": 3}),
        ("xl_edit", "set_cell", {"location": "E2", "value": 99}),
        ("xl_edit", "write_range", {"location": "D2", "data": [["z", 5], ["y", 6]]}),
        ("xl_edit", "set_cell", {"location": "E1", "value": "w"}),  # a column rename
    ]
    for i, case in enumerate(cases):
        p = _table_book(tmp_path / f"c{i}.xlsx")
        if case is not None:
            tool, op, args = case
            payload, err, _ = h.call(tool, {"op": op, "args": {"path": p, **args}})
            assert not err, (case, payload)
        verdict = srv.com_validate_opens_clean(p)
        assert verdict["opens_clean"], (case, verdict)


@pytest.mark.skipif(sys.platform != "win32", reason="Windows device namespace")
@pytest.mark.parametrize("path", ["\\\\.\\GLOBALROOT\\Device\\ConDrv\\Console",
                                  "\\\\?\\GLOBALROOT\\Device\\ConDrv\\Console",
                                  "\\\\.\\pipe\\x", "\\\\.\\PhysicalDrive0", "//./C:/x.xlsx",
                                  "\\??\\C:\\x.xlsx"])
def test_fd3_device_namespace_paths_refuse_at_once(path):
    h = fh.start(Settings(call_timeout_s=20.0, max_call_timeout_s=20.0))
    t0 = time.monotonic()
    payload, err, _ = h.call("xl_read", {"op": "read_range", "args": {"path": path,
                                                                      "location": "A1"}})
    assert err and payload["error"]["code"] == "BAD_PARAMS", payload
    assert "device path" in payload["error"]["message"]
    assert time.monotonic() - t0 < 10


@pytest.mark.skipif(sys.platform != "win32", reason="Windows long-path forms")
def test_fd3_the_long_path_forms_still_work(xlsx):
    long = "\\\\?\\" + os.path.abspath(xlsx)
    payload, err, _ = fh.start().call("xl_read", {"op": "read_range", "args": {
        "path": long, "location": "A1:B3"}})
    assert not err, payload
    from xlsx_mcp.core.sandbox import device_namespace
    assert not device_namespace(long)
    assert not device_namespace("\\\\?\\UNC\\server\\share\\b.xlsx")
    assert not device_namespace("C:\\x\\b.xlsx")
