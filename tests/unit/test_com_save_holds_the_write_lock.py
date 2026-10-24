"""A COM save holds the workbook's write lock from before Excel opens the
workbook until its copy is in place.

Excel edits the workbook as it was when Excel opened it. The COM tier took
no write lock, so a file-tier edit (by another server, or by another call in
this one) that landed after Excel opened the workbook and before Excel's copy
replaced it was overwritten by that copy, while both calls answered ok. Now
such an edit waits for the COM save and is then refused in the file tier's
own words (another kitchensink4xl process is modifying the workbook, or the
workbook changed after the edit read it); made after the save, it applies on
top of it. A program that does not use the lock is caught by a stamp check:
the workbook's size and modification time when the save began must still be
the file's when Excel has opened it and again just before the copy replaces
it. Otherwise the copy is dropped and the workbook keeps that program's
change.

The lock is taken and released in the calling thread. The COM worker thread
never takes it (it refuses to, so a mistake fails at once instead of
deadlocking), and only the calling thread puts the copy in place, so a COM
timeout releases the lock and an Excel that answers late never replaces the
workbook.

The first tests drive a fake Excel in this process and a real second server
(the Community launcher over stdio); the last one starts the real Excel.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import queue
import subprocess
import sys
import tempfile
import threading
import time
from pathlib import Path

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp.com import session as com_session
from xlsx_mcp.core import safesave
from xlsx_mcp.core.errors import ExcelBlocked, XlMcpError
from xlsx_mcp.ops import comtier

IS_WIN = sys.platform == "win32"
NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)
ROOT = Path(__file__).resolve().parents[2]

#: The file tier's own refusals for an edit that meets a COM save: the lock
#: wait ran out, or the workbook changed after the edit read it.
FILE_TIER_REFUSALS = {
    "WORKBOOK_LOCKED": "is being modified by another kitchensink4xl process",
    "CONFLICT": "changed on disk after this operation read it",
}


def _sha(p) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def _cell(p, ref):
    wb = openpyxl.load_workbook(p)
    try:
        return wb.active[ref].value
    finally:
        wb.close()


def _beside(book) -> list:
    return sorted(n for n in os.listdir(os.path.dirname(book))
                  if n not in (os.path.basename(book), safesave.BACKUP_DIR_NAME))


def _temp_left() -> list:
    return sorted(os.listdir(tempfile.gettempdir()))


@pytest.fixture
def book(tmp_path, monkeypatch):
    folder = tmp_path / "work"
    folder.mkdir()
    p = str(folder / "book.xlsx")
    wb = openpyxl.Workbook()
    wb.active["A1"] = "before"
    wb.save(p)
    temp = tmp_path / "temp"
    temp.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(temp))
    return p


# ------------------------------------------------------------- fake Excel


class _Worker:
    pid = 0
    app = object()


class _Manager:
    def acquire(self):
        return _Worker()

    def release_to_pool(self, w):
        pass


class _FakeWB:
    """A workbook Excel opened read-write. Like Excel, it holds the workbook
    as it was at the open; its copy is that plus Excel's edit."""

    def __init__(self, path, at_close=None, at_copy=None):
        self.path = path
        self.FullName = os.path.abspath(path)
        self.ReadOnly = False
        self.Password = ""
        self.opened = Path(path).read_bytes()
        self.at_close = at_close
        self.at_copy = at_copy

    def Save(self):
        raise AssertionError("wb.Save lets Excel replace the file itself")

    def SaveCopyAs(self, copy):
        wb = openpyxl.load_workbook(io.BytesIO(self.opened))
        wb.active["A1"] = "edited by Excel"
        wb.save(copy)
        if self.at_copy is not None:
            self.at_copy()  # the copy is written; Excel has not answered yet

    def Close(self, save):
        if self.at_close is not None:
            hook, self.at_close = self.at_close, None
            hook()  # Excel has closed the workbook; its copy is not in place yet


def _fake_excel(monkeypatch, *, at_open=None, at_close=None, at_copy=None, run_com=None):
    monkeypatch.setattr(com_session, "run_com", run_com or
                        (lambda label, job, timeout=None: job(_Manager())))
    monkeypatch.setattr(com_session, "_assert_hygiene", lambda app: ())
    monkeypatch.setattr(com_session, "_restore_hygiene", lambda app, *a: None)
    monkeypatch.setattr(com_session, "com_retry", lambda fn, label=None, **k: fn())

    def _open(app, p, password=None, **k):
        if at_open is not None:
            at_open(p)
        return _FakeWB(p, at_close=at_close, at_copy=at_copy)

    monkeypatch.setattr(com_session, "open_workbook", _open)


def _autofit(book, **kw):
    return comtier._run_mutation("autofit(book.xlsx)", book,
                                 lambda app, wb: {"autofit": True}, **kw)


# ------------------------------------------------- a second, real server


class _OtherServer:
    """A second kitchensink4xl server process: the Community launcher over
    real stdio, as an MCP client starts it. Its edits are file-tier ones."""

    def __init__(self):
        self.proc = subprocess.Popen(
            fh.launch_cmd(), stdin=subprocess.PIPE, stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL, env=fh._env(None), creationflags=NO_WINDOW)
        self._lines: queue.Queue = queue.Queue()
        threading.Thread(target=self._pump, daemon=True).start()
        self._id = 0
        self._request("initialize", {
            "protocolVersion": "2025-06-18", "capabilities": {},
            "clientInfo": {"name": "ks4xl-test", "version": "0"}})
        self._send({"jsonrpc": "2.0", "method": "notifications/initialized"})

    def _pump(self):
        for line in self.proc.stdout:
            self._lines.put(line)
        self._lines.put(None)

    def _send(self, msg):
        self.proc.stdin.write((json.dumps(msg) + "\n").encode("utf-8"))
        self.proc.stdin.flush()

    def _request(self, method, params, timeout=90.0):
        self._id += 1
        rid = self._id
        self._send({"jsonrpc": "2.0", "id": rid, "method": method, "params": params})
        end = time.monotonic() + timeout
        while True:
            line = self._lines.get(timeout=max(0.1, end - time.monotonic()))
            if line is None:
                raise RuntimeError("the other server closed its output")
            msg = json.loads(line)
            if msg.get("id") == rid:
                return msg

    def set_cell(self, path, cell, value, timeout=90.0) -> dict:
        msg = self._request("tools/call", {"name": "xl_edit", "arguments": {
            "op": "set_cell", "args": {"path": path, "location": cell, "value": value}}},
            timeout=timeout)
        r = msg.get("result") or {}
        content = r.get("content") or []
        text = content[0]["text"] if content else json.dumps(msg.get("error"))
        try:
            payload = json.loads(text)
        except ValueError:
            payload = {"_raw": text}
        return {"is_error": bool(r.get("isError")) or "error" in msg, "payload": payload}

    def close(self):
        try:
            self.proc.stdin.close()
        except OSError:
            pass
        try:
            self.proc.wait(timeout=20)
        except subprocess.TimeoutExpired:
            self.proc.kill()
            self.proc.wait(timeout=20)


@pytest.fixture
def other_server():
    srv = _OtherServer()
    yield srv
    srv.close()


def _edit_in_background(srv, book, value):
    answer: dict = {}

    def run():
        try:
            answer.update(srv.set_cell(book, "C3", value))
        except Exception as exc:  # noqa: BLE001
            answer["crashed"] = repr(exc)

    th = threading.Thread(target=run, daemon=True)
    th.start()
    return th, answer


def _assert_never_lost(answer, book, value):
    """The other server's edit is in the workbook, or it was refused in the
    file tier's own words. Never ok and then gone."""
    assert "crashed" not in answer, answer
    if not answer["is_error"]:
        assert _cell(book, "C3") == value, (
            "the other server answered ok, and Excel's copy then overwrote its edit")
        return None
    err = answer["payload"].get("error") or {}
    code, message = err.get("code"), err.get("message", "")
    assert code in FILE_TIER_REFUSALS, answer
    assert FILE_TIER_REFUSALS[code] in message, answer
    assert _cell(book, "C3") is None
    return code


# ---------------------------------------------------- the lost update


@pytest.mark.timeout(240)
def test_another_servers_edit_in_the_promote_window_waits_and_is_never_lost(
        book, monkeypatch, other_server):
    """The edit is made after Excel has closed the workbook and before its
    copy is in place, and the COM save waits for the edit's answer. Before, the
    edit answered ok and Excel's copy then overwrote it. Now the edit waits
    on the lock until its wait runs out (WORKBOOK_LOCKED, nothing written)."""
    box = {}

    def in_window():
        th, answer = _edit_in_background(other_server, book, "from the other server")
        th.join(60)
        box["answer"] = answer

    _fake_excel(monkeypatch, at_close=in_window)
    res = _autofit(book)
    assert res["ok"]
    assert _cell(book, "A1") == "edited by Excel"
    code = _assert_never_lost(box["answer"], book, "from the other server")
    assert code == "WORKBOOK_LOCKED", box["answer"]
    assert _beside(book) == [] and _temp_left() == []


@pytest.mark.timeout(240)
def test_another_servers_edit_that_waited_for_a_com_save_is_refused_as_stale(
        book, monkeypatch, other_server):
    """The COM save goes on while the edit is still waiting for the lock: the
    edit then finds the workbook changed since it read it (CONFLICT), so it
    is refused instead of writing over Excel's change. Before, it answered ok and
    lost."""
    box = {}

    def in_window():
        th, answer = _edit_in_background(other_server, book, "from the other server")
        th.join(6)  # under the other server's 10 s lock wait
        box["thread"], box["answer"] = th, answer

    _fake_excel(monkeypatch, at_close=in_window)
    res = _autofit(book)
    assert res["ok"]
    box["thread"].join(60)
    assert _cell(book, "A1") == "edited by Excel"
    code = _assert_never_lost(box["answer"], book, "from the other server")
    assert code == "CONFLICT", box["answer"]
    assert _beside(book) == [] and _temp_left() == []


@pytest.mark.timeout(240)
def test_an_edit_made_after_the_com_save_applies_on_top(book, monkeypatch, other_server):
    _fake_excel(monkeypatch)
    assert _autofit(book)["ok"]
    answer = other_server.set_cell(book, "C3", "after")
    assert not answer["is_error"], answer
    assert _cell(book, "A1") == "edited by Excel" and _cell(book, "C3") == "after"


# ------------------------------------------- writers that skip the lock


def _foreign_write(p, value):
    """A program that does not use this server's lock rewrites the file."""
    wb = openpyxl.load_workbook(p)
    wb.active["B2"] = value
    wb.save(p)
    st = os.stat(p)
    os.utime(p, ns=(st.st_atime_ns, st.st_mtime_ns + 2_000_000_000))


def test_a_program_that_ignores_the_lock_keeps_its_change(book, monkeypatch):
    """It writes the workbook after Excel opened it. Before, Excel's copy
    replaced the file and the program's change was gone."""
    slot = safesave.slot_dir(book) / safesave.PREV_SLOT
    slot_before = _sha(slot) if slot.exists() else None
    _fake_excel(monkeypatch, at_close=lambda: _foreign_write(book, "their change"))
    with pytest.raises(XlMcpError) as ei:
        _autofit(book)
    assert getattr(ei.value, "code", None) == "CONFLICT", ei.value
    assert "changed on disk" in str(ei.value)
    assert _cell(book, "B2") == "their change"
    assert _cell(book, "A1") == "before"
    assert (_sha(slot) if slot.exists() else None) == slot_before, "the prev slot was burned"
    assert _beside(book) == [] and _temp_left() == []


def test_a_write_while_excel_opens_the_workbook_is_refused_before_the_edit(book, monkeypatch):
    """A change that lands after the save took its stamp and before Excel
    has the workbook is caught as soon as Excel has opened it: the backup
    taken for this save would not hold the state Excel edits."""
    ran = []
    _fake_excel(monkeypatch, at_open=lambda p: _foreign_write(p, "their change"))
    with pytest.raises(XlMcpError) as ei:
        comtier._run_mutation("autofit(book.xlsx)", book,
                              lambda app, wb: ran.append(1) or {"autofit": True})
    assert getattr(ei.value, "code", None) == "CONFLICT", ei.value
    assert ran == [], "Excel edited a workbook that had changed under it"
    assert _cell(book, "B2") == "their change"
    assert _beside(book) == [] and _temp_left() == []


# ------------------------------------------------ threads and timeouts


def test_a_thread_marked_as_the_com_worker_refuses_the_write_lock(book):
    """The COM worker runs jobs for callers that hold the lock; a hold there
    would wait for itself. It fails at once instead."""
    out = {}

    def worker():
        safesave.forbid_write_lock_on_this_thread("a test COM worker")
        try:
            with safesave.write_lock(book):
                out["took"] = True
        except RuntimeError as exc:
            out["refused"] = str(exc)

    th = threading.Thread(target=worker)
    th.start()
    th.join(10)
    assert "refused" in out and "took" not in out, out
    with safesave.write_lock(book):  # any other thread still takes it
        pass


def _com_ok() -> bool:
    try:
        return bool(com_session.com_available()[0])
    except Exception:  # noqa: BLE001
        return False


@pytest.mark.skipif(not _com_ok(), reason="needs pywin32 and Excel's COM registration "
                    "(no Excel is started)")
def test_the_com_worker_thread_refuses_the_write_lock_instead_of_deadlocking(book, tmp_path):
    """The real executor's worker thread. With the caller holding the lock,
    a job that took it would wait for its own caller until the timeout."""
    ex = com_session.ComExecutor(journal_path=str(tmp_path / "journal.json"))

    def job(manager):
        with safesave.write_lock(book):
            return "took it"

    try:
        t0 = time.monotonic()
        with pytest.raises(RuntimeError):
            ex.submit("probe", job, timeout=20)
        with safesave.write_lock(book):
            with pytest.raises(RuntimeError):
                ex.submit("probe", job, timeout=20)
        assert time.monotonic() - t0 < 15, "the job waited for the lock"
    finally:
        ex.shutdown()


def _child_takes_the_lock(book) -> str:
    code = ("import sys; from xlsx_mcp.core import safesave as s; s.LOCK_WAIT_SECONDS = 3\n"
            "with s.write_lock(sys.argv[1]): print('took it')")
    env = dict(os.environ, PYTHONPATH=str(ROOT / "src"))
    r = subprocess.run([sys.executable, "-X", "utf8", "-c", code, book], capture_output=True,
                       text=True, timeout=60, env=env, stdin=subprocess.DEVNULL,
                       creationflags=NO_WINDOW)
    return (r.stdout + r.stderr).strip()


def test_a_com_timeout_releases_the_lock_and_a_late_excel_never_replaces_the_workbook(
        book, monkeypatch):
    """The job is still inside Excel's save when the call's timeout
    answers. The lock is released with the answer, and when the job comes
    back later its copy is dropped, never put in place (it used to be left in the
    temp folder)."""
    before = _sha(book)
    release = threading.Event()
    inside = threading.Event()
    jobs = []

    def at_copy():
        inside.set()
        release.wait(30)

    def late_job(job):
        try:
            jobs.append(("returned", job(_Manager())))
        except BaseException as exc:  # noqa: BLE001 (the answer was already given)
            jobs.append(("raised", str(exc)))

    def run_com(label, job, timeout=None):
        th = threading.Thread(target=late_job, args=(job,), daemon=True)
        th.start()
        assert inside.wait(30)
        run_com.thread = th
        raise ExcelBlocked(f"{label} exceeded the timeout (test)")

    _fake_excel(monkeypatch, at_copy=at_copy, run_com=run_com)
    with pytest.raises(ExcelBlocked):
        _autofit(book)
    assert _child_takes_the_lock(book) == "took it"
    release.set()
    run_com.thread.join(30)
    assert len(jobs) == 1 and jobs[0][0] == "raised", jobs  # it stopped at its copy
    assert _sha(book) == before
    assert _beside(book) == []
    assert _temp_left() == [], "the late job's copy was left in the temp folder"


# ------------------------------------------------------------- real Excel


def _shut_down_excel():
    try:
        com_session.get_executor().shutdown()
    except Exception:  # noqa: BLE001
        pass
    com_session._EXECUTOR = None


@pytest.mark.com
@pytest.mark.skipif(not IS_WIN, reason="Windows and Excel")
@pytest.mark.skipif(not _com_ok(), reason="needs Excel (the product's private hidden pool)")
@pytest.mark.timeout(300)
def test_live_recalculate_and_another_servers_edit_never_lose_it(book, monkeypatch,
                                                                 other_server):
    """recalculate(engine='com') in the real hidden Excel; the other server
    edits once Excel has saved its copy and closed the workbook, and the
    recalculation waits up to 6 s for that edit's answer before its copy goes
    in place. Before: ok, ok, and the edit was gone."""
    wb = openpyxl.Workbook()
    wb.active["A1"] = 2
    wb.active["A2"] = "=A1*21"
    wb.save(book)
    real_put = com_session.put_in_place
    box = {}

    def put_in_place(copy, path, **kw):
        th, answer = _edit_in_background(other_server, book, "from the other server")
        th.join(6)
        box["thread"], box["answer"] = th, answer
        return real_put(copy, path, **kw)

    monkeypatch.setattr(com_session, "put_in_place", put_in_place)
    try:
        res = comtier.recalculate(book, engine="com", timeout_seconds=90)
        assert res["ok"] and res["recalculated"]
        box["thread"].join(60)
        cached = openpyxl.load_workbook(book, data_only=True)
        assert cached.active["A2"].value == 42
        _assert_never_lost(box["answer"], book, "from the other server")
        assert [n for n in _beside(book) if not n.startswith("~$")] == []
    finally:
        _shut_down_excel()
