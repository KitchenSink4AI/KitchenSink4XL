"""Silent corruption, live since 1.2.6: set_filter on a table writes the
TABLE's own filter, never a sheet filter over the table.

set_filter with location {"table": "Fruit"}, the table's range, or a cell in
the table answered ok and verified, but wrote a sheet-level <autoFilter> over
the table's range plus an _xlnm._FilterDatabase name pointing into it, and
Excel will not open such a workbook ("Open method of Workbooks class
failed"). Excel filters a table with the table's own autoFilter (in its table
part). Now:

- a target inside a table filters that table (criteria in the table part,
  non-matching rows hidden), with no sheet filter and no _FilterDatabase;
- a range that reaches outside the table it overlaps refuses with the remedy;
- clear_filter on a table clears the table's criteria and unhides its rows,
  and removes a sheet filter that overlaps it (the repair for a file an
  older version broke);
- the post-save integrity check fails any save that leaves a sheet filter
  over a table or a _FilterDatabase pointing into one, so no route can ship
  this again (the file is restored, or never replaced).
"""
from __future__ import annotations

import re
import zipfile

import openpyxl
import pytest
from openpyxl.worksheet.table import Table

from xlsx_mcp.core import verify as _verify
from xlsx_mcp.core.errors import XlMcpError
from xlsx_mcp.ops import sortfilter


def _book(tmp_path, name="t.xlsx", totals=False):
    p = str(tmp_path / name)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for r in [["a", "b", "c"], [1, 2, 3], [4, 5, 6], [7, 8, 9], [10, 11, 12]]:
        ws.append(r)
    data = wb.create_sheet("Data")
    for r in [["Item", "Amount", "Date"], ["Apple", 3, 1], ["Pear", 7, 2], ["Plum", 9, 3]]:
        data.append(r)
    ref = "A1:C4"
    if totals:
        data.append(["Total", "=SUBTOTAL(109,[Amount])", None])
        ref = "A1:C5"
    t = Table(displayName="Fruit", ref=ref)
    if totals:
        t.totalsRowCount = 1
    data.add_table(t)
    wb.save(p)
    return p


def _parts(p):
    z = zipfile.ZipFile(p)
    names = z.namelist()
    sheet = z.read("xl/worksheets/sheet2.xml").decode()
    table = next(z.read(n).decode() for n in names if n.startswith("xl/tables/"))
    wbx = z.read("xl/workbook.xml").decode()
    return sheet, table, wbx


def _hidden_rows(sheet_xml):
    return sorted(int(m) for m in re.findall(r'<row r="(\d+)"[^>]*hidden="1"', sheet_xml))


LOCATIONS = {"table": {"table": "Fruit"},
             "range": {"range": "A1:C4", "sheet": "Data"},
             "cell": {"cell": "B2", "sheet": "Data"}}


@pytest.mark.parametrize("where", sorted(LOCATIONS))
def test_a_filter_on_a_table_is_the_tables_own_filter(tmp_path, where):
    p = _book(tmp_path)
    out = sortfilter.set_filter(p, LOCATIONS[where],
                                criteria=[{"column": "Amount", "op": "gt", "value": 5}])
    assert out["ok"] and out["changed"]["filter"]["table"] == "Fruit"
    sheet, table, wbx = _parts(p)
    assert "<autoFilter" not in sheet, "no sheet filter over a table"
    assert "_FilterDatabase" not in wbx
    assert re.search(r'<autoFilter ref="A1:C4"', table)
    assert "greaterThan" in table and 'colId="1"' in table
    assert _hidden_rows(sheet) == [2]
    assert _verify.package_problems(p) == set()


def test_a_filter_with_no_criteria_on_a_table_writes_no_sheet_filter(tmp_path):
    p = _book(tmp_path)
    sortfilter.set_filter(p, {"table": "Fruit"})
    sheet, table, wbx = _parts(p)
    assert "<autoFilter" not in sheet and "_FilterDatabase" not in wbx
    assert '<autoFilter ref="A1:C4"' in table


def test_the_totals_row_stays_out_of_the_filter(tmp_path):
    p = _book(tmp_path, totals=True)
    sortfilter.set_filter(p, {"table": "Fruit"},
                          criteria=[{"column": "Item", "op": "eq", "value": "Pear"}])
    sheet, table, _wbx = _parts(p)
    assert '<autoFilter ref="A1:C4"' in table
    assert _hidden_rows(sheet) == [2, 4]  # the totals row 5 is never hidden


def test_a_range_that_reaches_outside_the_table_refuses(tmp_path):
    p = _book(tmp_path)
    before = open(p, "rb").read()
    with pytest.raises(XlMcpError) as ei:
        sortfilter.set_filter(p, {"range": "A1:D6", "sheet": "Data"})
    assert "Fruit" in str(ei.value) and '{"table": "Fruit"}' in str(ei.value)
    assert open(p, "rb").read() == before


def test_a_plain_range_outside_any_table_keeps_the_sheet_filter(tmp_path):
    p = _book(tmp_path)
    sortfilter.set_filter(p, {"range": "A1:C5", "sheet": "Sheet1"},
                          criteria=[{"column": "a", "op": "gt", "value": 4}])
    z = zipfile.ZipFile(p)
    assert '<autoFilter ref="A1:C5"' in z.read("xl/worksheets/sheet1.xml").decode()
    assert _verify.package_problems(p) == set()


def test_clear_filter_on_a_table_clears_its_criteria_and_rows(tmp_path):
    p = _book(tmp_path)
    sortfilter.set_filter(p, {"table": "Fruit"},
                          criteria=[{"column": "Amount", "op": "gt", "value": 5}])
    out = sortfilter.clear_filter(p, {"table": "Fruit"})
    assert out["ok"] and out["changed"]["filter_cleared"]["table"] == "Fruit"
    sheet, table, wbx = _parts(p)
    assert _hidden_rows(sheet) == []
    assert "filterColumn" not in table and '<autoFilter ref="A1:C4"' in table
    assert "<autoFilter" not in sheet and "_FilterDatabase" not in wbx


def test_clear_filter_with_no_location_finds_the_one_filtered_table(tmp_path):
    p = _book(tmp_path)
    sortfilter.set_filter(p, {"table": "Fruit"},
                          criteria=[{"column": "Amount", "op": "gt", "value": 5}])
    out = sortfilter.clear_filter(p, sheet="Data")
    assert out["changed"]["filter_cleared"]["table"] == "Fruit"
    assert _hidden_rows(_parts(p)[0]) == []


def test_a_sort_inside_a_filtered_table_leaves_the_hidden_rows_in_place(tmp_path):
    """Excel sorts only the visible rows of a filtered range; the table's own
    filter counts as one now that set_filter writes it there."""
    p = _book(tmp_path)
    sortfilter.set_filter(p, {"table": "Fruit"},
                          criteria=[{"column": "Amount", "op": "gt", "value": 5}])
    out = sortfilter.sort_range(p, {"range": "A1:C4", "sheet": "Data"},
                                keys=[{"column": "Amount", "order": "desc"}])
    assert out["changed"]["sorted"].get("rows_pinned_hidden") == 1
    wb = openpyxl.load_workbook(p)
    ws = wb["Data"]
    assert ws["A2"].value == "Apple" and ws.row_dimensions[2].hidden
    assert [ws["A3"].value, ws["A4"].value] == ["Plum", "Pear"]


def _broken_by_an_older_version(tmp_path):
    """The shape set_filter used to write: a sheet filter over the table."""
    p = _book(tmp_path, "old.xlsx")
    wb = openpyxl.load_workbook(p)
    wb["Data"].auto_filter.ref = "A1:C4"
    wb.save(p)
    return p


def test_the_integrity_check_names_a_sheet_filter_over_a_table(tmp_path):
    p = _broken_by_an_older_version(tmp_path)
    problems = _verify.package_problems(p)
    assert any("Fruit" in s and "filter" in s for s in problems), problems
    assert any("_FilterDatabase" in s for s in problems), problems


def test_clear_filter_repairs_a_file_an_older_version_broke(tmp_path):
    p = _broken_by_an_older_version(tmp_path)
    out = sortfilter.clear_filter(p, {"table": "Fruit"})
    assert out["ok"]
    assert _verify.package_problems(p) == set()


def test_any_save_that_would_leave_a_sheet_filter_over_a_table_fails_closed(tmp_path, monkeypatch):
    """Fail closed: if any route produced the corruption again, the save is
    refused and the workbook is left exactly as it was."""
    from xlsx_mcp.core.errors import ValidationFailed
    p = _book(tmp_path)
    before = open(p, "rb").read()
    real = sortfilter._table_for_filter
    monkeypatch.setattr(sortfilter, "_table_for_filter", lambda ws, grid, what: None)
    with pytest.raises(ValidationFailed) as ei:
        sortfilter.set_filter(p, {"table": "Fruit"})
    assert "Fruit" in str(ei.value)
    assert open(p, "rb").read() == before
    assert real is not None
