"""Minor behaviours that real jobs exposed, with Excel as the judge where
Excel exists.

  * one formula form in every reader (=XLOOKUP, not =_xlfn.XLOOKUP)
  * recalculate's freshness note speaks the readers' label (cached)
  * pivot values take func (query_range's word) or function
  * engine='formulas' says it was removed and what to use instead
  * diagnose's route says what an edit will actually do
  * a pivot's location is the range get_pivot reads
  * a single-series chart is one colour
  * the default render takes in the charts; chart list names its ranges
  * an export of a workbook open in Excel says it used the saved copy
  * built-in styles change only what Excel's own change
  * a ceiling refusal does not blame the range's shape
  * sheet names clash without regard to case; History is reserved
  * recalculate names a circular reference
  * the 17-digit refusal no longer points at a tool that writes no cells
  * a location sheet and a different sheet argument refuse
"""

from __future__ import annotations

import zipfile

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp import server as srv
from xlsx_mcp.com import session as com_session
from xlsx_mcp.core import package as pkg_mod
from xlsx_mcp.surface import copy as copy_mod

_COM = com_session.com_available()[0]
com = pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool)")


def _call(tool, op, **args):
    payload, err, _ = fh.start().call(tool, {"op": op, "args": args})
    assert not err, payload
    return payload


def _refused(tool, op, **args):
    payload, err, _ = fh.start().call(tool, {"op": op, "args": args})
    assert err, payload
    return payload["error"]


def _sales(tmp_path, name="s.xlsx"):
    p = str(tmp_path / name)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Region", "Month", "Units"])
    for row in (("N", "Jan", 3), ("S", "Jan", 5), ("N", "Feb", 4), ("S", "Feb", 6)):
        ws.append(row)
    wb.save(p)
    return p




def test_x05_find_and_the_grid_view_show_the_formula_read_range_shows(tmp_path):
    p = _sales(tmp_path)
    _call("xl_edit", "set_formula", path=p, location="E2", sheet="Data",
          formula='=XLOOKUP("S",A2:A5,C2:C5)')
    shown = _call("xl_read", "read_range", path=p, location="E2", sheet="Data",
                  values="formula")["values"][0][0]
    assert shown == '=XLOOKUP("S",A2:A5,C2:C5)'
    found = _call("xl_read", "find_cells", path=p, query=shown, look_in="formulas",
                  match="exact")
    assert found["total"] == 1 and found["matches"][0]["formula"] == shown
    assert found["matches"][0]["formula_stored"].startswith("=_xlfn.XLOOKUP")
    view = _call("xl_read", "get_grid_view", path=p, sheet="Data", values="formula")
    assert shown in view["table"] and "_xlfn" not in view["table"]




def test_x10_pivot_values_refuse_an_unknown_key_without_excel(tmp_path):
    p = _sales(tmp_path)
    err = _refused("xl_app", "com_manage_pivot", path=p, action="create",
                   source_range="A1:C5", rows=["Region"],
                   values=[{"field": "Units", "fn": "sum"}])
    assert "func" in err["message"]




@pytest.mark.parametrize("engine", ["formulas", "FORMULAS"])
def test_x22_the_formulas_engine_says_it_was_removed(engine, tmp_path):
    p = _sales(tmp_path)
    err = _refused("xl_app", "recalculate", path=p, engine=engine)
    assert err["code"] == "CALC_UNAVAILABLE", err
    assert err["message"] == (
        "The formulas engine was removed in 2.0.0. Recalculation needs Windows with "
        "Microsoft Excel: use engine 'com' or 'auto'. Reads still return each formula's "
        "cached value, labelled cached.")
    assert "fallback" not in err["hint"], err




def test_x28_diagnose_says_what_a_cell_edit_on_a_chart_workbook_does(tmp_path):
    p = _sales(tmp_path)
    _call("xl_format", "manage_chart", path=p, action="create", chart_type="bar",
          data="C1:C5", sheet="Data")
    route = _call("xl_read", "diagnose_workbook", path=p)["routing"]["surgical_edit"]
    out = _call("xl_edit", "set_cell", path=p, location="A7", sheet="Data", value="x")
    assert out["ok"] is True
    assert route["route"] != "refuse", route




@pytest.mark.parametrize("kind,vary", [("line", "0"), ("bar", "0"), ("area", "0"),
                                       ("pie", None)])
def test_x31_a_single_series_chart_is_one_colour(kind, vary, tmp_path):
    p = _sales(tmp_path)
    _call("xl_format", "manage_chart", path=p, action="create", chart_type=kind,
          data="C1:C5", categories="A2:A5", sheet="Data")
    xml = zipfile.ZipFile(p).read("xl/charts/chart1.xml").decode()
    if vary is None:
        assert 'varyColors val="0"' not in xml
    else:
        assert 'varyColors val="0"' in xml




def test_x32_the_chart_list_names_the_ranges_it_plots(tmp_path):
    p = _sales(tmp_path)
    _call("xl_format", "manage_chart", path=p, action="create", chart_type="line",
          data="C1:C5", categories="A2:A5", sheet="Data", at_cell="F2")
    (ch,) = _call("xl_format", "manage_chart", path=p, action="list")["charts"]
    (series,) = ch["series"]
    assert series["values"] == "'Data'!$C$2:$C$5"
    assert series["categories"] == "'Data'!$A$2:$A$5"




def test_x52_an_export_of_an_open_workbook_says_it_used_the_saved_copy(tmp_path):
    p = _sales(tmp_path)
    (tmp_path / "~$s.xlsx").write_bytes(b"owner")
    out = _call("xl_workbook", "export_range", path=p, fmt="csv")
    assert any("open in Excel" in w for w in out["warnings"])
    assert "reads, exports and copies use the copy last saved" in pkg_mod.held_refusal(p)




def test_x53_good_leaves_the_number_format_and_alignment_alone(tmp_path):
    p = _sales(tmp_path)
    wb = openpyxl.load_workbook(p)
    c = wb["Data"]["C2"]
    c.number_format = "0.000"
    from openpyxl.styles import Alignment
    c.alignment = Alignment(horizontal="center")
    wb.save(p)
    _call("xl_format", "apply_style", path=p, style="Good", location="C2", sheet="Data")
    c = openpyxl.load_workbook(p)["Data"]["C2"]
    assert c.style == "Good" and c.number_format == "0.000"
    assert c.alignment.horizontal == "center"
    assert c.font.sz == 11  # the workbook's Normal size, as Excel's Good


def test_x53_currency_is_excels_own_format_and_changes_nothing_else(tmp_path):
    p = _sales(tmp_path)
    _call("xl_format", "apply_style", path=p, style="Currency", location="C2", sheet="Data")
    styles = zipfile.ZipFile(p).read("xl/styles.xml").decode()
    assert 'numFmtId="44"' in styles and '"$"_-' not in styles
    c = openpyxl.load_workbook(p)["Data"]["C2"]
    assert c.font.sz == 11


def test_x53_excels_own_names_for_headings_work(tmp_path):
    p = _sales(tmp_path)
    out = _call("xl_format", "apply_style", path=p, style="Heading 1", location="A1",
                sheet="Data")
    assert out["changed"]["style"]["name"] == "Headline 1"
    _call("xl_format", "apply_style", path=p, style="20% - Accent1", location="B1",
          sheet="Data")




def test_x55_a_ceiling_refusal_has_its_own_hint(tmp_path):
    p = _sales(tmp_path)
    err = _refused("xl_read", "read_range", path=p, location="A1:J30000", sheet="Data")
    assert err["code"] == "RANGE_OUT_OF_BOUNDS"
    assert err["hint"] == copy_mod.CEILING_HINT and "inverted" not in err["hint"]




def test_x57_sheet_names_clash_without_regard_to_case(tmp_path):
    p = str(tmp_path / "n.xlsx")
    srv.create_workbook(p, sheets=["Summary", "Data"])
    for args in ({"action": "add", "new_name": "summary"},
                 {"action": "rename", "sheet": "Data", "new_name": "SUMMARY"},
                 {"action": "copy", "sheet": "Data", "new_name": "summary"},
                 {"action": "add", "new_name": "History"}):
        err = _refused("xl_structure", "manage_worksheet", path=p, **args)
        assert "Summary" in err["message"] or "History" in err["message"]
    # A sheet may change only the case of its own name.
    _call("xl_structure", "manage_worksheet", path=p, action="rename", sheet="Summary",
          new_name="SUMMARY")
    assert openpyxl.load_workbook(p).sheetnames == ["SUMMARY", "Data"]
    err = _refused("xl_workbook", "create_workbook", path=str(tmp_path / "m.xlsx"),
                   sheets=["A", "a"])
    assert "duplicate" in err["message"]




def test_x60_a_17_digit_double_is_stored_whole(tmp_path):
    # A cell edit splices the cell
    # into the original package, which stores all 17 significant digits, so
    # the write that used to refuse the whole batch now simply lands.
    p = _sales(tmp_path)
    out = _call("xl_edit", "write_range", path=p, location="E2", sheet="Data",
                data=[[0.1 + 0.2]])
    assert out["ok"] is True
    assert openpyxl.load_workbook(p)["Data"]["E2"].value == 0.1 + 0.2


@pytest.mark.usefixtures("rewrite_route")
def test_x60_the_precision_refusal_gives_advice_that_works(tmp_path):
    p = _sales(tmp_path)
    err = _refused("xl_edit", "write_range", path=p, location="E2", sheet="Data",
                   data=[[0.1 + 0.2]])
    assert "15 significant digits" in err["message"] and "xl_app" not in err["message"]




def test_x61_a_location_sheet_and_a_different_sheet_refuse(tmp_path):
    p = _sales(tmp_path)
    wb = openpyxl.load_workbook(p)
    wb.create_sheet("Other")
    wb.save(p)
    err = _refused("xl_edit", "set_cell", path=p, location={"cell": "A1", "sheet": "Data"},
                   sheet="Other", value=1)
    assert "Data" in err["message"] and "Other" in err["message"]
    _call("xl_edit", "set_cell", path=p, location={"cell": "A9", "sheet": "Data"},
          sheet="Data", value=1)  # the same sheet twice is fine


# ------------------------------------------------------------- Excel judges


@com
@pytest.mark.com
@pytest.mark.timeout(300)
def test_x59_x07_recalculate_names_a_circular_reference(tmp_path):
    p = _sales(tmp_path)
    srv.set_formula(p, "E2", "=F2+1", sheet="Data")
    srv.set_formula(p, "F2", "=E2+1", sheet="Data")
    out = srv.recalculate(p, engine="com")
    assert out["circular_references"] and out["circular_references"][0].startswith("Data!")
    assert any("circular reference" in w for w in out["warnings"])
    assert "labelled cached" in out["freshness"]


@com
@pytest.mark.com
@pytest.mark.timeout(300)
def test_x29_the_pivot_location_is_the_range_get_pivot_reads(tmp_path):
    p = _sales(tmp_path)
    out = srv.com_manage_pivot(p, "create", source_range="A1:C5", dest_sheet="P",
                               dest_cell="A3", rows=["Region"], filters=["Month"],
                               values=[{"field": "Units", "function": "sum"}])
    made = out["changed"]
    (pv,) = srv.get_pivot(p)["pivots"]
    assert made["location"] == pv["location"] and "$" not in made["location"]
    assert made["location_with_filters"].startswith("A1")


@com
@pytest.mark.com
@pytest.mark.timeout(300)
def test_x32_the_default_render_takes_in_the_chart(tmp_path):
    from xlsx_mcp.ops import comtier
    if not comtier._clipboard_available():
        pytest.skip("rendering needs the Windows clipboard, which this session lacks")
    p = _sales(tmp_path)
    srv.manage_chart(p, "create", chart_type="line", data={"range": "C1:C5"},
                     sheet="Data", at_cell="H2")
    out = srv.com_render_sheet(p, str(tmp_path / "r.png"), sheet="Data")
    last_col = out["range"].split(":")[1].rstrip("0123456789")
    assert last_col >= "H" and len(last_col) == 1, out["range"]


@com
@pytest.mark.com
@pytest.mark.timeout(300)
def test_x53_excel_shows_currency_as_its_own_currency_style(tmp_path):
    """Excel's Range.Text for a cell styled Currency by the product equals
    Excel's own Currency style on the same value."""
    p = _sales(tmp_path)
    _call("xl_format", "apply_style", path=p, style="Currency", location="C2", sheet="Data")

    def job(manager):
        w = manager.acquire()
        prior = com_session._assert_hygiene(w.app)
        wb = fresh = None
        try:
            wb = com_session.open_workbook(w.app, p)
            ws = wb.Worksheets("Data")
            ws.Columns(3).ColumnWidth = 14
            ours = str(ws.Range("C2").Text)
            # The reference is Excel's own Currency style in a NEW workbook
            # (the file's own style would echo whatever the product wrote).
            fresh = w.app.Workbooks.Add()
            ref = fresh.Worksheets(1)
            ref.Columns(3).ColumnWidth = 14
            ref.Range("C2").Value = 3
            ref.Range("C2").Style = "Currency"
            return ours, str(ref.Range("C2").Text), str(ws.Range("C2").Style.Name)
        finally:
            if fresh is not None:
                fresh.Close(False)
            if wb is not None:
                wb.Close(False)
            com_session._restore_hygiene(w.app, *prior)
            manager.release_to_pool(w)

    ours, theirs, name = com_session.run_com("oracle", job, timeout=120)
    # The fill padding (*) depends on the font's metrics (the file's Calibri,
    # the new workbook's theme font), so spacing is compared collapsed.
    assert ours.split() == theirs.split() and name == "Currency", (ours, theirs)
