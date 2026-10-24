"""Structured-reference rename and storage defects. Red before the fix. The
assertions here read the saved file so they need no Excel."""

from __future__ import annotations

import re
import zipfile

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp import server as srv


def _table_book(p, header="v") -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["D1"], ws["E1"], ws["D2"], ws["E2"], ws["D3"], ws["E3"] = "k", header, "a", 1, "b", 2
    wb.save(p)
    srv.create_table(str(p), "D1:E3", "T")
    return str(p)


def _formula_xml(p, cell) -> str:
    xml = zipfile.ZipFile(p).read("xl/worksheets/sheet1.xml").decode()
    m = re.search(rf'<c r="{cell}"[^>]*>.*?<f>(.*?)</f>', xml)
    return m.group(1) if m else ""


# The F-T rename guard looks for the old column name in brackets
# with re.escape of the RAW name, so a structured reference that escapes a
# special character the way Excel requires (T[x'#], T[a'[b]) is not seen:
# the header rename is accepted and the formula, which Excel computed as 3
# before, opens as #REF! (Excel verdict, 2 of 2). The product does not rewrite
# formulas on a rename, so the guard is the only protection.

@pytest.mark.parametrize("header,ref", [("x#", "T[x'#]"), ("a[b", "T[a'[b]")])
def test_ft2_a_rename_never_orphans_an_escaped_structured_reference(header, ref, tmp_path):
    p = _table_book(tmp_path / "t.xlsx", header=header)
    srv.set_formula(p, "G1", f"=SUM({ref})")
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_cell", "args": {
        "path": p, "location": "E1", "value": "y"}})
    if not err:
        assert ref.split("[", 1)[1].rstrip("]") not in _formula_xml(p, "G1"), \
            f"renamed {header!r} to 'y' while G1 still says {_formula_xml(p, 'G1')!r}"


# F-S (MAJOR, pre-existing on 1.2.5) set_formula stores Excel's UI shorthand
# for structured references as typed ([@v], and spaces inside brackets). The
# file format needs T[[#This Row],[v]]; Excel opens the stored text as
# =#REF!*2 and keeps it (verdict: =T[@v]*2 -> "=#REF!*2"), while the product
# answered verified: true with no warning.

@pytest.mark.parametrize("formula", ["=T[@v]*2", "=SUM(T[ v ])"])
def test_fs_a_structured_reference_is_stored_in_file_form_or_refused(formula, tmp_path):
    p = _table_book(tmp_path / "t.xlsx")
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_formula", "args": {
        "path": p, "location": "G2", "formula": formula}})
    if not err:
        stored = _formula_xml(p, "G2")
        assert "[@" not in stored and "[ " not in stored, \
            f"stored {stored!r}, which Excel opens as #REF!"
