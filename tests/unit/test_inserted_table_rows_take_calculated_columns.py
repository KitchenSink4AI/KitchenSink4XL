"""A row inserted inside a table takes the table's calculated columns.

Excel's own Insert of a sheet row inside a table's data rows (or at its
totals row) puts the column formula into every calculated column of the new
row, and a new table row added with ListRows.Add does the same. The product
left those cells empty, so a Total column showed nothing for the new row
until someone copied the formula down. A row inserted at the header row or
below the last row is not the table's, and gets nothing. Expected results
are Excel's own, measured with Excel's Insert and ListRows.Add. All data
here is made up.
"""

from __future__ import annotations

import openpyxl
import pytest
from openpyxl.worksheet.table import Table, TableColumn, TableFormula

from xlsx_mcp.ops import structure, tables

TOTAL = "Sales[[#This Row],[Qty]]*Sales[[#This Row],[Price]]"


def _book(tmp_path, *, totals=False, rel=False, plain=False):
    p = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    headers = ["Item", "Qty", "Price", "Total"]
    for j, h in enumerate(headers, start=1):
        ws.cell(1, j, h)
    for i, (a, b, c) in enumerate([("Widget", 4, 2.5), ("Gadget", 3, 4.0),
                                   ("Doohickey", 8, 1.25)], start=2):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)
        ws.cell(i, 3, c)
        if plain:
            ws.cell(i, 4, "=B%d*C%d" % (i, i))
        elif rel:
            ws.cell(i, 4, "=B%d*C%d" % (i, i))
        else:
            ws.cell(i, 4, "=" + TOTAL)
    last = 4
    if totals:
        ws.cell(5, 1, "Total")
        ws.cell(5, 4, "=SUBTOTAL(109,Sales[Total])")
        last = 5
    t = Table(displayName="Sales", ref=f"A1:D{last}",
              totalsRowCount=1 if totals else None)
    cols = [TableColumn(id=j, name=h) for j, h in enumerate(headers, start=1)]
    if not plain:
        cols[3].calculatedColumnFormula = TableFormula(
            attr_text="B2*C2" if rel else TOTAL)
    if totals:
        cols[0].totalsRowLabel = "Total"
        cols[3].totalsRowFunction = "sum"
    t.tableColumns = cols
    ws.add_table(t)
    wb.save(p)
    return str(p)


def _state(path):
    wb = openpyxl.load_workbook(path)
    try:
        ws = wb["Data"]
        return ws.tables["Sales"].ref, {
            r: ws.cell(r, 4).value for r in range(1, 9)}
    finally:
        wb.close()


@pytest.mark.parametrize("at,count", [(2, 1), (3, 1), (4, 1), (3, 2)])
def test_a_row_inserted_inside_the_data_takes_the_column_formula(
        tmp_path, at, count):
    p = _book(tmp_path)
    structure.modify_grid_structure(p, "insert_rows", at, count=count,
                                    sheet="Data")
    ref, col = _state(p)
    assert ref == f"A1:D{4 + count}"
    for r in range(2, 5 + count):
        assert col[r] == "=" + TOTAL, (r, col[r])


@pytest.mark.parametrize("at,ref", [(1, "A2:D5"), (5, "A1:D4")])
def test_a_row_at_the_header_or_below_the_table_is_not_the_tables(
        tmp_path, at, ref):
    p = _book(tmp_path)
    structure.modify_grid_structure(p, "insert_rows", at, sheet="Data")
    got_ref, col = _state(p)
    assert got_ref == ref
    assert col[at] is None


def test_a_row_inserted_at_the_totals_row_joins_the_data(tmp_path):
    p = _book(tmp_path, totals=True)
    structure.modify_grid_structure(p, "insert_rows", 5, sheet="Data")
    ref, col = _state(p)
    assert ref == "A1:D6"
    assert col[5] == "=" + TOTAL
    assert col[6] == "=SUBTOTAL(109,Sales[Total])"


def test_a_column_formula_with_cell_references_follows_the_row(tmp_path):
    p = _book(tmp_path, rel=True)
    structure.modify_grid_structure(p, "insert_rows", 3, sheet="Data")
    _ref, col = _state(p)
    assert col[3] == "=B3*C3"
    assert col[4] == "=B4*C4" and col[5] == "=B5*C5"


def test_a_table_with_no_calculated_column_gets_nothing(tmp_path):
    p = _book(tmp_path, plain=True)
    structure.modify_grid_structure(p, "insert_rows", 3, sheet="Data")
    _ref, col = _state(p)
    assert col[3] is None


def test_a_new_table_row_takes_the_column_formula_unless_given_a_value(
        tmp_path):
    p = _book(tmp_path)
    tables.manage_table(p, "Sales", "add_row",
                        values=[["Gizmo", 2, 5.0], ["Thing", 1, 1.0, 42]])
    ref, col = _state(p)
    assert ref == "A1:D6"
    assert col[5] == "=" + TOTAL
    assert col[6] == 42


def test_a_new_row_under_a_totals_row_table_takes_it_too(tmp_path):
    p = _book(tmp_path, totals=True)
    tables.manage_table(p, "Sales", "add_row", values=["Gizmo", 2, 5.0])
    ref, col = _state(p)
    assert ref == "A1:D6"
    assert col[5] == "=" + TOTAL
