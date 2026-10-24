"""manage_table to_range refuses, before any change, when a conditional-format or data-validation rule names the
table, and the refusal says which kind of rule it is and which cells it covers, so it can be found on a sheet with
many rules."""
from __future__ import annotations

import openpyxl
import pytest
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table


@pytest.mark.parametrize("kind", ["conditional format", "data validation"])
def test_rule_refusal_names_the_cells_and_the_kind(tmp_path, kind):
    from xlsx_mcp.ops import tables
    p = str(tmp_path / "sales.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in [["Region", "Q1", "Q2"], ["North", 1, 2], ["South", 3, 4], ["East", 5, 6]]:
        ws.append(r)
    ws.add_table(Table(displayName="Sales", ref="A1:C4"))
    if kind == "conditional format":
        ws.conditional_formatting.add("B2:B4", FormulaRule(formula=["B2>AVERAGE(Sales[Q1])"],
                                                           fill=PatternFill("solid", fgColor="FFFF00")))
    else:
        dv = DataValidation(type="custom", formula1="B2<=MAX(Sales[Q2])")
        dv.add("B2:B4")
        ws.add_data_validation(dv)
    wb.save(p)
    before = open(p, "rb").read()
    with pytest.raises(Exception) as ei:
        tables.manage_table(p, "Sales", "to_range")
    msg = str(ei.value)
    assert "Data!B2:B4" in msg, f"the refusal does not name the cells the rule covers: {msg}"
    assert kind in msg and "conditional format, data validation or chart" not in msg, msg
    assert open(p, "rb").read() == before


def test_rule_refusal_quotes_a_sheet_name_that_needs_it(tmp_path):
    from xlsx_mcp.ops import tables
    p = str(tmp_path / "sales.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q1 Sales"
    for r in [["Region", "Q1"], ["North", 1], ["South", 3]]:
        ws.append(r)
    ws.add_table(Table(displayName="Sales", ref="A1:B3"))
    dv = DataValidation(type="custom", formula1="B2<=MAX(Sales[Q1])")
    dv.add("B2:B3")
    ws.add_data_validation(dv)
    wb.save(p)
    with pytest.raises(Exception) as ei:
        tables.manage_table(p, "Sales", "to_range")
    assert "'Q1 Sales'!B2:B3" in str(ei.value)


def test_rule_refusal_names_two_areas_with_a_comma_as_excel_writes_them(tmp_path):
    from xlsx_mcp.ops import tables
    p = str(tmp_path / "sales.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in [["Region", "Q1", "Q2"], ["North", 1, 2], ["South", 3, 4], ["East", 5, 6]]:
        ws.append(r)
    ws.add_table(Table(displayName="Sales", ref="A1:C4"))
    ws.conditional_formatting.add("B2:B4 C2:C4", FormulaRule(formula=["B2>AVERAGE(Sales[Q1])"],
                                                             fill=PatternFill("solid", fgColor="FFFF00")))
    wb.save(p)
    with pytest.raises(Exception) as ei:
        tables.manage_table(p, "Sales", "to_range")
    assert "Data!B2:B4,C2:C4" in str(ei.value), str(ei.value)
