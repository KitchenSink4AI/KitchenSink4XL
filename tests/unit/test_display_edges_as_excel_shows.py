"""Display and value edges, as Excel shows them.

  - General shows at most 11 characters (12 with a minus): 0.123456789012
    shows 0.123456789 and 9007199254740993 shows 9.0072E+15; a denormal
    number shows 0. Exports keep 15 digits: they are data, not display.
  - _x in a format leaves the width of x empty, and Excel's text holds one
    space there (0.00_) shows 1234.5678 as "1234.57 ").
  - A date past 31 December 9999 fills the cell with # in Excel; a read of
    its value gave the error #VALUE!, though the cell holds the number.
  - In a workbook that counts dates from 1904, serial 0 is 1 January 1904,
    and a negative date or time shows with a minus (-1 shows -1904-01-02).
  - A duration ([h]:mm) read as Python's "1 day, 18:00:00"; it reads as
    Excel's elapsed time, 42:00:00.

Expected texts are Excel's own Range.Text, measured in Excel. Made-up data.
"""

from __future__ import annotations

import openpyxl
import pytest

from xlsx_mcp.core import numfmt
from xlsx_mcp.ops import cells

#: (value, format, Excel's text)
EXCEL = [
    (9007199254740993, "General", "9.0072E+15"),
    (0.123456789012, "General", "0.123456789"),
    (5e-324, "General", "0"),
    (123456789012, "General", "1.23457E+11"),
    (123456789.123, "General", "123456789.1"),
    (-1.234e-06, "General", "-0.000001234"),
    (1e-10, "General", "1E-10"),
    (0.30000000000000004, "General", "0.3"),
    (1e15, "General", "1E+15"),
    (12345678901, "General", "12345678901"),
    (1234.5678, "0.00_)", "1234.57 "),
    (12, "0_);(0)", "12 "),
    (-12, "0_);(0)", "(12)"),
    (2958466, "yyyy-mm-dd", "########"),
]


@pytest.mark.parametrize("value,fmt,excel", EXCEL,
                         ids=[f"{v}|{f}" for v, f, _ in EXCEL])
def test_display_matches_excel(value, fmt, excel):
    assert numfmt.display(value, fmt) == excel


@pytest.mark.parametrize("value,fmt,excel", [
    (0, "yyyy-mm-dd", "1904-01-01"),
    (-1, "yyyy-mm-dd", "-1904-01-02"),
    (1, "yyyy-mm-dd", "1904-01-02"),
    (0.25, "[h]:mm", "6:00"),
    (-0.25, "[h]:mm", "-6:00"),
    (46073, "yyyy-mm-dd", "2030-02-21"),
])
def test_the_1904_date_system(value, fmt, excel):
    assert numfmt.display(value, fmt, date1904=True) == excel


def test_exports_keep_fifteen_digits():
    assert numfmt.general(0.123456789012) == "0.123456789012"


def _book(tmp_path, date1904=False):
    p = tmp_path / "edges.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 2958466
    ws["A1"].number_format = "yyyy-mm-dd"
    ws["A2"] = 1.75
    ws["A2"].number_format = "[h]:mm"
    ws["A3"] = 0
    ws["A3"].number_format = "yyyy-mm-dd"
    ws["A4"] = -0.25
    ws["A4"].number_format = "[h]:mm"
    if date1904:
        wb.epoch = openpyxl.utils.datetime.MAC_EPOCH
    wb.save(p)
    return str(p)


def test_a_number_past_the_last_date_reads_as_the_number(tmp_path):
    p = _book(tmp_path)
    got = cells.read_range(p, "A1:A2", sheet="Data")["values"]
    assert got == [[2958466], ["42:00:00"]]
    shown = cells.read_range(p, "A1:A2", sheet="Data",
                             values="displayed")["values"]
    assert shown == [["########"], ["42:00"]]


def test_a_1904_workbook_displays_as_excel_does(tmp_path):
    p = _book(tmp_path, date1904=True)
    shown = cells.read_range(p, "A3:A4", sheet="Data",
                             values="displayed")["values"]
    assert shown == [["1904-01-01"], ["-6:00"]]
