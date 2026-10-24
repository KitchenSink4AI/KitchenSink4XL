"""A relative defined name keeps Excel's meaning through inserts and deletes.

Excel stores a relative reference in a defined name relative to cell A1 and
resolves it from the cell that uses the name: "the cell above" is stored as
Data!A1048576 and "the cell to the left" as Data!XFD1. They are offsets, not
positions, so Excel's own row and column inserts and deletes leave them as
they are (measured with Excel's own edits, the oracle for this file). Only
the absolute ($) parts of a name are positions on the sheet, and only they
move.

The rewriter treated every part of a name as a position: an insert anywhere
on the sheet turned Data!A1048576 into Data!A1048577 (a row past the grid)
and Data!XFD1 into Data!XFE1, and a delete turned "the cell above" into "two
cells above", so a formula using the name read the wrong cell (a workbook's
=Prev read 0 where Excel shows 64). All data here is made up.
"""

from __future__ import annotations

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName

from xlsx_mcp.ops import structure

NAMES = {
    "Above": "Data!A1048576",        # the cell above, same column
    "Left": "Data!XFD1",             # the cell to the left, same row
    "Left3": "Data!XFB1:XFD1",       # the three cells to the left
    "AboveAbsCol": "Data!$A1048576",  # column A fixed, the row above
    "TopRowRel": "Data!A$1",         # row 1 fixed, the same column
    "Plain": "Data!$A$1:$A$6",       # fully absolute: a position
}


def _book(tmp_path):
    p = tmp_path / "relnames.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, 9):
        ws.cell(r, 1, r * 10)
        ws.cell(r, 2, r)
    ws["C3"] = "=Above"
    ws["C4"] = "=Left"
    ws["D5"] = "=SUM(Left3)"
    ws["E4"] = "=AboveAbsCol"
    ws["F2"] = "=TopRowRel"
    ws["G2"] = "=SUM(Plain)"
    calc = wb.create_sheet("Calc")
    calc["B5"] = "=Above"
    for name, text in NAMES.items():
        wb.defined_names.add(DefinedName(name=name, attr_text=text))
    wb.save(p)
    return str(p)


def _names(path):
    wb = openpyxl.load_workbook(path)
    try:
        return {n: wb.defined_names[n].attr_text for n in NAMES}
    finally:
        wb.close()


def _expect(**changed):
    out = dict(NAMES)
    out.update(changed)
    return out


CASES = [
    # (action, at, count, sheet, the names Excel stores afterwards)
    ("insert_rows", 10, 1, "Data", _expect()),               # below the data
    ("delete_rows", 10, 1, "Data", _expect()),               # an empty row
    ("insert_rows", 3, 1, "Data", _expect(Plain="Data!$A$1:$A$7")),
    ("delete_rows", 2, 1, "Data", _expect(Plain="Data!$A$1:$A$5")),
    ("insert_rows", 1, 1, "Data", _expect(TopRowRel="Data!A$2",
                                          Plain="Data!$A$2:$A$7")),
    ("insert_rows", 4, 2, "Data", _expect(Plain="Data!$A$1:$A$8")),
    ("insert_cols", "C", 1, "Data", _expect()),
    ("delete_cols", "B", 1, "Data", _expect()),
    ("insert_cols", "A", 1, "Data", _expect(AboveAbsCol="Data!$B1048576",
                                            Plain="Data!$B$1:$B$6")),
    ("insert_rows", 2, 1, "Calc", _expect()),                # another sheet
]


@pytest.mark.parametrize("action,at,count,sheet,want", CASES,
                         ids=[f"{c[0]}_{c[1]}x{c[2]}_{c[3]}" for c in CASES])
def test_names_keep_their_relative_offsets(tmp_path, action, at, count,
                                           sheet, want):
    p = _book(tmp_path)
    structure.modify_grid_structure(p, action, at, count=count, sheet=sheet)
    assert _names(p) == want


def test_deleting_an_empty_row_below_the_data_leaves_the_cell_above_alone(
        tmp_path):
    """The corpus shape: a name for "the cell above", and a row deleted far
    below the data. Nothing the name points at moved."""
    p = _book(tmp_path)
    structure.modify_grid_structure(p, "delete_rows", 10, sheet="Data")
    assert _names(p)["Above"] == "Data!A1048576"
    wb = openpyxl.load_workbook(p)
    try:
        assert wb["Data"]["C3"].value == "=Above"
        assert wb["Calc"]["B5"].value == "=Above"
    finally:
        wb.close()


def test_an_absolute_name_still_follows_its_cells(tmp_path):
    p = _book(tmp_path)
    structure.modify_grid_structure(p, "delete_rows", 1, sheet="Data")
    names = _names(p)
    assert names["Plain"] == "Data!$A$1:$A$5"
    assert names["TopRowRel"] == "Data!#REF!"   # its fixed row was deleted
    assert names["Above"] == "Data!A1048576"


def test_no_name_is_ever_written_past_the_grid(tmp_path):
    """Inserting a column used to write Data!XFE1, a column that does not
    exist, into the name for "the cell to the left"."""
    p = _book(tmp_path)
    structure.modify_grid_structure(p, "insert_cols", "A", sheet="Data")
    structure.modify_grid_structure(p, "insert_rows", 1, sheet="Data")
    for text in _names(p).values():
        assert "XFE" not in text and "1048577" not in text
