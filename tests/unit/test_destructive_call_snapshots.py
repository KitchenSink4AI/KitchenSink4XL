"""A destructive call keeps a
labelled copy of the workbook as it was, in the workbook's own backup
folder, and names it in its result; manage_backups action='restore'
source=<label> brings it back after any number of later edits (the prev
slot those edits rotate away is no longer the only undo).
"""

from __future__ import annotations

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp import server as srv
from xlsx_mcp.ops import backups


def _call(tool, op, **args):
    payload, err, _ = fh.start().call(tool, {"op": op, "args": args})
    assert not err, payload
    return payload


def _book(tmp_path):
    p = str(tmp_path / "s.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, 6):
        ws.append([f"r{r}", r, r * 10])
    wb.create_sheet("Notes")["A1"] = "keep"
    from openpyxl.worksheet.table import Table
    ws2 = wb.create_sheet("T")
    ws2.append(["Name", "Qty"])
    ws2.append(["a", 1])
    ws2.append(["b", 2])
    ws2.add_table(Table(displayName="Tbl", ref="A1:B3"))
    wb.save(p)
    return p


def _state(p):
    wb = openpyxl.load_workbook(p)
    return {ws.title: [[c.value for c in row] for row in ws.iter_rows()]
            for ws in wb.worksheets}, sorted(wb.sheetnames)


DESTRUCTIVE = [
    ("xl_structure", "manage_worksheet", {"action": "delete", "sheet": "Notes"}),
    ("xl_structure", "modify_grid_structure",
     {"action": "delete_rows", "at": 2, "count": 2, "sheet": "Data"}),
    ("xl_structure", "manage_table", {"name": "Tbl", "action": "delete_row", "index": 1}),
    ("xl_edit", "clear_range", {"location": "A1:C5", "sheet": "Data"}),
    ("xl_format", "set_merge",
     {"action": "merge", "location": "A1:C1", "sheet": "Data", "confirm_data_loss": True}),
    ("xl_edit", "replace_cells", {"find": "r", "replace": "x", "sheet": "Data"}),
    ("xl_workbook", "set_protection", {"action": "sheet", "sheet": "Data"}),
]


@pytest.mark.parametrize("tool,op,args", DESTRUCTIVE, ids=[d[1] for d in DESTRUCTIVE])
def test_b11_the_label_restores_the_exact_pre_op_workbook(tool, op, args, tmp_path):
    p = _book(tmp_path)
    before = _state(p)
    out = _call(tool, op, path=p, **args)
    label = out["snapshot"]
    assert label.startswith(f"pre-{op} ")
    # Three later edits: prev now holds none of the pre-op state.
    for i in range(3):
        _call("xl_edit", "set_cell", path=p, location="E1", sheet="T", value=i)
    if op == "set_protection":
        _call("xl_workbook", "set_protection", path=p, action="remove", sheet="Data")
    listed = _call("xl_workbook", "manage_backups", action="list", path=p)
    assert label in [s["label"] for s in listed["snapshots"]]
    _call("xl_workbook", "manage_backups", action="restore", path=p, source=label)
    assert _state(p) == before
    if op == "set_protection":
        assert not openpyxl.load_workbook(p)["Data"].protection.sheet


def test_b11_the_snapshot_lives_in_the_backup_folder_not_beside_the_workbook(tmp_path):
    p = _book(tmp_path)
    out = _call("xl_edit", "clear_range", path=p, location="A1", sheet="Data")
    assert sorted(x.name for x in tmp_path.iterdir()) == [".ks4xl-backups", "s.xlsx"]
    (entry,) = _call("xl_workbook", "manage_backups", action="list", path=p)["snapshots"]
    assert entry["label"] == out["snapshot"] and ".ks4xl-backups" in entry["path"]


@pytest.mark.parametrize("tool,op,args", [
    ("xl_edit", "set_cell", {"location": "A1", "sheet": "Data", "value": 1}),
    ("xl_structure", "manage_worksheet", {"action": "rename", "sheet": "Notes",
                                          "new_name": "N2"}),
    ("xl_structure", "modify_grid_structure", {"action": "insert_rows", "at": 2,
                                               "sheet": "Data"}),
    ("xl_edit", "replace_cells", {"find": "r", "replace": "x", "sheet": "Data",
                                  "dry_run": True}),
    ("xl_workbook", "set_protection", {"action": "status"}),
])
def test_b11_ordinary_edits_take_no_snapshot(tool, op, args, tmp_path):
    p = _book(tmp_path)
    out = _call(tool, op, path=p, **args)
    assert "snapshot" not in out
    assert _call("xl_workbook", "manage_backups", action="list", path=p)["snapshots"] == []


def test_b11_a_refused_call_leaves_no_snapshot(tmp_path):
    p = _book(tmp_path)
    payload, err, _ = fh.start().call("xl_format", {"op": "set_merge", "args": {
        "path": p, "action": "merge", "location": "A1:C1", "sheet": "Data"}})
    assert err  # absorbed values without confirm_data_loss refuse
    assert _call("xl_workbook", "manage_backups", action="list", path=p)["snapshots"] == []


def test_b11_the_newest_twenty_are_kept(tmp_path, monkeypatch):
    monkeypatch.setattr(backups, "AUTO_SNAPSHOT_KEEP", 3)
    p = _book(tmp_path)
    labels = [_call("xl_edit", "clear_range", path=p, location=f"A{i}", sheet="Data")["snapshot"]
              for i in range(1, 6)]
    kept = [s["label"] for s in
            _call("xl_workbook", "manage_backups", action="list", path=p)["snapshots"]]
    assert kept == labels[-3:]


def test_b11_an_unknown_source_names_the_choices(tmp_path):
    p = _book(tmp_path)
    payload, err, _ = fh.start().call("xl_workbook", {"op": "manage_backups", "args": {
        "action": "restore", "path": p, "source": "pre-clear_range nope"}})
    assert err and "snapshot" in payload["error"]["message"]


def test_b11_the_engine_call_alone_takes_no_snapshot(tmp_path):
    """Snapshots are a surface behaviour (the dispatcher arms them); direct
    engine calls, as the older tests make, are unchanged."""
    p = _book(tmp_path)
    out = srv.clear_range(p, "A1", sheet="Data")
    assert "snapshot" not in out
