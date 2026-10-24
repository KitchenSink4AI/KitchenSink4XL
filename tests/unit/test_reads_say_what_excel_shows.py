"""A read says what Excel shows where that differs from the stored value.

A weak model asked for the text Excel shows in a cell read the stored value
and reported it: 0.75 under the format # ?/? was told as "0.75", where Excel
shows " 3/4". The shown text was one argument away (values="displayed"), but
no help text, op line or argument list named that argument. Now:

  * read_range's help and op line name values="displayed", the other reads'
    help names it, and every read op that takes values accepts it by name;
  * the default reads (read_range, get_cells, get_grid_view, find_cells and
    get_table) add `shows` for each cell whose shown text differs from the
    value they return, and cells that read the same add nothing.

The shown text is held to Excel's own: the number-display fixture
(tests/fixtures/numfmt/excel_range_text.json, values under number formats
with the Range.Text a private hidden Excel showed) is written into a
workbook and read back with a default read, cell by cell. All content is
made up.
"""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.table import Table

import family_harness as fh
from xlsx_mcp.ops import cells, search, tables, view
from xlsx_mcp.surface import normalize

ORACLE = json.loads((Path(__file__).resolve().parents[1] / "fixtures" / "numfmt"
                     / "excel_range_text.json").read_text("utf-8"))


def _shown_book(tmp_path: Path) -> str:
    """The reported case: B2 0.75 under # ?/?, B3 1.25 under [h]:mm:ss."""
    p = tmp_path / "shown.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Times"
    ws["A1"], ws["B1"] = "What", "Amount"
    ws["A2"], ws["A3"] = "Share", "Hours"
    ws["B2"] = 0.75
    ws["B2"].number_format = "# ?/?"
    ws["B3"] = 1.25
    ws["B3"].number_format = "[h]:mm:ss"
    ws.add_table(Table(displayName="Amounts", ref="A1:B3"))
    wb.save(p)
    return str(p)


def _call(tool, op, **args):
    payload, err, _ = fh.start().call(tool, {"op": op, "args": args})
    assert not err, payload
    return payload


# ------------------------------------------------- the option has a name


def test_the_help_and_the_op_line_name_the_displayed_mode():
    payload, err, _ = fh.start().call("xl_help", {"op": "read_range"})
    assert not err, payload
    assert "'displayed'" in payload["detail"] and "shows" in payload["detail"], payload["detail"]
    assert "'displayed'" in payload["summary"], payload["summary"]
    tools = {t.name: t for t in fh.start().list_tools()}
    read = tools["xl_read"]
    text = read.description + json.dumps(read.inputSchema)
    line = next(ln for ln in text.replace("\\n", "\n").splitlines()
                if ln.startswith("- read_range("))
    assert "values='displayed'" in line, line
    for op in ("get_cells", "get_table", "query_range", "get_grid_view"):
        payload, err, _ = fh.start().call("xl_help", {"op": op})
        assert not err and "displayed" in payload["detail"], (op, payload)
    payload, err, _ = fh.start().call("xl_help", {"op": "find_cells"})
    assert not err and "shows" in payload["detail"], payload


@pytest.mark.parametrize("op", ["read_range", "get_cells", "get_table", "query_range",
                                "export_range", "export_file"])
def test_every_op_that_takes_displayed_lists_it_and_routes_its_spellings(op):
    assert "displayed" in normalize.VOCAB[(op, "values")]
    args, sent = normalize.normalize_args(op, {"values": "Displayed"}, ("values",))
    assert args["values"] == "displayed" and sent == {"values": "Displayed"}


def test_the_grid_views_refusal_of_displayed_points_at_the_reads_that_take_it(tmp_path):
    payload, err, _ = fh.start().call("xl_read", {"op": "get_grid_view", "args": {
        "path": _shown_book(tmp_path), "values": "displayed"}})
    assert err and "values='displayed'" in payload["error"]["message"], payload


# ------------------------------------------------- default reads say it


def test_a_default_read_says_what_a_fraction_shows(tmp_path):
    out = _call("xl_read", "read_range", path=_shown_book(tmp_path), location="B2:B3")
    assert out["values"] == [[0.75], ["30:00:00"]]
    assert out["shows"] == {"B2": " 3/4"}


def test_every_default_read_says_it(tmp_path):
    p = _shown_book(tmp_path)
    assert cells.read_range(p, "B2:B3", values="both")["shows"] == {"B2": " 3/4"}
    assert cells.get_cells(p, ["B2", "B3"])["shows"] == {"B2": " 3/4"}
    assert cells.get_cells(p, ["Times!B2", "A2"])["shows"] == {"B2": " 3/4"}
    assert view.get_grid_view(p)["shows"] == {"B2": " 3/4"}
    hit = search.find_cells(p, "0.75")["matches"][0]
    assert hit["cell"] == "B2" and hit["value"] == "0.75" and hit["shows"] == " 3/4"
    assert tables.get_table(p, "Amounts")["shows"] == {"B2": " 3/4"}


def test_a_formula_result_says_it_too_and_a_missing_result_says_nothing(tmp_path):
    p = tmp_path / "f.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"] = 3, 4
    ws["B1"] = "=A1/A2"
    ws["B1"].number_format = "0.0%"
    ws["B2"] = "=A1*A2"
    ws["B2"].number_format = "#,##0.00"
    wb.save(p)
    # B1 carries a cached result, as Excel leaves one; B2 has none yet
    import re
    import zipfile
    with zipfile.ZipFile(p) as z:
        raw = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    raw, n = re.subn(r"<f>A1/A2</f>(?:<v\s*/>|<v></v>)?", "<f>A1/A2</f><v>0.75</v>", raw)
    assert n == 1, raw
    q = tmp_path / "g.xlsx"
    with zipfile.ZipFile(p) as src, zipfile.ZipFile(q, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = raw.encode("utf-8") if item.filename == "xl/worksheets/sheet1.xml" \
                else src.read(item.filename)
            dst.writestr(item, data)
    out = cells.read_range(str(q), "B1:B2")
    assert out["values"] == [[0.75], [None]]
    assert out["shows"] == {"B1": "75.0%"}


def test_a_read_across_sheets_names_each_cells_sheet(tmp_path):
    p = _shown_book(tmp_path)
    wb = openpyxl.load_workbook(p)
    other = wb.create_sheet("Other")
    other["A1"] = 0.5
    other["A1"].number_format = "0%"
    wb.save(p)
    out = cells.get_cells(p, ["Times!B2", "Other!A1", "Times!A2"])
    assert out["shows"] == {"Times!B2": " 3/4", "Other!A1": "50%"}


def test_a_long_read_lists_the_first_cells_and_says_how_many_more(tmp_path, monkeypatch):
    from xlsx_mcp.ops import gridio
    monkeypatch.setattr(gridio, "SHOWS_LISTED", 3)
    p = tmp_path / "many.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 8):
        ws.cell(r, 1, r / 8).number_format = "0.0%"
    wb.save(p)
    out = cells.read_range(str(p), "A1:A7")
    assert out["shows"] == {"A1": "12.5%", "A2": "25.0%", "A3": "37.5%"}
    assert out["shows_more"] is True
    assert "first 3 cells" in out["shows_note"] and "values='displayed'" in out["shows_note"]
    # the scatter read and the grid view list the same way; find_cells marks every match
    assert cells.get_cells(str(p), [f"A{r}" for r in range(1, 8)])["shows_more"] is True
    assert len(search.find_cells(str(p), "0.", look_in="values")["matches"]) == 7
    assert all("shows" in m for m in search.find_cells(str(p), "0.")["matches"])


def test_cells_that_read_the_same_add_nothing(tmp_path):
    p = tmp_path / "plain.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", 12, 0.5, True, "=1+1"])
    ws.append(["Other", -7, 1234.25, False, None])
    ws["F1"] = dt.date(1999, 3, 4)
    ws["F1"].number_format = "yyyy-mm-dd"
    ws["G1"] = "00123"
    ws["G1"].number_format = "@"
    wb.save(p)
    p = str(p)
    for out in (cells.read_range(p, "A1:G2"), cells.read_range(p, "A1:G2", values="both"),
                cells.get_cells(p, ["A1", "B1", "C2", "D1", "F1", "G1"]), view.get_grid_view(p),
                search.find_cells(p, "1")):
        assert "shows" not in json.dumps(out), out


# ------------------------------------------------- held to Excel's own text


def _excel_number(v: float) -> str:
    """A number as Excel stores it: the shortest text that reads back as the
    same double (openpyxl writes 16 significant digits, which moves a value
    such as 0.041672453703703705 off the half second Excel measured)."""
    return str(int(v)) if float(v).is_integer() and abs(v) < 1e15 else repr(float(v))


def _oracle_book(tmp_path: Path) -> tuple[str, list]:
    import re
    import zipfile

    rows = [(g, v, f, want) for g, cs in sorted(ORACLE["groups"].items()) for v, f, want in cs]
    p = tmp_path / "oracle_16.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Shown"
    for i, (_g, v, f, _want) in enumerate(rows, start=1):
        ws.cell(i, 1, v).number_format = f
    wb.save(p)
    q = tmp_path / "oracle.xlsx"
    with zipfile.ZipFile(p) as src, zipfile.ZipFile(q, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text, n = re.subn(
                    r'(<c r="A(\d+)"[^>]*>)<v>[^<]*</v>',
                    lambda m: f"{m.group(1)}<v>{_excel_number(rows[int(m.group(2)) - 1][1])}</v>",
                    data.decode("utf-8"))
                assert n == len(rows)
                data = text.encode("utf-8")
            dst.writestr(item, data)
    return str(q), rows


def test_the_shown_text_is_excels_own(tmp_path):
    """Every fixture cell, read by default: shows holds Excel's text, or the
    cell carries no shows and its value reads as Excel's text. A time held
    to the millisecond at exactly half of its last shown unit (59.5 seconds
    under [ss], where the file holds 59.49999999999999) is settled from the
    number the file stores. values='displayed' gives the same text."""
    p, rows = _oracle_book(tmp_path)
    assert len(rows) == 2569
    wrong, by_shows, by_value, shown_wrong, step = [], 0, 0, [], 100  # one listing's worth
    for top in range(1, len(rows) + 1, step):
        bottom = min(len(rows), top + step - 1)
        out = cells.read_range(p, f"A{top}:A{bottom}", sheet="Shown")
        shows = out.get("shows", {})
        assert "shows_more" not in out
        shown = cells.read_range(p, f"A{top}:A{bottom}", sheet="Shown", values="displayed")
        for k, (value,) in enumerate(out["values"]):
            g, v, f, want = rows[top + k - 1]
            got = shows.get(f"A{top + k}")
            if got is not None:
                by_shows += 1
            else:
                got = str(value)
                by_value += 1
            if got != want:
                wrong.append((g, v, f, want, got))
            if shown["values"][k][0] != want:
                shown_wrong.append((g, v, f, want, shown["values"][k][0]))
    assert not wrong, f"{len(wrong)} of {len(rows)} differ from Excel's text: {wrong[:8]}"
    assert not shown_wrong, f"values='displayed': {len(shown_wrong)} differ: {shown_wrong[:8]}"
    assert by_shows > 2000 and by_value > 0, (by_shows, by_value)
