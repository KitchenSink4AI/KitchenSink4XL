"""A workbook whose ZIP also holds directory entries (names ending in '/') is
edited normally: those entries are not OOXML parts, so a full rewrite that does
not write them back is not a lost part."""
import zipfile

import openpyxl
import pytest

from xlsx_mcp.core import hazard
from xlsx_mcp.ops import cells, lifecycle

DIRS = ("_rels/", "docProps/", "xl/", "xl/_rels/", "xl/theme/", "xl/worksheets/")


@pytest.fixture
def book_with_directory_entries(tmp_path):
    plain = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "kept"
    wb.save(plain)
    path = tmp_path / "book-with-dirs.xlsx"
    with zipfile.ZipFile(plain) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in zin.namelist():
            zout.writestr(name, zin.read(name))
        for d in DIRS:
            zout.writestr(d, b"")
    return path


def test_the_part_inventory_leaves_directory_entries_out(book_with_directory_entries):
    rep = hazard.scan_path(str(book_with_directory_entries))
    assert not any(n.endswith("/") for n in rep.parts)
    assert not any(n.endswith("/") for n in rep.sizes)
    assert "xl/workbook.xml" in rep.sizes


def test_adding_a_worksheet_rewrites_the_workbook_and_verifies(book_with_directory_entries):
    out = lifecycle.manage_worksheet(str(book_with_directory_entries), action="add", sheet="Added")
    assert out["ok"] and out.get("verified") is True
    wb = openpyxl.load_workbook(book_with_directory_entries)
    assert wb.sheetnames == ["Sheet", "Added"]
    assert wb["Sheet"]["A1"].value == "kept"


def test_a_cell_edit_still_verifies(book_with_directory_entries):
    out = cells.set_cell(str(book_with_directory_entries), "B2", "edited")
    assert out["ok"] and out.get("verified") is True
    assert openpyxl.load_workbook(book_with_directory_entries)["Sheet"]["B2"].value == "edited"
