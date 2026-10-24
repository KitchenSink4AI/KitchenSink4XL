"""Deleting a table column turns every structured reference that named it
into #REF!, in cells and in defined names, as Excel stores it. Made-up
workbook.

Measured with a private hidden Excel (ListColumns("FY24").Delete() on the
same workbook), Excel stores:
    =SUM(Sales[FY24])            ->  =SUM(#REF!)
    =SUM(Sales[[FY23]:[FY24]])   ->  =SUM(#REF!)
    =Sales[[#Headers],[FY24]]    ->  =#REF!
    name FY24Tot SUM(Sales[FY24]) -> SUM(#REF!)
and leaves the references to the other columns as they were. Before, the
dead reference stayed as text, so a column added later under the same name
would have been picked up by formulas Excel had already broken.
"""

from __future__ import annotations

import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

import family_harness as fh


def _ok(tool, op, **args):
    payload, err, _ = fh.start().call(tool, {"op": op, "args": args})
    assert not err, (op, payload)
    return payload


def _book(path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Region", "FY23", "FY24", "Note"])
    ws.append(["North", 1, 2, "a"])
    ws.append(["South", 3, 4, "b"])
    t = Table(displayName="Sales", ref="A1:D3")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    s = wb.create_sheet("Summ")
    for i, f in enumerate(["=SUM(Sales[FY24])", "=SUM(Sales[FY23])+1",
                           "=SUM(Sales[[FY23]:[FY24]])", "=Sales[[#Headers],[FY24]]",
                           "=COUNTA(Sales[Note])", '="Sales[FY24]"&SUM(Sales[fy24])'], start=1):
        s.cell(i, 1, f)
    wb.defined_names["FY24Tot"] = DefinedName("FY24Tot", attr_text="SUM(Sales[FY24])")
    s["B1"] = "=FY24Tot"
    wb.save(path)
    return str(path)


def test_references_to_a_deleted_table_column_become_ref_errors(tmp_path):
    p = _book(tmp_path / "t.xlsx")
    _ok("xl_structure", "manage_table", path=p, name="Sales", action="delete_column",
        column="FY24")
    wb = openpyxl.load_workbook(p)
    s = wb["Summ"]
    assert [s.cell(i, 1).value for i in range(1, 7)] == [
        "=SUM(#REF!)", "=SUM(Sales[FY23])+1", "=SUM(#REF!)", "=#REF!",
        "=COUNTA(Sales[Note])", '="Sales[FY24]"&SUM(#REF!)']
    assert s["B1"].value == "=FY24Tot"
    assert wb.defined_names["FY24Tot"].attr_text == "SUM(#REF!)"
    assert [c.name for c in wb["Data"].tables["Sales"].tableColumns] == ["Region", "FY23", "Note"]
