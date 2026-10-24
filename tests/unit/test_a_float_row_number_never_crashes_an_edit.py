"""A row numbered "2.0" is read, edited and diagnosed, never a crash.

Some producers write the row number of a worksheet row as a decimal
(<row r="2.0">). openpyxl reads it as row 2. The byte-preserving route's
scanner parsed it with int() and raised a bare ValueError that no caller
expected, so a cell edit and diagnose_workbook on such a file failed with a
Python error instead of an answer. The rewrite route's carry step read the
original's row numbers the same way. The workbooks here are made up by
code: openpyxl writes them, and one byte patch turns their row numbers into
decimals.
"""

from __future__ import annotations

import re
import zipfile

import openpyxl
import pytest

from xlsx_mcp.core import surgical as _surgical
from xlsx_mcp.core.errors import WorkbookCorrupt, XlMcpError
from xlsx_mcp.ops import cells, lifecycle, structure

SHEET = "xl/worksheets/sheet1.xml"


def _decimal_rows(path: str, heights: bool) -> None:
    """Rewrite every <row r="N"> of the first sheet as <row r="N.0">."""
    with zipfile.ZipFile(path) as zin:
        parts = [(i, zin.read(i)) for i in zin.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in parts:
            if info.filename == SHEET:
                data = re.sub(rb'<row r="(\d+)"', rb'<row r="\1.0"', data)
                assert b'r="2.0"' in data
                if heights:
                    assert b'r="2.0" ht="' in data or b'ht="' in data
            zout.writestr(info, data)


def _book(tmp_path, heights: bool = False) -> str:
    p = str(tmp_path / "decimal_rows.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"], ws["B1"] = "Item", "Qty"
    ws["A2"], ws["B2"] = "bolts", 12
    ws["A3"], ws["B3"] = "nuts", 30
    ws["B4"] = "=SUM(B2:B3)"
    if heights:
        ws.row_dimensions[2].height = 21
        ws.row_dimensions[3].height = 24
    wb.save(p)
    wb.close()
    _decimal_rows(p, heights)
    return p


def test_the_fixture_really_has_decimal_row_numbers(tmp_path):
    p = _book(tmp_path)
    with zipfile.ZipFile(p) as z:
        xml = z.read(SHEET)
    assert b'<row r="1.0"' in xml and b'<row r="1"' not in xml
    assert openpyxl.load_workbook(p)["Data"]["B3"].value == 30


def test_the_splice_check_answers_instead_of_raising(tmp_path):
    p = _book(tmp_path)
    reason = _surgical.splice_check(p)
    assert reason is None or isinstance(reason, str)


def test_a_cell_edit_lands_on_a_decimal_row(tmp_path):
    # Decimal row numbers with no row dimensions: openpyxl reads them as
    # integers, and the byte-preserving route now does too.
    p = _book(tmp_path)
    r = cells.set_cell(p, {"cell": "B3", "sheet": "Data"}, 31)
    assert r["ok"] and r["verified"]
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["B3"].value == 31
    assert wb["Data"]["A2"].value == "bolts"
    wb.close()


def test_a_decimal_row_with_a_height_refuses_plainly_not_crashes(tmp_path):
    # openpyxl's own reader cannot load a <row r="1.0"> that also carries a
    # height (bind_row_dimensions does int("1.0")). The edit must refuse
    # plainly, never let the Python error escape.
    p = _book(tmp_path, heights=True)
    with pytest.raises(XlMcpError) as exc:
        cells.set_cell(p, {"cell": "B3", "sheet": "Data"}, 31)
    assert isinstance(exc.value, WorkbookCorrupt)
    # the file is untouched
    with zipfile.ZipFile(p) as z:
        assert b'<row r="2.0"' in z.read(SHEET)


def test_a_new_row_lands_after_decimal_rows(tmp_path):
    p = _book(tmp_path)
    r = cells.set_cell(p, {"cell": "A9", "sheet": "Data"}, "washers")
    assert r["ok"] and r["verified"]
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A9"].value == "washers"
    assert wb["Data"]["B3"].value == 30
    wb.close()


def test_diagnose_answers_for_a_decimal_row_workbook(tmp_path):
    p = _book(tmp_path)
    d = lifecycle.diagnose_workbook(p)
    assert d["readable"] is True
    assert d["routing"]["cell_edits"]["route"] in ("raw_ooxml", "openpyxl")


def test_a_row_insert_over_plain_decimal_rows(tmp_path):
    # A structural edit takes the rewrite route, which loads through
    # openpyxl; decimal rows with no dimensions load fine, so the insert
    # succeeds and the formula below it shifts.
    p = _book(tmp_path)
    r = structure.modify_grid_structure(p, "insert_rows", {"cell": "A3", "sheet": "Data"})
    assert r["ok"] and r["verified"] and r["route"] == "openpyxl"
    wb = openpyxl.load_workbook(p)
    ws = wb["Data"]
    assert ws["A4"].value == "nuts" and ws["B5"].value == "=SUM(B2:B4)"
    wb.close()
