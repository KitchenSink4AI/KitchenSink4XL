"""Writes honour sheet
and workbook protection, refusing with the new code PROTECTED.

Every call goes through the family surface the way a client makes it (the
in-process harness); the engine functions called in process keep their raw
behaviour, as in Word. The done-when sweep at the end runs every file-tier
writing op's own xl_help example on a fully protected copy of the op-sweep's
canary workbook.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp import envelope, registry
from xlsx_mcp import server as srv
from xlsx_mcp.surface import copy as copy_mod
from test_op_sweep import _args_for, _canary_book


def _sha(p) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def _book(tmp_path, *, options=None, unlock=None, structure=False, name="p.xlsx") -> str:
    p = tmp_path / name
    srv.create_workbook(str(p), sheets=["Data", "Other"])
    srv.write_range(str(p), "A1", [["Name", "Qty"], ["a", 1], ["b", 2], ["c", 3]],
                    sheet="Data")
    srv.set_formula(str(p), "B5", "=SUM(B2:B4)", sheet="Data")
    srv.set_cell(str(p), "A1", "=Data!B2", sheet="Other")
    if unlock:
        srv.set_protection(str(p), "unlock", sheet="Data", unlock_ranges=unlock)
    srv.set_protection(str(p), "sheet", sheet="Data", options=options or {})
    if structure:
        srv.set_protection(str(p), "workbook")
    return str(p)


def _call(tool, op, **args):
    return fh.start().call(tool, {"op": op, "args": args})


def _refused(payload, err) -> bool:
    return err and payload.get("error", {}).get("code") == "PROTECTED"


# --------------------------------------------------------------- the code


def test_protected_is_a_closed_code_with_a_hint():
    assert "PROTECTED" in envelope.CLOSED_CODES and "GUARD_TRIPPED" in envelope.CLOSED_CODES
    assert envelope.HINTS["PROTECTED"] == copy_mod.NEW_HINTS["PROTECTED"]
    assert "set_protection" in envelope.HINTS["PROTECTED"]


# ------------------------------------------------------------ locked cells


@pytest.mark.parametrize("tool,op,args", [
    ("xl_edit", "set_cell", {"location": "A2", "value": "x", "sheet": "Data"}),
    ("xl_edit", "set_cells", {"cells": [{"cell": "A2", "value": "x"}], "sheet": "Data"}),
    ("xl_edit", "write_range", {"location": "A2", "data": [["x", 9]], "sheet": "Data"}),
    ("xl_edit", "clear_range", {"location": "A2:B2", "sheet": "Data"}),
    ("xl_edit", "set_formula", {"location": "C2", "formula": "=B2*2", "sheet": "Data"}),
    ("xl_edit", "apply_edits", {"edits": [{"op": "set_value", "location": {
        "cell": "A2", "sheet": "Data"}, "value": "x"}]}),
    ("xl_edit", "replace_cells", {"find": "a", "replace": "z", "match": "exact",
                                  "sheet": "Data"}),
    ("xl_edit", "copy_range", {"source": {"range": "A2:B2", "sheet": "Data"},
                               "dest": {"cell": "D2", "sheet": "Data"}}),
    ("xl_edit", "move_range", {"source": {"range": "A2:B2", "sheet": "Data"},
                               "dest": {"cell": "D2", "sheet": "Data"}}),
    ("xl_edit", "import_data", {"source": "x,y\n1,2\n", "fmt": "csv",
                                "location": {"cell": "D1", "sheet": "Data"}}),
    ("xl_structure", "sort_range", {"location": {"range": "A1:B4", "sheet": "Data"},
                                    "keys": [{"column": "Qty", "order": "desc"}]}),
])
def test_a_locked_cell_of_a_protected_sheet_refuses_every_writer(tool, op, args, tmp_path):
    p = _book(tmp_path)
    before = _sha(p)
    payload, err, _ = _call(tool, op, path=p, **args)
    assert _refused(payload, err), payload
    assert _sha(p) == before
    assert payload["error"]["details"]["sheet"] == "Data"
    assert "Nothing was changed" in payload["error"]["message"]


def test_unlocked_cells_stay_editable_and_locked_ones_do_not(tmp_path):
    p = _book(tmp_path, unlock=["A2:A4"])
    payload, err, _ = _call("xl_edit", "set_cell", path=p, location="A3", value="ok",
                            sheet="Data")
    assert not err, payload
    assert openpyxl.load_workbook(p)["Data"]["A3"].value == "ok"
    payload, err, _ = _call("xl_edit", "set_cell", path=p, location="B3", value=5,
                            sheet="Data")
    assert _refused(payload, err)
    assert payload["error"]["details"]["cells"] == ["B3"]


def test_another_sheet_is_not_protected_by_this_one(tmp_path):
    p = _book(tmp_path)
    payload, err, _ = _call("xl_edit", "set_cell", path=p, location="C3", value=1,
                            sheet="Other")
    assert not err, payload


def test_a_reference_rewrite_on_a_protected_sheet_is_not_an_edit_of_it(tmp_path):
    """Rows inserted on Other move nothing on Data; Data's formulas that point
    at Other are rewritten, as Excel rewrites them on a protected sheet."""
    p = _book(tmp_path)
    wb = openpyxl.load_workbook(p)
    wb["Data"].protection.sheet = False
    wb["Data"]["D2"] = "=Other!A1"
    wb["Data"].protection.sheet = True
    wb.save(p)
    payload, err, _ = _call("xl_structure", "modify_grid_structure", path=p,
                            action="insert_rows", at=1, sheet="Other")
    assert not err, payload
    assert openpyxl.load_workbook(p)["Data"]["D2"].value == "=Other!A2"


def test_lifting_protection_is_always_allowed_and_then_edits_work(tmp_path):
    p = _book(tmp_path, structure=True)
    payload, err, _ = _call("xl_workbook", "set_protection", path=p, action="remove")
    assert not err, payload
    payload, err, _ = _call("xl_edit", "set_cell", path=p, location="A2", value="x",
                            sheet="Data")
    assert not err, payload


def test_in_process_engine_calls_keep_their_raw_behaviour(tmp_path):
    """As in Word: the backstop is armed only during a dispatched call."""
    p = _book(tmp_path)
    srv.set_cell(p, "A2", "engine", sheet="Data")
    assert openpyxl.load_workbook(p)["Data"]["A2"].value == "engine"


# ----------------------------------------------------- the sheet's options


@pytest.mark.parametrize("tool,op,args,option", [
    ("xl_format", "format_cells", {"location": "A2", "font": {"bold": True}}, "format_cells"),
    ("xl_format", "set_merge", {"action": "merge", "location": "D1:E1"}, "format_cells"),
    ("xl_format", "manage_conditional_format", {
        "action": "add", "location": "B2:B4", "cf_type": "cell_is",
        "params": {"operator": "greaterThan", "formula": ["1"], "fill": "FFC7CE"}},
     "format_cells"),
    ("xl_format", "set_dimensions", {"column_widths": {"A": 30}}, "format_columns"),
    ("xl_format", "set_dimensions", {"row_heights": {"2": 30}}, "format_rows"),
    ("xl_structure", "modify_grid_structure", {"action": "insert_rows", "at": 7},
     "insert_rows"),
    ("xl_structure", "modify_grid_structure", {"action": "insert_cols", "at": "F"},
     "insert_columns"),
    ("xl_structure", "set_filter", {"location": "A1:B4"}, "auto_filter"),
    ("xl_format", "manage_comment", {"action": "add", "location": "D9", "text": "n"},
     "objects"),
    ("xl_format", "manage_hyperlink", {"action": "add", "location": "D9",
                                       "target": "https://example.com"},
     "insert_hyperlinks"),
])
def test_each_option_decides_its_kind_of_change(tool, op, args, option, tmp_path):
    blocked = _book(tmp_path, unlock=["D1:E20", "A2"], name="blocked.xlsx")
    payload, err, _ = _call(tool, op, path=blocked, sheet="Data", **args)
    assert _refused(payload, err), (op, payload)
    allowed = _book(tmp_path, unlock=["D1:E20", "A2"], options={option: True},
                    name="allowed.xlsx")
    payload, err, _ = _call(tool, op, path=allowed, sheet="Data", **args)
    assert not err or payload["error"]["code"] != "PROTECTED", (op, option, payload)


@pytest.mark.parametrize("tool,op,args", [
    ("xl_structure", "manage_data_validation", {"action": "add", "location": "D2",
                                                "dv_type": "list", "values": ["a", "b"]}),
    ("xl_structure", "create_table", {"location": "D1:E3", "name": "New"}),
])
def test_validation_and_tables_never_change_on_a_protected_sheet(tool, op, args, tmp_path):
    p = _book(tmp_path, unlock=["D1:E20"],
              options={o: True for o in ("format_cells", "insert_rows", "objects")})
    payload, err, _ = _call(tool, op, path=p, sheet="Data", **args)
    assert _refused(payload, err), payload


def test_deleting_rows_needs_the_option_and_no_locked_value_in_them(tmp_path):
    p = _book(tmp_path, options={"delete_rows": True})
    payload, err, _ = _call("xl_structure", "modify_grid_structure", path=p,
                            action="delete_rows", at=2, sheet="Data")
    assert _refused(payload, err) and "A2" in payload["error"]["message"], payload
    payload, err, _ = _call("xl_structure", "modify_grid_structure", path=p,
                            action="delete_rows", at=9, sheet="Data")
    assert not err, payload


def test_sorting_needs_the_sort_option_even_for_unlocked_cells(tmp_path):
    p = _book(tmp_path, unlock=["A1:B4"])
    args = {"location": {"range": "A1:B4", "sheet": "Data"},
            "keys": [{"column": "Qty", "order": "desc"}]}
    payload, err, _ = _call("xl_structure", "sort_range", path=p, **args)
    assert _refused(payload, err), payload
    q = _book(tmp_path, unlock=["A1:B4"], options={"sort": True}, name="sortable.xlsx")
    payload, err, _ = _call("xl_structure", "sort_range", path=q, **args)
    assert not err, payload


# ------------------------------------------------------- workbook structure


@pytest.mark.parametrize("args", [
    {"action": "add", "new_name": "New"},
    {"action": "rename", "sheet": "Other", "new_name": "Renamed"},
    {"action": "delete", "sheet": "Other"},
    {"action": "copy", "sheet": "Other"},
    {"action": "reorder", "sheet": "Other", "index": 0},
    {"action": "hide", "sheet": "Other", "state": "hidden"},
])
def test_a_structure_protected_workbook_refuses_every_sheet_change(args, tmp_path):
    p = _book(tmp_path, structure=True)
    before = _sha(p)
    payload, err, _ = _call("xl_structure", "manage_worksheet", path=p, **args)
    assert _refused(payload, err), payload
    assert payload["error"]["details"]["protection"] == "structure"
    assert _sha(p) == before


def test_structure_protection_alone_leaves_cells_editable(tmp_path):
    p = _book(tmp_path, structure=True)
    payload, err, _ = _call("xl_edit", "set_cell", path=p, location="C3", value=1,
                            sheet="Other")
    assert not err, payload


# ------------------------------------------------------- the done-when sweep

#: Writing ops Excel allows on a protected workbook: they change no cell,
#: format, row, column or sheet (document properties, page setup, headers,
#: view, defined names), lift the protection, restore a backup, or write a
#: new file without changing the source.
ALLOWED_ON_PROTECTED = {"set_workbook_properties", "set_page_layout", "set_header_footer",
                        "set_view", "manage_name", "set_protection", "manage_backups",
                        "create_workbook", "copy_workbook", "export_range", "export_file"}


def _protected_canary(path: Path) -> str:
    p = _canary_book(path)
    wb = openpyxl.load_workbook(p)
    names = wb.sheetnames
    for name in names:
        srv.set_protection(p, "sheet", sheet=name)
    srv.set_protection(p, "workbook")
    return p


def _cells(p) -> dict:
    wb = openpyxl.load_workbook(p)
    return {(ws.title, c.coordinate): (c.value, c.style_id)
            for ws in wb.worksheets for row in ws.iter_rows() for c in row}


def _fingerprint(p) -> tuple:
    """What protection guards, read independently of the product: sheets and
    their order and state, every cell's value, style, link and comment,
    merges, row and column dimensions, conditional formats, validations,
    tables, objects and the autofilter."""
    wb = openpyxl.load_workbook(p)
    out = [tuple((ws.title, ws.sheet_state) for ws in wb.worksheets)]
    for ws in wb.worksheets:
        out.append(tuple(sorted(
            (c.coordinate, repr(c.value), c.style_id,
             c.hyperlink.target if c.hyperlink else None,
             c.comment.text if c.comment else None)
            for row in ws.iter_rows() for c in row
            if c.value is not None or c.has_style or c.hyperlink or c.comment)))
        out.append(tuple(sorted(str(r) for r in ws.merged_cells.ranges)))
        out.append(tuple(sorted((k, d.ht, d.hidden) for k, d in ws.row_dimensions.items())))
        out.append(tuple(sorted((k, d.width, d.hidden) for k, d in ws.column_dimensions.items())))
        out.append(tuple(sorted(str(cf.sqref) for cf in ws.conditional_formatting)))
        out.append(tuple(sorted(str(dv.sqref) for dv in ws.data_validations.dataValidation)))
        out.append(tuple(sorted(ws.tables.items())))
        out.append((len(ws._images), len(ws._charts), ws.auto_filter.ref))
    return tuple(out)


def test_b1_no_write_op_changes_a_protected_workbook_silently(tmp_path):
    source = tmp_path / "canary-protected.xlsx"
    _protected_canary(source)
    names = sorted(n for n, s in registry.public_ops().items()
                   if s.writes and not s.runs_live)
    h = fh.start()
    silent = []
    for name in names:
        work = tmp_path / name
        work.mkdir()
        args = _args_for(name, work, source)
        book = work / f"{name}.xlsx"
        before = _fingerprint(book) if book.exists() else None
        payload, err, _ = h.call(registry.REGISTRY[name].group, {"op": name, "args": args})
        if _refused(payload, err):
            continue
        if before is None or not book.exists() or _fingerprint(book) == before:
            continue  # nothing protection guards changed (a re-save is not a change)
        if name in ALLOWED_ON_PROTECTED:
            assert _fingerprint(book)[1:] == before[1:] or name == "manage_name",                 f"{name} changed what protection guards"
            continue
        silent.append(name)
    assert silent == [], f"these ops changed a protected workbook: {silent}"


# ------------------------------------------- a protected model, replayed from a real job


def test_x44_the_timesheet_job_replayed(tmp_path):
    """The hunt's job J12 (calls.jsonl seq 347-357), replayed: a timesheet
    whose inputs C2:G3 stay editable under a password-protected sheet and a
    password-locked structure. Before the fix every step answered ok: the locked
    rate B2 took 99, rows were inserted and a sheet was added."""
    p = str(tmp_path / "timesheet.xlsx")
    h = fh.start()
    steps = [
        ("xl_workbook", "create_workbook", {"sheets": ["Timesheet"]}),
        ("xl_edit", "write_range", {"location": {"cell": "A1", "sheet": "Timesheet"}, "data": [
            ["Employee", "Rate", "Mon", "Tue", "Wed", "Thu", "Fri", "Hours", "Pay"],
            ["Kim", 42.5, 8, 8, 7.5, 8, 6, "=SUM(C2:G2)", "=H2*B2"],
            ["Lee", 38, 8, 8, 8, 8, 8, "=SUM(C3:G3)", "=H3*B3"],
            ["Total", None, None, None, None, None, None, "=SUM(H2:H3)", "=SUM(I2:I3)"]]}),
        ("xl_workbook", "set_protection", {"action": "sheet", "sheet": "Timesheet",
                                           "password": "s3cret", "unlock_ranges": ["C2:G3"]}),
        ("xl_workbook", "set_protection", {"action": "workbook", "password": "s3cret"}),
    ]
    for tool, op, args in steps:
        payload, err, _ = h.call(tool, {"op": op, "args": {"path": p, **args}})
        assert not err, (op, payload)
    status, err, _ = h.call("xl_workbook", {"op": "set_protection", "args": {
        "path": p, "action": "status"}})
    assert not err and status["sheets"][0]["unlocked_ranges"] == ["C2:G3"], status
    ok, err, _ = h.call("xl_edit", {"op": "set_cell", "args": {
        "path": p, "location": {"cell": "C2", "sheet": "Timesheet"}, "value": 9}})
    assert not err, ok
    for tool, op, args in [
        ("xl_edit", "set_cell", {"location": {"cell": "B2", "sheet": "Timesheet"},
                                 "value": 99}),
        ("xl_structure", "modify_grid_structure", {"sheet": "Timesheet",
                                                   "action": "insert_rows", "at": 3}),
        ("xl_structure", "manage_worksheet", {"action": "add", "new_name": "Extra"}),
    ]:
        payload, err, _ = h.call(tool, {"op": op, "args": {"path": p, **args}})
        assert _refused(payload, err), (op, payload)
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Timesheet"] and wb["Timesheet"]["B2"].value == 42.5
    assert wb["Timesheet"]["C2"].value == 9


def test_x43_an_argument_the_action_ignores_is_refused(tmp_path):
    p = _book(tmp_path)
    payload, err, _ = _call("xl_workbook", "set_protection", path=p, action="workbook",
                            unlock_ranges=["A1"])
    assert err and payload["error"]["code"] == "BAD_PARAMS", payload
    assert "does not use unlock_ranges" in payload["error"]["message"]
