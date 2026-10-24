"""A first note or link on a sheet whose root already uses the prefix r.

The relationship id of a <legacyDrawing> or <hyperlink> lives in the
officeDocument relationships namespace. When the worksheet root binds no
prefix to that namespace, the byte-preserving route declares one. A producer
may already bind r to a namespace of its own (marked ignorable), or bind the
relationships namespace with single quotes; declaring xmlns:r a second time
makes the sheet malformed XML ("Attribute xmlns:r redefined") and Excel cannot
open it. The declaration must reuse a correct binding or take a free prefix.

Every workbook here is made up and built by code.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook

from xlsx_mcp.core import surgical as _surgical
from xlsx_mcp.ops import annotations

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
FOREIGN = "urn:ks4corpus:foreign"
SHEET = "xl/worksheets/sheet1.xml"

ROOTS = {
    # r is taken by a foreign, ignorable namespace that the root itself uses
    "foreign_r": (f'<worksheet xmlns="{MAIN}" xmlns:mc="{MC}" xmlns:r="{FOREIGN}" '
                  f'mc:Ignorable="r" r:batch="made-up-7">'),
    # the relationships namespace is already bound, with single quotes
    "single_quoted_rel": f"<worksheet xmlns='{MAIN}' xmlns:r='{REL}'>",
    # r foreign, and the next candidate r0 is declared deeper in the sheet
    "foreign_r_and_r0_inside": (f'<worksheet xmlns="{MAIN}" xmlns:mc="{MC}" xmlns:r="{FOREIGN}" '
                                f'mc:Ignorable="r" r:batch="made-up-8">'),
}


def _book(tmp_path: Path, kind: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger"
    ws["A1"] = "Item"
    ws["B1"] = "Qty"
    ws["A2"] = "Widget"
    ws["B2"] = 7
    src = tmp_path / "src.xlsx"
    wb.save(src)
    with zipfile.ZipFile(src) as z:
        parts = {i.filename: z.read(i.filename) for i in z.infolist()}
    sheet = parts[SHEET].decode("utf-8")
    root = re.search(r"<worksheet\b[^>]*>", sheet)
    assert root is not None
    sheet = sheet[:root.start()] + ROOTS[kind] + sheet[root.end():]
    if kind == "foreign_r_and_r0_inside":
        sheet = sheet.replace("<sheetFormatPr ",
                              f'<sheetFormatPr xmlns:r0="{FOREIGN}" ', 1)
        assert 'xmlns:r0="' in sheet
    parts[SHEET] = sheet.encode("utf-8")
    dst = tmp_path / f"{kind}.xlsx"
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return str(dst)


def _sheet_xml(path: str) -> bytes:
    with zipfile.ZipFile(path) as z:
        return z.read(SHEET)


def _rel_ids(path: str, local: str) -> list[str]:
    root = ET.fromstring(_sheet_xml(path))
    return [el.get(f"{{{REL}}}id") for el in root.iter(f"{{{MAIN}}}{local}")]


@pytest.mark.parametrize("kind", sorted(ROOTS))
def test_first_note_over_a_taken_r_prefix_stays_well_formed(tmp_path, kind):
    p = _book(tmp_path, kind)
    r = annotations.manage_comment(p, "add", {"cell": "C2", "sheet": "Ledger"},
                                   text="check the count")
    assert r["ok"] and r["verified"]
    assert r["route"] == _surgical.ROUTE_RAW, r.get("route_note")
    data = _sheet_xml(p)
    ET.fromstring(data)  # well-formed, no attribute declared twice
    rids = _rel_ids(p, "legacyDrawing")
    assert len(rids) == 1 and rids[0], "the drawing's id is in the relationships namespace"
    listed = annotations.manage_comment(p, "list")
    assert [(c["cell"], c["text"]) for c in listed["comments"]] == [("C2", "check the count")]
    if kind.startswith("foreign_r"):
        assert f'xmlns:r="{FOREIGN}"'.encode() in data, "the foreign binding is kept"
        assert b'r:batch="made-up-' in data


@pytest.mark.parametrize("kind", sorted(ROOTS))
def test_first_link_over_a_taken_r_prefix_stays_well_formed(tmp_path, kind):
    p = _book(tmp_path, kind)
    r = annotations.manage_hyperlink(p, "add", {"cell": "A2", "sheet": "Ledger"},
                                     target="https://example.com/widgets?id=7")
    assert r["ok"] and r["verified"]
    assert r["route"] == _surgical.ROUTE_RAW, r.get("route_note")
    data = _sheet_xml(p)
    ET.fromstring(data)
    rids = _rel_ids(p, "hyperlink")
    assert len(rids) == 1 and rids[0]
    links = annotations.manage_hyperlink(p, "list")["hyperlinks"]
    assert [(h["cell"], h["target"]) for h in links] == [("A2", "https://example.com/widgets?id=7")]
    # and the link comes out again, taking its relationship with it
    r = annotations.manage_hyperlink(p, "remove", {"cell": "A2", "sheet": "Ledger"})
    assert r["ok"]
    assert annotations.manage_hyperlink(p, "list")["hyperlinks"] == []
    with zipfile.ZipFile(p) as z:
        rels = z.read("xl/worksheets/_rels/sheet1.xml.rels") \
            if "xl/worksheets/_rels/sheet1.xml.rels" in z.namelist() else b""
    assert b"/hyperlink" not in rels


def test_note_then_link_share_one_declared_prefix(tmp_path):
    p = _book(tmp_path, "foreign_r")
    annotations.manage_comment(p, "add", {"cell": "C2", "sheet": "Ledger"}, text="n")
    annotations.manage_hyperlink(p, "add", {"cell": "A2", "sheet": "Ledger"},
                                 target="https://example.com/w")
    data = _sheet_xml(p)
    ET.fromstring(data)
    root_tag = re.search(rb"<worksheet\b[^>]*>", data).group(0)
    assert root_tag.count(REL.encode()) == 1, root_tag
    assert all(_rel_ids(p, "legacyDrawing")) and all(_rel_ids(p, "hyperlink"))


def test_rel_prefix_helper_reads_only_the_root_and_both_quote_styles():
    root = f"<worksheet xmlns='{MAIN}' xmlns:rel='{REL}'><sheetData/></worksheet>".encode()
    data, pfx, shift = _surgical.worksheet_rel_prefix(root)
    assert (data, pfx, shift) == (root, b"rel", 0)
    # a binding on a child element is not in scope for a new top-level child
    nested = (f'<worksheet xmlns="{MAIN}"><sheetPr xmlns:q="{REL}"/><sheetData/>'
              f"</worksheet>").encode()
    data, pfx, shift = _surgical.worksheet_rel_prefix(nested)
    assert pfx != b"q" and shift > 0
    ET.fromstring(data)
    assert f'xmlns:{pfx.decode()}="{REL}"'.encode() in data[:data.index(b">") + 1]
    # r taken by a foreign namespace, r0 taken deeper in the document: r1
    taken = (f'<worksheet xmlns="{MAIN}" xmlns:r="{FOREIGN}"><x xmlns:r0="{FOREIGN}"/>'
             f"<sheetData/></worksheet>").encode()
    data, pfx, shift = _surgical.worksheet_rel_prefix(taken)
    assert pfx == b"r1"
    ET.fromstring(data)
