"""validate's structure check names what makes Excel refuse a file.

The structure check answered passed=true for workbooks Excel will not open:
a font name longer than Excel's 31 characters, a number stored as "+12", a
formula past 8,192 stored characters, a macro part that is not a macro
project. It now reads the stored values against their cell types (the same
check as the gate after every save), the font names, the formula lengths and
the macro part, and names each problem. The gate after every save refuses a
produced file whose font name Excel refuses, as it refuses the other
problems, and format_cells refuses such a font name before writing.
All content is made up.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import openpyxl
import pytest

from xlsx_mcp.core import verify
from xlsx_mcp.core.errors import XlMcpError
from xlsx_mcp.ops import format as fmt
from xlsx_mcp.ops import validation

LONG_FONT = "Made Up Display Font For A Test Case"  # 36 characters
FONT_31 = "Made Up Display Font For A Test"         # 31 characters


def _book(tmp_path: Path, name: str = "book.xlsx") -> Path:
    p = tmp_path / name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Region"
    ws["A2"] = 12
    ws["B2"] = "=A2*2"
    wb.save(p)
    return p


def _patch(p: Path, part: str, fn) -> None:
    with zipfile.ZipFile(p) as zin:
        items = [(i, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            if info.filename == part:
                data = fn(data.decode("utf-8")).encode("utf-8")
            zout.writestr(info, data)


def _font_named(p: Path, name: str) -> None:
    _patch(p, "xl/styles.xml",
           lambda x: re.sub(r'<name val="[^"]*"\s*/>', f'<name val="{name}"/>',
                            x, count=1))


def _structure(p: Path) -> dict:
    return validation.validate(str(p), checks=["structure"])["results"][
        "structure"]


def test_a_clean_book_passes(tmp_path):
    res = _structure(_book(tmp_path))
    assert res["passed"] is True, res["findings"]["issues"]


def test_a_font_name_over_31_characters_is_named(tmp_path):
    p = _book(tmp_path)
    _font_named(p, LONG_FONT)
    res = _structure(p)
    assert res["passed"] is False
    text = " ".join(res["findings"]["issues"])
    assert LONG_FONT in text and "31" in text


def test_a_font_name_of_31_characters_passes(tmp_path):
    p = _book(tmp_path)
    _font_named(p, FONT_31)
    res = _structure(p)
    assert res["passed"] is True, res["findings"]["issues"]


def test_a_number_excel_cannot_read_is_named(tmp_path):
    p = _book(tmp_path)
    _patch(p, "xl/worksheets/sheet1.xml",
           lambda x: x.replace("<v>12</v>", "<v>+12</v>", 1))
    res = _structure(p)
    assert res["passed"] is False
    assert "Data!A2" in " ".join(res["findings"]["issues"])


def test_a_formula_past_8192_characters_is_named(tmp_path):
    p = _book(tmp_path)
    long_body = "1234567890+" * 745 + "1234567890"   # 8,205 characters
    _patch(p, "xl/worksheets/sheet1.xml",
           lambda x: x.replace("<f>A2*2</f>", f"<f>{long_body}</f>", 1))
    res = _structure(p)
    assert res["passed"] is False
    text = " ".join(res["findings"]["issues"])
    assert "Data!B2" in text and "8,192" in text


def test_a_macro_part_that_is_not_a_macro_project_is_named(tmp_path):
    p = tmp_path / "macros.xlsm"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.save(p)
    with zipfile.ZipFile(p) as zin:
        items = [(i, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            zout.writestr(info, data)
        zout.writestr("xl/vbaProject.bin", b"not a macro project at all")
    res = _structure(p)
    assert res["passed"] is False
    assert "vbaProject.bin" in " ".join(res["findings"]["issues"])


def test_the_gate_after_a_save_refuses_a_font_name_excel_refuses(tmp_path):
    p = _book(tmp_path)
    _font_named(p, LONG_FONT)
    res = verify.verify_after_write(str(p), pre_parts=[])
    assert not res.ok
    text = " ".join(res.reasons)
    assert LONG_FONT in text and "31" in text


def test_the_gate_after_a_save_refuses_a_formula_past_the_limit(tmp_path):
    p = _book(tmp_path)
    long_body = "1234567890+" * 745 + "1234567890"
    _patch(p, "xl/worksheets/sheet1.xml",
           lambda x: x.replace("<f>A2*2</f>", f"<f>{long_body}</f>", 1))
    res = verify.verify_after_write(str(p), pre_parts=[])
    assert not res.ok
    assert "Data!B2" in " ".join(res.reasons)


@pytest.mark.parametrize("where", ["format_cells", "define_style"])
def test_format_refuses_a_font_name_excel_refuses_before_writing(tmp_path,
                                                                  where):
    p = _book(tmp_path)
    before = p.read_bytes()
    with pytest.raises(XlMcpError) as ei:
        if where == "format_cells":
            fmt.format_cells(str(p), "A1", font={"name": LONG_FONT})
        else:
            fmt.apply_style(str(p), "Banner",
                            define={"font": {"name": LONG_FONT}})
    assert "31" in str(ei.value)
    assert p.read_bytes() == before


def test_format_takes_a_font_name_of_31_characters(tmp_path):
    p = _book(tmp_path)
    fmt.format_cells(str(p), "A1", font={"name": FONT_31})
    wb = openpyxl.load_workbook(p)
    assert wb["Data"]["A1"].font.name == FONT_31
