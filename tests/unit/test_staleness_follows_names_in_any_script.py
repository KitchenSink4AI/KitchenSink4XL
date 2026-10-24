"""The staleness report after a cell edit names every formula whose cached
result the edit may have changed, including the ones that reach the edited
cell through a defined name written in another script (Tasa_ñ, 売上Rate), a
name holding ? or a backslash (Excel accepts a?b and a\\b), a sheet's own name
used with its sheet (=Loc!LocalRate*5), and a table named in another script
(Ventas_ñ[Amt]). Made-up workbook.

Before, the name scan read ASCII names only and skipped every name after a
'!', so those dependents were never listed and, when nothing else forced a
recalculation, the workbook kept their old results while the report said it
was complete. Measured with a private hidden Excel on the same workbook:
after Excel's own edit of each input, exactly the formula listed here
changed value, and after the tools' edit Excel opened the old result for it
(no recalculation was asked for).
"""

from __future__ import annotations

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

from xlsx_mcp import server as srv
from xlsx_mcp.core import surgical


def _name(wb, name, text):
    wb.defined_names[name] = DefinedName(name, attr_text=text)


def _book(path, with_table=False) -> str:
    wb = openpyxl.Workbook()
    inp = wb.active
    inp.title = "Inputs"
    for r in range(2, 10):
        inp.cell(r, 2, r * 10)
    loc = wb.create_sheet("Loc")
    loc["A1"] = 1
    summ = wb.create_sheet("Summ")
    _name(wb, "Tasa_ñ", "Inputs!$B$2")
    _name(wb, "売上Rate", "Inputs!$B$3")
    _name(wb, "a?b", "Inputs!$B$4")
    _name(wb, "a\\b", "Inputs!$B$5")
    loc.defined_names["LocalRate"] = DefinedName("LocalRate", attr_text="Inputs!$B$6")
    forms = ["=Tasa_ñ*2", "=売上Rate+1", "=a?b*3", "=a\\b+4", "=Loc!LocalRate*5",
             "=Inputs!B8*6", "=Inputs!B9+0"]
    if with_table:
        inp["D1"], inp["E1"] = "Amt", "Qty"
        inp["D2"], inp["E2"] = "=$B$7", 1
        inp["D3"], inp["E3"] = 5, 2
        t = Table(displayName="Ventas_ñ", ref="D1:E3")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        inp.add_table(t)
        forms.append("=SUM(Ventas_ñ[Amt])")
    for i, f in enumerate(forms, start=1):
        summ.cell(i, 1, f)
    wb.save(path)
    return str(path)


DEPENDS = [  # (edited input, the one formula that reads it)
    ("B2", "A1"),   # Tasa_ñ
    ("B3", "A2"),   # 売上Rate
    ("B4", "A3"),   # a?b
    ("B5", "A4"),   # a\b
    ("B6", "A5"),   # Loc!LocalRate, a sheet's own name with its sheet
    ("B8", "A6"),   # the ASCII control
]


@pytest.mark.parametrize("cell,dependent", DEPENDS)
def test_the_walk_finds_the_formula_through_each_name(tmp_path, cell, dependent):
    wb = openpyxl.load_workbook(_book(tmp_path / "s.xlsx"))
    r, c = openpyxl.utils.cell.coordinate_to_tuple(cell)
    stale = surgical.stale_formulas(wb, {"Inputs": {(r, c)}})
    assert stale.cells == [("Summ", dependent)]
    assert stale.complete


def test_an_input_nothing_reads_leaves_nothing_stale(tmp_path):
    wb = openpyxl.load_workbook(_book(tmp_path / "s.xlsx"))
    assert surgical.stale_formulas(wb, {"Inputs": {(7, 2)}}).cells == []


def test_a_table_named_in_another_script_is_bounded_like_any_table(tmp_path):
    wb = openpyxl.load_workbook(_book(tmp_path / "t.xlsx", with_table=True))
    # B7 feeds the table (D2 = $B$7), and the table feeds Summ!A8 only.
    stale = surgical.stale_formulas(wb, {"Inputs": {(7, 2)}})
    assert stale.cells == [("Inputs", "D2"), ("Summ", "A8")]
    # An edit the table does not read leaves its sum alone (it was listed
    # for every edit, since the table name was not read at all).
    assert surgical.stale_formulas(wb, {"Inputs": {(2, 2)}}).cells == [("Summ", "A1")]


@pytest.mark.parametrize("cell,dependent", DEPENDS[:5])
def test_the_edit_reports_the_dependent_and_asks_excel_to_recalculate(tmp_path, cell, dependent):
    p = _book(tmp_path / "s.xlsx")
    r = srv.set_cell(p, {"cell": cell, "sheet": "Inputs"}, 99)
    cached = r["cached_values"]
    assert cached.get("possibly_stale_cells") == [f"Summ!{dependent}"], cached
    if r.get("route") == surgical.ROUTE_RAW:
        assert cached.get("recalculates_on_open") is True


def test_sheet_qualified_names_and_strings_are_read_apart():
    names = {(None, "TASA_Ñ"): (None, "=Inputs!$B$2"),
             ("loc", "LOCALRATE"): ("Loc", "=Inputs!$B$6"),
             (None, "LOCALRATE"): (None, "=Inputs!$B$9")}
    regs = surgical._formula_regions("=Loc!LocalRate+'Loc'!LocalRate", "Summ", names, {})
    assert regs == [("inputs", 6, 2, 6, 2), ("inputs", 6, 2, 6, 2)]
    # inside a string literal a name is text
    assert surgical._formula_regions('="Tasa_ñ"&"x"', "Summ", names, {}) == []
    assert surgical._formula_regions("=Tasa_ñ", "Summ", names, {}) == [("inputs", 2, 2, 2, 2)]
