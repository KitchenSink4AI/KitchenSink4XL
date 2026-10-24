"""A table damaged by an earlier set_filter.

Until 2.0.0rc2, set_filter on a table wrote a SHEET filter over the table's
range (plus an _xlnm._FilterDatabase name), and Excel will not open such a
workbook; 1.2.6, the released version, writes the same shape. This version stopped
writing it and clear_filter repairs it, but set_filter on the same table and
any other edit (set_cell) answered ok and verified and left the workbook as
Excel still refuses it: the gate failed only on damage a save ADDED.

Now every file-tier save repairs the damage the workbook already carried
(core.tablefilter: the sheet filter's criteria move into the table's own
filter when they cover exactly its range, or the sheet filter is dropped;
the name written beside it goes), the gate fails closed on a sheet filter
over a table whatever its age, and the COM routes refuse such a workbook
before Excel starts, naming clear_filter. Damage a save would ADD is never
repaired (test_table_filter_corruption_guard's fail-closed test still holds), and an
_FilterDatabase name on its own, which Excel opens and keeps itself, is left
alone.
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.filters import AutoFilter, FilterColumn, Filters
from openpyxl.worksheet.table import Table

from xlsx_mcp.core import tablefilter
from xlsx_mcp.core import verify as _verify
from xlsx_mcp.core.errors import UnsupportedStructure, ValidationFailed
from xlsx_mcp.ops import cells, sortfilter


def _damaged(tmp_path, name="old.xlsx", sheet_ref="A1:C4", col=1, table_filter=True):
    """The shape set_filter wrote in older versions: a sheet filter over table Fruit
    (criteria on Amount = 7, the rows that fail it hidden) and the
    _FilterDatabase name openpyxl writes beside it."""
    p = str(tmp_path / name)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in [["Item", "Amount", "Date"], ["Apple", 3, 1], ["Pear", 7, 2], ["Plum", 9, 3]]:
        ws.append(r)
    t = Table(displayName="Fruit", ref="A1:C4")
    if table_filter:
        t.autoFilter = AutoFilter(ref="A1:C4")
    ws.add_table(t)
    ws.auto_filter.ref = sheet_ref
    ws.auto_filter.filterColumn = [FilterColumn(colId=col, filters=Filters(filter=["7"]))]
    ws.row_dimensions[2].hidden = True
    ws.row_dimensions[4].hidden = True
    wb.save(p)
    assert tablefilter.scan(p).filters, "fixture should carry the old damage"
    return p


def _xml(p):
    with zipfile.ZipFile(p) as z:
        names = z.namelist()
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
        table = next(z.read(n).decode() for n in names if n.startswith("xl/tables/"))
        wbx = z.read("xl/workbook.xml").decode()
    return sheet, table, wbx


def _sheet_filter(sheet_xml):
    tail = sheet_xml[sheet_xml.find("</sheetData>"):]
    return "<autoFilter" in tail


def _hidden_rows(sheet_xml):
    return sorted(int(m) for m in re.findall(r'<row r="(\d+)"[^>]*hidden="1"', sheet_xml))


def test_any_edit_repairs_the_damage_and_says_so(tmp_path):
    p = _damaged(tmp_path)
    r = cells.set_cell(p, {"cell": "E1", "sheet": "Data"}, "note")
    assert r["ok"] and r["verified"]
    sheet, table, wbx = _xml(p)
    assert not _sheet_filter(sheet) and "_FilterDatabase" not in wbx
    assert _verify.package_problems(p) == set()
    # the filter the user asked for is now the table's own, over the same hidden rows
    assert "<filterColumn" in table and 'colId="1"' in table and 'val="7"' in table
    assert _hidden_rows(sheet) == [2, 4]
    text = " ".join(r["warnings"])
    assert "Fruit" in text and "filter" in text
    assert openpyxl.load_workbook(p)["Data"]["E1"].value == "note"


def test_a_rewrite_route_edit_repairs_it_too(tmp_path):
    from xlsx_mcp.ops import format as fmt
    p = _damaged(tmp_path)
    r = fmt.format_cells(p, {"range": "A1:C1", "sheet": "Data"}, font={"bold": True})
    assert r["ok"]
    assert _verify.package_problems(p) == set()


def test_set_filter_on_the_damaged_table_repairs_it(tmp_path):
    p = _damaged(tmp_path)
    r = sortfilter.set_filter(p, {"table": "Fruit"},
                              criteria=[{"column": "Amount", "op": "gt", "value": 5}])
    assert r["ok"]
    assert r["changed"]["filter"]["removed_sheet_filter_over_table"] == "A1:C4"
    sheet, table, wbx = _xml(p)
    assert not _sheet_filter(sheet) and "_FilterDatabase" not in wbx
    assert _verify.package_problems(p) == set()


def test_a_sheet_filter_that_is_not_the_tables_range_is_dropped(tmp_path):
    """Criteria relative to another range cannot move: the sheet filter goes,
    the rows stay as they were, and the answer says how to show them."""
    p = _damaged(tmp_path, sheet_ref="B2:B3", col=0)
    r = cells.set_cell(p, {"cell": "E1", "sheet": "Data"}, 1)
    sheet, table, wbx = _xml(p)
    assert not _sheet_filter(sheet) and "<filterColumn" not in table
    assert _hidden_rows(sheet) == [2, 4]
    assert any("clear_filter" in w for w in r["warnings"])
    assert _verify.package_problems(p) == set()


def test_a_table_without_its_own_filter_gets_the_sheet_filter_dropped(tmp_path):
    p = _damaged(tmp_path, table_filter=False)
    r = cells.set_cell(p, {"cell": "E1", "sheet": "Data"}, 1)
    assert r["ok"]
    assert _verify.package_problems(p) == set()


def test_the_gate_fails_closed_on_old_damage_left_in_place(tmp_path, monkeypatch):
    """If the repair did not happen, the save is refused and the workbook is
    left exactly as it was: never ok and verified for a file Excel refuses."""
    p = _damaged(tmp_path)
    before = Path(p).read_bytes()
    monkeypatch.setattr(tablefilter, "repair_package", lambda *a, **k: [])
    with pytest.raises(ValidationFailed) as ei:
        cells.set_cell(p, {"cell": "E1", "sheet": "Data"}, "note")
    assert "Fruit" in str(ei.value) and "clear_filter" in str(ei.value)
    assert Path(p).read_bytes() == before


def test_the_com_routes_refuse_it_before_excel_starts(tmp_path, monkeypatch):
    from xlsx_mcp.com import session as S
    from xlsx_mcp.ops import comtier
    monkeypatch.setattr(S, "run_com", lambda *a, **k: (_ for _ in ()).throw(
        AssertionError("Excel was started on a workbook it will not open")))
    p = _damaged(tmp_path)
    before = Path(p).read_bytes()
    calls = [lambda: comtier.com_autofit(p),
             lambda: comtier.com_save_with_password(p, "secret")]
    if S.com_available()[0]:
        calls.append(lambda: comtier.recalculate(p, engine="com"))
    for call in calls:
        with pytest.raises(UnsupportedStructure) as ei:
            call()
        assert "clear_filter" in str(ei.value) and "Fruit" in str(ei.value)
    assert Path(p).read_bytes() == before


def test_a_filter_database_name_on_its_own_is_left_alone(tmp_path):
    """Excel keeps an _FilterDatabase name inside a table when a filtered
    range becomes a table, and opens it; an edit must not refuse it or
    change it."""
    p = tmp_path / "excel_made.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in [["Item", "Amount"], ["Apple", 3], ["Pear", 7]]:
        ws.append(r)
    ws.add_table(Table(displayName="Fruit", ref="A1:B3"))
    wb.save(p)
    with zipfile.ZipFile(p) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    wbx = parts["xl/workbook.xml"].decode()
    wbx = re.sub(r"(</sheets>)", r'\1<definedNames><definedName name="_xlnm._FilterDatabase" '
                 r'localSheetId="0" hidden="1">Data!$A$1:$B$3</definedName></definedNames>',
                 wbx, count=1)
    parts["xl/workbook.xml"] = wbx.encode()
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(n, data)
    assert tablefilter.scan(str(p)).names and not tablefilter.scan(str(p)).filters
    r = cells.set_cell(str(p), {"cell": "D1", "sheet": "Data"}, 1)
    assert r["ok"] and not any("_FilterDatabase" in w for w in r.get("warnings", []))
    assert "_FilterDatabase" in _xml(str(p))[2]


def test_damage_an_edit_would_add_is_not_repaired(tmp_path, monkeypatch):
    """The repair covers only what the workbook already carried: a route that
    wrote the shape again still fails closed (the fail-closed guarantee)."""
    p = str(tmp_path / "clean.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in [["Item", "Amount"], ["Apple", 3], ["Pear", 7]]:
        ws.append(r)
    t = Table(displayName="Fruit", ref="A1:B3")
    t.autoFilter = AutoFilter(ref="A1:B3")
    ws.add_table(t)
    wb.save(p)
    before = Path(p).read_bytes()
    monkeypatch.setattr(sortfilter, "_table_for_filter", lambda ws, grid, what: None)
    with pytest.raises(ValidationFailed):
        sortfilter.set_filter(p, {"table": "Fruit"},
                              criteria=[{"column": "Amount", "op": "gt", "value": 5}])
    assert Path(p).read_bytes() == before
