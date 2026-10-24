"""Surface behaviours carried over from the Word editor: lenient on routing
and shape, strict on meaning, the wrong-tool reroute, and the help, hint,
op-line, view and token-bill items.

Every call goes through the family surface as a client makes it.
"""

from __future__ import annotations

import json

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp import registry
from xlsx_mcp import server as srv
from xlsx_mcp.surface import normalize


@pytest.fixture
def book(tmp_path):
    p = str(tmp_path / "s.xlsx")
    srv.create_workbook(p)
    srv.write_range(p, "A1", [["Name", "Qty"], ["a", 3], ["b", 1], ["c", 2]])
    return p


def _call(tool, op, **args):
    return fh.start().call(tool, {"op": op, "args": args})




def test_c1_no_two_ops_fold_to_the_same_name():
    normalize.op_keys(registry.public_ops())  # raises on a clash


@pytest.mark.parametrize("spelling", ["SetCell", "setCell", "set-cell", "SET_CELL"])
def test_c1_an_op_routes_by_its_folded_spelling_and_says_so(spelling, book):
    payload, err, _ = fh.start().call("xl_edit", {"op": spelling, "args": {
        "path": book, "location": "C1", "value": 1}})
    assert not err, payload
    assert payload["normalized_from"]["op"] == spelling


def test_c1_anything_fuzzier_never_routes(book):
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_cel", "args": {
        "path": book, "location": "C1", "value": 1}})
    assert err and payload["error"]["code"] == "NOT_FOUND"
    assert "nothing was run" in payload["error"]["message"]
    assert openpyxl.load_workbook(book).active["C1"].value is None




def test_c3_an_unknown_op_suggests_at_most_three_with_tool_and_purpose(book):
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_cel", "args": {}})
    near = payload["error"]["details"]["did_you_mean"]
    assert 1 <= len(near) <= 3 and near[0]["op"] == "set_cell"
    assert near[0]["tool"] == "xl_edit" and near[0]["purpose"]


def test_c3_a_destructive_op_is_never_suggested_for_a_name_without_its_verb():
    ops = list(registry.public_ops())
    for probe in ("range", "cells", "backup", "filter", "sheet"):
        for name in normalize.suggestions(probe, ops):
            assert not any(v in name for v in normalize.DESTRUCTIVE_VERBS), (probe, name)


def test_c3_help_gives_the_same_suggestions():
    payload, err, _ = fh.start().call("xl_help", {"op": "read_rnge"})
    assert err and payload["error"]["details"]["did_you_mean"][0]["op"] == "read_range"


def test_c1_help_routes_the_folded_spelling():
    payload, err, _ = fh.start().call("xl_help", {"op": "ReadRange"})
    assert not err and payload["op"] == "read_range"




def test_c2_flat_arguments_move_into_args(book):
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_cell", "path": book,
                                                  "location": "C2", "value": 5})
    assert not err, payload
    assert set(payload["normalized_from"]["flat_args"]) == {"path", "location", "value"}
    assert openpyxl.load_workbook(book).active["C2"].value == 5


@pytest.mark.parametrize("key", ["file_path", "file"])
def test_c2_word_and_powerpoint_file_spellings_name_the_path(key, book):
    payload, err, _ = _call("xl_read", "read_range", **{key: book}, location="A1")
    assert not err, payload
    assert payload["normalized_from"]["args"][key] == book


def test_c2_unknown_keys_are_still_refused(book):
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_cell", "colour": "red",
                                                  "args": {"path": book}})
    assert err and payload["error"]["code"] == "BAD_PARAMS"




@pytest.mark.parametrize("op,args,field,canonical", [
    ("clear_range", {"location": "A2", "what": "Contents"}, "what", "contents"),
    ("modify_grid_structure", {"action": "insert_columns", "at": "C"}, "action", "insert_cols"),
    ("modify_grid_structure", {"action": "InsertRows", "at": 2}, "action", "insert_rows"),
    ("manage_worksheet", {"action": "Add", "new_name": "X"}, "action", "add"),
    ("manage_worksheet", {"action": "hide", "sheet": "Sheet1", "state": "veryHidden"},
     "state", "very_hidden"),
])
def test_e2_near_miss_values_are_accepted_and_echoed(op, args, field, canonical, book):
    if op == "manage_worksheet" and args.get("action") == "hide":
        srv.manage_worksheet(book, "add", new_name="Other")
    tool = registry.REGISTRY[op].group
    payload, err, _ = _call(tool, op, path=book, **args)
    assert not err, payload
    assert payload["normalized_from"]["args"][field] == args[field]


def test_e2_unknown_values_still_refuse(book):
    payload, err, _ = _call("xl_edit", "clear_range", path=book, location="A2", what="most")
    assert err and payload["error"]["code"] == "BAD_PARAMS"


def test_e2_a_bare_search_string_is_a_search_location(book):
    payload, err, _ = _call("xl_read", "read_range", path=book, location={"search": "Qty"})
    assert not err, payload
    assert payload["normalized_from"]["args"]["location"] == {"search": "Qty"}


# --------------------------------------------------------- naming


@pytest.mark.parametrize("sent", ["greaterThan", "greater_than", "Greater Than"])
def test_s3_cf_operator_values_are_snake_case_and_camel_still_works(sent, book):
    payload, err, _ = _call("xl_format", "manage_conditional_format", path=book,
                            action="add", location="B2:B4", cf_type="cell_is",
                            params={"operator": sent, "formula": ["1"], "fill": "FFC7CE"})
    assert not err, payload
    if sent != "greater_than":
        assert payload["normalized_from"]["args"]["params.operator"] == sent
    listed, _, _ = _call("xl_format", "manage_conditional_format", path=book, action="list")
    assert listed["rules"][0]["operator"] == "greater_than"


def test_s3_chart_cell_is_at_cell_and_anchor_still_works(book):
    payload, err, _ = _call("xl_format", "manage_chart", path=book, action="create",
                            chart_type="bar", data="B1:B4", anchor="E2")
    assert not err, payload
    assert payload["normalized_from"]["args"]["anchor"] == "E2"
    assert "at_cell" in registry.REGISTRY["manage_chart"].args.model_fields
    assert "anchor" not in registry.REGISTRY["manage_chart"].args.model_fields


@pytest.mark.parametrize("sent", ["textLength", "text_length"])
def test_s3_validation_type_is_snake_case(sent, book):
    payload, err, _ = _call("xl_structure", "manage_data_validation", path=book,
                            action="add", location="C2", dv_type=sent,
                            operator="lessThan", formula1="10")
    assert not err, payload
    listed, _, _ = _call("xl_structure", "manage_data_validation", path=book, action="list")
    v = listed["validations"][0]
    assert v["type"] == "text_length" and v["operator"] == "less_than"


# ----------------------------------------------------------- the reroute


def test_a_read_op_sent_to_a_writing_tool_runs_and_says_so(book):
    payload, err, _ = _call("xl_edit", "read_range", path=book, location="A1")
    assert not err, payload
    assert payload["normalized_from"]["tool"] == "xl_edit"


def test_a_write_never_runs_through_the_read_only_tool(book):
    payload, err, _ = _call("xl_read", "set_cell", path=book, location="C1", value=1)
    assert err and payload["error"]["code"] == "BAD_PARAMS"
    assert payload["error"]["details"]["tool"] == "xl_edit"
    assert openpyxl.load_workbook(book).active["C1"].value is None
