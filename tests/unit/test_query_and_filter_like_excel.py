"""Reading and filtering the way Excel does.

  * text comparisons ignore case (eq, ne, in, gt...), as Excel's do
  * an aggregate over a column holding text says which cells it left out
  * sort keys and order_by read one direction vocabulary; nothing falls
    back to ascending silently
  * every filter criterion is stored in the autofilter, as Excel stores
    it, so Excel's Reapply shows the same rows (Excel is the judge)
"""

from __future__ import annotations

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp import server as srv
from xlsx_mcp.com import session as com_session

_COM = com_session.com_available()[0]


def _sales(p) -> str:
    p = str(p)
    srv.create_workbook(p)
    rows = [["Region", "Product", "Amount"]]
    for i in range(1, 21):
        rows.append(["North" if i % 2 else "South", "Gizmo Pro" if i % 3 else "Widget",
                     i * 100])
    srv.write_range(p, "A1", rows)
    return p


def _call(tool, op, **args):
    payload, err, _ = fh.start().call(tool, {"op": op, "args": args})
    assert not err, payload
    return payload


@pytest.mark.parametrize("op,value,expected", [
    ("eq", "gizmo pro", 14), ("ne", "GIZMO PRO", 6), ("in", ["widget"], 6)])
def test_x51_text_comparisons_ignore_case(op, value, expected, tmp_path):
    p = _sales(tmp_path / "s.xlsx")
    out = _call("xl_read", "query_range", path=p, where=[
        {"column": "Product", "op": op, "value": value}])
    assert out["matched"] == expected


def test_x51_ordered_text_comparisons_ignore_case(tmp_path):
    p = str(tmp_path / "t.xlsx")
    srv.create_workbook(p)
    srv.write_range(p, "A1", [["Name"], ["apple"], ["Banana"], ["cherry"]])
    out = _call("xl_read", "query_range", path=p, where=[
        {"column": "Name", "op": "gt", "value": "b"}])
    assert out["matched"] == 2  # Banana and cherry, as Excel orders text


def test_x15_text_left_out_of_a_sum_is_reported_with_its_cells(tmp_path):
    p = str(tmp_path / "m.xlsx")
    srv.create_workbook(p)
    srv.write_range(p, "A1", [["Amount"], [100], ["1,234.56"], [50], ["n/a"]])
    out = _call("xl_read", "query_range", path=p,
                aggregate=[{"column": "Amount", "func": "sum"}])
    assert out["groups"][0]["aggregates"]["sum_Amount"] == 150
    assert out["partial"] is True
    assert out["excluded_text"] == {"count": 2, "cells": ["A3", "A5"]}
    assert "A3" in out["warning"]


@pytest.mark.parametrize("key", [{"column": "Amount", "descending": True},
                                 {"column": "Amount", "order": "down"},
                                 {"column": "Amount", "dir": "Descending"},
                                 {"column": "Amount", "order": "z-a"}])
def test_x26_every_spelling_of_descending_sorts_descending(key, tmp_path):
    p = _sales(tmp_path / "s.xlsx")
    _call("xl_structure", "sort_range", path=p, location="A1:C21", keys=[key])
    ws = openpyxl.load_workbook(p).active
    assert [ws[f"C{r}"].value for r in (2, 3, 21)] == [2000, 1900, 100]


@pytest.mark.parametrize("key,needle", [
    ({"column": "Amount", "order": "sideways"}, "not understood"),
    ({"column": "Amount", "descending": "yes"}, "true or false"),
    ({"column": "Amount", "order": "asc", "descending": True}, "two different directions"),
    ({"column": "Amount", "reverse": True}, "not one of its keys"),
])
def test_x26_an_unknown_direction_refuses_instead_of_sorting_ascending(key, needle,
                                                                       tmp_path):
    p = _sales(tmp_path / "s.xlsx")
    before = open(p, "rb").read()
    payload, err, _ = fh.start().call("xl_structure", {"op": "sort_range", "args": {
        "path": p, "location": "A1:C21", "keys": [key]}})
    assert err and needle in payload["error"]["message"], payload
    assert open(p, "rb").read() == before


def test_x26_order_by_reads_the_same_vocabulary(tmp_path):
    p = _sales(tmp_path / "s.xlsx")
    out = _call("xl_read", "query_range", path=p, columns=["Amount"],
                order_by=[{"column": "Amount", "descending": True}], limit=2)
    assert out["rows"] == [[2000], [1900]]
    payload, err, _ = fh.start().call("xl_read", {"op": "query_range", "args": {
        "path": p, "order_by": [{}]}})
    assert err and payload["error"]["code"] == "BAD_PARAMS"  # was INTERNAL_ERROR


def _stored_filter(p):
    ws = openpyxl.load_workbook(p).active
    return ws.auto_filter


def test_x27_every_criterion_is_stored(tmp_path):
    p = _sales(tmp_path / "s.xlsx")
    out = _call("xl_structure", "set_filter", path=p, location="A1:C21", criteria=[
        {"column": "Region", "op": "eq", "value": "North"},
        {"column": "Amount", "op": "gt", "value": 1000}])
    af = _stored_filter(p)
    cols = {fc.colId: fc for fc in af.filterColumn}
    assert set(cols) == {0, 2}
    assert [f for f in cols[0].filters.filter] == ["North"]
    cf = cols[2].customFilters.customFilter
    assert [(c.operator, c.val) for c in cf] == [("greaterThan", "1000")]
    assert out["changed"]["filter"]["data_rows_visible"] == 5


def test_x27_a_criterion_excel_cannot_hold_is_said(tmp_path):
    p = _sales(tmp_path / "s.xlsx")
    out = _call("xl_structure", "set_filter", path=p, location="A1:C21", criteria=[
        {"column": "Product", "op": "regex", "value": "^Giz"}])
    assert any("cannot be stored" in w for w in out["warnings"])


def _excel_visible_rows(path: str) -> int:
    def job(manager):
        w = manager.acquire()
        prior = com_session._assert_hygiene(w.app)
        wb = None
        try:
            wb = com_session.open_workbook(w.app, path)
            ws = wb.Worksheets(1)
            ws.AutoFilter.ApplyFilter()
            return sum(1 for r in range(2, 22) if not ws.Rows(r).Hidden)
        finally:
            if wb is not None:
                wb.Close(False)
            com_session._restore_hygiene(w.app, *prior)
            manager.release_to_pool(w)
    return com_session.run_com("oracle", job, timeout=120)


@pytest.mark.com
@pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool)")
@pytest.mark.timeout(300)
@pytest.mark.parametrize("criteria", [
    [{"column": "Region", "op": "eq", "value": "north"},
     {"column": "Amount", "op": "gt", "value": 1000}],
    [{"column": "Product", "op": "contains", "value": "pro"},
     {"column": "Amount", "op": "le", "value": 900}],
    [{"column": "Product", "op": "ne", "value": "widget"}],
])
def test_excel_reapplies_the_same_rows(criteria, tmp_path):
    p = _sales(tmp_path / "s.xlsx")
    out = _call("xl_structure", "set_filter", path=p, location="A1:C21", criteria=criteria)
    assert _excel_visible_rows(p) == out["changed"]["filter"]["data_rows_visible"]
