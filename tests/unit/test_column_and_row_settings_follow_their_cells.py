"""Row and column settings follow their rows and columns through inserts and
deletes, on a made-up workbook: a wide column, one style on a span of three
columns, a hidden column, a collapsed outline group, a nested group, and a
hidden span reaching the sheet's last column (16384); and on the row side a
tall row, hidden rows and a collapsed group.

openpyxl moves the cells but not the settings, so before this a width, a
hidden column or a group stayed at its old letter and landed on whatever
moved there. Every expected setting below is what Excel itself stored after
the same edit (Columns("D:D").Insert() and the rest, made by a private hidden
Excel on the same workbook). Measured there too: an inserted column or row
takes the width or height, the style and the outline level of its neighbour
to the left or above, never its hidden state.
"""

from __future__ import annotations

import shutil

import openpyxl
import pytest
from openpyxl.styles import PatternFill
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.dimensions import ColumnDimension

import family_harness as fh
from xlsx_mcp.com import session as com_session

_COM = com_session.com_available()[0]
FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
LAST = 16384


def _ok(tool, op, **args):
    payload, err, _ = fh.start().call(tool, {"op": op, "args": args})
    assert not err, (op, payload)
    return payload


def _span(ws, first, last, **kw):
    cd = ColumnDimension(ws, index=first, **kw)
    cd.min, cd.max = column_index_from_string(first), last
    return cd


def _workbook(path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cols"
    for r in range(1, 6):
        for c in range(1, 17):
            ws.cell(r, c, r * 100 + c)
    ws["R1"] = "=SUM(A1:P1)"
    cd = ws.column_dimensions
    cd["B"].width = 20
    styled = _span(ws, "C", 5, width=15)
    styled.number_format = "0.00"
    styled.fill = FILL
    cd["C"] = styled
    cd["F"].hidden = True
    cd["F"].width = 12
    cd["G"] = _span(ws, "G", 8, outline_level=1, hidden=True, width=9)
    cd["I"].collapsed = True
    cd["I"].width = 13
    cd["J"] = _span(ws, "J", 12, outline_level=1, width=11)
    cd["M"] = _span(ws, "M", 13, outline_level=2, width=7)
    cd["T"] = _span(ws, "T", LAST, hidden=True, width=9)
    rows = wb.create_sheet("Rows")
    for r in range(1, 16):
        rows.cell(r, 1, r)
    rd = rows.row_dimensions
    rd[2].height = 30
    rd[4].hidden = True
    rd[5].hidden = True
    for r in (6, 7):
        rd[r].outline_level = 1
        rd[r].hidden = True
    rd[8].collapsed = True
    rd[9].height = 22
    wb.save(path)
    return str(path)


def _columns(path) -> dict:
    """Column number -> (width, hidden, outline level, number format) for every
    column a setting covers, the last column included."""
    ws = openpyxl.load_workbook(path)["Cols"]
    out = {}
    for cd in ws.column_dimensions.values():
        lo = cd.min or column_index_from_string(cd.index)
        hi = cd.max or lo
        for c in list(range(lo, min(hi, 30) + 1)) + ([LAST] if hi == LAST else []):
            out[c] = (cd.width, bool(cd.hidden), cd.outline_level, cd.number_format)
    return out


def _rows(path) -> dict:
    ws = openpyxl.load_workbook(path)["Rows"]
    return {r: (rd.height, bool(rd.hidden), rd.outline_level)
            for r, rd in ws.row_dimensions.items() if rd.height or rd.hidden or rd.outline_level}


def _expect(spec) -> dict:
    out = {}
    for (lo, hi), setting in spec.items():
        for c in list(range(lo, min(hi, 30) + 1)) + ([LAST] if hi == LAST else []):
            out[c] = setting
    return out


W = "General"


def test_a_column_inserted_inside_a_span_widens_it_and_moves_the_rest(tmp_path):
    p = _workbook(tmp_path / "w.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_cols", at="D", sheet="Cols")
    assert _columns(p) == _expect({
        (2, 2): (20, False, 0, W),
        (3, 6): (15, False, 0, "0.00"),      # C:E was one span; the new D joins it
        (7, 7): (12, True, 0, W),
        (8, 9): (9, True, 1, W),
        (10, 10): (13, False, 0, W),
        (11, 13): (11, False, 1, W),
        (14, 14): (7, False, 2, W),
        (21, LAST): (9, True, 0, W),          # still reaches the last column
    })


def test_a_column_deleted_inside_a_span_shortens_it_and_moves_the_rest(tmp_path):
    p = _workbook(tmp_path / "w.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="delete_cols", at="D", sheet="Cols")
    assert _columns(p) == _expect({
        (2, 2): (20, False, 0, W),
        (3, 4): (15, False, 0, "0.00"),
        (5, 5): (12, True, 0, W),
        (6, 7): (9, True, 1, W),
        (8, 8): (13, False, 0, W),
        (9, 11): (11, False, 1, W),
        (12, 12): (7, False, 2, W),
        (19, LAST): (9, True, 0, W),          # the freed last column stays hidden
    })


def test_a_column_inserted_inside_a_hidden_group_is_in_the_group_but_shown(tmp_path):
    p = _workbook(tmp_path / "w.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_cols", at="H", sheet="Cols")
    cols = _columns(p)
    assert cols[7] == (9, True, 1, W)
    assert cols[8] == (9, False, 1, W)        # Excel: width and level from G, not hidden
    assert cols[9] == (9, True, 1, W)
    assert cols[10] == (13, False, 0, W)
    assert cols[14] == (7, False, 2, W)
    assert cols[21] == cols[LAST] == (9, True, 0, W)


def test_two_columns_deleted_across_settings(tmp_path):
    p = _workbook(tmp_path / "w.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="delete_cols", at="F", count=2,
        sheet="Cols")
    cols = _columns(p)
    assert cols[5] == (15, False, 0, "0.00")
    assert cols[6] == (9, True, 1, W)         # what was H
    assert cols[7] == (13, False, 0, W)       # what was I
    assert cols[18] == cols[LAST] == (9, True, 0, W)


def test_an_inserted_row_takes_the_height_of_the_row_above(tmp_path):
    p = _workbook(tmp_path / "w.xlsx")
    _ok("xl_structure", "modify_grid_structure", path=p, action="insert_rows", at=3, sheet="Rows")
    assert _rows(p) == {2: (30, False, 0), 3: (30, False, 0), 5: (None, True, 0),
                        6: (None, True, 0), 7: (None, True, 1), 8: (None, True, 1),
                        10: (22, False, 0)}


@pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool)")
@pytest.mark.timeout(600)
def test_excel_shows_the_same_columns_after_the_same_edit(tmp_path):
    """The same edits made by Excel itself, read back through Excel."""
    def read(path):
        out = {}

        def job(manager):
            w = manager.acquire()
            prior = com_session._assert_hygiene(w.app)
            wb = None
            try:
                wb = com_session.open_workbook(w.app, path)
                ws = wb.Worksheets("Cols")
                for c in list(range(2, 25)) + [LAST]:
                    col = ws.Columns(c)
                    out[c] = (round(col.ColumnWidth, 2), bool(col.Hidden), col.OutlineLevel,
                              ws.Cells(1000, c).NumberFormat)
            finally:
                if wb is not None:
                    wb.Close(False)
                com_session._restore_hygiene(w.app, *prior)
                manager.release_to_pool(w)
        com_session.run_com("oracle", job, timeout=120)
        return out

    def excel_edit(path, action):
        def job(manager):
            w = manager.acquire()
            prior = com_session._assert_hygiene(w.app)
            wb = None
            try:
                wb = com_session.open_workbook(w.app, path)
                action(wb.Worksheets("Cols"))
                wb.Save()
            finally:
                if wb is not None:
                    wb.Close(False)
                com_session._restore_hygiene(w.app, *prior)
                manager.release_to_pool(w)
        com_session.run_com("oracle", job, timeout=120)

    for i, (action, at, count, excel) in enumerate([
            ("insert_cols", "D", 1, lambda ws: ws.Columns("D:D").Insert()),
            ("delete_cols", "D", 1, lambda ws: ws.Columns("D:D").Delete()),
            ("insert_cols", "H", 1, lambda ws: ws.Columns("H:H").Insert()),
            ("delete_cols", "J", 3, lambda ws: ws.Columns("J:L").Delete())]):
        ours = _workbook(tmp_path / f"ours{i}.xlsx")
        theirs = str(tmp_path / f"excel{i}.xlsx")
        shutil.copy2(ours, theirs)
        _ok("xl_structure", "modify_grid_structure", path=ours, action=action, at=at,
            count=count, sheet="Cols")
        excel_edit(theirs, excel)
        assert read(ours) == read(theirs), (action, at, count)
