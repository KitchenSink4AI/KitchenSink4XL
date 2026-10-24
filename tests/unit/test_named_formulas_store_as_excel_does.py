"""A defined name's formula is stored the way Excel itself stores it.

Excel keeps a defined name's formula in the file with the same prefixes as a
cell formula: _xlfn. on its newer functions (LAMBDA, LET, XLOOKUP and the
rest), _xlpm. on every LET or LAMBDA name, and _xlop. on an optional LAMBDA
parameter where it is declared. A name stored without them reads as a call to
an unknown add-in function: a LAMBDA name calculated #NAME? wherever a cell
called it, and Excel then saved it as _xludf.LAMBDA.

MEASURED in Excel 16 (hidden, on a private desktop): names added through
Excel's own object model are stored as

    term     _xlfn.LAMBDA(_xlpm.a,_xlpm.a+1)
    twice    _xlfn.LAMBDA(_xlpm.a,_xlop.b,IF(_xlfn.ISOMITTED(_xlpm.b),_xlpm.a*2,_xlpm.a*_xlpm.b))
    bump     _xlfn.LET(_xlpm.k,10,_xlpm.k+Data!$A$1)
    lastrow  _xlfn.XLOOKUP(2,Data!$A:$A,Data!$A:$A,0)

and a cell that calls one keeps the bare call (term(A1)). The same file
written with the name bare opens, and =term(A1) is #NAME?; written with the
prefixes, it is 2.

No Excel here except the tests marked LIVE, which start the product's hidden
Excel (run them on a private desktop).
"""
from __future__ import annotations

import zipfile

import openpyxl
import pytest

from xlsx_mcp.com import session as com_session
from xlsx_mcp.core.errors import XlMcpError
from xlsx_mcp.ops import names as nm

LIVE = pytest.mark.skipif(not com_session.com_available()[0], reason="needs Excel (COM)")

#: (name, as typed, as Excel stores it), measured as above.
EXCEL_STORES = [
    ("term", "LAMBDA(a,a+1)", "_xlfn.LAMBDA(_xlpm.a,_xlpm.a+1)"),
    ("twice", "LAMBDA(a,[b],IF(ISOMITTED(b),a*2,a*b))",
     "_xlfn.LAMBDA(_xlpm.a,_xlop.b,IF(_xlfn.ISOMITTED(_xlpm.b),_xlpm.a*2,_xlpm.a*_xlpm.b))"),
    ("bump", "LET(k,10,k+Data!$A$1)", "_xlfn.LET(_xlpm.k,10,_xlpm.k+Data!$A$1)"),
    ("lastrow", "XLOOKUP(2,Data!$A:$A,Data!$A:$A,0)", "_xlfn.XLOOKUP(2,Data!$A:$A,Data!$A:$A,0)"),
]


@pytest.fixture()
def book(tmp_path):
    p = tmp_path / "names.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 1
    ws["A2"] = 2
    wb.create_sheet("Notes")["A1"] = "notes"
    wb.save(p)
    return p


def _stored_names(path) -> str:
    with zipfile.ZipFile(path) as z:
        xml = z.read("xl/workbook.xml").decode("utf8")
    i, j = xml.find("<definedNames>"), xml.find("</definedNames>")
    return xml[i:j] if i >= 0 else ""


@pytest.mark.parametrize("name,typed,stored", EXCEL_STORES, ids=[n for n, _t, _s in EXCEL_STORES])
def test_a_name_is_stored_as_excel_stores_it(book, name, typed, stored):
    out = nm.manage_name(str(book), "add", name=name, refers_to=typed)
    assert out["ok"] and out["verified"]
    assert f'<definedName name="{name}">{stored}</definedName>' in _stored_names(book)
    assert "_xludf" not in _stored_names(book)


@pytest.mark.parametrize("name,typed,stored", EXCEL_STORES, ids=[n for n, _t, _s in EXCEL_STORES])
def test_an_equals_sign_and_update_store_the_same(book, name, typed, stored):
    nm.manage_name(str(book), "add", name=name, refers_to="Data!$A$1")
    nm.manage_name(str(book), "update", name=name, refers_to="=" + typed)
    assert f'<definedName name="{name}">{stored}</definedName>' in _stored_names(book)


def test_a_definition_already_in_excels_form_is_stored_once(book):
    stored = EXCEL_STORES[1][2]
    nm.manage_name(str(book), "add", name="twice", refers_to=stored)
    assert f'<definedName name="twice">{stored}</definedName>' in _stored_names(book)


@pytest.mark.parametrize("plain", ["Data!$A$1:$A$2", "0.07", "SUM(Data!$A$1:$A$2)", "'Notes'!$A$1"])
def test_a_definition_without_newer_functions_is_stored_as_given(book, plain):
    nm.manage_name(str(book), "add", name="plain", refers_to=plain)
    wb = openpyxl.load_workbook(book)
    assert wb.defined_names["plain"].value == plain


def test_reads_show_the_definition_as_written(book):
    from xlsx_mcp.ops import lifecycle
    for name, typed, _stored in EXCEL_STORES:
        nm.manage_name(str(book), "add", name=name, refers_to=typed)
    listing = {n["name"]: n["refers_to"] for n in nm.manage_name(str(book), "list")["names"]}
    meta = {n["name"]: n["value"] for n in lifecycle.get_workbook_metadata(str(book))["defined_names"]}
    for name, typed, _stored in EXCEL_STORES:
        assert listing[name] == typed
        assert meta[name] == typed


def test_the_answer_shows_the_definition_as_written(book):
    out = nm.manage_name(str(book), "add", name="term", refers_to="LAMBDA(a,a+1)")
    assert out["changed"]["name"]["refers_to"] == "LAMBDA(a,a+1)"


def test_a_definition_excel_cannot_store_is_refused_and_nothing_changes(book):
    before = book.read_bytes()
    with pytest.raises(XlMcpError):
        nm.manage_name(str(book), "add", name="term", refers_to="LAMBDA (a,a+1)")
    assert book.read_bytes() == before


# ------------------------------------------------- cell formulas, the same rule

CELL_STORES = [
    ("=LET(x,A1,x+1)", "=_xlfn.LET(_xlpm.x,A1,_xlpm.x+1)"),
    ("=LAMBDA(a,a+1)(A1)", "=_xlfn.LAMBDA(_xlpm.a,_xlpm.a+1)(A1)"),
    ("=XLOOKUP(1,A1:A2,A1:A2)", "=_xlfn.XLOOKUP(1,A1:A2,A1:A2)"),
]


@pytest.mark.parametrize("typed,stored", CELL_STORES, ids=["let", "lambda", "xlookup"])
@pytest.mark.parametrize("route", ["set_formula", "set_cell", "write_range", "apply_edits"])
def test_every_cell_route_stores_the_prefixes(book, route, typed, stored):
    from xlsx_mcp.ops import cells, formulas, view
    if route == "set_formula":
        formulas.set_formula(str(book), "C1", typed, sheet="Data")
    elif route == "set_cell":
        cells.set_cell(str(book), "C1", typed, sheet="Data")
    elif route == "write_range":
        cells.write_range(str(book), "C1", [[typed]], sheet="Data")
    else:
        view.apply_edits(str(book), [{"op": "set_formula", "location": "C1", "sheet": "Data",
                                      "formula": typed}])
    wb = openpyxl.load_workbook(book)
    v = wb["Data"]["C1"].value
    assert (v.text if hasattr(v, "text") else v) == stored


# ------------------------------------------------------------- in real Excel

@LIVE
@pytest.mark.timeout(300)
def test_a_named_lambda_and_let_calculate_in_excel(book):
    """The names are added and called through the product, Excel recalculates
    the workbook, and every call has its value (none is #NAME?); the file
    Excel saved still holds the names as Excel stores them."""
    from xlsx_mcp.ops import comtier, formulas
    for name, typed, _stored in EXCEL_STORES:
        nm.manage_name(str(book), "add", name=name, refers_to=typed)
    for cell, f in (("B1", "=term(A1)"), ("B2", "=twice(A1)"), ("B3", "=twice(A1,5)"),
                    ("B4", "=bump"), ("B5", "=LET(y,term(A2),y*3)"), ("B6", "=lastrow")):
        formulas.set_formula(str(book), cell, f, sheet="Data")
    try:
        out = comtier.recalculate(str(book), engine="com", timeout_seconds=180)
        assert out.get("ok", True), out
    finally:
        com_session.get_executor().shutdown()
    wb = openpyxl.load_workbook(book, data_only=True)
    ws = wb["Data"]
    assert [ws[c].value for c in ("B1", "B2", "B3", "B4", "B5", "B6")] == [2, 2, 5, 11, 9, 2]
    names = _stored_names(book)
    assert "_xludf" not in names
    for name, _typed, stored in EXCEL_STORES:
        assert f'<definedName name="{name}">{stored}</definedName>' in names
