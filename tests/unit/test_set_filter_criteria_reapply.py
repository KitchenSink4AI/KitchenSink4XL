"""set_filter's stored criteria re-apply in Excel to the
rows the server showed.

A trial judged 19 table-filter cases against Excel's own Reapply; 6
disagreed: Excel matches a value list (eq, in) against the text a cell SHOWS
(1000 under #,##0 shows "1,000", 0.5 under 0% shows "50%"), keeps text codes
apart ("007" is not "7"), and reads * and ? in custom filters as wildcards.
Now eq / in store the shown text of the rows they pick and keep a row when its
shown text is on that list; ne / not_in follow Excel's notEqual (a criterion
that reads as a number removes only cells holding that number); wildcards are
escaped with Excel's ~. Judged by Excel itself on its own workbook: 30 of 30
cases now re-apply to the same rows. The expected rows below are what Excel's
Reapply showed for the same data.
"""
from __future__ import annotations

import zipfile

import openpyxl
import pytest
from openpyxl.worksheet.table import Table

from xlsx_mcp.ops import sortfilter

ROWS = [  # Name, Amount, Pct, Code, Flag, Note, Price   (sheet rows 2..11)
    ("apple", 1000, 0.5, "007", True, "x", 1.5),
    ("Apple", 1500, 0.25, "7", False, None, 2),
    ("Banana", 999, 0.5, 7, True, "y", 1.25),
    ("banana split", 2500, 0.75, "007", False, None, 1.5),
    ("Cherry", 1000, 0.5, "8", True, "z", 3),
    ("apple pie", 50, 0.1, "7", False, "", 1.5),
    ("Date?", 0, 0.5, "9", True, "w", 2),
    ("*star", -5, 1, "007", False, None, 1.25),
    (None, 1000, 0.5, "7", True, "v", 2),
    ("cherry", 12345.5, 0.3, "07", False, "u", 0.5),
]


@pytest.fixture
def book(tmp_path):
    p = str(tmp_path / "tf.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Amount", "Pct", "Code", "Flag", "Note", "Price"])
    for row in ROWS:
        ws.append(list(row))
    for r in range(2, 12):
        ws.cell(r, 2).number_format = "#,##0"
        ws.cell(r, 3).number_format = "0%"
        if isinstance(ws.cell(r, 4).value, str):
            ws.cell(r, 4).number_format = "@"
        ws.cell(r, 7).number_format = "0.00"
    ws.add_table(Table(displayName="T", ref="A1:G11"))
    wb.save(p)
    return p


def _hidden(p):
    ws = openpyxl.load_workbook(p)["Data"]
    return [r for r in range(2, 12) if ws.row_dimensions[r].hidden]


def _table_xml(p):
    with zipfile.ZipFile(p) as z:
        return next(z.read(n).decode() for n in z.namelist() if n.startswith("xl/tables/"))


CASES = {
    # the six found re-applying differently
    "amount_eq_1000": ([{"column": "Amount", "op": "eq", "value": 1000}], [3, 4, 5, 7, 8, 9, 11]),
    "pct_eq_0.5": ([{"column": "Pct", "op": "eq", "value": 0.5}], [3, 5, 7, 9, 11]),
    "price_in": ([{"column": "Price", "op": "in", "value": [1.5, 2]}], [4, 6, 9, 11]),
    "code_eq_007": ([{"column": "Code", "op": "eq", "value": "007"}], [3, 4, 6, 7, 8, 10, 11]),
    "code_eq_7": ([{"column": "Code", "op": "eq", "value": "7"}], [2, 5, 6, 8, 9, 11]),
    "name_contains_q": ([{"column": "Name", "op": "contains", "value": "?"}],
                        [2, 3, 4, 5, 6, 7, 9, 10, 11]),
    # neighbours that must keep agreeing
    "code_eq_int_7": ([{"column": "Code", "op": "eq", "value": 7}], [2, 5, 6, 8, 9, 11]),
    "code_ne_7": ([{"column": "Code", "op": "ne", "value": "7"}], [4]),
    "code_ne_007": ([{"column": "Code", "op": "ne", "value": "007"}], [4]),
    "code_not_in_7_8": ([{"column": "Code", "op": "not_in", "value": ["7", "8"]}], [4]),
    "amount_ne_1000": ([{"column": "Amount", "op": "ne", "value": 1000}], [2, 6, 10]),
    "price_ne_1.5": ([{"column": "Price", "op": "ne", "value": 1.5}], [2, 5, 7]),
    "name_ne_date_q": ([{"column": "Name", "op": "ne", "value": "Date?"}], [8]),
    "name_not_in": ([{"column": "Name", "op": "not_in", "value": ["Date?", "*star"]}], [8, 9]),
    "name_eq_star": ([{"column": "Name", "op": "eq", "value": "*star"}],
                     [2, 3, 4, 5, 6, 7, 8, 10, 11]),
    "name_contains_star": ([{"column": "Name", "op": "contains", "value": "*"}],
                           [2, 3, 4, 5, 6, 7, 8, 10, 11]),
    "name_endswith_q": ([{"column": "Name", "op": "endswith", "value": "?"}],
                        [2, 3, 4, 5, 6, 7, 9, 10, 11]),
    "name_eq_apple": ([{"column": "Name", "op": "eq", "value": "apple"}], [4, 5, 6, 7, 8, 9, 10, 11]),
    "name_in": ([{"column": "Name", "op": "in", "value": ["apple", "Cherry"]}], [4, 5, 7, 8, 9, 10]),
    "amount_between": ([{"column": "Amount", "op": "ge", "value": 50},
                        {"column": "Amount", "op": "le", "value": 1000}], [3, 5, 8, 9, 11]),
    "note_is_blank": ([{"column": "Note", "op": "is_blank"}], [2, 4, 6, 8, 10, 11]),
}


@pytest.mark.parametrize("case", sorted(CASES))
def test_the_rows_shown_are_the_rows_excel_reapplies(book, case):
    criteria, excel_hidden = CASES[case]
    sortfilter.set_filter(book, {"table": "T"}, criteria, backup=False)
    assert _hidden(book) == excel_hidden


def test_a_value_list_stores_the_text_each_cell_shows(book):
    sortfilter.set_filter(book, {"table": "T"},
                          [{"column": "Amount", "op": "eq", "value": 1000}], backup=False)
    assert 'val="1,000"' in _table_xml(book)
    sortfilter.set_filter(book, {"table": "T"},
                          [{"column": "Price", "op": "in", "value": [1.5, 2]}], backup=False)
    xml = _table_xml(book)
    assert 'val="1.50"' in xml and 'val="2.00"' in xml


def test_wildcards_are_escaped_where_excel_reads_them(book):
    sortfilter.set_filter(book, {"table": "T"},
                          [{"column": "Name", "op": "contains", "value": "?"}], backup=False)
    assert 'val="*~?*"' in _table_xml(book)
    sortfilter.set_filter(book, {"table": "T"},
                          [{"column": "Name", "op": "ne", "value": "Date?"}], backup=False)
    assert 'val="Date~?"' in _table_xml(book)


def test_the_escape_itself():
    assert sortfilter._escape_wildcards("a*b?c~") == "a~*b~?c~~"
