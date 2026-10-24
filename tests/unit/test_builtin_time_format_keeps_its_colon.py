"""Excel's built-in number format 47 shows mm:ss.0, colon included.

openpyxl's table of built-in formats spells id 47 "mmss.0", without its
colon. Excel's own save of an mm:ss.0 cell writes numFmtId 47, and Excel
shows 59.5 seconds under it as 00:59.5 (measured with a private hidden Excel).
Every read that gives the shown text must give Excel's. No Excel here; every
workbook is made up and built by code.
"""
from __future__ import annotations

import zipfile

import openpyxl
import pytest

from xlsx_mcp.core import numfmt
from xlsx_mcp.ops import cells


def _book(tmp_path, value):
    p = tmp_path / "made.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Shown"
    ws.column_dimensions["A"].width = 60
    ws["A1"] = value
    ws["A1"].number_format = "mmss.0"  # openpyxl stores this text as built-in id 47
    wb.save(p)
    return str(p)


@pytest.mark.parametrize("seconds,excel", [(59.5, "00:59.5"), (0.5, "00:00.5"), (3599.5, "59:59.5"),
                                           (30, "00:30.0")])
def test_builtin_format_47_reads_as_excel_shows_it(tmp_path, seconds, excel):
    p = _book(tmp_path, seconds / 86400)
    with zipfile.ZipFile(p) as z:
        assert 'numFmtId="47"' in z.read("xl/styles.xml").decode()
    shown = cells.read_range(p, "A1", sheet="Shown", values="displayed")["values"][0][0]
    default = cells.read_range(p, "A1", sheet="Shown")
    told = default.get("shows", {}).get("A1", str(default["values"][0][0]))
    assert (shown, told) == (excel, excel)


def test_the_display_maps_the_spelling_itself():
    assert numfmt.display(59.5 / 86400, "mmss.0") == "00:59.5"
