"""The full rewrite keeps a sensitivity label, and (found while
fixing a related defect) the password to modify.

A Microsoft Purview sensitivity label lives in docMetadata/LabelInfo.xml,
reached from the ROOT relationships (type .../2020/02/relationships/
classificationlabels). The rewrite dropped it, so every rewrite-route edit on a
labelled workbook failed "unexplained part loss", even with allow_loss, while
diagnose said the rewrite keeps everything. The rewrite now carries it
byte-for-byte, and diagnose agrees with the save.

The rewrite also dropped <fileSharing> from workbook.xml, silently
removing a workbook's password to modify. It is carried back.
"""
from __future__ import annotations

import zipfile

import openpyxl
import pytest

from xlsx_mcp.core import hazard

LABEL = (b'<?xml version="1.0" encoding="utf-8" standalone="yes"?>'
         b'<clbl:labelList xmlns:clbl="http://schemas.microsoft.com/office/2020/mipLabelMetadata">'
         b'<clbl:label id="{11111111-2222-3333-4444-555555555555}" enabled="1" method="Standard" siteId='
         b'"{66666666-7777-8888-9999-000000000000}" removed="0"/></clbl:labelList>')
LEGACY = b'<fileSharing userName="Owner" reservationPassword="CC3D"/>'
MODERN = (b'<fileSharing readOnlyRecommended="1" userName="Owner" '
          b'algorithmName="SHA-512" hashValue="aGFzaA==" saltValue="c2FsdA==" '
          b'spinCount="100000"/>')


def _build(tmp_path, *, label=False, sharing=None, name="w.xlsx"):
    base = str(tmp_path / ("base_" + name))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, 6):
        ws.cell(r, 1, r)
        ws.cell(r, 2, f"=A{r}*2")
    wb.save(base)
    out = str(tmp_path / name)
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            b = zin.read(i.filename)
            if label and i.filename == "_rels/.rels":
                b = b.replace(b"</Relationships>",
                              b'<Relationship Id="rIdLbl" Type="http://schemas.microsoft.com/office/2020/02/'
                              b'relationships/classificationlabels" Target="docMetadata/LabelInfo.xml"/>'
                              b"</Relationships>")
            if label and i.filename == "[Content_Types].xml":
                b = b.replace(b"</Types>", b'<Override PartName="/docMetadata/LabelInfo.xml" ContentType='
                              b'"application/vnd.ms-office.classificationlabels+xml"/></Types>')
            if sharing and i.filename == "xl/workbook.xml":
                k = b.index(b"<workbookPr")
                b = b[:k] + sharing + b[k:]
            zout.writestr(i, b)
        if label:
            zout.writestr("docMetadata/LabelInfo.xml", LABEL)
    return out


def _insert(p, **kw):
    from xlsx_mcp.ops import structure
    return structure.modify_grid_structure(p, "insert_rows", 2, count=1, sheet="Data", **kw)


def _rename(p):
    from xlsx_mcp.ops import lifecycle
    return lifecycle.manage_worksheet(p, "rename", sheet="Data", new_name="Data2")


def _title(p):
    from xlsx_mcp.ops import properties
    return properties.set_workbook_properties(p, title="New title")


def _widths(p):
    from xlsx_mcp.ops import format as fmt
    return fmt.set_dimensions(p, sheet="Data", column_widths={"A": 20})


EDITS = {"insert_rows": _insert, "insert_rows_allow_loss": lambda p: _insert(p, allow_loss=True),
         "rename_sheet": _rename, "set_title": _title, "column_widths": _widths}


@pytest.mark.parametrize("edit", sorted(EDITS))
def test_rewrite_edits_keep_the_sensitivity_label(tmp_path, edit):
    p = _build(tmp_path, label=True)
    out = EDITS[edit](p)
    assert out.get("ok") is not False
    with zipfile.ZipFile(p) as z:
        assert z.read("docMetadata/LabelInfo.xml") == LABEL
        rels = z.read("_rels/.rels")
        assert b"classificationlabels" in rels and b"docMetadata/LabelInfo.xml" in rels
        assert b"/docMetadata/LabelInfo.xml" in z.read("[Content_Types].xml")


def test_insert_rows_route_was_the_rewrite(tmp_path):
    out = _insert(_build(tmp_path, label=True))
    assert out.get("route") == "openpyxl"


def test_diagnose_agrees_with_the_save(tmp_path):
    from xlsx_mcp.ops import lifecycle
    p = _build(tmp_path, label=True)
    route = lifecycle.diagnose_workbook(p)["routing"]["structural_edits"]["route"]
    _insert(p)  # raises if the save refuses
    assert route != "refuse"


@pytest.mark.parametrize("sharing", [LEGACY, MODERN], ids=["legacy", "modern"])
@pytest.mark.parametrize("edit", ["insert_rows", "rename_sheet", "set_title"])
def test_rewrite_edits_keep_the_password_to_modify(tmp_path, sharing, edit):
    p = _build(tmp_path, sharing=sharing)
    EDITS[edit](p)
    with zipfile.ZipFile(p) as z:
        wb = z.read("xl/workbook.xml")
    assert sharing in wb
    # schema order: fileSharing comes before workbookPr
    assert wb.index(b"<fileSharing") < wb.index(b"<workbookPr")
    assert hazard.write_reservation(p)["password_to_modify"]
    openpyxl.load_workbook(p)  # still a well-formed workbook
