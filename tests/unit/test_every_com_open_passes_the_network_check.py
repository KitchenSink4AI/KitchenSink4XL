"""Every COM operation opens its workbook through one network check.

Excel recalculates a workbook as it opens when the workbook asks for it, and
a workbook written by openpyxl or edited by this server asks. A formula such
as WEBSERVICE then fetches from the network, carrying whatever the workbook
puts in its address. So every route that opens a workbook in Excel passes the
same check (com.session.network_gate inside open_workbook): a workbook whose
formulas or data connections can reach the network is refused before Excel
opens it, unless the call passes allow_network=true.

NO test here starts Excel. The COM run is replaced by a stand-in Excel that
records every Workbooks.Open; the check must refuse before any open. Every
workbook is made up and built by code.
"""
from __future__ import annotations

import zipfile

import openpyxl
import pytest

from xlsx_mcp.com import session as com_session
from xlsx_mcp.core import netreach
from xlsx_mcp.core.errors import FormulaRejected
from xlsx_mcp.ops import comtier


class _Cell:
    def __init__(self, formula, address="$B$1"):
        self.Formula2 = formula
        self.Formula = formula
        self.Address = address


class _Range:
    def __init__(self, formulas):
        self._f = formulas

    def Find(self, What, LookIn, LookAt, MatchCase):
        for a, f in self._f.items():
            if What.lower() in f.lower():
                return _Cell(f, a)
        return None

    def FindNext(self, after):
        return None


class _Sheet:
    def __init__(self, formulas):
        self.Name = "Rates"
        self.UsedRange = _Range(formulas)
        self.Cells = type("C", (), {"FormatConditions": type("F", (), {"Count": 0})()})()


class _Book:
    def __init__(self, formulas):
        self.Worksheets = [_Sheet(formulas)]
        self.Names = []
        self.Connections = type("K", (), {"Count": 0})()
        self.FullName = ""
        self.Name = "book"

    def Close(self, *_a):
        pass


class _Books:
    def __init__(self, app):
        self.app = app

    def Add(self):
        self.app.log.append(("Add",))
        return _Book({})

    def Open(self, *args):
        self.app.log.append(("Open", args[0], args[2], self.app.Calculation))
        if self.app.inside is not None:
            return _Book(self.app.inside)
        raise RuntimeError("the stand-in Excel opened the workbook")

    @property
    def Count(self):
        return 0


class _App:
    def __init__(self, inside=None):
        self.log = []
        self.inside = inside
        self.DisplayAlerts = False
        self.EnableEvents = False
        self.Calculation = -4105
        self.Workbooks = _Books(self)


class _Worker:
    pid = 4242

    def __init__(self, app):
        self.app = app


class _Manager:
    def __init__(self, app):
        self.w = _Worker(app)

    def acquire(self):
        return self.w

    def release_to_pool(self, _w):
        pass


@pytest.fixture
def excel(monkeypatch):
    app = _App()
    monkeypatch.setattr(com_session, "com_available", lambda **_: (True, "a stand-in for the test"))
    monkeypatch.setattr(com_session, "run_com", lambda label, fn, timeout=None: fn(_Manager(app)))
    monkeypatch.setattr(comtier, "_clipboard_available", lambda: True)
    return app


def _fetching(tmp_path, name="rates.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws["A1"] = "made-up value"
    ws["B1"] = '=_xlfn.WEBSERVICE("http://127.0.0.1:9/x?v="&A1)'
    p = tmp_path / name
    wb.save(p)
    return str(p)


def _ops(tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    return {
        "com_validate_opens_clean": lambda p, **k: comtier.com_validate_opens_clean(p, **k),
        "com_validate_opens_clean_password": lambda p, **k: comtier.com_validate_opens_clean(p, password="x", **k),
        "com_export_pdf": lambda p, **k: comtier.com_export_pdf(p, str(out / "a.pdf"), **k),
        "com_render_sheet": lambda p, **k: comtier.com_render_sheet(p, str(out / "a.png"), **k),
        "com_convert_format": lambda p, **k: comtier.com_convert_format(p, str(out / "a.xlsb"), **k),
        "com_autofit": lambda p, **k: comtier.com_autofit(p, **k),
        "com_save_with_password": lambda p, **k: comtier.com_save_with_password(p, "made-up-pw", **k),
        "com_manage_pivot_list": lambda p, **k: comtier.com_manage_pivot(p, "list", **k),
        "com_manage_pivot_refresh": lambda p, **k: comtier.com_manage_pivot(p, "refresh", **k),
        "com_set_sparkline_list": lambda p, **k: comtier.com_set_sparkline(p, "list", **k),
        "com_set_sparkline_create": lambda p, **k: comtier.com_set_sparkline(
            p, "create", location="C1", source="A1:B1", **k),
        "recalculate": lambda p, **k: comtier.recalculate(p, engine="com", **k),
        "com_goal_seek": lambda p, **k: comtier.com_goal_seek(p, "B1", 1, "A1", **k),
    }


OPS = ["com_validate_opens_clean", "com_validate_opens_clean_password", "com_export_pdf", "com_render_sheet",
       "com_convert_format", "com_autofit", "com_save_with_password", "com_manage_pivot_list",
       "com_manage_pivot_refresh", "com_set_sparkline_list", "com_set_sparkline_create", "recalculate",
       "com_goal_seek"]


@pytest.mark.parametrize("op", OPS)
def test_no_com_op_opens_a_fetching_workbook_unasked(tmp_path, excel, op):
    p = _fetching(tmp_path)
    with pytest.raises(FormulaRejected) as ei:
        _ops(tmp_path)[op](p)
    assert "network_reach" in ei.value.detail
    assert "allow_network=true" in str(ei.value)
    assert [e for e in excel.log if e[0] == "Open"] == []


@pytest.mark.parametrize("op", OPS)
def test_allow_network_lets_every_com_op_open_it(tmp_path, excel, op):
    p = _fetching(tmp_path)
    try:
        _ops(tmp_path)[op](p, allow_network=True)
    except FormulaRejected as exc:
        pytest.fail(f"{op} refused with allow_network=true: {exc}")
    except Exception:  # noqa: BLE001 - the stand-in Excel cannot do the work itself
        pass
    assert [e for e in excel.log if e[0] == "Open"], f"{op} never reached Excel's open"


def test_a_workbook_with_nothing_that_fetches_opens_as_before(tmp_path, excel):
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.active["B1"] = "=A1*2"
    p = tmp_path / "plain.xlsx"
    wb.save(p)
    try:
        comtier.com_export_pdf(str(p), str(tmp_path / "a.pdf"))
    except FormulaRejected as exc:
        pytest.fail(str(exc))
    except Exception:  # noqa: BLE001
        pass
    opens = [e for e in excel.log if e[0] == "Open"]
    assert len(opens) == 1 and ("Add",) not in excel.log  # no check in Excel: the package was read


def test_a_text_file_with_a_fetching_formula_is_refused(tmp_path, excel):
    # Excel enters a text file's formulas as it loads them, so it would fetch even with calculation off
    p = tmp_path / "rates.csv"
    p.write_text('name,value\nmade-up,"=WEBSERVICE(""http://127.0.0.1:9/x"")"\n', encoding="utf-8")
    with pytest.raises(FormulaRejected):
        comtier.com_convert_format(str(p), str(tmp_path / "rates.xlsx"))
    assert [e for e in excel.log if e[0] == "Open"] == []


def test_a_text_file_without_one_is_not_refused(tmp_path):
    p = tmp_path / "plain.csv"
    p.write_text("name,value\nmade-up,=1+2\nimage of a cat,3\n", encoding="utf-8")
    assert netreach.file_reach(str(p)) == (netreach.READ_TEXT, None)


def _xlsb_like(tmp_path):
    """A package whose workbook part is binary (an .xlsb): its formulas are not text."""
    p = tmp_path / "book.xlsb"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/'
                   'package/2006/content-types"><Default Extension="rels" ContentType="application/'
                   'vnd.openxmlformats-package.relationships+xml"/><Override PartName="/xl/workbook.bin" '
                   'ContentType="application/vnd.ms-excel.sheet.binary.macroEnabled.main"/></Types>')
        z.writestr("_rels/.rels", '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/'
                   'package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
                   'officeDocument/2006/relationships/officeDocument" Target="xl/workbook.bin"/></Relationships>')
        z.writestr("xl/workbook.bin", b"\x83\x01\x00")
    return str(p)


def test_a_binary_workbook_is_checked_by_excel_with_calculation_off_first(tmp_path, monkeypatch):
    p = _xlsb_like(tmp_path)
    assert netreach.file_reach(p) == (netreach.READ_IN_EXCEL, None)
    app = _App(inside={"$B$1": '=WEBSERVICE("http://127.0.0.1:9/x")'})
    with pytest.raises(FormulaRejected) as ei:
        com_session.open_workbook(app, p, read_only=True)
    assert "calculation turned off" in str(ei.value)
    # a blank workbook first, calculation manual, then a read-only open: and only that one open
    assert app.log[0] == ("Add",)
    opens = [e for e in app.log if e[0] == "Open"]
    assert len(opens) == 1 and opens[0][2] is True and opens[0][3] == -4135


def test_a_binary_workbook_with_nothing_that_fetches_then_opens_as_usual(tmp_path):
    p = _xlsb_like(tmp_path)
    app = _App(inside={"$A$1": "=1+2"})
    com_session.open_workbook(app, p, read_only=True)
    opens = [e for e in app.log if e[0] == "Open"]
    assert len(opens) == 2  # the check (calculation off), then the open the caller asked for


def test_the_check_sees_formulas_from_the_package_alone(tmp_path):
    how, reach = netreach.file_reach(_fetching(tmp_path))
    assert how == netreach.READ_PACKAGE
    assert reach["formulas"][0]["functions"] == ["WEBSERVICE"]


def test_the_shared_open_refuses_by_itself_whatever_route_reaches_it(tmp_path):
    # the routes refuse early too; the shared open is the check that holds whichever route calls it
    app = _App()
    p = _fetching(tmp_path)
    for read_only in (True, False):
        with pytest.raises(FormulaRejected) as ei:
            com_session.open_workbook(app, p, read_only=read_only)
        assert "did not open it" in str(ei.value)
    assert [e for e in app.log if e[0] == "Open"] == []
    with pytest.raises(RuntimeError):
        com_session.open_workbook(app, p, read_only=True, allow_network=True)
    assert len([e for e in app.log if e[0] == "Open"]) == 1
