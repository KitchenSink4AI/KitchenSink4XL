"""Structured references, with Excel's own behaviour as the table of truth.

The expected forms below were measured with a private hidden Excel
(Range.Formula given the typed text, then the
saved sheet XML read back). test_excel_reads_the_stored_forms checks the
product's stored forms against Excel again, end to end, where Excel exists.
"""

from __future__ import annotations

import re
import zipfile

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp import server as srv
from xlsx_mcp.com import session as com_session
from xlsx_mcp.core import structrefs as sr
from xlsx_mcp.core.errors import XlMcpError

_COM = com_session.com_available()[0]

TABLES = {"t": sr.TableInfo("T", "Sheet", "D1:H3", ["k", "v", "x#", "a[b", "w"])}

#: typed -> what Excel stores.
EXCEL_STORES = [
    ("=T[@v]*2", "=T[[#This Row],[v]]*2"),
    ("=T[@[v]]+1", "=T[[#This Row],[v]]+1"),
    ("=SUM(T[x'#])", "=SUM(T[x'#])"),
    ("=SUM(T[a'[b])", "=SUM(T[a'[b])"),
    ("=SUM(T[[#Headers],[v]])", "=SUM(T[[#Headers],[v]])"),
    ("=COUNTA(T[@])", "=COUNTA(T[#This Row])"),
    ("=SUM( T[ [v]:[x'#] ] )", "=SUM( T[ [v]:[x'#] ] )"),
    ("=SUM(t[V])", "=SUM(T[v])"),
]


@pytest.mark.parametrize("typed,stored", EXCEL_STORES)
def test_the_file_form_is_what_excel_stores(typed, stored):
    assert sr.to_file_form(typed, TABLES, "Sheet", 2, 10) == stored


@pytest.mark.parametrize("typed,stored", [("=[@v]*3", "=T[[#This Row],[v]]*3"),
                                          ("=[v]", "=T[[#This Row],[v]]")])
def test_an_unqualified_reference_inside_the_table_is_qualified_as_excel_does(typed, stored):
    assert sr.to_file_form(typed, TABLES, "Sheet", 2, 8) == stored  # H2, inside D1:H3


@pytest.mark.parametrize("typed,why", [
    ("=SUM(T[ v ])", "no column ' v '"),
    ("=SUM(U[v])", "no table named 'U'"),
    ("=SUM(T[nope])", "no column 'nope'"),
    ("=[@v]*2", "with no table name"),
    ("=SUM(T[#Bogus])", "not one of #All"),
])
def test_what_excel_refuses_is_refused(typed, why):
    with pytest.raises(XlMcpError, match=re.escape(why)):
        sr.to_file_form(typed, TABLES, "Sheet", 2, 10)  # J2, outside the table


@pytest.mark.parametrize("text", [
    '=A1&"T[@v]"', "='[x] y'!A1", "=[1]Sheet1!A1", "=[Book.xlsx]Sheet1!A1",
    "=Book.xlsx!T[v]", "=SUM(A1:B2)", "=INDIRECT(\"R[1]C\",FALSE)",
    # LAMBDA's optional parameters (the full suite caught these on dd97738).
    "=LAMBDA(x,[y],x)", "=LAMBDA(x, [y] ,IF(ISOMITTED(y),x,x+y))(1,2)",
    "=_xlfn.LAMBDA(a,[f],f(a))(2,LAMBDA(q,q*3))",
])
def test_text_that_is_not_a_reference_to_this_workbooks_tables_is_left_alone(text):
    assert sr.to_file_form(text, TABLES, "Sheet", 2, 10) == text


def test_a_bracket_in_a_lambda_body_is_still_a_reference():
    # The last LAMBDA argument is its body, never an optional parameter.
    with pytest.raises(XlMcpError, match="no table name"):
        sr.to_file_form("=LAMBDA(x,[v])(1)", TABLES, "Sheet", 2, 10)
    assert sr.to_file_form("=LAMBDA(x,x*T[@v])(2)", TABLES, "Sheet", 2, 5) ==         "=LAMBDA(x,x*T[[#This Row],[v]])(2)"


def test_the_spaces_hint_names_the_form_excel_accepts():
    with pytest.raises(XlMcpError, match=re.escape("Write T[v]")):
        sr.to_file_form("=SUM(T[ v ])", TABLES, "Sheet", 2, 10)


# ------------------------------------------------------- F-T2 through the surface


def _book(p, header="x#") -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["D1"], ws["E1"], ws["D2"], ws["E2"], ws["D3"], ws["E3"] = "k", header, "a", 1, "b", 2
    ws["G1"], ws["H1"], ws["G2"], ws["H2"] = "k2", header, "c", 5
    wb.save(p)
    srv.create_table(str(p), "D1:E3", "T")
    srv.create_table(str(p), "G1:H2", "U")
    return str(p)


@pytest.mark.parametrize("header,formula", [
    ("x#", "=SUM(T[x'#])"),
    ("a[b", "=SUM(T[a'[b])"),
    ("x#", "=T[[#This Row],[x'#]]"),
    ("it's", "=SUM(T[it''s])"),
    ("x#", "=SUM(t[X'#])"),
])
def test_ft2_a_rename_is_refused_while_any_form_of_reference_names_the_column(
        header, formula, tmp_path):
    p = _book(tmp_path / "t.xlsx", header)
    wb = openpyxl.load_workbook(p)
    wb.active["J2"] = formula
    wb.save(p)
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_cell", "args": {
        "path": p, "location": "E1", "value": "y"}})
    assert err and "pointing at a column that no longer exists" in payload["error"]["message"], \
        payload
    assert openpyxl.load_workbook(p).active["E1"].value == header


def test_ft2_an_unqualified_reference_inside_the_table_counts(tmp_path):
    p = _book(tmp_path / "t.xlsx")
    wb = openpyxl.load_workbook(p)
    wb.active["F2"] = "=T[[#This Row],[x'#]]"  # outside: qualified
    wb.save(p)
    wb = openpyxl.load_workbook(p)
    ws = wb.active
    ws["F2"] = None
    ws["E2"] = "=[@[x'#]]"  # the unqualified file-less form, inside T
    wb.save(p)
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_cell", "args": {
        "path": p, "location": "E1", "value": "y"}})
    assert err, payload


def test_ft2_a_defined_name_counts(tmp_path):
    from openpyxl.workbook.defined_name import DefinedName
    p = _book(tmp_path / "t.xlsx")
    wb = openpyxl.load_workbook(p)
    wb.defined_names["Hashes"] = DefinedName("Hashes", attr_text="T[x'#]")
    wb.save(p)
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_cell", "args": {
        "path": p, "location": "E1", "value": "y"}})
    assert err and "Hashes" in payload["error"]["message"], payload


def test_ft2_a_reference_to_another_tables_column_does_not_block_the_rename(tmp_path):
    """U has a column x# too; a formula naming U[x'#] is not about T."""
    p = _book(tmp_path / "t.xlsx")
    wb = openpyxl.load_workbook(p)
    wb.active["J2"] = "=SUM(U[x'#])"
    wb.save(p)
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_cell", "args": {
        "path": p, "location": "E1", "value": "y"}})
    assert not err, payload
    wb = openpyxl.load_workbook(p)
    assert [c.name for c in wb.active.tables["T"].tableColumns] == ["k", "y"]


# ------------------------------------------------------- F-S through the surface


def _stored(p, cell) -> str:
    xml = zipfile.ZipFile(p).read("xl/worksheets/sheet1.xml").decode()
    m = re.search(rf'<c r="{cell}"[^>]*>.*?<f>(.*?)</f>', xml)
    return m.group(1) if m else ""


@pytest.mark.parametrize("tool,op,args", [
    ("xl_edit", "set_formula", {"location": "J2", "formula": "=T[@k]&\"!\""}),
    ("xl_edit", "set_cell", {"location": "J2", "value": "=T[@k]&\"!\""}),
    ("xl_edit", "write_range", {"location": "J2", "data": [["=T[@k]&\"!\""]]}),
    ("xl_edit", "apply_edits", {"edits": [{"op": "set_formula", "location": "J2",
                                           "formula": "=T[@k]&\"!\""}]}),
])
def test_fs_every_formula_writer_stores_the_file_form_and_says_so(tool, op, args, tmp_path):
    p = _book(tmp_path / "t.xlsx", "v")
    payload, err, _ = fh.start().call(tool, {"op": op, "args": {"path": p, **args}})
    assert not err, payload
    assert _stored(p, "J2") == "T[[#This Row],[k]]&amp;\"!\""
    assert "structured_references_stored_as" in json_dump(payload)


def json_dump(payload) -> str:
    import json
    return json.dumps(payload)


def test_fs_a_reference_excel_refuses_is_refused_before_anything_is_written(tmp_path):
    p = _book(tmp_path / "t.xlsx", "v")
    before = open(p, "rb").read()
    payload, err, _ = fh.start().call("xl_edit", {"op": "set_formula", "args": {
        "path": p, "location": "J2", "formula": "=SUM(T[ v ])"}})
    assert err and payload["error"]["code"] == "BAD_PARAMS", payload
    assert "Write T[v]" in payload["error"]["message"]
    assert open(p, "rb").read() == before


@pytest.mark.com
@pytest.mark.skipif(not _COM, reason="needs Excel (private hidden pool)")
@pytest.mark.timeout(300)
def test_excel_reads_the_stored_forms(tmp_path):
    """Excel as the judge: the workbook opens clean and computes every
    formula the product stored from typed shorthand."""
    p = _book(tmp_path / "t.xlsx", "v")
    srv.set_formula(p, "J1", "=SUM(t[V])")
    srv.set_formula(p, "J2", "=T[@v]*2")  # rows 2 and 3 are the table's rows
    srv.set_formula(p, "J3", "=COUNTA(T[@])")
    res = srv.com_validate_opens_clean(p)
    assert res["opens_clean"] is True, res
    srv.recalculate(p, engine="com")
    got = srv.read_range(p, "J1:J3", values="cached")
    assert [r[0] for r in got["values"]] == [3, 2, 2], got
