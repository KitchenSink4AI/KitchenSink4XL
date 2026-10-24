"""Long answers lead with their totals and paging fields.

A client may show only the first ~2,000 characters of an answer. The totals
of a long read (how many matches, rows, names or sheets there are) and its
paging block (truncated, what was left out, how to get the rest) sat after
the long list, where such a client never shows them, so a model saw a page
of matches and took it for all of them. Every long read now puts them first.
All content is made up.
"""

from __future__ import annotations

import asyncio
import json

import openpyxl
import pytest
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from xlsx_mcp import server as _server  # noqa: F401 (registers the ops)
from xlsx_mcp.settings import Settings
from xlsx_mcp.surface import dispatch, groups
from xlsx_mcp.surface import plan as _plan

ROWS = 2500
HEAD = 2000


@pytest.fixture(scope="module")
def big(tmp_path_factory):
    p = tmp_path_factory.mktemp("long") / "long_answers.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Region", "Units", "Price", "Total", "Note"])
    regions = ["North", "South", "East", "West"]
    for i in range(2, ROWS + 2):
        ws.append([regions[i % 4], i % 37, round(1.5 + i % 11, 2),
                   f"=B{i}*C{i}", f"Made-up note number {i} for the test"])
    ws.add_table(Table(displayName="Orders", ref=f"A1:E{ROWS + 1}"))
    for i in range(2, 202):
        wb.defined_names[f"Point_{i}"] = DefinedName(
            f"Point_{i}", attr_text=f"Data!$B${i}")
        ws.cell(i, 7).hyperlink = f"https://www.example.com/item/{i}"
    for i in range(2, 62):
        ws.conditional_formatting.add(
            f"B{i}", CellIsRule(operator="greaterThan", formula=[str(i)],
                                fill=PatternFill("solid", start_color="FFFF00")))
    for k in range(1, 121):
        wb.create_sheet(f"Region sheet {k:03d}")["A1"] = k
    wb.save(p)
    return str(p)


def _call(op: str, args: dict) -> tuple[str, dict]:
    pl = _plan.build(Settings())
    payload, is_error = asyncio.run(dispatch.dispatch(
        pl, groups.group_of(op), {"op": op, "args": args}))
    assert not is_error, payload
    return json.dumps(payload, ensure_ascii=False), payload


CASES = {
    "find_cells": ({"query": "North"}, ["total", "truncated", "next_offset"]),
    "query_range": ({"location": f"Data!A1:E{ROWS + 1}", "limit": 400},
                    ["matched", "truncated", "next_offset"]),
    "read_range": ({"location": f"Data!A1:E{ROWS + 1}"},
                   ["page", "omitted", "how_to_continue"]),
    "get_table": ({"name": "Orders"}, ["row_count", "page"]),
    "manage_name": ({"action": "list"}, ["count"]),
    "manage_conditional_format": ({"action": "list"}, ["count"]),
    "manage_hyperlink": ({"action": "list"}, ["count"]),
    "get_workbook_metadata": ({}, ["sheet_count"]),
}


@pytest.mark.parametrize("op", sorted(CASES))
def test_the_first_2000_characters_hold_the_totals_and_paging(big, op):
    args, keys = CASES[op]
    text, payload = _call(op, {"path": big, **args})
    assert len(text) > HEAD, f"{op}: the answer is short; the test needs a long one"
    head = text[:HEAD]
    for key in keys:
        assert f'"{key}"' in head, f"{op}: {key} is not in the first {HEAD} characters"
