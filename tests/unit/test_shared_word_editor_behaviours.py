"""Behaviours the Excel editor shares with the Word editor: neutral
metadata, the default note author, threaded comments and the migrate-from-v1
map.

  * a new workbook carries neutral metadata (no "openpyxl" author)
  * a note is signed with the user's own name by default, not the product
  * threaded comments read as threads (author, text, replies)
  * no served string names a 1.x-only route; migrate-from-v1 carries
    the full map, generated from the live tables
"""

from __future__ import annotations

import json
import re
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest

import family_harness as fh
from xlsx_mcp import envelope, registry
from xlsx_mcp import server as srv
from xlsx_mcp.surface import copy as copy_mod
from xlsx_mcp.surface import help as help_mod

ROOT = Path(__file__).resolve().parents[2]


def _call(tool, op, **args):
    payload, err, _ = fh.start().call(tool, {"op": op, "args": args})
    assert not err, payload
    return payload


def test_b4_a_new_workbook_names_no_author(tmp_path):
    p = str(tmp_path / "n.xlsx")
    _call("xl_workbook", "create_workbook", path=p)
    core = zipfile.ZipFile(p).read("docProps/core.xml").decode()
    assert "openpyxl" not in core.lower()
    assert "<dc:creator" not in core and "lastModifiedBy" not in core
    assert "dcterms:created" in core


def test_x33_the_default_note_author_is_not_the_product(tmp_path, monkeypatch):
    from xlsx_mcp.core import authorship
    monkeypatch.setattr(authorship, "_office_value",
                        lambda name: "Dana Reyes" if name == "UserName" else None)
    p = str(tmp_path / "c.xlsx")
    srv.create_workbook(p)
    _call("xl_format", "manage_comment", path=p, action="add", location="A1", text="note")
    assert openpyxl.load_workbook(p).active["A1"].comment.author == "Dana Reyes"
    _call("xl_format", "manage_comment", path=p, action="add", location="A2", text="n",
          author="Reviewer")
    assert openpyxl.load_workbook(p).active["A2"].comment.author == "Reviewer"


def _with_thread(p: Path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B8"] = 1
    from openpyxl.comments import Comment
    ws["B8"].comment = Comment("[Threaded comment]\n\nYour version of Excel allows you to "
                               "read this threaded comment", "tc={T1}")
    wb.save(p)
    person = ('<personList xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/'
              'threadedcomments"><person displayName="Ada Reviewer" id="{P1}" userId="x" '
              'providerId="None"/></personList>')
    threads = ('<ThreadedComments xmlns="http://schemas.microsoft.com/office/spreadsheetml/'
               '2018/threadedcomments"><threadedComment ref="B8" dT="2026-10-04T17:01:26.41" '
               'personId="{P1}" id="{T1}"><text>Please confirm the pick logic</text>'
               '</threadedComment><threadedComment ref="B8" personId="{P1}" id="{T2}" '
               'parentId="{T1}"><text>Confirmed</text></threadedComment></ThreadedComments>')
    tmp = str(p) + ".tmp"
    with zipfile.ZipFile(p) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/_rels/workbook.xml.rels":
                data = data.decode().replace("</Relationships>", (
                    '<Relationship Id="rIdP" Type="http://schemas.microsoft.com/office/2017/10/'
                    'relationships/person" Target="persons/person.xml"/></Relationships>')).encode()
            if item.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                data = data.decode().replace("</Relationships>", (
                    '<Relationship Id="rIdT" Type="http://schemas.microsoft.com/office/2017/10/'
                    'relationships/threadedComment" '
                    'Target="../threadedComments/threadedComment1.xml"/></Relationships>')).encode()
            zout.writestr(item, data)
        zout.writestr("xl/persons/person.xml", person)
        zout.writestr("xl/threadedComments/threadedComment1.xml", threads)
    shutil.move(tmp, p)
    return str(p)


def test_x40_a_threaded_comment_reads_as_its_thread(tmp_path):
    p = _with_thread(tmp_path / "t.xlsx")
    out = _call("xl_format", "manage_comment", path=p, action="list")
    (c,) = out["comments"]
    assert c["threaded"] is True and c["cell"] == "B8"
    assert c["author"] == "Ada Reviewer" and c["text"] == "Please confirm the pick logic"
    assert [r["text"] for r in c["replies"]] == ["Confirmed"]
    assert "tc=" not in json.dumps(out) and "Your version of Excel" not in json.dumps(out)




_STALE = re.compile(r"com pack|\bpacks?\b|in v1\b|enable_tools|disable_tools|live=|"
                    r"forthcoming|not offered in v1|is queued", re.IGNORECASE)


def _served() -> list[tuple[str, str]]:
    out = [("copy", s) for s in copy_mod.all_strings()]
    out += [(f"hint:{k}", v) for k, v in envelope.HINTS.items()]
    out.append(("recovery", envelope.RECOVERY_HINT))
    for name, spec in registry.public_ops().items():
        out.append((f"help:{name}", help_mod._detail(spec)))
    from xlsx_mcp.ops import workflows
    for task in workflows.WORKFLOWS:
        wf = workflows.get_workflows(task)
        wf.pop("map", None)
        out.append((f"workflow:{task}", json.dumps(wf)))
    out.append(("workflows", json.dumps(workflows.get_workflows())))
    return out


def test_b7_no_served_string_names_a_1x_only_route():
    hits = [(where, m.group(0)) for where, text in _served()
            for m in _STALE.finditer(text)
            if not (where == "workflow:migrate-from-v1")]
    assert hits == [], hits


def test_b7_migrate_from_v1_maps_every_released_tool_with_its_changes():
    released = json.loads((ROOT / "migration" / "released_tools_1.2.5.json")
                          .read_text("utf-8"))["tools"]
    out = _call("xl_read", "get_workflows", task="migrate")
    assert out["task"] == "migrate-from-v1"
    rows = {r["old"]: r for r in out["map"]}
    assert set(released) <= set(rows)
    assert rows["enable_tools"]["op"] is None
    assert rows["manage_chart"]["changed"]["anchor"] == "at_cell"
    assert rows["manage_data_validation"]["changed"]["dv_type=textLength"] == \
        "dv_type=text_length"
    for r in rows.values():
        if r["op"]:
            assert r["op"] in registry.public_ops() and r["tool"]
