"""The silent data-loss defects of released Excel 1.2.5 that the 2.0
family surface fixes. Each one is proved with a regression test.

Every test here calls the ENGINE functions (srv.set_cell and friends), whose
names and signatures are the same in 1.2.5 and 2.0, and judges the file on
disk, so this one file runs RED against the released 1.2.5 wheel and GREEN on
2.0. It imports nothing that exists only in 2.0.

  * a write into one cell of a legacy array formula silently deleted the
    array, or was reported verified while Excel kept the old value
  * an extensionless PDF output silently replaced the user's <name>.pdf
  * a write into a table header (a number, or the column renamed in the
    cell only) left a workbook Excel will not open without a repair
  * Excel's typing shorthand for structured references (T[@v], T[ v ])
    was stored as typed; Excel opens it as #REF!
"""

from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula

from xlsx_mcp import server as srv


def _sha(p) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def _com_ok() -> bool:
    try:
        from xlsx_mcp.com import session
        return bool(session.com_available()[0])
    except Exception:  # noqa: BLE001
        return False


# ------------------------------------------------------------------ a write into part of an array formula


def _array_book(p: Path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B1"], ws["B2"] = 1, 2
    ws["A1"] = ArrayFormula("A1:A2", "=B1:B2*2")
    wb.save(p)
    return str(p)


@pytest.mark.parametrize("cell", ["A1", "A2"])
def test_fa_a_write_into_part_of_an_array_never_lands_silently(cell, tmp_path):
    p = _array_book(tmp_path / "a.xlsx")
    before = _sha(p)
    try:
        srv.set_cell(p, cell, 5)
    except Exception:  # noqa: BLE001 (a refusal is the right answer)
        assert _sha(p) == before, "refused, but the file changed anyway"
        return
    wb = openpyxl.load_workbook(p)
    a1 = wb.active["A1"].value
    pytest.fail(f"writing {cell} inside the array A1:A2 was accepted; A1 is now {a1!r} "
                "(Excel: the array is deleted, or the written value is discarded)")


# ------------------------------------------------------------------ an extensionless PDF output


@pytest.mark.com
@pytest.mark.skipif(not _com_ok(), reason="needs Excel (the product's private hidden pool)")
@pytest.mark.timeout(300)
def test_fe_an_extensionless_pdf_output_never_replaces_the_users_pdf(tmp_path):
    book = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "figures"
    wb.save(book)
    users = tmp_path / "report.pdf"
    users.write_bytes(b"%PDF-1.4 the user's own report, not to be replaced\n")
    before = _sha(users)
    try:
        srv.com_export_pdf(str(book), str(tmp_path / "report"))
    except Exception:  # noqa: BLE001 (refusing is fine; replacing is not)
        pass
    assert _sha(users) == before, "report.pdf was replaced by an export to 'report'"


# ------------------------------------------------------------------ a write into a table header


def _table_book(p: Path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["D1"], ws["E1"], ws["D2"], ws["E2"], ws["D3"], ws["E3"] = "k", "v", "a", 1, "b", 2
    wb.save(p)
    srv.create_table(str(p), "D1:E3", "T")
    return str(p)


def _table_is_openable(p: str) -> tuple[bool, str]:
    """Excel's rule: every header cell is non-empty, unique text equal to
    its table column's name."""
    wb = openpyxl.load_workbook(p)
    ws = wb.active
    t = ws.tables["T"]
    names = [c.name for c in t.tableColumns]
    heads = [ws.cell(1, c).value for c in (4, 5)]
    ok = all(isinstance(h, str) and h for h in heads) and heads == names
    return ok, f"headers {heads!r} vs column names {names!r}"


@pytest.mark.parametrize("value", [5, "=1+1"])
def test_ft_a_header_write_never_leaves_a_table_excel_must_repair(value, tmp_path):
    p = _table_book(tmp_path / "t.xlsx")
    before = _sha(p)
    try:
        srv.set_cell(p, "E1", value)
    except Exception:  # noqa: BLE001
        assert _sha(p) == before
        return
    ok, why = _table_is_openable(p)
    assert ok, f"accepted, and Excel will not open the table without a repair: {why}"


def test_ft_a_header_renamed_in_its_cell_renames_the_column_too(tmp_path):
    p = _table_book(tmp_path / "t.xlsx")
    try:
        srv.set_cell(p, "E1", "units")
    except Exception:  # noqa: BLE001
        return  # a refusal is safe
    ok, why = _table_is_openable(p)
    assert ok, why


# ------------------------------------------------------------------ structured-reference shorthand


def _stored_formula(p: str, cell: str) -> str:
    xml = zipfile.ZipFile(p).read("xl/worksheets/sheet1.xml").decode()
    m = re.search(rf'<c r="{cell}"[^>]*>.*?<f>(.*?)</f>', xml)
    return m.group(1) if m else ""


@pytest.mark.parametrize("formula", ["=T[@v]*2", "=SUM(T[ v ])", "=T[@[v]]+1"])
def test_fs_structured_reference_shorthand_is_stored_in_file_form(formula, tmp_path):
    p = _table_book(tmp_path / "t.xlsx")
    before = _sha(p)
    try:
        srv.set_formula(p, "G2", formula)
    except Exception:  # noqa: BLE001
        assert _sha(p) == before
        return
    stored = _stored_formula(p, "G2")
    assert "[@" not in stored and "[ " not in stored and " ]" not in stored, \
        f"stored {stored!r}, which Excel opens as #REF!"
