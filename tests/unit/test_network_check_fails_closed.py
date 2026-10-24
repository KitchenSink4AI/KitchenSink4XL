"""The network check fails closed: what it cannot read with certainty counts as reaching the network.

The check reads every member of a workbook package with a real XML parser (and the raw text as well), so a call is
found however the markup spells it and wherever in the package it sits. When the file cannot be read with certainty,
the answer is never "nothing found": a member that cannot be read, a checksum that does not match, XML a parser
rejects, an encoding other than UTF-8 or UTF-16, two names for one part, a relationship that cannot be resolved, a
file too large to read in full. Such a file is refused unless the call passes allow_network=true, and the refusal
says why. Excel 16 opens the files below that hide a call and fetches from them (measured with a listener on
127.0.0.1; the live twins are in test_xl_app_reaches_no_network_unasked.py).

Every file is made up by code. No test here starts Excel.
"""
from __future__ import annotations

import re
import struct
import zipfile

import openpyxl
import pytest

from xlsx_mcp.core import netreach
from xlsx_mcp.core.errors import FormulaRejected, XlMcpError
from xlsx_mcp.core.netreach import file_reach
from xlsx_mcp.ops import comtier

URL = "http://127.0.0.1:9/made-up"
CALL = f'_xlfn.WEBSERVICE("{URL}")'
OFFICE_DOC = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
SHEET = "xl/worksheets/sheet1.xml"


def _items(tmp_path, b1="=1"):
    p = tmp_path / "src.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws["A1"] = "made-up"
    ws["B1"] = b1
    wb.save(p)
    with zipfile.ZipFile(p) as z:
        return [(i.filename, z.read(i.filename)) for i in z.infolist()]


def _write(tmp_path, items, name="book.xlsx"):
    out = tmp_path / name
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in items:
            z.writestr(n, d)
    return str(out)


def _with(tmp_path, part, fn, b1="=1", name="book.xlsx"):
    items = [(n, fn(d.decode("utf-8")) if n == part else d) for n, d in _items(tmp_path, b1)]
    return _write(tmp_path, [(n, d.encode("utf-8") if isinstance(d, str) else d) for n, d in items], name)


def _cells(cells):
    def fn(t):
        t = re.sub(r"^\s*<\?xml[^>]*\?>\s*", "", t)
        return re.sub(r"<sheetData>.*</sheetData>|<sheetData\s*/>", lambda m: "<sheetData>" + cells + "</sheetData>",
                      t, flags=re.S)
    return fn


def _found(reach, fn="WEBSERVICE"):
    return bool(reach) and any(fn in f["functions"] for f in reach["formulas"])


def _unread(reach, words):
    return bool(reach) and any(words in u for u in reach.get("unread", []))


def _patch_crc(path, member):
    """Store a wrong CRC-32 for one member, in its local header and in the central directory."""
    name = member.encode()
    data = bytearray(open(path, "rb").read())
    with zipfile.ZipFile(path) as z:
        info = z.getinfo(member)
    good, bad = struct.pack("<I", info.CRC), struct.pack("<I", info.CRC ^ 0xFFFFFFFF)
    lh = info.header_offset
    assert data[lh + 14:lh + 18] == good
    data[lh + 14:lh + 18] = bad
    i = data.find(b"PK\x01\x02")
    while i >= 0:
        n = struct.unpack("<H", data[i + 28:i + 30])[0]
        if bytes(data[i + 46:i + 46 + n]) == name:
            data[i + 16:i + 20] = bad
            break
        i = data.find(b"PK\x01\x02", i + 4)
    open(path, "wb").write(bytes(data))


# ---- a call is found however the markup or the package hides it ---------------------------------------------------

def test_a_name_split_by_an_instruction_whose_target_starts_with_xml_is_found(tmp_path):
    p = _with(tmp_path, SHEET, _cells(f'<row r="1"><c r="B1"><f>_xlfn.WEB<?xml-x?>SERVICE("{URL}")</f></c></row>'))
    how, reach = file_reach(p)
    assert how == netreach.READ_PACKAGE
    assert reach["formulas"][0] == {"where": "Rates!B1", "functions": ["WEBSERVICE"]}, reach


def test_a_name_split_by_a_child_element_is_found(tmp_path):
    p = _with(tmp_path, SHEET, _cells(f'<row r="1"><c r="B1"><f>_xlfn.WEB<x/>SERVICE("{URL}")</f></c></row>'))
    assert _found(file_reach(p)[1])


def test_a_part_declared_utf7_is_refused_and_its_call_still_named(tmp_path):
    def fn(t):
        t = _cells(f'<row r="1"><c r="B1"><f>_xlfn.+AFc-EBSERVICE("{URL}")</f></c></row>')(t)
        return ('<?xml version="1.0" encoding="UTF-7"?>' + t).encode("ascii")
    how, reach = file_reach(_with(tmp_path, SHEET, fn))
    assert _found(reach), reach
    assert _unread(reach, "declares the encoding UTF-7")


def test_a_commented_out_workbook_relationship_points_the_check_nowhere(tmp_path):
    items = _items(tmp_path, f"={CALL}")
    decoy = f'<!--<Relationship Id="rIdD" Type="{OFFICE_DOC}" Target="xl/decoy.xml"/>-->'
    out = []
    for n, d in items:
        if n == "_rels/.rels":
            d = re.sub(r"(<Relationships\b[^>]*>)", lambda m: m.group(1) + decoy, d.decode(), count=1).encode()
        out.append((n, d))
    out.append(("xl/decoy.xml", dict(items)["xl/workbook.xml"]))
    how, reach = file_reach(_write(tmp_path, out))
    assert reach["formulas"][0] == {"where": "Rates!B1", "functions": ["WEBSERVICE"]}, reach
    assert "unread" not in reach


def test_the_workbook_part_is_found_whatever_the_case_of_its_name(tmp_path):
    out = []
    for n, d in _items(tmp_path, f"={CALL}"):
        if n == "_rels/.rels":
            d = d.replace(b"xl/workbook.xml", b"xl/book.xml")
        elif n == "[Content_Types].xml":
            d = d.replace(b"/xl/workbook.xml", b"/xl/book.xml")
        elif n == "xl/workbook.xml":
            out.append(("xl/Book.xml", d))  # the real workbook part; xl/workbook.xml stays as a decoy
        elif n == "xl/_rels/workbook.xml.rels":
            n = "xl/_rels/Book.xml.rels"
        out.append((n, d))
    how, reach = file_reach(_write(tmp_path, out))
    assert reach["formulas"][0] == {"where": "Rates!B1", "functions": ["WEBSERVICE"]}, reach


def test_a_call_in_a_part_no_relationship_reaches_is_found(tmp_path):
    items = _items(tmp_path) + [("xl/orphan.xml", f"<x><y>{CALL}</y></x>".encode())]
    how, reach = file_reach(_write(tmp_path, items))
    assert {"where": "the part xl/orphan.xml", "functions": ["WEBSERVICE"]} in reach["formulas"], reach


def test_a_call_in_an_attribute_of_an_unknown_element_is_found(tmp_path):
    p = _with(tmp_path, SHEET, lambda t: t.replace("</worksheet>", f'<extLst><ext v="{CALL}"/></extLst></worksheet>'))
    assert _found(file_reach(p)[1])


def test_a_data_connection_in_any_part_is_listed(tmp_path):
    conn = ('<connections xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<connection id="1" name="Made-up feed" type="4"/></connections>')
    items = _items(tmp_path) + [("xl/elsewhere/conn.xml", conn.encode())]
    how, reach = file_reach(_write(tmp_path, items))
    assert reach["connections"] == ["Made-up feed"], reach


# ---- what the check cannot read with certainty is refused ---------------------------------------------------------

def test_a_part_with_a_wrong_checksum_is_refused_and_still_read(tmp_path):
    p = _write(tmp_path, _items(tmp_path, f"={CALL}"))
    _patch_crc(p, SHEET)
    how, reach = file_reach(p)
    assert _unread(reach, f"the part {SHEET}: its stored checksum does not match")
    assert _found(reach), reach


def test_a_wrong_checksum_alone_is_enough_to_refuse(tmp_path):
    p = _write(tmp_path, _items(tmp_path))
    _patch_crc(p, SHEET)
    assert _unread(file_reach(p)[1], "checksum")


def test_a_picture_with_a_wrong_checksum_is_refused_too(tmp_path):
    items = _items(tmp_path) + [("xl/media/image1.png", b"\x89PNG\r\n\x1a\n" + bytes(range(256)) * 4)]
    p = _write(tmp_path, items)
    _patch_crc(p, "xl/media/image1.png")
    assert _unread(file_reach(p)[1], "xl/media/image1.png: its stored checksum")


def test_an_encrypted_member_is_refused(tmp_path):
    p = _write(tmp_path, _items(tmp_path))
    data = bytearray(open(p, "rb").read())
    with zipfile.ZipFile(p) as z:
        info = z.getinfo(SHEET)
    lh = info.header_offset
    data[lh + 6] |= 1  # the "encrypted" flag, local header
    i = data.find(b"PK\x01\x02")
    while i >= 0:
        n = struct.unpack("<H", data[i + 28:i + 30])[0]
        if bytes(data[i + 46:i + 46 + n]) == SHEET.encode():
            data[i + 8] |= 1  # and central directory
            break
        i = data.find(b"PK\x01\x02", i + 4)
    open(p, "wb").write(bytes(data))
    assert _unread(file_reach(p)[1], f"the part {SHEET}")


@pytest.mark.parametrize("encoding", ["ISO-8859-1", "windows-1252", "US-ASCII", "UTF-7"])
def test_a_part_declared_in_another_encoding_is_refused(tmp_path, encoding):
    def fn(t):
        t = re.sub(r"^\s*<\?xml[^>]*\?>\s*", "", t)
        return (f'<?xml version="1.0" encoding="{encoding}"?>' + t).encode("ascii")
    assert _unread(file_reach(_with(tmp_path, SHEET, fn))[1], f"declares the encoding {encoding}")


@pytest.fixture
def no_parser_encoding(monkeypatch):
    """The parser reports no declared encoding at the root's start event (as
    some libxml2 versions do until the whole part is read, and for an
    encoding they do not support): the declaration is all the check has."""
    import types

    from xlsx_mcp.core import xmlsafe

    real = xmlsafe.pull_parser

    class _Root:
        def getroottree(self):
            return types.SimpleNamespace(docinfo=types.SimpleNamespace(encoding=None))

    class _Pull:
        def __init__(self, *a, **kw):
            self.inner = real(*a, **kw)

        def feed(self, data):
            return self.inner.feed(data)

        def close(self):
            return self.inner.close()

        def read_events(self):
            for event, _element in self.inner.read_events():
                yield event, _Root()

    monkeypatch.setattr(xmlsafe, "pull_parser", _Pull)


_BODY = '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData/></worksheet>'


def _scan():
    return netreach._XmlScan(SHEET, netreach._STORED_CALL_RE, kind="sheet", title="Rates")


def _declared(declaration):
    def fn(t):
        return (declaration + re.sub(r"^\s*<\?xml[^>]*\?>\s*", "", t)).encode("ascii")
    return fn


@pytest.mark.parametrize("parser_says", [True, False], ids=["parser-encoding", "no-parser-encoding"])
def test_a_declaration_longer_than_the_bound_is_refused_not_read_as_none(tmp_path, request,
                                                                         parser_says):
    """A valid declaration whose white space runs past the first chunk: the
    encoding it names is past the bound, so the part is refused as unread
    whether or not the parser reports an encoding."""
    if not parser_says:
        request.getfixturevalue("no_parser_encoding")
    decl = '<?xml version="1.0"' + " " * netreach._CHUNK + 'encoding="ISO-8859-1"?>'
    reach = file_reach(_with(tmp_path, SHEET, _declared(decl)))[1]
    assert _unread(reach, "its XML declaration does not end within its first 64 KiB"), reach


def test_a_declaration_split_across_chunks_is_read(no_parser_encoding):
    data = ('<?xml version="1.0" encoding="ISO-8859-1"?>' + _BODY).encode("ascii")
    for cut in ((4,), (4, 9, 30), tuple(range(1, 45))):
        bounds = (0, *cut, len(data))
        chunks = [data[a:b] for a, b in zip(bounds, bounds[1:])]
        doubts = netreach._read_xml(chunks, _scan(), netreach._STORED_CALL_RE)
        assert any("declares the encoding ISO-8859-1" in d for d in doubts), (cut, doubts)


def test_a_declaration_split_across_chunks_of_a_package_part_is_read(tmp_path, monkeypatch,
                                                                     no_parser_encoding):
    monkeypatch.setattr(netreach, "_CHUNK", 16)
    p = _with(tmp_path, SHEET, _declared('<?xml version="1.0"   encoding="windows-1252"?>'))
    assert _unread(file_reach(p)[1], "declares the encoding windows-1252")


@pytest.mark.parametrize("declaration", ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
                                         "<?xml version='1.0'\tencoding = 'utf-8'?>",
                                         '<?xml version="1.0"?>', ""],
                         ids=["utf8", "single-quotes", "no-encoding", "no-declaration"])
@pytest.mark.parametrize("parser_says", [True, False], ids=["parser-encoding", "no-parser-encoding"])
def test_an_ordinary_part_is_read_with_or_without_the_parsers_encoding(tmp_path, request,
                                                                       declaration, parser_says):
    if not parser_says:
        request.getfixturevalue("no_parser_encoding")
    monkeypatch = request.getfixturevalue("monkeypatch")
    monkeypatch.setattr(netreach, "_CHUNK", 16)  # the declaration over several chunks
    p = _with(tmp_path, SHEET, _declared(declaration))
    assert file_reach(p) == (netreach.READ_PACKAGE, None)


def test_a_part_stored_as_utf32_is_refused(tmp_path):
    p = _with(tmp_path, SHEET, lambda t: t.encode("utf-32"))
    assert _unread(file_reach(p)[1], "UTF-32")


def test_utf16_parts_with_and_without_a_mark_are_read(tmp_path):
    for mark in (True, False):
        def fn(t):
            t = re.sub(r"^\s*<\?xml[^>]*\?>\s*", "", t).replace("=1", "")
            t = '<?xml version="1.0" encoding="UTF-16"?>' + t
            return (b"\xff\xfe" if mark else b"") + t.encode("utf-16-le")
        p = _with(tmp_path, SHEET, fn, b1=f"={CALL}", name=f"u16_{mark}.xlsx")
        how, reach = file_reach(p)
        assert reach["formulas"] == [{"where": "Rates!B1", "functions": ["WEBSERVICE"]}], (mark, reach)


@pytest.mark.parametrize("damage", [
    lambda t: t.replace("</worksheet>", ""),  # never closed
    lambda t: t.replace("<sheetData>", "<sheetData>&made-up;"),  # an entity nothing declares
    lambda t: t.replace("<sheetData>", "<sheetData><c r=\"A9\" s=\"<\"/>"),  # "<" in an attribute value
    lambda t: t + "<extra/>",  # a second root
])
def test_xml_the_parser_rejects_is_refused(tmp_path, damage):
    assert _unread(file_reach(_with(tmp_path, SHEET, damage))[1], "not well-formed XML")


def test_a_call_in_xml_the_parser_rejects_is_still_named(tmp_path):
    p = _with(tmp_path, SHEET, lambda t: t.replace("</worksheet>", ""), b1=f"={CALL}")
    reach = file_reach(p)[1]
    assert _found(reach) and _unread(reach, "not well-formed")


def test_two_copies_of_a_part_name_are_refused(tmp_path):
    items = _items(tmp_path)
    p = tmp_path / "twice.xlsx"
    with zipfile.ZipFile(p, "w") as z:
        for n, d in items:
            z.writestr(n, d)
        with pytest.warns(UserWarning):
            z.writestr(SHEET, dict(items)[SHEET])
    with pytest.raises(XlMcpError):
        file_reach(str(p))


def test_names_that_differ_only_by_an_escape_are_refused(tmp_path):
    items = _items(tmp_path) + [("xl/worksheets/sheet%31.xml", dict(_items(tmp_path))[SHEET])]
    assert _unread(file_reach(_write(tmp_path, items))[1], "which name the same part")


def test_a_relationship_to_a_missing_part_is_refused(tmp_path):
    def fn(t):
        return t.replace("</Relationships>", '<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/'
                         'officeDocument/2006/relationships/worksheet" Target="worksheets/gone.xml"/>'
                         "</Relationships>")
    p = _with(tmp_path, "xl/_rels/workbook.xml.rels", fn)
    assert _unread(file_reach(p)[1], "points at worksheets/gone.xml, which is not in the package")


def test_a_relationship_without_an_id_is_refused(tmp_path):
    def fn(t):
        return t.replace("</Relationships>", '<Relationship Type="x" Target="styles.xml"/></Relationships>')
    assert _unread(file_reach(_with(tmp_path, "xl/_rels/workbook.xml.rels", fn))[1], "without a unique Id")


def test_two_workbook_parts_named_at_the_root_are_refused(tmp_path):
    def fn(t):
        return t.replace("</Relationships>", f'<Relationship Id="rIdX" Type="{OFFICE_DOC}" '
                         'Target="xl/workbook.xml"/></Relationships>')
    assert _unread(file_reach(_with(tmp_path, "_rels/.rels", fn))[1], "exactly one workbook part")


def test_a_sheet_whose_relationship_cannot_be_resolved_is_refused(tmp_path):
    p = _with(tmp_path, "xl/workbook.xml", lambda t: t.replace('r:id="rId1"', 'r:id="rId404"'))
    assert _unread(file_reach(p)[1], "the sheet Rates names a relationship that cannot be resolved")


def test_a_package_without_its_content_types_is_refused(tmp_path):
    items = [(n, d) for n, d in _items(tmp_path) if n != "[Content_Types].xml"]
    assert _unread(file_reach(_write(tmp_path, items))[1], "content types")


def test_a_file_too_large_to_read_in_full_is_refused_never_half_read(tmp_path, monkeypatch):
    p = _write(tmp_path, _items(tmp_path))
    monkeypatch.setattr(netreach, "_MAX_READ", 1000)
    assert _unread(file_reach(p)[1], "more than this check reads in full")
    csv = tmp_path / "big.csv"
    csv.write_text("a,b\n" * 1000, encoding="utf-8")
    assert _unread(file_reach(str(csv))[1], "more than this check reads in full")


def test_a_package_whose_zip_structure_cannot_be_read_is_left_to_excel_to_look(tmp_path):
    # only Excel can say what it makes of it: Excel looks with calculation turned off, as for an .xlsb,
    # and a file Excel cannot open is reported as one (com_validate_opens_clean: opens_clean false)
    p = _write(tmp_path, _items(tmp_path, f"={CALL}"))
    data = open(p, "rb").read()
    open(p, "wb").write(data[: len(data) // 2])
    assert file_reach(p) == (netreach.READ_IN_EXCEL, None)
    assert netreach.network_reach(p) is None  # the shared COM open has Excel look


# ---- text files ---------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("raw,words", [
    ("﻿name,value\n".encode("utf-32"), "UTF-32"),
    (b"n\x00a\x00m\x00e\x00,\x00v\x00\n\x00", "zero bytes"),
    (b"\xff\xfe" + "name,value\n".encode("utf-16-le") + b"\x00\xd8", "does not decode"),
])
def test_text_the_check_cannot_decode_as_excel_does_is_refused(tmp_path, raw, words):
    p = tmp_path / "rates.csv"
    p.write_bytes(raw)
    assert _unread(file_reach(str(p))[1], words)


def test_utf16_text_with_a_mark_is_read(tmp_path):
    p = tmp_path / "rates.csv"
    p.write_bytes(b"\xff\xfe" + f'name,value\nmade-up,"=WEBSERVICE(""{URL}"")"\n'.encode("utf-16-le"))
    assert _found(file_reach(str(p))[1])


def test_text_in_the_system_code_page_is_read_byte_for_byte(tmp_path):
    p = tmp_path / "plain.csv"
    p.write_bytes("name,café\nmade-up,3\n".encode("cp1252"))
    assert file_reach(str(p)) == (netreach.READ_TEXT, None)
    p.write_bytes(f'café,"=WEBSERVICE(""{URL}"")"\n'.encode("cp1252"))
    assert _found(file_reach(str(p))[1])


def test_text_that_starts_like_xml_but_is_not_well_formed_is_refused(tmp_path):
    p = tmp_path / "flat.fods"
    p.write_text('<?xml version="1.0"?><root><a>never closed', encoding="utf-8")
    assert _unread(file_reach(str(p))[1], "not well-formed")


# ---- ordinary workbooks are not refused ---------------------------------------------------------------------------

def test_ordinary_text_that_mentions_a_function_is_not_refused(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Notes"
    ws["A1"] = "image (1) and Image(s)"
    ws["A2"] = "translate (later), webservice (old)"
    ws["A3"] = "=SUM(1,2)"
    ws.append(["copy(1)", "happy (yes)"])
    p = tmp_path / "notes.xlsx"
    wb.save(p)
    assert file_reach(str(p)) == (netreach.READ_PACKAGE, None)


def test_an_ordinary_workbook_with_rules_names_tables_and_notes_is_not_refused(tmp_path):
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Units", "Price"])
    for i in range(1, 30):
        ws.append([f"Made-up {i}", i, i * 2.5])
    ws["E1"] = "=SUM(B2:B30)"
    ws["A1"].comment = Comment("A made-up note", "Tester")
    ws.add_table(Table(displayName="Orders", ref="A1:C30"))
    ws.conditional_formatting.add("B2:B30", CellIsRule(operator="greaterThan", formula=["10"]))
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="100")
    ws.add_data_validation(dv)
    dv.add("B2:B30")
    wb.defined_names["Units"] = openpyxl.workbook.defined_name.DefinedName("Units", attr_text="Sales!$B$2:$B$30")
    wb.create_sheet("Second sheet")["A1"] = "=Sales!E1*2"
    p = tmp_path / "ordinary.xlsx"
    wb.save(p)
    assert file_reach(str(p)) == (netreach.READ_PACKAGE, None)


# ---- the refusal ----------------------------------------------------------------------------------------------------

def test_the_refusal_says_why_and_names_the_override(tmp_path):
    p = _write(tmp_path, _items(tmp_path))
    _patch_crc(p, SHEET)
    with pytest.raises(FormulaRejected) as ei:
        comtier._refuse_network_before_excel(p, False, "export_pdf")
    msg = str(ei.value)
    assert "could not read with certainty" in msg and "checksum" in msg and "allow_network=true" in msg
    assert ei.value.detail["network_reach"]["unread"]
    comtier._refuse_network_before_excel(p, True, "export_pdf")  # asked: nothing is refused here
