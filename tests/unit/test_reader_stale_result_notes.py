"""get_cells in its default mode and get_table say when
the cached results they return may be out of date, as read_range does.

After an edit that leaves formula results stale (the edit names them and
flags the workbook to recalculate on open), read_range, query_range, the
exports, the grid view and find_cells carry cached_note. get_cells in its
default mode did not: its formula probe streams the sheet read-only, where a
blank cell is an EmptyCell with no address, and reading that address raised
and was swallowed, so the probe "could not answer" on any row with a gap and
the labels fell back to plain values. get_table had no note at all.
"""
from __future__ import annotations

import re
import zipfile

import openpyxl
from openpyxl.worksheet.table import Table

from xlsx_mcp.ops import cells, gridio, tables


def _patch(src, out, edits):
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            b = zin.read(i.filename)
            for part, fn in edits:
                if i.filename == part:
                    b = fn(b)
            zout.writestr(i, b)


def _no_pending(b):
    return re.sub(rb"<calcPr[^>]*/>", b'<calcPr calcId="191029"/>', b)


def _cache(b, cell, formula, value):
    pat = (rb'<c r="' + cell.encode() + rb'"([^>]*)><f>' + re.escape(formula.encode())
           + rb"</f>(?:<v\s*/>|<v></v>)?</c>")
    return re.sub(pat, rb'<c r="' + cell.encode() + rb'"\1><f>' + formula.encode()
                  + rb"</f><v>" + str(value).encode() + rb"</v></c>", b)


def _book(tmp_path):
    """Sales!A1:C3 is table Tb with a formula column; E2 holds a formula
    with an EMPTY D2 before it (the gap the default-mode probe tripped on).
    Every formula carries a cached result and nothing is owed a recalc."""
    base = str(tmp_path / "base.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Item", "Qty", "Amount"])
    ws.append(["a", 2, "=B2*10"])
    ws.append(["b", 3, "=B3*10"])
    ws["E2"] = "=B2*100"
    ws.add_table(Table(displayName="Tb", ref="A1:C3"))
    wb.save(base)

    def caches(b):
        for cell, f, v in (("C2", "B2*10", 20), ("C3", "B3*10", 30), ("E2", "B2*100", 200)):
            b = _cache(b, cell, f, v)
        return b
    p = str(tmp_path / "book.xlsx")
    _patch(base, p, [("xl/worksheets/sheet1.xml", caches), ("xl/workbook.xml", _no_pending)])
    assert cells.read_range(p, "C2:E2")["values"] == [[20, None, 200]]
    return p


def test_the_default_mode_probe_reads_a_row_with_gaps(tmp_path):
    p = _book(tmp_path)
    mask = gridio.formula_mask_cells(p, "Sales", [(2, 5), (2, 1)])
    assert mask is not None
    assert (2, 5) in mask and (2, 1) not in mask


def test_get_cells_default_mode_says_results_may_be_stale(tmp_path):
    p = _book(tmp_path)
    out = cells.set_cell(p, "B2", 99)
    assert "Sales!E2" in (out.get("cached_values") or {}).get("possibly_stale_cells", [])
    got = cells.get_cells(p, ["E2"])
    assert got["cells"][0][1] == 200  # the pre-edit cached result
    assert got.get("labels", {}).get("cached") == ["E2"]
    assert got.get("cached_note") == gridio.PENDING_NOTE


def test_get_table_says_results_may_be_stale(tmp_path):
    p = _book(tmp_path)
    cells.set_cell(p, "B2", 99)
    got = tables.get_table(p, "Tb")
    assert got["rows"] == [["a", 99, 20], ["b", 3, 30]]
    assert got.get("cached_note") == gridio.PENDING_NOTE
    for mode in ("both", "displayed"):
        assert tables.get_table(p, "Tb", values=mode).get("cached_note") == gridio.PENDING_NOTE, mode


def test_no_note_when_nothing_is_owed_a_recalculation(tmp_path):
    p = _book(tmp_path)
    assert "cached_note" not in cells.get_cells(p, ["E2", "C3"])
    assert "cached_note" not in tables.get_table(p, "Tb")


def test_no_note_when_the_cells_read_hold_no_formula(tmp_path):
    p = _book(tmp_path)
    cells.set_cell(p, "B2", 99)
    assert "cached_note" not in cells.get_cells(p, ["A2", "B3"])
    assert "cached_note" not in tables.get_table(p, "Tb", columns=["Item", "Qty"])
    assert "cached_note" not in tables.get_table(p, "Tb", values="formula")
