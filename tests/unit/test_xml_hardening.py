"""Every XML part is read with one hardened parser, and a workbook whose XML
declares a document type is refused plainly and fast.

A workbook package's XML never carries a document type declaration (the
Open Packaging Conventions forbid one, and Excel never writes one). A DTD is
how a part expands text without limit ("billion laughs", quadratic blowup)
or reads another file (an external entity). Before this, openpyxl's
worksheet and shared-string reader (the standard library's expat) expanded a
small internal entity into the cell text without a word, its other readers
dropped an entity's text from an element but expanded one inside an
attribute, and a billion-laughs part came back as the parser's own message,
under BAD_PARAMS or INTERNAL_ERROR.

Every workbook here is made up in code (hostile_xlsx.py). External targets
are a made-up file in the test's folder or a listener on 127.0.0.1, and a
socket guard refuses any other connection, so nothing here can reach an
outside host. Excel cannot start in this file: the COM tier's only way to
start it is replaced by a failure."""

from __future__ import annotations

import ast
import http.server
import io
import json
import socket
import sys
import threading
import time
import zipfile
from pathlib import Path

import pytest

import hostile_xlsx as hx

SRC = Path(__file__).resolve().parents[2] / "src" / "xlsx_mcp"
SCRIPTS = Path(__file__).resolve().parents[2] / "scripts"

#: What every refusal says, whichever entry point met the part.
SENTENCE = "document type declaration"


# ------------------------------------------------------------------ fixtures

class _Listener:
    def __init__(self):
        self.hits: list[str] = []
        listener = self

        class Handler(http.server.BaseHTTPRequestHandler):
            def do_GET(self):  # noqa: N802
                listener.hits.append(self.path)
                body = (f'<!ENTITY leak "{hx.MARKER}">' if self.path.endswith(".dtd")
                        else hx.MARKER).encode()
                self.send_response(200)
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)

            def log_message(self, *args):
                pass

        self.server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.base = f"http://127.0.0.1:{self.server.server_address[1]}"
        threading.Thread(target=self.server.serve_forever, daemon=True).start()


@pytest.fixture
def net(monkeypatch):
    """A listener on 127.0.0.1 that records every request an external entity
    could make, and a socket guard that refuses (and records) any connection
    to anything but loopback."""
    outside: list[str] = []
    real_connect = socket.socket.connect

    def guarded(sock, addr):
        host = addr[0] if isinstance(addr, tuple) else str(addr)
        if host not in ("127.0.0.1", "::1", "localhost"):
            outside.append(repr(addr))
            raise OSError("the test's socket guard allows loopback only")
        return real_connect(sock, addr)

    monkeypatch.setattr(socket.socket, "connect", guarded)
    lis = _Listener()
    lis.outside = outside
    yield lis
    lis.server.shutdown()


@pytest.fixture
def targets(tmp_path):
    folder = tmp_path / "targets"
    folder.mkdir()
    return hx.made_up_targets(folder)


@pytest.fixture
def no_excel(monkeypatch):
    """Excel cannot start: the COM tier's spawn fails the test instead."""
    from xlsx_mcp.com import instances

    def refuse(*_a, **_k):
        raise AssertionError("Excel must not start in this test")

    monkeypatch.setattr(instances.ExcelInstanceManager, "_new_app", refuse)
    try:
        import win32com.client as w32
    except ImportError:
        return
    monkeypatch.setattr(w32, "DispatchEx", refuse)
    monkeypatch.setattr(w32, "Dispatch", refuse)


def _hostile_file(tmp_path, kind, attack, targets, net, encoding="utf-8", name="attack.xlsx"):
    part, _root, _slot, _st, chart = hx.KINDS[kind]
    target, dtd = targets
    data = hx.hostile(kind, attack, target_url=target.as_uri(), dtd_url=dtd.as_uri(),
                      http_base=net.base, encoding=encoding)
    return hx.make_xlsx(tmp_path / name, parts={part: data}, chart=chart)


def _assert_refusal_text(msg: str, part: str | None = None):
    assert SENTENCE in msg, msg
    assert "Nothing was changed" in msg, msg
    if part is not None:
        assert part in msg, msg
    # Never the parser's own words.
    for raw in ("amplification", "line 1, column", "Entity '", "PEReference",
                "undefined entity", "not well-formed"):
        assert raw not in msg, msg
    assert hx.MARKER not in msg and hx.EXPANSION not in msg


# ------------------------------------------------- the in-process entry points

@pytest.mark.parametrize("kind", sorted(hx.KINDS))
@pytest.mark.parametrize("attack", hx.ATTACKS)
def test_opening_a_workbook_whose_part_declares_a_dtd_refuses_fast_and_plainly(
        attack, kind, tmp_path, targets, net):
    from xlsx_mcp.core.errors import WorkbookCorrupt
    from xlsx_mcp.core.package import WorkbookPackage

    f = _hostile_file(tmp_path, kind, attack, targets, net)
    t0 = time.perf_counter()
    with pytest.raises(WorkbookCorrupt) as err:
        pkg = WorkbookPackage.open(str(f))
        pkg.workbook  # noqa: B018 (the model load is the read)
    assert time.perf_counter() - t0 < 2.0
    _assert_refusal_text(str(err.value), hx.KINDS[kind][0])
    assert net.hits == [] and net.outside == []


@pytest.mark.parametrize("kind", hx.READ_KINDS)
@pytest.mark.parametrize("attack", hx.ATTACKS)
def test_the_read_model_refuses_before_openpyxl_reads_a_part(
        attack, kind, tmp_path, targets, net):
    from xlsx_mcp.core.errors import WorkbookCorrupt
    from xlsx_mcp.ops import gridio

    f = _hostile_file(tmp_path, kind, attack, targets, net)
    t0 = time.perf_counter()
    with pytest.raises(WorkbookCorrupt) as err:
        gridio.open_wb(str(f))
    assert time.perf_counter() - t0 < 2.0
    _assert_refusal_text(str(err.value), hx.KINDS[kind][0])
    assert net.hits == [] and net.outside == []


@pytest.mark.parametrize("kind", ["workbook", "sheet", "sharedStrings", "styles", "comments",
                                  "content_types"])
@pytest.mark.parametrize("attack", ["billion_laughs", "small_internal_entity", "dtd_no_entities"])
@pytest.mark.parametrize("encoding", ["utf-16-le", "utf-16-be"])
def test_a_utf16_part_that_declares_a_dtd_is_refused(
        encoding, attack, kind, tmp_path, targets, net):
    from xlsx_mcp.core.errors import WorkbookCorrupt
    from xlsx_mcp.ops import gridio

    f = _hostile_file(tmp_path, kind, attack, targets, net, encoding=encoding)
    with pytest.raises(WorkbookCorrupt) as err:
        gridio.open_wb(str(f))
    _assert_refusal_text(str(err.value), hx.KINDS[kind][0])


def test_a_declaration_after_a_long_prolog_is_still_refused(tmp_path, targets, net):
    from xlsx_mcp.core.errors import WorkbookCorrupt
    from xlsx_mcp.core.package import WorkbookPackage

    data = hx.hostile("sheet", "small_internal_entity").decode("utf-8")
    pad = "<!-- " + "made-up padding " * 20000 + "--><?ks4 made-up?>"
    data = data.replace("<!DOCTYPE", pad + "<!DOCTYPE", 1)
    f = hx.make_xlsx(tmp_path / "pad.xlsx", parts={"xl/worksheets/sheet1.xml": data})
    with pytest.raises(WorkbookCorrupt, match=SENTENCE):
        WorkbookPackage.open(str(f))


def test_a_declaration_hidden_by_a_declared_utf7_encoding_is_refused(tmp_path):
    """In UTF-7, "<!DOCTYPE" need not be the ASCII bytes a byte scan looks
    for, and libxml2 decodes it. The check reads the prolog as the parser
    does. (Before it, openpyxl named the sheet with the expanded entity.)"""
    from xlsx_mcp.core.errors import WorkbookCorrupt
    from xlsx_mcp.ops import gridio

    xml = hx.base_parts()["xl/workbook.xml"]
    head, rest = xml.split("?>", 1)
    head = head.replace('encoding="UTF-8"', 'encoding="UTF-7"')
    # "<!DOCTYPE workbook [<!ENTITY e "made-up words">]>" with every
    # delimiter in UTF-7's base64 form.
    hidden = ("+ADwAIQ-DOCTYPE workbook +AFsAPA-+ACE-ENTITY e +ACI-made-up words"
              "+ACIAPgBd-+AD4-")
    rest = rest.replace('name="Sheet1"', 'name="&e;"', 1)
    data = (head + "?>").encode("ascii") + hidden.encode("ascii") + rest.encode("utf-7")
    f = hx.make_xlsx(tmp_path / "utf7.xlsx", parts={"xl/workbook.xml": data})
    with pytest.raises(WorkbookCorrupt, match=SENTENCE):
        gridio.open_wb(str(f))


def test_doctype_text_in_a_comment_is_not_a_declaration(tmp_path):
    import family_harness as fh

    parts = hx.base_parts()
    sheet = parts["xl/worksheets/sheet1.xml"].replace(
        "?>", "?><!-- made-up note: <!DOCTYPE is only text here -->", 1)
    sst = parts["xl/sharedStrings.xml"].replace(
        "<si>", "<!-- made-up: <!DOCTYPE inside the body --><si>", 1)
    f = hx.make_xlsx(tmp_path / "c.xlsx", parts={"xl/worksheets/sheet1.xml": sheet,
                                                 "xl/sharedStrings.xml": sst})
    payload, is_error, _ = fh.start().call("xl_read", {"op": "read_range", "args": {
        "path": str(f), "location": "A1:B2"}})
    assert not is_error, payload
    assert payload["values"] == [["Made-up label", 42], ["Made-up inline text", 7.5]]


# ----------------------------------------------------- the public tool surface

@pytest.mark.parametrize("attack", hx.ATTACKS)
@pytest.mark.parametrize("kind", ["sheet", "sharedStrings", "workbook", "styles", "comments"])
def test_the_read_surface_refuses_with_unsupported_content(
        kind, attack, tmp_path, targets, net):
    import family_harness as fh

    f = _hostile_file(tmp_path, kind, attack, targets, net)
    t0 = time.perf_counter()
    payload, is_error, _ = fh.start().call(
        "xl_read", {"op": "read_range", "args": {"path": str(f), "location": "A1:B2",
                                                 "values": "both"}})
    assert time.perf_counter() - t0 < 10.0
    assert is_error and payload["error"]["code"] == "UNSUPPORTED_CONTENT", payload
    _assert_refusal_text(payload["error"]["message"], hx.KINDS[kind][0])
    assert hx.MARKER not in json.dumps(payload) and hx.EXPANSION not in json.dumps(payload)
    assert net.hits == [] and net.outside == []


#: Ops that read no workbook part: they make, copy or list files, or answer
#: without a file (the update ops take no path at all).
_NO_PART_READ = {"create_workbook", "copy_workbook", "manage_backups", "get_server_info",
                 "get_workflows", "com_status", "update_now", "update_notices"}


def _examples() -> dict[str, tuple[str, dict]]:
    import family_harness as fh
    from xlsx_mcp import registry

    h = fh.start()
    out = {}
    for name in sorted(registry.public_ops()):
        if name in _NO_PART_READ:
            continue
        pl, _err, _ = h.call("xl_help", {"op": name})
        out[name] = (pl["tool"], pl["example"]["args"])
    return out


def test_every_op_that_reads_a_workbook_refuses_one_whose_part_declares_a_dtd(
        tmp_path, targets, net, no_excel):
    """The same refusal from every op that reads the file, the COM tier
    included (before Excel would start), and the file is unchanged."""
    import family_harness as fh

    f = _hostile_file(tmp_path, "sharedStrings", "small_internal_entity", targets, net)
    before = f.read_bytes()
    h = fh.start()
    wrong = {}
    for op, (tool, example) in _examples().items():
        args = dict(example, path=str(f))
        for key in ("output",):
            if key in args:
                args[key] = str(tmp_path / ("out_" + op + Path(args[key]).suffix))
        payload, is_error, _ = h.call(tool, {"op": op, "args": args})
        err = payload.get("error") if isinstance(payload, dict) else None
        if not (is_error and err and err.get("code") == "UNSUPPORTED_CONTENT"
                and SENTENCE in err.get("message", "")):
            wrong[op] = payload
        elif hx.EXPANSION in json.dumps(payload):
            wrong[op] = payload
    assert wrong == {}
    assert f.read_bytes() == before
    assert net.hits == [] and net.outside == []


# -------------------------------------------------- openpyxl never sees a DTD

def _openpyxl_parse_spies(monkeypatch) -> list[tuple[str, bool]]:
    """Wrap every parse entry point openpyxl's readers hold, recording for
    each input whether it declares a DTD. The wrappers then call the real
    parser, so a load behaves as before."""
    import importlib

    import openpyxl.xml.functions as functions
    from xlsx_mcp.core import xmlsafe

    for mod in ("openpyxl.reader.excel", "openpyxl.reader.workbook", "openpyxl.reader.strings",
                "openpyxl.reader.drawings", "openpyxl.worksheet._reader",
                "openpyxl.worksheet._read_only", "openpyxl.packaging.relationship",
                "openpyxl.packaging.manifest", "openpyxl.styles.stylesheet",
                "openpyxl.workbook.external_link.external", "openpyxl.chart.reader",
                "openpyxl.comments.comment_sheet", "openpyxl.pivot.table", "openpyxl.pivot.cache"):
        try:
            importlib.import_module(mod)
        except ImportError:
            pass
    seen: list[tuple[str, bool]] = []
    real_fromstring, real_iterparse = functions.fromstring, functions.iterparse

    def fromstring(xml, *a, **k):
        data = xml.encode("utf-8") if isinstance(xml, str) else bytes(xml)
        seen.append(("fromstring", xmlsafe.declares_dtd(data)))
        return real_fromstring(xml, *a, **k)

    def iterparse(source, *a, **k):
        data = source.read() if hasattr(source, "read") else Path(source).read_bytes()
        seen.append(("iterparse", xmlsafe.declares_dtd(data)))
        return real_iterparse(io.BytesIO(data), *a, **k)

    for name, mod in list(sys.modules.items()):
        if not name.startswith("openpyxl") or mod is None:
            continue
        if getattr(mod, "fromstring", None) is real_fromstring:
            monkeypatch.setattr(mod, "fromstring", fromstring)
        if getattr(mod, "iterparse", None) is real_iterparse:
            monkeypatch.setattr(mod, "iterparse", iterparse)
    return seen


def test_the_spies_sit_in_openpyxls_read_path(tmp_path, monkeypatch):
    from xlsx_mcp.ops import gridio

    seen = _openpyxl_parse_spies(monkeypatch)
    gridio.open_wb(str(hx.make_xlsx(tmp_path / "plain.xlsx", chart=True)))
    kinds = {k for k, _ in seen}
    assert kinds == {"fromstring", "iterparse"} and len(seen) >= 8, seen
    assert not any(dtd for _, dtd in seen)


@pytest.mark.parametrize("kind", hx.READ_KINDS)
def test_openpyxl_never_parses_a_part_that_declares_a_dtd(kind, tmp_path, targets, net,
                                                          monkeypatch):
    from xlsx_mcp.core.errors import WorkbookCorrupt
    from xlsx_mcp.ops import gridio

    seen = _openpyxl_parse_spies(monkeypatch)
    f = _hostile_file(tmp_path, kind, "small_internal_entity", targets, net)
    with pytest.raises(WorkbookCorrupt, match=SENTENCE):
        gridio.open_wb(str(f))
    assert not any(dtd for _, dtd in seen), seen


def _oddly_named_workbook(tmp_path) -> Path:
    """The workbook part under a name no extension check would take for XML:
    openpyxl finds it through [Content_Types].xml and parses it anyway."""
    parts = hx.base_parts()
    wb = parts.pop("xl/workbook.xml").replace(
        "?>", '?><!DOCTYPE workbook [<!ENTITY e "made-up words">]>', 1).replace(
        'name="Sheet1"', 'name="&e;"', 1)
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        "/xl/workbook.xml", "/xl/workbook.dat")
    parts["_rels/.rels"] = parts["_rels/.rels"].replace('"xl/workbook.xml"', '"xl/workbook.dat"')
    parts["xl/workbook.dat"] = wb
    path = tmp_path / "odd.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return path


@pytest.mark.parametrize("read_only", [False, True])
def test_the_loader_checks_each_part_it_reads_whatever_its_name(read_only, tmp_path, monkeypatch):
    from xlsx_mcp.core import xmlsafe
    from xlsx_mcp.core.errors import WorkbookCorrupt

    seen = _openpyxl_parse_spies(monkeypatch)
    f = _oddly_named_workbook(tmp_path)
    with pytest.raises(WorkbookCorrupt, match="xl/workbook.dat"):
        xmlsafe.load_workbook(str(f), read_only=read_only)
    assert not any(dtd for _, dtd in seen), seen


def _oddly_named_complete_workbook(tmp_path) -> Path:
    """_oddly_named_workbook with the workbook's relationships part renamed
    to match (xl/_rels/workbook.dat.rels), so a plain openpyxl load reads
    the whole package and the entity becomes the sheet's name."""
    odd = _oddly_named_workbook(tmp_path)
    with zipfile.ZipFile(odd) as z:
        members = [("xl/_rels/workbook.dat.rels" if n == "xl/_rels/workbook.xml.rels" else n,
                    z.read(n)) for n in z.namelist()]
    return hx.make_package(tmp_path / "odd-complete.xlsx", members)


# ------------------------------- recalculate has no formulas engine (2.0.0)
# 2.0.0 removed engine='formulas' (the `formulas` library). Explicit
# 'formulas' and 'auto' without Excel both refuse, and nothing reaches the
# library, whatever other workbooks the file names. Each workbook below is
# made up in code; the UNC host is a made-up .invalid name, and the
# file-system calls that would resolve it are intercepted.

REMOVED = (
    "The formulas engine was removed in 2.0.0. Recalculation needs Windows with "
    "Microsoft Excel: use engine 'com' or 'auto'. Reads still return each formula's "
    "cached value, labelled cached.")
NO_EXCEL = (
    "Recalculation needs Windows with Microsoft Excel, and Excel is not available "
    "here. Reads still return each formula's cached value, labelled cached.")
UNC_HOST = "ks4-nohost.invalid"

_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_ODR = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _main_workbook(path: Path, formula: str, link_target: str | None = None) -> Path:
    """A clean workbook whose A1 names another workbook: in the formula, or
    through an externalLink part whose target is link_target."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"], ws["B1"], ws["C1"] = formula, 5, "=B1*2"
    wb.save(path)
    if link_target is None:
        return path
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    items["xl/externalLinks/externalLink1.xml"] = (
        f'<?xml version="1.0"?><externalLink xmlns="{_MAIN}" xmlns:r="{_ODR}">'
        '<externalBook r:id="rId1"><sheetNames><sheetName val="Prices"/></sheetNames>'
        '<sheetDataSet><sheetData sheetId="0"><row r="1"><cell r="A1" t="n"><v>1</v></cell>'
        '</row></sheetData></sheetDataSet></externalBook></externalLink>').encode()
    items["xl/externalLinks/_rels/externalLink1.xml.rels"] = (
        '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/'
        f'package/2006/relationships"><Relationship Id="rId1" Type="{_ODR}/externalLinkPath" '
        f'Target="{link_target}" TargetMode="External"/></Relationships>').encode()
    wbx = items["xl/workbook.xml"].decode().replace(
        "</sheets>", f'</sheets><externalReferences><externalReference xmlns:r="{_ODR}" '
        'r:id="rIdExt1"/></externalReferences>', 1)
    items["xl/workbook.xml"] = wbx.encode()
    items["xl/_rels/workbook.xml.rels"] = items["xl/_rels/workbook.xml.rels"].decode().replace(
        "</Relationships>", f'<Relationship Id="rIdExt1" Type="{_ODR}/externalLink" '
        'Target="externalLinks/externalLink1.xml"/></Relationships>').encode()
    items["[Content_Types].xml"] = items["[Content_Types].xml"].decode().replace(
        "</Types>", '<Override PartName="/xl/externalLinks/externalLink1.xml" ContentType='
        '"application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/>'
        '</Types>').encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in items.items():
            z.writestr(n, data)
    return path


def _workbook(kind: str, tmp_path: Path) -> Path:
    if kind == "plain":
        return hx.make_xlsx(tmp_path / "plain.xlsx")
    if kind == "odd-complete":
        return _oddly_named_complete_workbook(tmp_path)
    if kind == "linked-to-odd-complete":
        _oddly_named_complete_workbook(tmp_path)  # beside it, as the formula names it
        return _main_workbook(tmp_path / "main.xlsx", "='[odd-complete.xlsx]made-up words'!B1+1")
    assert kind == "unc-link"
    return _main_workbook(tmp_path / "main.xlsx", "=[1]Prices!A1*3",
                          link_target=f"\\\\{UNC_HOST}\\share\\other.xlsx")


@pytest.fixture(params=["imported", "importable"])
def formulas_tripwire(request, monkeypatch):
    """A stand-in for the `formulas` library that records and refuses any
    use: as a module already imported (every attribute, the __spec__ an
    import statement reads included), or as one an import would find."""
    import importlib.abc
    import types

    touched: list[tuple[str, str]] = []

    class _Tripwire(types.ModuleType):
        def __getattribute__(self, name):
            touched.append(("attribute", name))
            raise AssertionError(f"formulas.{name} was reached")

    class _Finder(importlib.abc.MetaPathFinder):
        def find_spec(self, name, path=None, target=None):
            if name == "formulas" or name.startswith("formulas."):
                touched.append(("import", name))
                raise AssertionError(f"import {name} was attempted")
            return None

    for name in [n for n in sys.modules if n == "formulas" or n.startswith("formulas.")]:
        monkeypatch.delitem(sys.modules, name)
    if request.param == "imported":
        monkeypatch.setitem(sys.modules, "formulas", _Tripwire("formulas"))
    else:
        monkeypatch.setattr(sys, "meta_path", [_Finder(), *sys.meta_path])
    return touched


@pytest.fixture
def excel_unavailable(monkeypatch):
    """No Excel: COM answers unavailable (each probe counted), and every
    route that would start Excel fails the test."""
    from xlsx_mcp.ops import comtier

    probes: list[int] = []

    def refuse(*_a, **_k):
        raise AssertionError("Excel must not start in this test")

    def unavailable(*_a, **_k):
        probes.append(1)
        return False, "no Excel in this test"

    monkeypatch.setattr(comtier._session, "com_available", unavailable)
    monkeypatch.setattr(comtier._session, "run_com", refuse)
    monkeypatch.setattr(comtier, "_run_mutation", refuse)
    return probes


def _file_and_path_spies(monkeypatch) -> tuple[list[str], list[tuple[str, str]]]:
    """Record every workbook a ZIP open, openpyxl.load_workbook or openpyxl's
    reader is handed, and every path given to the calls ntpath.realpath
    makes. A UNC path is answered "not found" without any I/O; other paths
    go to the real call."""
    import ntpath
    import os

    import openpyxl
    import openpyxl.reader.excel as reader_excel

    opened: list[str] = []
    resolved: list[tuple[str, str]] = []

    def name_of(f) -> str:
        return os.fspath(f) if isinstance(f, (str, os.PathLike)) else f"<{type(f).__name__}>"

    real_zip = zipfile.ZipFile.__init__

    def zip_init(self, file, *a, **k):
        opened.append(name_of(file))
        return real_zip(self, file, *a, **k)

    real_load = openpyxl.load_workbook

    def load_workbook(filename, *a, **k):
        opened.append(name_of(filename))
        return real_load(filename, *a, **k)

    real_reader = reader_excel.ExcelReader.__init__

    def reader_init(self, fn, *a, **k):
        opened.append(name_of(fn))
        return real_reader(self, fn, *a, **k)

    monkeypatch.setattr(zipfile.ZipFile, "__init__", zip_init)
    monkeypatch.setattr(openpyxl, "load_workbook", load_workbook)
    monkeypatch.setattr(reader_excel, "load_workbook", load_workbook)
    monkeypatch.setattr(reader_excel.ExcelReader, "__init__", reader_init)

    def intercept(name, real):
        def call(p, *a, **k):
            text = os.fspath(p) if isinstance(p, (str, os.PathLike)) else str(p)
            resolved.append((name, text))
            if text.startswith("\\\\") and not text.startswith("\\\\?\\") \
                    or UNC_HOST in text or text.upper().startswith("\\\\?\\UNC\\"):
                err = OSError(2, "intercepted by the test: no such file")
                err.winerror = 2
                raise err
            return real(p, *a, **k)
        return call

    for name in ("_getfinalpathname", "_nt_readlink", "_findfirstfile"):
        real = getattr(ntpath, name, None)
        if real is not None:
            monkeypatch.setattr(ntpath, name, intercept(name, real))
    return opened, resolved


def _is_unc(text: str) -> bool:
    return (text.startswith("\\\\") and not text.startswith("\\\\?\\")) \
        or text.upper().startswith("\\\\?\\UNC\\") or UNC_HOST in text


@pytest.mark.parametrize("kind", ["plain", "odd-complete", "linked-to-odd-complete", "unc-link"])
@pytest.mark.parametrize("engine", ["formulas", "auto"])
def test_recalculate_never_reaches_the_formulas_library(
        engine, kind, tmp_path, formulas_tripwire, excel_unavailable, monkeypatch):
    """Explicit engine='formulas', and 'auto' with Excel unavailable: the
    call refuses CALC_UNAVAILABLE, the library is never imported or
    touched, no other workbook is opened, no UNC path is resolved, and no
    file changes. engine='formulas' is refused before the file is read and
    before COM is probed."""
    import os

    from xlsx_mcp.core.errors import CalcUnavailable
    from xlsx_mcp.ops import comtier

    f = _workbook(kind, tmp_path)
    before = {p.name: p.read_bytes() for p in tmp_path.iterdir()}
    seen = _openpyxl_parse_spies(monkeypatch)
    opened, resolved = _file_and_path_spies(monkeypatch)
    with pytest.raises(CalcUnavailable) as ei:
        comtier.recalculate(str(f), engine=engine)
    assert str(ei.value) == (REMOVED if engine == "formulas" else NO_EXCEL)
    assert formulas_tripwire == []
    assert [p for _, p in resolved if _is_unc(p)] == []
    assert not any(dtd for _, dtd in seen), seen
    named = os.path.normcase(os.path.abspath(f))
    assert {os.path.normcase(os.path.abspath(p)) for p in opened if not p.startswith("<")} <= {named}, opened
    if engine == "formulas":
        assert opened == [] and excel_unavailable == [], (opened, excel_unavailable)
    assert {p.name: p.read_bytes() for p in tmp_path.iterdir()} == before


def test_the_name_based_check_alone_does_not_see_the_odd_workbook(tmp_path):
    """The control for the odd-complete case above: the name-based package
    check passes it, and a plain openpyxl load expands its entity, so the
    refusal above is the engine's, not the package check's."""
    import openpyxl

    from xlsx_mcp.core import xmlsafe

    f = _oddly_named_complete_workbook(tmp_path)
    xmlsafe.check_file(str(f))
    plain = openpyxl.load_workbook(str(f))
    assert plain.sheetnames == ["made-up words"]
    plain.close()


def test_a_read_only_load_checks_a_sheet_before_openpyxl_reads_it(tmp_path, targets, net,
                                                                  monkeypatch):
    """The streaming (read-only) load reads a sheet through the same check,
    whether openpyxl opens it at load or at the first row."""
    import contextlib

    from xlsx_mcp.core import xmlsafe
    from xlsx_mcp.core.errors import WorkbookCorrupt

    seen = _openpyxl_parse_spies(monkeypatch)
    f = _hostile_file(tmp_path, "sheet", "small_internal_entity", targets, net)
    with pytest.raises(WorkbookCorrupt, match=SENTENCE):
        wb = xmlsafe.load_workbook(str(f), read_only=True)
        with contextlib.closing(wb):
            list(wb["Sheet1"].iter_rows(values_only=True))
    assert not any(dtd for _, dtd in seen), seen


# ------------------------------------------------------------- the parser

def test_the_parser_never_expands_or_loads(targets, net):
    from lxml import etree

    from xlsx_mcp.core import xmlsafe

    p = xmlsafe.parser()
    for opt in ("resolve_entities", "no_network", "load_dtd", "dtd_validation",
                "attribute_defaults"):
        with pytest.raises(ValueError):
            xmlsafe.parser(**{opt: True})
    calls = []

    class Spy(etree.Resolver):
        def resolve(self, url, pubid, context):
            calls.append(url)

    p.resolvers.add(Spy())
    target, dtd = targets
    for attack in ("external_entity_file", "external_entity_http", "parameter_entity_file",
                   "parameter_entity_http"):
        xml = hx.hostile("sheet", attack, target_url=target.as_uri(), dtd_url=dtd.as_uri(),
                         http_base=net.base)
        # Below the refusal: the raw parser itself reads nothing.
        root = etree.fromstring(xml, parser=p)
        assert hx.MARKER not in "".join(root.itertext())
    assert calls == [] and net.hits == []


@pytest.mark.parametrize("attack", hx.ATTACKS)
def test_fromstring_and_walk_refuse_every_attack_in_memory(attack, targets, net):
    from xlsx_mcp.core import xmlsafe
    from xlsx_mcp.core.errors import WorkbookCorrupt

    target, dtd = targets
    data = hx.hostile("sheet", attack, target_url=target.as_uri(), dtd_url=dtd.as_uri(),
                      http_base=net.base)
    with pytest.raises(WorkbookCorrupt, match=SENTENCE):
        xmlsafe.fromstring(data, doc="made-up.xlsx", part="xl/worksheets/sheet1.xml")
    with pytest.raises(WorkbookCorrupt, match=SENTENCE):
        xmlsafe.walk(data, doc="made-up.xlsx", part="xl/worksheets/sheet1.xml")
    assert net.hits == []


def test_a_part_that_is_not_well_formed_raises_the_parse_error_it_did_before():
    from xml.etree.ElementTree import ParseError

    from xlsx_mcp.core import xmlsafe

    with pytest.raises(ParseError):
        xmlsafe.fromstring(b"<a><b></a>")
    with pytest.raises(ParseError):
        xmlsafe.walk(b"<a><b></a>")


def test_comments_and_instructions_are_not_part_of_the_tree():
    """The parse sites were written against ElementTree, which leaves XML
    comments and processing instructions out; so does the hardened parse."""
    from xlsx_mcp.core import xmlsafe

    root = xmlsafe.fromstring(b"<a><!-- made-up --><?ks4 made-up?><b/><c/></a>")
    assert [el.tag for el in root] == ["b", "c"]


# ---------------------------------------------- no raw parse outside the module

#: The modules whose import means an XML parser.
_XML_MODULES = ("xml.etree", "xml.dom", "xml.sax", "xml.parsers", "pyexpat", "_elementtree",
                "lxml", "defusedxml")
#: Calls that parse XML whatever object they hang off.
_XML_CALLS = {"fromstring", "fromstringlist", "XML", "XMLID", "iterparse", "XMLParser",
              "XMLPullParser", "parseString", "ParserCreate", "make_parser"}
#: openpyxl's way into a package: only xmlsafe.load_workbook calls it.
_LOADER_CALLS = {"load_workbook", "ExcelReader"}


def _dotted(node) -> str:
    parts = []
    while isinstance(node, ast.Attribute):
        parts.append(node.attr)
        node = node.value
    if isinstance(node, ast.Name):
        parts.append(node.id)
    return ".".join(reversed(parts))


def _offenders(root: Path, *, loaders: bool) -> list[str]:
    out = []
    for py in sorted(root.rglob("*.py")):
        if py.name == "xmlsafe.py" and py.parent.name == "core":
            continue
        tree = ast.parse(py.read_text(encoding="utf-8"))
        from_xmlsafe: set[str] = set()
        where = f"{py.relative_to(root.parent)}"
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                for alias in node.names:
                    if alias.name.startswith(_XML_MODULES) or alias.name == "xml":
                        out.append(f"{where}:{node.lineno} import {alias.name}")
                    if loaders and alias.name.startswith("openpyxl.reader"):
                        out.append(f"{where}:{node.lineno} import {alias.name}")
            elif isinstance(node, ast.ImportFrom):
                mod = node.module or ""
                names = {a.name for a in node.names}
                if mod.endswith("xmlsafe"):
                    from_xmlsafe |= {a.asname or a.name for a in node.names}
                    continue
                if mod.startswith(_XML_MODULES) or (mod == "xml" and names - {"__name__"}):
                    out.append(f"{where}:{node.lineno} from {mod} import {', '.join(sorted(names))}")
                if loaders and (mod.startswith("openpyxl.reader") or (
                        mod.startswith("openpyxl") and names & (_LOADER_CALLS | _XML_CALLS))):
                    out.append(f"{where}:{node.lineno} from {mod} import {', '.join(sorted(names))}")
        for node in ast.walk(tree):
            if not isinstance(node, ast.Call):
                continue
            func = node.func
            watched = _XML_CALLS | (_LOADER_CALLS if loaders else set())
            if isinstance(func, ast.Attribute) and func.attr in watched:
                receiver = _dotted(func.value)
                if receiver not in ("xmlsafe", "_xmlsafe"):
                    out.append(f"{where}:{node.lineno} {receiver}.{func.attr}")
            elif isinstance(func, ast.Name) and func.id in watched and func.id not in from_xmlsafe:
                out.append(f"{where}:{node.lineno} {func.id}")
    return out


def test_no_module_parses_xml_outside_the_hardened_parser():
    """src may parse XML, and build an openpyxl model from a file, only
    through core/xmlsafe.py. scripts/ (development tools, not shipped) may
    not import a parser either; their openpyxl loads read files the scripts
    or the fixture corpus made, so the loader rule covers src."""
    offenders = _offenders(SRC, loaders=True) + _offenders(SCRIPTS, loaders=False)
    assert offenders == []


def test_the_scan_catches_what_it_is_for(tmp_path):
    """The AST rule itself, on made-up modules: every raw form is caught."""
    pkg = tmp_path / "pkg"
    pkg.mkdir()
    (pkg / "a.py").write_text(
        "import xml.etree.ElementTree as ET\nET.fromstring(b'<a/>')\n", encoding="utf-8")
    (pkg / "b.py").write_text(
        "from lxml import etree\netree.XMLParser()\n", encoding="utf-8")
    (pkg / "c.py").write_text(
        "import openpyxl\nopenpyxl.load_workbook('x.xlsx')\n", encoding="utf-8")
    (pkg / "d.py").write_text(
        "from openpyxl import load_workbook\nload_workbook('x.xlsx')\n", encoding="utf-8")
    (pkg / "e.py").write_text(
        "import xml.parsers.expat\nxml.parsers.expat.ParserCreate()\n", encoding="utf-8")
    (pkg / "ok.py").write_text(
        "from .core import xmlsafe as _xmlsafe\n_xmlsafe.fromstring(b'<a/>')\n"
        "_xmlsafe.load_workbook('x.xlsx')\n", encoding="utf-8")
    found = _offenders(pkg, loaders=True)
    for name in ("a.py", "b.py", "c.py", "d.py", "e.py"):
        assert any(name in line for line in found), (name, found)
    assert not any("ok.py" in line for line in found), found


# ---------------------------------------------------------------- controls

def test_a_normal_workbook_reads_and_edits_as_before(tmp_path):
    import family_harness as fh

    f = hx.make_xlsx(tmp_path / "normal.xlsx", chart=True)
    h = fh.start()
    payload, is_error, _ = h.call("xl_read", {"op": "read_range", "args": {
        "path": str(f), "location": "A1:B2"}})
    assert not is_error and payload["values"] == [["Made-up label", 42],
                                                  ["Made-up inline text", 7.5]]
    payload, is_error, _ = h.call("xl_format", {"op": "manage_comment", "args": {
        "path": str(f), "action": "list"}})
    assert not is_error and payload["comments"][0]["text"] == "Made-up note"
    payload, is_error, _ = h.call("xl_edit", {"op": "set_cell", "args": {
        "path": str(f), "location": "B1", "value": 43}})
    assert not is_error and payload["saved"] is True, payload
    payload, is_error, _ = h.call("xl_read", {"op": "read_range", "args": {
        "path": str(f), "location": "B1"}})
    assert not is_error and payload["values"] == [[43]]


def test_a_comment_inside_a_styles_list_counts_as_before(tmp_path):
    from xlsx_mcp.ops import format as fmt

    styles = hx.base_parts()["xl/styles.xml"].replace(
        "<cellXfs count=\"2\">", "<cellXfs count=\"2\"><!-- made-up note -->", 1)
    f = hx.make_xlsx(tmp_path / "s.xlsx", parts={"xl/styles.xml": styles})
    assert fmt._styles_xml_counts(str(f))["cell_formats"] == 2
