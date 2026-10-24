"""Output-budget pins: what a result COSTS, and that it still says everything.

The input surface of this server was fat-trimmed on purpose in September 2026
and the reasoning behind it is recorded in core/schemas.py: pay on error, not on
load. The output side had never had the same pass. A fat audit
measured it: a 200-row read cost 22,666 tokens to deliver 6,479 tokens of
payload, because every result crossed the wire twice, once pretty-printed and
once again whole in structuredContent, and because the label matrix spent 52%
of a read repeating the word "value" about ordinary cells.

Every pin here runs in BOTH directions, which is the only way a token cut is
allowed to be judged:

  - the BUDGET direction: the payload is actually smaller, measured, not
    asserted by inspection;
  - the COMPLETENESS direction: the smaller payload carries the identical
    information, reconstructed and compared against what the dense form said.

A trim that fails the second direction is not a trim, it is data loss.

The runtime pin (get_cells opening the workbook once per cell) lives here too
because it is the same fat, spent in seconds instead of tokens.
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest
from conftest import cell_rows, dense_labels, label_at

from xlsx_mcp.core import calc as _calc
from xlsx_mcp.ops import cells, gridio


@pytest.fixture
def grid(tmp_path):
    """200 rows x 8 columns, one column of never-calculated formulas: the
    shape the audit measured."""
    p = tmp_path / "grid.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Region", "Q1", "Q2", "Q3", "Q4", "Total", "Note", "Flag"])
    for i in range(2, 202):
        ws.append([f"R{i}", i, i * 2, i * 3, i * 4,
                   f"=SUM(B{i}:E{i})", "note", True])
    wb.save(p)
    wb.close()
    return str(p)


def _tokens(payload) -> int:
    return round(len(json.dumps(payload, default=str,
                                separators=(",", ":"))) / 4)


# ------------------------------------------------------------ runtime budget


def test_get_cells_opens_the_workbook_a_fixed_number_of_times(grid,
                                                              monkeypatch):
    """500 scattered cells cost the SAME opens as 5.

    get_cells reached read_matrix once per cell, and read_matrix runs the
    honest-label formula probe, which is a full openpyxl parse. 500 cells
    therefore cost 501 opens and 5.3 seconds where the equivalent read_range
    cost 2 opens and 31 ms; MAX_SCATTER_CELLS is 1,000, so a PERMITTED call
    sat there for ten seconds on a small file. Counting opens rather than
    timing is deliberate: a time threshold is a flake on a loaded machine,
    and the opens are the actual defect.
    """
    opens = {"n": 0}
    real = openpyxl.load_workbook

    def counted(*a, **k):
        opens["n"] += 1
        return real(*a, **k)

    monkeypatch.setattr(openpyxl, "load_workbook", counted)

    small = [f"A{r}" for r in range(2, 7)]
    big = [f"{c}{r}" for r in range(2, 127) for c in "ABCD"][:500]

    opens["n"] = 0
    cells.get_cells(grid, small)
    few = opens["n"]

    opens["n"] = 0
    cells.get_cells(grid, big)
    many = opens["n"]

    assert few == many, (
        f"opens scale with the cell count: {few} for 5 cells, {many} for 500")
    assert many <= 3, f"a scatter read took {many} workbook opens"

    opens["n"] = 0
    cells.read_range(grid, "A1:H201")
    assert many <= opens["n"] + 1, (
        f"get_cells x500 ({many} opens) is not comparable to read_range "
        f"({opens['n']} opens)")


def test_scatter_probe_is_bounded_by_the_rows_asked_for(grid):
    """A scatter of two far-apart cells does not stream the space between
    them: the probe clamps to the sheet's own last row."""
    mask = gridio.formula_mask_cells(grid, "Data", [(2, 6), (201, 6)])
    assert mask == {(2, 6): "=SUM(B2:E2)", (201, 6): "=SUM(B201:E201)"}
    beyond = gridio.formula_mask_cells(grid, "Data", [(900_000, 1)])
    assert beyond == {}


def test_the_hoisted_probe_labels_exactly_what_the_per_cell_probe_did(grid):
    """COMPLETENESS: one probe over many cells, same answers as one probe per
    cell. The labels are the honest-calc contract; only the opens changed."""
    addrs = ["A2", "F2", "B3", "F3", "G4", "F201"]
    hoisted = cell_rows(cells.get_cells(grid, addrs))
    per_cell = [cell_rows(cells.get_cells(grid, [a]))[0] for a in addrs]
    assert [(r["cell"], r["value"], r["label"]) for r in hoisted] == \
           [(r["cell"], r["value"], r["label"]) for r in per_cell]
    assert {r["label"] for r in hoisted} == {"value", _calc.LABEL_ABSENT}


# -------------------------------------------------------------- label budget


def test_sparse_labels_carry_the_dense_matrix_exactly(grid):
    """COMPLETENESS: the grouped form rebuilds the parallel matrix cell for
    cell, so nothing a caller could have read is gone."""
    out = cells.read_range(grid, "A1:H201", values="both")
    rebuilt = dense_labels(out, rows=out["rows"], cols=out["cols"])

    # what read_matrix itself says, dense, straight from the engine
    fwb = gridio.open_wb(grid, data_only=False)
    cwb = gridio.open_wb(grid, data_only=True)
    try:
        g = gridio.resolve(fwb, "A1:H201", path=grid)
        _v, dense, _hf = gridio.read_matrix(
            g, mode="both", formula_wb=fwb, cached_wb=cwb, path=grid)
    finally:
        fwb.close()
        cwb.close()

    assert rebuilt == dense
    assert out["labels_default"] == _calc.LABEL_VALUE


def test_sparse_labels_are_smaller_than_the_matrix_they_replace(grid):
    """BUDGET: measured, on the shape the audit measured."""
    out = cells.read_range(grid, "A1:H201", values="both")
    dense = dense_labels(out, rows=out["rows"], cols=out["cols"])
    sparse_cost = _tokens(out["labels"]) + _tokens(out["labels_default"])
    dense_cost = _tokens(dense)
    assert sparse_cost < dense_cost / 2, (
        f"labels cost {sparse_cost} tokens against {dense_cost} dense")
    # and the whole payload is dominated by values again, not by labels
    assert sparse_cost < _tokens(out["values"]) / 2


def test_an_all_ordinary_read_sends_no_labels_at_all(tmp_path):
    """A sheet with no formulas has nothing to label, and says so by omission
    rather than by a matrix of the word "value"."""
    p = tmp_path / "plain.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r in range(1, 21):
        ws.append([r, f"row{r}"])
    wb.save(p)
    wb.close()
    out = cells.read_range(str(p), "A1:B20", values="both")
    assert "labels" not in out
    assert label_at(out, "A1") == _calc.LABEL_VALUE


def test_labels_name_the_formula_cells_and_only_those(grid):
    """The information the labels exist to carry: WHICH cells are not plain
    values. Column F is the formula column; nothing else is listed."""
    out = cells.read_range(grid, "A1:H201", values="both")
    listed = {a for addrs in out["labels"].values() for a in addrs}
    assert listed == {f"F{r}" for r in range(2, 202)}


# ----------------------------------------------------------- get_cells shape


def test_get_cells_rows_carry_every_value_the_dicts_carried(grid):
    """COMPLETENESS: address, value and label survive the array encoding."""
    addrs = ["A2", "B2", "F2", "H2"]
    out = cells.get_cells(grid, addrs, values="both")
    assert out["fields"] == ["cell", "value"]
    assert out["sheet"] == "Data"
    rows = cell_rows(out)
    assert [r["cell"] for r in rows] == addrs
    assert [r["value"] for r in rows] == ["R2", 2, None, True]
    assert [r["label"] for r in rows] == [
        "value", "value", _calc.LABEL_ABSENT, "value"]
    assert all(r["sheet"] == "Data" for r in rows)


def test_get_cells_qualifies_addresses_when_the_sheet_is_not_uniform(tmp_path):
    """The envelope can only hoist ONE sheet. Cells from two sheets keep
    their sheet, in the address, the way a formula would write it."""
    p = tmp_path / "two.xlsx"
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "One"
    a["A1"] = 1
    b = wb.create_sheet("Sheet Two")
    b["A1"] = 2
    wb.save(p)
    wb.close()
    out = cells.get_cells(str(p), [{"cell": "A1", "sheet": "One"},
                                   {"cell": "A1", "sheet": "Sheet Two"}])
    assert "sheet" not in out
    assert [r[0] for r in out["cells"]] == ["One!A1", "'Sheet Two'!A1"]
    rows = cell_rows(out)
    assert [(r["sheet"], r["cell"], r["value"]) for r in rows] == [
        ("One", "A1", 1), ("Sheet Two", "A1", 2)]


def test_get_cells_payload_is_a_fraction_of_the_dict_encoding(grid):
    """BUDGET: the audit's own case, 500 cells, measured against the shape it
    replaced ({sheet, cell, value, label} per cell)."""
    addrs = [f"{c}{r}" for r in range(2, 127) for c in "ABCD"][:500]
    out = cells.get_cells(grid, addrs)
    rows = cell_rows(out)
    old_shape = {
        "count": out["count"], "value_mode": out["value_mode"],
        "cells": [{"sheet": r["sheet"], "cell": r["cell"],
                   "value": r["value"], "label": r["label"]} for r in rows],
    }
    new_cost, old_cost = _tokens(out), _tokens(old_shape)
    assert new_cost < old_cost / 3, (
        f"500 cells cost {new_cost} tokens against {old_cost} before")
    # and what is left is the values plus the addresses that name them. The
    # audit measured 32x over bare values; the floor for a SCATTER read is
    # not 1x, because the caller asked about 500 specific addresses and the
    # addresses are what say which value is which.
    bare = _tokens([r["value"] for r in rows])
    assert new_cost < bare * 4, (
        f"{new_cost} tokens to deliver {bare} tokens of values")


# ------------------------------------------------------------ wire envelope


def test_a_success_crosses_the_wire_once_and_compact():
    """BUDGET + COMPLETENESS at the MCP boundary: content only, no duplicate
    structuredContent, no indentation, and the same object either way."""
    import asyncio

    from fastmcp import Client

    from xlsx_mcp import server

    async def run():
        async with Client(server.create_server()) as c:
            return await c.call_tool("xl_read", {"op": "get_server_info"})

    res = asyncio.run(run())
    assert res.structured_content is None, (
        "success results still ship a duplicate structuredContent copy")
    text = res.content[0].text
    payload = json.loads(text)
    assert text == json.dumps(payload, ensure_ascii=False,
                              separators=(",", ":")), (
        "success payload is still pretty-printed")
    assert payload["ok"] is True
    assert payload["version"]


def test_a_refusal_keeps_its_structure_and_loses_only_the_whitespace():
    """The refusal contract is untouchable: isError, the closed code, the
    message and the hint all stay. Family surface: the refusal now
    crosses the wire once as well, compact, with no structuredContent copy."""
    from xlsx_mcp import envelope
    from xlsx_mcp.core.errors import WorkbookNotFound

    payload = envelope.refusal(WorkbookNotFound("no such workbook: nope.xlsx"))
    r = envelope.tool_result(payload, is_error=True)
    assert r.is_error is True
    assert r.structured_content is None
    assert payload["ok"] is False
    assert payload["error"]["code"] == "NOT_FOUND"
    assert payload["error"]["hint"]
    text = r.content[0].text
    assert "\n" not in text
    assert json.loads(text) == payload


def test_no_tool_declares_a_contentless_output_schema():
    """A schema of {"type": "object", "additionalProperties": true} validates
    nothing and obliged every client to expect a second copy of every
    result. Family surface: no tool declares an output schema at all."""
    import asyncio

    from fastmcp import Client

    from xlsx_mcp import server

    async def run():
        async with Client(server.create_server()) as c:
            return await c.list_tools()

    offenders = [t.name for t in asyncio.run(run()) if t.outputSchema is not None]
    assert not offenders, f"tools declaring an outputSchema: {offenders}"


def test_the_published_token_figures_are_the_measured_ones():
    """The estimator counts what tools/list actually sends.

    It once summed description + inputSchema and dropped annotations, title
    and _meta, so the published figures understated the wire by ~16%. A
    server that sells honest numbers measures what it publishes. Family
    surface: scripts/measure_surface.py counts each tool exactly as the
    client receives it over real stdio."""
    import asyncio
    import importlib.util

    from fastmcp import Client

    from xlsx_mcp import server

    spec = importlib.util.spec_from_file_location(
        "_measure_surface", Path(__file__).resolve().parents[2] / "scripts" /
        "measure_surface.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)

    async def run():
        async with Client(server.create_server()) as c:
            return await c.list_tools()

    t = {x.name: x for x in asyncio.run(run())}["xl_read"]
    wire = t.model_dump(exclude_none=True, by_alias=True, mode="json")
    assert mod.approx_tokens(wire) == round(len(json.dumps(
        wire, ensure_ascii=False, separators=(",", ":"))) / 4)
    fields_only = len(t.description or "") + len(
        json.dumps(t.inputSchema or {}, separators=(",", ":")))
    assert mod.approx_tokens(wire) > round(fields_only / 4), (
        "the estimator is back to counting the flattering subset")
