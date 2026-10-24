"""The network check reads a network function however the file spells it.

Excel fetches from WEBSERVICE (and the other network functions) in more spellings than the plain name. Measured in
Excel 16 on made-up files with a listener on 127.0.0.1:

- an OpenDocument spreadsheet (.ods) stores the function after a namespace, COM.MICROSOFT.WEBSERVICE, and a bare
  name works there too;
- in an .xlsx, the function name may be written with XML character references (decimal, hex, zero padded), with a
  comment inside the name, in a CDATA section, or in a part stored as UTF-16 with or without a byte order mark.

The check must see all of these without opening the file in Excel. These tests build the files by code and read them
with the file check alone; the live twin, which asks Excel, is in test_xl_app_reaches_no_network_unasked.py.
"""
from __future__ import annotations

import zipfile

import openpyxl
import pytest

from xlsx_mcp.core import netreach
from xlsx_mcp.core.errors import XlMcpError
from xlsx_mcp.ops import comtier

ODS_NS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
          'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
          'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2"')


def _content(formula: str) -> str:
    return ('<?xml version="1.0" encoding="UTF-8"?>'
            f'<office:document-content {ODS_NS}><office:body><office:spreadsheet>'
            '<table:table table:name="Rates"><table:table-row>'
            '<table:table-cell office:value-type="string"><text:p>made-up</text:p></table:table-cell>'
            f'<table:table-cell table:formula="{formula}" office:value-type="string"><text:p></text:p>'
            '</table:table-cell></table:table-row></table:table></office:spreadsheet></office:body>'
            '</office:document-content>')


def _ods(path, formula: str):
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.spreadsheet")
        z.writestr("content.xml", _content(formula))
        z.writestr("META-INF/manifest.xml",
                   '<?xml version="1.0" encoding="UTF-8"?><manifest:manifest '
                   'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"/>')
    return str(path)


URL = "&quot;http://127.0.0.1:9/x?v=&quot;&amp;[.A1]"


@pytest.mark.parametrize("call,fn", [
    ("of:=COM.MICROSOFT.WEBSERVICE(" + URL + ")", "WEBSERVICE"),
    ("of:=WEBSERVICE(" + URL + ")", "WEBSERVICE"),
    ("of:=ORG.LIBREOFFICE.WEBSERVICE(" + URL + ")", "WEBSERVICE"),
    ("of:=_xlfn.WEBSERVICE(" + URL + ")", "WEBSERVICE"),
    ("of:=COM.MICROSOFT.IMAGE(&quot;http://127.0.0.1:9/a.png&quot;)", "IMAGE"),
    ("of:=COM.MICROSOFT.RTD(&quot;made.up.prog&quot;;&quot;&quot;;&quot;t&quot;)", "RTD"),
    ("of:=COM.MICROSOFT.CUBEVALUE(&quot;made-up&quot;)", "CUBEVALUE"),
    ("of:=COM.MICROSOFT.TRANSLATE(&quot;made-up&quot;;&quot;fr&quot;)", "TRANSLATE"),
    ("of:=_xlfn._xlws.PY(&quot;made-up&quot;)", "PY"),
    ("of:=COM.MICROSOFT.PY(&quot;made-up&quot;)", "PY"),
    ("of:=COM.MICROSOFT.WEB&#83;ERVICE(" + URL + ")", "WEBSERVICE"),
])
def test_an_opendocument_spreadsheet_is_read_whatever_the_prefix(tmp_path, call, fn):
    how, reach = netreach.file_reach(_ods(tmp_path / "rates.ods", call))
    assert how == netreach.READ_TEXT
    assert reach and fn in reach["formulas"][0]["functions"], (call, reach)


def test_the_xl_app_check_refuses_an_opendocument_spreadsheet_before_excel(tmp_path):
    p = _ods(tmp_path / "rates.ods", "of:=COM.MICROSOFT.WEBSERVICE(" + URL + ")")
    with pytest.raises(XlMcpError) as ei:
        comtier._refuse_network_before_excel(p, False, "export_pdf")
    assert "allow_network" in str(ei.value)


def test_a_flat_opendocument_file_and_a_csv_are_read_the_same_way(tmp_path):
    fods = tmp_path / "rates.fods"
    fods.write_text(_content("of:=COM.MICROSOFT.WEBSERVICE(" + URL + ")"), encoding="utf-8")
    assert netreach.file_reach(str(fods))[1]
    for text in ('=COM.MICROSOFT.WEBSERVICE("http://127.0.0.1:9/x")', '=_xlfn.WEBSERVICE("http://127.0.0.1:9/x")',
                 '=_xlfn._xlws.PY("made-up")'):
        csv = tmp_path / "rates.csv"
        csv.write_text('name,value\nmade-up,"' + text.replace('"', '""') + '"\n', encoding="utf-8")
        assert netreach.file_reach(str(csv))[1], text


def test_a_utf16_text_file_without_a_byte_order_mark_is_read(tmp_path):
    p = tmp_path / "rates.fods"
    p.write_bytes(_content("of:=COM.MICROSOFT.WEBSERVICE(" + URL + ")").replace("UTF-8", "UTF-16").encode("utf-16-le"))
    assert netreach.file_reach(str(p))[1]


@pytest.mark.parametrize("text", [
    'name,value\nmade-up,=1+2\n',
    'name,value\nrun,"see script.py(1) and notes.txt"\n',
    'name,value\nmade-up,"xWEBSERVICE(1) and NOTIMAGE(2)"\n',
    'name,value\nmade-up,"the word image is not a call, nor is translate"\n',
])
def test_text_that_calls_nothing_that_fetches_is_not_refused(tmp_path, text):
    p = tmp_path / "plain.csv"
    p.write_text(text, encoding="utf-8")
    assert netreach.file_reach(str(p)) == (netreach.READ_TEXT, None)


def test_an_opendocument_spreadsheet_without_one_is_not_refused(tmp_path):
    p = _ods(tmp_path / "plain.ods", "of:=COM.MICROSOFT.ENCODEURL(&quot;a b&quot;)")
    assert netreach.file_reach(p) == (netreach.READ_TEXT, None)


# ---- an .xlsx whose markup spells the name another way ------------------------------------------------------------

FORMULA = '_xlfn.WEBSERVICE("http://127.0.0.1:9/x?v="&amp;A1)'


def _book(tmp_path):
    p = tmp_path / "rates.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws["A1"] = "made-up"
    ws["B1"] = '=_xlfn.WEBSERVICE("http://127.0.0.1:9/x?v="&A1)'
    wb.save(p)
    return p


def _rewrite(src, dst, part, fn):
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == part:
                data = fn(data.decode("utf-8"))
                data = data if isinstance(data, bytes) else data.encode("utf-8")
            zo.writestr(item, data)
    return str(dst)


SPELLINGS = {
    "decimal reference": lambda t: t.replace("_xlfn.WEBSERVICE(", "_xlfn.WEB&#83;ERVICE("),
    "hex reference": lambda t: t.replace("_xlfn.WEBSERVICE(", "_xlfn.WEB&#x53;ERVICE("),
    "zero padded reference": lambda t: t.replace("_xlfn.WEBSERVICE(", "_xlfn.WEB&#" + "0" * 80 + "83;ERVICE("),
    "comment inside the name": lambda t: t.replace("_xlfn.WEBSERVICE(", "_xlfn.WEB<!-- x -->SERVICE("),
    "cdata section": lambda t: t.replace(FORMULA, "<![CDATA[" + FORMULA.replace("&amp;", "&") + "]]>"),
    "utf-16 with a mark": lambda t: b"\xff\xfe" + t.replace("UTF-8", "UTF-16").encode("utf-16-le"),
    "utf-16 without a mark": lambda t: t.replace("UTF-8", "UTF-16").encode("utf-16-le"),
}


@pytest.mark.parametrize("how", list(SPELLINGS))
def test_a_formula_written_another_way_in_the_markup_is_still_found(tmp_path, how):
    p = _rewrite(_book(tmp_path), tmp_path / "other.xlsx", "xl/worksheets/sheet1.xml", SPELLINGS[how])
    kind, reach = netreach.file_reach(p)
    assert kind == netreach.READ_PACKAGE
    assert reach and reach["formulas"][0]["where"] == "Rates!B1", (how, reach)
    assert reach["formulas"][0]["functions"] == ["WEBSERVICE"]


def test_the_same_spellings_are_found_in_a_defined_name_and_a_rule(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws["A1"] = 1
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.workbook.defined_name import DefinedName
    ws.conditional_formatting.add("A1", FormulaRule(formula=['LEN(_xlfn.WEBSERVICE("http://127.0.0.1:9/x"))>0']))
    wb.defined_names["Fetch"] = DefinedName("Fetch", attr_text='_xlfn.WEBSERVICE("http://127.0.0.1:9/n")')
    src = tmp_path / "rules.xlsx"
    wb.save(src)
    mid = _rewrite(src, tmp_path / "mid.xlsx", "xl/worksheets/sheet1.xml",
                   lambda t: t.replace("_xlfn.WEBSERVICE(", "_xlfn.WEB&#83;ERVICE("))
    out = _rewrite(mid, tmp_path / "out.xlsx", "xl/workbook.xml",
                   lambda t: t.replace("_xlfn.WEBSERVICE(", "_xlfn.WEB&#x53;ERVICE("))
    reach = netreach.network_reach(out)
    wheres = [f["where"] for f in reach["formulas"]]
    assert "the defined name Fetch" in wheres
    assert any("rule on Rates" in w for w in wheres), wheres


def test_a_name_that_is_not_a_call_in_the_markup_is_not_reported(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws["A1"] = "WEBSERVICE( is only text here"
    ws["B1"] = '="WEBSERVICE("&"x"'
    ws["C1"] = "=1+2"
    src = tmp_path / "plain.xlsx"
    wb.save(src)
    p = _rewrite(src, tmp_path / "plain2.xlsx", "xl/worksheets/sheet1.xml",
                 lambda t: t.replace("1+2", "1&#43;2"))
    assert netreach.file_reach(p) == (netreach.READ_PACKAGE, None)


def test_a_dotted_name_in_an_xlsx_formula_is_not_a_call(tmp_path):
    # measured: Excel reads Foo.WEBSERVICE( and COM.MICROSOFT.WEBSERVICE( in an .xlsx as unknown names (#NAME?)
    # and fetches nothing, so the package check keeps treating them as another function
    src = _book(tmp_path)
    p = _rewrite(src, tmp_path / "dotted.xlsx", "xl/worksheets/sheet1.xml",
                 lambda t: t.replace("_xlfn.WEBSERVICE(", "COM.MICROSOFT.WEBSERVICE("))
    assert netreach.file_reach(p) == (netreach.READ_PACKAGE, None)
